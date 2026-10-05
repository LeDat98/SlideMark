from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree
import copy

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x2E, 0x6F, 0xD8)
TEAL = RGBColor(0x1A, 0x9C, 0x8F)
AMBER = RGBColor(0xF2, 0xA9, 0x00)
RED = RGBColor(0xD6, 0x45, 0x45)
GREEN = RGBColor(0x2E, 0x9E, 0x5B)
GRAY = RGBColor(0x5F, 0x6B, 0x7A)
LIGHT = RGBColor(0xF1, 0xF4, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x22)


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def style_run(r, size, bold=False, color=DARK, italic=False):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def fill_tf(tf, paras, size=16, bold=False, color=DARK, align=None, bullet=False, space=4):
    """paras: str or list of str / (str, dict(size,bold,color,align))."""
    if isinstance(paras, str):
        paras = [paras]
    tf.word_wrap = True
    for i, p in enumerate(paras):
        opts = {}
        if isinstance(p, tuple):
            p, opts = p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = para.add_run()
        r.text = p
        style_run(r, opts.get("size", size), opts.get("bold", bold), opts.get("color", color))
        a = opts.get("align", align)
        if a is not None:
            para.alignment = a
        para.space_after = Pt(space)
        if bullet:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", "285750")
            pPr.set("indent", "-285750")
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")


def text(slide, x, y, w, h, paras, size=16, bold=False, color=DARK, align=None,
         anchor=MSO_ANCHOR.TOP, bullet=False):
    tbx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tbx.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    fill_tf(tf, paras, size, bold, color, align, bullet)
    return tbx


def shape(slide, kind, x, y, w, h, fill=LIGHT, line=None, paras=None, size=14, bold=False,
          color=DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, bullet=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    if paras is not None:
        fill_tf(tf, paras, size, bold, color, align, bullet)
    return s


def title(slide, t, sub=None):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 1.1, fill=NAVY)
    text(slide, 0.6, 0.15, 12.1, 0.8, t, size=30, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(slide, 0.6, 1.2, 12.1, 0.5, sub, size=18, color=GRAY)


def footnote(slide, t):
    text(slide, 0.6, 6.95, 12.1, 0.35, t, size=10, color=GRAY)


def line(slide, x1, y1, x2, y2, color=GRAY, width=2, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
    return c


def table(slide, x, y, w, rows, col_w=None, row_h=0.45, size=14, header_fill=NAVY,
          align_cols=None, highlight_rows=(), highlight_fill=RGBColor(0xFF, 0xF3, 0xCC)):
    nr, nc = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    t = gs.table
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    for r in range(nr):
        t.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = t.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            elif r in highlight_rows:
                cell.fill.fore_color.rgb = highlight_fill
            else:
                cell.fill.fore_color.rgb = LIGHT if r % 2 == 0 else WHITE
            tf = cell.text_frame
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(rows[r][c])
            style_run(run, size, bold=(r == 0 or r in highlight_rows), color=WHITE if r == 0 else DARK)
            if align_cols and c in align_cols and r > 0:
                p.alignment = align_cols[c]
            if r == 0:
                p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
    return t


def chart(slide, kind, x, y, w, h, cats, series, legend=True, labels=True, number_format=None,
          colors=None, font_size=12, title_text=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.size = Pt(font_size)
    ch.font.name = FONT
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    if title_text:
        ch.has_title = True
        ch.chart_title.text_frame.text = title_text
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        style_run(r, font_size + 2, bold=True, color=DARK)
    else:
        ch.has_title = False
    if labels:
        pl = ch.plots[0]
        pl.has_data_labels = True
        if number_format:
            pl.data_labels.number_format = number_format
            pl.data_labels.number_format_is_linked = False
    if colors and kind not in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.PIE):
        for s, c in zip(ch.plots[0].series, colors):
            if kind in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
                s.format.line.color.rgb = c
                s.format.line.width = Pt(3)
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = c
    return ch


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t

prs = new_prs()

# 1. Cover
s = blank(prs)
shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, fill=NAVY)
shape(s, MSO_SHAPE.RECTANGLE, 0.8, 3.55, 1.6, 0.08, fill=AMBER)
text(s, 0.8, 2.1, 11.5, 1.4, "LedgerAI", size=60, bold=True, color=WHITE)
text(s, 0.8, 3.8, 11.5, 0.8, "AI bookkeeping for small and medium businesses", size=26, color=WHITE)
text(s, 0.8, 5.8, 11.5, 0.5, "Seed pitch deck", size=16, color=RGBColor(0xB8, 0xC4, 0xDA))

# 2. Problem
s = blank(prs)
title(s, "The problem: bookkeeping drains SMBs")
items = [
    ("20+ hrs", "Owners spend over 20 hours a month on manual data entry and receipts."),
    ("1 in 4", "One in four books contains errors that surface only at tax time."),
    ("$500+", "Accountants charge $500+ per month, which is out of reach for micro-businesses."),
]
for i, (big, desc) in enumerate(items):
    x = 0.6 + i * 4.1
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, 1.9, 3.8, 4.2, fill=LIGHT)
    text(s, x + 0.2, 2.2, 3.4, 1.2, big, size=48, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    text(s, x + 0.3, 3.8, 3.2, 2.0, desc, size=18, color=DARK, align=PP_ALIGN.CENTER)

# 3. Solution
s = blank(prs)
title(s, "The solution: books that keep themselves")
text(s, 0.6, 1.4, 12.1, 0.6, "LedgerAI reads, categorizes and reconciles transactions automatically.", size=20, color=GRAY)
steps = [("1", "Connect", "Link bank, card and invoicing accounts in two minutes."),
         ("2", "Automate", "AI categorizes every transaction and matches receipts."),
         ("3", "Review", "Approve exceptions and export reports for your accountant.")]
for i, (n, h, d) in enumerate(steps):
    x = 0.6 + i * 4.3
    shape(s, MSO_SHAPE.OVAL, x, 2.6, 0.9, 0.9, fill=TEAL, paras=n, size=28, bold=True, color=WHITE)
    text(s, x, 3.65, 3.5, 0.6, h, size=24, bold=True, color=NAVY)
    text(s, x, 4.3, 3.5, 1.5, d, size=18)
    if i < 2:
        line(s, x + 3.5, 3.05, x + 4.2, 3.05, color=GRAY, width=3)

# 4. Traction
s = blank(prs)
title(s, "Traction: MRR up 5x in four quarters")
chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 0.6, 1.4, 8.2, 5.4, ["Q1", "Q2", "Q3", "Q4"],
      [("MRR (USD)", (10000, 18000, 31000, 52000))], legend=False, number_format='$#,##0',
      colors=[BLUE], font_size=14, title_text="Monthly recurring revenue (USD)")
shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 9.3, 2.2, 3.4, 3.2, fill=LIGHT,
      paras=[("$52k", {"size": 44, "bold": True, "color": BLUE}),
             ("MRR in Q4", {"size": 18}),
             ("+68% QoQ in Q4", {"size": 18, "bold": True, "color": GREEN})])
prs.save("01-pitch.pptx")
