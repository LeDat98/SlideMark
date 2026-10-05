from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree
import copy

FONT = "Calibri"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
BLUE = RGBColor(0x2E, 0x6F, 0xD8)
TEAL = RGBColor(0x1A, 0x9C, 0x8F)
AMBER = RGBColor(0xF2, 0xA9, 0x00)
RED = RGBColor(0xD6, 0x45, 0x45)
GREEN = RGBColor(0x2E, 0x9E, 0x5B)
GRAY = RGBColor(0x5F, 0x6B, 0x7A)
LIGHT = RGBColor(0xF1, 0xF4, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x22, 0x22, 0x22)


def new_prs():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def style_run(r, size, bold=False, color=DARK, italic=False):
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = FONT
    rPr = r._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def fill_tf(tf, paras, size=16, bold=False, color=DARK, align=None, bullet=False, space=4):
    """paras: str or list of str / (str, dict(size,bold,color,align))."""
    if isinstance(paras, str):
        paras = [paras]
    tf.word_wrap = True
    for i, p in enumerate(paras):
        opts = {}
        if isinstance(p, tuple):
            p, opts = p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = para.add_run()
        r.text = p
        style_run(r, opts.get("size", size), opts.get("bold", bold), opts.get("color", color))
        a = opts.get("align", align)
        if a is not None:
            para.alignment = a
        para.space_after = Pt(space)
        if bullet:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", "285750")
            pPr.set("indent", "-285750")
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")


def text(slide, x, y, w, h, paras, size=16, bold=False, color=DARK, align=None,
         anchor=MSO_ANCHOR.TOP, bullet=False):
    tbx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tbx.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    fill_tf(tf, paras, size, bold, color, align, bullet)
    return tbx


def shape(slide, kind, x, y, w, h, fill=LIGHT, line=None, paras=None, size=14, bold=False,
          color=DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, bullet=False):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    tf = s.text_frame
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    if paras is not None:
        fill_tf(tf, paras, size, bold, color, align, bullet)
    return s


def title(slide, t, sub=None):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 1.1, fill=NAVY)
    text(slide, 0.6, 0.15, 12.1, 0.8, t, size=30, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(slide, 0.6, 1.2, 12.1, 0.5, sub, size=18, color=GRAY)


def footnote(slide, t):
    text(slide, 0.6, 6.95, 12.1, 0.35, t, size=10, color=GRAY)


def line(slide, x1, y1, x2, y2, color=GRAY, width=2, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if arrow:
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
    return c


def table(slide, x, y, w, rows, col_w=None, row_h=0.45, size=14, header_fill=NAVY,
          align_cols=None, highlight_rows=(), highlight_fill=RGBColor(0xFF, 0xF3, 0xCC)):
    nr, nc = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    t = gs.table
    if col_w:
        for i, cw in enumerate(col_w):
            t.columns[i].width = Inches(cw)
    for r in range(nr):
        t.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = t.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.fill.solid()
            if r == 0:
                cell.fill.fore_color.rgb = header_fill
            elif r in highlight_rows:
                cell.fill.fore_color.rgb = highlight_fill
            else:
                cell.fill.fore_color.rgb = LIGHT if r % 2 == 0 else WHITE
            tf = cell.text_frame
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = str(rows[r][c])
            style_run(run, size, bold=(r == 0 or r in highlight_rows), color=WHITE if r == 0 else DARK)
            if align_cols and c in align_cols and r > 0:
                p.alignment = align_cols[c]
            if r == 0:
                p.alignment = PP_ALIGN.CENTER if c else PP_ALIGN.LEFT
    return t


def chart(slide, kind, x, y, w, h, cats, series, legend=True, labels=True, number_format=None,
          colors=None, font_size=12, title_text=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.size = Pt(font_size)
    ch.font.name = FONT
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    if title_text:
        ch.has_title = True
        ch.chart_title.text_frame.text = title_text
        r = ch.chart_title.text_frame.paragraphs[0].runs[0]
        style_run(r, font_size + 2, bold=True, color=DARK)
    else:
        ch.has_title = False
    if labels:
        pl = ch.plots[0]
        pl.has_data_labels = True
        if number_format:
            pl.data_labels.number_format = number_format
            pl.data_labels.number_format_is_linked = False
    if colors and kind not in (XL_CHART_TYPE.DOUGHNUT, XL_CHART_TYPE.PIE):
        for s, c in zip(ch.plots[0].series, colors):
            if kind in (XL_CHART_TYPE.LINE_MARKERS, XL_CHART_TYPE.LINE):
                s.format.line.color.rgb = c
                s.format.line.width = Pt(3)
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = c
    return ch


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t

FONT = "Yu Gothic"
prs = new_prs()
s = blank(prs)
title(s, "2026年度上期 業績サマリー")
text(s, 0.6, 1.2, 12.1, 0.5, "売上・利益ともに過去最高を更新、解約率は低水準を維持", size=18, color=GRAY)
kpis = [("売上高", "62.4億円"), ("営業利益", "8.1億円"), ("新規顧客", "312社"), ("解約率", "1.6%")]
for i, (k, v) in enumerate(kpis):
    x = 0.6 + i * 3.08
    shape(s, MSO_SHAPE.RECTANGLE, x, 1.9, 2.9, 1.7, fill=LIGHT)
    shape(s, MSO_SHAPE.RECTANGLE, x, 1.9, 0.1, 1.7, fill=BLUE)
    text(s, x + 0.25, 2.0, 2.5, 0.5, k, size=16, color=GRAY)
    text(s, x + 0.25, 2.5, 2.5, 0.9, v, size=34, bold=True, color=NAVY)
text(s, 0.6, 3.85, 12.1, 0.45, "事業部別実績", size=18, bold=True, color=NAVY)
rows = [["事業部", "売上（億円）", "前年比"],
        ["SaaS", "31.2", "+18.5%"],
        ["コンサル", "20.8", "+6.2%"],
        ["ハードウェア", "10.4", "-3.1%"]]
table(s, 0.6, 4.35, 12.1, rows, col_w=[5.1, 3.5, 3.5], row_h=0.5, size=16,
      align_cols={1: PP_ALIGN.RIGHT, 2: PP_ALIGN.RIGHT})
footnote(s, "出所：社内管理会計データ（2026年4月〜9月）")
prs.save("02-jp-summary.pptx")
