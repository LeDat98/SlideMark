"""Acceptance check for an agent-built deck against its brief (docs/AGENT_COST.md, "Acceptance script").

    python bench/agent_accept.py bench/briefs/b01-qbr/accept.json RUN_DIR/deck.pptx

``accept.json``: {"slides": 3, "size": "16:9", "charts": 1, "tables": 1, "notes": [3],
"strings": ["exact text", ...]}. Strings are searched (NFKC, whitespace-collapsed) in text frames, table
cells, group shapes, notes and chart XML (titles, categories, series names, values).
Prints one line per failed
requirement and a final ``accepted`` / ``rejected`` line; exit code 0 / 1.
"""

from __future__ import annotations

import json
import re
import sys
import unicodedata
from pathlib import Path


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", s)).strip().lower()


def _shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:  # group
            yield from _shapes(sh.shapes)


def facts(pptx: str | Path) -> dict:
    from pptx import Presentation
    from pptx.util import Emu

    prs = Presentation(str(pptx))
    w, h = Emu(prs.slide_width), Emu(prs.slide_height)
    ratio = w / h
    size = "16:9" if abs(ratio - 16 / 9) < 0.02 else "4:3" if abs(ratio - 4 / 3) < 0.02 else f"{ratio:.2f}"
    texts, charts, tables, notes = [], 0, 0, []
    for i, slide in enumerate(prs.slides, 1):
        for sh in _shapes(slide.shapes):
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                tables += 1
                for row in sh.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text)
            if getattr(sh, "has_chart", False) and sh.has_chart:
                charts += 1
                xml = sh.chart.part.blob.decode("utf-8", "replace")
                texts.extend(re.findall(r"<(?:a:t|c:v)>([^<]*)<", xml))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            notes.append(i)
            texts.append(slide.notes_slide.notes_text_frame.text)
    return {
        "slides": len(prs.slides),
        "size": size,
        "charts": charts,
        "tables": tables,
        "notes": notes,
        "text": _norm("\n".join(texts)),
    }


def check(spec: dict, pptx: str | Path) -> list[str]:
    """Failed requirements, one line each (empty = accepted)."""
    try:
        f = facts(pptx)
    except Exception as e:
        return [f"cannot open {pptx}: {type(e).__name__}: {e}"]
    bad = []
    if "slides" in spec and f["slides"] != spec["slides"]:
        bad.append(f"slides: {f['slides']} != {spec['slides']}")
    if spec.get("size", "16:9") != f["size"]:
        bad.append(f"size: {f['size']} != {spec.get('size', '16:9')}")
    for k in ("charts", "tables"):
        if f[k] < spec.get(k, 0):
            bad.append(f"{k}: {f[k]} < {spec[k]}")
    for n in spec.get("notes", []):
        if n not in f["notes"]:
            bad.append(f"notes missing on slide {n}")
    flat = f["text"].replace(" ", "")
    for s in spec.get("strings", []):
        if _norm(s) not in f["text"] and _norm(s).replace(" ", "") not in flat:
            bad.append(f"missing string: {s}")
    return bad


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 2:
        print(__doc__.strip().splitlines()[2].strip(), file=sys.stderr)
        return 2
    spec = json.loads(Path(argv[0]).read_text(encoding="utf-8"))
    bad = check(spec, argv[1])
    for b in bad:
        print(b)
    print("accepted" if not bad else f"rejected ({len(bad)} failed)")
    return 0 if not bad else 1


if __name__ == "__main__":
    raise SystemExit(main())
