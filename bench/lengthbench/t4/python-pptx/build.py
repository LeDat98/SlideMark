from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from lxml import etree

NAVY = RGBColor(0x12, 0x2B, 0x4A)
BLUE = RGBColor(0x1F, 0x5F, 0xA8)
TEAL = RGBColor(0x1B, 0x8A, 0x8F)
AMBER = RGBColor(0xE0, 0x8A, 0x1E)
INK = RGBColor(0x22, 0x2B, 0x36)
GREY = RGBColor(0x6B, 0x75, 0x82)
LIGHT = RGBColor(0xEE, 0xF2, 0xF7)
LINE = RGBColor(0xCF, 0xD8, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Meiryo"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = 13.333


def set_font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)


def rect(slide, x, y, w, h, fill=LIGHT, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    return s


def text(slide, x, y, w, h, parts, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.MIDDLE, shape=None, margin=0.08, spacing=None):
    """parts: str, or list of paragraphs; a paragraph is str or list of (text, size, bold, color)."""
    tb = shape or slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    paras = parts if isinstance(parts, list) else [parts]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        if spacing:
            para.space_after = Pt(spacing)
        runs = p if isinstance(p, list) else [(p, size, bold, color)]
        for t, sz, b, c in runs:
            r = para.add_run()
            r.text = t
            set_font(r, sz, b, c)
    return tb


def base(title, n):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, 1.15, NAVY)
    rect(s, 0, 1.15, W, 0.06, AMBER)
    text(s, 0.6, 0.12, W - 1.2, 0.9, title, size=28, color=WHITE, bold=True)
    text(s, 0.6, 7.05, 8, 0.3, "北斗クラウド株式会社　取締役会資料", size=10, color=GREY)
    text(s, W - 1.6, 7.05, 1.0, 0.3, str(n), size=10, color=GREY, align=PP_ALIGN.RIGHT)
    rect(s, 0.6, 7.02, W - 1.2, 0.01, LINE)
    return s


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


def style_chart_text(chart, size=14):
    chart.font.size = Pt(size)
    chart.font.name = FONT
    chart.font.color.rgb = INK


# ---------- 1 title
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, 7.5, NAVY)
rect(s, 0.9, 2.2, 0.12, 2.3, AMBER)
text(s, 1.25, 2.1, 11, 1.6, "北斗クラウド 中期経営計画 2027–2029", size=42, color=WHITE, bold=True,
     anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.25, 3.8, 11, 0.7, "北斗クラウド株式会社 取締役会資料 2027年4月", size=20,
     color=RGBColor(0xC9, 0xD6, 0xE6))
rect(s, 0, 6.9, W, 0.6, BLUE)

# ---------- 2 president message
s = base("社長メッセージ", 2)
items = [("3年で売上を2倍、営業利益率を20%へ", BLUE, "01"),
         ("顧客の業務時間を年間1,000万時間削減する", TEAL, "02"),
         ("「現場が使い続けるSaaS」をつくる", AMBER, "03")]
for i, (t, c, num) in enumerate(items):
    x = 0.6 + i * 4.1
    rect(s, x, 1.8, 3.9, 4.0, LIGHT)
    rect(s, x, 1.8, 3.9, 0.12, c)
    text(s, x + 0.3, 2.1, 3.3, 0.9, num, size=44, bold=True, color=c, anchor=MSO_ANCHOR.TOP)
    text(s, x + 0.3, 3.2, 3.3, 2.9, t, size=30, bold=True, color=NAVY, anchor=MSO_ANCHOR.TOP)

# ---------- 3 market
s = base("市場環境", 3)
cards = [
    ("2.1兆円", [("国内の中小企業向けSaaS市場は2026年の1.2兆円から2029年に2.1兆円へ", 20, True, NAVY)], BLUE),
    ("38%", [("クラウド会計の導入率は中小企業で38%", 20, True, NAVY)], TEAL),
    ("+22%", [("人手不足を理由にITを導入した企業は前年比+22%", 20, True, NAVY)], AMBER),
    ("41%", [("競合上位3社のシェア合計は41%", 20, True, NAVY)], GREY),
]
for i, (big, body, c) in enumerate(cards):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.2, 1.6 + row * 2.65
    rect(s, x, y, 5.95, 2.45, LIGHT)
    rect(s, x, y, 0.12, 2.45, c)
    text(s, x + 0.35, y + 0.1, 5.4, 1.0, big, size=48, bold=True, color=c)
    text(s, x + 0.35, y + 1.25, 5.4, 1.1, [body], anchor=MSO_ANCHOR.TOP)

# ---------- 4 targets
s = base("2029年度の数値目標", 4)
rows = [("売上高 ", "420億円", "（2026年度 210億円）"),
        ("営業利益率 ", "20%", "（同 11%）"),
        ("ARR ", "380億円", ""),
        ("解約率 ", "年0.8%以下", "（同 1.4%）"),
        ("有料顧客数 ", "48,000社", "（同 26,000社）")]
for i, (a, b, c) in enumerate(rows):
    y = 1.5 + i * 1.04
    rect(s, 0.6, y, W - 1.2, 0.92, LIGHT if i % 2 == 0 else WHITE, line=LINE)
    rect(s, 0.6, y, 0.12, 0.92, BLUE)
    runs = [(a, 22, True, NAVY), (b, 34, True, BLUE)]
    if c:
        runs.append((c, 18, False, GREY))
    text(s, 1.0, y, W - 2.0, 0.92, [runs])

# ---------- 5 segment chart
s = base("顧客セグメント別の売上構成（2026年度）", 5)
cd = CategoryChartData()
cats = ["製造業", "卸・小売", "建設・不動産", "サービス業", "その他"]
vals = [31, 24, 18, 15, 12]
cd.categories = cats
cd.add_series("売上構成比", vals)
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.5), Inches(8.2), Inches(5.3), cd)
ch = gf.chart
style_chart_text(ch, 16)
ch.has_legend = False
ch.has_title = False
pl = ch.plots[0]
pl.gap_width = 45
pl.has_data_labels = True
dl = pl.data_labels
dl.number_format = '0"%"'
dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(18)
dl.font.bold = True
ser = pl.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = BLUE
ser.points[0].format.fill.solid()
ser.points[0].format.fill.fore_color.rgb = AMBER
ch.category_axis.reverse_order = True
ch.category_axis.format.line.color.rgb = LINE
ch.category_axis.tick_labels.font.size = Pt(16)
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 36
ch.value_axis.minimum_scale = 0
# right panel
rect(s, 9.1, 1.6, 3.6, 5.1, LIGHT)
text(s, 9.3, 1.75, 3.2, 0.5, "構成比（2026年度）", size=14, bold=True, color=GREY)
for i, (c, v) in enumerate(zip(cats, vals)):
    y = 2.35 + i * 0.85
    text(s, 9.3, y, 3.2, 0.7, [[(c + " ", 20, False, INK), (f"{v}%", 26, True, AMBER if i == 0 else BLUE)]])

# ---------- 6 roadmap
s = base("3年間のプロダクト計画", 6)
road = [("2027年度", "在庫管理と請求書の自動照合をリリース", "2027年度：", BLUE),
        ("2028年度", "AIによる資金繰り予測を全プランに搭載", "2028年度：", TEAL),
        ("2029年度", "業種別テンプレートを30業種に拡大", "2029年度：", AMBER)]
rect(s, 0.6, 2.25, W - 1.2, 0.06, LINE)
for i, (yr, body, lead, c) in enumerate(road):
    x = 0.6 + i * 4.1
    arrow = rect(s, x, 1.55, 3.9, 0.75, c, shape=MSO_SHAPE.PENTAGON)
    text(s, 0, 0, 0, 0, [[(lead, 24, True, WHITE)]], shape=arrow, margin=0.25)
    rect(s, x, 2.5, 3.9, 2.7, LIGHT)
    text(s, x + 0.25, 2.65, 3.4, 2.4, body, size=26, bold=True, color=NAVY, anchor=MSO_ANCHOR.TOP)
rect(s, 0.6, 5.5, W - 1.2, 1.05, NAVY)
text(s, 0.9, 5.5, W - 1.8, 1.05, "各年度の第1四半期に大型リリース、第3四半期に改善リリース", size=26,
     bold=True, color=WHITE)

# ---------- 7 comparison table
s = base("競合との比較", 7)
hdr = ["", "月額", "導入日数", "サポート満足度", "API連携数"]
data = [("北斗クラウド：", ["月額9,800円", "導入日数3日", "サポート満足度92%", "API連携数120"]),
        ("A社：", ["月額12,000円", "導入日数10日", "サポート満足度85%", "API連携数80"]),
        ("B社：", ["月額7,500円", "導入日数5日", "サポート満足度78%", "API連携数45"])]
tbl_shape = s.shapes.add_table(4, 5, Inches(0.6), Inches(1.6), Inches(W - 1.2), Inches(4.6))
tbl = tbl_shape.table
tblPr = tbl._tbl.tblPr
sid = tblPr.find(qn("a:tableStyleId"))
if sid is not None:
    sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # no style, table grid
tbl.columns[0].width = Inches(2.9)
for j in range(1, 5):
    tbl.columns[j].width = Inches((W - 1.2 - 2.9) / 4)
tbl.rows[0].height = Inches(0.8)
for r in range(1, 4):
    tbl.rows[r].height = Inches(1.26)


def cell(r, c, parts, fill, color=INK, size=20, bold=False, align=PP_ALIGN.CENTER):
    ce = tbl.cell(r, c)
    ce.fill.solid()
    ce.fill.fore_color.rgb = fill
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = ce.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for t, sz, b, col in parts:
        rn = p.add_run()
        rn.text = t
        set_font(rn, sz, b, col)


for j, h in enumerate(hdr):
    cell(0, j, [(h, 18, True, WHITE)], NAVY)
for i, (name, vs) in enumerate(data):
    r = i + 1
    hl = i == 0
    fill = RGBColor(0xFD, 0xF1, 0xDE) if hl else (LIGHT if i % 2 else WHITE)
    cell(r, 0, [(name, 20, True, NAVY)], fill, align=PP_ALIGN.LEFT)
    for j, v in enumerate(vs):
        lab = hdr[j + 1]
        num = v[len(lab):]
        cell(r, j + 1, [(lab, 15, False, GREY), (num, 26, True, BLUE if not hl else AMBER)], fill)
    # put label above value
for r in range(1, 4):
    for j in range(1, 5):
        p = tbl.cell(r, j).text_frame.paragraphs[0]
        runs = p.runs
        # split into two paragraphs: label then number
        lab_run, num_run = runs[0], runs[1]
        p._p.remove(num_run._r)
        p2 = tbl.cell(r, j).text_frame.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2._p.append(num_run._r)

# ---------- 8 pricing
s = base("価格改定の方針", 8)
rect(s, 0.6, 1.6, 5.3, 5.1, NAVY)
text(s, 0.9, 1.8, 4.7, 0.5, "スタンダードプラン", size=16, color=RGBColor(0xC9, 0xD6, 0xE6))
text(s, 0.9, 2.6, 4.7, 1.0, "9,800円", size=40, bold=True, color=WHITE)
text(s, 0.9, 3.55, 4.7, 0.7, "▼", size=28, color=AMBER, align=PP_ALIGN.LEFT)
text(s, 0.9, 4.3, 4.7, 1.1, "10,800円", size=48, bold=True, color=AMBER)
text(s, 0.9, 5.6, 4.7, 0.6, "2027年10月から", size=18, color=WHITE)
pts = ["2027年10月からスタンダードプランを9,800円から10,800円へ",
       "既存顧客は12か月間据え置き",
       "年間契約の割引率を10%から15%へ拡大",
       "解約率への影響は+0.2ptと想定し、サポート体制で吸収する"]
for i, t in enumerate(pts):
    y = 1.6 + i * 1.3
    rect(s, 6.2, y, 6.5, 1.15, LIGHT)
    n = rect(s, 6.2, y, 0.8, 1.15, [BLUE, TEAL, AMBER, GREY][i])
    text(s, 0, 0, 0, 0, str(i + 1), size=26, bold=True, color=WHITE, align=PP_ALIGN.CENTER, shape=n)
    text(s, 7.15, y, 5.4, 1.15, t, size=21, bold=True, color=NAVY)
notes(s, "スタンダードプランの値上げは既存顧客への12か月の猶予と年間契約の割引拡大で痛みを和らげ、解約率の上昇は+0.2ptにとどめる方針です。")

# ---------- 9 organization
s = base("組織の変更", 9)
org = ["カスタマーサクセス部を営業本部から独立させ、2027年7月に発足",
       "AI開発室を新設、2027年度中に20名体制",
       "地方拠点（福岡・札幌）を営業所から支社へ格上げ",
       "執行役員にCS担当を1名追加"]
cols = [BLUE, TEAL, AMBER, GREY]
for i, t in enumerate(org):
    col, row = i % 2, i // 2
    x, y = 0.6 + col * 6.2, 1.6 + row * 2.65
    rect(s, x, y, 5.95, 2.45, LIGHT)
    rect(s, x, y, 5.95, 0.12, cols[i])
    text(s, x + 0.3, y + 0.3, 1.2, 0.8, f"{i + 1:02d}", size=32, bold=True, color=cols[i])
    text(s, x + 0.3, y + 1.0, 5.35, 1.35, t, size=23, bold=True, color=NAVY, anchor=MSO_ANCHOR.TOP)

# ---------- 10 risks table
s = base("主なリスクと対策", 10)
risks = [("大手の参入による価格競争", "業種特化機能で差別化"),
         ("採用難", "リファラル採用比率を40%へ、年収テーブル改定"),
         ("セキュリティ事故", "ISMAP取得を2028年度に完了"),
         ("為替・クラウド費用の上昇", "複数クラウド契約で年5%の原価低減")]
t2 = s.shapes.add_table(5, 3, Inches(0.6), Inches(1.6), Inches(W - 1.2), Inches(5.0)).table
sid = t2._tbl.tblPr.find(qn("a:tableStyleId"))
if sid is not None:
    sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
t2.columns[0].width = Inches(5.0)
t2.columns[1].width = Inches(0.7)
t2.columns[2].width = Inches(W - 1.2 - 5.7)
t2.rows[0].height = Inches(0.7)
for r in range(1, 5):
    t2.rows[r].height = Inches(1.07)


def cell2(r, c, t, fill, color, size, bold, align=PP_ALIGN.LEFT):
    ce = t2.cell(r, c)
    ce.fill.solid()
    ce.fill.fore_color.rgb = fill
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    ce.margin_left = Inches(0.2)
    p = ce.text_frame.paragraphs[0]
    p.alignment = align
    ce.text_frame.word_wrap = True
    rn = p.add_run()
    rn.text = t
    set_font(rn, size, bold, color)


cell2(0, 0, "リスク", NAVY, WHITE, 18, True)
cell2(0, 1, "", NAVY, WHITE, 18, True)
cell2(0, 2, "対策", NAVY, WHITE, 18, True)
for i, (a, b) in enumerate(risks):
    f = LIGHT if i % 2 == 0 else WHITE
    cell2(i + 1, 0, a + "：", f, NAVY, 22, True)
    cell2(i + 1, 1, "▶", f, AMBER, 20, True, PP_ALIGN.CENTER)
    cell2(i + 1, 2, b, f, INK, 22, False)

# ---------- 11 investment
s = base("投資配分（3年間 合計150億円）", 11)
cd = CategoryChartData()
cats = ["プロダクト開発", "営業・マーケティング", "人材採用・育成", "セキュリティ・基盤"]
vals = [70, 40, 25, 15]
cd.categories = cats
cd.add_series("投資額（億円）", vals)
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.5), Inches(8.4), Inches(5.3), cd)
ch = gf.chart
style_chart_text(ch, 16)
ch.has_legend = False
ch.has_title = False
pl = ch.plots[0]
pl.gap_width = 45
pl.has_data_labels = True
pl.data_labels.number_format = '0"億円"'
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size = Pt(18)
pl.data_labels.font.bold = True
pl.series[0].format.fill.solid()
pl.series[0].format.fill.fore_color.rgb = TEAL
pl.series[0].points[0].format.fill.solid()
pl.series[0].points[0].format.fill.fore_color.rgb = BLUE
ch.category_axis.reverse_order = True
ch.category_axis.tick_labels.font.size = Pt(16)
ch.category_axis.format.line.color.rgb = LINE
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.minimum_scale = 0
ch.value_axis.maximum_scale = 85
rect(s, 9.3, 1.6, 3.4, 5.1, LIGHT)
text(s, 9.5, 1.8, 3.0, 0.5, "3年間 合計", size=14, bold=True, color=GREY)
text(s, 9.5, 2.3, 3.0, 1.2, "150億円", size=38, bold=True, color=NAVY)
for i, (c, v) in enumerate(zip(cats, vals)):
    y = 3.7 + i * 0.7
    text(s, 9.5, y, 3.0, 0.6, [[(c + " ", 13, False, INK), (f"{v}億円", 18, True, BLUE)]])

# ---------- 12 hiring
s = base("採用計画", 12)
cd = CategoryChartData()
cd.categories = ["2027年度", "2028年度", "2029年度"]
cd.add_series("エンジニア", [60, 80, 90])
cd.add_series("その他", [60, 70, 90])
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(6.9), Inches(1.5), Inches(5.9), Inches(5.3), cd)
ch = gf.chart
style_chart_text(ch, 14)
ch.has_title = False
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(14)
pl = ch.plots[0]
pl.gap_width = 60
pl.has_data_labels = True
pl.data_labels.font.size = Pt(16)
pl.data_labels.font.bold = True
pl.data_labels.font.color.rgb = WHITE
pl.series[0].format.fill.solid()
pl.series[0].format.fill.fore_color.rgb = BLUE
pl.series[1].format.fill.solid()
pl.series[1].format.fill.fore_color.rgb = GREY
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False
ch.value_axis.maximum_scale = 200
ch.category_axis.tick_labels.font.size = Pt(16)
hire = [("2027年度 ", "120名", "（うちエンジニア60名）"),
        ("2028年度 ", "150名", "（うちエンジニア80名）"),
        ("2029年度 ", "180名", "（うちエンジニア90名）")]
for i, (a, b, c) in enumerate(hire):
    y = 1.6 + i * 1.15
    rect(s, 0.6, y, 6.0, 1.0, LIGHT)
    rect(s, 0.6, y, 0.12, 1.0, BLUE)
    text(s, 0.9, y, 5.6, 1.0, [[(a, 18, True, NAVY), (b, 28, True, BLUE), (c, 17, False, GREY)]])
rect(s, 0.6, 5.2, 6.0, 1.5, NAVY)
text(s, 0.9, 5.2, 5.5, 1.5, "女性管理職比率を2029年度に25%へ（2026年度 14%）", size=24, bold=True, color=WHITE)

# ---------- 13 voice
s = base("お客様の声", 13)
rect(s, 0.6, 1.7, W - 1.2, 4.9, LIGHT)
rect(s, 0.6, 1.7, 0.15, 4.9, AMBER)
text(s, 1.0, 1.7, 1.5, 1.6, "“", size=96, bold=True, color=AMBER, anchor=MSO_ANCHOR.TOP)
text(s, 1.6, 2.5, W - 3.2, 2.5, "「月末の締め作業が3日から半日になった。現場が自分で使えるのが一番の違いです」",
     size=38, bold=True, color=NAVY)
text(s, 1.6, 5.3, W - 3.2, 0.8, "— 株式会社三浦製作所 経理部長", size=24, color=GREY, align=PP_ALIGN.RIGHT)

# ---------- 14 regional
s = base("地域展開", 14)
reg = [("2027年度：", "福岡・札幌で地域密着の導入支援を開始", BLUE),
       ("2028年度：", "名古屋・広島に拠点を追加", TEAL),
       ("2029年度：", "台湾で日系企業向けに提供開始", AMBER)]
for i, (a, b, c) in enumerate(reg):
    x = 0.6 + i * 4.1
    arrow = rect(s, x, 1.6, 3.9, 0.75, c, shape=MSO_SHAPE.PENTAGON)
    text(s, 0, 0, 0, 0, a.rstrip("："), size=24, bold=True, color=WHITE, shape=arrow, margin=0.25)
    rect(s, x, 2.55, 3.9, 2.5, LIGHT)
    text(s, x + 0.25, 2.7, 3.4, 2.1, [[(a, 26, True, c), (b, 26, True, NAVY)]], anchor=MSO_ANCHOR.TOP)
rect(s, 0.6, 5.3, W - 1.2, 1.3, NAVY)
text(s, 0.9, 5.3, W - 1.8, 1.3, "各拠点で地元の会計事務所と提携し、紹介経由の契約を全体の30%へ", size=26,
     bold=True, color=WHITE)

# ---------- 15 requests
s = base("取締役会への依頼事項", 15)
ask = ["中期経営計画の承認", "投資枠150億円の承認", "価格改定（2027年10月実施）の承認",
       "組織変更（CS部独立・AI開発室新設）の承認"]
for i, t in enumerate(ask):
    y = 1.6 + i * 1.3
    rect(s, 0.6, y, W - 1.2, 1.12, LIGHT, line=LINE)
    n = rect(s, 0.6, y, 1.1, 1.12, [BLUE, TEAL, AMBER, NAVY][i])
    text(s, 0, 0, 0, 0, str(i + 1), size=32, bold=True, color=WHITE, align=PP_ALIGN.CENTER, shape=n)
    text(s, 2.0, y, W - 3.0, 1.12, t, size=30, bold=True, color=NAVY)
notes(s, "以上四点が本日ご承認をお願いしたい事項であり、承認後は2027年4月から計画を実行に移します。")

prs.save("deck.pptx")
