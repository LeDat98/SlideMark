"""User templates: resolve a deck ``theme`` (built-in name, .yaml, .pptx/.potx) and open a template as base.

Owner: renderer workstream. Nothing here raises on bad input: problems become ``bad-theme`` diagnostics.
"""

from __future__ import annotations

import copy
import io
import re
import uuid
import zipfile
from pathlib import Path

from lxml import etree

from .ir import Diagnostic
from .theme import DEFAULT, Fonts, Theme, apply_slide_ink, apply_tokens, get_theme, theme_from_data

__all__ = [
    "deck_theme",
    "resolve_theme",
    "template_size",
    "footer_top",
    "open_template",
    "pick_layout",
    "clone_footer",
]

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS = {"a": _A, "p": _P}
_CT_TEMPLATE = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
_CT_PRESENTATION = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
_SECTION_EXT = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"


# --------------------------------------------------------------------------- theme resolution


def qn(tag: str) -> str:
    """``pptx.oxml.ns.qn``, imported lazily: python-pptx costs ~50 ms and only rendering needs it."""
    from pptx.oxml.ns import qn as _qn

    return _qn(tag)


def _bad(path: str, why: str, hint: str) -> Diagnostic:
    return Diagnostic(level="warning", message=f"theme {path!r}: {why}", rule="bad-theme", hint=hint)


def resolve_theme(name: str, base_dir: str | Path | None = None) -> tuple[Theme, list[Diagnostic]]:
    """Built-in name -> registry theme; a .pptx/.potx/.yaml path (relative to ``base_dir``) -> Theme."""
    name = (name or "default").strip()
    try:
        return get_theme(name), []
    except KeyError:
        pass
    suffix = Path(name).suffix.lower()
    if suffix not in (".pptx", ".potx", ".yaml", ".yml"):
        return DEFAULT, [
            _bad(
                name,
                "unknown theme",
                "use a built-in name (see `slidemark themes`) or a .pptx/.potx/.yaml path",
            )
        ]
    path = Path(name)
    if not path.is_absolute():
        path = Path(base_dir or ".") / path
    try:
        path = path.resolve()
        if not path.is_file():
            return DEFAULT, [
                _bad(name, "file not found", "give the path relative to the .md file; using default")
            ]
        theme = _from_yaml(path) if suffix in (".yaml", ".yml") else _from_office(path)
        return theme, []
    except Exception as e:  # corrupt file, bad YAML, pydantic errors: never raise
        first = str(e).strip().splitlines()[0][:120] if str(e).strip() else type(e).__name__
        return DEFAULT, [
            _bad(
                name,
                f"{type(e).__name__}: {first}",
                "fix or re-save the file (PowerPoint: Save As .pptx); using default",
            )
        ]


def deck_theme(deck, base_dir: str | Path | None = None) -> tuple[Theme, list[Diagnostic]]:
    """The deck's theme (preset, file or template) with its inline header tokens applied."""
    theme, diags = resolve_theme(deck.theme, base_dir)
    if deck.tokens:
        theme, more = apply_tokens(theme, deck.tokens)
        diags = diags + more
    apply_slide_ink(deck, theme)  # dark/light slide classes and dark `bg=` flip the text color
    return theme, diags


def _hex_rgb(h: str) -> tuple[int, int, int]:
    h = h.lstrip("#")
    return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)


def _mix(a: str, b: str, t: float) -> str:
    """``a`` moved ``t`` (0..1) of the way towards ``b``."""
    ra, rb = _hex_rgb(a), _hex_rgb(b)
    return "#" + "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ra, rb, strict=True))


def _scheme_color(clr) -> str | None:
    if clr is None:
        return None
    for child in clr:
        val = child.get("val") if child.tag == f"{{{_A}}}srgbClr" else child.get("lastClr")
        if val and re.fullmatch(r"[0-9A-Fa-f]{6}", val):
            return "#" + val.upper()
    return None


def _theme_xml(zf: zipfile.ZipFile) -> etree._Element:
    names = zf.namelist()
    # the theme related to the master is normally theme1.xml
    target = "ppt/theme/theme1.xml" if "ppt/theme/theme1.xml" in names else None
    if target is None:
        raise ValueError("no ppt/theme/theme1.xml (not a PowerPoint file?)")
    return etree.fromstring(zf.read(target))


def _from_office(path: Path) -> Theme:
    with zipfile.ZipFile(path) as zf:
        root = _theme_xml(zf)
    scheme = root.find(".//a:clrScheme", _NS)
    if scheme is None:
        raise ValueError("theme1.xml has no a:clrScheme")
    c = {
        k: _scheme_color(scheme.find(f"a:{k}", _NS))
        for k in (
            "dk1",
            "lt1",
            "dk2",
            "lt2",
            "accent1",
            "accent2",
            "accent3",
            "accent4",
            "accent5",
            "accent6",
        )
    }
    bg, fg = c["lt1"] or "#FFFFFF", c["dk1"] or "#1F2937"
    colors = dict(DEFAULT.colors)
    colors.update(
        bg=bg,
        fg=fg,
        primary=c["accent1"] or colors["primary"],
        secondary=c["accent2"] or colors["secondary"],
        accent=c["accent3"] or c["accent2"] or colors["accent"],
        muted=c["dk2"] or _mix(fg, bg, 0.5),
        surface=_mix(bg, fg, 0.06),
        border=_mix(bg, fg, 0.22),
    )
    fonts = Fonts(**DEFAULT.fonts.model_dump())
    fs = root.find(".//a:fontScheme", _NS)
    if fs is not None:
        for key, field in (("majorFont", "heading"), ("minorFont", "body")):
            latin = fs.find(f"a:{key}/a:latin", _NS)
            face = latin.get("typeface") if latin is not None else None
            if face and not face.startswith("+"):
                setattr(fonts, field, face)
        ea = fs.find("a:minorFont/a:ea", _NS)
        if ea is not None and ea.get("typeface") and not ea.get("typeface").startswith("+"):
            fonts.ea = ea.get("typeface")
    return DEFAULT.model_copy(
        update={"name": path.stem, "colors": colors, "fonts": fonts, "template": str(path)}
    )


def _merge(base: dict, over: dict) -> dict:
    out = dict(base)
    for k, v in over.items():
        out[k] = _merge(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
    return out


def _from_yaml(path: Path) -> Theme:
    import yaml

    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise ValueError("YAML theme must be a mapping of Theme fields")
    data.setdefault("extends", "default")  # a theme file extends `default` unless it says otherwise
    tpl = data.get("template")
    if tpl and not Path(tpl).is_absolute():
        data["template"] = str((path.parent / tpl).resolve())
    return theme_from_data(data, path.stem)


def template_size(theme: Theme) -> str | None:
    """The template's slide size as ``"<w>emux<h>emu"`` (accepted by ``units.slide_size``), or None."""
    if not theme.template:
        return None
    try:
        with zipfile.ZipFile(theme.template) as zf:
            sz = etree.fromstring(zf.read("ppt/presentation.xml")).find("p:sldSz", _NS)
        return f"{int(sz.get('cx'))}emux{int(sz.get('cy'))}emu"
    except Exception:
        return None


_FOOTER_TOP: dict[tuple[str, int, int], int | None] = {}


def _ph_tops(root) -> dict[str, int]:
    """y (EMU) of the dt / ftr / sldNum placeholders of a master or layout XML that carry their own xfrm."""
    out: dict[str, int] = {}
    for sp in root.iter(f"{{{_P}}}sp"):
        ph = sp.find(".//p:nvPr/p:ph", _NS)
        off = sp.find("p:spPr/a:xfrm/a:off", _NS)
        if ph is not None and off is not None and ph.get("type") in _FOOTERISH:
            try:
                out[ph.get("type")] = int(off.get("y"))
            except (TypeError, ValueError):
                pass
    return out


def footer_top(theme: Theme) -> int | None:
    """Top (EMU) of the template's footer zone: the smallest y of its date / footer / slide-number
    placeholders on the master and the layouts, or None without a template or without such placeholders.

    Read from the file once per path (cached by mtime and size). Never raises.
    """
    if not theme.template:
        return None
    try:
        st = Path(theme.template).stat()
        key = (str(theme.template), st.st_mtime_ns, st.st_size)
    except OSError:
        return None
    if key in _FOOTER_TOP:
        return _FOOTER_TOP[key]
    tops: list[int] = []
    try:
        with zipfile.ZipFile(theme.template) as zf:
            for name in zf.namelist():
                if re.fullmatch(r"ppt/slide(Master|Layout)s/slide(Master|Layout)\d+\.xml", name):
                    tops += _ph_tops(etree.fromstring(zf.read(name))).values()
    except Exception:
        tops = []
    _FOOTER_TOP[key] = min(tops) if tops else None
    return _FOOTER_TOP[key]


# --------------------------------------------------------------------------- opening a template


def open_template(path: str | Path):
    """Open ``path`` (.pptx or .potx) as a Presentation with all of its slides removed."""
    raw = Path(path).read_bytes()
    with zipfile.ZipFile(io.BytesIO(raw)) as zin:
        ct = zin.read("[Content_Types].xml").decode("utf-8")
        if _CT_TEMPLATE in ct:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        data = ct.replace(_CT_TEMPLATE, _CT_PRESENTATION).encode("utf-8")
                    zout.writestr(item, data)
            raw = buf.getvalue()
    from pptx import Presentation

    prs = Presentation(io.BytesIO(raw))
    _drop_slides(prs)
    return prs


def _drop_slides(prs) -> None:
    lst = prs.slides._sldIdLst
    for sld_id in list(lst):
        prs.part.drop_rel(sld_id.get(qn("r:id")))
        lst.remove(sld_id)
    root = prs.part._element
    for tag in ("p:custShowLst",):  # references removed slides
        for el in root.findall(qn(tag)):
            root.remove(el)
    for ext in root.findall(f".//{qn('p:ext')}"):  # section list references slide ids
        if ext.get("uri") == _SECTION_EXT:
            ext.getparent().remove(ext)


# --------------------------------------------------------------------------- layouts

_ORDER = {
    "cover": (("title",), ("title slide", "cover")),
    "section": (("secHead",), ("section header", "section")),
    "title": (("titleOnly",), ("title only",)),
    "blank": (("blank",), ("blank",)),
}
_FOOTERISH = {"dt", "ftr", "sldNum"}


def _ph_types(layout) -> list[str]:
    return [ph.get("type", "obj") for ph in layout._element.findall(".//" + qn("p:ph"))]


def _has_title(layout) -> bool:
    return any(t in ("title", "ctrTitle") for t in _ph_types(layout))


def pick_layout(prs, kind: str):
    """Layout for ``kind`` in cover | section | title | blank. Matches by type, then name, then fallback."""
    layouts = list(prs.slide_layouts)
    types, names = _ORDER[kind]
    for lo in layouts:
        if lo._element.get("type") in types:
            return lo
    for lo in layouts:
        if lo.name.strip().lower() in names:
            return lo
    for lo in layouts:
        if any(n in lo.name.lower() for n in names):
            return lo
    if kind == "blank":
        pool = layouts
        score = lambda lo: len([t for t in _ph_types(lo) if t not in _FOOTERISH])  # noqa: E731
    else:
        pool = [lo for lo in layouts if _has_title(lo)] or layouts
        score = lambda lo: len([t for t in _ph_types(lo) if t not in _FOOTERISH | {"title", "ctrTitle"}])  # noqa: E731
    return min(pool, key=score)


# --------------------------------------------------------------------------- footer / slide number


def clone_footer(slide, kind: str, text: str) -> bool:
    """Put a native footer (``kind="ftr"``) or slide-number (``"sldNum"``) placeholder on ``slide``.

    Uses the layout's placeholder (inherits geometry and style), or the master's with its geometry copied.
    Returns False when neither has one (caller keeps its own text box).
    """
    layout = slide.slide_layout
    src, from_layout = None, True
    for holder, is_layout in ((layout, True), (layout.slide_master, False)):
        for sp in holder._element.findall(".//" + qn("p:sp")):
            ph = sp.find(".//" + qn("p:ph"))
            if ph is not None and ph.get("type") == kind:
                src, from_layout = sp, is_layout
                break
        if src is not None:
            break
    if src is None:
        return False
    sp_tree = slide.shapes._spTree
    next_id = max([int(i) for i in sp_tree.xpath("//p:cNvPr/@id")] + [1]) + 1
    ph = src.find(".//" + qn("p:ph"))
    name = "Footer" if kind == "ftr" else "Slide Number"
    body = (
        f'<a:fld id="{{{str(uuid.uuid4()).upper()}}}" type="slidenum">'
        f'<a:rPr lang="en-US"/><a:t>{text}</a:t></a:fld>'
        if kind == "sldNum"
        else f'<a:r><a:rPr lang="en-US"/><a:t>{_esc(text)}</a:t></a:r>'
    )
    xml = (
        f'<p:sp xmlns:p="{_P}" xmlns:a="{_A}"><p:nvSpPr><p:cNvPr id="{next_id}" name="{name}"/>'
        f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr/></p:nvSpPr><p:spPr/>'
        f"<p:txBody><a:bodyPr/><a:lstStyle/><a:p>{body}</a:p></p:txBody></p:sp>"
    )
    sp = etree.fromstring(xml)
    sp.find(".//" + qn("p:nvPr")).append(copy.deepcopy(ph))
    if not from_layout:
        xfrm = src.find(qn("p:spPr") + "/" + qn("a:xfrm"))
        if xfrm is not None:
            sp.find(qn("p:spPr")).append(copy.deepcopy(xfrm))
    sp_tree.append(sp)
    return True


def _esc(s: str) -> str:
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
