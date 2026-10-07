"""Tables, charts, images and code -> native PowerPoint objects."""

from __future__ import annotations

import math
from pathlib import Path

from lxml import etree
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import (
    XL_CHART_TYPE,
    XL_LABEL_POSITION,
    XL_LEGEND_POSITION,
    XL_MARKER_STYLE,
    XL_TICK_LABEL_POSITION,
)
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.text.text import Font, TextFrame
from pptx.util import Emu, Pt
from pygments import lex
from pygments.lexers import TextLexer, get_lexer_by_name
from pygments.token import Comment, Keyword, Name, Number, Operator, String

from ..ir import Chart, Code, Image, Paragraph, Placed, Run, Series, Style, Table
from ..layout import chartnote, measure
from ..layout.css import border_spec, cell_insets
from ..layout.tablehl import hl_cell, hl_cols, hl_names, hl_rows, join_names
from ..layout.tables import column_widths, compact_header, pinned, table_grid
from ..theme import Theme
from ..units import EMU_PER_PT, to_emu
from . import waterfall as wfall
from .axis import (
    axis_shown,
    label_collisions,
    label_pt,
    legend_pt,
    line_axis,
    pie_label_pt,
    pie_percent,
    resolve_axis,
)
from .charthl import apply_hl, pin_plot
from .effects import apply_fill, apply_shadow
from .text import _ANCHOR, fill_text
from .util import RenderCtx, emu, hex6, rgb

CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked-bar": XL_CHART_TYPE.BAR_STACKED,
    "stacked-column": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "radar": XL_CHART_TYPE.RADAR_MARKERS,
    "waterfall": XL_CHART_TYPE.COLUMN_STACKED,
}


def flag(v, default: bool = False) -> bool:
    if v is None:
        return default
    if isinstance(v, str):
        return v.strip().lower() in ("1", "true", "on", "yes", "y")
    return bool(v)


# --------------------------------------------------------------------------- table


def _cell_border(tcPr, color: str, width: int = 6350, dash: str | None = None) -> None:
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag))
        ln.set("w", str(width))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", color)
        if dash in _DASH_VAL:
            etree.SubElement(ln, qn("a:prstDash")).set("val", _DASH_VAL[dash])
        tcPr.insert(i, ln)


_DASH_VAL = {"dash": "dash", "dot": "sysDot"}
_FILL_TAGS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")


def _cell_side(tcPr, side: str, spec: tuple[float, str, str] | None, theme: Theme) -> None:
    """Replace one border of a cell (``a:lnL`` / ``lnR`` / ``lnT`` / ``lnB``); ``spec`` None = no line."""
    tag = {"left": "a:lnL", "right": "a:lnR", "top": "a:lnT", "bottom": "a:lnB"}[side]
    old = tcPr.find(qn(tag))
    ln = etree.Element(qn(tag))
    if spec is None:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    else:
        width, dash, color = spec
        ln.set("w", str(round(width * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", hex6(theme, color))
        if dash in _DASH_VAL:
            etree.SubElement(ln, qn("a:prstDash")).set("val", _DASH_VAL[dash])
    if old is not None:
        old.addprevious(ln)
        tcPr.remove(old)
    else:
        tcPr.insert(0, ln)


def _cell_borders(rc: RenderCtx, tcPr, st: Style | None) -> None:
    """CSS borders of one cell: the uniform ``border`` first, then per-side ``border-*`` over it."""
    if st is None:
        return
    uniform = None
    if st.line or st.line_width is not None or st.line_dash:
        if st.line_width is not None and st.line_width <= 0:
            for side in ("left", "right", "top", "bottom"):
                _cell_side(tcPr, side, None, rc.theme)
        else:
            uniform = (
                st.line_width if st.line_width is not None else 0.5,
                st.line_dash or "solid",
                st.line or rc.theme.table_border,
            )
            for side in ("left", "right", "top", "bottom"):
                _cell_side(tcPr, side, uniform, rc.theme)
    for side in ("left", "right", "top", "bottom"):
        v = getattr(st, f"border_{side}")
        if v is not None:
            _cell_side(tcPr, side, border_spec(v), rc.theme)


def _cell_fill(rc: RenderCtx, cell, st: Style) -> None:
    """Cell background from a (CSS) style: solid, #RRGGBBAA, opacity or a gradient."""
    if st.fill and st.fill.lower().startswith("url("):
        rc.diag(
            "css-unsupported",
            "a background image on a table cell is not drawn",
            "use a color or gradient for table cells",
        )
        return
    tmp = etree.Element(qn("a:spPr"))
    tcPr = cell._tc.get_or_add_tcPr()
    if not apply_fill(rc, tmp, st):
        for tag in _FILL_TAGS:
            for old in tcPr.findall(qn(tag)):
                tcPr.remove(old)
        anchor = (
            tcPr.find(qn("a:headers"))
            if tcPr.find(qn("a:headers")) is not None
            else tcPr.find(qn("a:extLst"))
        )
        nofill = etree.Element(qn("a:noFill"))
        anchor.addprevious(nofill) if anchor is not None else tcPr.append(nofill)
        return
    new = tmp[0]
    for tag in _FILL_TAGS:
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    last = None
    for child in tcPr:
        if etree.QName(child).localname in ("lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D"):
            last = child
    if last is not None:
        last.addnext(new)
    else:
        tcPr.insert(0, new)


def _mix(a: str, b: str, k: float) -> str:
    """``k`` x color ``a`` + (1 - k) x color ``b`` (both ``RRGGBB``), as ``#RRGGBB``."""
    return "#" + "".join(
        f"{round(int(a[i : i + 2], 16) * k + int(b[i : i + 2], 16) * (1 - k)):02X}" for i in (0, 2, 4)
    )


def _gantt_grid(tbl, t: Table, theme: Theme, border: str, fill_of: dict, nrows: int, ncols: int) -> None:
    """Gantt tables: the vertical lines between period columns of the body are lightened (the bars carry the
    chart); ``render.gantt_grid`` is the share of the border color that stays."""
    k = theme.render.gantt_grid
    for r in range(t.header_rows, nrows):
        for c in range(1, ncols):
            fill = fill_of.get((r, c), "bg")
            light = _mix(hex6(theme, border), hex6(theme, fill), k)
            tcPr = tbl.cell(r, c)._tc.get_or_add_tcPr()
            if c > 1:
                _cell_side(tcPr, "left", (0.5, "solid", light), theme)
            if c < ncols - 1:
                _cell_side(tcPr, "right", (0.5, "solid", light), theme)


def add_table(rc: RenderCtx, slide, pl: Placed, name: str) -> None:
    t: Table = pl.element  # type: ignore[assignment]
    theme = rc.theme
    nrows, ncols, anchors = table_grid(t)
    gf = slide.shapes.add_table(nrows, ncols, Emu(pl.x), Emu(pl.y), Emu(pl.w), Emu(max(pl.h, nrows * 100000)))
    hl = hl_cell(t)
    cols = sorted(hl_cols(t))
    gf.name = (
        name
        + (f" hl={join_names(hl_names(t))}" if hl_rows(t) else "")
        + (f" hlcol={','.join(str(c + 1) for c in cols)}" if cols else "")
    )  # the importer reads it back
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.attrib.pop(attr, None)
    for child in list(tblPr):
        if etree.QName(child).localname == "tableStyleId":
            tblPr.remove(child)
    cw = pl.element.attrs.get("_col_w") or column_widths(t, ncols, anchors, pl.w)
    rh = pl.element.attrs.get("_row_h") or [max(pl.h // nrows, 1)] * nrows
    if pl.element.attrs.get("_row_h") and not pinned(
        t
    ):  # a stretched table: header rows stay compact, the body takes the slack
        rh = compact_header(t, anchors, cw, pl.style, pl.font_scale, rh, measure.tokens().table_header_max)
    for i, w in enumerate(cw[:ncols]):
        tbl.columns[i].width = Emu(int(w))
    for i, h in enumerate(rh[:nrows]):
        tbl.rows[i].height = Emu(int(h))
    if pinned(t):  # the frame is exactly its pinned rows tall
        gf.height = Emu(int(sum(rh[:nrows])))
    border = hex6(theme, theme.table_border)
    body_fill = theme.table_body_fill_of(pl.style.fill)  # CSS `table { background }` = the body cell fill
    t_border = None
    if pl.style.line or pl.style.line_width is not None:  # CSS `table { border }`: every cell border
        t_border = (
            pl.style.line_width if pl.style.line_width is not None else 0.5,
            pl.style.line_dash or "solid",
        )
        border = hex6(theme, pl.style.line or theme.table_border)
    zebra = "zebra" in t.classes or flag(t.attrs.get("zebra"))
    fill_of: dict[tuple[int, int], str] = {}
    # fill every grid cell first (merged-away cells included) so nothing falls back to a default style
    for r in range(nrows):
        for c in range(ncols):
            cell = tbl.cell(r, c)
            hdr = r < t.header_rows
            fill = theme.table_cell_fill(r, c, t.header_rows, t.header_cols, body_fill, zebra, hl(r, c))
            fill_of[(r, c)] = fill
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                rgb(theme, fill)
                if not fill.lower().startswith(("linear", "radial"))
                else rgb(theme, "surface")
            )
            _cell_border(
                cell._tc.get_or_add_tcPr(),
                border,
                round(t_border[0] * 12700) if t_border else 6350,
                t_border[1] if t_border else None,
            )
            if fill.lower().startswith(("linear", "radial")):
                _cell_fill(rc, cell, Style(fill=fill))
    if "gantt" in t.classes and not t_border and 0 <= theme.render.gantt_grid < 1:
        _gantt_grid(tbl, t, theme, border, fill_of, nrows, ncols)
    grouped = t.header_rows > 0 and any(ct.colspan > 1 for r, _c, ct in anchors if r < t.header_rows)
    for r, c, ct in anchors:
        cell = tbl.cell(r, c)
        rs = min(max(ct.rowspan, 1), nrows - r)
        cs = min(max(ct.colspan, 1), ncols - c)
        if rs > 1 or cs > 1:
            try:
                cell.merge(tbl.cell(r + rs - 1, c + cs - 1))
            except Exception:
                rc.diag(
                    "table-merge",
                    f"cannot merge cell at row {r + 1}, col {c + 1}",
                    "check the < and ^ markers",
                )
        hdr = r < t.header_rows
        st = theme.table_cell_style(
            pl.style,
            hdr,
            c < t.header_cols,
            cs,
            ct.style,
            theme.table_hl_fill_of(body_fill) if hl(r, c) else None,
        )
        if grouped and c >= t.header_cols and not (ct.style and ct.style.align) and (hdr or cs > 1):
            st = st.merged(
                Style(align="center")
            )  # grouped header: sub-headers and spanning bars are centered
        if ct.style and (ct.style.fill or ct.style.opacity is not None):
            _cell_fill(rc, cell, Style(fill=ct.style.fill or fill_of[(r, c)], opacity=ct.style.opacity))
        for rr in range(r, r + rs):  # a merged cell: the covered cells share its borders
            for cc in range(c, c + cs):
                _cell_borders(rc, tbl.cell(rr, cc)._tc.get_or_add_tcPr(), ct.style)
        cl, ctp, cr, cb = cell_insets(ct.style, *measure.cell_pad())
        if (  # `table.num_pad=0.6in`: right-aligned (numeric) cells keep this inset on the right
            theme.table_num_pad is not None
            and (st.align == "right" or (ct.style and ct.style.align == "right"))
            and not (ct.style and (ct.style.padding is not None or ct.style.padding_right is not None))
        ):
            cr = min(max(to_emu(theme.table_num_pad), 0), max(sum(cw[c : c + cs]) // 2, 0))
        fill_text(
            rc,
            cell.text_frame,
            ct.paragraphs,
            st.merged(Style(valign=st.valign or "middle")),
            pl.font_scale,
            inset=0,
            box_w=round(sum(cw[c : c + cs]) - cl - cr),
        )
        cell.vertical_anchor = _ANCHOR.get(st.valign or "middle", MSO_ANCHOR.MIDDLE)
        cell.margin_left, cell.margin_right = Emu(cl), Emu(cr)
        cell.margin_top, cell.margin_bottom = Emu(ctp), Emu(cb)


# --------------------------------------------------------------------------- chart


def _num(v) -> float | None:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if f == f and abs(f) != float("inf") else None


_ZERO_BASE = ("column", "bar", "stacked-column", "stacked-bar", "area")


def _all_nonneg(ch: Chart) -> bool:
    vals = [v for s in ch.series for v in s.values if v is not None]
    return bool(vals) and min(vals) >= 0


def _chart_font(chart, theme, face: str, size_pt: float, color) -> None:
    """Chart-wide text: size, color, latin + East Asian typeface."""
    f = chart.font
    f.size = Pt(size_pt)
    f.color.rgb = color
    f.name = face
    d_rpr = chart._chartSpace.find(qn("c:txPr")).find(qn("a:p")).find(qn("a:pPr")).find(qn("a:defRPr"))
    latin = d_rpr.find(qn("a:latin"))
    ea = etree.Element(qn("a:ea"))
    ea.set("typeface", theme.fonts.ea)
    if latin is not None:
        latin.addnext(ea)
    else:
        d_rpr.append(ea)


def _legend_pos(opts: dict, pie: bool, nseries: int):
    legend = opts.get("legend")
    pos_map = {
        "bottom": XL_LEGEND_POSITION.BOTTOM,
        "top": XL_LEGEND_POSITION.TOP,
        "left": XL_LEGEND_POSITION.LEFT,
        "right": XL_LEGEND_POSITION.RIGHT,
    }
    default = XL_LEGEND_POSITION.RIGHT if pie else XL_LEGEND_POSITION.BOTTOM
    if isinstance(legend, str) and legend.strip().lower() in pos_map:
        return pos_map[legend.strip().lower()]
    if isinstance(legend, str) and legend.strip().lower() == "none":
        return None
    if legend is not None:
        return default if flag(legend, True) else None
    return default if (pie or nseries > 1) else None


def _positive_axis_ids(chart) -> None:
    """python-pptx templates write negative ``c:axId`` / ``c:crossAx`` (``xs:unsignedInt`` in ECMA-376).

    Every negative id becomes a positive one, the same for ``axId`` and ``crossAx`` references in the chart.
    """
    els = [e for e in chart._chartSpace.iter(qn("c:axId"), qn("c:crossAx"))]
    ids = []
    for e in els:
        try:
            ids.append(int(e.get("val", "0")))
        except ValueError:
            ids.append(0)
    used = {i for i in ids if i >= 0}
    mapping: dict[int, int] = {}
    for old in dict.fromkeys(i for i in ids if i < 0):
        new = min(abs(old), 2**32 - 2)
        while new in used:
            new += 1
        used.add(new)
        mapping[old] = new
    for e, i in zip(els, ids, strict=True):
        if i in mapping:
            e.set("val", str(mapping[i]))


def _luminance(hex_rgb: str) -> float:
    def lin(v: int) -> float:
        x = v / 255
        return x / 12.92 if x <= 0.03928 else ((x + 0.055) / 1.055) ** 2.4

    r, g, b = (int(hex_rgb[i : i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * lin(r) + 0.7152 * lin(g) + 0.0722 * lin(b)


def label_color_on(fill_hex: str, theme: Theme | None = None) -> str:
    """Light ink on dark fills, dark ink on light ones (WCAG-style: the higher contrast wins).

    The inks are ``theme.render.ink_light`` / ``ink_dark`` ('RRGGBB' is returned).
    """
    theme = theme or Theme(name="none")
    light, dark_hex = hex6(theme, theme.render.ink_light), hex6(theme, theme.render.ink_dark)
    lum = _luminance(fill_hex)
    dark = _luminance(dark_hex)
    return light if 1.05 / (lum + 0.05) > (lum + 0.05) / (dark + 0.05) else dark_hex


def _ink_hex(theme: Theme, fill: str, fallback: str, *backs: str) -> str:
    """'RRGGBB' ink readable on ``fill`` (and on every ``backs`` color when one ink can do both)."""
    return theme.chart_label_ink(fill, *backs) or fallback


_PIE_POS = {
    "center": XL_LABEL_POSITION.CENTER,
    "inside_end": XL_LABEL_POSITION.INSIDE_END,
    "outside_end": XL_LABEL_POSITION.OUTSIDE_END,
    "best_fit": XL_LABEL_POSITION.BEST_FIT,
}


_LABEL_POS = {
    "outside": XL_LABEL_POSITION.OUTSIDE_END,
    "inside": XL_LABEL_POSITION.INSIDE_END,
    "center": XL_LABEL_POSITION.CENTER,
    "above": XL_LABEL_POSITION.ABOVE,
    "below": XL_LABEL_POSITION.BELOW,
    "left": XL_LABEL_POSITION.LEFT,
    "right": XL_LABEL_POSITION.RIGHT,
}
_PIE_WORDS = {"outside": "outside_end", "inside": "inside_end", "center": "center", "best": "best_fit"}


def _pie_pos(theme: Theme):
    """Default label position of a pie wedge (``render.chart_pie_label_pos``; unknown words = best fit)."""
    return _PIE_POS.get(str(theme.render.chart_pie_label_pos).strip().lower().replace("-", "_"))


def _pie_point_labels(ser, pal, n, theme, kind, fg, size, nf, lab_pct, vals=(), exact=None) -> None:
    """Per-slice ``c:dLbl`` so each label has its own readable ink (slices differ in fill).

    A doughnut label always sits on its slice. A pie label sits where ``render.chart_pie_label_pos`` says
    (inside the wedge: ink chosen for the slice); a wedge below ``render.chart_pie_label_min`` of the total
    goes outside (ink for the page), and best fit may land either way (one ink that reads on both).
    """
    rt = theme.render
    pt = pie_label_pt(size, rt, exact)
    dls = ser.data_labels  # series-level dLbls override the plot-level ones: repeat the shared settings
    dls.font.size = Pt(pt)
    dls.font.bold = rt.chart_pie_label_bold
    dls.font.color.rgb = fg
    if lab_pct:
        dls.show_value, dls.show_percentage = False, True
        dls.number_format = nf if (nf and "%" in nf) else "0%"
        dls.number_format_is_linked = False
    else:
        dls.show_value = True
        if nf:
            dls.number_format = nf
            dls.number_format_is_linked = False
    pos = _pie_pos(theme)
    if kind == "pie" and pos is not None:
        dls.position = pos
    bg, fg_hex = hex6(theme, "bg"), hex6(theme, "fg")
    nums = [_num(v) for v in vals]
    total = sum(v for v in nums if v and v > 0)
    for pi in range(n):
        fill = pal[pi % len(pal)]
        share = (nums[pi] / total) if pi < len(nums) and nums[pi] and total > 0 else None
        out = kind == "pie" and (
            pos == XL_LABEL_POSITION.OUTSIDE_END
            or (share is not None and share < rt.chart_pie_label_min and pos != XL_LABEL_POSITION.BEST_FIT)
        )
        both = kind == "pie" and pos in (None, XL_LABEL_POSITION.BEST_FIT)
        ink = fg_hex if out else _ink_hex(theme, fill, fg_hex, *([bg] if both else []))
        dl = ser.points[pi].data_label
        dl.font.size = Pt(pt)
        dl.font.bold = rt.chart_pie_label_bold
        dl.font.color.rgb = RGBColor.from_string(ink)
        el = dl._dLbl
        if el is None:
            continue
        for tag, val in (("c:showVal", not lab_pct), ("c:showPercent", lab_pct)):
            e = el.find(qn(tag))
            if e is not None:
                e.set("val", "1" if val else "0")
        if lab_pct or nf:
            fmt = (nf if (nf and "%" in nf) else "0%") if lab_pct else nf
            nfe = el.makeelement(qn("c:numFmt"), {"formatCode": fmt, "sourceLinked": "0"})
            anchor = el.find(qn("c:spPr"))
            if anchor is None:
                anchor = el.find(qn("c:txPr"))
            anchor.addprevious(nfe)
        if kind == "pie":
            if out:
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            elif pos is not None:
                dl.position = pos


_BAR_WORDS = {"outside", "inside", "center"}


def _point_labels(
    ser, pos: dict[int, XL_LABEL_POSITION], default, size_pt: float, ink, nf, inside_ink=None
) -> None:
    """Data labels with a position of their own per point (``labels=above,below,...``, collisions).

    A series-level ``c:dLbls`` overrides the plot-level one, so its shared settings are repeated. ``ink``
    is the label color; ``inside_ink(i)`` (bars) gives the ink of point ``i`` when its label sits inside the
    bar (chosen for the fill under it)."""
    dls = ser.data_labels
    dls.show_value = True
    dls.font.size = Pt(size_pt)
    dls.font.color.rgb = ink
    dls.position = default
    if nf:
        dls.number_format, dls.number_format_is_linked = nf, False
    for i, where in sorted(pos.items()):
        dl = ser.points[i].data_label
        dl.font.size = Pt(size_pt)
        inside = where in (
            XL_LABEL_POSITION.INSIDE_END,
            XL_LABEL_POSITION.CENTER,
            XL_LABEL_POSITION.INSIDE_BASE,
        )
        dl.font.color.rgb = (
            RGBColor.from_string(inside_ink(i)) if (inside and inside_ink is not None) else ink
        )
        dl.position = where
        el = dl._dLbl
        if el is not None and nf:
            nfe = el.makeelement(qn("c:numFmt"), {"formatCode": nf, "sourceLinked": "0"})
            anchor = el.find(qn("c:spPr"))
            if anchor is None:
                anchor = el.find(qn("c:txPr"))
            anchor.addprevious(nfe)


def _below_labels(ser, idxs: list[int], theme: Theme, size_pt: float, fg, nf) -> None:
    """Move the data labels of the points ``idxs`` of a line series below the point (the series default is
    above). A series-level ``c:dLbls`` overrides the plot-level one, so its shared settings are repeated."""
    dls = ser.data_labels
    dls.show_value = True
    dls.font.size = Pt(size_pt)
    dls.font.color.rgb = fg
    dls.position = XL_LABEL_POSITION.ABOVE
    if nf:
        dls.number_format, dls.number_format_is_linked = nf, False
    for i in idxs:
        dl = ser.points[i].data_label
        dl.font.size = Pt(size_pt)
        dl.font.color.rgb = fg
        dl.position = XL_LABEL_POSITION.BELOW
        el = dl._dLbl
        if el is not None and nf:
            nfe = el.makeelement(qn("c:numFmt"), {"formatCode": nf, "sourceLinked": "0"})
            anchor = el.find(qn("c:spPr"))
            if anchor is None:
                anchor = el.find(qn("c:txPr"))
            anchor.addprevious(nfe)


def _waterfall_plan(ch: Chart, ser: Series, opts: dict, theme: Theme) -> dict:
    """Bars, the six native stacked series, the axis and the label texts of a waterfall."""
    totals = [int(i) for i in ch.options.get("totals", []) if isinstance(i, (int, float))]
    vals = [_num(v) for v in ser.values]
    bs = wfall.bars(vals, totals)
    got = wfall.axis(bs, theme.render)
    ax = got or (0.0, 0.0, 0.0)
    span = (ax[1] - ax[0]) or max((b[2] - b[1] for b in bs), default=1.0) or 1.0
    cols = wfall.columns(bs, span * 0.07)
    rt = theme.render
    colors = {
        wfall.UP: hex6(theme, rt.waterfall_up),
        wfall.DOWN: hex6(theme, rt.waterfall_down),
        wfall.TOTAL: hex6(theme, rt.waterfall_total),
    }
    colors[wfall.CROSS] = colors[wfall.UP]
    return {
        "bars": bs,
        "axis": got,
        "top": max((b[2] for b in bs if b[0] != "gap"), default=0.0),
        "colors": colors,
        "series": [
            Series(name=(ser.name if i == wfall.PAD else n), values=c)
            for i, (n, c) in enumerate(zip(wfall.SERIES, cols, strict=True))
        ],
    }


def _style_waterfall_series(ser, si: int, plan: dict, theme: Theme, lab_pt: float, fg, lab_on, nf) -> None:
    if si == wfall.BASE or si == wfall.PAD:  # invisible: no fill, no line
        ser.format.fill.background()
        ser.format.line.fill.background()
    else:
        ser.format.fill.solid()
        color = plan["colors"][si]
        ser.format.fill.fore_color.rgb = RGBColor.from_string(color)
        if si == wfall.CROSS:  # the below-zero part of a crossing bar takes the color of that bar
            for i, b in enumerate(plan["bars"]):
                if b[3] is not None and b[0] == "down":
                    pt = ser.points[i]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = RGBColor.from_string(plan["colors"][wfall.DOWN])
        ser.format.line.fill.background()
    if si != wfall.PAD or not lab_on:
        return
    ink = fg
    for i, (kind, _lo, _hi, shown) in enumerate(plan["bars"]):
        if kind == "gap":
            continue
        dl = ser.points[i].data_label
        tf = dl.text_frame
        tf.text = wfall.fmt_num(shown, nf, sign=kind in ("up", "down"))
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(lab_pt)
        run.font.color.rgb = ink
        dl.position = XL_LABEL_POSITION.INSIDE_BASE


def _hide_labels(sdl, idxs: list[int]) -> None:
    """Delete the data labels of the points ``idxs`` (a label that cannot fit its segment is hidden)."""
    dls = sdl._element
    for k, i in enumerate(sorted(idxs)):
        dl = parse_xml(
            f'<c:dLbl xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
            f'<c:idx val="{i}"/><c:delete val="1"/></c:dLbl>'
        )
        dls.insert(k, dl)


def _hide_legend_entry(chart, idx: int) -> None:
    leg = chart._chartSpace.find(".//" + qn("c:legend"))
    if leg is None:
        return
    anchor = leg.find(qn("c:legendPos"))
    e = leg.makeelement(qn("c:legendEntry"), {})
    etree.SubElement(e, qn("c:idx")).set("val", str(idx))
    etree.SubElement(e, qn("c:delete")).set("val", "1")
    anchor.addnext(e) if anchor is not None else leg.insert(0, e)


def _waterfall_legend(chart) -> None:
    """Hide the legend entries of the helper series (base, below zero, label carrier)."""
    leg = chart._chartSpace.find(".//" + qn("c:legend"))
    if leg is None:
        return
    anchor = leg.find(qn("c:legendPos"))
    for idx in (wfall.BASE, wfall.CROSS, wfall.PAD):
        e = leg.makeelement(qn("c:legendEntry"), {})
        etree.SubElement(e, qn("c:idx")).set("val", str(idx))
        etree.SubElement(e, qn("c:delete")).set("val", "1")
        anchor.addnext(e) if anchor is not None else leg.insert(0, e)


TOTAL_CARRIER = "total (hidden)"  # name of the invisible series that carries the stack-total labels
BLANK_CARRIER = (
    " "  # ... on a horizontal stack: LibreOffice deletes the wrong legend entry there, so it is unnamed
)


def stack_totals(ch: Chart, series: list[Series], lab_on: bool, opts: dict, rt) -> list[Series]:
    """``series`` plus a hidden carrier on top whose per-point labels read each stack's sum.

    Wanted for stacked kinds with segment labels (``totals=off`` disables, ``totals=on`` forces); skipped
    when a value is negative (the sum is not the stack's end) or no stack has a positive sum.
    """
    t = str(opts.get("totals", "")).strip().lower()
    on = flag(t) if t in ("1", "true", "on", "yes", "y", "0", "false", "off", "no", "n") else lab_on
    if not on or ch.kind not in ("stacked-bar", "stacked-column") or not series or not _all_nonneg(ch):
        return series
    n = max(len(s.values) for s in series)
    sums = [round(sum(_num(s.values[i]) or 0.0 for s in series if i < len(s.values)), 10) for i in range(n)]
    top = max(sums, default=0.0)
    if top <= 0:
        return series
    pad = top * rt.chart_total_pad
    name = BLANK_CARRIER if ch.kind == "stacked-bar" else TOTAL_CARRIER
    return [*series, Series(name=name, values=[pad if v > 0 else None for v in sums])]


def _style_total_carrier(ser, sums: list[float], theme: Theme, lab_pt: float, fg, nf) -> None:
    ser.format.fill.background()
    ser.format.line.fill.background()
    for i, v in enumerate(sums):
        if v <= 0:
            continue
        dl = ser.points[i].data_label
        tf = dl.text_frame
        tf.text = wfall.fmt_num(v, nf, sign=False)
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(lab_pt)
        run.font.bold = True
        run.font.color.rgb = fg
        dl.position = XL_LABEL_POSITION.INSIDE_BASE


def _wedge_line(pt, theme: Theme, opts: dict) -> None:
    """Outline of a pie / doughnut wedge: ``slice.line=`` else ``render.chart_pie_line`` (+ ``_width``)."""
    from .util import parse_color

    rt = theme.render
    word = str(opts.get("slice_line") or rt.chart_pie_line)
    width = _num(opts.get("slice_line_w"))
    if width is None:
        width = rt.chart_pie_line_width
    hexv, alpha = parse_color(theme, word, "bg")
    if word.strip().lower() == "none" or alpha == 0:
        pt.format.line.fill.background()
        return
    pt.format.line.color.rgb = RGBColor.from_string(hexv)
    if width and width > 0:
        pt.format.line.width = Pt(width)


_DLBL_ORDER = ("showLegendKey", "showVal", "showCatName", "showSerName", "showPercent", "showBubbleSize")


def _set_flag(el, tag: str, val: bool) -> None:
    """Set ``c:<tag>`` of a ``c:dLbls`` / ``c:dLbl``, creating it in schema order when missing."""
    e = el.find(qn(f"c:{tag}"))
    if e is None:
        e = el.makeelement(qn(f"c:{tag}"), {})
        after = [el.find(qn(f"c:{t}")) for t in _DLBL_ORDER[: _DLBL_ORDER.index(tag)]]
        prev = next((a for a in reversed(after) if a is not None), None)
        if prev is not None:
            prev.addnext(e)
        else:
            el.insert(0, e)
    e.set("val", "1" if val else "0")


def _label_font(el) -> Font:
    return TextFrame(el.get_or_add_txPr(), None).paragraphs[0].font


def style_labels(chart, opts: dict, theme: Theme, pie: bool) -> None:
    """``labels.bold=`` / ``labels.color=`` on every data label of the chart and ``labels=...+name`` on a pie
    (category name, a newline, then the value or share). Hidden (deleted) labels are left alone."""
    bold = opts.get("label_bold")
    color = opts.get("label_color")
    name = bool(opts.get("label_name")) and pie
    if bold is None and not color and not name:
        return
    rgb_ = RGBColor.from_string(hex6(theme, str(color))) if color else None
    for el in chart._chartSpace.iter(qn("c:dLbls"), qn("c:dLbl")):
        if el.find(qn("c:delete")) is not None:
            continue
        fonts = [Font(r) for r in el.iterfind(".//" + qn("a:rPr"))]  # a label with its own text (waterfall)
        if el.find(qn("c:txPr")) is not None or not fonts:
            fonts.append(_label_font(el))
        for f in fonts:
            if bold is not None:
                f.bold = bool(bold)
            if rgb_ is not None:
                f.color.rgb = rgb_
        if name:
            _set_flag(el, "showCatName", True)
            if el.find(qn("c:separator")) is None:
                sep = el.makeelement(qn("c:separator"), {})
                sep.text = "\n"
                anchor = el.find(qn("c:showBubbleSize"))
                if anchor is None:
                    anchor = el.find(qn("c:showPercent"))
                if anchor is not None:
                    anchor.addnext(sep)


def add_chart(rc: RenderCtx, slide, pl: Placed, name: str) -> None:
    ch: Chart = pl.element  # type: ignore[assignment]
    theme = rc.theme
    opts = {str(k).lower().replace("-", "_"): v for k, v in ch.options.items()}
    kind = ch.kind
    lpos = str(opts.get("label_pos") or "")
    if kind in ("pie", "doughnut") and lpos in _PIE_WORDS:  # `labels=outside` on a pie
        rt = theme.render.model_copy(update={"chart_pie_label_pos": _PIE_WORDS[lpos]})
        theme = theme.model_copy(update={"render": rt})
    ctype = CHART_TYPES.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    pie = kind in ("pie", "doughnut")
    series = ch.series
    cats_in = list(ch.categories)
    if pie and len(series) > 1 and len(cats_in) <= 1 and all(len(s.values) == 1 for s in series):
        # one slice per CSV row ("APAC,48"): a pie wants one series with a category per slice
        cats_in = [s.name for s in series]
        series = [
            Series(name=ch.categories[0] if ch.categories else "", values=[s.values[0] for s in series])
        ]
    wf = None
    if kind == "waterfall" and series:
        wf = _waterfall_plan(ch, series[0], opts, theme)
        series = wf["series"]
    labels = opts.get("labels", opts.get("data_labels"))
    lab_word = labels.strip().lower() if isinstance(labels, str) else ""
    lab_on = lab_word in ("percent", "value") or flag(labels)
    real_series = len(series)
    series = stack_totals(ch, series, lab_on, opts, theme.render) if not wf else series
    carrier = len(series) > real_series
    ncat = max([len(cats_in), *(len(s.values) for s in series)])
    cats = [str(c) for c in cats_in] or [str(i + 1) for i in range(ncat)]
    cats += [str(i + 1) for i in range(len(cats), ncat)]  # more values than categories
    if kind == "scatter":
        data = XyChartData()
        xs = []
        for i, c in enumerate(cats):
            x = _num(str(c).replace(",", ""))
            xs.append(x if x is not None else float(i + 1))
        for s in series:
            ser = data.add_series(s.name)
            for x, v in zip(xs, s.values, strict=False):
                if _num(v) is not None:
                    ser.add_data_point(x, _num(v))
        if not series:
            data.add_series("")
    else:
        data = CategoryChartData()
        data.categories = cats or [" "]
        for s in series:
            vals = [_num(v) for v in s.values]
            vals = (vals + [None] * len(cats))[: len(cats) or 1]
            data.add_series(s.name, vals)
        if not series:
            data.add_series("", [None] * (len(cats) or 1))
    note_plan = chartnote.plan(ch, pl, theme)
    gf = slide.shapes.add_chart(ctype, Emu(pl.x), Emu(pl.y), Emu(pl.w), Emu(pl.h), data)
    gf.name = name
    chart = gf.chart
    _positive_axis_ids(chart)
    size = chartnote.chart_size(pl, theme)
    lab_exact = chartnote.opt_pt(pl, "label_size")  # `size=16,14`: data labels exactly 16 ...
    tick_exact = chartnote.opt_pt(pl, "tick_size")  # ... value-axis numbers exactly 14
    fg = rgb(theme, pl.style.color or "fg")
    _chart_font(chart, theme, pl.style.font or theme.fonts.body, size, fg)
    if ch.title:
        chart.has_title = True
        tf = chart.chart_title.text_frame
        tf.text = ch.title
        run = tf.paragraphs[0].runs[0]
        run.font.size = Pt(size * theme.render.chart_title_scale)
        run.font.bold = True
        run.font.color.rgb = rgb(theme, "fg")
        chart.chart_title.include_in_layout = False
    else:
        chart.has_title = False
    pos = _legend_pos(opts, pie, 1 if wf else real_series)
    chart.has_legend = pos is not None
    if pos is not None:
        chart.legend.position = pos
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(legend_pt(size, theme.render, chartnote.opt_pt(pl, "legend_size")))
    # colors: explicit list (theme names / hex) else the theme palette
    cl = opts.get("colors")
    if isinstance(cl, str):
        cl = [p.strip() for p in cl.replace(";", ",").split(",") if p.strip()]
    pal = theme.chart_palette(cl)
    # one series of bars with a color per bar (`colors=a,b,c`, one per category): a `c:dPt` fill per point
    point_cols = (
        pal
        if (
            kind in ("column", "bar")
            and not wf
            and real_series == 1
            and isinstance(cl, list)
            and ncat > 1
            and len(cl) == ncat
        )
        else None
    )
    point_words = opts.get("label_points")  # `labels=above,below,...`: a position per category
    point_words = [str(w) for w in point_words] if isinstance(point_words, (list, tuple)) else None
    # number formats
    pct_flag = flag(opts.get("percent"))
    nf = opts.get("fmt") or opts.get("number_format") or opts.get("format")
    nf = str(nf) if nf else ('0"%"' if pct_flag else "#,##0" if flag(opts.get("grouped")) else None)
    lab_pct = pie and lab_on and pie_percent(labels, nf, theme.render)
    plot = chart.plots[0]
    if lab_on and kind not in ("scatter", "waterfall"):  # python-pptx has no data labels for XY series
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(label_pt(kind, ncat, size, theme.render, lab_exact))
        dl.font.color.rgb = fg
        if lab_pct and pie:
            dl.show_value, dl.show_percentage = False, True
            dl.number_format = nf if (nf and "%" in nf) else "0%"
            dl.number_format_is_linked = False
        else:
            dl.show_value = True
            if nf:
                dl.number_format = nf
                dl.number_format_is_linked = False
        if pie:
            dl.font.bold = theme.render.chart_pie_label_bold
        if kind == "pie":
            if (pos := _pie_pos(theme)) is not None:
                dl.position = pos
        elif kind in ("column", "bar"):
            dl.position = _LABEL_POS.get(lpos, XL_LABEL_POSITION.OUTSIDE_END)
        elif kind == "line":
            dl.position = _LABEL_POS.get(lpos, XL_LABEL_POSITION.ABOVE)
    seg_fits_cache: dict = {}
    for plot in chart.plots:
        for si, ser in enumerate(plot.series):
            color = pal[si % len(pal)]
            if carrier and si == real_series:
                tot = [
                    round(sum(_num(s.values[i]) or 0.0 for s in series[:-1] if i < len(s.values)), 10)
                    for i in range(len(series[-1].values))
                ]
                _style_total_carrier(
                    ser, tot, theme, lab_exact or size * theme.render.chart_label_scale, fg, nf
                )
            elif wf:
                _style_waterfall_series(
                    ser, si, wf, theme, lab_exact or size * theme.render.chart_label_scale, fg, lab_on, nf
                )
            elif pie:
                for pi in range(len(cats)):
                    pt = ser.points[pi]
                    pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = RGBColor.from_string(pal[pi % len(pal)])
                    _wedge_line(pt, theme, opts)
                if lab_on:
                    _pie_point_labels(
                        ser, pal, len(cats), theme, kind, fg, size, nf, lab_pct, series[0].values, lab_exact
                    )
            elif kind in ("line", "radar", "scatter"):
                if kind == "scatter":
                    ser.format.line.fill.background()
                else:
                    ser.format.line.color.rgb = RGBColor.from_string(color)
                    ser.format.line.width = Pt(theme.render.chart_line_width)
                if hasattr(ser, "smooth"):
                    ser.smooth = False
                try:
                    ser.marker.format.fill.solid()
                    ser.marker.format.fill.fore_color.rgb = RGBColor.from_string(color)
                    ser.marker.format.line.color.rgb = RGBColor.from_string(color)
                    if (msz := _num(opts.get("marker"))) is not None:  # `marker=9`
                        ser.marker.style = XL_MARKER_STYLE.CIRCLE
                        ser.marker.size = int(msz)
                except Exception:
                    pass
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = RGBColor.from_string(color)
                if point_cols and si == 0:
                    for pi, pc in enumerate(point_cols):
                        pt = ser.points[pi]
                        pt.format.fill.solid()
                        pt.format.fill.fore_color.rgb = RGBColor.from_string(pc)
                if lab_on and kind in ("bar", "column") and lpos in ("inside", "center"):
                    sdl = ser.data_labels  # inside the bar: the ink is chosen per series fill
                    sdl.show_value = True
                    sdl.font.size = Pt(label_pt(kind, ncat, size, theme.render, lab_exact))
                    sdl.font.color.rgb = RGBColor.from_string(
                        _ink_hex(theme, color, label_color_on(color, theme))
                    )
                    sdl.position = _LABEL_POS[lpos]
                    if nf:
                        sdl.number_format = nf
                        sdl.number_format_is_linked = False
                    if point_cols:  # each bar has its own fill: its label gets its own ink
                        _point_labels(
                            ser,
                            dict.fromkeys(range(ncat), _LABEL_POS[lpos]),
                            _LABEL_POS[lpos],
                            label_pt(kind, ncat, size, theme.render),
                            fg,
                            nf,
                            lambda i: _ink_hex(theme, point_cols[i], label_color_on(point_cols[i], theme)),
                        )
                if lab_on and kind in ("stacked-bar", "stacked-column"):  # labels sit inside the fill
                    sdl = ser.data_labels
                    sdl.show_value = True
                    seg_pt = label_pt(kind, ncat, size, theme.render, lab_exact)
                    sdl.font.size = Pt(seg_pt)
                    sdl.font.color.rgb = RGBColor.from_string(
                        _ink_hex(theme, color, label_color_on(color, theme))
                    )
                    sdl.position = _LABEL_POS.get(lpos, XL_LABEL_POSITION.CENTER)
                    # a zero-width segment gets no (clipped) label: the zero section of the format is empty
                    sdl.number_format = (
                        f"{nf};-{nf};;" if nf and ";" not in nf else (nf or "General;-General;;")
                    )
                    sdl.number_format_is_linked = False
                    fits = seg_fits_cache.setdefault("f", chartnote.seg_fits(ch, pl, theme, seg_pt))
                    _hide_labels(sdl, [ci for ci in range(ncat) if fits.get((si, ci)) is False])
    if point_words and lab_on and kind in ("column", "bar"):  # `labels=outside,inside,...`: per category
        for si, ser in enumerate(chart.plots[0].series):
            if si >= real_series:
                continue
            own = pal[si % len(pal)]

            def ink_in(i: int, own=own) -> str:
                c = point_cols[i] if point_cols else own
                return _ink_hex(theme, c, label_color_on(c, theme))

            where = {i: _LABEL_POS[w] for i, w in enumerate(point_words[:ncat]) if w in _LABEL_POS}
            _point_labels(
                ser,
                where,
                _LABEL_POS.get(lpos, XL_LABEL_POSITION.OUTSIDE_END),
                label_pt(kind, ncat, size, theme.render),
                fg,
                nf,
                ink_in,
            )
    if kind == "doughnut":
        hole = chart.plots[0]._element.find(qn("c:holeSize"))
        if hole is not None:
            hole.set("val", "55")
    if wf and pos is not None:  # a legend lists only the visible bars
        _waterfall_legend(chart)
    if carrier and pos is not None and kind == "stacked-column":  # the legend lists only the visible series
        _hide_legend_entry(chart, real_series)
    if kind in ("bar", "column", "stacked-bar", "stacked-column", "waterfall"):
        try:
            rt = theme.render
            gw = _num(opts.get("gap_width"))
            if gw is None:
                gw = rt.chart_gap_few if ncat <= rt.chart_gap_few_cats else rt.chart_gap
            chart.plots[0].gap_width = max(0, min(500, round(gw)))
            if (ov := _num(opts.get("overlap"))) is not None and kind in ("bar", "column"):
                chart.plots[0].overlap = max(-100, min(100, round(ov)))
        except Exception:
            pass
    if hl_names := apply_hl(ch, chart, cats, real_series, wf, theme, pal):
        gf.name = f"{name} hl={','.join(hl_names)}"  # the importer reads it back
    if note_plan and note_plan.plot:
        pin_plot(chart, note_plan.plot)
    if lab_on:
        style_labels(chart, opts, theme, pie)
    if pie:
        return
    try:
        va = chart.value_axis
        if not axis_shown(opts, kind, ncat, lab_on, theme.render):
            va.visible = False
            va.has_major_gridlines = False
        else:
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = rgb(theme, theme.render.chart_grid)
            va.format.line.fill.background()
            if tick_exact:
                va.tick_labels.font.size = Pt(tick_exact)
            if nf:
                va.tick_labels.number_format = nf
                va.tick_labels.number_format_is_linked = False
        lo, hi = _num(opts.get("min")), _num(opts.get("max"))
        lo, hi, unit = resolve_axis(
            kind,
            [s.values for s in series] if not wf else [],
            theme.render,
            lo,
            hi,
            wf["axis"] if wf else None,
            _all_nonneg(ch),
            wf["top"] if wf else None,
            ncat=ncat,
            tight=not axis_shown(opts, kind, ncat, lab_on, theme.render),
        )
        if kind == "line" and lo is None and hi is None and not unit:
            la = line_axis([s.values for s in series], theme.render)
            # data far from zero: the axis starts above it (a flat 0-9 axis hides the trend)
            if la and la[0] > 0:
                lo, hi, unit = la
        if note_plan and note_plan.axis:  # the pointer of a `note=` needs the exact scale
            lo, hi, unit = note_plan.axis
        if (step := _num(opts.get("step"))) and step > 0:  # `step=200`: the author's gridline distance
            span = (hi if hi is not None else 0.0) - (lo or 0.0)
            if hi is not None and span / step > 60:
                rc.diag(
                    "bad-chart-option",
                    f"step={step:g} would draw {span / step:.0f} gridlines on this axis",
                    "use a larger step, or drop step= for the automatic one",
                    line=ch.line,
                )
            else:
                unit = step
                if hi is not None and _num(opts.get("max")) is None and not (note_plan and note_plan.axis):
                    hi = round((lo or 0.0) + math.ceil(span / step - 1e-9) * step, 10)  # ends on a gridline
        if kind == "line" and lab_on:  # close lines: their labels alternate above / below
            lpt = label_pt(kind, ncat, size, theme.render, lab_exact)
            close = label_collisions([s.values for s in series], lo, hi, lpt, pl.h / EMU_PER_PT, theme.render)
            manual = {i: _LABEL_POS[w] for i, w in enumerate(point_words or []) if w in _LABEL_POS}
            if manual:  # the author named the position of these categories: the collision rule yields
                for si, ser in enumerate(chart.plots[0].series):
                    where = {i: XL_LABEL_POSITION.BELOW for i in close.get(si, []) if i not in manual}
                    _point_labels(ser, {**where, **manual}, XL_LABEL_POSITION.ABOVE, lpt, fg, nf)
            else:
                for si, idxs in close.items():
                    _below_labels(chart.plots[0].series[si], idxs, theme, lpt, fg, nf)
        if wf and wf["axis"] and wf["axis"][0] < 0:
            chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
        if unit:
            va.major_unit = unit
        if lo is not None:
            va.minimum_scale = lo
        if hi is not None:
            va.maximum_scale = hi
        if kind in ("bar", "stacked-bar"):  # first category on top, value axis stays at the bottom
            chart.category_axis.reverse_order = True
            va._element.find(qn("c:crosses")).set("val", "max")
        chart.category_axis.format.line.color.rgb = rgb(theme, "border")
        if kind in _ZERO_BASE and any(
            v < 0 for s in ch.series for v in s.values if v is not None
        ):  # negative bars would cover the category labels
            chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    except Exception:
        pass  # radar / scatter axes differ; defaults are fine


# --------------------------------------------------------------------------- image


def resolve_image(rc: RenderCtx, src: str) -> Path | None:
    if "://" in src or src.startswith("data:"):
        return None
    p = Path(src).expanduser()
    if not p.is_absolute():
        p = Path(rc.base_dir) / p
    return p if p.is_file() else None


def _style_picture(rc: RenderCtx, pic, st: Style, w: int, h: int) -> None:
    """CSS on a picture: border (``line``), ``border-radius``, ``box-shadow`` and ``opacity``."""
    spPr = pic._element.spPr
    if st.radius:
        geom = spPr.find(qn("a:prstGeom"))
        if geom is not None:
            geom.set("prst", "roundRect")
            av = geom.find(qn("a:avLst"))
            if av is None:
                av = etree.SubElement(geom, qn("a:avLst"))
            gd = etree.SubElement(av, qn("a:gd"))
            gd.set("name", "adj")
            gd.set("fmla", f"val {round(min(emu(st.radius) / max(min(w, h), 1), 0.5) * 100000)}")
    if st.line and (st.line_width is None or st.line_width > 0):
        pic.line.color.rgb = rgb(rc.theme, st.line)
        pic.line.width = Pt(st.line_width if st.line_width is not None else rc.theme.render.line_width)
        if st.line_dash in ("dash", "dot"):
            pic.line.dash_style = MSO_LINE.DASH if st.line_dash == "dash" else MSO_LINE.ROUND_DOT
    if st.shadow:
        apply_shadow(rc, spPr, st, w, h)
    if st.opacity is not None and 0 <= st.opacity < 1:
        blip = pic._element.find(".//" + qn("a:blip"))
        if blip is not None:
            etree.SubElement(blip, qn("a:alphaModFix")).set("amt", str(round(st.opacity * 100000)))


def add_image(rc: RenderCtx, slide, pl: Placed, name: str) -> bool:
    """Add the picture; returns False when the image cannot be used (caller draws a placeholder)."""
    im: Image = pl.element  # type: ignore[assignment]
    from .svg import add_svg, is_svg_src

    if is_svg_src(im.src):
        return add_svg(rc, slide, pl, name)
    path = resolve_image(rc, im.src)
    if path is None:
        hint = (
            "download it and use a local path"
            if "://" in im.src
            else "check the path, relative to the .md file"
        )
        rc.diag("image-missing", f"image not found: {im.src}", hint, line=im.line)
        return False
    try:
        from PIL import Image as PILImage

        with PILImage.open(path) as pil:
            iw, ih = pil.size
        x, y, w, h = pl.x, pl.y, pl.w, pl.h
        crop = None
        if im.fit == "contain" and iw and ih:
            s = min(w / iw, h / ih)
            dw, dh = round(iw * s), round(ih * s)
            top = pl.style.valign == "top"  # beside text: tops line up
            x, y, w, h = x + (w - dw) // 2, y + (0 if top else (h - dh) // 2), dw, dh
        elif im.fit == "cover" and iw and ih:
            ia, ba = iw / ih, w / h
            crop = ((1 - ba / ia) / 2, 0.0) if ia > ba else (0.0, (1 - ia / ba) / 2)
        pic = slide.shapes.add_picture(str(path), Emu(x), Emu(y), Emu(w), Emu(h))
        if crop:
            pic.crop_left = pic.crop_right = crop[0]
            pic.crop_top = pic.crop_bottom = crop[1]
        pic.name = name
        pic._element.nvPicPr.cNvPr.set("descr", im.alt or "")
        _style_picture(rc, pic, pl.style, w, h)
        return True
    except Exception as e:
        rc.diag(
            "image-unreadable",
            f"cannot read image {im.src}: {type(e).__name__}",
            "use PNG, JPEG or GIF",
            line=im.line,
        )
        return False


# --------------------------------------------------------------------------- code


def _token_color(ttype) -> tuple[str | None, bool]:  # default style: theme color names
    if ttype in Comment:
        return "muted", True
    if ttype in Keyword:
        return "primary", False
    if ttype in String:
        return "success", False
    if ttype in Number:
        return "accent", False
    if ttype in Name.Function or ttype in Name.Class or ttype in Name.Builtin or ttype in Name.Decorator:
        return "secondary", False
    if ttype in Operator:
        return "fg", False
    return None, False


def _pygments_colors(name: str):
    """``ttype -> (color, italic)`` from a pygments style (``render.code_style``), or None if unknown."""
    try:
        from pygments.styles import get_style_by_name

        style = get_style_by_name(name)
    except Exception:
        return None

    def color(ttype) -> tuple[str | None, bool]:
        st = style.style_for_token(ttype)
        return (st["color"].upper() if st["color"] else None), bool(st["italic"])

    return color


def code_paragraphs(c: Code, style_name: str = "default", diag=None) -> list[Paragraph]:
    """Syntax-highlighted lines. ``style_name`` "default" colors tokens with theme colors; any other value
    is a pygments style name (``render.code_style``); an unknown one falls back with a diagnostic."""
    pick = _token_color
    if style_name and style_name != "default":
        pick = _pygments_colors(style_name)
        if pick is None:
            if diag:
                diag(
                    "bad-token",
                    f"unknown code style {style_name!r}",
                    "use a pygments style such as monokai, friendly or default",
                )
            pick = _token_color
    try:
        lexer = get_lexer_by_name(c.lang) if c.lang else TextLexer()
    except Exception:
        lexer = TextLexer()
    text = c.text.expandtabs(4).rstrip("\n")
    lines: list[list[Run]] = [[]]
    for ttype, value in lex(text, lexer):
        color, italic = pick(ttype)
        parts = value.split("\n")
        for i, part in enumerate(parts):
            if i > 0:
                lines.append([])
            if part:
                lines[-1].append(Run(text=part, code=True, color=color, italic=italic))
    if len(lines) > 1 and not lines[-1]:
        lines.pop()
    return [Paragraph(runs=runs) for runs in lines]
