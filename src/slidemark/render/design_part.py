"""The deck's design source (theme, tokens, css rules, html slides) stored inside the .pptx.

Stored as a ``customXml`` item (``/customXml/itemN.xml`` + ``itemPropsN.xml``), which PowerPoint keeps on save
and does not flag for repair. The item is ``<sm:design xmlns:sm="urn:slidemark:design">`` whose text is JSON:

    {"v": 1, "theme": "none", "tokens": {"colors.primary": "#7C5CFF"},
     "css": [["h1, h2", "color: #fff; font-size: 54pt"]],
     "slides": [{"id": 256, "css": [...], "html": "<div>..</div>", "title": "Launch", "cls": ["dark"],
                 "dec": ["noemph"], "el": [["Signal", "hero", "#intro"]],
                 "fences": [{"src": "<div>..</div>", "info": "{render=native}", "first": 1}]}]}

``el`` lists the ``##`` boxes of a slide that carry a class or id the design refers to (a token class, a css
selector): ``[heading, name, ...]`` in document order, ``name`` being ``hero`` for ``.hero`` or ``#intro``.

Slides are keyed by their ``p:sldId`` id, so reordering or deleting slides in PowerPoint keeps the match.
The importer reads it back (``slidemark import``); nothing here ever raises.
"""

from __future__ import annotations

import json
import re
import uuid
from typing import Any
from xml.sax.saxutils import escape, unescape

from ..ir import CssRule, Deck, Diagnostic, Style

DECISION_CLASSES = ("noemph", "defaults")  # slide directives that only record an author's decision

__all__ = ["style_to_css", "write_design_part", "read_design_part", "design_payload"]

NS = "urn:slidemark:design"
RT_CUSTOM_XML = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
RT_CUSTOM_PROPS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps"
CT_ITEM = "application/xml"
CT_PROPS = "application/vnd.openxmlformats-officedocument.customXmlProperties+xml"
VERSION = 1


# --------------------------------------------------------------------------- Style -> CSS


def _n(x: float | int | str) -> str:
    if isinstance(x, str):
        return x
    s = f"{float(x):.3f}".rstrip("0").rstrip(".")
    return "0" if s in ("", "-0") else s


def _len(x: float | int | str) -> str:
    return x if isinstance(x, str) else f"{_n(x)}pt"


_TRANSFORM = {"upper": "uppercase", "lower": "lowercase", "capitalize": "capitalize", "none": "none"}
_DASH = {"solid": "solid", "dash": "dashed", "dot": "dotted"}


def _grid(spec: str) -> tuple[str, str]:
    if re.search(r"[a-z.]", spec):
        rows = " ".join('"' + " ".join(r) + '"' for r in spec.split("/"))
        return "grid-template-areas", rows
    if ":" in spec or not spec.isdigit():
        return "grid-template-columns", " ".join(f"{p}fr" for p in spec.split(":"))
    return "grid-template-columns", f"repeat({spec}, 1fr)"


def style_to_css(style: Style) -> str:
    """Canonical CSS declarations (``a: b; c: d``) that parse back to ``style``."""
    s = style
    d: list[tuple[str, str]] = []
    if s.color is not None:
        d.append(("color", s.color))
    if s.fill is not None:
        d.append(("background", s.fill))
    if s.line is not None:
        d.append(("border-color", s.line))
    if s.line_width is not None:
        d.append(("border-width", _len(s.line_width)))
    if s.line_dash is not None:
        d.append(("border-style", _DASH[s.line_dash]))
    for side in ("top", "right", "bottom", "left"):
        v = getattr(s, f"border_{side}")
        if v is not None:  # canonical "1pt dash #hex" -> css keywords (dashed, dotted)
            d.append((f"border-{side}", re.sub(r"\b(dash|dot)\b", lambda m: _DASH[m.group(1)], v)))
    if s.radius is not None:
        d.append(("border-radius", _len(s.radius)))
    if s.shadow is not None and s.shadow is not True:
        if s.shadow is False:
            d.append(("box-shadow", "none"))
        else:
            parts = s.shadow.split()
            nums = [p if re.search(r"[a-z%]", p) else f"{p}pt" for p in parts[:-1]]
            d.append(("box-shadow", " ".join([*nums, parts[-1]]) if parts else "none"))
    if s.opacity is not None:
        d.append(("opacity", _n(s.opacity)))
    if s.font is not None:
        d.append(("font-family", '"' + s.font.replace('"', "") + '"'))
    if s.font_size is not None:
        d.append(("font-size", _len(s.font_size)))
    if s.bold is not None:
        d.append(("font-weight", "bold" if s.bold else "normal"))
    if s.italic is not None:
        d.append(("font-style", "italic" if s.italic else "normal"))
    if s.letter_spacing is not None:
        d.append(("letter-spacing", _len(s.letter_spacing)))
    if s.line_spacing is not None:
        d.append(("line-height", _n(s.line_spacing)))
    if s.align is not None:
        d.append(("text-align", s.align))
    if s.valign is not None:
        d.append(("vertical-align", s.valign))
    if s.text_transform is not None:
        d.append(("text-transform", _TRANSFORM[s.text_transform]))
    if s.underline is not None or s.strike is not None:
        deco = [w for w, on in (("underline", s.underline), ("line-through", s.strike)) if on]
        d.append(("text-decoration", " ".join(deco) or "none"))
    sides = {k: getattr(s, f"padding_{k}") for k in ("top", "right", "bottom", "left")}
    if s.padding is not None and all(v == s.padding for v in sides.values()):
        d.append(("padding", _len(s.padding)))
    else:
        if s.padding is not None:
            d.append(("padding", _len(s.padding)))
        for k, v in sides.items():
            if v is not None and v != s.padding:
                d.append((f"padding-{k}", _len(v)))
    if s.margin is not None:
        d.append(("margin", _len(s.margin)))
    if s.gap is not None:
        d.append(("gap", _len(s.gap)))
    if s.grid is not None:
        d.append(_grid(s.grid))
    if s.rotation is not None:
        d.append(("transform", f"rotate({_n(s.rotation)}deg)"))
    if s.width is not None:
        d.append(("width", s.width))
    return "; ".join(f"{k}: {v}" for k, v in d)


def rules_payload(rules: list[CssRule]) -> list[list[str]]:
    """``[[selectors, declarations], ...]`` in source order; neighbours with equal declarations merge."""
    out: list[list[str]] = []
    for r in rules:
        decl = style_to_css(r.style)
        if not decl:
            continue
        if out and out[-1][1] == decl:
            out[-1][0] += ", " + r.selector
        else:
            out.append([r.selector, decl])
    return out


# --------------------------------------------------------------------------- payload


def _referenced(deck: Deck) -> set[str]:
    """Class names (``hero``) and ids (``#intro``) that token paths or css selectors of the deck mention."""
    names: set[str] = set()
    for path in deck.tokens or {}:
        parts = path.split(".")
        if len(parts) >= 3 and parts[0] == "classes":
            names.add(parts[1])
    for sl in [deck, *deck.slides]:
        for r in sl.css or []:
            names.update(re.findall(r"\.([A-Za-z_][\w-]*)", r.selector))
            names.update("#" + m for m in re.findall(r"#([A-Za-z_][\w-]*)", r.selector))
    return names


def _boxes(elements: list[Any], names: set[str], out: list[list[str]], every: bool = False) -> None:
    """Append ``[heading, name, ...]`` for each container whose classes or id are in ``names``.

    ``every`` also records the boxes without such a name (``[heading]``): their heading text is stored
    because a CSS ``text-transform`` changed its casing on the slide.
    """
    for el in elements:
        if getattr(el, "type", None) != "container":
            continue
        mine = [c for c in el.classes if c in names]
        if el.id and "#" + el.id in names:
            mine.append("#" + el.id)
        if (mine or every) and el.title is not None:
            out.append(["".join(p.plain for p in el.title.paragraphs), *mine])
        _boxes(el.children, names, out, every)


def _texts(elements: list[Any], names: set[str], out: list[list[str]]) -> None:
    """Append ``[first paragraph, name, ...]`` for each text block that carries a class in ``names``."""
    for el in elements:
        t = getattr(el, "type", None)
        if t == "container":
            _texts(el.children, names, out)
        elif t == "text" and el.paragraphs and el.role == "body":
            mine = [c for c in el.classes if c in names]
            if el.id and "#" + el.id in names:
                mine.append("#" + el.id)
            if mine:
                out.append([el.paragraphs[0].plain, *mine])


def _transforms(deck: Deck) -> bool:
    """True when some css rule or token of the deck sets ``text-transform`` (a literal re-casing of text)."""
    for sl in [deck, *deck.slides]:
        if any(r.style.text_transform not in (None, "none") for r in sl.css or []):
            return True
    return any("text-transform" in k or "text_transform" in k for k in deck.tokens or {})


def _fences(elements: list[Any], out: list[dict[str, Any]], top: bool = True) -> None:
    """Append ``{src, info, first}`` for each html fence kept as Raw (``first``: nothing before it)."""
    for k, el in enumerate(elements):
        t = getattr(el, "type", None)
        if t == "raw" and getattr(el, "kind", "") == "html" and el.attrs.get("html_src"):
            ent: dict[str, Any] = {"src": el.attrs["html_src"]}
            if el.attrs.get("html_info"):
                ent["info"] = el.attrs["html_info"]
            if top and k == 0:
                ent["first"] = 1
            out.append(ent)
        elif t == "container":
            _fences(el.children, out, False)


def design_payload(deck: Deck, slide_ids: list[int | None]) -> dict[str, Any]:
    """The JSON-able design source of ``deck``; ``{}`` when there is nothing worth storing."""
    data: dict[str, Any] = {}
    if deck.theme and deck.theme != "default":
        data["theme"] = deck.theme
    if deck.tokens:
        data["tokens"] = dict(deck.tokens)
    if deck.css:
        data["css"] = rules_payload(deck.css)
    slides: list[dict[str, Any]] = []
    names = _referenced(deck)
    recase = _transforms(deck)
    for i, sl in enumerate(deck.slides):
        ent: dict[str, Any] = {}
        if sl.css:
            ent["css"] = rules_payload(sl.css)
        if names or recase:
            els: list[list[str]] = []
            _boxes(sl.elements, names, els, recase)
            if els:
                ent["el"] = els
        if names:
            txs: list[list[str]] = []
            _texts(sl.elements, names, txs)
            if txs:
                ent["tx"] = txs
        if dec := [c for c in DECISION_CLASSES if c in sl.classes]:
            ent["dec"] = dec  # `@noemph` / `@defaults`: no visual trace, kept for the importer
        if recase and sl.title is not None and sl.html is None:
            ent["ttl"] = "".join(p.plain for p in sl.title.paragraphs)
        fences: list[dict[str, Any]] = []
        _fences(sl.elements, fences)
        if fences:
            ent["fences"] = fences
        if sl.html is not None:
            ent["html"] = sl.html
            if sl.title is not None:
                ent["title"] = "".join(p.plain for p in sl.title.paragraphs)
            if sl.classes:
                ent["cls"] = list(sl.classes)
        if ent:
            ent["i"] = i
            sid = slide_ids[i] if i < len(slide_ids) else None
            if sid is not None:
                ent["id"] = sid
            slides.append(ent)
    if slides:
        data["slides"] = slides
    if data:
        data = {"v": VERSION, **data}
    return data


# --------------------------------------------------------------------------- package part


def write_design_part(prs, deck: Deck, slides: list[Any]) -> None:
    """Add the design item to ``prs`` (a python-pptx Presentation). Never raises; omitted when empty."""
    try:
        ids = [getattr(s, "slide_id", None) for s in slides]
        data = design_payload(deck, ids)
        if not data:
            return
        from pptx.opc.package import Part
        from pptx.opc.packuri import PackURI

        text = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
        xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<sm:design xmlns:sm="{NS}">{escape(text)}</sm:design>'
        )
        pkg = prs.part.package
        item_uri = pkg.next_partname("/customXml/item%d.xml")
        n = re.search(r"(\d+)\.xml$", str(item_uri))
        props_uri = PackURI(f"/customXml/itemProps{n.group(1) if n else 99}.xml")
        guid = "{" + str(uuid.uuid4()).upper() + "}"
        props = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<ds:datastoreItem xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml" '
            f'ds:itemID="{guid}"><ds:schemaRefs><ds:schemaRef ds:uri="{NS}"/></ds:schemaRefs>'
            "</ds:datastoreItem>"
        )
        item = Part(item_uri, CT_ITEM, pkg, xml.encode("utf-8"))
        pprops = Part(props_uri, CT_PROPS, pkg, props.encode("utf-8"))
        item.relate_to(pprops, RT_CUSTOM_PROPS)
        prs.part.relate_to(item, RT_CUSTOM_XML)
    except Exception as e:  # never raise: a deck without its design part is still a valid deck
        deck.diagnostics.append(
            Diagnostic(
                level="warning",
                message=f"design part not written: {type(e).__name__}",
                rule="design-part",
                hint="the .pptx is fine; `slidemark import` will not restore tokens, css and html slides",
            )
        )


def read_design_part(prs) -> dict[str, Any] | None:
    """The stored design payload of ``prs`` or ``None`` (absent, foreign or unreadable). Never raises."""
    try:
        for rel in prs.part.rels.values():
            if rel.is_external or rel.reltype != RT_CUSTOM_XML:
                continue
            blob = rel.target_part.blob
            m = re.search(rb"<sm:design[^>]*>(.*)</sm:design>", blob, re.S)
            if not m or NS.encode() not in blob[:300]:
                continue
            data = json.loads(unescape(m.group(1).decode("utf-8")))
            if isinstance(data, dict) and data.get("v") == VERSION:
                return data
    except Exception:  # noqa: S110 - best effort
        return None
    return None
