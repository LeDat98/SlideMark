"""Paragraphs and runs -> DrawingML text, with native bullets, fonts (latin + ea) and hyperlinks."""

from __future__ import annotations

import re
import uuid

from lxml import etree
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from ..ir import Paragraph, Run, Style
from ..layout import jbreak, measure
from ..layout.css import insets, transform_text
from ..units import EMU_PER_PT
from .util import RenderCtx, hex6, rgb

_RPR_ORDER = [
    "ln",
    "noFill",
    "solidFill",
    "gradFill",
    "blipFill",
    "pattFill",
    "grpFill",
    "effectLst",
    "effectDag",
    "highlight",
    "uLnTx",
    "uLn",
    "uFillTx",
    "uFill",
    "latin",
    "ea",
    "cs",
    "sym",
    "hlinkClick",
    "hlinkMouseOver",
    "rtl",
    "extLst",
]
_ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
_LANG = {
    "ja": "ja-JP",
    "vi": "vi-VN",
    "en": "en-US",
    "zh": "zh-CN",
    "ko": "ko-KR",
    "fr": "fr-FR",
    "de": "de-DE",
}


_RANGE = re.compile(r"(?<=\d)([\u2013\u2014\-/~\uff5e])(?=\d)")


def _join_ranges(text: str) -> str:
    """Keep number ranges (``2027–2029``, ``10-12``, ``1/2027``) on one line: viewers break after a dash or
    slash, so a WORD JOINER (U+2060, invisible) sits on both sides when the mark is between two digits."""
    return _RANGE.sub("\u2060\\1\u2060", text) if any(c.isdigit() for c in text) else text


def _cjk_lang(lang: str | None) -> bool:
    return (lang or "").lower().split("-")[0] in ("ja", "zh", "ko")


_BADGE_PAD = "\u3000"  # full-width space: renders flush inside the highlight (Latin spaces leave a gap)
_BULLETS = ["•", "–", "•", "–"]
_NUMBERING = ["arabicPeriod", "alphaLcParenR", "romanLcPeriod", "arabicPeriod"]


def insert_rpr_child(rpr, child) -> None:
    """Insert ``child`` into an rPr honoring the schema sequence."""
    name = etree.QName(child).localname
    idx = _RPR_ORDER.index(name)
    for existing in list(rpr):
        en = etree.QName(existing).localname
        if en in _RPR_ORDER and _RPR_ORDER.index(en) > idx:
            existing.addprevious(child)
            return
    rpr.append(child)


def _lang_for(text: str, deck_lang: str | None) -> str:
    if deck_lang:
        return _LANG.get(deck_lang.lower().split("-")[0], deck_lang)
    for ch in text:
        o = ord(ch)
        if 0x3040 <= o <= 0x30FF:
            return "ja-JP"
        if 0xAC00 <= o <= 0xD7AF:
            return "ko-KR"
        if 0x4E00 <= o <= 0x9FFF:
            return "ja-JP"
    return "en-US"


def _format_run(
    rc: RenderCtx,
    r,
    run: Run,
    style: Style,
    size_pt: float,
    text: str,
    scale: float = 1.0,
    squeeze: float = 0.0,
) -> None:
    theme = rc.theme
    f = r.font
    f.size = Pt(size_pt)
    bold = bool(run.bold or style.bold or run.highlight)
    if run.highlight and not style.bold and not theme.render.badge_cjk_bold and measure.has_cjk(text):
        bold = False  # synthetic bold smears small CJK glyphs inside the highlight
    italic = bool(run.italic or style.italic)
    f.bold = True if bold else None
    f.italic = True if italic else None
    if run.underline or style.underline:
        f.underline = True
    color = theme.run_color(run.color, run.highlight, style.color, size_pt, style.fill)  # shared with lint
    f.color.rgb = rgb(theme, color, "fg")
    rpr = r._r.get_or_add_rPr()
    if style.opacity is not None and 0 <= style.opacity < 1 and not style.fill:
        clr = rpr.find(qn("a:solidFill")).find(qn("a:srgbClr"))  # CSS opacity on text without a fill
        etree.SubElement(clr, qn("a:alpha")).set("val", str(round(style.opacity * 100000)))
    if run.strike or style.strike:
        rpr.set("strike", "sngStrike")
    if style.letter_spacing:
        rpr.set("spc", str(round(style.letter_spacing * scale * 100)))  # 1/100 pt
    elif squeeze:
        rpr.set("spc", str(round(squeeze * 100)))  # CJK orphan squeeze (measure.paragraph_squeeze)
    if run.sup:
        rpr.set("baseline", "30000")
    elif run.sub:
        rpr.set("baseline", "-25000")
    lang = _lang_for(text, rc.deck.lang)
    rpr.set("lang", lang)
    rpr.set("altLang", "en-US")
    if run.highlight:
        hl = etree.SubElement(rpr, qn("a:highlight"))
        etree.SubElement(hl, qn("a:srgbClr")).set("val", hex6(theme, run.highlight, theme.render.highlight))
        rpr.remove(hl)
        insert_rpr_child(rpr, hl)
    latin = theme.fonts.mono if run.code else (style.font or theme.fonts.body)
    ea = style.font_ea or theme.fonts.ea
    for tag, face in (("a:latin", latin), ("a:ea", ea)):
        el = etree.Element(qn(tag))
        el.set("typeface", face)
        insert_rpr_child(rpr, el)
    if run.link:
        if run.link.startswith("#"):
            rc.links.append((r._r, rpr, run.link[1:], rc.slide_index))
        else:
            r.hyperlink.address = run.link


def _soft_break(para, size_pt: float) -> None:
    para.add_line_break()
    br = para._p[-1]
    rpr = etree.SubElement(br, qn("a:rPr"))
    rpr.set("sz", str(round(size_pt * 100)))
    rpr.set("bmk", jbreak.SOFT_BREAK_MARK)


def _set_bullet(para, p: Paragraph, size_pt: float) -> None:
    pPr = para._p.get_or_add_pPr()
    pPr.set("eaLnBrk", "1")  # kinsoku line breaking, no hanging punctuation: matches layout.measure
    pPr.set("hangingPunct", "0")
    if p.marker:
        marL, indent = measure.list_indent(size_pt, p.level)
        pPr.set("marL", str(marL))
        pPr.set("indent", str(indent))
        lvl = min(p.level, 3)
        if p.marker == "bullet":
            bf = etree.SubElement(pPr, qn("a:buFont"))
            bf.set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", _BULLETS[lvl])
        else:
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "+mj-lt")
            etree.SubElement(pPr, qn("a:buAutoNum")).set("type", _NUMBERING[lvl])
    else:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))


def fill_text(
    rc: RenderCtx,
    tf,
    paragraphs: list[Paragraph],
    style: Style,
    scale: float = 1.0,
    *,
    para_gap: bool = True,
    field: str | None = None,
    inset: int | None = None,
    gap_em: float | None = None,
    box_w: int | None = None,
    squeeze: bool = True,
) -> None:
    """Write ``paragraphs`` into text frame ``tf`` using the (already merged) ``style``.

    ``gap_em`` is the space before every paragraph but the first, x font size (default ``measure.para_gap()``;
    the layout passes a larger value for roomy cards, as ``attrs["para_gap"]`` of the text).
    ``box_w`` (EMU, the text width before the frame margins) turns on the CJK orphan squeeze and the phrase
    breaks, which need the line width; ``squeeze=False`` keeps only the phrase breaks (chevrons).
    """
    gap = measure.para_gap() if gap_em is None else gap_em
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    if inset is not None:
        pl = pt = pr = pb = inset
    else:  # padding + per-side padding_* + per-side border widths (CSS box model)
        pl, pt, pr, pb = insets(style)
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = (Emu(pl), Emu(pt), Emu(pr), Emu(pb))
    tf.vertical_anchor = _ANCHOR.get(style.valign or "top", MSO_ANCHOR.TOP)
    paras = paragraphs or [Paragraph()]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pst = style.merged(p.style)
        size = (pst.font_size or 18) * scale
        para.alignment = _ALIGN.get(pst.align or "left", PP_ALIGN.LEFT)
        if pst.line_spacing:
            para.line_spacing = pst.line_spacing
        para.space_before = Pt(size * gap) if (i > 0 and para_gap) else Pt(0)
        para.space_after = Pt(0)
        _set_bullet(para, p, size)
        sq = 0.0
        cuts: dict[int, list[int]] = {}
        if box_w is not None and not field:
            wpt = (box_w - pl - pr - (measure.list_indent(size, p.level)[0] if p.marker else 0)) / EMU_PER_PT
            sq = measure.paragraph_squeeze(p, pst, wpt, size) if squeeze else 0.0
            if plan := jbreak.plan_breaks(p, pst, wpt, size, sq):
                cuts, sq = plan
        first_text = ""
        prev = ""
        prev_r = None
        bound = [r.text for r in p.runs] if field else measure.bound_texts(p.runs)
        if box_w is not None and not field:  # a binding never forces a mid-word break
            bound = measure.loosen_wide_bindings([r.text for r in p.runs], bound, wpt, size, pst.font)
        for ri, (run, run_text) in enumerate(zip(p.runs, bound, strict=True)):
            segs = run_text.replace("\r", "").replace("\v", "\n").split("\n")
            for k, seg in enumerate(segs):
                if k > 0:
                    para.add_line_break()
                    prev = ""
                    prev_r = None
                if seg == "" and len(segs) > 1:
                    continue
                seg = transform_text(seg, pst.text_transform)
                # badge: padding keeps bold CJK glyphs inside the highlight (LibreOffice clips them otherwise)
                pad = ""
                if run.highlight:
                    tk = measure.tokens()
                    pad = _BADGE_PAD * tk.badge_pad_cjk if measure.has_cjk(seg) else "\u00a0" * tk.badge_pad
                    if tk.badge_gap and prev_r is not None and not prev.endswith((" ", "\u00a0", "\u3000")):
                        prev_r.text = (
                            prev_r.text + " "
                        )  # a plain space keeps the badge off the text before it
                ins = sorted({o for o in cuts.get(ri, []) if 0 < o < len(seg)}) if len(segs) == 1 else []
                offs = [0, *ins, len(seg)]
                for a, b in zip(offs, offs[1:], strict=False):
                    if a > 0:
                        _soft_break(para, size)  # phrase break (jbreak): the importer drops it
                    r = para.add_run()
                    r.text = f"{pad}{_join_ranges(seg[a:b])}{pad}"
                    _format_run(rc, r, run, pst, size, seg[a:b], scale, sq)
                prev = seg
                prev_r = None if run.highlight else r  # a badge carries its own padding
                first_text = first_text or seg
                if field == "slide_number":
                    fld = r._r
                    fld.tag = qn("a:fld")
                    fld.set("id", "{" + str(uuid.uuid4()).upper() + "}")
                    fld.set("type", "slidenum")
        end = etree.SubElement(para._p, qn("a:endParaRPr"))
        end.set("lang", _lang_for(first_text, rc.deck.lang))
        end.set("sz", str(round(size * 100)))
