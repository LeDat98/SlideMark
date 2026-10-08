"""Vector icons as native custom-geometry shapes (editable, recolorable in PowerPoint)."""

from __future__ import annotations

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Emu

from .. import icons
from ..ir import Placed, Shape, Style
from ..layout import css, measure
from ..units import EMU_PER_PT
from .effects import apply_fill
from .util import RenderCtx, rgb

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_SCALE = 1000  # path units per grid unit (the path box is 24000 x 24000)


def _pt(x: float, y: float) -> str:
    return f'<a:pt x="{round(x * _SCALE)}" y="{round(y * _SCALE)}"/>'


def _cust_geom(layers) -> etree._Element:
    size = icons.GRID * _SCALE
    paths = []
    for layer in layers:
        body = []
        for op, pts in layer:
            if op == "M":
                body.append(f"<a:moveTo>{_pt(*pts[0])}</a:moveTo>")
            elif op == "L":
                body.append(f"<a:lnTo>{_pt(*pts[0])}</a:lnTo>")
            elif op == "C":
                body.append(f"<a:cubicBezTo>{''.join(_pt(*p) for p in pts)}</a:cubicBezTo>")
            elif op == "Q":
                body.append(f"<a:quadBezTo>{''.join(_pt(*p) for p in pts)}</a:quadBezTo>")
            else:
                body.append("<a:close/>")
        paths.append(f'<a:path w="{size}" h="{size}">{"".join(body)}</a:path>')
    xml = (
        f'<a:custGeom xmlns:a="{_A}"><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/>'
        '<a:rect l="0" t="0" r="r" b="b"/>'
        f"<a:pathLst>{''.join(paths)}</a:pathLst></a:custGeom>"
    )
    return etree.fromstring(xml)


def _file_icon(rc: RenderCtx, slide, pl: Placed, name: str, src: str) -> bool:
    """``icon=file.svg``: native geometry for simple filled SVG, else the SVG as a picture."""
    from ..ir import Image
    from .objects import resolve_image
    from .svg import add_svg

    path = resolve_image(rc, src)
    if path is None:
        rc.diag("icon-missing", f"icon file not found: {src}", "check the path, relative to the .md file")
        return False
    try:
        layers = icons.svg_layers(path.read_text(encoding="utf-8-sig", errors="replace"))
    except OSError:
        layers = None
    if layers:
        _draw(rc, slide, pl, f"icon {path.name}", layers)
        return True
    pic = Placed(
        element=Image(src=str(path), alt=path.stem, fit="contain"),
        x=pl.x,
        y=pl.y,
        w=pl.w,
        h=pl.h,
        style=pl.style,
    )
    return add_svg(rc, slide, pic, name)


def _draw(rc: RenderCtx, slide, pl: Placed, name: str, layers) -> None:
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(pl.x), Emu(pl.y), Emu(pl.w), Emu(pl.h))
    shp.name = name
    el = shp._element
    if (style_el := el.find(qn("p:style"))) is not None:
        el.remove(style_el)
    sp_pr = el.spPr
    prst = sp_pr.find(qn("a:prstGeom"))
    prst.addprevious(_cust_geom(layers))
    sp_pr.remove(prst)
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(rc.theme, pl.style.fill or pl.style.color or "primary")
    shp.line.fill.background()
    shp.shadow.inherit = False


def _disc(rc: RenderCtx, slide, pl: Placed):
    """The native ``Icon disc`` (circle / rounded square / square) filling ``pl``; ``None`` if no disc."""
    attrs = pl.element.attrs
    color = attrs.get("disc")
    if not color:
        return None
    shape = str(attrs.get("disc_shape") or rc.theme.icon_disc_shape)
    kind = {
        "circle": MSO_SHAPE.OVAL,
        "rounded": MSO_SHAPE.ROUNDED_RECTANGLE,
        "square": MSO_SHAPE.RECTANGLE,
    }.get(shape)
    if kind is None:
        rc.diag(
            "bad-attr",
            f"unknown disc shape {shape!r}: drew a circle",
            "use icon.disc.shape=circle, rounded or square",
        )
        kind = MSO_SHAPE.OVAL
    shp = slide.shapes.add_shape(kind, Emu(pl.x), Emu(pl.y), Emu(pl.w), Emu(pl.h))
    shp.name = "Icon disc"
    el = shp._element
    if (style_el := el.find(qn("p:style"))) is not None:
        el.remove(style_el)
    if not apply_fill(rc, el.spPr, Style(fill=str(color))):
        shp.fill.background()
    shp.line.fill.background()
    shp.shadow.inherit = False
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            share = float(attrs.get("disc_round", rc.theme.layout.icon_disc_round))
            shp.adjustments[0] = min(max(share, 0.0), 0.5)
        except (IndexError, ValueError):
            pass
    return shp


def add_icon(rc: RenderCtx, slide, pl: Placed, name: str) -> bool:
    """Draw ``pl.element.attrs["icon"]``; return False (nothing drawn) for an unknown icon name or file.

    With ``attrs["disc"]`` the placed rectangle is the disc: a native ``Icon disc`` shape, and the glyph
    centred in it at ``attrs["disc_glyph"]`` of its diameter."""
    raw = str(pl.element.attrs.get("icon", "")).strip()
    file_icon = icons.is_file(raw)
    layers = None if file_icon else icons.commands(raw.lower())
    if not file_icon and not layers:
        return False
    disc = _disc(rc, slide, pl)
    glyph = pl
    if disc is not None:
        ratio = float(pl.element.attrs.get("disc_glyph", rc.theme.layout.icon_disc_glyph))
        g = max(round(min(pl.w, pl.h) * min(max(ratio, 0.05), 1.0)), 1)
        glyph = pl.model_copy(
            update={"x": pl.x + (pl.w - g) // 2, "y": pl.y + (pl.h - g) // 2, "w": g, "h": g}
        )
    if file_icon:
        ok = _file_icon(rc, slide, glyph, name, raw)
    else:
        _draw(rc, slide, glyph, f"icon {raw.lower()}", layers)
        ok = True
    if not ok and disc is not None:
        disc._element.getparent().remove(disc._element)  # no glyph: no empty disc either
    return ok


def add_chevron_icon(rc: RenderCtx, slide, pl: Placed, st: Style) -> None:
    """The icon of a chevron / ``@steps`` arrow (``attrs["icon"]``), drawn from the arrow's FINAL rectangle.

    It sits just before the centred heading text, vertically centred; the layout only reserved the room
    (``icon_side`` the mark, ``icon_inset`` the mark plus its gap). A pass that moved or resized the arrow
    therefore never leaves the icon behind."""
    el = pl.element
    side = int(el.attrs.get("icon_side") or 0)
    if not side or not el.attrs.get("icon"):
        return
    adj = float(el.attrs.get("adj", rc.theme.layout.chevron_adj))
    depth = round(adj * min(pl.w, pl.h))
    pad = css.insets(st, 0)[0]
    inset = int(el.attrs.get("icon_inset") or 0)
    left = pl.x + depth + pad
    avail = max(pl.w - 2 * depth - 2 * pad - inset, 1)
    size = (st.font_size or 18) * pl.font_scale
    line = max((measure.text_em(p.plain, bold=True) * size * EMU_PER_PT for p in el.paragraphs), default=0)
    off = max(round((avail - min(line, avail)) / 2), 0)
    attrs: dict = {"icon": el.attrs["icon"]}
    if el.attrs.get("icon_disc"):
        attrs["disc"] = el.attrs["icon_disc"]
    ink = el.attrs.get("icon_ink") or st.color or "primary"
    mark = Placed(
        element=Shape(shape="icon", attrs=attrs),
        x=left + off,
        y=pl.y + (pl.h - side) // 2,
        w=side,
        h=side,
        style=Style(fill=str(ink)),
    )
    add_icon(rc, slide, mark, f"icon {el.attrs['icon']}")


def add_bar_icon(rc: RenderCtx, slide, pl: Placed) -> None:
    """The icon of the conclusion bar (``conclusion.icon``), vertically centred at the bar's left end.

    The layout stored the mark's size and its distance from the left edge on the text; it is drawn from the
    bar's final rectangle, so a bar a pass moved or stretched keeps its icon."""
    el = pl.element
    side = int(el.attrs.get("icon_side") or 0)
    if not side or not el.attrs.get("icon"):
        return
    attrs: dict = {"icon": el.attrs["icon"]}
    for key in ("disc_shape", "disc_glyph", "disc_round"):
        if key in el.attrs:
            attrs[key] = el.attrs[key]
    if el.attrs.get("icon_disc"):
        attrs["disc"] = el.attrs["icon_disc"]
    mark = Placed(
        element=Shape(shape="icon", attrs=attrs),
        x=pl.x + int(el.attrs.get("icon_dx") or 0),
        y=pl.y + (pl.h - side) // 2,
        w=side,
        h=side,
        style=Style(fill=str(el.attrs.get("icon_ink") or pl.style.color or "bg")),
    )
    add_icon(rc, slide, mark, f"icon {el.attrs['icon']}")
