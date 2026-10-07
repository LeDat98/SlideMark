"""Native shapes of the composition vocabulary (DL3b) that python-pptx has no autoshape for.

- ``Shape(shape="arc")``: an ``arc`` preset (a line along the ellipse inscribed in the box) from
  ``attrs["start"]`` to ``attrs["end"]`` degrees (clockwise from 3 o'clock), with an arrowhead at the end
  (``attrs["head"] == "arrow"``) or at the start (``"back"``). The line comes from the style (``line`` /
  ``line_width``); the arc has no fill.
- ``Shape(shape="trap")``: a trapezoid with exact insets as custom geometry: ``attrs["top"]`` and
  ``attrs["bot"]`` are the left and right insets (shares of the width) of the top and bottom edge, so a
  funnel or pyramid stage can taper by any amount. Its text rectangle is ``attrs["text"]`` of the width
  (default: the mean width).
"""

from __future__ import annotations

from collections.abc import Callable

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from ..ir import Placed
from .text import fill_text
from .util import RenderCtx

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_ANGLE = 60000  # OOXML angle unit: 1/60000 degree


def _angle(deg: float) -> int:
    return round((deg % 360) * _ANGLE)


def draw_arc(rc: RenderCtx, slide, pl: Placed, name: str, autoshape: Callable) -> None:
    el = pl.element
    shp = autoshape(rc, slide, pl, MSO_SHAPE.ARC, name)
    geom = shp._element.spPr.find(qn("a:prstGeom"))
    av = geom.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(geom, qn("a:avLst"))
    for child in list(av):
        av.remove(child)
    for gd, deg in (("adj1", el.attrs.get("start", 0)), ("adj2", el.attrs.get("end", 90))):
        e = etree.SubElement(av, qn("a:gd"))
        e.set("name", gd)
        e.set("fmla", f"val {_angle(float(deg))}")
    head = el.attrs.get("head")
    if head in ("arrow", "back"):
        ln = shp.line._get_or_add_ln()
        end = etree.SubElement(ln, qn("a:tailEnd" if head == "arrow" else "a:headEnd"))
        end.set("type", "triangle")
        end.set("w", "med")
        end.set("len", "med")


def draw_trapezoid(rc: RenderCtx, slide, pl: Placed, name: str, autoshape: Callable) -> None:
    el = pl.element
    shp = autoshape(rc, slide, pl, MSO_SHAPE.RECTANGLE, name)
    w, h = max(pl.w, 1), max(pl.h, 1)
    top = min(max(float(el.attrs.get("top", 0.0)), 0.0), 0.5)
    bot = min(max(float(el.attrs.get("bot", 0.0)), 0.0), 0.5)
    pts = [(round(top * w), 0), (round((1 - top) * w), 0), (round((1 - bot) * w), h), (round(bot * w), h)]
    text = el.attrs.get("text")  # share of the width the text may use (centred); default: the mean width
    inset = round(((1 - float(text)) / 2 if text is not None else (top + bot) / 2) * 100000)
    path = "".join(
        f'<a:{op}><a:pt x="{x}" y="{y}"/></a:{op}>'
        for op, (x, y) in zip(("moveTo", "lnTo", "lnTo", "lnTo"), pts, strict=True)
    )
    xml = (
        f'<a:custGeom xmlns:a="{_A}"><a:avLst/>'
        f'<a:gdLst><a:gd name="il" fmla="*/ w {inset} 100000"/><a:gd name="ir" fmla="+- r 0 il"/></a:gdLst>'
        '<a:ahLst/><a:cxnLst/><a:rect l="il" t="t" r="ir" b="b"/>'
        f'<a:pathLst><a:path w="{w}" h="{h}">{path}<a:close/></a:path></a:pathLst></a:custGeom>'
    )
    sp_pr = shp._element.spPr
    prst = sp_pr.find(qn("a:prstGeom"))
    prst.addprevious(etree.fromstring(xml))
    sp_pr.remove(prst)
    if el.paragraphs:
        fill_text(rc, shp.text_frame, el.paragraphs, pl.style, pl.font_scale, box_w=None, squeeze=False)
