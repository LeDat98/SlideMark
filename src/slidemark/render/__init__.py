"""IR + Placed -> .pptx. Owner: renderer workstream."""

from __future__ import annotations

import re
import uuid
from math import cos, radians, sin
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from ..ir import (
    Chart,
    Code,
    Container,
    Deck,
    Diagnostic,
    Image,
    Media,
    Placed,
    Raw,
    Shape,
    Slide,
    Style,
    Table,
    Text,
)
from ..layout import measure
from ..layout.css import has_side_borders, side_borders, slide_style
from ..template import clone_footer, open_template, pick_layout
from ..theme import DEFAULT_SIZES, Theme
from ..units import slide_size, to_emu
from .anim import build_timing
from .design_part import write_design_part
from .effects import apply_fill, apply_shadow, cover_crop, set_picture
from .htmlimg import add_html_image, add_html_native, close_html
from .icons import add_icon
from .math import add_math
from .media import add_media, finish_timing
from .objects import add_chart, add_image, add_table, code_paragraphs, resolve_image
from .text import fill_text, insert_rpr_child
from .util import RenderCtx, emu, is_gradient, parse_color, rgb

__all__ = ["render"]

_SHAPES = {
    "rect": MSO_SHAPE.RECTANGLE,
    "rectangle": MSO_SHAPE.RECTANGLE,
    "rounded-rect": MSO_SHAPE.ROUNDED_RECTANGLE,
    "rounded-rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_SHAPE.OVAL,
    "oval": MSO_SHAPE.OVAL,
    "circle": MSO_SHAPE.OVAL,
    "arrow-right": MSO_SHAPE.RIGHT_ARROW,
    "arrow-left": MSO_SHAPE.LEFT_ARROW,
    "arrow-up": MSO_SHAPE.UP_ARROW,
    "arrow-down": MSO_SHAPE.DOWN_ARROW,
    "chevron": MSO_SHAPE.CHEVRON,
    "pentagon": MSO_SHAPE.PENTAGON,
    "diamond": MSO_SHAPE.DIAMOND,
    "triangle": MSO_SHAPE.ISOSCELES_TRIANGLE,
    "hexagon": MSO_SHAPE.HEXAGON,
    "star": MSO_SHAPE.STAR_5_POINT,
    "cylinder": MSO_SHAPE.CAN,
    "parallelogram": MSO_SHAPE.PARALLELOGRAM,
    "trapezoid": MSO_SHAPE.TRAPEZOID,
    "cloud": MSO_SHAPE.CLOUD,
    "plus": MSO_SHAPE.MATH_PLUS,
    "callout": MSO_SHAPE.RECTANGULAR_CALLOUT,
}
_TRANSITIONS = {
    "fade": "<p:fade/>",
    "push": '<p:push dir="l"/>',
    "wipe": '<p:wipe dir="l"/>',
    "split": '<p:split orient="horz" dir="out"/>',
    "cover": '<p:cover dir="l"/>',
    "zoom": '<p:zoom dir="in"/>',
    "morph": '<p159:morph option="byObject"/>',  # fallback: fade
}
_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS_MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"
_NS_P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
_NS_P159 = "http://schemas.microsoft.com/office/powerpoint/2015/09/main"
_SECTION_URI = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"


def render(deck: Deck, placed: list[list[Placed]], theme: Theme, out: str | Path) -> Path:
    """Write ``deck`` to ``out``. ``placed[i]`` is the layout result for ``deck.slides[i]``."""
    out = Path(out)
    prs = None
    measure.set_tokens(theme.layout)
    rc = RenderCtx(deck=deck, theme=theme, base_dir=str(deck.attrs.get("base_dir", ".")))
    if theme.template:
        try:
            prs = open_template(theme.template)
            rc.template = True  # the template's own slide size wins
        except Exception as e:
            deck.diagnostics.append(
                Diagnostic(
                    level="warning",
                    message=f"template {theme.template!r} could not be opened: {type(e).__name__}",
                    rule="bad-theme",
                    hint="re-save the template as .pptx; the default design was used",
                )
            )
    if prs is None:
        prs = Presentation()
        try:
            prs.slide_width, prs.slide_height = (Emu(v) for v in slide_size(deck.size))
        except ValueError:
            prs.slide_width, prs.slide_height = (Emu(v) for v in slide_size("16:9"))
    for i, slide in enumerate(deck.slides):
        rc.slide_index = i
        items = placed[i] if i < len(placed) else []
        try:
            rc.slides.append(_render_slide(rc, prs, slide, items))
        except Exception as e:  # never raise on bad input: keep an empty slide and say why
            rc.slides.append(
                prs.slides.add_slide(pick_layout(prs, "blank") if rc.template else prs.slide_layouts[6])
            )
            rc.diag(
                "render-error",
                f"slide failed to render: {type(e).__name__}: {e}",
                "simplify this slide or report a bug",
                "error",
            )
    close_html(rc)
    _resolve_links(rc)
    _sections(rc, prs)
    _core_properties(prs, deck)
    write_design_part(prs, deck, rc.slides)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


# --------------------------------------------------------------------------- slide


def _render_slide(rc: RenderCtx, prs, slide: Slide, items: list[Placed]):
    has_title = any(isinstance(p.element, Text) and p.element.role == "title" for p in items)
    if rc.template:
        kind = _slide_kind(slide, rc.slide_index)
        if kind == "content":
            kind = "title" if has_title else "blank"
        s = prs.slides.add_slide(pick_layout(prs, kind))
    else:
        s = prs.slides.add_slide(prs.slide_layouts[5 if has_title else 6])
    _background(rc, prs, s, slide)
    counters: dict[str, int] = {}
    used_title = False
    shape_ids: list[list[int]] = []  # per Placed index: ids of the shapes it produced (for animations)
    for pl in items:
        seen = {int(i) for i in s.shapes._spTree.xpath("./*/*[1]/p:cNvPr/@id")}
        shape_ids.append([])
        try:
            use_ph = (
                has_title and not used_title and isinstance(pl.element, Text) and pl.element.role == "title"
            )
            used_title = used_title or use_ph
            if rc.template and _native_footer(rc, s, pl):
                continue
            _render_item(rc, s, pl, counters, use_ph)
            now = [int(i) for i in s.shapes._spTree.xpath("./*/*[1]/p:cNvPr/@id")]
            shape_ids[-1] = [i for i in now if i not in seen]
        except Exception as e:
            rc.diag(
                "render-error",
                f"{pl.element.type} failed to render: {type(e).__name__}: {e}",
                "simplify this element",
                "error",
            )
    if has_title and not used_title and s.shapes.title is not None:
        s.shapes.title._element.getparent().remove(s.shapes.title._element)
    if rc.template:  # drop unused template placeholders (empty "Click to add ..." prompts)
        keep = s.shapes.title._element if used_title and s.shapes.title is not None else None
        for shp in list(s.placeholders):
            if shp._element is not keep and not shp._element.xpath(".//p:ph[@type='ftr' or @type='sldNum']"):
                shp._element.getparent().remove(shp._element)
    if slide.notes:
        s.notes_slide.notes_text_frame.text = slide.notes
    if slide.hidden:
        s._element.set("show", "0")
    _transition(rc, s, slide)
    try:
        _timing(s, slide, items, shape_ids)
    except Exception as e:  # animations are optional: keep the slide static
        rc.diag("anim", f"build animation skipped: {type(e).__name__}: {e}", "report a bug", "warning")
    finish_timing(rc, s)
    return s


def _slide_kind(slide: Slide, index: int) -> str:
    """Same inference as the layout engine: cover | section | blank | content."""
    if slide.layout in ("cover", "section", "blank"):
        return slide.layout
    if slide.layout is None and slide.title and not slide.elements and not slide.conclusion:
        return "cover" if index == 0 else "section"
    return "content"


def _native_footer(rc: RenderCtx, s, pl: Placed) -> bool:
    """With a template, footer / slide number become the layout's native placeholders (fields)."""
    el = pl.element
    field = el.attrs.get("field") if isinstance(el, Text) else None
    if field == "footer":
        return clone_footer(s, "ftr", "".join(p.plain for p in el.paragraphs))
    if field == "slide_number":
        return clone_footer(s, "sldNum", str(rc.slide_index + 1))
    return False


def _background(rc: RenderCtx, prs, s, slide: Slide) -> None:
    """Slide background: ``Slide.background``, else CSS ``slide { background }`` (color, gradient, url)."""
    bg = slide.background
    if not bg:
        css_fill = slide_style(rc.deck, slide, rc.slide_index).fill
        bg = css_fill if css_fill and parse_color(rc.theme, css_fill)[1] != 0 else None
    if rc.template and not bg:
        return  # keep the template's own background
    theme = rc.theme
    fill = s.background.fill
    if bg and (m := re.fullmatch(r"\s*url\(\s*(.*?)\s*\)\s*", bg, re.I | re.S)):
        bg = m.group(1).strip("\"'")
    if bg and is_gradient(bg):
        fill.solid()
        tmp = etree.Element(qn("a:spPr"))
        if apply_fill(rc, tmp, Style(fill=bg)):
            bgpr = s._element.find(qn("p:cSld")).find(qn("p:bg")).find(qn("p:bgPr"))
            for old in bgpr.findall(qn("a:solidFill")):
                old.addprevious(tmp[0])
                bgpr.remove(old)
            return
    if bg and (resolve_image(rc, bg) is not None or re.search(r"\.(png|jpe?g|gif|bmp)$", bg, re.I)):
        fill.solid()
        fill.fore_color.rgb = rgb(theme, "bg")
        path = resolve_image(rc, bg)
        if path is None:
            rc.diag(
                "image-missing",
                f"background image not found: {bg}",
                "check the path, relative to the .md file",
            )
            return
        pic = s.shapes.add_picture(str(path), 0, 0, prs.slide_width, prs.slide_height)
        pic.name = "Background"
        tree = s.shapes._spTree
        tree.remove(pic._element)
        tree.insert(2, pic._element)
        return
    fill.solid()
    fill.fore_color.rgb = rgb(theme, bg or "bg", theme.render.slide_bg)


def _transition(rc: RenderCtx, s, slide: Slide) -> None:
    name, _, dur = (slide.transition or "").strip().lower().partition(":")
    if not name or name == "none":
        return
    inner = _TRANSITIONS.get(name)
    if inner is None:
        rc.diag(
            "transition", f"unknown transition {name!r}", f"use one of: {', '.join(_TRANSITIONS)}", "info"
        )
        return
    try:
        secs = min(max(float(dur), 0.1), 10.0) if dur else None
    except ValueError:
        secs = None
    if name == "morph" and secs is None:
        secs = 2.0
    spd = "med" if secs is None else "fast" if secs <= 0.5 else "med" if secs <= 0.75 else "slow"
    if secs is None:
        xml = f'<p:transition xmlns:p="{_NS_P}" spd="{spd}">{inner}</p:transition>'
    else:
        fallback_inner = "<p:fade/>" if name == "morph" else inner
        xml = (
            f'<mc:AlternateContent xmlns:mc="{_NS_MC}" xmlns:p="{_NS_P}" xmlns:p14="{_NS_P14}" '
            f'xmlns:p159="{_NS_P159}"><mc:Choice Requires="{"p159" if name == "morph" else "p14"}">'
            f'<p:transition spd="{spd}" p14:dur="{round(secs * 1000)}">{inner}</p:transition></mc:Choice>'
            f'<mc:Fallback><p:transition spd="{spd}">{fallback_inner}</p:transition></mc:Fallback>'
            "</mc:AlternateContent>"
        )
    sld = s._element
    anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    anchor.addnext(etree.fromstring(xml))


def _timing(s, slide: Slide, items: list[Placed], shape_ids: list[list[int]]) -> None:
    """Add ``p:timing`` after the transition (schema order: cSld, clrMapOvr, transition, timing, extLst)."""
    counts = {shp.shape_id: len(shp.text_frame.paragraphs) for shp in s.shapes if shp.has_text_frame}
    timing = build_timing(items, shape_ids, "build" in slide.classes, counts)
    if timing is None:
        return
    sld = s._element
    anchor = sld.find(qn("p:transition"))
    if anchor is None:
        for child in sld:
            if child.tag == f"{{{_NS_MC}}}AlternateContent":
                anchor = child
    if anchor is None:
        anchor = sld.find(qn("p:clrMapOvr"))
    if anchor is None:
        anchor = sld.find(qn("p:cSld"))
    anchor.addnext(timing)


def _sections(rc: RenderCtx, prs) -> None:
    """``p14:sectionLst``: one PowerPoint section per section slide (default on, header ``sections: off``)."""
    deck = rc.deck
    if str(deck.attrs.get("sections", "on")).lower() == "off":
        return
    kinds = [_slide_kind(sl, i) for i, sl in enumerate(deck.slides)]
    sld_ids = [e.get("id") for e in prs.slides._sldIdLst]
    if "section" not in kinds or len(sld_ids) != len(deck.slides):
        return
    groups: list[tuple[str, list[str]]] = []
    for i, sl in enumerate(deck.slides):
        title = " ".join(p.plain for p in sl.title.paragraphs).strip() if sl.title else ""
        if kinds[i] == "section":
            groups.append((title or f"Section {len(groups) + 1}", []))
        elif not groups:
            groups.append((title if kinds[i] == "cover" and title else "Default", []))
        groups[-1][1].append(str(sld_ids[i]))
    parts = []
    for k, (name, ids) in enumerate(groups):
        guid = str(uuid.uuid5(uuid.NAMESPACE_URL, f"slidemark-section-{k}-{name}")).upper()
        sec = etree.Element(f"{{{_NS_P14}}}section", nsmap={"p14": _NS_P14})
        sec.set("name", name)
        sec.set("id", "{" + guid + "}")
        lst = etree.SubElement(sec, f"{{{_NS_P14}}}sldIdLst")
        for sid in ids:
            etree.SubElement(lst, f"{{{_NS_P14}}}sldId").set("id", sid)
        parts.append(sec)
    ext = etree.Element(qn("p:ext"))
    ext.set("uri", _SECTION_URI)
    sl = etree.SubElement(ext, f"{{{_NS_P14}}}sectionLst", nsmap={"p14": _NS_P14})
    sl.extend(parts)
    pres = prs.part._element
    ext_lst = pres.find(qn("p:extLst"))
    if ext_lst is None:
        ext_lst = etree.SubElement(pres, qn("p:extLst"))
    for old in ext_lst.findall(qn("p:ext")):
        if old.get("uri") == _SECTION_URI:
            ext_lst.remove(old)
    ext_lst.insert(0, ext)


# --------------------------------------------------------------------------- items


def _name(pl: Placed, counters: dict[str, int]) -> str:
    el = pl.element
    if getattr(el, "id", None):
        return str(el.id)
    if isinstance(el, Text):
        field = el.attrs.get("field")
        if field == "footer":
            return "Footer"
        if field == "slide_number":
            return "Slide Number"
        base = {"title": "Title", "subtitle": "Subtitle", "lead": "Lead", "conclusion": "Conclusion"}.get(
            el.role
        )
        if base:
            return base
        base = {"heading": "Heading", "footnote": "Footnote"}.get(el.role, "Text")
    elif isinstance(el, Container):
        base = "Card"
    else:
        base = (
            "Pill"
            if isinstance(el, Shape) and "pill" in el.classes
            else {
                "table": "Table",
                "chart": "Chart",
                "image": "Image",
                "media": "Media",
                "code": "Code",
                "raw": "Raw",
                "shape": "Shape",
            }[el.type]
        )
    counters[base] = counters.get(base, 0) + 1
    return f"{base} {counters[base]}"


_DASH = {"dash": MSO_LINE.DASH, "dot": MSO_LINE.ROUND_DOT}


def _line(rc: RenderCtx, line, spec: tuple[float, str, str]) -> None:
    width, dash, color = spec
    line.color.rgb = rgb(rc.theme, color)
    line.width = Pt(width)
    if dash in _DASH:
        line.dash_style = _DASH[dash]


def _border_plan(rc: RenderCtx, st: Style):
    """(closed outline of the shape or None, per-side lines to draw or None).

    Without per-side borders this is the old ``line`` / ``line_width``. With them, sides that differ from the
    uniform line are drawn as separate lines on top of it; if a side is removed (``border-left: none``) the
    outline goes and every remaining side is drawn as a line."""
    dw = rc.theme.render.line_width
    uniform = None
    if st.line and (st.line_width is None or st.line_width > 0):
        uniform = (st.line_width if st.line_width is not None else dw, st.line_dash or "solid", st.line)
    if not has_side_borders(st):
        return uniform, None
    sides = side_borders(st, dw)
    diff = {k: v for k, v in sides.items() if v != uniform}
    if not diff:
        return uniform, None
    if uniform is not None and all(v is not None for v in diff.values()):
        return uniform, diff
    return None, {k: v for k, v in sides.items() if v is not None}


def _picture_fill(rc: RenderCtx, shp, st: Style, pl: Placed | None) -> bool:
    """``background: url(path)`` as a native picture fill (cover); False (with a diagnostic) when unusable."""
    m = re.fullmatch(r"\s*url\(\s*(.*?)\s*\)\s*", st.fill or "", re.I | re.S)
    src = m.group(1).strip("\"'") if m else ""
    path = resolve_image(rc, src) if src else None
    if path is None:
        rc.diag(
            "image-missing",
            f"background image not found: {src or st.fill}",
            "check the path, relative to the .md file; the shape has no fill",
        )
        return False
    try:
        from PIL import Image as PILImage

        with PILImage.open(path) as pil:
            iw, ih = pil.size
        _part, rid = shp.part.get_or_add_image_part(str(path))
        crop = cover_crop(iw, ih, pl.w if pl else 0, pl.h if pl else 0)
        op = st.opacity if st.opacity is not None and 0 <= st.opacity < 1 else None
        set_picture(shp._element.spPr, rid, crop, op)
        return True
    except Exception as e:
        rc.diag(
            "image-unreadable",
            f"cannot read background image {src}: {type(e).__name__}",
            "use PNG, JPEG or GIF",
        )
        return False


def _style_shape(rc: RenderCtx, shp, st: Style, pl: Placed | None = None) -> None:
    spPr = shp._element.spPr
    if st.fill and st.fill.lower().startswith("url("):
        if not _picture_fill(rc, shp, st, pl):
            shp.fill.background()
    elif not apply_fill(rc, spPr, st):
        shp.fill.background()
    outline, _sides = _border_plan(rc, st)
    if outline:
        _line(rc, shp.line, outline)
    else:
        shp.line.fill.background()
    if st.shadow:
        apply_shadow(rc, spPr, st, pl.w if pl else 0, pl.h if pl else 0)
    else:
        shp.shadow.inherit = False


def _turn(x: int, y: int, center: tuple[float, float], deg: float) -> tuple[int, int]:
    """Point (x, y) rotated clockwise by ``deg`` about ``center``."""
    rad = radians(deg)
    dx, dy = x - center[0], y - center[1]
    return round(center[0] + dx * cos(rad) - dy * sin(rad)), round(center[1] + dx * sin(rad) + dy * cos(rad))


def _side_lines(rc: RenderCtx, slide, pl: Placed, st: Style, name: str) -> None:
    """Per-side borders (``border-bottom: 2px solid #0DD``) as thin straight lines on the edges of ``pl``."""
    _outline, sides = _border_plan(rc, st)
    for side, spec in (sides or {}).items():
        w = max(round(spec[0] * 12700), 1)
        half = w // 2
        if side == "top":
            x0, y0, x1, y1 = pl.x, pl.y + half, pl.x + pl.w, pl.y + half
        elif side == "bottom":
            x0, y0, x1, y1 = pl.x, pl.y + pl.h - half, pl.x + pl.w, pl.y + pl.h - half
        elif side == "left":
            x0, y0, x1, y1 = pl.x + half, pl.y, pl.x + half, pl.y + pl.h
        else:
            x0, y0, x1, y1 = pl.x + pl.w - half, pl.y, pl.x + pl.w - half, pl.y + pl.h
        if st.rotation:  # the item is rotated about its center: so are its edges
            c = (pl.x + pl.w / 2, pl.y + pl.h / 2)
            (x0, y0), (x1, y1) = _turn(x0, y0, c, st.rotation), _turn(x1, y1, c, st.rotation)
        cx = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x0), Emu(y0), Emu(x1), Emu(y1))
        cx.name = f"{name} border {side}"
        style_el = cx._element.find(qn("p:style"))
        if style_el is not None:
            cx._element.remove(style_el)
        _line(rc, cx.line, spec)


def _round(shp, st: Style, pl: Placed) -> None:
    if st.radius and shp.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE:
        short = max(min(pl.w, pl.h), 1)
        shp.adjustments[0] = min(emu(st.radius) / short, 0.5)


def _autoshape(rc: RenderCtx, slide, pl: Placed, kind, name: str):
    shp = slide.shapes.add_shape(kind, Emu(pl.x), Emu(pl.y), Emu(pl.w), Emu(pl.h))
    shp.name = name
    style_el = shp._element.find(qn("p:style"))
    if (
        style_el is not None
    ):  # drop theme style refs: LibreOffice/PowerPoint would add a shadow and white text
        shp._element.remove(style_el)
    _style_shape(rc, shp, pl.style, pl)
    _round(shp, pl.style, pl)
    return shp


def _render_item(rc: RenderCtx, s, pl: Placed, counters: dict[str, int], use_placeholder: bool) -> None:
    """Draw ``pl``; then its per-side borders (CSS) and its rotation (``transform: rotate()``)."""
    before = len(s.shapes)
    _render_item0(rc, s, pl, counters, use_placeholder)
    el = pl.element
    st = pl.style
    if has_side_borders(st) and isinstance(el, (Container, Text, Shape, Code)):
        first = list(s.shapes)[before:]
        _side_lines(rc, s, pl, st, first[0].name if first else _name(pl, {}))
    if st.rotation:
        for shp in list(s.shapes)[before:]:
            tag = etree.QName(shp._element).localname
            if tag == "graphicFrame":
                rc.diag(
                    "css-unsupported",
                    f"rotate() on a {el.type} is not drawn (PowerPoint cannot rotate tables and charts)",
                    "remove the transform from tables and charts",
                )
            elif tag != "cxnSp":
                shp.rotation = st.rotation % 360


def _render_item0(rc: RenderCtx, s, pl: Placed, counters: dict[str, int], use_placeholder: bool) -> None:
    el = pl.element
    st = pl.style
    name = _name(pl, counters)
    if isinstance(el, Container):
        if st.fill or st.line:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if st.radius else MSO_SHAPE.RECTANGLE
            _autoshape(rc, s, pl, kind, name)
    elif isinstance(el, Text):
        field = el.attrs.get("field")
        if use_placeholder:
            shp = s.shapes.title
            shp.left, shp.top, shp.width, shp.height = Emu(pl.x), Emu(pl.y), Emu(pl.w), Emu(pl.h)
            shp.name = name
            tree = shp._element.getparent()
            tree.remove(shp._element)
            tree.append(shp._element)  # keep layout z-order
            tf = shp.text_frame
        elif st.fill or st.line:
            kind = MSO_SHAPE.ROUNDED_RECTANGLE if st.radius else MSO_SHAPE.RECTANGLE
            shp = _autoshape(rc, s, pl, kind, name)
            tf = shp.text_frame
        else:
            shp = s.shapes.add_textbox(Emu(pl.x), Emu(pl.y), Emu(pl.w), Emu(pl.h))
            shp.name = name
            tf = shp.text_frame
        gap = el.attrs.get("para_gap")
        fill_text(
            rc,
            tf,
            el.paragraphs,
            st,
            pl.font_scale,
            field=field,
            gap_em=float(gap) if isinstance(gap, (int, float)) else None,
            box_w=None if field else pl.w,
        )
        if el.attrs.get("nowrap"):  # one line measured by Chromium: a wider substitute font must not wrap it
            tf.word_wrap = False
        if "callout" in el.classes and st.line and not use_placeholder:
            _accent_bar(rc, s, pl, name)
    elif isinstance(el, Shape) and el.shape == "line":
        _connector(rc, s, pl, name)
    elif isinstance(el, Shape) and el.shape == "icon":
        add_icon(rc, s, pl, name)  # unknown names draw nothing (the parser warns)
    elif isinstance(el, Shape):
        kind = _SHAPES.get(el.shape)
        if kind is None:
            rc.diag(
                "shape",
                f"unknown shape {el.shape!r}",
                f"use one of: {', '.join(sorted(_SHAPES))}",
                line=el.line,
            )
            kind = MSO_SHAPE.RECTANGLE
        shp = _autoshape(rc, s, pl, kind, name)
        if el.paragraphs:
            chev = el.shape == "chevron"
            pill = bool(el.attrs.get("pill"))
            box_w: int | None = pl.w
            if chev:  # the preset text rectangle starts a point depth inside both ends
                adj = float(el.attrs.get("adj", rc.theme.layout.chevron_adj))
                box_w = pl.w - 2 * round(adj * min(pl.w, pl.h)) - int(el.attrs.get("icon_inset", 0))
            elif el.attrs.get("icon_inset") or pill:
                box_w = None
            fill_text(rc, shp.text_frame, el.paragraphs, st, pl.font_scale, box_w=box_w, squeeze=not chev)
            if pill:  # a pill is one line: a wider substitute font spills over, never wraps
                shp.text_frame.word_wrap = False
            if inset := el.attrs.get("icon_inset"):  # room for the icon the layout placed before the text
                shp.text_frame.margin_left = Emu(shp.text_frame.margin_left + int(inset))
        if el.shape == "chevron":
            shp.adjustments[0] = float(el.attrs.get("adj", rc.theme.layout.chevron_adj))
    elif isinstance(el, Table):
        add_table(rc, s, pl, name)
    elif isinstance(el, Chart):
        add_chart(rc, s, pl, name)
    elif isinstance(el, Image):
        if not add_image(rc, s, pl, name):
            _placeholder(rc, s, pl, name, f"[image: {el.alt or el.src[:60]}]")
    elif isinstance(el, Media):
        if not add_media(rc, s, pl, name):
            _placeholder(rc, s, pl, name, f"[{el.kind}: {el.alt or el.src}]")
    elif isinstance(el, Code):
        shp = _autoshape(rc, s, pl, MSO_SHAPE.ROUNDED_RECTANGLE if st.radius else MSO_SHAPE.RECTANGLE, name)
        fill_text(
            rc,
            shp.text_frame,
            code_paragraphs(el, rc.theme.render.code_style, rc.diag),
            st,
            pl.font_scale,
            para_gap=False,
        )
    elif isinstance(el, Raw) and el.kind == "math" and add_math(rc, s, pl, name):
        pass
    elif (
        isinstance(el, Raw)
        and el.kind == "html"
        and add_html_native(rc, pl, lambda p: _render_item(rc, s, p, counters, False))
    ):
        pass
    elif isinstance(el, Raw) and el.kind == "html" and add_html_image(rc, s, pl):
        pass
    elif isinstance(el, Raw) and el.kind == "html":
        _placeholder(rc, s, pl, name, "[html]")
    elif isinstance(el, Raw):
        _placeholder(rc, s, pl, name, f"[{el.kind}]")
        rc.diag(
            "raw",
            f"{el.kind} content is not rendered yet",
            "replace it with an image or native elements",
            "info",
            line=el.line,
        )


def _accent_bar(rc: RenderCtx, slide, pl: Placed, name: str) -> None:
    """Callout cards get a thick accent bar on the left edge, in the border color."""
    bar = pl.model_copy(
        update={
            "w": min(to_emu(rc.theme.layout.callout_bar_w), pl.w),
            "style": Style(fill=pl.style.line, line=None),
        }
    )
    _autoshape(rc, slide, bar, MSO_SHAPE.RECTANGLE, f"{name} accent")


def _connector(rc: RenderCtx, slide, pl: Placed, name: str) -> None:
    """A connector: straight (corner to opposite corner, flips pick the diagonal) or an elbow.

    Elbows are ``bentConnector3``: horizontal-vertical-horizontal for ``route="h"``; for ``route="v"`` the
    shape is rotated by 90 degrees (vertical-horizontal-vertical). ``adj`` places the middle segment.
    """
    el = pl.element
    fh, fv = bool(el.attrs.get("flip_h")), bool(el.attrs.get("flip_v"))
    x0, x1 = (pl.x + pl.w, pl.x) if fh else (pl.x, pl.x + pl.w)
    y0, y1 = (pl.y + pl.h, pl.y) if fv else (pl.y, pl.y + pl.h)
    elbow = bool(el.attrs.get("elbow"))
    kind = MSO_CONNECTOR.ELBOW if elbow else MSO_CONNECTOR.STRAIGHT
    cx = slide.shapes.add_connector(kind, Emu(x0), Emu(y0), Emu(x1), Emu(y1))
    cx.name = name
    style_el = cx._element.find(qn("p:style"))
    if style_el is not None:
        cx._element.remove(style_el)
    if elbow:
        _elbow_geometry(cx, el.attrs, pl, fh, fv)
    _glue(slide, cx, el.attrs, fh, fv)
    st = pl.style
    cx.line.color.rgb = rgb(rc.theme, st.line or "primary")
    cx.line.width = Pt(st.line_width if st.line_width is not None else rc.theme.render.connector_width)
    if el.attrs.get("head") == "arrow":
        ln = cx.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")


def _find_box(slide, box) -> object | None:
    """The first rectangle / rounded rectangle / diamond on ``slide`` with exactly this (x, y, w, h).

    All three have the same connection sites (0 top, 1 left, 2 bottom, 3 right). A diamond used to be skipped,
    so a connector leaving a decision node stayed unglued and LibreOffice drew its rotated geometry wrongly
    (the bend ran along the top edge of the target).
    """
    if not box:
        return None
    for shp in slide.shapes:
        if shp.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        if (shp.left, shp.top, shp.width, shp.height) == tuple(box):
            if shp.auto_shape_type in (MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE, MSO_SHAPE.DIAMOND):
                return shp
    return None


def _glue(slide, cx, attrs: dict, fh: bool, fv: bool) -> None:
    """Attach both ends to the connected boxes (connection sites 0 top, 1 left, 2 bottom, 3 right).

    Glued connectors follow their boxes when the slide is edited, and LibreOffice routes them by the sides.
    """
    route = attrs.get("route")
    src, dst = _find_box(slide, attrs.get("src_box")), _find_box(slide, attrs.get("dst_box"))
    if route not in ("v", "h") or src is None or dst is None:
        return
    if route == "v":
        s_idx, e_idx = (0, 2) if fv else (2, 0)
    else:
        s_idx, e_idx = (1, 3) if fh else (3, 1)
    c_nv = cx._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
    for tag, shp, idx in (("a:stCxn", src, s_idx), ("a:endCxn", dst, e_idx)):
        el = etree.SubElement(c_nv, qn(tag))
        el.set("id", str(shp.shape_id))
        el.set("idx", str(idx))


def _elbow_geometry(cx, attrs: dict, pl: Placed, fh: bool, fv: bool) -> None:
    """Set the bend position and, for vertical routes, the 90 degree rotation of a ``bentConnector3``."""
    adj = attrs.get("adj")
    geom = cx._element.spPr.find(qn("a:prstGeom"))
    if adj is not None and geom is not None and abs(adj - 0.5) > 1e-4:
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj1")
        gd.set("fmla", f"val {round(adj * 100000)}")
    if attrs.get("route") != "v":
        return
    # vertical route: rotate the unrotated (w=|dy|, h=|dx|) box by 90 degrees around the bounding box center.
    # With rot=90 the path starts top-right; flipV moves the start to the left, flipH to the bottom.
    xfrm = cx._element.spPr.find(qn("a:xfrm"))
    mx, my = pl.x + pl.w // 2, pl.y + pl.h // 2
    xfrm.set("rot", "5400000")
    for k in ("flipH", "flipV"):
        xfrm.attrib.pop(k, None)
    if fv:
        xfrm.set("flipH", "1")
    if not fh:
        xfrm.set("flipV", "1")
    xfrm.find(qn("a:off")).set("x", str(mx - pl.h // 2))
    xfrm.find(qn("a:off")).set("y", str(my - pl.w // 2))
    xfrm.find(qn("a:ext")).set("cx", str(pl.h))
    xfrm.find(qn("a:ext")).set("cy", str(pl.w))


def _placeholder(rc: RenderCtx, s, pl: Placed, name: str, label: str) -> None:
    st = Style(
        font_size=rc.theme.sizes.get("caption", DEFAULT_SIZES["caption"]),
        color="muted",
        fill="surface",
        line="border",
        line_width=1,
        align="center",
        valign="middle",
    ).merged(pl.style)
    shp = _autoshape(rc, s, pl.model_copy(update={"style": st}), MSO_SHAPE.RECTANGLE, name)
    shp.line.dash_style = 4  # dash
    from ..ir import Paragraph, Run

    fill_text(rc, shp.text_frame, [Paragraph(runs=[Run(text=label)])], st, 1.0)


# --------------------------------------------------------------------------- second pass


def _resolve_links(rc: RenderCtx) -> None:
    ids = {sl.id: i for i, sl in enumerate(rc.deck.slides) if sl.id}
    for _run_el, rpr, target, src_idx in rc.links:
        idx: int | None = None
        if target.isdigit():
            idx = int(target) - 1
        elif target in ids:
            idx = ids[target]
        rc.slide_index = src_idx
        if idx is None or not 0 <= idx < len(rc.slides):
            rc.diag(
                "bad-jump",
                f"slide link #{target} does not exist; link dropped",
                "use #<n> with n from 1 to the slide count, or a slide id= set with '@id=name'",
                "info",
            )
            continue
        src = rc.slides[src_idx]
        rid = src.part.relate_to(rc.slides[idx].part, RT.SLIDE)
        h = etree.Element(qn("a:hlinkClick"))
        h.set(qn("r:id"), rid)
        h.set("action", "ppaction://hlinksldjump")
        insert_rpr_child(rpr, h)


def _core_properties(prs, deck: Deck) -> None:
    cp = prs.core_properties
    title = deck.title
    if not title:
        for sl in deck.slides:
            if sl.title and sl.title.paragraphs:
                title = " ".join(p.plain for p in sl.title.paragraphs)
                break
    # core properties are limited to 255 characters (python-pptx raises beyond that)
    if title:
        cp.title = title[:255]
    if deck.author:
        cp.author = deck.author[:255]
    if deck.lang:
        cp.language = deck.lang[:255]
