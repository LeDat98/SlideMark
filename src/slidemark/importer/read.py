"""python-pptx shapes -> flat ``Item`` lists (absolute EMU boxes, parsed text); no layout decisions."""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field
from typing import Any

from lxml import etree
from pptx.oxml.ns import qn

from ..ir import Diagnostic

_MISSING = re.compile(r"\[(image|video|audio): (.*)\]", re.S)
MONO = ("consolas", "courier", "menlo", "monaco", "mono", "source code")


@dataclass
class RunT:
    text: str
    bold: bool = False
    italic: bool = False
    strike: bool = False
    sup: bool = False
    sub: bool = False
    code: bool = False
    color: str | None = None  # RRGGBB upper
    badge: str | None = None  # highlight color RRGGBB when the run is a badge
    link: str | None = None
    size: float | None = None


@dataclass
class ParaT:
    runs: list[RunT] = field(default_factory=list)
    marker: str | None = None  # "bullet" | "number"
    level: int = 0
    size: float | None = None
    marl: int | None = None
    lvl_attr: int | None = None
    align: str | None = None  # a:pPr algn: l | ctr | r | just

    @property
    def plain(self) -> str:
        return "".join(r.text for r in self.runs)

    @property
    def all_bold(self) -> bool:
        txt = [r for r in self.runs if r.text.strip()]
        return bool(txt) and all(r.bold for r in txt)


@dataclass
class CellT:
    paras: list[ParaT] = field(default_factory=list)
    hmerge: bool = False
    vmerge: bool = False


@dataclass
class ChartT:
    kind: str
    title: str | None
    categories: list[str]
    series: list[tuple[str, list[float | None]]]
    options: dict[str, str]
    totals: list[int] = field(default_factory=list)  # waterfall: indexes of the `=` total bars


@dataclass
class Item:
    kind: str  # text | shape | table | chart | image | line
    x: int
    y: int
    w: int
    h: int
    name: str = ""
    ph: str | None = None  # placeholder type (title, ctrTitle, subTitle, body, ftr, sldNum, dt, obj)
    fill: str | None = None  # RRGGBB, "x" (non-solid / theme style) or None
    line: bool = False
    prst: str | None = None
    paras: list[ParaT] = field(default_factory=list)
    rows: list[list[CellT]] = field(default_factory=list)
    chart: ChartT | None = None
    img: tuple[bytes, str] | None = None
    alt: str = ""
    has_slidenum: bool = False
    role: str | None = None  # set by structure: title/lead/conclusion/footnote/footer/decor/callout:<kind>
    uid: int = 0
    radius: float | None = None  # corner radius in pt of a rounded rectangle
    sid: int = 0  # shape id in the slide (what ``a:stCxn``/``a:endCxn`` point at)
    col_w: list[int] = field(default_factory=list)  # table: column widths (EMU)
    row_h: list[int] = field(default_factory=list)  # table: row heights (EMU)
    gantt: bool = False  # table: bars (filled shapes over body cells) were folded into it (`{.gantt}`)
    cropped: bool = False  # image: a:srcRect crop (``fit=cover``)
    latex: str | None = None  # math: the equation as LaTeX (``` math fence)
    missing: tuple[str, str] | None = None  # image: a "[image: label]" placeholder of a file that was absent
    line_color: str | None = None  # RRGGBB of a solid outline (box classes color the card border)

    @property
    def cx(self) -> float:
        return self.x + self.w / 2

    @property
    def cy(self) -> float:
        return self.y + self.h / 2

    @property
    def area(self) -> int:
        return max(self.w, 0) * max(self.h, 0)

    @property
    def text(self) -> str:
        return "\n".join(p.plain for p in self.paras).strip()

    @property
    def max_size(self) -> float | None:
        sizes = [p.size for p in self.paras if p.size]
        return max(sizes) if sizes else None


@dataclass
class ConnT:
    """A connector: start/end point in slide EMU, the shape ids it is glued to and its arrowheads."""

    x0: float
    y0: float
    x1: float
    y1: float
    st: int | None = None  # a:stCxn id
    en: int | None = None  # a:endCxn id
    arrow_end: bool = False  # a:tailEnd (the end of the line; what the renderer draws for ``a>b``)
    arrow_start: bool = False  # a:headEnd
    order: int = 0


@dataclass
class SlideData:
    items: list[Item] = field(default_factory=list)
    conns: list[ConnT] = field(default_factory=list)
    notes: str | None = None
    hidden: bool = False
    connectors: int = 0
    num_field: bool = False
    transition: str | None = None  # ``t=`` value: "fade", "push:0.5", ...
    build: bool = False  # click-by-click appear animations on shapes
    chart_notes: list[tuple[tuple[int, int, int, int], str]] = field(default_factory=list)  # `note=` callouts


@dataclass
class ReadCtx:
    accent: str  # theme accent RRGGBB
    diags: list[Diagnostic] = field(default_factory=list)
    slide_no: int = 0
    slide_index: dict[Any, int] = field(default_factory=dict)  # slide part -> 1-based number
    next_uid: int = 0

    def skip(self, what: str, hint: str = "rebuild it as native shapes, text or a picture") -> None:
        self.diags.append(
            Diagnostic(
                level="info",
                message=f"skipped {what}",
                slide=self.slide_no,
                rule="import-skipped",
                hint=hint,
            )
        )


# --------------------------------------------------------------------------- text


def _hex(el) -> str | None:
    if el is None:
        return None
    clr = el.find(qn("a:srgbClr"))
    if clr is not None and clr.get("val"):
        return clr.get("val").upper()
    return None


def _read_run(r, ctx: ReadCtx, part) -> RunT:
    t = r.find(qn("a:t"))
    text = ((t.text or "") if t is not None else "").replace("\u2060", "")  # render's range joiners
    run = RunT(text=text.replace(" ", " "))  # no-break spaces: the renderer's orphan control
    rpr = r.find(qn("a:rPr"))
    if rpr is None:
        return run
    run.bold = rpr.get("b") in ("1", "true")
    run.italic = rpr.get("i") in ("1", "true")
    run.strike = (rpr.get("strike") or "noStrike") != "noStrike"
    base = rpr.get("baseline")
    if base and base.lstrip("-").isdigit():
        run.sup = int(base) > 0
        run.sub = int(base) < 0
    if rpr.get("sz", "").isdigit():
        run.size = int(rpr.get("sz")) / 100
    fill = rpr.find(qn("a:solidFill"))
    run.color = _hex(fill)
    hl = rpr.find(qn("a:highlight"))
    if hl is not None:
        run.badge = _hex(hl) or "x"
        run.text = text.strip("　 ")  # badge padding: full-width / no-break spaces
    latin = rpr.find(qn("a:latin"))
    if latin is not None and any(m in (latin.get("typeface") or "").lower() for m in MONO):
        run.code = True
    link = rpr.find(qn("a:hlinkClick"))
    if link is not None:
        run.link = _link_target(link, ctx, part)
    return run


def _link_target(link, ctx: ReadCtx, part) -> str | None:
    rid = link.get(qn("r:id"))
    if not rid or part is None:
        return None
    try:
        rel = part.rels[rid]
    except KeyError:
        return None
    if rel.is_external:
        return rel.target_ref
    n = ctx.slide_index.get(getattr(rel, "target_part", None))
    return f"#{n}" if n else None


def read_paras(
    txbody, ctx: ReadCtx, part, default_bullets: bool = False, keep_empty: bool = False
) -> tuple[list[ParaT], bool]:
    """Paragraphs of a txBody; the flag says a slide-number field was seen."""
    out: list[ParaT] = []
    slidenum = False
    if txbody is None:
        return out, slidenum
    for p in txbody.findall(qn("a:p")):
        ppr = p.find(qn("a:pPr"))
        marker = "bullet" if default_bullets else None
        lvl_attr = marl = None
        if ppr is not None:
            if ppr.find(qn("a:buNone")) is not None:
                marker = None
            elif ppr.find(qn("a:buAutoNum")) is not None:
                marker = "number"
            elif ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buBlip")) is not None:
                marker = "bullet"
            if ppr.get("lvl", "").isdigit():
                lvl_attr = int(ppr.get("lvl"))
            if ppr.get("marL", "").lstrip("-").isdigit():
                marl = int(ppr.get("marL"))
        runs: list[RunT] = []
        for ch in p:
            tag = etree.QName(ch).localname
            if tag == "r":
                runs.append(_read_run(ch, ctx, part))
            elif tag == "br":
                runs.append(RunT(text="\n"))
            elif tag == "fld":
                if (ch.get("type") or "").startswith("slidenum"):
                    slidenum = True
                else:
                    runs.append(_read_run(ch, ctx, part))
        while runs and runs[-1].text == "\n":
            runs.pop()
        if not "".join(r.text for r in runs).strip() and not keep_empty:
            continue
        size = next((r.size for r in runs if r.size), None)
        if size is None:
            end = p.find(qn("a:endParaRPr"))
            if end is not None and end.get("sz", "").isdigit():
                size = int(end.get("sz")) / 100
        algn = ppr.get("algn") if ppr is not None else None
        out.append(ParaT(runs=runs, marker=marker, size=size, marl=marl, lvl_attr=lvl_attr, align=algn))
    _levels(out)
    return out, slidenum


def _levels(paras: list[ParaT]) -> None:
    if any(p.lvl_attr for p in paras):
        for p in paras:
            p.level = p.lvl_attr or 0
        return
    marls = sorted({p.marl for p in paras if p.marker and p.marl is not None})
    for p in paras:
        if p.marker and p.marl is not None:
            p.level = marls.index(p.marl)


# --------------------------------------------------------------------------- geometry / style


def _fill_of(el) -> tuple[str | None, bool]:
    sppr = el.find(qn("p:spPr"))
    style = el.find(qn("p:style"))
    fill: str | None = None
    line = False
    if sppr is not None:
        if sppr.find(qn("a:noFill")) is not None:
            fill = None
        elif sppr.find(qn("a:solidFill")) is not None:
            fill = _hex(sppr.find(qn("a:solidFill"))) or "x"
        elif any(sppr.find(qn(t)) is not None for t in ("a:gradFill", "a:blipFill", "a:pattFill")):
            fill = "x"
        elif style is not None and _ref_idx(style, "a:fillRef"):
            fill = "x"
        ln = sppr.find(qn("a:ln"))
        if ln is not None:
            line = ln.find(qn("a:noFill")) is None and (
                ln.find(qn("a:solidFill")) is not None or (style is not None and _ref_idx(style, "a:lnRef"))
            )
        elif style is not None and _ref_idx(style, "a:lnRef"):
            line = True
    return fill, line


def _ref_idx(style, tag: str) -> bool:
    ref = style.find(qn(tag))
    return ref is not None and (ref.get("idx") or "0") != "0"


class Tf:
    """Affine map from a group's child space to slide EMU: ``abs = a + b * local`` per axis."""

    def __init__(self, ax=0.0, bx=1.0, ay=0.0, by=1.0):
        self.ax, self.bx, self.ay, self.by = ax, bx, ay, by

    def box(self, x, y, w, h) -> tuple[int, int, int, int]:
        return (
            round(self.ax + self.bx * x),
            round(self.ay + self.by * y),
            round(self.bx * w),
            round(self.by * h),
        )

    def child(self, el) -> Tf:
        xf = el.find(qn("p:grpSpPr") + "/" + qn("a:xfrm"))
        if xf is None:
            return self
        try:
            off, ext = xf.find(qn("a:off")), xf.find(qn("a:ext"))
            choff, chext = xf.find(qn("a:chOff")), xf.find(qn("a:chExt"))
            ox, oy, cx, cy = int(off.get("x")), int(off.get("y")), int(ext.get("cx")), int(ext.get("cy"))
            qx, qy = int(choff.get("x")), int(choff.get("y"))
            qw, qh = int(chext.get("cx")) or cx or 1, int(chext.get("cy")) or cy or 1
        except (AttributeError, TypeError, ValueError):
            return self
        sx, sy = cx / qw, cy / qh
        # local -> group space: ox + (x - qx) * sx ; then self
        gx_a, gx_b = ox - qx * sx, sx
        gy_a, gy_b = oy - qy * sy, sy
        return Tf(
            self.ax + self.bx * gx_a,
            self.bx * gx_b,
            self.ay + self.by * gy_a,
            self.by * gy_b,
        )


# --------------------------------------------------------------------------- charts

_KINDS = (
    ("DOUGHNUT", "doughnut"),
    ("PIE", "pie"),
    ("RADAR", "radar"),
    ("XY_SCATTER", "scatter"),
    ("BUBBLE", "scatter"),
    ("LINE", "line"),
    ("AREA", "area"),
)


def chart_kind(name: str) -> str | None:
    for key, kind in _KINDS:
        if key in name:
            return kind
    if "COL" in name or "BAR" in name:
        horiz = "COL" not in name
        stacked = "STACKED" in name
        return ("stacked-" if stacked else "") + ("bar" if horiz else "column")
    return None


def _num(v) -> float | None:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return f if f == f and abs(f) != float("inf") else None


def _xtext(t: str) -> str:
    f = _num(t)
    if f is None:
        return t
    return str(int(f)) if f == int(f) else repr(f)


def read_chart(shape, ctx: ReadCtx) -> ChartT | None:
    chart = shape.chart
    try:
        ct = chart.chart_type
        name = ct.name if hasattr(ct, "name") else str(ct)
    except Exception:
        ctx.skip("a chart of unknown type")
        return None
    kind = chart_kind(name)
    if kind is None:
        ctx.skip(f"chart type {name.lower()}", "use one of the 10 chart kinds: bar, column, line, pie, ...")
        return None
    if len(chart.plots) > 1:
        ctx.diags.append(
            Diagnostic(
                level="info",
                message="combo chart: only the first plot was imported",
                slide=ctx.slide_no,
                rule="import-skipped",
                hint="split the chart if both plots matter",
            )
        )
    plot = chart.plots[0]
    series = list(plot.series)
    if kind == "stacked-column" and series and (series[0].name or "") == _wf().SERIES[0]:
        got = _read_waterfall(chart, series)
        if got:
            return got
    if kind == "scatter":
        xs: list[str] = []
        if series:
            xs = [
                _xtext(v.text or "")
                for v in series[0]._element.xpath("./c:xVal//c:pt/c:v")  # type: ignore[attr-defined]
            ]
        cats = xs
    else:
        cats = [str(c) for c in plot.categories]
    data = [(s.name or "", [_num(v) for v in s.values]) for s in series]
    carrier = False
    if kind in ("stacked-bar", "stacked-column") and len(series) > 1 and _is_total_carrier(series[-1]):
        carrier = True  # the hidden series of the stack-total labels: the build recreates it
        series = series[:-1]
        full_data = data
        data = data[:-1]
    else:
        full_data = data
    opts: dict[str, str] = {}
    title = None
    try:
        if chart.has_title:
            title = chart.chart_title.text_frame.text.strip() or None
    except Exception:
        title = None
    try:
        nser = len(data)
        pie = kind in ("pie", "doughnut")
        default = None if (nser <= 1 and not pie) else ("right" if pie else "bottom")
        if nser > 1:
            default = "bottom"
        elif pie:
            default = None  # documented default: none for a single series
        if chart.has_legend:
            pos = {1: "bottom", -4107: "bottom", -4131: "left", -4152: "right", -4160: "top"}.get(
                int(chart.legend.position), "right"
            )
            if pos != default:
                opts["legend"] = pos
        elif default is not None:
            opts["legend"] = "none"
    except Exception:
        pass
    fmt = None
    try:
        if kind in ("stacked-bar", "stacked-column") and carrier != bool(plot.has_data_labels):
            nonneg = all(v is None or v >= 0 for _, vals in data for v in vals)
            if carrier or nonneg:  # totals follow the segment labels by default
                opts["totals"] = "on" if carrier else "off"
        if plot.has_data_labels:
            dl = plot.data_labels
            if kind in ("pie", "doughnut") and dl.show_percentage:
                opts["labels"] = "percent"
            elif dl.show_value:
                opts["labels"] = "on"
            if "labels" in opts and not dl.number_format_is_linked and dl.number_format != "General":
                fmt = dl.number_format
    except Exception:
        pass
    try:
        va = chart.value_axis
        if not va.visible:
            opts["axis"] = "off"
        if fmt is None and not va.tick_labels.number_format_is_linked:
            nf = va.tick_labels.number_format
            if nf and nf != "General":
                fmt = nf
        zero_base = kind in ("column", "bar", "stacked-column", "stacked-bar", "area") and all(
            v is None or v >= 0 for _, vals in data for v in vals
        )
        neg = _auto_neg(kind, data, va)
        if (
            va.minimum_scale is not None
            and not (zero_base and va.minimum_scale == 0)
            and not (neg and abs(neg[0] - va.minimum_scale) < 1e-9)
        ):
            opts["min"] = f"{va.minimum_scale:g}"  # a zero base on bars is the build default
        if (
            va.maximum_scale is not None
            and not _is_auto_max(kind, full_data, va)
            and not (neg and abs(neg[1] - va.maximum_scale) < 1e-9)
        ):
            opts["max"] = f"{va.maximum_scale:g}"
    except Exception:
        pass
    if fmt:
        opts["fmt"] = fmt
    return ChartT(kind=kind, title=title, categories=cats, series=data, options=opts)


def _is_total_carrier(ser) -> bool:
    """The last series is SlideMark's invisible stack-total carrier (no fill, a name marker)."""
    from ..render.objects import BLANK_CARRIER, TOTAL_CARRIER

    el = ser._element
    sp = el.find(qn("c:spPr"))
    return (ser.name or "") in (TOTAL_CARRIER, BLANK_CARRIER) and (
        sp is not None and sp.find(qn("a:noFill")) is not None
    )


def _wf():
    from ..render import waterfall

    return waterfall


def _read_waterfall(chart, series) -> ChartT | None:
    """A waterfall SlideMark drew (stacked columns with an invisible first series) -> one series + totals."""
    wf = _wf()
    if len(series) != len(wf.SERIES) or [s.name for s in series[:-1]] != list(wf.SERIES[:-1]):
        return None
    try:
        cols = [[_num(v) for v in s.values] for s in series]
        values, totals = wf.decode(cols)
        cats = [str(c) for c in chart.plots[0].categories]
        opts: dict[str, str] = {}
        title = chart.chart_title.text_frame.text.strip() or None if chart.has_title else None
        if chart.has_legend:
            pos = {1: "bottom", -4107: "bottom", -4131: "left", -4152: "right", -4160: "top"}.get(
                int(chart.legend.position), "right"
            )
            opts["legend"] = pos
        if series[wf.PAD]._element.find(qn("c:dLbls")) is not None:
            opts["labels"] = "on"
        va = chart.value_axis
        if not va.visible:
            opts["axis"] = "off"
        if not va.tick_labels.number_format_is_linked and va.tick_labels.number_format != "General":
            opts["fmt"] = va.tick_labels.number_format
        from ..theme import RenderTokens

        auto = wf.axis(wf.bars(values, totals), RenderTokens())
        lo, hi = va.minimum_scale, va.maximum_scale
        if lo is not None and not (auto and abs(lo - auto[0]) < 1e-9):
            opts["min"] = f"{lo:g}"
        if hi is not None and not (auto and abs(hi - auto[1]) < 1e-9):
            opts["max"] = f"{hi:g}"
        return ChartT(
            kind="waterfall",
            title=title,
            categories=cats,
            series=[(series[wf.PAD].name or "", values)],
            options=opts,
            totals=totals,
        )
    except Exception:
        return None


def _auto_neg(kind: str, data, va):
    """The (min, max, unit) the renderer picks for bars with negatives, else None."""
    try:
        from ..render.axis import neg_axis
        from ..theme import RenderTokens

        return neg_axis(kind, [vals for _, vals in data], RenderTokens())
    except Exception:
        return None


def _is_auto_max(kind: str, data, va) -> bool:
    """True when ``c:max`` is the nice max the renderer writes by itself (no `max=` needed on re-build)."""
    try:
        from ..render.axis import auto_axis
        from ..theme import RenderTokens

        if kind not in ("column", "bar", "stacked-column", "stacked-bar", "area"):
            return False
        if va.minimum_scale not in (None, 0):
            return False
        auto = auto_axis(kind, [vals for _, vals in data], RenderTokens())
        return bool(auto) and abs(auto[0] - va.maximum_scale) < 1e-9
    except Exception:
        return False


# --------------------------------------------------------------------------- shapes


def _alt(el, name: str) -> str:
    for c in el.iter(qn("p:cNvPr")):
        return c.get("descr") or ""
    return ""


_TRANS = ("fade", "push", "wipe", "split", "cover", "zoom", "morph")


def read_transition(sld) -> str | None:
    """The ``t=`` token of a slide: transition name plus ``:seconds`` when a duration was written."""
    best = None
    for tr in sld.iter():
        if not isinstance(tr.tag, str) or etree.QName(tr).localname != "transition":
            continue
        if best is None or any("dur" in etree.QName(a).localname for a in tr.attrib):
            best = tr
    if best is None:
        return None
    name = next(
        (etree.QName(c).localname for c in best if etree.QName(c).localname in _TRANS),
        None,
    )
    if name is None:
        return None
    dur = next((v for a, v in best.attrib.items() if etree.QName(a).localname == "dur"), None)
    if dur and dur.isdigit():
        secs = int(dur) / 1000
        if not (name == "morph" and secs == 2.0):
            return f"{name}:{secs:g}"
    return name


def read_sections(prs) -> list[tuple[str, list[int]]]:
    """PowerPoint sections as (name, [1-based slide numbers])."""
    try:
        ids = [int(e.get("id")) for e in prs.slides._sldIdLst]
        out = []
        for sec in prs.part._element.iter():
            if isinstance(sec.tag, str) and etree.QName(sec).localname == "section":
                nums = [
                    ids.index(int(s.get("id"))) + 1
                    for s in sec.iter()
                    if isinstance(s.tag, str)
                    and etree.QName(s).localname == "sldId"
                    and int(s.get("id")) in ids
                ]
                out.append((sec.get("name") or "", nums))
        return out
    except Exception:
        return []


CHART_NOTE = "ChartNote"  # shape names the renderer gives a `note=` callout and its pointer
_HL = re.compile(r" hl=(.+)$")  # the renderer appends the `hl=` categories to a chart's shape name


def read_slide(slide, ctx: ReadCtx) -> SlideData:
    data = SlideData()
    data.hidden = slide._element.get("show") in ("0", "false")
    part = slide.part
    _walk(slide.shapes, Tf(), data, ctx, part)
    _attach_notes(data)
    try:
        data.transition = read_transition(slide._element)
        data.build = bool(
            slide._element.xpath(
                ".//*[local-name()='cTn'][@nodeType='clickEffect'][@presetClass='entr']"
                "[.//*[local-name()='attrName' and text()='style.visibility']]"
            )
        )
    except Exception:
        pass
    try:
        if slide.has_notes_slide:
            txt = (slide.notes_slide.notes_text_frame.text or "").strip()
            data.notes = txt or None
    except Exception:
        data.notes = None
    return data


def _attach_notes(data: SlideData) -> None:
    """The callouts SlideMark draws for ``note=`` belong to the chart whose frame holds them."""
    for box, text in data.chart_notes:
        cx, cy = box[0] + box[2] / 2, box[1] + box[3] / 2
        for it in data.items:
            if it.kind == "chart" and it.chart and it.x <= cx <= it.x + it.w and it.y <= cy <= it.y + it.h:
                it.chart.options["note"] = text
                if (
                    it.chart.kind == "line"
                ):  # a note pins the line axis: that scale is not an author's min / max
                    _drop_pinned_axis(it.chart)
                break


def _drop_pinned_axis(ch: ChartT) -> None:
    from ..render.axis import line_axis
    from ..theme import RenderTokens

    got = line_axis([v for _, v in ch.series], RenderTokens())
    if got and "min" in ch.options and abs(float(ch.options["min"]) - got[0]) < 1e-9:
        del ch.options["min"]
    if got and "max" in ch.options and abs(float(ch.options["max"]) - got[1]) < 1e-9:
        del ch.options["max"]


def _walk(shapes, tf: Tf, data: SlideData, ctx: ReadCtx, part) -> None:
    for sh in shapes:
        try:
            _one(sh, tf, data, ctx, part)
        except Exception as e:  # never raise: a shape we cannot read is skipped
            ctx.skip(f"shape {getattr(sh, 'name', '?')!r} ({type(e).__name__})")


def _new(ctx: ReadCtx, kind: str, box, sid: int = 0, **kw) -> Item:
    ctx.next_uid += 1
    return Item(kind=kind, x=box[0], y=box[1], w=box[2], h=box[3], uid=ctx.next_uid, sid=sid, **kw)


def _radius(geom, box) -> float | None:
    """Corner radius (pt) from the ``adj`` guide of a ``roundRect``; None when the guide is absent."""
    for gd in geom.iter(qn("a:gd")):
        m = (gd.get("fmla") or "").split()
        if gd.get("name") == "adj" and len(m) == 2 and m[1].lstrip("-").isdigit():
            return int(m[1]) / 100000 * min(box[2], box[3]) / 12700
    return None


def _int(v) -> int | None:
    return int(v) if isinstance(v, str) and v.lstrip("-").isdigit() else None


def _arrow(ln, tag: str) -> bool:
    end = ln.find(qn(tag)) if ln is not None else None
    return end is not None and (end.get("type") or "none") != "none"


def _read_conn(el, tf: Tf, order: int) -> ConnT | None:
    """Endpoints of a ``p:cxnSp`` (path start = top-left of its box, flips and rotation applied)."""
    xf = el.find(qn("p:spPr") + "/" + qn("a:xfrm"))
    if xf is None or xf.find(qn("a:off")) is None or xf.find(qn("a:ext")) is None:
        return None
    try:
        x, y = int(xf.find(qn("a:off")).get("x")), int(xf.find(qn("a:off")).get("y"))
        w, h = int(xf.find(qn("a:ext")).get("cx")), int(xf.find(qn("a:ext")).get("cy"))
        rot = (int(xf.get("rot") or 0) / 60000) % 360
    except (TypeError, ValueError):
        return None
    mx, my = x + w / 2, y + h / 2
    pts = []
    for px, py in ((x, y), (x + w, y + h)):
        if xf.get("flipH") in ("1", "true"):
            px = 2 * mx - px
        if xf.get("flipV") in ("1", "true"):
            py = 2 * my - py
        if rot:
            c, s = math.cos(math.radians(rot)), math.sin(math.radians(rot))
            dx, dy = px - mx, py - my
            px, py = mx + dx * c - dy * s, my + dx * s + dy * c
        pts.append((tf.ax + tf.bx * px, tf.ay + tf.by * py))
    cnv = el.find(qn("p:nvCxnSpPr") + "/" + qn("p:cNvCxnSpPr"))
    st = en = None
    if cnv is not None:
        a, b = cnv.find(qn("a:stCxn")), cnv.find(qn("a:endCxn"))
        st = _int(a.get("id")) if a is not None else None
        en = _int(b.get("id")) if b is not None else None
    ln = el.find(qn("p:spPr") + "/" + qn("a:ln"))
    return ConnT(
        pts[0][0],
        pts[0][1],
        pts[1][0],
        pts[1][1],
        st,
        en,
        _arrow(ln, "a:tailEnd"),
        _arrow(ln, "a:headEnd"),
        order,
    )


def _one(sh, tf: Tf, data: SlideData, ctx: ReadCtx, part) -> None:
    el = sh._element
    tag = etree.QName(el).localname
    name = sh.name or ""
    if tag == "grpSp":
        _walk(sh.shapes, tf.child(el), data, ctx, part)
        return
    if name.startswith(CHART_NOTE):  # the pointer of a `note=` is not content
        if tag == "sp":
            text = " ".join("".join(el.xpath(".//a:t/text()")).replace("\u2060", "").split())
            box = tf.box(sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0)
            if text and name.startswith(CHART_NOTE) and not name.startswith(CHART_NOTE + "Line"):
                data.chart_notes.append((box, text))
        return
    if tag == "cxnSp":
        data.connectors += 1
        conn = _read_conn(el, tf, len(data.conns))
        if conn is not None:
            data.conns.append(conn)
        return
    try:
        box = tf.box(sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0)
    except Exception:
        box = (0, 0, 0, 0)
    if tag == "sp":
        ph = el.find(qn("p:nvSpPr") + "/" + qn("p:nvPr") + "/" + qn("p:ph"))
        ph_type = (ph.get("type") or "obj") if ph is not None else None
        paras_m = el.xpath(".//*[local-name()='oMathPara']")
        if paras_m or el.xpath(".//*[local-name()='oMath']"):
            from .mathml import omml_to_latex

            src = paras_m[0] if paras_m else el.xpath(".//*[local-name()='oMath']")[0]
            latex, exact = omml_to_latex(src)
            if not latex:
                ctx.skip(f"equation {name!r}", "write the formula as text or an image")
                return
            if not exact:
                ctx.diags.append(
                    Diagnostic(
                        level="info",
                        message=f"equation {name!r} converted approximately",
                        slide=ctx.slide_no,
                        rule="import-math",
                        hint="check the ```math block against the original formula",
                    )
                )
            data.items.append(_new(ctx, "math", box, sid=sh.shape_id, name=name, latex=latex))
            return
        fill, line = _fill_of(el)
        ln_el = el.find(qn("p:spPr") + "/" + qn("a:ln"))
        ln_fill = ln_el.find(qn("a:solidFill")) if ln_el is not None else None
        line_color = _hex(ln_fill) if line and ln_fill is not None else None
        geom = el.find(qn("p:spPr") + "/" + qn("a:prstGeom"))
        prst = geom.get("prst") if geom is not None else None
        radius = _radius(geom, box) if prst == "roundRect" else None
        ph_m = _MISSING.fullmatch(" ".join(t.strip() for t in el.xpath(".//a:t/text()")) or "")
        if ph_m and re.match(r"(Image|Media) \d+$", name):  # the library's placeholder for an absent file
            data.items.append(
                _new(ctx, "image", box, sid=sh.shape_id, name=name, missing=(ph_m.group(1), ph_m.group(2)))
            )
            return
        default_bullets = ph_type in ("body", "obj")
        paras, num = read_paras(
            el.find(qn("p:txBody")), ctx, part, default_bullets, keep_empty=name.lower().startswith("code")
        )
        while paras and not paras[-1].plain.strip():
            paras.pop()
        data.num_field = data.num_field or num
        it = _new(
            ctx,
            "text" if paras else "shape",
            box,
            sid=sh.shape_id,
            name=name,
            ph=ph_type,
            fill=fill,
            line=line,
            line_color=line_color,
            prst=prst,
            radius=radius,
            paras=paras,
            has_slidenum=num,
        )
        data.items.append(it)
    elif tag == "pic":
        links = el.xpath(".//*[local-name()='videoFile' or local-name()='audioFile']/@*[local-name()='link']")
        if links:  # movie/sound: the media part becomes ![alt](images/N-k.mp4), like pictures
            try:
                part = sh.part.related_part(links[0])
                blob, ext = part.blob, part.partname.ext.lstrip(".")
            except Exception:
                ctx.skip(f"media {name!r}", "media part could not be read")
                return
            data.items.append(
                _new(ctx, "image", box, sid=sh.shape_id, name=name, img=(blob, ext), alt=_alt(el, name))
            )
            return
        blob_ext = None
        svg_ids = el.xpath(".//*[local-name()='svgBlip']/@*[local-name()='embed']")
        if svg_ids:  # Office SVG picture: keep the vector, not the PNG fallback
            try:
                part = sh.part.related_part(svg_ids[0])
                blob_ext = (part.blob, "svg")
            except Exception:
                blob_ext = None
        if blob_ext is None:
            img = sh.image
            blob_ext = (img.blob, img.ext)
        it = _new(ctx, "image", box, sid=sh.shape_id, name=name, img=blob_ext, alt=_alt(el, name))
        src = el.find(qn("p:blipFill") + "/" + qn("a:srcRect"))
        it.cropped = src is not None and any(int(src.get(k) or 0) != 0 for k in ("l", "t", "r", "b"))
        data.items.append(it)
    elif tag == "graphicFrame":
        if sh.has_table:
            rows: list[list[CellT]] = []
            for tr in sh.table._tbl.tr_lst:
                row = []
                for tc in tr.tc_lst:
                    paras, _ = read_paras(tc.find(qn("a:txBody")), ctx, part)
                    row.append(
                        CellT(
                            paras=paras,
                            hmerge=tc.get("hMerge") in ("1", "true"),
                            vmerge=tc.get("vMerge") in ("1", "true"),
                        )
                    )
                rows.append(row)
            tit = _new(ctx, "table", box, sid=sh.shape_id, name=name, rows=rows)
            try:
                sx = box[2] / (sum(c.width for c in sh.table.columns) or 1)
                sy = box[3] / (sum(r.height for r in sh.table.rows) or 1)
                tit.col_w = [round(c.width * sx) for c in sh.table.columns]
                tit.row_h = [round(r.height * sy) for r in sh.table.rows]
            except Exception:
                pass
            data.items.append(tit)
        elif sh.has_chart:
            ch = read_chart(sh, ctx)
            if ch is not None:
                m = _HL.search(name)
                if m:
                    ch.options["hl"] = m.group(1)
                data.items.append(_new(ctx, "chart", box, sid=sh.shape_id, name=name, chart=ch))
        else:
            ctx.skip(f"object {name!r} (SmartArt, OLE or other)")
    else:
        ctx.skip(f"{tag} {name!r}")
