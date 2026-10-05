"""Brand-brief gaps (eval run 8): .kpi CSS, element-token aliases, file icons, more glyphs, CSS width."""

from __future__ import annotations

from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import icons
from slidemark.build import build_deck
from slidemark.parser import parse

HEAD = "theme: none\ncolors: bg=#FFFFFF fg=#222222 primary=#2255CC\n"
EMU_PT = 12700


def _build(md: str, tmp_path, name="o.pptx"):
    deck = parse(md)
    out = tmp_path / name
    build_deck(deck, out, tmp_path)
    return Presentation(str(out)), deck


def _runs(prs, slide=0):
    out = {}
    for sh in prs.slides[slide].shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    out[r.text] = r
    return out


KPI = "# Out\n## A {.kpi}\n-24%\nvs prior\n## B {.kpi}\n94%\nsurvey\n"


def test_kpi_css_styles_the_value(tmp_path):
    prs, _ = _build(
        HEAD + "```css\n.kpi { font-size: 40pt; color: #15803D; font-weight: bold }\n```\n" + KPI, tmp_path
    )
    r = _runs(prs)
    assert r["-24%"].font.size.pt == 40 and str(r["-24%"].font.color.rgb) == "15803D"
    assert r["vs prior"].font.size.pt < 20 and str(r["vs prior"].font.color.rgb) != "15803D"


def test_kpi_caption_selector(tmp_path):
    css = "```css\n.kpi .caption { color: #CC0000; font-size: 14pt }\n.kpi h2 { color: #0000CC }\n```\n"
    prs, _ = _build(HEAD + css + KPI, tmp_path)
    r = _runs(prs)
    assert str(r["vs prior"].font.color.rgb) == "CC0000" and r["vs prior"].font.size.pt == 14
    assert str(r["A"].font.color.rgb) == "0000CC"


def test_element_token_aliases(tmp_path):
    md = "theme: none\nstyle: title.weight=bold heading.italic=on body.size=11\n\n# Hi\n## Box\nbody text\n"
    deck = parse(md)
    assert not [d for d in deck.diagnostics if d.rule == "unknown-token"]
    sels = {r.selector: r.style for r in deck.css}
    assert sels["h1"].bold is True and sels["h2"].italic is True and sels["p"].font_size == 11


def test_element_token_hints():
    deck = parse("theme: none\nstyle: title.wieght=bold bold=on\n\n# Hi\n")
    hints = [d.hint for d in deck.diagnostics if d.rule == "unknown-token"]
    assert any("h1.font-weight" in h for h in hints) and len(hints) == 2


def test_title_weight_token_renders_bold(tmp_path):
    prs, _ = _build("theme: none\nstyle: title.weight=bold\n\n# Hello\n- x\n", tmp_path)
    assert _runs(prs)["Hello"].font.bold is True


def test_new_glyphs_parse_and_alias():
    for n in ("headphones", "battery", "music", "sparkles", "smile", "tooth", "house", "lightning"):
        assert icons.path(n) and icons.commands(n)
    deck = parse("# T\n## a {icon=headphones}\n- x\n")
    assert not deck.diagnostics


def test_unknown_icon_hint_lists_three_and_file():
    deck = parse("# T\n## a {icon=headphone-x}\n- x\n")
    (d,) = [d for d in deck.diagnostics if d.rule == "unknown-icon"]
    assert "icon=file.svg" in d.hint and "headphones" in d.hint
    deck = parse("# T\n## a {icon=zzzzzz}\n- x\n")
    assert "icon=file.svg" in deck.diagnostics[0].hint


def test_svg_file_icon_native_geometry(tmp_path):
    (tmp_path / "s.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><path d="M1 1h8v8h-8z"/>'
        '<circle cx="5" cy="5" r="2" fill="#fff"/></svg>'
    )
    prs, deck = _build(HEAD + "\n# T\n## a {icon=s.svg}\n- x\n", tmp_path)
    assert not [d for d in deck.diagnostics if d.level != "info"]
    icon = next(s for s in prs.slides[0].shapes if s.name.startswith("icon s.svg"))
    assert icon._element.spPr.find(qn("a:custGeom")) is not None


def test_svg_file_icon_stroke_becomes_picture(tmp_path):
    (tmp_path / "t.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="#c00">'
        '<circle cx="12" cy="12" r="9"/></svg>'
    )
    prs, _ = _build(HEAD + "\n# T\n## a {icon=t.svg}\n- x\n", tmp_path)
    assert any(s.shape_type == 13 for s in prs.slides[0].shapes)  # PICTURE


def test_svg_file_icon_missing_warns(tmp_path):
    _, deck = _build(HEAD + "\n# T\n## a {icon=nope.svg}\n- x\n", tmp_path)
    (d,) = [d for d in deck.diagnostics if d.rule == "icon-missing"]
    assert d.hint


def test_svg_layers_never_raise():
    for bad in ("", "<svg", "<svg xmlns='http://www.w3.org/2000/svg'><path d='M1 1 L'/></svg>", "x" * 10):
        assert icons.svg_layers(bad) is None


def test_css_width_fit_content_hugs_text(tmp_path):
    css = (
        "```css\n.sticker { width: fit-content; background: #C2410C; color: #fff; padding: 6px 12px }\n```\n"
    )
    md = HEAD + css + "\n# T\n@1\n{.sticker}\nNew edition\n"
    prs, _ = _build(md, tmp_path)
    sh = next(s for s in prs.slides[0].shapes if s.has_text_frame and "New edition" in s.text_frame.text)
    assert sh.width < 0.25 * prs.slide_width and sh.left < 0.1 * prs.slide_width


def test_css_width_percent_and_inline_block(tmp_path):
    css = "```css\n.half { width: 40% }\n.tag { display: inline-block; text-align: right }\n```\n"
    md = HEAD + css + "\n# T\n@1\n{.half}\nHalf\n\n{.tag}\nTag\n"
    prs, _ = _build(md, tmp_path)
    by = {s.text_frame.text: s for s in prs.slides[0].shapes if s.has_text_frame}
    full = prs.slide_width
    assert abs(by["Half"].width / full - 0.4 * (1 - 0.1)) < 0.1 or by["Half"].width < 0.5 * full
    assert by["Tag"].width < 0.2 * full and by["Tag"].left > 0.5 * full


def test_css_width_bad_value_is_a_diagnostic():
    deck = parse(HEAD + "```css\n.x { width: banana; display: none }\n```\n\n# T\n- a\n")
    assert sum(1 for d in deck.diagnostics if d.rule.startswith("css-")) >= 2
