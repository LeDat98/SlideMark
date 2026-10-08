"""Wave 2026-10-08 lane D: ``cover.art`` (a motif of native shapes on the cover and the closing slide).

Layout (inside the slide, never over the title block, deterministic by seed), render (the .pptx is reopened:
``Cover art N`` names, the accent node, opacity as alpha), tokens (checked with a hint), importer (the shapes
are dropped, the kind is read back) and fuzz-safety."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import build
from slidemark.importer import import_pptx
from slidemark.layout.coverart import art
from slidemark.theme import Theme, normalize_token

HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#14272F primary=#0B2A3C secondary=#1B8A8F accent=#F2A33A muted=#5A6F78\n"
    "fonts: heading=Arial body=Arial\n"
    "sizes: title=36 heading=20 body=16 caption=12\n"
)
DECK = (
    "\n# AI: từ nền tảng đến ứng dụng\n@cover bg=primary dark\n## Giới thiệu cho người mới bắt đầu\n"
    "\n# Nội dung\n- Một\n- Hai\n"
    "\n# Cảm ơn bạn đã theo dõi\n@cover bg=primary dark\n## Câu hỏi và trao đổi\n"
)
E = 914400


def make(tmp_path: Path, style: str = "", text: str = DECK, name: str = "d"):
    head = HEAD + (f"style: {style}\n" if style else "")
    deck = build(head + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def arts(prs, i: int):
    return [s for s in prs.slides[i].shapes if re.fullmatch(r"Cover art \d+", s.name)]


def bad(deck):
    return [
        d for d in deck.diagnostics if d.level in ("warning", "error") and not d.rule.startswith("design")
    ]


# ---------------------------------------------------------------------------------------------- layout


def theme(**kw) -> Theme:
    return Theme(
        name="none", colors={"primary": "#0B2A3C", "secondary": "#1B8A8F", "accent": "#F2A33A"}, **kw
    )


@pytest.mark.parametrize("kind", ["network", "dots", "rings"])
def test_every_shape_is_inside_the_art_area_and_named_in_order(kind):
    th = theme(cover_art=kind)
    W, H = 12192000, 6858000
    x0 = round(W * th.cover_art_split)
    items = art(th, th.layout, W, H, x0)
    assert items
    for n, (shape, r, _st) in enumerate(items, 1):
        assert shape.attrs["shape_name"] == f"Cover art {n}"
        assert r.x >= x0 - 1 and r.x + r.w <= W and r.y >= 0 and r.y + r.h <= H, (kind, n, r)


def test_network_has_8_to_12_dots_joined_by_lines_and_one_accent_node():
    th = theme(cover_art="network")
    items = art(th, th.layout, 12192000, 6858000, 7315200)
    dots = [i for i in items if i[0].shape == "ellipse"]
    lines = [i for i in items if i[0].shape == "line"]
    assert 8 <= len(dots) <= 12 and len(lines) >= len(dots) - 1
    assert sum(1 for _s, _r, s in dots if s.fill == "accent") == 1
    assert all(s.fill == "secondary" for _s, _r, s in dots if s.fill != "accent")


def test_seed_is_deterministic_and_changes_the_motif():
    def sig(seed):
        th = theme(cover_art="network", cover_art_seed=seed)
        return [(s.shape, r.x, r.y, r.w, r.h) for s, r, _st in art(th, th.layout, 12192000, 6858000, 7315200)]

    assert sig(7) == sig(7) and sig(7) != sig(8)


def test_rings_are_three_outline_circles_and_dots_fade_towards_the_title():
    th = theme(cover_art="rings")
    items = art(th, th.layout, 12192000, 6858000, 7315200)
    rings = [i for i in items if i[2].fill is None]
    assert len(rings) == 3 and len({r.w for _s, r, _st in rings}) == 3
    assert all(r.w == r.h for _s, r, _st in rings)
    d = theme(cover_art="dots")
    discs = [i for i in art(d, d.layout, 12192000, 6858000, 7315200) if i[2].fill != "accent"]
    left = [s.opacity for _sh, r, s in discs if r.x < 8500000]
    right = [s.opacity for _sh, r, s in discs if r.x > 10500000]
    assert max(left) < max(right)


def test_none_and_a_narrow_area_draw_nothing():
    th = theme()
    assert art(th, th.layout, 12192000, 6858000, 7315200) == []
    n = theme(cover_art="network")
    assert art(n, n.layout, 12192000, 6858000, 11800000) == []


# ---------------------------------------------------------------------------------------------- render


def test_cover_and_closing_slide_get_the_motif_and_the_title_keeps_its_share(tmp_path):
    deck, prs = make(tmp_path, "cover.art=network")
    assert not bad(deck)
    first, mid, last = arts(prs, 0), arts(prs, 1), arts(prs, 2)
    assert len(first) >= 15 and not mid and len(last) == len(first)
    W = prs.slide_width
    for i in (0, 2):
        title = next(s for s in prs.slides[i].shapes if s.name == "Title")
        assert title.left + title.width <= W * 0.6 + 2000  # cover.art.split = 0.6
        for s in arts(prs, i):
            overlap = min(s.left + s.width, title.left + title.width) - max(s.left, title.left)
            assert overlap <= 0, s.name  # the motif never touches the title block


def test_accent_node_and_opacity_are_written_as_native_alpha(tmp_path):
    _, prs = make(tmp_path, "cover.art=network cover.art.opacity=40%")
    shapes = arts(prs, 0)
    accents = [s for s in shapes if s.shape_type is not None and s.name and "1B8A8F" not in _fill_xml(s)]
    ellipses = [
        s
        for s in shapes
        if s._element.spPr.find(qn("a:prstGeom")) is not None
        and s._element.spPr.find(qn("a:prstGeom")).get("prst") == "ellipse"
    ]
    assert any("F2A33A" in _fill_xml(s) for s in ellipses) and accents
    soft = [s for s in ellipses if "1B8A8F" in _fill_xml(s)]
    assert soft and all('alpha val="64000"' in _fill_xml(s) for s in soft)  # 0.4 x 1.6 (dots read stronger)
    lines = [s for s in shapes if s.shape_type is not None and s._element.tag.endswith("cxnSp")]
    assert lines and all('alpha val="40000"' in _ln_xml(s) for s in lines)


def _fill_xml(shape) -> str:
    from lxml import etree

    return etree.tostring(shape._element.spPr).decode()


def _ln_xml(shape) -> str:
    from lxml import etree

    ln = shape._element.spPr.find(qn("a:ln"))
    return etree.tostring(ln).decode() if ln is not None else ""


def test_rings_outline_carries_the_alpha(tmp_path):
    _, prs = make(tmp_path, "cover.art=rings cover.art.opacity=0.5")
    rings = [s for s in arts(prs, 0) if 'alpha val="50000"' in _ln_xml(s)]
    assert len(rings) == 3


def test_color_split_and_seed_tokens_take_effect(tmp_path):
    _, a = make(tmp_path, "cover.art=dots cover.art.color=#FF0000", name="a")
    assert any("FF0000" in _fill_xml(s) for s in arts(a, 0))
    _, wide = make(tmp_path, "cover.art=dots cover.art.split=0.8", name="w")
    title = next(s for s in wide.slides[0].shapes if s.name == "Title")
    assert title.left + title.width > wide.slide_width * 0.6
    _, s1 = make(tmp_path, "cover.art=network cover.art.seed=1", name="s1")
    _, s2 = make(tmp_path, "cover.art=network cover.art.seed=2", name="s2")
    assert [(s.left, s.top) for s in arts(s1, 0)] != [(s.left, s.top) for s in arts(s2, 0)]


def test_no_token_no_shapes_and_no_warning(tmp_path):
    deck, prs = make(tmp_path)
    assert not arts(prs, 0) and not bad(deck)


def test_token_without_a_cover_is_an_attr_ignored_warning(tmp_path):
    deck, _ = make(tmp_path, "cover.art=network", text="\n# Nội dung\n- Một\n- Hai\n")
    assert [d.rule for d in deck.diagnostics if d.rule == "attr-ignored"] == ["attr-ignored"]


def test_bad_values_are_checked_with_a_hint(tmp_path):
    deck, _ = make(tmp_path, "cover.art=spiral cover.art.opacity=3 cover.art.split=0.1 cover.art.seed=x")
    errs = [d for d in deck.diagnostics if d.rule == "bad-token"]
    assert len(errs) == 4 and all(d.hint for d in errs)
    assert normalize_token("cover_art", "off", None) == [("cover_art", "none")]
    with pytest.raises(Exception, match="spiral"):
        normalize_token("cover_art", "spiral", None)


# ---------------------------------------------------------------------------------------------- importer


def test_importer_drops_the_motif_and_reads_the_kind(tmp_path):
    from slidemark.importer.read import ReadCtx, read_slide
    from slidemark.importer.vocab import cover_art_tokens

    make(tmp_path, "cover.art=network")
    text, _diags = import_pptx(tmp_path / "d.pptx")
    assert "Cover art" not in text and "AI: từ nền tảng đến ứng dụng" in text
    assert "Giới thiệu cho người mới bắt đầu" in text  # the subtitle survives the dropped shapes
    for kind in ("network", "rings", "dots"):
        _, prs = make(tmp_path, f"cover.art={kind}", name=kind)
        first = read_slide(prs.slides[0], ReadCtx(accent="F2A33A"))
        assert cover_art_tokens(first) == {"cover.art": kind}
    _, plain = make(tmp_path, name="plain")
    assert cover_art_tokens(read_slide(plain.slides[0], ReadCtx(accent="F2A33A"))) == {}


# ---------------------------------------------------------------------------------------------- fuzz

_WORD = st.text(alphabet="aAbáàạăâđêôơưÀ 0123456789", min_size=0, max_size=30)


@settings(
    max_examples=25,
    deadline=None,
    derandomize=True,
    suppress_health_check=[HealthCheck.too_slow, HealthCheck.function_scoped_fixture],
)
@given(
    kind=st.sampled_from(["network", "dots", "rings", "none", "bogus"]),
    seed=st.sampled_from(["1", "99", "-3", "x", "2.5"]),
    split=st.sampled_from(["0.3", "0.6", "0.9", "1.5", "none"]),
    title=_WORD,
    sub=_WORD,
    lay=st.sampled_from(["@cover", "@cover bg=primary dark", "@section", ""]),
)
def test_fuzz_cover_art_never_raises(kind, seed, split, title, sub, lay, tmp_path_factory):
    text = f"\n# {title or 'T'}\n{lay}\n## {sub or 'S'}\n\n# Nội dung\n- a\n"
    out = tmp_path_factory.mktemp("fz") / "d.pptx"
    deck = build(
        HEAD + f"style: cover.art={kind} cover.art.seed={seed} cover.art.split={split}\n" + text, out
    )
    assert out.exists()
    for d in deck.diagnostics:
        if d.level in ("warning", "error"):
            assert d.hint or d.rule.startswith("design"), str(d)
