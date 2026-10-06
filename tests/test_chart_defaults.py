"""Chart defaults of design wave 2: accent-free palettes, pie share labels, legend size, line label
collisions and the `chart-scale` info line (info, not warning: it must not cost an agent a rebuild).
Numbers are tokens; the .pptx is reopened."""

from __future__ import annotations

from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import build
from slidemark.ir import Chart, Deck, Paragraph, Run, Series, Slide, Text
from slidemark.layout import layout_slide
from slidemark.layout.chartnote import scale_warning
from slidemark.parser import parse
from slidemark.render import render
from slidemark.render.axis import label_collisions, pie_label_pt, pie_percent
from slidemark.template import deck_theme
from slidemark.theme import RenderTokens, get_theme


def chart_of(tmp_path, kind, cats, series, theme="jp-business", **opts):
    title = Text(role="title", paragraphs=[Paragraph(runs=[Run(text="T")])])
    ch = Chart(kind=kind, categories=cats, series=[Series(name=n, values=v) for n, v in series], options=opts)
    d = Deck(slides=[Slide(title=title, elements=[ch])], theme=theme)
    th = get_theme(theme)
    out = tmp_path / "c.pptx"
    render(d, [layout_slide(s, d, th, i) for i, s in enumerate(d.slides)], th, out)
    prs = Presentation(str(out))
    return [s for s in prs.slides[0].shapes if s.name.startswith("Chart")][0].chart, th


PIE = ("pie", ["a", "b", "c", "d"], [("share", [36, 25, 21, 18])])


def _flags(el):
    return {t: el.find(qn(f"c:{t}")).get("val") for t in ("showVal", "showPercent")}


def _dlbls(chart, si=0):
    return chart.plots[0].series[si]._element.find(qn("c:dLbls"))


# --------------------------------------------------------------------------- 1. palettes


def test_presets_keep_the_accent_out_of_the_series_palette():
    for name in ("jp-business", "default", "midnight"):
        th = get_theme(name)
        accent = th.hexval("accent").lstrip("#")
        pal = th.chart_palette(None)
        assert accent not in pal, name
        assert len(set(pal)) == len(pal) >= 5, name


def test_hl_on_a_chart_still_uses_the_accent(tmp_path):
    chart, th = chart_of(tmp_path, *PIE, hl="b")
    pt = chart.plots[0].series[0].points[1]
    assert str(pt.format.fill.fore_color.rgb) == th.hexval("accent").lstrip("#")


def test_third_series_is_not_the_accent(tmp_path):
    chart, th = chart_of(
        tmp_path, "line", ["x", "y"], [("a", [1, 2]), ("b", [2, 3]), ("c", [3, 4])], labels="on"
    )
    cols = [str(s.format.line.color.rgb) for s in chart.plots[0].series]
    assert th.hexval("accent").lstrip("#") not in cols and len(set(cols)) == 3


# --------------------------------------------------------------------------- 2. pie labels and legend


def test_pie_labels_on_show_the_share_once(tmp_path):
    chart, th = chart_of(tmp_path, *PIE, labels="on")
    dls = _dlbls(chart)
    assert _flags(dls) == {"showVal": "0", "showPercent": "1"}
    assert dls.find(qn("c:numFmt")).get("formatCode") == "0%"
    want = round(pie_label_pt(chart.font.size.pt, th.render) * 100)
    points = dls.findall(qn("c:dLbl"))
    assert len(points) == 4
    for dl in points:
        assert _flags(dl) == {"showVal": "0", "showPercent": "1"}
        rpr = dl.find(qn("c:txPr")).find(".//" + qn("a:defRPr"))
        assert rpr.get("b") == "1" and int(rpr.get("sz")) == want
        assert dl.find(qn("c:dLblPos")).get("val") == "inEnd"
    assert want > round(chart.font.size.pt * th.render.chart_label_scale * 100)


def test_pie_labels_value_percent_and_fmt_keep_working(tmp_path):
    chart, _ = chart_of(tmp_path, *PIE, labels="value")
    assert _flags(_dlbls(chart))["showVal"] == "1"
    chart, _ = chart_of(tmp_path, *PIE, labels="percent")
    assert _flags(_dlbls(chart))["showPercent"] == "1"
    chart, _ = chart_of(tmp_path, *PIE, labels="on", fmt='0"%"')  # values already percentages: raw + format
    dls = _dlbls(chart)
    assert _flags(dls) == {"showVal": "1", "showPercent": "0"}
    assert dls.find(qn("c:numFmt")).get("formatCode") == '0"%"'


def test_pie_percent_rule_and_token():
    rt = RenderTokens()
    assert pie_percent("on", None, rt) and not pie_percent("on", '0"%"', rt)
    assert not pie_percent("on", None, RenderTokens(chart_pie_labels="value"))
    assert pie_percent("percent", None, RenderTokens(chart_pie_labels="value"))


def test_small_wedge_label_goes_outside(tmp_path):
    chart, _ = chart_of(tmp_path, "pie", ["a", "b", "c"], [("s", [90, 7, 3])], labels="on")
    pos = [d.find(qn("c:dLblPos")).get("val") for d in _dlbls(chart).findall(qn("c:dLbl"))]
    assert pos == ["inEnd", "inEnd", "outEnd"]


def test_legend_is_larger_than_the_chart_text(tmp_path):
    chart, th = chart_of(tmp_path, *PIE, labels="on")
    assert abs(chart.legend.font.size.pt - chart.font.size.pt * th.render.chart_legend_scale) < 0.02
    assert th.render.chart_legend_scale > 1


def test_chart_defaults_are_overridable_tokens():
    deck = parse("style: render.chart_legend_scale=2 render.chart_collide_em=0\n\n# T\n")
    th, _ = deck_theme(deck, ".")
    assert th.render.chart_legend_scale == 2 and th.render.chart_collide_em == 0


# --------------------------------------------------------------------------- 3. line label collisions


LINE = [("タイ", [70, 82, 95, 105]), ("北米", [20, 30, 45, 60]), ("台湾", [40, 50, 60, 55])]


def _positions(chart):
    """{series: {point index: position}} of the per-point label overrides."""
    out = {}
    for si in range(len(chart.plots[0].series)):
        dls = _dlbls(chart, si)
        out[si] = (
            {
                int(d.find(qn("c:idx")).get("val")): d.find(qn("c:dLblPos")).get("val")
                for d in dls.findall(qn("c:dLbl"))
            }
            if dls is not None
            else {}
        )
    return out


def test_close_line_ends_alternate_above_and_below(tmp_path):
    chart, _ = chart_of(tmp_path, "line", ["24", "25", "26", "27"], LINE, labels="on", legend="bottom")
    pos = _positions(chart)
    assert pos[2] == {3: "b"}  # 台湾 55 goes under 北米 60 at 2027, only there
    assert pos[0] == {} and pos[1] == {}  # the higher label and the far line keep the plot default (above)
    assert _dlbls(chart, 2).find(qn("c:dLblPos")).get("val") == "t"  # the series default stays above


def test_far_lines_keep_every_label_above(tmp_path):
    chart, _ = chart_of(
        tmp_path, "line", ["a", "b"], [("x", [10, 20]), ("y", [60, 90])], labels="on", legend="bottom"
    )
    assert _positions(chart) == {0: {}, 1: {}}


def test_label_collisions_unit_and_token_off():
    rt = RenderTokens()
    vals = [[100, 100], [20, 60], [40, 55]]  # chart 400pt tall, labels 14pt: 60 vs 55 collide, 40 does not
    assert label_collisions(vals, 0, 120, 14, 400, rt) == {2: [1]}
    assert label_collisions(vals, 0, 120, 14, 400, RenderTokens(chart_collide_em=0)) == {}
    assert label_collisions([[1, 2]], 0, 3, 14, 400, rt) == {}
    assert label_collisions([[10], [90]], 0, 100, 14, 400, rt) == {}
    equal = label_collisions([[5], [5]], 0, 10, 14, 400, rt)
    assert equal == {1: [0]}  # equal values: the later series goes below


# --------------------------------------------------------------------------- 4. chart-scale


def _chart(series, kind="column"):
    return Chart(kind=kind, categories=["a", "b"], series=[Series(name=n, values=v) for n, v in series])


def test_scale_warning_names_the_crushed_series():
    th = get_theme("jp-business")
    msg, hint = scale_warning(_chart([("売上高", [1020, 1280]), ("営業利益", [61, 96])]), th)
    assert "'営業利益'" in msg and "'売上高'" in msg and "10%" in msg
    assert "two charts" in hint


def test_scale_warning_quiet_cases():
    th = get_theme("jp-business")
    assert scale_warning(_chart([("a", [100, 120]), ("b", [60, 90])]), th) is None
    assert scale_warning(_chart([("a", [1000, 1200]), ("b", [10, 90])], "stacked-column"), th) is None
    assert scale_warning(_chart([("a", [1000, 1200])]), th) is None
    off = th.model_copy(update={"render": th.render.model_copy(update={"chart_scale_ratio": 0})})
    assert scale_warning(_chart([("a", [1000, 1200]), ("b", [1, 2])]), off) is None


def test_scale_warning_reaches_the_build_diagnostics(tmp_path):
    md = "# T\n\n```column {labels=on}\n,2026,2027\n売上高,1020,1280\n営業利益,61,96\n```\n"
    deck = build(md, tmp_path / "s.pptx")
    got = [d for d in deck.diagnostics if d.rule == "chart-scale"]
    assert len(got) == 1 and got[0].slide == 1 and got[0].hint
