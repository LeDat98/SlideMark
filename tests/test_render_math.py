"""LaTeX -> OMML: native equations in a text box (reopened .pptx)."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from slidemark.ir import Deck, Paragraph, Raw, Run, Slide, Text
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.render.math import A14_NS, M_NS, latex_to_omml
from slidemark.theme import get_theme

NS = {"m": M_NS, "a14": A14_NS}
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def build(tmp_path: Path, src: str) -> tuple[Path, Deck]:
    title = Text(role="title", paragraphs=[Paragraph(runs=[Run(text="Math")])])
    d = Deck(slides=[Slide(title=title, elements=[Raw(kind="math", source=src)])])
    th = get_theme("default")
    out = tmp_path / "m.pptx"
    render(d, [layout_slide(s, d, th, i) for i, s in enumerate(d.slides)], th, out)
    return out, d


def omath(src: str):
    return etree.fromstring(f'<r xmlns:m="{M_NS}" xmlns:a="{A_NS}">{latex_to_omml(src)}</r>')


def text_of(t) -> str:
    return "".join(x.text for x in t.iter(f"{{{M_NS}}}t"))


@pytest.mark.parametrize(
    "src,tag",
    [
        (r"\frac{a}{b}", "m:f"),
        (r"x^2", "m:sSup"),
        (r"x_i", "m:sSub"),
        (r"x_i^2", "m:sSubSup"),
        (r"\sqrt{x}", "m:rad"),
        (r"\sqrt[3]{x}", "m:rad"),
        (r"\sum_{i=1}^{n} i^2", "m:nary"),
        (r"\int_0^1 f(x) dx", "m:nary"),
        (r"\prod_{k} k", "m:nary"),
        (r"\left( a + b \right)", "m:d"),
    ],
)
def test_structures(src, tag):
    assert omath(src).findall(f".//{tag}", NS)


def test_symbols_text_and_lines():
    t = omath(r"\alpha \times \beta \leq \infty \rightarrow \text{ok} \\ \pm \neq \approx \cdot \mathrm{d}")
    assert len(t.findall(".//m:oMath", NS)) == 2
    for ch in "αβ×≤∞→±≠≈⋅":
        assert ch in text_of(t)
    assert "ok" in text_of(t)


def test_sqrt_degree_and_nary_char():
    t = omath(r"\sqrt[3]{x} + \sum_{i=1}^n i")
    assert "3" in "".join(t.find(".//m:rad/m:deg", NS).itertext())
    assert t.find(".//m:nary/m:naryPr/m:chr", NS).get(f"{{{M_NS}}}val") == "∑"


@pytest.mark.parametrize(
    "src",
    ["\\foo{x}", "{{{", "}}}", "x^", "\\", "\\left(", "\\frac{1", "a_", "\\right)", "$$", "\\sqrt[", "é"],
)
def test_odd_input_never_raises(src):
    latex_to_omml(src)


def test_unknown_command_is_literal():
    assert "\\foo" in text_of(omath(r"\foo"))


def test_equation_in_deck_is_native(tmp_path):
    out, d = build(tmp_path, r"E = mc^2 \\ \frac{a}{b} = \sqrt{x}")
    prs = Presentation(str(out))
    root = prs.slides[0].shapes._spTree
    xml = etree.tostring(root).decode()
    assert "AlternateContent" in xml and 'Requires="a14"' in xml and "Fallback" in xml
    assert "[math]" not in xml
    assert root.findall(".//m:sSup", NS) and root.findall(".//m:f", NS) and root.findall(".//m:rad", NS)
    assert not [x for x in d.diagnostics if x.rule == "raw"]


def test_empty_equation_falls_back_to_placeholder(tmp_path):
    out, _ = build(tmp_path, "  ")
    prs = Presentation(str(out))
    assert "[math]" in etree.tostring(prs.slides[0].shapes._spTree).decode()


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_math_deck_opens_in_libreoffice(tmp_path):
    out, _ = build(tmp_path, r"\sum_{i=1}^{n} i = \frac{n(n+1)}{2}")
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(out)],
        capture_output=True,
        timeout=180,
    )
    assert (tmp_path / "m.pdf").exists(), r.stderr
