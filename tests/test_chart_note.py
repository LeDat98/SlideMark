"""Chart takeaways: ``hl=`` (emphasised points) and ``note=`` (native callout + pointer)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import build
from slidemark.importer import import_pptx
from slidemark.layout import layout_slide
from slidemark.lint import lint
from slidemark.parser import parse
from slidemark.theme import get_theme

CATS = ",駅前小型,都市型SM,郊外SM,郊外大型\n"


def fence(kind, attrs, body=None):
    body = body or CATS + "営業利益率,4.8,3.1,1.2,-2.6"
    return f"# T\n\n```{kind} {{{attrs}}}\n{body}\n```\n"


def build_md(tmp_path, md, name="d"):
    src = tmp_path / f"{name}.md"
    src.write_text(md, encoding="utf-8")
    out = tmp_path / f"{name}.pptx"
    deck = build(src, out)
    return deck, Presentation(str(out))


def chart_of(prs):
    return next(s for s in prs.slides[0].shapes if s.has_chart)


def points(chart):
    """{series index: {point index: fill hex}} of the explicit data points."""
    out = {}
    for si, ser in enumerate(chart.plots[0].series):
        for dpt in ser._element.findall(qn("c:dPt")):
            idx = int(dpt.find(qn("c:idx")).get("val"))
            clr = dpt.find(".//" + qn("a:srgbClr"))
            out.setdefault(si, {})[idx] = clr.get("val") if clr is not None else None
    return out


def accent_hex():
    return get_theme("default").hexval("accent").lstrip("#")


def shapes_named(prs, prefix):
    return [s for s in prs.slides[0].shapes if s.name.startswith(prefix)]


def the_note(prs):
    (note,) = [s for s in shapes_named(prs, "ChartNote") if not s.name.startswith("ChartNoteLine")]
    return note


def inside(inner, outer, tol=1000):
    return (
        inner.left >= outer.left - tol
        and inner.top >= outer.top - tol
        and inner.left + inner.width <= outer.left + outer.width + tol
        and inner.top + inner.height <= outer.top + outer.height + tol
    )


def manual_layout(chart):
    man = chart._chartSpace.find(".//" + qn("c:plotArea") + "/" + qn("c:layout") + "/" + qn("c:manualLayout"))
    assert man is not None
    assert man.find(qn("c:layoutTarget")).get("val") == "inner"
    return {t: float(man.find(qn(f"c:{t}")).get("val")) for t in "xywh"}


# --------------------------------------------------------------------------- parser


def test_parser_reads_hl_and_note():
    deck = parse(fence("bar", 'hl=郊外大型,駅前小型 note="赤字の半数は郊外"'))
    ch = deck.slides[0].elements[0]
    assert ch.options["hl"] == ["郊外大型", "駅前小型"]
    assert ch.options["note"] == "赤字の半数は郊外"
    assert not [d for d in deck.diagnostics if d.level != "info"]


def test_parser_unknown_hl_category_warns_with_valid_list():
    deck = parse(fence("bar", "hl=nope,郊外SM"))
    ch = deck.slides[0].elements[0]
    assert ch.options["hl"] == ["郊外SM"]
    d = next(d for d in deck.diagnostics if d.rule == "chart-hl")
    assert d.level == "warning" and "郊外大型" in d.hint


def test_parser_hl_is_case_insensitive_and_pie_rows():
    ch = parse(fence("column", "hl=q2", ",Q1,Q2\na,1,2")).slides[0].elements[0]
    assert ch.options["hl"] == ["Q2"]
    pie = parse(fence("pie", "hl=APAC", ",v\nAPAC,48\nEMEA,30")).slides[0].elements[0]
    assert pie.options["hl"] == ["APAC"]


@pytest.mark.parametrize("kind", ["scatter", "radar", "area"])
def test_parser_hl_unsupported_kind_warns(kind):
    deck = parse(fence(kind, "hl=郊外SM"))
    assert any(d.rule == "chart-hl" for d in deck.diagnostics)
    assert "hl" not in deck.slides[0].elements[0].options


def test_parser_empty_values_do_not_raise():
    deck = parse(fence("bar", 'hl="" note=""'))
    assert any(d.rule == "bad-chart-option" for d in deck.diagnostics)


# --------------------------------------------------------------------------- hl points


def test_hl_single_series_bar_colors_only_the_point(tmp_path):
    _, prs = build_md(tmp_path, fence("bar", "hl=郊外大型"))
    assert points(chart_of(prs).chart) == {0: {3: accent_hex()}}


def test_hl_token_overrides_color(tmp_path):
    md = "style: render.chart_hl=#123456\n\n" + fence("column", "hl=郊外SM")
    _, prs = build_md(tmp_path, md)
    assert points(chart_of(prs).chart) == {0: {2: "123456"}}


def test_hl_multi_series_outlines_every_series(tmp_path):
    body = CATS + "a,1,2,3,4\nb,2,3,4,5"
    _, prs = build_md(tmp_path, fence("column", "hl=郊外SM", body))
    for ser in chart_of(prs).chart.plots[0].series:
        dpt = ser._element.find(qn("c:dPt"))
        assert dpt.find(qn("c:idx")).get("val") == "2"
        assert dpt.find(".//" + qn("a:ln") + "//" + qn("a:srgbClr")).get("val") == accent_hex()
        fill = dpt.find(qn("c:spPr")).find(qn("a:solidFill") + "/" + qn("a:srgbClr")).get("val")
        assert fill != accent_hex()  # keeps its series color (the outline carries the emphasis)


def test_hl_line_marker_and_pie_slice(tmp_path):
    _, prs = build_md(tmp_path, fence("line", "hl=郊外SM", CATS + "a,1,2,3,4"))
    dpt = chart_of(prs).chart.plots[0].series[0]._element.find(qn("c:dPt"))
    assert dpt.find(qn("c:marker")).find(qn("c:size")).get("val") == "11"
    _, prs = build_md(tmp_path, fence("pie", "hl=APAC", ",v\nAPAC,48\nEMEA,30"), "pie")
    assert points(chart_of(prs).chart)[0][0] == accent_hex()


def test_hl_waterfall_colors_bar_and_keeps_others(tmp_path):
    body = ",2026,再編,PB,2029\n営業,44,18,-9,="
    _, prs = build_md(tmp_path, fence("waterfall", "hl=PB", body))
    pts = points(chart_of(prs).chart)
    assert [s for s in pts.values() if accent_hex() in s.values()] == [{2: accent_hex()}]  # the PB bar only
    assert all(i == 2 for s in pts.values() for i in s)


def test_hl_names_the_chart_shape(tmp_path):
    _, prs = build_md(tmp_path, fence("bar", "hl=郊外大型,駅前小型"))
    assert chart_of(prs).name.endswith(" hl=郊外大型,駅前小型")


# --------------------------------------------------------------------------- note: shapes, geometry


def test_note_is_native_shape_inside_chart_frame(tmp_path):
    deck, prs = build_md(tmp_path, fence("bar", 'hl=郊外大型 note="赤字42店の半数は郊外大型"'))
    note = the_note(prs)
    assert note.text_frame.text.replace("\u2060", "") == "赤字42店の半数は郊外大型"
    assert inside(note, chart_of(prs))
    assert not [d for d in deck.diagnostics if d.rule == "chart-note"]
    assert note.fill.type is not None  # tinted callout, not a bare text box


def test_note_pointer_hits_the_hl_bar(tmp_path):
    _, prs = build_md(tmp_path, fence("bar", 'hl=郊外大型 note="赤字の半数は郊外"'))
    gf = chart_of(prs)
    lay = manual_layout(gf.chart)
    va = gf.chart.value_axis
    lo, hi = va.minimum_scale, va.maximum_scale
    line = shapes_named(prs, "ChartNoteLine")[0]
    ex, ey = line.end_x, line.end_y
    px, pw = gf.left + lay["x"] * gf.width, lay["w"] * gf.width
    py, ph = gf.top + lay["y"] * gf.height, lay["h"] * gf.height
    row = 3  # 郊外大型, the last of 4 rows (first on top)
    assert py + row * ph / 4 < ey < py + (row + 1) * ph / 4
    zero_x = px + pw * (0 - lo) / (hi - lo)
    assert abs(ex - zero_x) < 0.1 * pw  # the negative bar ends at the zero line, facing the note


def _bar_rects(prs, vals):
    """Absolute (left, top, right, bottom) EMU of the bars of a horizontal bar chart, first on top."""
    gf = chart_of(prs)
    lay = manual_layout(gf.chart)
    va = gf.chart.value_axis
    lo, hi = va.minimum_scale, va.maximum_scale
    px, pw = gf.left + lay["x"] * gf.width, lay["w"] * gf.width
    py, ph = gf.top + lay["y"] * gf.height, lay["h"] * gf.height
    n = len(vals)
    out = []
    for i, v in enumerate(vals):
        cy, half = py + (i + 0.5) * ph / n, ph / n * 0.3
        x0, x1 = sorted((px + pw * (0 - lo) / (hi - lo), px + pw * (v - lo) / (hi - lo)))
        out.append((x0, cy - half, x1, cy + half))
    return out


LAUNCH = ",North,Coast,Valley,Metro\nVisits,64,108,72,166"


def test_note_pointer_ends_on_the_longest_horizontal_bar(tmp_path):
    attrs = 'labels=on hl=Metro note="Metro alone: 40% of visits" title="Expected visits per year (k)"'
    _, prs = build_md(tmp_path, fence("bar", attrs, LAUNCH))
    line = shapes_named(prs, "ChartNoteLine")[0]
    ex, ey = line.end_x, line.end_y
    x0, y0, x1, y1 = _bar_rects(prs, [64, 108, 72, 166])[3]
    tol = 4 * 12700
    assert x0 - tol <= ex <= x1 + tol, "the pointer ends beside the value label, not on the bar"
    assert y0 - tol <= ey <= y1 + tol
    assert ex < x1  # lands inside the bar's end, not past it


def test_note_beside_horizontal_bars_covers_no_bar(tmp_path):
    attrs = 'labels=on hl=Metro note="Metro alone: 40% of visits" title="Expected visits per year (k)"'
    _, prs = build_md(tmp_path, fence("bar", attrs, LAUNCH))
    note = the_note(prs)
    for i, (x0, y0, x1, y1) in enumerate(_bar_rects(prs, [64, 108, 72, 166])):
        hit = note.left < x1 and x0 < note.left + note.width and note.top < y1 and y0 < note.top + note.height
        assert not hit, f"note covers bar {i}"
    assert note.top + note.height <= _bar_rects(prs, [64, 108, 72, 166])[3][1]  # above the highlighted bar


def test_note_pointer_of_the_top_horizontal_bar_comes_from_below(tmp_path):
    body = ",Metro,North,Coast\nVisits,166,64,108"
    _, prs = build_md(tmp_path, fence("bar", 'labels=on hl=Metro note="Metro alone: 40%"', body))
    line = shapes_named(prs, "ChartNoteLine")[0]
    x0, y0, x1, y1 = _bar_rects(prs, [166, 64, 108])[0]
    tol = 4 * 12700
    assert x0 - tol <= line.end_x <= x1 + tol and y0 - tol <= line.end_y <= y1 + tol


def test_note_column_pointer_lands_above_the_bar(tmp_path):
    body = ",Q1,Q2,Q3,Q4\n売上,10,40,25,18"
    _, prs = build_md(tmp_path, fence("column", 'hl=Q2 note="Q2 is the peak"', body))
    gf = chart_of(prs)
    lay = manual_layout(gf.chart)
    va = gf.chart.value_axis
    line = shapes_named(prs, "ChartNoteLine")[0]
    ex, ey = line.end_x, line.end_y
    px, pw = gf.left + lay["x"] * gf.width, lay["w"] * gf.width
    py, ph = gf.top + lay["y"] * gf.height, lay["h"] * gf.height
    assert abs(ex - (px + 1.5 * pw / 4)) < 0.02 * pw  # centre of the second column
    bar_top = py + ph * (1 - (40 - va.minimum_scale) / (va.maximum_scale - va.minimum_scale))
    assert bar_top - 0.15 * ph < ey <= bar_top  # just above the bar (its label sits there)


def test_note_does_not_cover_the_columns(tmp_path):
    vals = [10, 40, 25, 18, 31]
    body = ",Q1,Q2,Q3,Q4,Q5\n売上," + ",".join(map(str, vals))
    _, prs = build_md(tmp_path, fence("column", 'hl=Q2 note="Q2 is the peak of the year" labels=on', body))
    gf, note = chart_of(prs), the_note(prs)
    lay = manual_layout(gf.chart)
    va = gf.chart.value_axis
    py, ph = gf.top + lay["y"] * gf.height, lay["h"] * gf.height
    px, pw = gf.left + lay["x"] * gf.width, lay["w"] * gf.width
    for i, v in enumerate(vals):
        cx0, cx1 = px + i * pw / 5, px + (i + 1) * pw / 5
        top = py + ph * (1 - (v - va.minimum_scale) / (va.maximum_scale - va.minimum_scale))
        overlap_x = note.left < cx1 and cx0 < note.left + note.width
        assert not (overlap_x and note.top + note.height > top), f"note covers column {i}"


def test_note_without_hl_has_no_pointer(tmp_path):
    _, prs = build_md(tmp_path, fence("column", 'note="Only a takeaway"', ",A,B\nx,1,2"))
    the_note(prs)
    assert not shapes_named(prs, "ChartNoteLine")
    manual_layout(chart_of(prs).chart)


def test_chart_without_note_keeps_auto_layout(tmp_path):
    _, prs = build_md(tmp_path, fence("bar", "hl=郊外大型"))
    area = chart_of(prs).chart._chartSpace.find(".//" + qn("c:plotArea") + "/" + qn("c:layout"))
    assert area is None or area.find(qn("c:manualLayout")) is None
    assert not shapes_named(prs, "ChartNote")


def test_long_note_shrinks_then_wraps_then_warns(tmp_path):
    _, prs = build_md(tmp_path, fence("bar", f'hl=郊外SM note="{"x" * 60}"'))
    the_note(prs)
    deck, prs = build_md(tmp_path, fence("bar", f'hl=郊外SM note="{"あ" * 200}"'), "huge")
    assert any(d.rule == "chart-note" for d in deck.diagnostics)
    assert inside(the_note(prs), chart_of(prs))


@pytest.mark.parametrize(
    "kind", ["bar", "column", "stacked-bar", "stacked-column", "line", "area", "waterfall"]
)
def test_note_on_every_plotted_kind(tmp_path, kind):
    body = ",a,b,c\nx,3,5,-2" if kind in ("waterfall", "line") else ",a,b,c\nx,3,5,2\ny,1,2,3"
    attrs = 'hl=b note="takeaway" labels=on' if kind != "area" else 'note="takeaway" labels=on'
    _, prs = build_md(tmp_path, fence(kind, attrs, body), kind)
    assert inside(the_note(prs), chart_of(prs))


@pytest.mark.parametrize("kind", ["pie", "doughnut", "scatter", "radar"])
def test_note_on_unplotted_kinds_still_draws(tmp_path, kind):
    _, prs = build_md(tmp_path, fence(kind, 'note="takeaway"', ",a,b\nx,3,5"), kind)
    assert inside(the_note(prs), chart_of(prs))


def test_fuzz_never_raises(tmp_path):
    cases = [
        ('hl="" note=""', ",a,b\nx,1,2"),
        ("hl=zz note=y", ",a,b\nx,1,2"),
        ("hl=a,a,a note=n", ",a,b\nx,,"),
        ('hl=a note="n" min=5 max=1', ",a,b\nx,1,2"),
        ('hl=a note="ねーむ" legend=right title="t" axis=off', ",a,b\nx,1,2\ny,3,4"),
        ('hl=a note="n"', ",a\nx,1"),
        ('hl=a note="n"', ",a,b\nx,,"),
    ]
    for i, (attrs, body) in enumerate(cases):
        for kind in ("bar", "column", "line", "waterfall", "pie", "stacked-column"):
            build_md(tmp_path, fence(kind, attrs, body), f"f{i}{kind}")


def test_lint_has_no_overlap_between_note_and_chart(tmp_path):
    deck, _ = build_md(tmp_path, fence("bar", 'hl=郊外大型 note="赤字の半数は郊外"'))
    th = get_theme("default")
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    assert not [d for d in lint(deck, placed, th) if d.rule == "overlap"]


# --------------------------------------------------------------------------- importer round trip


def test_import_keeps_hl_and_note(tmp_path):
    _, prs = build_md(tmp_path, fence("bar", 'hl=郊外大型 note="赤字42店の半数は郊外大型"'))
    text, _ = import_pptx(tmp_path / "d.pptx")
    assert "hl=郊外大型" in text and "note=赤字42店の半数は郊外大型" in text
    assert text.count("赤字42店の半数は郊外大型") == 1  # the callout is not also emitted as content
    _, prs2 = build_md(tmp_path, text, "again")
    the_note(prs2)
    assert shapes_named(prs2, "ChartNoteLine")
    assert points(chart_of(prs2).chart) == points(chart_of(prs).chart)


def test_import_line_note_does_not_invent_min_max(tmp_path):
    build_md(tmp_path, fence("line", 'hl=b note="peak"', ",a,b,c\nx,3,5,4"))
    text, _ = import_pptx(tmp_path / "d.pptx")
    assert "min=" not in text and "max=" not in text and "hl=b" in text


def _note_pt(note) -> float:
    return min(int(r.get("sz")) for r in note._element.iter(qn("a:rPr")) if r.get("sz")) / 100


def test_note_is_never_smaller_than_the_axis_labels(tmp_path):
    md = (
        "# Demand\n\n@1:1\n"
        '```bar {title="Visits (k)" labels=on hl=Metro note="Metro alone: 40% of visits"}\n'
        ",North,Coast,Valley,Metro\nVisits,64,108,72,166\n```\n## Why\n- a\n- b\n"
    )
    deck, prs = build_md(tmp_path, md)
    note = the_note(prs)
    axis = {
        int(e.get("sz")) / 100 for e in chart_of(prs).chart._chartSpace.iter(qn("a:defRPr")) if e.get("sz")
    }
    assert _note_pt(note) >= max(axis)
    assert inside(note, chart_of(prs))
