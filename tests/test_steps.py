"""``@steps``: an arrow row (headings) with an outcome card under every arrow, in the same columns."""

from __future__ import annotations

import pytest
from pptx import Presentation

from slidemark import build
from slidemark.importer import import_pptx
from slidemark.parser import parse

DECK = """# 導入ステップと成果
> 3か月で全店に展開する
@4 steps head=arrow
## 現状分析
- 全店の業務を棚卸し
- 課題を30件に整理
## 設計
- 標準フローを策定
## 試行
## 全店展開
- 全120店へ順次導入
- 研修を月2回実施
※ 出所: 社内PMO
"""


def _shapes(md: str, tmp_path, name: str = "a.pptx"):
    out = tmp_path / name
    build(md, out)
    return {s.name: s for s in Presentation(str(out)).slides[0].shapes}


def test_parser_makes_one_steps_group_from_the_boxes():
    deck = parse(DECK)
    sl = deck.slides[0]
    assert deck.diagnostics == [] and "steps" not in sl.classes
    (g,) = sl.elements
    assert g.classes == ["plain", "group", "steps"] and g.grid == "4"
    assert [b.title.paragraphs[0].plain for b in g.children] == ["現状分析", "設計", "試行", "全店展開"]
    assert sl.conclusion is None and len(sl.footnotes) == 1


def test_blocks_after_the_steps_stay_below_and_ratios_are_kept():
    md = "# T\n@1:2:1 steps\n## A\n- a\n## B\n- b\n## C\n- c\n@end\n```table\nx,y\n1,2\n```\n"
    sl = parse(md).slides[0]
    assert [e.type for e in sl.elements] == ["container", "table"] and sl.grid == "1x2"
    assert sl.elements[0].grid == "1:2:1"


def test_too_few_steps_warns_and_never_raises():
    for md in ("# T\n@steps\n## Only\n- a\n", "# T\n@steps\nplain text\n", "# T\n@3 steps\n"):
        deck = parse(md)
        assert [d.rule for d in deck.diagnostics] == ["steps-few"], md
        assert "## Heading" in deck.diagnostics[0].hint


def test_unusable_grid_is_replaced_with_a_note():
    deck = parse("# T\n@2x2 steps\n## A\n- a\n## B\n- b\n## C\n- c\n")
    assert [d.rule for d in deck.diagnostics] == ["steps-grid"]
    assert deck.slides[0].elements[0].grid == "3"


def test_cards_sit_under_their_arrows_in_the_same_columns(tmp_path):
    sh = _shapes(DECK, tmp_path)
    arrows = [sh[f"Step {i} arrow"] for i in range(1, 5)]
    cards = {i: sh[f"Step {i} card"] for i in (1, 2, 4)}
    assert "Step 3 card" not in sh  # a step without bullets only shows its arrow
    assert len({a.top for a in arrows}) == 1 and len({a.height for a in arrows}) == 1
    for i, card in cards.items():
        a = arrows[i - 1]
        assert (card.left, card.width) == (a.left, a.width)
        assert card.top >= a.top + a.height  # under the arrow ...
        assert card.top - (a.top + a.height) < 914400 // 2  # ... right under it
    assert len({c.height for c in cards.values()}) == 1  # cards of a row share one height
    foot = next(s for s in sh.values() if s.name.startswith("Footnote"))
    assert max(c.top + c.height for c in cards.values()) <= foot.top
    assert arrows[0].top < cards[1].top < cards[1].top + cards[1].height


def test_cards_grow_down_to_the_conclusion_bar_when_there_is_text(tmp_path):
    md = (
        "# Plan\n> Outcome per stage\n@3 steps\n"
        + "".join(
            f"## Stage {i}\n- first outcome of {i}\n- second outcome of {i}\n- third one\n" for i in (1, 2, 3)
        )
        + "@end\n> Result: 20% effort saved\n"
    )
    sh = _shapes(md, tmp_path)
    card, bar = sh["Step 1 card"], sh["Conclusion"]
    assert bar.top - (card.top + card.height) < 914400 * 0.7  # no dead band between the cards and the bar
    assert card.height > 914400 * 1.5


def test_arrow_text_is_not_smaller_than_the_card_text(tmp_path):
    sh = _shapes(DECK, tmp_path)

    def size(shape) -> float:
        runs = [r for p in shape.text_frame.paragraphs for r in p.runs if r.text.strip()]
        return max(r.font.size.pt for r in runs if r.font.size)

    assert size(sh["Step 1 arrow"]) >= max(size(sh[n]) for n in sh if n.startswith("Text"))


def test_style_tokens_change_the_arrows_and_cards(tmp_path):
    head = "style: steps-arrow.fill=#FF5A1F steps-card.fill=#EEF2FF layout.steps_gap=0.4in\n\n"
    sh = _shapes(head + DECK, tmp_path)
    assert str(sh["Step 1 arrow"].fill.fore_color.rgb) == "FF5A1F"
    assert str(sh["Step 1 card"].fill.fore_color.rgb) == "EEF2FF"
    a, c = sh["Step 1 arrow"], sh["Step 1 card"]
    assert abs((c.top - (a.top + a.height)) - 914400 * 0.4) < 914400 * 0.02


def test_dense_japanese_five_steps_fit_without_warnings(tmp_path):
    step = "## {i}. 市場調査と顧客インタビュー\n- 顧客インタビュー 20社を実施\n- 競合 8社の機能・価格を比較\n"
    md = "theme: jp-business\nlang: ja\n\n# 新規事業の立ち上げプロセスと成果物\n@5 steps\n"
    md += "".join(step.format(i=i) for i in range(1, 6))
    deck = parse(md)
    out = tmp_path / "a.pptx"
    build(md, out)
    assert [d for d in deck.diagnostics if d.level != "info"] == []
    sh = {s.name: s for s in Presentation(str(out)).slides[0].shapes}
    right = max(s.left + s.width for n, s in sh.items() if n.endswith(("arrow", "card")))
    assert right <= Presentation(str(out)).slide_width


def test_round_trip_gives_steps_back(tmp_path):
    a = tmp_path / "a.pptx"
    build(DECK, a)
    text, _ = import_pptx(a, tmp_path)
    assert "@4 steps" in text and "## 試行" in text and "- 研修を月2回実施" in text
    b = tmp_path / "b.pptx"
    build(text, b)
    n1 = {s.name for s in Presentation(str(a)).slides[0].shapes}
    n2 = {s.name for s in Presentation(str(b)).slides[0].shapes}
    assert {n for n in n1 if n.startswith("Step")} == {n for n in n2 if n.startswith("Step")}


@pytest.mark.parametrize(
    "md",
    [
        "# T\n@steps\n## A\n## B\n",
        "# T\n@steps\n## A\n- x\n## B\n- y\n@end\n@2\n## C\n- z\n## D\n- w\n",
        "# T\n@steps chevron\n## A {.danger icon=check}\n- x\n## B {.kpi}\n1\n",
        "# T\n@2 steps\n## A\n### sub\n- x\n## B\n```mermaid\ngraph LR\nA-->B\n```\n",
        "# T\n@9 steps\n## A\n- x\n## B\n- y\n",
        "# T\n@steps a>b\n## A\n- x\n## B\n- y\n",
    ],
)
def test_fuzz_never_raises(md, tmp_path):
    parse(md)
    build(md, tmp_path / "f.pptx")


SPARSE = """theme: jp-business
lang: ja

# 年間スケジュール
@4 steps head=arrow
## 4–6月
- 価格改定
## 7–9月
- タイ工場着工
## 10–12月
- 北米発売
## 1–3月
- 効果検証
"""


def _max_pt(shape) -> float:
    return max(r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.text.strip())


def test_sparse_steps_use_the_body(tmp_path):
    out = tmp_path / "a.pptx"
    build(SPARSE, out)
    prs = Presentation(str(out))
    sh = {s.name: s for s in prs.slides[0].shapes}
    arrows = [sh[f"Step {i} arrow"] for i in range(1, 5)]
    cards = [sh[f"Step {i} card"] for i in range(1, 5)]
    texts = [s for n, s in sh.items() if n.startswith("Text")]
    assert min(_max_pt(t) for t in texts) >= 16  # 11pt theme text grew to the ceiling (layout.grow_max 1.5 x)
    assert min(_max_pt(a) for a in arrows) >= 18  # the label grew and stayed on one line
    assert all(len(a.text_frame.text.splitlines()) == 1 for a in arrows)
    span = max(c.top + c.height for c in cards) - min(a.top for a in arrows)
    body = prs.slide_height - sh["Title"].top - sh["Title"].height
    assert span >= 0.5 * body  # the group covers most of the body, not a strip at the top
    top_gap = min(a.top for a in arrows) - (sh["Title"].top + sh["Title"].height)
    bottom_gap = prs.slide_height - max(c.top + c.height for c in cards)
    assert top_gap < bottom_gap  # centred a little above the middle
    for a, c in zip(arrows, cards, strict=True):
        assert (a.left, a.width) == (c.left, c.width)


def test_sparse_composition_is_a_token(tmp_path):
    out = tmp_path / "a.pptx"
    build("style: layout.steps_sparse=off\n" + SPARSE, out)
    sh = {s.name: s for s in Presentation(str(out)).slides[0].shapes}
    assert _max_pt(sh["Text 1"]) <= 14


# --- wave 2 lane D: the card look (`steps.head=card`, the default)

CARD_HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#14272F primary=#0B2A3C secondary=#1B8A8F accent=#F2A33A surface=#EEF4F6 "
    "muted=#5A6F78\n"
    "fonts: heading=Arial body=Arial\n"
    "sizes: title=34 heading=20 body=16 caption=12\n"
)
CARDS = """
# Mô hình viết từng token
@4 steps num
## Token hóa {icon=code}
- Cắt văn bản thành các token
- Một từ có thể thành vài token
## Embedding {icon=database}
- Mỗi token thành một vector số
## Attention {icon=search}
- Cân nhắc token nào liên quan
- Hiểu ngữ cảnh của cả câu
## Dự đoán {icon=bolt}
@end
> Một token mỗi lần
"""
E = 914400


def _card_deck(tmp_path, style: str = "", text: str = CARDS, name: str = "c.pptx"):
    head = CARD_HEAD + (f"style: {style}\n" if style else "")
    out = tmp_path / name
    deck = build(head + text, out)
    prs = Presentation(str(out))
    return deck, prs, {s.name: s for s in prs.slides[0].shapes}


def _kids(prs, sh, card):
    """The shapes of the slide inside ``card`` (a name-keyed dict would lose the repeated ``Icon disc``)."""
    c = sh[card]
    return sorted(
        (
            s
            for s in prs.slides[0].shapes
            if s is not c
            and s.left >= c.left - 5
            and s.top >= c.top - 5
            and s.left + s.width <= c.left + c.width + 5
            and s.top + s.height <= c.top + c.height + 5
        ),
        key=lambda s: s.top,
    )


def test_the_heading_moves_into_the_card_and_the_arrow_shows_the_number(tmp_path):
    deck, _prs, sh = _card_deck(tmp_path, 'icon.disc=secondary steps.caption="Bước {n}"')
    assert [
        d for d in deck.diagnostics if d.level in ("warning", "error") and not d.rule.startswith("design")
    ] == []
    for i, word in enumerate(("Token hóa", "Embedding", "Attention", "Dự đoán"), 1):
        assert sh[f"Step {i} arrow"].text_frame.text == str(i)
        head = sh[f"Heading {i}"]
        assert head.text_frame.paragraphs[0].text == f"Bước {i}"
        assert head.text_frame.paragraphs[1].text == word
        assert f"Step {i} card" in sh  # (the last step has no bullets: its card still holds the heading)


def test_card_content_is_top_anchored_in_reading_order_without_rules(tmp_path):
    _d, prs, sh = _card_deck(tmp_path, 'icon.disc=secondary steps.caption="Bước {n}"')
    assert not [s for s in prs.slides[0].shapes if s.name == "Rule"]  # no spread rules between the paragraphs
    for i in (1, 2, 3):
        card = sh[f"Step {i} card"]
        kids = _kids(prs, sh, f"Step {i} card")
        assert {k.name for k in kids} >= {"Icon disc", f"Heading {i}"}
        disc = next(k for k in kids if k.name == "Icon disc")
        head, text = sh[f"Heading {i}"], sh[f"Text {i}"]
        assert card.top < disc.top < head.top < text.top  # disc, then caption + heading, then the body
        assert text.top + text.height <= card.top + card.height
        assert disc.top - card.top < 0.4 * E  # top-anchored: the content starts at the card's top padding
    heights = {sh[f"Step {i} card"].height for i in (1, 2, 3, 4)}
    assert len(heights) == 1  # one card height per row
    bar = sh["Conclusion"]
    assert (
        bar.top - (sh["Step 1 card"].top + sh["Step 1 card"].height) < 0.7 * E
    )  # the free height went to the cards


def test_text_never_shrinks_below_the_body_size_and_may_grow(tmp_path):
    _d, _p, sh = _card_deck(tmp_path)

    def size(shape):
        return max(
            r.font.size.pt
            for p in shape.text_frame.paragraphs
            for r in p.runs
            if r.font.size and r.text.strip()
        )

    for i in (1, 2, 3):
        assert size(sh[f"Text {i}"]) >= 16 and size(sh[f"Heading {i}"]) >= 20
        assert size(sh[f"Text {i}"]) <= 16 * 1.5 + 0.01  # layout.grow_max
    assert len({size(sh[f"Text {i}"]) for i in (1, 2, 3)}) == 1  # one growth factor for the row


def test_disc_sits_in_the_card_only_with_icon_disc_else_the_glyph_stays_in_the_arrow(tmp_path):
    _d, _p, with_disc = _card_deck(tmp_path, "icon.disc=secondary", name="a.pptx")
    _d, _p, bare = _card_deck(tmp_path, name="b.pptx")
    card, disc = with_disc["Step 1 card"], with_disc["Icon disc"]
    assert card.top <= disc.top and disc.top + disc.height <= card.top + card.height
    arrow = bare["Step 1 arrow"]
    glyph = next(s for n, s in bare.items() if n.startswith("icon "))
    assert (
        arrow.left < glyph.left < arrow.left + arrow.width
        and arrow.top <= glyph.top < arrow.top + arrow.height
    )
    assert arrow.text_frame.text == "1"  # (the number is next to the glyph)


def test_head_arrow_is_the_old_look_by_token_and_by_slide(tmp_path):
    _d, _p, tok = _card_deck(tmp_path, "steps.head=arrow", name="t.pptx")
    _d, _p, one = _card_deck(
        tmp_path, "", CARDS.replace("@4 steps num", "@4 steps head=arrow"), name="s.pptx"
    )
    for sh in (tok, one):
        assert sh["Step 1 arrow"].text_frame.text == "Token hóa"
        assert "Step 4 card" not in sh  # a step without bullets shows only its arrow
        assert "Heading 1" not in sh
    _d, _p, mixed = _card_deck(
        tmp_path, "steps.head=arrow", CARDS.replace("@4 steps num", "@4 steps head=card"), name="m.pptx"
    )
    assert mixed["Step 1 arrow"].text_frame.text == "1"  # the slide wins over the deck token


def test_chevron_row_turned_into_steps_keeps_its_headings_in_the_arrows(tmp_path):
    text = "\n# Lộ trình\n@4 chevron\n" + "".join(f"## Giai đoạn {i}\n- việc số {i}\n" for i in range(1, 5))
    _d, _p, sh = _card_deck(tmp_path, text=text)
    assert "Heading 1" not in sh  # no card heading: the author wrote the heading in the arrow
    heads = [s for s in sh.values() if s.has_text_frame and s.text_frame.text.startswith("Giai đoạn")]
    assert len(heads) == 4 and all(
        "chevron" in s._element.xml or "homePlate" in s._element.xml for s in heads
    )


def test_card_align_and_valign_tokens(tmp_path):
    _d, _p, sh = _card_deck(tmp_path, "steps-card.align=center", name="c.pptx")
    assert sh["Heading 1"].text_frame.paragraphs[1].alignment is not None
    from pptx.enum.text import PP_ALIGN

    assert sh["Heading 1"].text_frame.paragraphs[1].alignment == PP_ALIGN.CENTER
    _d, _p, top = _card_deck(tmp_path, name="t.pptx")
    _d, _p, mid = _card_deck(tmp_path, "steps-card.valign=middle", name="m.pptx")
    assert mid["Heading 2"].top >= top["Heading 2"].top  # the block moved down in its card, never up
    assert (
        mid["Heading 2"].top + mid["Heading 2"].height <= mid["Step 2 card"].top + mid["Step 2 card"].height
    )


def test_steps_head_token_and_key_are_checked(tmp_path):
    deck, _p, _sh = _card_deck(tmp_path, "steps.head=middle")
    bad = [d for d in deck.diagnostics if d.rule == "bad-token"]
    assert len(bad) == 1 and bad[0].hint
    deck2 = parse("# T\n@3 steps head=wat\n## A\n- a\n## B\n- b\n## C\n- c\n")
    assert [d.rule for d in deck2.diagnostics] == ["bad-attr"] and deck2.diagnostics[0].hint
    deck3 = parse("# T\n@3 steps head=arrow\n## A\n- a\n## B\n- b\n## C\n- c\n")
    assert deck3.diagnostics == []


def test_card_round_trip_gives_headings_icons_and_head_back(tmp_path):
    a = tmp_path / "a.pptx"
    build(CARD_HEAD + "style: icon.disc=secondary\n" + CARDS, a)
    text, _ = import_pptx(a, tmp_path)
    assert "@4 steps" in text and "head=arrow" not in text
    assert (
        "## Token hóa {icon=code}" in text
        and "- Cắt văn bản thành các token" in text
        and "## Dự đoán" in text
    )
    b = tmp_path / "b.pptx"
    build(text, b)

    def names(p):
        return {
            s.name for s in Presentation(str(p)).slides[0].shapes if s.name.startswith(("Step", "Heading"))
        }

    assert names(a) == names(b)
    c = tmp_path / "c.pptx"  # the old look keeps its word
    build(CARD_HEAD + CARDS.replace("@4 steps num", "@4 steps head=arrow"), c)
    text2, _ = import_pptx(c, tmp_path)
    assert (
        "head=arrow"
        in text2.splitlines()[
            text2.splitlines().index(next(x for x in text2.splitlines() if x.startswith("@4 steps")))
        ]
    )


@pytest.mark.parametrize(
    "md",
    [
        "# T\n@steps\n## A {icon=nope}\n- x\n## B\n- y\n",
        "# T\n@2 steps\n## A\n### sub\n- x\n## B\n```mermaid\ngraph LR\nA-->B\n```\n",
        "# T\n@3 steps\n## A\n| a | b |\n|-|-|\n| 1 | 2 |\n## B\n- y\n## C\n- z\n",
        "# T\n@6 steps\n"
        + "".join(
            f"## S{i} {{icon=bolt}}\n- a long line of text that wraps over several lines {i}\n"
            for i in range(6)
        ),
        "# T\n@2 steps\n## A {size=40}\n- x\n## B\n- y\n@end\n> bar\n^ note\n",
        "# T\n@3 steps\n## 現状分析\n- 全店の業務を棚卸し\n## 設計\n## 試行\n- x\n",
    ],
)
def test_card_look_fuzz_never_raises(md, tmp_path):
    for head in ("", 'style: icon.disc=primary steps.caption="Bước {n}" steps-card.valign=bottom\n'):
        deck = build(CARD_HEAD + head + md, tmp_path / "f.pptx")
        for d in deck.diagnostics:
            if d.level in ("warning", "error"):
                assert d.hint or d.rule.startswith("design"), str(d)


def test_caption_on_a_heading_band_keeps_the_band_ink_and_bodies_line_up(tmp_path):
    md = (
        "theme: jp-business\nlang: ja\n\n# 年間スケジュール\n@4 steps num\n"
        "## 4–6月と、少し長い見出しが二行になる場合\n- 価格改定\n## 7–9月\n- タイ工場着工\n"
        "## 10–12月\n- 北米発売\n## 1–3月\n- 効果検証\n"
    )
    out = tmp_path / "j.pptx"
    deck = build(md, out)
    assert [d for d in deck.diagnostics if d.rule == "contrast"] == []  # STEP n readable on the band
    sh = {s.name: s for s in Presentation(str(out)).slides[0].shapes}
    assert len({sh[f"Heading {i}"].height for i in (1, 2, 3, 4)}) == 1  # the tallest heading sets the row
    assert len({sh[f"Text {i}"].top for i in (1, 2, 3, 4)}) == 1  # so the bodies start at the same y


def test_pinned_card_geometry_with_the_card_look(tmp_path):
    _d, prs, sh = _card_deck(
        tmp_path, "steps-card.h=3.6in steps-arrow.h=0.5in steps.gap=0.2in icon.disc=secondary"
    )
    cards = [sh[f"Step {i} card"] for i in (1, 2, 3, 4)]
    assert {c.height for c in cards} == {round(3.6 * E)}
    for i in (1, 2, 3):
        card, head, text = sh[f"Step {i} card"], sh[f"Heading {i}"], sh[f"Text {i}"]
        assert card.top <= head.top and text.top + text.height <= card.top + card.height
    assert max(c.top + c.height for c in cards) <= prs.slide_height
