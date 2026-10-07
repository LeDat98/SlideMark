"""DL3b part 2 (lane K): ``@iconlist`` ``@quote`` ``@split`` ``@proscons`` ``@progress``
``@harvey`` ``@heatmap`` ``@pins``. Every form: parser (the word and its secondary attributes reach
the IR, a bad key warns), tokens (``<word>.*`` accepted, a bad value is a ``bad-token`` warning),
render (the .pptx is reopened and its named shapes asserted), importer round trip (named shapes fold
back to the same word), fuzz-safety (a bad input is a diagnostic with a hint, never an exception).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from slidemark import build, parse
from slidemark.importer import import_pptx

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "examples"
HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#222B36 primary=#142B4D secondary=#1F5FA8 accent=#E09F1F success=#2E9E6B "
    "danger=#D1495B muted=#59626E surface=#EEF2F7 border=#C9D2DE\n"
    "fonts: heading=Arial body=Arial\n"
)
MAP = "![map](assets/region-map.png)"


def make(tmp_path, text: str, name: str = "d", head: str = HEAD):
    deck = build(head + text, tmp_path / f"{name}.pptx", base_dir=ASSETS)
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def by_name(prs, i, pattern: str = "", exact: bool = False):
    sh = prs.slides[i].shapes
    return [s for s in sh if (s.name == pattern if exact else s.name.startswith(pattern))]


def one(prs, i, name: str):
    got = by_name(prs, i, name, exact=True)
    assert len(got) == 1, (name, [s.name for s in prs.slides[i].shapes])
    return got[0]


def txt(shp) -> str:
    """The text of a shape with the layout's orphan binding (U+00A0) and joiners (U+2060) normalised."""
    return shp.text_frame.text.replace("\u00a0", " ").replace("\u2060", "")


def fill_hex(shp) -> str:
    return str(shp.fill.fore_color.rgb)


def rules(deck, rule: str):
    return [d for d in deck.diagnostics if d.rule == rule]


def bad(deck):
    """Warnings other than the advisory design ones."""
    return [d for d in deck.diagnostics if d.level == "warning" and not d.rule.startswith("design-")]


def imported(tmp_path, name: str = "d") -> str:
    text, _ = import_pptx(tmp_path / f"{name}.pptx")
    return text


# --------------------------------------------------------------------------- parser


@pytest.mark.parametrize(
    ("line", "cls", "attrs"),
    [
        ("@iconlist cols=2", "iconlist", {"cols": "2"}),
        ("@quote align=center size=36", "quote", {"align": "center", "size": "36"}),
        ("@split side=right ratio=2:3 bleed=on fit=contain", "split", {"side": "right", "ratio": "2:3"}),
        ("@proscons fill=surface", "proscons", {"fill": "surface"}),
        ("@pros", "pros", {}),
        ("@progress max=10", "progress", {"max": "10"}),
        ("@harvey size=20", "harvey", {"size": "20"}),
        ("@heatmap min=0 max=100 colors=#F3F6FA,primary text=off", "heatmap", {"colors": "#F3F6FA,primary"}),
        ("@pins legend=bottom", "pins", {"legend": "bottom"}),
    ],
)
def test_word_and_secondary_attributes_reach_the_slide(line, cls, attrs):
    deck = parse(f"# T\n{line}\n- a\n")
    s = deck.slides[0]
    assert cls in s.classes and attrs.items() <= s.attrs.items()
    assert not [d for d in deck.diagnostics if d.rule in ("unknown-token", "bad-grid")]


def test_a_key_the_form_does_not_take_warns_with_the_valid_keys():
    deck = parse("# T\n@quote cols=2\nhello\n")
    (d,) = rules(deck, "unknown-token")
    assert "cols" in d.message and "align" in d.hint
    deck = parse("# T\n@cols=2\n- a\n")  # no form to act on
    assert "no form" in rules(deck, "unknown-token")[0].message


# --------------------------------------------------------------------------- tokens


def test_tokens_are_accepted_and_bad_values_warn():
    ok = (
        "style: iconlist.icon.size=0.6in iconlist.icon.color=accent iconlist.gap=0.3in quote.size=40 "
        "quote.mark.color=primary split.gap=0.4in proscons.plus.color=teal progress.h=0.2in "
        "progress.label.size=20 harvey.size=0.5in harvey.line=none heatmap.colors=#F3F6FA,primary "
        "heatmap.text=off pins.fill=accent pins.size=0.4in pins.legend=bottom\n"
    )
    deck = parse(HEAD + ok + "\n# T\n@quote\nx\n")
    assert not [d for d in deck.diagnostics if d.rule in ("bad-token", "unknown-token")], deck.diagnostics
    deck = parse(
        HEAD
        + "style: pins.legend=up heatmap.text=maybe heatmap.colors=red harvey.fill=notacolor1 "
        + "split.ratio=big\n"
        + "\n# T\n@quote\nx\n"
    )
    assert len(rules(deck, "bad-token")) == 5
    assert all(d.hint for d in rules(deck, "bad-token"))


def test_style_tokens_without_their_form_are_flagged(tmp_path):
    deck, _ = make(tmp_path, "style: quote.size=40 pins.fill=accent\n\n# A\n- x\n")
    assert len(rules(deck, "attr-ignored")) == 2
    deck, _ = make(tmp_path, "style: quote.size=40\n\n# A\n@quote\nhello\n", "q")
    assert not rules(deck, "attr-ignored")


# --------------------------------------------------------------------------- @iconlist

ICONS = (
    "\n# Why\n> reasons\n@iconlist cols=2\n"
    "- icon=bolt **Faster** median 38 s\n- :shield: **Safer** signed\n"
    "- icon=users **Shared** one board\n- icon=chart **Honest** numbers\n"
)


def test_iconlist_draws_icon_title_and_text_in_columns(tmp_path):
    deck, prs = make(tmp_path, "style: iconlist.icon.size=0.6in iconlist.icon.color=accent\n" + ICONS)
    texts = [s for s in by_name(prs, 0, "Iconlist ") if s.name.count(" ") == 1]
    assert len(texts) == 4
    assert len({t.left for t in texts}) == 2 and len({t.top for t in texts}) == 2  # 2 columns x 2 rows
    icons = by_name(prs, 0, "icon ")
    assert sorted(i.name for i in icons) == ["icon bolt", "icon chart", "icon shield", "icon users"]
    assert all(i.width == i.height == Emu(int(0.6 * 914400)) for i in icons)
    assert all(fill_hex(i) == "E09F1F" for i in icons)  # iconlist.icon.color
    first = texts[0].text_frame.paragraphs
    assert first[0].runs[0].font.bold and first[0].text == "Faster" and "38" in first[1].text
    assert not bad(deck)


def test_iconlist_cols_gap_fill_and_bad_values(tmp_path):
    _, one_col = make(tmp_path, ICONS.replace("cols=2", "cols=1"), "c1")
    texts = [s for s in by_name(one_col, 0, "Iconlist ") if s.name.count(" ") == 1]
    assert len({t.left for t in texts}) == 1
    _, carded = make(tmp_path, ICONS.replace("cols=2", "cols=2 fill=surface"), "fill")
    assert all(fill_hex(s) == "EEF2F7" for s in by_name(carded, 0, "Iconlist ") if s.name.count(" ") == 1)
    deck, _ = make(tmp_path, ICONS.replace("cols=2", "cols=9 size=big") + "- icon=nope **X** y\n", "bad")
    assert {d.rule for d in deck.diagnostics} >= {"bad-attr", "unknown-icon"}
    assert any("did you mean" in d.hint or "slidemark docs" in d.hint for d in rules(deck, "unknown-icon"))


def test_iconlist_size_is_pinned_and_gap_is_the_gutter(tmp_path):
    _, small = make(tmp_path, ICONS.replace("cols=2", "cols=2 size=16 gap=1in"), "s")
    texts = [s for s in by_name(small, 0, "Iconlist ") if s.name.count(" ") == 1]
    sizes = {r.font.size.pt for t in texts for p in t.text_frame.paragraphs[:1] for r in p.runs}
    assert sizes == {16.0}
    xs = sorted({t.left for t in texts})
    assert xs[1] - xs[0] > Emu(914400)  # the column pitch holds the 1in gutter


def test_iconlist_round_trip(tmp_path):
    make(tmp_path, ICONS, "rt")
    text = imported(tmp_path, "rt")
    assert "@iconlist cols=2" in text
    assert "- icon=bolt **Faster** median 38 s" in text and "- icon=shield **Safer** signed" in text


# --------------------------------------------------------------------------- @quote

QUOTE = '\n# Q\n@quote align=center\n> "We stopped arguing about the deck."\n> — Mika Tanaka, Northwind\n'


def test_quote_splits_the_attribution_and_centres(tmp_path):
    deck, prs = make(tmp_path, "style: quote.mark.color=primary quote.size=36\n" + QUOTE)
    text, by, mark = (one(prs, 0, n) for n in ("Quote text", "Quote by", "Quote mark"))
    assert txt(text).startswith('"We stopped')
    assert txt(by) == "— Mika Tanaka, Northwind"
    assert str(mark.text_frame.paragraphs[0].runs[0].font.color.rgb) == "142B4D"
    assert text.text_frame.paragraphs[0].runs[0].font.size.pt == 36.0  # quote.size
    assert mark.top < text.top < by.top
    assert abs((text.left + text.width / 2) - prs.slide_width / 2) < Emu(914400)  # align=center
    assert not bad(deck)


def test_quote_align_size_fill_and_plain_body_text(tmp_path):
    _, prs = make(tmp_path, "\n# Q\n@quote align=right size=30 fill=surface\nBe yourself.\n— Oscar\n")
    text = one(prs, 0, "Quote text")
    assert text.text_frame.paragraphs[0].alignment is not None
    assert text.text_frame.paragraphs[0].runs[0].font.size.pt == 30.0
    assert fill_hex(one(prs, 0, "Quote panel")) == "EEF2F7"
    assert txt(one(prs, 0, "Quote by")) == "— Oscar"
    deck, _ = make(tmp_path, "\n# Q\n@quote align=weird\nx\n", "bad")
    assert rules(deck, "bad-attr")


def test_quote_round_trip(tmp_path):
    make(tmp_path, QUOTE, "rt")
    text = imported(tmp_path, "rt")
    assert "@quote align=center" in text and "> — Mika Tanaka, Northwind" in text
    assert '> "We stopped arguing about the deck."' in text


# --------------------------------------------------------------------------- @split

SPLIT = "\n# S\n@split side=right ratio=1:1 bleed=on\n> A line\n![studio](assets/studio.png)\n- one\n- two\n"


def test_split_bleed_image_runs_to_the_edge_and_text_sits_on_the_other_side(tmp_path):
    deck, prs = make(tmp_path, SPLIT)
    pic = by_name(prs, 0, "Image")[0]
    W, H = prs.slide_width, prs.slide_height
    assert (
        pic.top == 0 and pic.height == H and pic.left + pic.width == W and abs(pic.width - W / 2) < Emu(20000)
    )
    text_shapes = [s for s in prs.slides[0].shapes if s.has_text_frame and txt(s).strip()]
    assert text_shapes and all(s.left + s.width <= pic.left + Emu(1) for s in text_shapes)
    assert any(txt(s) == "S" for s in text_shapes)  # the title is on the text side
    assert not bad(deck)


def test_split_inset_left_ratio_and_fill_block(tmp_path):
    _, prs = make(tmp_path, SPLIT.replace("side=right ratio=1:1 bleed=on", "ratio=2:1"), "in")
    pic = by_name(prs, 0, "Image")[0]
    assert pic.left > 0 and pic.top > 0 and pic.top + pic.height < prs.slide_height  # inset by the margins
    assert pic.width > prs.slide_width * 0.45  # 2:1 gives the image the larger share
    text = [s for s in prs.slides[0].shapes if s.has_text_frame and txt(s).startswith("one")][0]
    assert text.left > pic.left + pic.width
    _, block = make(tmp_path, "\n# S\n@split fill=primary side=right\n- a\n- b\n", "fill")
    blk = one(block, 0, "Split fill")
    assert fill_hex(blk) == "142B4D" and blk.left > block.slide_width / 2 - Emu(914400)
    assert not by_name(block, 0, "Image")


def test_split_bad_attributes_warn_and_never_crash(tmp_path):
    deck, _ = make(tmp_path, f"\n# S\n@split side=up ratio=big bleed=maybe fit=zzz\n{MAP}\n- a\n")
    assert len(rules(deck, "bad-attr")) == 4
    deck, _ = make(tmp_path, "\n# S\n@split\n- a\n")  # no image: a colour block and an info line
    assert any(d.rule == "form-empty" for d in deck.diagnostics)


def test_split_round_trip(tmp_path):
    make(tmp_path, SPLIT, "rt")
    text = imported(tmp_path, "rt")
    assert "@split side=right ratio=1:1 bleed=on" in text and "- one" in text and "> A line" in text


# --------------------------------------------------------------------------- @proscons

PROS = (
    "\n# P\n@proscons\n## Why buy\n- Live in eight weeks\n- Tax rules kept current\n"
    "## Why build\n- Nine months of work\n− We own every outage\n> Verdict: buy now\n"
)


def test_proscons_cards_heads_glyphs_and_verdict(tmp_path):
    deck, prs = make(tmp_path, "style: proscons.plus.color=accent proscons.minus.color=primary\n" + PROS)
    c1, c2 = one(prs, 0, "Proscons 1"), one(prs, 0, "Proscons 2")
    assert c1.left < c2.left and c1.width == c2.width and c1.top == c2.top
    assert fill_hex(one(prs, 0, "Proscons 1 head")) == "E09F1F"
    assert fill_hex(one(prs, 0, "Proscons 2 head")) == "142B4D"
    g1, g2 = one(prs, 0, "Proscons 1 item 1 glyph"), one(prs, 0, "Proscons 2 item 2 glyph")
    assert txt(g1) == "+" and txt(g2) == "−"
    assert fill_hex(g1) == "E09F1F" and fill_hex(g2) == "142B4D"
    assert txt(one(prs, 0, "Proscons 2 item 2")) == "We own every outage"  # the glyph is drawn
    assert any(s.name == "Conclusion" for s in prs.slides[0].shapes)
    assert not bad(deck)


def test_proscons_alias_and_fewer_boxes(tmp_path):
    _, prs = make(tmp_path, PROS.replace("@proscons", "@pros"), "alias")
    assert one(prs, 0, "Proscons 1")
    deck, prs = make(tmp_path, "\n# P\n@proscons\n## Only\n- a\n", "one")
    assert rules(deck, "form-empty") and not by_name(prs, 0, "Proscons")


def test_proscons_round_trip(tmp_path):
    make(tmp_path, PROS, "rt")
    text = imported(tmp_path, "rt")
    assert "@proscons" in text and "## Why buy" in text and "+ Live in eight weeks" in text
    assert (
        "- Nine months of work" in text and "- We own every outage" in text and "> Verdict: buy now" in text
    )


# --------------------------------------------------------------------------- @progress

PROG = "\n# P\n@progress max=100\n- Build 90%\n- Test 45%\n- Docs 9%\n"


def test_progress_bars_are_proportional_and_take_the_tokens(tmp_path):
    deck, prs = make(tmp_path, "style: progress.fill=accent progress.track=border progress.h=0.2in\n" + PROG)
    for n, frac in ((1, 0.9), (2, 0.45), (3, 0.09)):
        track, fill = one(prs, 0, f"Progress {n} track"), one(prs, 0, f"Progress {n} fill")
        assert abs(fill.width / track.width - frac) < 0.005 and fill.left == track.left
        assert fill.height == track.height == Emu(int(0.2 * 914400))
        assert fill_hex(fill) == "E09F1F" and fill_hex(track) == "C9D2DE"
    assert txt(one(prs, 0, "Progress 1 value")) == "90%"
    assert txt(one(prs, 0, "Progress 2 label")) == "Test"
    assert not bad(deck)


def test_progress_max_table_source_and_bad_rows(tmp_path):
    _, prs = make(tmp_path, "\n# P\n@progress max=10\n| item | v |\n|-|-|\n| a | 7 |\n| b | 12 |\n", "tab")
    t, f = one(prs, 0, "Progress 1 track"), one(prs, 0, "Progress 1 fill")
    assert abs(f.width / t.width - 0.7) < 0.005
    assert one(prs, 0, "Progress 2 fill").width == one(prs, 0, "Progress 2 track").width  # clamped at max
    deck, _ = make(tmp_path, "\n# P\n@progress max=0\n- no number\n- x 5%\n", "bad")
    assert {d.rule for d in deck.diagnostics} >= {"bad-attr", "progress-value"}


def test_progress_round_trip_recovers_max(tmp_path):
    make(tmp_path, "\n# P\n@progress max=10\n- a 7\n- b 3\n", "rt")
    text = imported(tmp_path, "rt")
    assert "@progress max=10" in text and "- a 7" in text and "- b 3" in text
    make(tmp_path, PROG, "pct")
    assert "@progress\n" in imported(tmp_path, "pct") and "- Build 90%" in imported(tmp_path, "pct")


# --------------------------------------------------------------------------- @harvey

HARVEY = (
    "\n# H\n@harvey\n| Vendor | Price | Support |\n|-|-|-|\n| Ledgerly | 4 | 2 |\n| Paybridge | 50% | n/a |\n"
)


def test_harvey_balls_are_native_quarter_shapes(tmp_path):
    deck, prs = make(tmp_path, "style: harvey.fill=accent harvey.line=primary harvey.size=0.5in\n" + HARVEY)
    full = one(prs, 0, "Harvey 2,2 q4")  # names count table rows and columns from 1 (the header is row 1)
    half = one(prs, 0, "Harvey 3,2 q2")
    assert fill_hex(full) == "E09F1F" and full.width == full.height == Emu(int(0.5 * 914400))
    assert str(half.line.color.rgb) == "142B4D"
    pie = one(prs, 0, "Harvey 3,2 pie")
    assert pie.adjustments[0] == pytest.approx(162.0) and pie.adjustments[1] == pytest.approx(
        54.0
    )  # 270 -> 90 deg
    assert (pie.left, pie.top, pie.width) == (half.left, half.top, half.width)
    assert not by_name(prs, 0, "Harvey 2,2 pie")  # a full ball needs no wedge
    assert txt(one(prs, 0, "Harvey row 3")) == "Paybridge"
    assert txt(one(prs, 0, "Harvey col 2")) == "Price"
    assert txt(one(prs, 0, "Harvey text 3,3")) == "n/a"
    assert rules(deck, "harvey-value")  # n/a is not 0-4: left as text, with a hint
    assert rules(deck, "harvey-value")[0].hint


def test_harvey_percent_values_and_round_trip(tmp_path):
    make(tmp_path, HARVEY, "rt")
    text = imported(tmp_path, "rt")
    assert "@harvey" in text and "| Ledgerly | 4 | 2 |" in text and "| Paybridge | 2 | n/a |" in text
    assert "| Vendor | Price | Support |" in text


# --------------------------------------------------------------------------- @heatmap

HEAT = (
    "\n# M\n@heatmap min=0 max=100 colors=#000000,#FFFFFF\n"
    "| R | A | B |\n|-|-|-|\n| x | 0 | 50 |\n| y | 100 | 25 |\n"
)


def cell(prs, r, c):
    tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
    return tbl.cell(r, c)


def test_heatmap_fills_interpolate_and_ink_follows(tmp_path):
    deck, prs = make(tmp_path, HEAT)
    assert str(cell(prs, 1, 1).fill.fore_color.rgb) == "000000"
    assert str(cell(prs, 2, 1).fill.fore_color.rgb) == "FFFFFF"
    assert str(cell(prs, 1, 2).fill.fore_color.rgb) in ("7F7F7F", "808080")  # 50 -> the middle
    assert str(cell(prs, 2, 2).fill.fore_color.rgb) == "404040"  # 25
    ink = lambda c: str(c.text_frame.paragraphs[0].runs[0].font.color.rgb)  # noqa: E731
    assert ink(cell(prs, 1, 1)) != ink(cell(prs, 2, 1))  # light on dark, dark on light
    assert cell(prs, 1, 0).fill.type is None or str(cell(prs, 1, 0).fill.fore_color.rgb) != "000000"
    assert not bad(deck)


def test_heatmap_scale_defaults_and_modes(tmp_path):
    _, prs = make(
        tmp_path,
        "style: heatmap.colors=#000000,#FFFFFF\n\n# M\n@heatmap\n| R | A |\n|-|-|\n| x | 10 |\n| y | 30 |\n",
        "d",
    )
    assert (
        str(cell(prs, 1, 1).fill.fore_color.rgb) == "000000"
        and str(cell(prs, 2, 1).fill.fore_color.rgb) == "FFFFFF"
    )
    _, off = make(tmp_path, HEAT.replace("colors=#000000,#FFFFFF", "text=off"), "off")
    assert txt(cell(off, 1, 2)).strip() == ""
    _, three = make(
        tmp_path, HEAT.replace("#000000,#FFFFFF", "#000000,#FF0000,#FFFFFF").replace("0 | 50", "0 | 50"), "3"
    )
    assert str(cell(three, 1, 2).fill.fore_color.rgb) == "FF0000"  # the middle stop at the middle value
    deck, _ = make(
        tmp_path,
        "\n# M\n@heatmap min=9 max=1 colors=red text=zzz\n| R | A |\n|-|-|\n| x | 3 |\n| y | 4 |\n",
        "bad",
    )
    assert len(rules(deck, "bad-attr")) == 3
    deck, _ = make(tmp_path, "\n# M\n@heatmap\n| R | A |\n|-|-|\n| x | y |\n", "nonum")
    assert rules(deck, "form-empty")


def test_heatmap_round_trip_keeps_the_scale(tmp_path):
    make(tmp_path, HEAT, "rt")
    text = imported(tmp_path, "rt")
    assert "@heatmap min=0 max=100 colors=#000000,#FFFFFF" in text
    assert "| x | 0 | 50 |" in text


# --------------------------------------------------------------------------- @pins

PINS = f"\n# P\n@pins legend=right\n{MAP}\n- x=25% y=40% **Harbor** HQ\n- x=75% y=80% **Delta** lab\n"


def test_pins_sit_at_the_image_positions_and_the_legend_follows(tmp_path):
    deck, prs = make(tmp_path, "style: pins.fill=primary pins.size=0.4in\n" + PINS)
    pic = by_name(prs, 0, "Image")[0]
    p1, p2 = one(prs, 0, "Pin 1"), one(prs, 0, "Pin 2")
    assert p1.width == p1.height == Emu(int(0.4 * 914400)) and fill_hex(p1) == "142B4D"
    cx = lambda s: s.left + s.width / 2  # noqa: E731
    cy = lambda s: s.top + s.height / 2  # noqa: E731
    assert (
        abs((cx(p1) - pic.left) / pic.width - 0.25) < 0.003
        and abs((cy(p1) - pic.top) / pic.height - 0.40) < 0.003
    )
    assert (
        abs((cx(p2) - pic.left) / pic.width - 0.75) < 0.003
        and abs((cy(p2) - pic.top) / pic.height - 0.80) < 0.003
    )
    leg = one(prs, 0, "Pins legend 1")
    assert leg.left > pic.left + pic.width and txt(leg).startswith("Harbor")
    assert txt(one(prs, 0, "Pin 2 legend")) == "2"
    assert not bad(deck)


def test_pins_legend_positions_and_bad_positions(tmp_path):
    _, bottom = make(tmp_path, PINS.replace("legend=right", "legend=bottom"), "b")
    pic = by_name(bottom, 0, "Image")[0]
    assert one(bottom, 0, "Pins legend 1").top >= pic.top + pic.height - Emu(10)
    _, off = make(tmp_path, PINS.replace("legend=right", "legend=off"), "o")
    assert not by_name(off, 0, "Pins legend") and one(off, 0, "Pin 1")
    deck, _ = make(
        tmp_path, PINS.replace("legend=right", "legend=sideways") + "- x=140% y=2% far\n- nothing\n", "bad"
    )
    assert {d.rule for d in deck.diagnostics} >= {"bad-attr", "pins-position"}


def test_pins_round_trip(tmp_path):
    make(tmp_path, PINS, "rt")
    text = imported(tmp_path, "rt")
    assert "@pins" in text and "- x=25% y=40% **Harbor** HQ" in text and "- x=75% y=80% **Delta** lab" in text
    make(tmp_path, PINS.replace("legend=right", "legend=bottom"), "rtb")
    assert "legend=bottom" in imported(tmp_path, "rtb")


# --------------------------------------------------------------------------- fit line, example, fuzz


def test_fit_lines_name_the_forms(tmp_path):
    md = ICONS + QUOTE + PROS + PROG + HARVEY + HEAT + PINS + SPLIT
    deck, _ = make(tmp_path, md)
    lines = deck.attrs["fit_map"]
    joined = "\n".join(lines)
    for word in (
        "icon list 4 items in 2 cols",
        "quote ",
        "pros / cons 2 boxes",
        "progress bars 3 bars",
        "harvey balls 3 balls",
        "heatmap 3x3, scale 0..100",
        "map pins 2 pins",
        "split right bleed",
    ):
        assert word in joined, (word, joined)


def test_example_24_builds_with_no_warnings(tmp_path):
    deck = build(ROOT / "examples" / "24-vocabulary-2.md", tmp_path / "x.pptx")
    assert not [d for d in deck.diagnostics if d.level == "warning"], [str(d) for d in deck.diagnostics]
    words = {c for s in deck.slides for c in s.classes}
    assert {"iconlist", "quote", "split", "proscons", "progress", "harvey", "heatmap", "pins"} <= words


FUZZ = [
    "@iconlist",
    "@iconlist\n| a | b |\n|-|-|\n| 1 | 2 |",
    "@iconlist cols=9 size=x fill=zz\n- icon=nope **A** b\n- :zz: c\n- plain",
    "@iconlist cols=3\n" + "".join(f"- icon=bolt **Item {i}** some words\n" for i in range(30)),
    "@quote",
    "@quote align=weird\n- a\n- b",
    "@quote\n> " + "long words " * 200 + "\n> — Someone",
    "@quote\n> — Someone",
    "@split",
    f"@split ratio=bad side=up bleed=maybe fit=zzz\n{MAP}\n{MAP}\n- a",
    "@split bleed=on\n![x](nope.png)\n- a",
    "@split fill=primary side=right\n- a\n## Box\n- c",
    "@proscons\n## A\n- x",
    "@pros\n## A\n- x\n## B\n- y\n## C\n- z",
    "@proscons\n## A\n## B",
    "@progress",
    "@progress max=0\n- no number\n- x 150%\n- y -3\n- z 7/0",
    "@progress\n" + "".join(f"- Row {i} {i * 3}%\n" for i in range(30)),
    "@harvey\n- a",
    "@harvey\n| a | b | c |\n|-|-|-|\n| x | 9 | n/a |\n| y | 50% | 3 |\n| z |  | 100 |",
    "@harvey\n| a |\n|-|\n| 1 |",
    "@heatmap\n- a",
    "@heatmap text=off fill=accent\n| a | b | c |\n|-|-|-|\n| x | 3 | 1,200 |\n| y | -4 | 12% |",
    "@heatmap colors=primary,accent,success\n| a | b |\n|-|-|\n| x | 3 |",
    "@pins\n- x=1% y=2% a",
    f"@pins\n{MAP}",
    f"@pins legend=sideways\n{MAP}\n- x=140% y=2% a\n- nothing\n- x=0.5 y=0.5 half",
    "@pins legend=bottom\n![m](nope.png)\n- x=10% y=10% a\n- x=50% y=50% b\n- x=90% y=90% c\n- x=20% y=80% d",
    "@iconlist dense dark\n- icon=bolt **A** b",
    "@quote\n> a\n> — b\n@end",
]


@pytest.mark.parametrize("body", FUZZ)
def test_fuzz_never_crashes_and_never_fails_the_layout(tmp_path, body):
    deck, prs = make(tmp_path, f"\n# T\n{body}\n")
    assert len(prs.slides) == 1
    assert not [d for d in deck.diagnostics if d.level == "error" or d.rule in ("layout-error", "form-error")]
