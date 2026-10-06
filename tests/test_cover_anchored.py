"""Anchored cover: token-driven band at the top, title block bottom-aligned in it, rule and footer caption."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from slidemark import build
from slidemark.ir import Shape, Text
from slidemark.layout import layout_slide
from slidemark.parser import parse
from slidemark.template import deck_theme

EX = Path(__file__).resolve().parent.parent / "examples"
H = 720 * 9525
W = 1280 * 9525


def _cover(head: str = "theme: jp-business\nfooter: ACME Corp\n", style: str = "", sub: str = "FY2027 board"):
    src = head + (f"style: {style}\n" if style else "") + f"\n# Logistics DX Plan\n{sub}\n"
    deck = parse(src)
    theme, diags = deck_theme(deck, EX)
    deck.diagnostics += diags
    return layout_slide(deck.slides[0], deck, theme, 0), deck, theme


def _by_id(placed, ident):
    return [p for p in placed if isinstance(p.element, Shape) and p.element.id == ident]


def _role(placed, r):
    return [p for p in placed if isinstance(p.element, Text) and p.element.role == r]


def test_band_is_anchored_to_the_top_and_title_sits_bottom_aligned_in_it():
    placed, _, th = _cover()
    (band,) = _by_id(placed, "band")
    assert (band.x, band.y, band.w) == (0, 0, W)
    assert abs(band.h - th.cover_band_h * H) <= 1
    (t,) = _role(placed, "title")
    (s,) = _role(placed, "subtitle")
    assert band.y <= t.y and s.y + s.h <= band.y + band.h
    assert t.y + t.h <= s.y  # subtitle below the title, with air
    assert s.y - (t.y + t.h) >= 0.5 * (0.2 * 914400)
    assert band.h - (s.y + s.h - band.y) >= 0.9 * 0.5 * 914400  # token pad under the block


def test_rule_runs_along_the_band_edge_and_can_be_switched_off():
    placed, _, _ = _cover()
    (band,) = _by_id(placed, "band")
    (rule,) = _by_id(placed, "rule")
    assert rule.y == band.h and rule.w == W and rule.style.fill == "accent"
    placed, _, _ = _cover(style="cover.rule=none")
    assert not _by_id(placed, "rule")


def test_footer_shows_as_caption_on_the_cover_and_can_be_switched_off():
    placed, _, _ = _cover()
    cap = [p for p in _role(placed, "caption") if p.element.attrs.get("field") == "footer"]
    assert len(cap) == 1 and cap[0].y > 0.85 * H
    placed, _, _ = _cover(style="cover.footer=off")
    assert not _role(placed, "caption")


def test_tokens_parse_and_drive_the_geometry():
    placed, deck, th = _cover(style="cover.band_h=45% cover.pad=1in cover.gap=0.4in cover.rule_h=0.1in")
    assert not [d for d in deck.diagnostics if d.rule == "bad-token"]
    assert th.cover_band_h == 0.45
    (band,) = _by_id(placed, "band")
    assert abs(band.h - 0.45 * H) <= 1
    (rule,) = _by_id(placed, "rule")
    assert rule.h == 914400 // 10


def test_band_h_zero_restores_the_legacy_cover():
    placed, _, _ = _cover(style="cover.band_h=0")
    (band,) = _by_id(placed, "band")
    assert band.y > 0
    assert not _by_id(placed, "rule")


def test_no_band_color_still_anchors_the_block_with_a_margin_wide_rule():
    placed, _, th = _cover(head="theme: default\nfooter: ACME Corp\n")
    assert not _by_id(placed, "band")
    (rule,) = _by_id(placed, "rule")
    (t,) = _role(placed, "title")
    assert t.y + t.h < rule.y and rule.w < W


def test_explicit_cover_styling_keeps_the_legacy_composition():
    src = "theme: jp-business\n\n# Plan {size=30}\nSub\n"
    deck = parse(src)
    theme, _ = deck_theme(deck, EX)
    placed = layout_slide(deck.slides[0], deck, theme, 0)
    assert not _by_id(placed, "rule")


def test_bad_values_never_raise():
    placed, deck, _ = _cover(style="cover.band_h=banana cover.pad=wide cover.rule=#12")
    assert placed


def test_renders_into_a_reopened_pptx(tmp_path):
    out = tmp_path / "c.pptx"
    build("theme: jp-business\nfooter: ACME Corp\n\n# Plan\nFY2027\n\n# Next\ntext\n- a\n", out)
    sl = Presentation(str(out)).slides[0]
    texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame]
    assert "Plan" in texts and "ACME Corp" in texts
    assert any(sh.top == 0 and sh.width == W for sh in sl.shapes)
