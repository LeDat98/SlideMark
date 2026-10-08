"""A table beside a taller chart, or the last block above a footnote, fills down to there (vfill)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from slidemark.layout import layout_slide
from slidemark.parser import parse
from slidemark.render import render
from slidemark.template import deck_theme

FULL = """# T

| a | b | c |
|---|---|---|
| x | 1 | 2 |
| y | 2 | 3 |
| z | 3 | 4 |

^ footnote text
"""

CHART = """# T
@1:1

```bar
,p,q,r
a,4,1,-2
```

| a | b |
|---|---|
| x | 1 |
| y | 2 |
| z | 3 |
"""


def _lay(src: str, tmp_path: Path):
    d = parse(src)
    th = deck_theme(d, str(tmp_path))
    th = th[0] if isinstance(th, tuple) else th
    return d, th, layout_slide(d.slides[0], d, th, 0)


def test_table_beside_chart_takes_chart_height(tmp_path):
    d, th, items = _lay(CHART, tmp_path)
    tab = next(p for p in items if type(p.element).__name__ == "Table")
    ch = next(p for p in items if type(p.element).__name__ == "Chart")
    assert tab.y <= ch.y + 2
    assert tab.h > 0.6 * ch.h  # was a short block hugging the top
    assert tab.y + tab.h <= ch.y + ch.h + 2
    out = tmp_path / "o.pptx"
    render(d, [items], th, out)
    assert Presentation(str(out)).slides[0].shapes


def test_full_width_table_reaches_footnote(tmp_path):
    _d, _th, items = _lay(FULL, tmp_path)
    tab = next(p for p in items if type(p.element).__name__ == "Table")
    foot = [p for p in items if getattr(p.element, "role", "") == "footnote"]
    assert foot
    assert tab.y + tab.h <= foot[0].y
    assert foot[0].y - (tab.y + tab.h) < 0.6 * 914400  # within about a gutter plus padding


def test_fit_is_safe_on_odd_tables(tmp_path):
    for src in ("# T\n\n|a|\n|-|\n\n^ f\n", "# T\n@1:1\n![x](a.png)\n|a|b|\n|-|-|\n"):
        _lay(src, tmp_path)
