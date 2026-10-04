from pptx import Presentation

from slidemark import build


def test_end_to_end_table_merge(tmp_path):
    src = "# T\n| a | b | c |\n|-|-|-|\n| x | y | < |\n| z | w | v |\n"
    out = tmp_path / "o.pptx"
    deck = build(src, out)
    assert not [d for d in deck.diagnostics if d.level == "error"]
    tbl = next(s for s in Presentation(out).slides[0].shapes if s.has_table).table
    assert tbl.cell(1, 1).span_width == 2
    assert tbl.cell(1, 0).text == "x"


def test_examples_build(tmp_path):
    from pathlib import Path

    for md in sorted(Path(__file__).parent.parent.glob("examples/*.md")):
        deck = build(md, tmp_path / f"{md.stem}.pptx")
        assert not [d for d in deck.diagnostics if d.level == "error"], md
