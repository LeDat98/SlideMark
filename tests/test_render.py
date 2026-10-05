from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from slidemark.ir import (
    Box,
    Cell,
    Chart,
    Code,
    Container,
    Deck,
    Image,
    Paragraph,
    Raw,
    Run,
    Series,
    Shape,
    Slide,
    Style,
    Table,
    Text,
)
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.theme import get_theme


def T(s, role="body", **kw):
    return Text(role=role, paragraphs=[Paragraph(runs=[Run(text=s)])], **kw)


def bullets(*items, marker="bullet"):
    return Text(paragraphs=[Paragraph(runs=[Run(text=i)], marker=marker) for i in items])


def build(deck: Deck, tmp_path: Path, theme="default", hug: bool = True) -> Presentation:
    th = get_theme(theme).model_copy(deep=True)
    th.layout.hug_cards = hug
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    out = tmp_path / "out.pptx"
    render(deck, placed, th, out)
    return Presentation(str(out))


def shapes_named(slide, prefix):
    return [s for s in slide.shapes if s.name.startswith(prefix)]


def test_title_runs_and_bullets(tmp_path):
    s = Slide(
        title=T("Hello", "title"),
        elements=[
            Text(
                paragraphs=[
                    Paragraph(
                        runs=[
                            Run(text="plain "),
                            Run(text="bold", bold=True),
                            Run(text="ital", italic=True, underline=True, strike=True),
                            Run(text="x", sup=True),
                            Run(text="y", sub=True),
                            Run(text="code", code=True),
                            Run(text="hl", highlight="#FFFF00", color="danger"),
                            Run(text="link", link="https://example.com"),
                        ],
                        marker="bullet",
                    ),
                    Paragraph(runs=[Run(text="nested")], marker="bullet", level=1),
                    Paragraph(runs=[Run(text="one")], marker="number"),
                ]
            )
        ],
    )
    prs = build(Deck(slides=[s]), tmp_path)
    sl = prs.slides[0]
    assert sl.shapes.title is not None and sl.shapes.title.text_frame.text == "Hello"
    assert sl.shapes.title.name == "Title"
    body = shapes_named(sl, "Text")[0]
    xml = body._element.xml
    assert "<a:buChar" in xml and 'char="•"' in xml and "<a:buAutoNum" in xml
    assert 'lvl="1"' in xml or 'marL="' in xml
    runs = body.text_frame.paragraphs[0].runs
    assert runs[1].font.bold and runs[2].font.italic and runs[2].font.underline
    assert 'strike="sngStrike"' in xml and 'baseline="30000"' in xml and 'baseline="-25000"' in xml
    assert "<a:highlight>" in xml
    assert runs[7].hyperlink.address == "https://example.com"
    assert "Consolas" in xml


def test_cjk_has_ea_and_lang(tmp_path):
    s = Slide(title=T("日本語のタイトル", "title"), elements=[bullets("売上は前年比で8%増加しました")])
    prs = build(Deck(slides=[s], lang="ja"), tmp_path, "jp-business")
    xml = shapes_named(prs.slides[0], "Text")[0]._element.xml
    assert "<a:ea " in xml and 'lang="ja-JP"' in xml
    assert "<a:ea " in prs.slides[0].shapes.title._element.xml


def test_table_merges_and_fills(tmp_path):
    def c(t, **kw):
        return Cell(paragraphs=[Paragraph(runs=[Run(text=t)])], **kw)

    t = Table(
        rows=[
            [c("A"), c("B", colspan=2)],
            [c("1"), c("2"), c("3")],
            [c("tall", rowspan=2), c("x"), c("y")],
            [c("p"), c("q")],
        ]
    )
    prs = build(Deck(slides=[Slide(title=T("T", "title"), elements=[t])]), tmp_path)
    gf = shapes_named(prs.slides[0], "Table")[0]
    tbl = gf.table
    assert len(tbl.rows) == 4 and len(tbl.columns) == 3
    assert tbl.cell(0, 1).is_merge_origin and tbl.cell(0, 1).span_width == 2
    assert tbl.cell(2, 0).span_height == 2
    assert tbl.cell(3, 1).text_frame.text == "p"
    assert "tableStyleId" not in gf._element.xml
    assert "<a:solidFill>" in tbl.cell(1, 1)._tc.xml


@pytest.mark.parametrize(
    "kind,xl",
    [
        ("column", "COLUMN_CLUSTERED"),
        ("bar", "BAR_CLUSTERED"),
        ("stacked-bar", "BAR_STACKED"),
        ("stacked-column", "COLUMN_STACKED"),
        ("line", "LINE_MARKERS"),
        ("area", "AREA"),
        ("pie", "PIE"),
        ("doughnut", "DOUGHNUT"),
        ("scatter", "XY_SCATTER"),
        ("radar", "RADAR_MARKERS"),
    ],
)
def test_charts(tmp_path, kind, xl):
    ch = Chart(
        kind=kind,
        title="Sales",
        categories=["1", "2", "3"],
        series=[Series(name="a", values=[1, 2, 3]), Series(name="b", values=[3, None, 1])],
        options={"legend": "bottom", "labels": True},
    )
    d = Deck(slides=[Slide(title=T("C", "title"), elements=[ch])])
    prs = build(d, tmp_path)
    gf = shapes_named(prs.slides[0], "Chart")[0]
    assert gf.has_chart
    assert gf.chart.chart_type.name == xl
    assert gf.chart.chart_title.text_frame.text == "Sales"
    assert not [x for x in d.diagnostics if x.level == "error"]


def test_image_modes_and_missing(tmp_path):
    from PIL import Image as PI

    PI.new("RGB", (200, 100), "red").save(tmp_path / "a.png")
    d = Deck(
        slides=[
            Slide(
                elements=[
                    Image(src="a.png", alt="A", fit="cover"),
                    Image(src="a.png", alt="B", fit="contain"),
                    Image(src="missing.png", alt="M"),
                ],
                layout="blank",
            )
        ],
        attrs={"base_dir": str(tmp_path)},
    )
    prs = build(d, tmp_path)
    pics = [s for s in prs.slides[0].shapes if s.shape_type == 13]
    assert len(pics) == 2
    assert pics[0].crop_left > 0 or pics[0].crop_top > 0
    assert pics[1].crop_left == 0
    assert pics[1]._element.nvPicPr.cNvPr.get("descr") == "B"
    assert any(x.rule == "image-missing" for x in d.diagnostics)
    assert any("image" in s.name.lower() and not s.shape_type == 13 for s in prs.slides[0].shapes)


def test_code_shapes_container_raw_notes_hidden_props(tmp_path):
    s = Slide(
        title=T("Misc", "title"),
        elements=[
            Code(lang="python", text="def f(x):\n    return x + 1  # hi\n"),
            Container(title=T("Box", "heading"), children=[T("inside")]),
            Shape(shape="ellipse", paragraphs=[Paragraph(runs=[Run(text="O")])], style=Style(fill="accent")),
            Raw(kind="mermaid", source="graph TD"),
        ],
        notes="speaker notes",
        hidden=True,
        background="#112233",
        transition="fade",
    )
    d = Deck(slides=[s], title="My deck", author="Me", footer="ACME", slide_number=True)
    prs = build(d, tmp_path)
    sl = prs.slides[0]
    assert sl.notes_slide.notes_text_frame.text == "speaker notes"
    assert sl._element.get("show") == "0"
    assert "p:transition" in sl._element.xml and "<p:fade" in sl._element.xml
    assert prs.core_properties.title == "My deck" and prs.core_properties.author == "Me"
    code = shapes_named(sl, "Code")[0]
    assert len(code.text_frame.paragraphs) == 2
    assert 'typeface="Consolas"' in code._element.xml
    assert code._element.xml.count("<a:solidFill>") > 2  # token colors
    card = shapes_named(sl, "Card")[0]
    assert card.shape_type == 1
    assert [x for x in sl.shapes if x.name == "Footer"][0].text_frame.text == "ACME"
    num = [x for x in sl.shapes if x.name == "Slide Number"][0]
    assert 'type="slidenum"' in num._element.xml
    bg = sl._element.xml
    assert "112233" in bg
    assert any(x.rule == "raw" for x in d.diagnostics)


def test_internal_links(tmp_path):
    a = Slide(
        id="intro",
        title=T("A", "title"),
        elements=[
            Text(
                paragraphs=[
                    Paragraph(
                        runs=[
                            Run(text="go", link="#2"),
                            Run(text="id", link="#intro"),
                            Run(text="bad", link="#9"),
                        ]
                    )
                ]
            )
        ],
    )
    b = Slide(title=T("B", "title"))
    d = Deck(slides=[a, b])
    prs = build(d, tmp_path)
    xml = shapes_named(prs.slides[0], "Text")[0]._element.xml
    assert xml.count("hlinksldjump") == 2
    rels = [r for r in prs.slides[0].part.rels.values() if r.reltype.endswith("/slide")]
    assert len(rels) == 2
    assert any(x.rule == "bad-jump" for x in d.diagnostics)


def test_flow_and_chevron_shapes(tmp_path):
    s = Slide(
        title=T("Steps", "title"),
        classes=["chevron"],
        grid="3",
        elements=[T("one"), T("two"), T("three")],
    )
    prs = build(Deck(slides=[s]), tmp_path)
    ch = [
        x
        for x in prs.slides[0].shapes
        if x.shape_type == 1 and x.auto_shape_type is not None and "CHEVRON" in str(x.auto_shape_type)
    ]
    assert len(ch) == 3
    assert ch[0].text_frame.text == "one"


def test_never_raises_on_odd_input(tmp_path):
    s = Slide(
        title=T("odd", "title"),
        grid="zzz!!",
        elements=[
            Shape(shape="nope"),
            Table(rows=[]),
            Chart(kind="pie"),
            Container(children=[Text(box=Box(x="bad", y="1in"))]),
        ],
    )
    d = Deck(slides=[s], size="weird")
    prs = build(d, tmp_path)
    assert len(prs.slides) == 1


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_opens_in_libreoffice(tmp_path):
    def c(t, **kw):
        return Cell(paragraphs=[Paragraph(runs=[Run(text=t)])], **kw)

    s1 = Slide(
        title=T("日本語 and English", "title"),
        lead=T("lead", "lead"),
        elements=[
            Container(title=T("現状", "heading"), children=[bullets("売上 12億円", "顧客数 1,240社")]),
            Container(
                title=T("Chart", "heading"),
                children=[
                    Chart(kind="column", categories=["Q1", "Q2"], series=[Series(name="s", values=[1, 2])])
                ],
            ),
            Table(rows=[[c("a"), c("b", colspan=2)], [c("1"), c("2"), c("3")]]),
        ],
        conclusion=T("結論", "conclusion"),
        footnotes=[T("※ 出所", "footnote")],
    )
    s2 = Slide(title=T("Code", "title"), elements=[Code(lang="python", text="print(1)")])
    d = Deck(slides=[s1, s2], footer="f", slide_number=True)
    th = get_theme("jp-business")
    placed = [layout_slide(s, d, th, i) for i, s in enumerate(d.slides)]
    out = tmp_path / "lo.pptx"
    render(d, placed, th, out)
    r = subprocess.run(
        [
            "soffice",
            "--headless",
            "-env:UserInstallation=file:///tmp/lo-b",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(out),
        ],
        capture_output=True,
        timeout=120,
    )
    assert (tmp_path / "lo.pdf").exists(), r.stderr
    assert (tmp_path / "lo.pdf").stat().st_size > 1000
    assert Emu(1)  # keep import used


def _space_before_pts(
    tmp_path, items, theme="default"
):  # cards keep spare height (hug off) so paragraphs spread
    boxes = [Container(title=T(h, "heading"), children=[bullets(*items)]) for h in ("a", "b")]
    d = Deck(slides=[Slide(title=T("t", "title"), grid="2", elements=boxes)])
    prs = build(d, tmp_path, theme, hug=False)
    out = []
    for shp in prs.slides[0].shapes:
        if shp.has_text_frame and shp.text_frame.text.startswith(items[0]):
            for para in shp.text_frame.paragraphs:
                sb = para._p.pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}spcBef")
                pts = sb[0].get("val")
                out.append((int(pts), para.runs[0].font.size.pt))
            break
    return out


def test_roomy_card_paragraphs_get_a_larger_spc_bef(tmp_path):
    from slidemark.layout import measure

    rows = _space_before_pts(tmp_path, ("one", "two", "three"))
    assert rows[0][0] == 0
    for val, size in rows[1:]:
        assert val / 100 > size * measure.PARA_GAP + 0.5  # spread: more than the default gap
        assert val / 100 <= size * 0.6 + 0.02  # ... up to about half a line


def test_full_card_paragraphs_keep_the_default_spc_bef(tmp_path):
    from slidemark.layout import measure

    rows = _space_before_pts(tmp_path, tuple(f"line number {i} of a very full card" for i in range(9)))
    for val, size in rows[1:]:
        assert abs(val / 100 - size * measure.PARA_GAP) < 0.02
