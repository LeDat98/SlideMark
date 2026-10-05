"""Sparse rows of short boxes: one row, text and cards grow together, a small empty band."""

from __future__ import annotations

from pathlib import Path

from slidemark.ir import Container, Placed
from slidemark.layout import layout_slide
from slidemark.parser import parse
from slidemark.template import deck_theme

from .helpers import top_anchored

ROOT = Path(__file__).resolve().parent.parent


def _slide(name: str, number: int) -> tuple[list[Placed], int]:
    md = ROOT / "examples" / f"{name}.md"
    deck = parse(md.read_text(encoding="utf-8"))
    theme, _ = deck_theme(deck, md.parent)
    items = layout_slide(deck.slides[number - 1], deck, theme, number - 1)
    return items, 6492240  # slide bottom minus the page margin (EMU)


def _cards(items: list[Placed]) -> list[Placed]:
    return [p for p in items if isinstance(p.element, Container)]


def _lead_tails(items: list[Placed]) -> list[float]:
    """Card tail (empty share below the last content line) of the row leaders, as bench/whitespace.py."""
    import sys

    sys.path.insert(0, str(ROOT / "bench"))
    from whitespace import card_tails

    return [t for _n, t, lead in card_tails(items) if lead]


def _band(items: list[Placed], height: int) -> float:
    """Empty height between the cards and the next element below them (or the slide bottom), per body."""
    cards = _cards(items)
    top, bottom = min(c.y for c in cards), max(c.y + c.h for c in cards)
    below = [p.y for p in items if p.y >= bottom and not isinstance(p.element, Container)]
    limit = min(below) if below else height
    return (limit - bottom) / max(limit - top, 1)


def test_three_short_boxes_stay_one_row_and_fill_the_body():  # 13-brand-aurora slide 2
    items, h = _slide("13-brand-aurora", 2)
    cards = _cards(items)
    assert len(cards) == 3 and len({c.y for c in cards}) == 1  # not stacked into wide bars
    assert max(c.w for c in cards) < 0.4 * 12192000
    assert _lead_tails(items) and max(_lead_tails(items)) <= 0.25  # cards hug; the leftover is one band


def test_row_above_a_conclusion_bar_loses_most_of_the_band():  # 01-basics slide 3
    items, h = _slide("01-basics", 3)
    cards = _cards(items)
    assert len(cards) == 3 and len({c.y for c in cards}) == 1
    assert max(_lead_tails(items)) <= 0.25  # cards hug their text (were 3x taller)


def test_html_mixed_card_row_has_no_big_band():  # 15-html-mixed slide 2
    items, h = _slide("15-html-mixed", 2)
    cards = _cards(items)
    assert len(cards) == 3 and len({c.y for c in cards}) == 1
    assert max(_lead_tails(items)) <= 0.25


def test_dense_two_box_slide_hugs_its_text():  # 11-jp-consulting slide 9
    items, h = _slide("11-jp-consulting", 9)
    cards = _cards(items)
    assert len(cards) == 2 and len({c.y for c in cards}) == 1
    import sys

    sys.path.insert(0, str(ROOT / "bench"))
    from whitespace import card_fills

    fills = [f for _t, f in card_fills(items)]
    assert sum(fills) / len(fills) >= 0.6  # cards hug their text; the leftover is one band under them
    assert max(_lead_tails(items)) <= 0.25


def test_text_grows_before_cards_stretch():  # 13-brand-aurora slide 2: cards were ~3x their text
    import sys

    sys.path.insert(0, str(ROOT / "bench"))
    from whitespace import card_fills

    items, h = _slide("13-brand-aurora", 2)
    fills = [f for _t, f in card_fills(items)]
    assert fills and max(fills) >= 0.5 and sum(fills) / len(fills) >= 0.38
    body = [p for p in items if getattr(p.element, "role", None) == "body"]
    assert 1.2 <= min(p.font_scale for p in body) <= 1.5 + 1e-6  # text grew first, within ``sparse_step_max``


def test_dense_two_box_slide_cards_are_half_filled():  # 11-jp-consulting slide 9: short CJK lines cannot grow
    import sys

    sys.path.insert(0, str(ROOT / "bench"))
    from whitespace import card_fills

    items, _h = _slide("11-jp-consulting", 9)
    fills = [f for _t, f in card_fills(items)]
    assert len(fills) == 2 and sum(fills) / len(fills) >= 0.5 and min(fills) >= 0.4


def test_no_row_leader_card_has_a_tail_over_a_quarter_on_any_example():
    import sys

    sys.path.insert(0, str(ROOT / "bench"))
    from whitespace import TAIL_MAX

    bad = []
    for md in sorted((ROOT / "examples").glob("*.md")):
        deck = parse(md.read_text(encoding="utf-8"))
        theme, _ = deck_theme(deck, md.parent)
        for i, s in enumerate(deck.slides):
            tails = _lead_tails(layout_slide(s, deck, theme, i))
            if tails and max(tails) > TAIL_MAX:
                bad.append((md.stem, i + 1, round(max(tails), 2)))
    assert not bad, bad


def test_rendered_cards_hug_their_text(tmp_path):
    from pptx import Presentation

    from slidemark.render import render

    md = ROOT / "examples" / "16-jp-strategy.md"
    deck = parse(md.read_text(encoding="utf-8"))
    theme, _ = deck_theme(deck, md.parent)
    theme = top_anchored(theme)
    placed = [layout_slide(s, deck, theme, i) for i, s in enumerate(deck.slides)]
    out = tmp_path / "d.pptx"
    render(deck, placed, theme, out)
    slide = Presentation(str(out)).slides[9]
    boxes = [
        sh
        for sh in slide.shapes
        if sh.has_text_frame and "ご決議事項" in sh.text_frame.text and sh.height > 0
    ]
    cards = [sh for sh in slide.shapes if sh.shape_type == 1 and sh.height > 1_000_000 and sh.top > 1_000_000]
    assert boxes and cards
    assert max(sh.top + sh.height for sh in cards) < 0.55 * Presentation(str(out)).slide_height
