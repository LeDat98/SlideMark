"""A badge in a ``{.gantt}`` bar label becomes a native rounded pill inside the bar."""

from __future__ import annotations

from pptx import Presentation

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.ir import Shape
from slidemark.layout import layout_slide
from slidemark.parser import parse
from slidemark.theme import get_theme

PILL = """# Plan

{.gantt header=2}
| 施策 | 2027 | < | 2028 | < |
|-|-|-|-|-|
| ^ | 上期 | 下期 | 上期 | 下期 |
| IoT | 新機種へ搭載 [開始]{.badge} | < | 既設機 | ― |
| AI | - | 25社体制 [完了]{.badge .success} | < | < |
"""


def _pills(md: str):
    deck = parse(md)
    placed = layout_slide(deck.slides[0], deck, get_theme("default"), 0)
    return [p for p in placed if isinstance(p.element, Shape) and "pill" in p.element.classes]


def test_bar_badge_becomes_a_native_pill_inside_the_bar(tmp_path):
    out = tmp_path / "p.pptx"
    build(PILL, out)
    shapes = list(Presentation(out).slides[0].shapes)
    pills = [s for s in shapes if s.name.startswith("Pill")]
    assert {s.text_frame.text.strip() for s in pills} == {"開始", "完了"}
    for pill in pills:
        assert "ROUNDED" in str(pill.auto_shape_type)
        bar = next(
            s
            for s in shapes
            if s is not pill
            and s.has_text_frame
            and s.left <= pill.left
            and pill.left + pill.width <= s.left + s.width
            and s.top <= pill.top
            and pill.top + pill.height <= s.top + s.height
            and s.text_frame.text.replace("\u2060", "").strip() in {"新機種へ搭載", "25社体制"}
        )
        assert pill.height < 0.8 * bar.height
        assert pill.left > bar.left
    assert not [
        s
        for s in shapes
        if not s.name.startswith("Pill") and s.has_text_frame and "開始" in s.text_frame.text
    ]


def test_bar_badge_stays_a_highlight_when_it_does_not_fit():
    md = PILL.replace("新機種へ搭載 [開始]{.badge}", "新機種へ搭載と既設機の後付けと保守 [開始]{.badge}")
    assert len(_pills(md)) == 1  # only the "25社体制" bar still fits


def test_bar_pill_import_round_trip(tmp_path):
    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    build(PILL, a)
    text, _ = import_pptx(a)
    assert "新機種へ搭載 [開始]{.badge" in text and "25社体制 [完了]{.badge" in text
    build(text, b)
    text2, _ = import_pptx(b)
    assert text2 == text


def test_fuzz_safe_odd_badges():
    for cell in ("[x]{.badge}", "a [b]{.badge} c", "a [b]{.badge}[c]{.badge .danger}", "[ ]{.badge}"):
        _pills(PILL.replace("新機種へ搭載 [開始]{.badge}", cell))
