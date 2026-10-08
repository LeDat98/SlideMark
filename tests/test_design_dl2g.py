"""DL2 lane G: the last two KPI forms, ``kpi.band`` (a header band behind the label of every ``.kpi``
card) and ``kpi.unit.size`` (the trailing unit of the value line, 億円 / 名 / %, drawn smaller than the
digits). Each form: parser (token reaches the theme), render (the .pptx is reopened and its shapes and run
sizes asserted), importer round trip, ``attr-ignored`` gating, and fuzz-safety (a bad value is a
``bad-token`` warning with a hint, never a raise)."""

from __future__ import annotations

import pytest
from pptx import Presentation

from slidemark import build, parse
from slidemark.importer import import_pptx
from slidemark.ir import Paragraph, Run
from slidemark.layout.kpiunit import split_unit
from slidemark.template import deck_theme

HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#222B36 primary=#142B4D secondary=#1F5FA8 accent=#E09F1F teal=#2A9D8F "
    "muted=#59626E surface=#EEF2F7 border=#C9D2DE sand=#F2E8CF\n"
)
KPIS = (
    "\n# 目標\n> lead\n@kpi\n"
    "## 売上高\n1,280億円\n前年比 +8%\n"
    "## 営業利益\n96億円\n前年比 +12%\n"
    "## 営業利益率\n7.5%\n+0.3pt\n"
    "## 比率\n3.8\n+0.2\n"
)


def make(tmp_path, text: str, name: str = "d"):
    deck = build(HEAD + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def shapes(prs, i, prefix=""):
    return [s for s in prs.slides[i].shapes if s.name.lower().startswith(prefix.lower())]


def warns(deck, rule):
    return [d for d in deck.diagnostics if d.rule == rule]


def runs(shp):
    return [r for p in shp.text_frame.paragraphs for r in p.runs]


def value_runs(prs, i=0):
    """The runs of the first paragraph of every KPI text box (number and unit), per card."""
    return [
        [r for r in s.text_frame.paragraphs[0].runs]
        for s in shapes(prs, i, "Text")
        if s.text_frame.text.strip()
    ]


def fill_hex(shp) -> str:
    return str(shp.fill.fore_color.rgb)


# --------------------------------------------------------------------------- 1. kpi.band


def test_band_tokens_reach_the_theme():
    d = parse(HEAD + "style: kpi.band=primary kpi.band.color=sand kpi.band.size=18\n" + KPIS)
    th, _ = deck_theme(d, ".")
    assert th.kpi_band == "primary" and th.kpi_band_color == "sand" and th.kpi_band_size == 18
    assert not [x for x in d.diagnostics if x.rule in ("bad-token", "attr-ignored", "unknown-token")]


def test_band_is_off_by_default(tmp_path):
    _, prs = make(tmp_path, KPIS)
    assert not [s for s in shapes(prs, 0, "Heading") if s.fill.type == 1]


def test_band_is_a_filled_label_at_the_top_of_every_card(tmp_path):
    deck, prs = make(tmp_path, "style: kpi.band=primary\n" + KPIS)
    assert not warns(deck, "overflow") and not warns(deck, "contrast") and not warns(deck, "attr-ignored")
    cards, heads = shapes(prs, 0, "Card"), shapes(prs, 0, "Heading")
    assert len(cards) == len(heads) == 4
    for card, h in zip(cards, heads, strict=True):
        assert fill_hex(h) == "142B4D"
        assert (h.left, h.top, h.width) == (card.left, card.top, card.width)  # flush on the card's top edge
        assert h.height < card.height / 2
        assert str(runs(h)[0].font.color.rgb) == "FFFFFF"  # the ink reads on the band
    texts = [s for s in shapes(prs, 0, "Text") if s.text_frame.text.strip()]
    for card, h, t in zip(cards, heads, texts, strict=True):
        assert t.top >= h.top + h.height - 2  # number and caption sit under the band
        assert t.top + t.height <= card.top + card.height + 2
    assert len({h.top for h in heads}) == 1 and len({h.height for h in heads}) == 1  # one row, one band


def test_band_ink_and_size_tokens(tmp_path):
    _, prs = make(tmp_path, "style: kpi.band=sand kpi.band.color=teal kpi.band.size=15\n" + KPIS)
    h = shapes(prs, 0, "Heading")[0]
    assert fill_hex(h) == "F2E8CF" and runs(h)[0].font.size.pt == 15
    # teal on sand is too pale: the ink falls back to something readable instead of the stated color
    assert str(runs(h)[0].font.color.rgb) != "FFFFFF"


def test_band_size_is_pinned_in_a_lone_row(tmp_path):
    one = "\n# K\n## A {.kpi}\n1,280億円\nvs +8%\n## B {.kpi}\n96億円\nvs +12%\n"
    _, prs = make(tmp_path, "style: kpi.band=primary kpi.band.size=14\n" + one)
    assert {runs(h)[0].font.size.pt for h in shapes(prs, 0, "Heading")} == {14.0}
    cards, heads = shapes(prs, 0, "Card"), shapes(prs, 0, "Heading")
    assert all(h.top == c.top and h.width == c.width for c, h in zip(cards, heads, strict=True))


def test_band_over_a_lone_kpi_row_keeps_the_label_on_the_card_top(tmp_path):
    one = "\n# K\n## A {.kpi}\n1,280\nvs +8%\n## B {.kpi}\n96\nvs +12%\n"
    deck, prs = make(tmp_path, "style: kpi.band=primary\n" + one)
    assert not warns(deck, "overflow")
    for c, h in zip(shapes(prs, 0, "Card"), shapes(prs, 0, "Heading"), strict=True):
        assert h.top == c.top and h.top + h.height < c.top + c.height / 2


def test_band_with_an_icon_puts_the_icon_under_the_band(tmp_path):
    deck, prs = make(
        tmp_path, "style: kpi.band=primary\n\n# K\n## A {.kpi icon=chart}\n12\ncap\n## B {.kpi}\n3\nc\n"
    )
    assert not warns(deck, "overflow")
    head = shapes(prs, 0, "Heading")[0]
    icons = [s for s in prs.slides[0].shapes if "icon" in s.name.lower()]
    assert icons and all(i.top >= head.top + head.height - 2 for i in icons)


def test_band_works_with_stripe_and_rule(tmp_path):
    deck, prs = make(tmp_path, "style: kpi.band=primary kpi.stripe=accent kpi.rule=border\n" + KPIS)
    assert not warns(deck, "overflow") and not warns(deck, "contrast")
    assert len(shapes(prs, 0, "KPI value")) == 4
    assert len([s for s in shapes(prs, 0, "Heading") if s.fill.type == 1]) == 4


def test_band_only_in_the_slide_that_states_it(tmp_path):
    text = KPIS + "\n# 2\n@kpi\nstyle: kpi.band=secondary\n## A\n1\nc\n## B\n2\nc\n"
    _, prs = make(tmp_path, "style: kpi.band=primary\n" + text)
    assert {fill_hex(h) for h in shapes(prs, 0, "Heading")} == {"142B4D"}
    assert {fill_hex(h) for h in shapes(prs, 1, "Heading")} == {"1F5FA8"}


def test_band_round_trip_folds_back_to_a_kpi_card(tmp_path):
    make(tmp_path, "style: kpi.band=primary kpi.unit.size=22\n" + KPIS, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert text.count("{.kpi") == 4
    assert "kpi_band=primary" in text  # the design part carries the token
    lines = text.splitlines()
    k = next(i for i, ln in enumerate(lines) if ln.startswith("## 売上高"))
    assert lines[k + 1 : k + 3] == ["1,280億円", "前年比 +8%"]  # the band is the label, not a box of its own
    assert "<br" not in text and text.count("## ") == 4


def test_band_foreign_import_keeps_label_and_value(tmp_path):
    """A deck without the design part (a band is only a filled label box) still imports as KPI cards."""
    make(tmp_path, "style: kpi.band=primary kpi.unit.size=18\n" + KPIS, "fo")
    prs = Presentation(str(tmp_path / "fo.pptx"))
    for rid, rel in list(prs.part.rels.items()):
        if rel.reltype.endswith("/customXml"):
            prs.part.drop_rel(rid)
    prs.save(tmp_path / "foreign.pptx")
    text, _ = import_pptx(tmp_path / "foreign.pptx")
    assert "kpi_band" not in text  # really without the design part
    assert text.count("{.kpi") == 4
    lines = text.splitlines()
    k = next(i for i, ln in enumerate(lines) if ln.startswith("## 売上高"))
    assert lines[k + 1 : k + 3] == ["1,280億円", "前年比 +8%"]


# --------------------------------------------------------------------------- 2. kpi.unit.size


def test_unit_tokens_reach_the_theme():
    d = parse(HEAD + "style: kpi.unit.size=22 kpi.unit.color=muted\n" + KPIS)
    th, _ = deck_theme(d, ".")
    assert th.kpi_unit_size == 22 and th.kpi_unit_color == "muted"
    assert not [x for x in d.diagnostics if x.rule in ("bad-token", "attr-ignored", "unknown-token")]


def test_unit_is_off_by_default(tmp_path):
    _, prs = make(tmp_path, KPIS)
    assert all(len(rs) == 1 for rs in value_runs(prs))


def test_unit_run_is_smaller_than_the_digits(tmp_path):
    deck, prs = make(tmp_path, "style: kpi.unit.size=18\n" + KPIS)
    assert not warns(deck, "attr-ignored") and not warns(deck, "overflow")
    vals = value_runs(prs)
    assert [[r.text for r in rs] for rs in vals] == [["1,280", "億円"], ["96", "億円"], ["7.5", "%"], ["3.8"]]
    for rs in vals[:3]:
        assert rs[1].font.size.pt == 18 and rs[0].font.size.pt > 18
        assert rs[1].font.bold == rs[0].font.bold  # same weight, same face
    shp = next(s for s in shapes(prs, 0, "Text") if s.text_frame.text.startswith("1,280"))
    assert shp.text_frame.text.replace("⁠", "").splitlines()[0] == "1,280億円"  # still one value line


def test_unit_never_exceeds_the_number(tmp_path):
    _, prs = make(tmp_path, "style: kpi.unit.size=90\n" + KPIS)
    for rs in value_runs(prs)[:3]:
        assert rs[1].font.size.pt <= rs[0].font.size.pt + 0.01


def test_unit_color(tmp_path):
    _, prs = make(tmp_path, "style: kpi.unit.size=18 kpi.unit.color=muted\n" + KPIS)
    rs = value_runs(prs)[0]
    assert str(rs[1].font.color.rgb) == "59626E" and str(rs[0].font.color.rgb) != "59626E"


def test_unit_scales_with_the_slide_shrink(tmp_path):
    many = "".join(f"## K{i}\n{i * 111},000億円\nc\n" for i in range(1, 7))
    _, prs = make(tmp_path, "style: kpi.unit.size=20\n\n# K\n@kpi\n@6\n" + many)
    sizes = {(rs[0].font.size.pt, rs[1].font.size.pt) for rs in value_runs(prs)}
    assert all(u <= n for n, u in sizes)


def test_unit_in_a_slide_line_changes_that_slide_only(tmp_path):
    text = KPIS + "\n# 2\n@kpi\nstyle: kpi.unit.size=16\n## A\n12名\nc\n## B\n5名\nc\n"
    _, prs = make(tmp_path, text)
    assert all(len(rs) == 1 for rs in value_runs(prs, 0))
    assert all(len(rs) == 2 and rs[1].font.size.pt == 16 for rs in value_runs(prs, 1))


def test_unit_with_the_divider_rule(tmp_path):
    deck, prs = make(tmp_path, "style: kpi.unit.size=18 kpi.rule=border\n" + KPIS)
    assert not warns(deck, "overflow")
    values = shapes(prs, 0, "KPI value")
    assert [[r.text for r in runs(v)] for v in values][0] == ["1,280", "億円"]


def test_unit_round_trip_joins_the_runs_into_the_value(tmp_path):
    make(tmp_path, "style: kpi.unit.size=18 kpi.unit.color=muted\n" + KPIS, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    lines = text.splitlines()
    k = next(i for i, ln in enumerate(lines) if ln.startswith("## 売上高"))
    assert lines[k + 1 : k + 3] == ["1,280億円", "前年比 +8%"]
    assert "7.5%" in lines and "3.8" in lines
    assert text.count("{.kpi") == 4 and "{.muted}" not in text


def test_unit_only_trailing_non_digits():
    def u(text, size=20.0):
        out = split_unit(Paragraph(runs=[Run(text=text)]), size)
        return [(r.text, r.size) for r in out.runs]

    assert u("1,280億円") == [("1,280", None), ("億円", 20.0)]
    assert u("7.5%") == [("7.5", None), ("%", 20.0)]
    assert u("約100億円") == [("約100", None), ("億円", 20.0)]
    assert u("100〜120億円") == [("100〜120", None), ("億円", 20.0)]
    assert u("-20億円") == [("-20", None), ("億円", 20.0)]
    assert u("1.5x") == [("1.5", None), ("x", 20.0)]
    assert u("12 pts") == [("12 ", None), ("pts", 20.0)]
    for plain in ("3.8", "N/A", "", "Q3", "あ"):  # no digits, or digits at the end: left alone
        assert u(plain) == [(plain, None)]


def test_unit_keeps_run_marks_across_the_cut():
    p = Paragraph(
        runs=[Run(text="12", bold=True, color="accent"), Run(text="億", italic=True), Run(text="円")]
    )
    out = split_unit(p, 14.0)
    assert [(r.text, r.size, r.bold, r.italic, r.color) for r in out.runs] == [
        ("12", None, True, False, "accent"),
        ("億", 14.0, False, True, None),
        ("円", 14.0, False, False, None),
    ]
    one = split_unit(Paragraph(runs=[Run(text="12億円", bold=True)]), 14.0)
    assert [(r.text, r.size, r.bold) for r in one.runs] == [("12", None, True), ("億円", 14.0, True)]


# --------------------------------------------------------------------------- 3. gating and fuzz-safety


@pytest.mark.parametrize(
    "tok", ["kpi.band=primary", "kpi.unit.size=20", "kpi.band.size=14", "kpi.unit.color=muted"]
)
def test_kpi_tokens_without_a_kpi_card_warn(tmp_path, tok):
    deck, _ = make(tmp_path, f"style: {tok}\n\n# A\n## B\n- x\n")
    bad = warns(deck, "attr-ignored")
    assert bad and "no `.kpi` card" in bad[0].hint


BAD = [
    "kpi.band=nope-color",
    "kpi.band.color=nope-color",
    "kpi.band.size=big",
    "kpi.band.size=-4",
    "kpi.band.size=0",
    "kpi.unit.size=big",
    "kpi.unit.size=-3",
    "kpi.unit.size=0",
    "kpi.unit.color=nope-color",
    "layout.kpi_band_pad=-3",
    "layout.kpi_band_pad=wide",
]


@pytest.mark.parametrize("tok", BAD)
def test_bad_values_never_raise(tmp_path, tok):
    deck, prs = make(tmp_path, f"style: kpi.band=primary kpi.unit.size=20 {tok}\n" + KPIS)
    assert prs.slides[0].shapes
    bad = warns(deck, "bad-token")
    for d in bad:
        assert d.hint


def test_value_without_digits_and_odd_values_build(tmp_path):
    odd = "\n# K\n@kpi\n## A\nN/A\nc\n## B\n億円\nc\n## C\n¥1,280\nc\n## D\n+8%\nc\n## E\n12 pts\nc\n"
    deck, prs = make(tmp_path, "style: kpi.band=primary kpi.unit.size=16\n" + odd)
    assert prs.slides[0].shapes and not warns(deck, "overflow")


def test_unreadable_unit_color_names_its_own_token(tmp_path):
    deck, _ = make(tmp_path, "style: kpi.unit.size=18 kpi.unit.color=#EEEEEE\n" + KPIS)
    hints = [d.hint for d in warns(deck, "contrast")]
    assert hints and any("kpi.unit.color" in h for h in hints), hints


# --------------------------------------------------------------------------- 4. sizes: kpi=


def test_sizes_kpi_sets_the_number_size(tmp_path):
    three = "\n# K\n@kpi\n## A\n12\nc\n## B\n34\nc\n## C\n56\nc\n"
    _, base = make(tmp_path, three, "base")
    _, big = make(tmp_path, "sizes: kpi=48\n" + three, "big")
    sz = lambda prs: max(r.font.size.pt for rs in value_runs(prs) for r in rs)  # noqa: E731
    assert sz(big) == 48 and sz(base) != 48  # a stated size is never grown; the default one is
    _, own = make(tmp_path, "sizes: kpi=48\nstyle: kpi.size=40\n" + three, "own")
    assert sz(own) == 40  # the class token is the more specific one


def test_sizes_kpi_in_a_slide_changes_that_slide_only(tmp_path):
    text = "\n# A\n@kpi\n## A\n1\nc\n## B\n2\nc\n\n# B\n@kpi\nsizes: kpi=50\n## A\n1\nc\n## B\n2\nc\n"
    _, prs = make(tmp_path, text)
    assert 50.0 not in {r.font.size.pt for rs in value_runs(prs, 0) for r in rs}
    assert {r.font.size.pt for rs in value_runs(prs, 1) for r in rs} == {50.0}


def test_sizes_kpi_reaches_the_number_with_a_unit(tmp_path):
    three = (
        "\n# K\n@kpi\nsizes: kpi=60\nstyle: kpi.unit.size=32\n## A\n465億円\nc\n## B\n18%\nc\n## C\n88%\nc\n"
    )
    _, prs = make(tmp_path, three)
    rs = value_runs(prs)[0]
    assert rs[0].font.size.pt == 60 and rs[1].font.size.pt == 32  # the unit is measured at its own size
