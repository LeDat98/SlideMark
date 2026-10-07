"""Raw HTML blocks -> native shapes measured by Chromium (skipped without Playwright/Chromium)."""

from __future__ import annotations

import base64
import io
import sys

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

from slidemark.build import build
from slidemark.htmlnative import convertible, html_to_placed
from slidemark.ir import Image, Shape, Text
from slidemark.render.htmlimg import HtmlRenderer
from slidemark.theme import Theme

from .helpers import needs_soffice

PX = 9525
CARD = (
    '<div style="background:#eef;border:1px solid #33c;border-radius:8px;padding:12px">'
    "<h3>Title</h3><p>Body <b>bold</b></p></div>"
)
GRID = (
    '<div style="display:grid;grid-template-columns:1fr 1fr;gap:20px">'
    '<div style="background:#fee;padding:8px">A</div><div style="background:#efe;padding:8px">B</div></div>'
)
THEME = Theme(name="t")


@pytest.fixture()
def renderer():
    pytest.importorskip("playwright.sync_api")
    r = HtmlRenderer()
    if r.render("<p>x</p>", 40, 20, THEME) is None:
        r.close()
        pytest.skip("Chromium cannot be launched")
    yield r
    r.close()


@pytest.fixture(scope="module")
def chromium():
    """Skip when Chromium is missing. The probe browser is closed: only one sync Playwright per process."""
    pytest.importorskip("playwright.sync_api")
    r = HtmlRenderer()
    try:
        ok = r.render("<p>x</p>", 40, 20, THEME) is not None
    finally:
        r.close()
    if not ok:
        pytest.skip("Chromium cannot be launched")


def _place(renderer, html, w=600, h=300):
    res = html_to_placed(html, 100 * PX, 50 * PX, w * PX, h * PX, THEME, ".", renderer=renderer)
    assert res is not None
    return res


def _deck(body: str, attr: str = "{render=native}") -> str:
    return f"# T\n\n```html{attr}\n{body}\n```\n"


def test_convertible_rules():
    assert convertible(CARD)
    assert convertible(GRID)
    for bad in ("<canvas></canvas>", "<video src=x></video>", "<iframe></iframe>"):
        assert not convertible(bad), bad
    # handled per element now: gradients, transforms, svg, text-transform
    for ok in (
        '<svg width="5"></svg>',
        '<div style="background:linear-gradient(red,blue)">x</div>',
        '<div style="transform:rotate(3deg)">x</div>',
        '<p style="text-transform:uppercase">x</p>',
    ):
        assert convertible(ok), ok


def test_card_becomes_rounded_rect_and_text(renderer):
    items, diags = _place(renderer, CARD)
    assert diags == []
    rect = items[0]
    assert isinstance(rect.element, Shape) and rect.element.shape == "rounded-rect"
    assert rect.style.fill == "#EEEEFF" and rect.style.line == "#3333CC"
    assert rect.style.radius == pytest.approx(6.0)  # 8 px
    assert rect.x == 100 * PX and rect.y == 50 * PX
    texts = [p for p in items if isinstance(p.element, Text)]
    assert [t.element.paragraphs[0].plain for t in texts] == ["Title", "Body bold"]
    title, body = texts
    assert title.element.paragraphs[0].runs[0].bold  # <h3>
    assert [r.bold for r in body.element.paragraphs[0].runs] == [False, True]
    assert title.style.font_size == pytest.approx(18 * 1.17, abs=0.01)  # h3 = 1.17em of the 18 pt theme body
    assert title.style.padding == 0
    for t in texts:  # inside the card, padding respected
        assert t.x >= rect.x + 12 * PX - 1 and t.x + t.w <= rect.x + rect.w + 1
        assert rect.y <= t.y and t.y + t.h <= rect.y + rect.h + 1
    assert title.y < body.y


def test_grid_two_cards_side_by_side(renderer):
    items, _ = _place(renderer, GRID)
    rects = [p for p in items if isinstance(p.element, Shape)]
    assert len(rects) == 2
    a, b = rects
    assert b.x > a.x + a.w - 1 and a.y == b.y
    assert a.style.fill == "#FFEEEE" and b.style.fill == "#EEFFEE"


def test_list_items_are_bullets(renderer):
    items, _ = _place(renderer, "<ul><li>One</li><li>Two<ul><li>Sub</li></ul></li></ul><ol><li>x</li></ol>")
    texts = [p.element for p in items if isinstance(p.element, Text)]
    paras = texts[0].paragraphs
    assert [(p.plain, p.marker, p.level) for p in paras] == [
        ("One", "bullet", 0),
        ("Two", "bullet", 0),
        ("Sub", "bullet", 1),
    ]
    assert texts[1].paragraphs[0].marker == "number"


def test_image_data_uri_and_table_cells(renderer):
    from PIL import Image as PILImage

    buf = io.BytesIO()
    PILImage.new("RGB", (4, 4), "red").save(buf, "PNG")
    uri = "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()
    html = f'<img src="{uri}" width="40" height="40"><table><tr><th>H</th><td>12</td></tr></table>'
    items, _ = _place(renderer, html)
    imgs = [p for p in items if isinstance(p.element, Image)]
    assert len(imgs) == 1 and imgs[0].w == 40 * PX
    cells = [p.element.paragraphs[0].plain for p in items if isinstance(p.element, Text)]
    assert cells == ["H", "12"]


def test_unreachable_content_is_cut_with_warning(renderer):
    _, diags = _place(renderer, '<div style="height:900px;background:#ccc">x</div>', h=100)
    assert [d.rule for d in diags] == ["html-native-overflow"]


def test_empty_html_returns_none(renderer):
    assert html_to_placed("<div></div>", 0, 0, 100 * PX, 50 * PX, THEME, ".", renderer=renderer) is None


def _slide_shapes(path):
    return list(Presentation(str(path)).slides[0].shapes)


def test_build_native_shapes_and_ea_font(tmp_path, chromium):
    out = tmp_path / "a.pptx"
    body = CARD.replace("Body", "日本語テキスト")
    deck = build(_deck(body), out)
    shapes = _slide_shapes(out)
    assert not [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    autos = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert any(s.auto_shape_type == MSO_SHAPE.ROUNDED_RECTANGLE for s in autos)
    boxes = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    assert [b.text_frame.text for b in boxes] == ["Title", "日本語テキスト bold"]
    assert boxes[0].text_frame.margin_left == 0
    xml = boxes[1]._element.xml
    assert "<a:ea " in xml and 'b="1"' in xml
    assert not [d for d in deck.diagnostics if d.level != "info" and not (d.rule or "").startswith("design-")]


def test_grid_build_positions(tmp_path, chromium):
    out = tmp_path / "g.pptx"
    build(_deck(GRID), out)
    rects = [s for s in _slide_shapes(out) if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    assert len(rects) == 2 and rects[1].left > rects[0].left and rects[0].top == rects[1].top


def test_render_image_forces_picture(tmp_path, chromium):
    out = tmp_path / "i.pptx"
    build(_deck(CARD, "{render=image}"), out)
    shapes = _slide_shapes(out)
    assert [s.shape_type for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert not [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]


def test_canvas_falls_back_to_picture(tmp_path, chromium):
    out = tmp_path / "c.pptx"
    build(_deck('<canvas width="100" height="50"></canvas><p>x</p>', ""), out)
    assert [s for s in _slide_shapes(out) if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_default_mode_uses_native_for_convertible_raw(tmp_path, chromium):
    out = tmp_path / "d.pptx"
    build(_deck("<form><p>Only <b>text</b></p></form>", ""), out)  # <form> keeps the subset parser out
    shapes = _slide_shapes(out)
    assert not [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert any(s.has_text_frame and "Only" in s.text_frame.text for s in shapes)


def test_without_playwright_keeps_placeholder(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "playwright", None)
    monkeypatch.setitem(sys.modules, "playwright.sync_api", None)
    out = tmp_path / "n.pptx"
    deck = build(_deck(CARD), out)
    assert [d for d in deck.diagnostics if d.rule == "html-image-unavailable"]
    assert not [s for s in _slide_shapes(out) if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_schema_valid(tmp_path, chromium):
    from slidemark.xsd import validate_pptx

    out = tmp_path / "x.pptx"
    build(_deck(CARD + GRID + "<ul><li>a</li></ul>"), out)
    errors = validate_pptx(out)
    if errors is None:
        pytest.skip("ECMA-376 schemas not available")
    assert errors == []


@needs_soffice
def test_libreoffice_opens(tmp_path, chromium):
    from slidemark.preview import pptx_to_pngs

    out = tmp_path / "l.pptx"
    build(_deck(CARD + GRID), out)
    assert len(pptx_to_pngs(out, tmp_path / "png")) == 1
