"""DL3d lane A: the syntax a foreign deck (the python-pptx open-brief deck) uses and SlideMark lacked.

Mixed run sizes / colours in one line (``[420]{size=28 bold=true color=primary}``), per-card stripe colours
(``box.stripe=a,b,c``), a big number as heading (``@4 num=text``, ``## 01 {.num}``), number tiles
(``@kpi tile``), a band on the cover's bottom edge, and a dark statement panel. Every form: parser (the
words reach the IR), render (the .pptx is reopened and its shapes / runs asserted), importer round trip and
fuzz-safety (a bad value is a diagnostic or ignored, never an exception).
"""

from __future__ import annotations

import re
import tempfile
from pathlib import Path

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from pptx import Presentation
from pptx.util import Emu, Inches

from slidemark import build, parse
from slidemark.importer import import_pptx
from slidemark.importer.runs import mixed_colors, mixed_sizes, span
from slidemark.layout import measure
from slidemark.template import deck_theme

TMP = Path(tempfile.mkdtemp(prefix="sm-dl3d-"))
HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#222B36 primary=#142B4D secondary=#1F5FA8 accent=#B35C00 teal=#1B8A8F "
    "muted=#59626E surface=#EEF2F7 border=#C9D2DE\n"
)


def make(tmp_path, text: str, name: str = "d"):
    deck = build(HEAD + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def warns(deck, rule):
    return [d for d in deck.diagnostics if d.rule == rule]


def shapes(prs, i, name):
    return [s for s in prs.slides[i].shapes if s.name == name]


def fill_hex(shp) -> str:
    return str(shp.fill.fore_color.rgb)


def all_runs(prs, i):
    out = []
    for s in prs.slides[i].shapes:
        if s.has_text_frame:
            out += [r for p in s.text_frame.paragraphs for r in p.runs]
        if s.has_table:
            for row in s.table.rows:
                for c in row.cells:
                    out += [r for p in c.text_frame.paragraphs for r in p.runs]
    return out


def norm(text: str) -> str:
    """Text as the author wrote it: the renderer adds U+2060 number joiners and no-break spaces."""
    return text.replace("\u2060", "").replace("\u00a0", " ")


def sizes(prs, i, text: str) -> list[float]:
    return [r.font.size.pt for r in all_runs(prs, i) if text in norm(r.text) and r.font.size]


# --------------------------------------------------------------------------- 1. spans: parser

SPANS = (
    "[売上高]{size=14 bold=true} [420億円]{size=28 bold=true color=primary} [（2026年度）]{size=12 .muted}"
)


def runs_of(text: str):
    deck = parse(
        HEAD + "\n# T\n- first item\n- " + text + "\n"
    )  # (a lone paragraph would be a cover subtitle)
    el = deck.slides[0].elements[0]
    return [r for r in el.paragraphs[-1].runs]


def test_span_takes_size_bold_color_and_class():
    by = {r.text: r for r in runs_of(SPANS)}
    big = by["420億円"]
    assert (big.size, big.exact, big.bold, big.color) == (28.0, True, True, "primary")
    assert by["売上高"].size == 14 and by["売上高"].bold
    note = by["（2026年度）"]
    assert (note.size, note.exact, note.color) == (12.0, True, "muted")
    assert not by[" "].exact if " " in by else True


def test_span_size_accepts_pt_and_italic():
    r = runs_of("[x]{size=18pt italic=on}")[0]
    assert (r.size, r.exact, r.italic) == (18.0, True, True)
    assert runs_of("[x]{size=7.5}")[0].size == 7.5


@pytest.mark.parametrize("bad", ["size=", "size=12pt pt", "size=a b"])
def test_an_invalid_attribute_list_stays_text(bad):
    r = runs_of(f"[x]{{{bad}}}")[0]
    assert r.size is None and not r.exact and "[x]" in r.text


@pytest.mark.parametrize("bad", ["abc", "0", "-3", "9999", "1e999", "nan"])
def test_a_bad_span_size_is_ignored(bad):
    r = runs_of(f"[x]{{size={bad}}}")[0]
    assert r.size is None and not r.exact and r.text == "x"


def test_plain_text_and_other_spans_stay_not_exact():
    assert not any(r.exact for r in runs_of("plain **bold** [x]{.danger} ==y=="))


def test_a_heading_ending_in_a_span_keeps_its_span_and_its_own_attributes():
    deck = parse(
        HEAD
        + "\n# T\n## 売上高 [420億円]{size=28} [（同 11%）]{size=12 .muted}\n本文\n## B [x]{size=9} {.hero}\n"
    )
    a, b = deck.slides[0].elements
    assert [r.size for r in a.title.paragraphs[0].runs if r.size] == [28.0, 12.0]
    assert "hero" not in a.classes and "muted" not in a.classes
    assert "hero" in b.classes and [r.size for r in b.title.paragraphs[0].runs if r.size] == [9.0]


def test_a_br_in_a_table_cell_does_not_warn_but_outside_a_table_it_does():
    ok = parse(HEAD + "\n# T\n| a | b |\n|-|-|\n| x<br>y | z |\n")
    assert not warns(ok, "html-br")
    assert warns(parse(HEAD + "\n# T\nline<br>two\n"), "html-br")
    cell = ok.slides[0].elements[0].rows[1][0]
    assert [r.text for r in cell.paragraphs[0].runs] == ["x", "\n", "y"]


# --------------------------------------------------------------------------- 1. spans: render and layout


def test_spans_are_drawn_at_their_exact_size_in_a_paragraph(tmp_path):
    _, prs = make(tmp_path, "\n# T\n" + SPANS + "\n")
    assert sizes(prs, 0, "420億円") == [28.0] and sizes(prs, 0, "売上高") == [14.0]
    assert sizes(prs, 0, "（2026年度）") == [12.0]


def test_a_sparse_slide_does_not_grow_a_text_with_pinned_spans(tmp_path):
    plain = "\n# T\n- 売上高 420億円\n- 営業利益率 20%\n"
    deck, a = make(tmp_path, plain, "plain")
    _, b = make(tmp_path, "\n# T\n- [売上高]{size=14} 420億円\n- 営業利益率 20%\n", "pin")
    body = deck_theme(deck)[0].sizes["body"]
    assert max(sizes(a, 0, "営業利益率")) > body  # the sparse-slide pass grows the plain list ...
    assert sizes(b, 0, "売上高") == [14.0]  # ... a pinned span stays exact ...
    assert max(sizes(b, 0, "営業利益率")) == pytest.approx(body)  # ... and the text it sits in is not grown


def test_spans_in_a_list_row_a_table_cell_a_box_heading_and_a_kpi_value(tmp_path):
    md = (
        "\n# T\n@rows plain\n- [売上高]{size=14} [420億円]{size=28 bold=true}\n"
        "- [利益]{size=14} [20%]{size=28}\n"
        "\n# U\n| a | b |\n|-|-|\n"
        "| x | [月額]{size=10 .muted}<br>[9,800円]{size=22 bold=true color=secondary} |\n"
        "\n# V\n@1\n## [売上高]{size=14} [420億円]{size=30 color=secondary}\n"
        "\n# W\n@kpi\n## 売上\n[1,280]{size=40}[億円]{size=14}\n前年比 +8%\n"
    )
    _, prs = make(tmp_path, md)
    assert sizes(prs, 0, "420億円") == [28.0] and sizes(prs, 0, "売上高") == [14.0]
    assert sizes(prs, 1, "9,800円") == [22.0] and sizes(prs, 1, "月額") == [10.0]
    assert sizes(prs, 2, "420億円") == [30.0] and sizes(prs, 2, "売上高") == [14.0]
    assert sizes(prs, 3, "1,280") == [40.0] and sizes(prs, 3, "億円") == [14.0]


def test_a_line_is_as_tall_as_its_largest_pinned_span():
    from slidemark.ir import Paragraph, Run, Style

    st_ = Style(font_size=12)
    small = [Paragraph(runs=[Run(text="売上高 420億円")])]
    big = [Paragraph(runs=[Run(text="売上高 "), Run(text="420億円", size=36, exact=True)])]
    w = 400 * 12700
    assert measure.paragraphs_height(big, w, st_) > 2 * measure.paragraphs_height(small, w, st_)
    # a paragraph whose every run is pinned is measured at the largest pinned size, not at the style size
    allp = [Paragraph(runs=[Run(text="a", size=30, exact=True), Run(text="b", size=10, exact=True)])]
    only = [Paragraph(runs=[Run(text="ab", size=30, exact=True)])]
    assert measure.paragraphs_height(allp, w, st_) == pytest.approx(
        measure.paragraphs_height(only, w, st_), rel=0.2
    )
    # the wider pinned run wraps earlier than the same text at the style size
    long = "売" * 30
    narrow = 200 * 12700
    wide = [Paragraph(runs=[Run(text=long, size=24, exact=True)])]
    thin = [Paragraph(runs=[Run(text=long)])]
    assert measure.paragraphs_height(wide, narrow, st_) > measure.paragraphs_height(thin, narrow, st_)


def test_a_pinned_span_below_the_minimum_size_warns_tiny_text(tmp_path):
    deck, _ = make(tmp_path, "\n# T\ntext [fine print]{size=5}\n")
    assert warns(deck, "tiny-text")


def test_a_small_pinned_span_is_judged_at_its_own_size_for_contrast(tmp_path):
    big, _ = make(tmp_path, "\n# T\n- first\n- [big]{size=30 color=#B8BCC4}\n", "big")
    small, _ = make(tmp_path, "\n# T\n- first\n- [small]{size=12 color=#B8BCC4}\n", "small")
    hint_big = [d.hint for d in warns(big, "contrast") if "big" in d.message]
    hint_small = [d.hint for d in warns(small, "contrast") if "small" in d.message]
    assert hint_big and hint_small  # both are unreadable ...
    assert (
        "3:1" in hint_big[0] and "4.5:1" in hint_small[0]
    )  # ... the fix is sized for the span, not the line


# --------------------------------------------------------------------------- 1. spans: importer writer


def test_importer_span_writer_pins_mixed_sizes_and_names_colours():
    from slidemark.importer.read import RunT

    runs = [
        RunT("売上高", bold=True, size=14, color="222B36"),
        RunT("420億円", bold=True, size=28, color="1F5FA8"),
    ]
    assert mixed_sizes(runs) and mixed_colors(runs) and not mixed_sizes(runs[:1])
    names = {"1F5FA8": "secondary"}
    assert span("420億円", runs[1], pin=True, tint=True, names=names, fg="222B36") == (
        "[420億円]{size=28 bold=true color=secondary}"
    )
    assert (
        span("売上高", runs[0], pin=True, tint=True, names=names, fg="222B36")
        == "[売上高]{size=14 bold=true}"
    )
    assert span("x", RunT("x"), pin=False) is None  # a uniform run needs no span


def test_spans_round_trip_through_import(tmp_path):
    md = (
        "\n# T\n@rows plain\n- "
        + SPANS
        + "\n\n# U\n| a | b |\n|-|-|\n"
        + "| x | [月額]{size=10 .muted}<br>[9,800円]{size=22 bold=true color=secondary} |\n"
    )
    make(tmp_path, md, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "[420億円]{size=28 bold=true color=primary}" in text
    assert "[（2026年度）]{size=12 .muted}" in text
    assert "[9,800円]{size=22 bold=true color=secondary}" in text
    again = build(text, tmp_path / "rt2.pptx")
    assert not [
        d for d in again.diagnostics if d.level == "warning" and d.rule not in ("design-slide", "design-none")
    ]
    prs = Presentation(str(tmp_path / "rt2.pptx"))
    assert sizes(prs, 0, "420億円") == [28.0] and sizes(prs, 1, "9,800円") == [22.0]


def test_a_kpi_value_keeps_its_unit_run_without_spans(tmp_path):
    make(tmp_path, "style: kpi.unit.size=20\n\n# T\n@kpi\n## 売上\n1,280億円\n前年比 +8%\n", "u")
    text, _ = import_pptx(tmp_path / "u.pptx")
    assert "1,280億円" in text and "{size=" not in text.split("# T")[1]


# --------------------------------------------------------------------------- 2. stripes

BOXES = "\n# T\n@3\n## A\n- a\n## B\n- b\n## C\n- c\n"


def test_box_stripes_cycle_over_the_cards_in_order(tmp_path):
    deck, prs = make(
        tmp_path, "style: box.stripe=primary,secondary box.stripe_h=8pt\n" + BOXES + "## D\n- d\n"
    )
    st_ = shapes(prs, 0, "Stripe")
    assert [fill_hex(s) for s in st_] == ["142B4D", "1F5FA8", "142B4D", "1F5FA8"]
    cards = [s for s in prs.slides[0].shapes if s.name.startswith("Card")]
    assert all(
        s.height == Emu(Inches(8 / 72)) and s.width == c.width and s.left == c.left
        for s, c in zip(st_, cards, strict=True)
    )
    assert not warns(deck, "attr-ignored")


def test_stripe_side_left_and_a_card_override_keeps_its_place_in_the_cycle(tmp_path):
    md = "style: box.stripe=primary,secondary,teal box.stripe.side=left\n" + BOXES.replace(
        "## B", "## B {stripe=accent}"
    )
    _, prs = make(tmp_path, md)
    st_ = shapes(prs, 0, "Stripe")
    assert [fill_hex(s) for s in st_] == ["142B4D", "B35C00", "1B8A8F"]
    cards = [s for s in prs.slides[0].shapes if s.name.startswith("Card")]
    assert all(
        s.width < s.height and s.top == c.top and s.height == c.height
        for s, c in zip(st_, cards, strict=True)
    )


def test_item_cards_and_kpi_cards_have_their_own_stripes(tmp_path):
    md = (
        "style: item.stripe=primary,teal kpi.stripe=secondary,accent kpi.stripe.side=left\n"
        "\n# T\n@1 items\n## L\n- a\n- b\n- c\n\n# U\n@kpi\n## A\n1\nx\n## B\n2\ny\n"
    )
    _, prs = make(tmp_path, md)
    assert [fill_hex(s) for s in shapes(prs, 0, "Stripe")] == ["142B4D", "1B8A8F", "142B4D"]
    kp = shapes(prs, 1, "rule")
    assert [fill_hex(s) for s in kp] == ["1F5FA8", "B35C00"] and kp[0].width < kp[0].height


def test_a_single_kpi_stripe_is_unchanged(tmp_path):
    _, prs = make(tmp_path, "style: kpi.stripe=secondary\n\n# T\n@kpi\n## A\n1\nx\n## B\n2\ny\n")
    assert [fill_hex(s) for s in shapes(prs, 0, "rule")] == ["1F5FA8", "1F5FA8"]


def test_a_card_without_fill_or_border_gets_no_stripe_and_plain_boxes_are_skipped(tmp_path):
    _, prs = make(tmp_path, "style: box.stripe=primary\n\n# T\n@2\n## A {.plain}\n- a\n## B\n- b\n")
    assert len(shapes(prs, 0, "Stripe")) == 1


@pytest.mark.parametrize(
    "tok",
    [
        "box.stripe=nope",
        "box.stripe=",
        "box.stripe.side=middle",
        "kpi.stripe.side=",
        "item.stripe=primary,xx",
    ],
)
def test_a_bad_stripe_token_warns_with_a_hint_and_builds(tmp_path, tok):
    deck, prs = make(tmp_path, f"style: {tok}\n" + BOXES)
    bad = warns(deck, "bad-token")
    assert bad and bad[0].hint and len(prs.slides) == 1


def test_a_stripe_on_something_that_is_not_a_card_is_reported(tmp_path):
    deck, _ = make(tmp_path, "\n# T\n{stripe=teal}\n| a |\n|-|\n| b |\n")
    assert [d for d in warns(deck, "attr-ignored") if "stripe" in d.message]
    ok, _ = make(tmp_path, "\n# T\n## A {stripe=teal}\n- a\n", "ok")
    assert not [d for d in warns(ok, "attr-ignored") if "stripe" in d.message]


def test_a_stripe_token_without_cards_is_reported(tmp_path):
    deck, _ = make(tmp_path, "style: box.stripe=primary\n\n# T\n- a\n- b\n")
    assert [d for d in warns(deck, "attr-ignored") if "box.stripe" in d.message]


def test_stripe_shapes_are_dropped_by_the_importer_and_tokens_survive(tmp_path):
    make(tmp_path, "style: box.stripe=primary,secondary\n" + BOXES, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "box_stripe=primary,secondary" in text.split("# T")[0] and "Stripe" not in text
    deck = build(text, tmp_path / "rt2.pptx")
    assert len(shapes(Presentation(str(tmp_path / "rt2.pptx")), 0, "Stripe")) == 3
    assert not [d for d in deck.diagnostics if d.level == "warning" and not d.rule.startswith("design")]


# --------------------------------------------------------------------------- 3. number headings and tiles

NUMS = "\n# T\n@3 num=text\n## 売上を2倍\nbody a\n## 時間を削減\nbody b\n## 現場\nbody c\n"


def test_num_text_words_reach_the_slide():
    deck = parse(HEAD + NUMS)
    assert {"num", "num-text"} <= set(deck.slides[0].classes)
    assert "num-text" not in parse(HEAD + NUMS.replace("num=text", "num")).slides[0].classes
    assert "num-text" not in parse(HEAD + NUMS.replace("num=text", "num=badge")).slides[0].classes
    bad = parse(HEAD + NUMS.replace("num=text", "num=huge"))
    assert [d for d in bad.diagnostics if d.rule == "bad-attr"]


def test_num_text_draws_big_coloured_numbers_that_cycle(tmp_path):
    md = "style: box.num.size=36 box.num.color=primary,secondary box.stripe=teal\n" + NUMS
    deck, prs = make(tmp_path, md)
    heads = shapes(prs, 0, "Number text")
    assert len(heads) == 3
    first = [p.runs[0] for s in heads for p in s.text_frame.paragraphs[:1]]
    assert [norm(r.text) for r in first] == ["01", "02", "03"]
    assert all(r.font.size.pt >= 36 for r in first[:1])
    assert [str(r.font.color.rgb) for r in first] == ["142B4D", "1F5FA8", "142B4D"]
    second = heads[0].text_frame.paragraphs[1].runs[0]
    assert norm(second.text) == "売上を2倍" and second.font.size.pt < first[0].font.size.pt
    assert not [s for s in prs.slides[0].shapes if s.name.startswith("Num ")]  # no circle badges
    assert not warns(deck, "attr-ignored")


def test_num_text_takes_the_stripe_colour_and_the_format_token(tmp_path):
    _, prs = make(tmp_path, 'style: box.stripe=teal,accent box.num.text="第{n}章"\n' + NUMS)
    first = [s.text_frame.paragraphs[0].runs[0] for s in shapes(prs, 0, "Number text")]
    assert [norm(r.text) for r in first] == ["第1章", "第2章", "第3章"]
    assert [str(r.font.color.rgb) for r in first] == ["1B8A8F", "B35C00", "1B8A8F"]


def test_a_num_class_heading_is_the_number_itself(tmp_path):
    md = "\n# T\n@2\n## 01 {.num}\nbody\n## 2.1兆円 {.num}\nbody 2\n"
    deck, prs = make(tmp_path, "style: box.num.size=40 box.num.color=secondary\n" + md)
    heads = shapes(prs, 0, "Number heading")
    assert [norm(h.text_frame.text) for h in heads] == ["01", "2.1兆円"]
    r = heads[0].text_frame.paragraphs[0].runs[0]
    assert r.font.size.pt == 40 and str(r.font.color.rgb) == "1F5FA8"
    assert not warns(deck, "attr-ignored")


def test_num_headings_round_trip_through_import(tmp_path):
    make(tmp_path, "style: box.stripe=primary,secondary\n" + NUMS, "n1")
    t1, _ = import_pptx(tmp_path / "n1.pptx")
    assert "num=text" in t1 and "## 売上を2倍" in t1 and "## 01" not in t1
    make(tmp_path, "\n# T\n@2\n## 01 {.num}\nbody\n## 02 {.num}\nbody 2\n", "n2")
    t2, _ = import_pptx(tmp_path / "n2.pptx")
    assert "## 01 {.num}" in t2 and "## 02 {.num}" in t2


TILES = (
    "\n# T\n@2x2 kpi tile\n## 2.1兆円\n国内の市場は2.1兆円へ\n## 38%\n導入率は38%\n"
    "## +22%\n前年比+22%\n## 41%\nシェア41%\n"
)


def test_a_kpi_tile_is_a_number_over_its_description(tmp_path):
    md = "style: kpi.stripe=teal,secondary kpi.stripe.side=left kpi.stripe_h=0.08in kpi.size=40\n" + TILES
    deck, prs = make(tmp_path, md)
    tiles = shapes(prs, 0, "Tile")
    assert len(tiles) == 4
    num = tiles[0].text_frame.paragraphs[0].runs[0]
    assert num.text.replace(" ", "").replace("⁠", "") == "2.1兆円" and num.font.size.pt >= 30
    assert [str(t.text_frame.paragraphs[0].runs[0].font.color.rgb) for t in tiles] == [
        "1B8A8F",
        "1F5FA8",
        "1B8A8F",
        "1F5FA8",
    ]
    cap = tiles[0].text_frame.paragraphs[1].runs[0]
    assert cap.font.size.pt < num.font.size.pt
    assert [fill_hex(s) for s in shapes(prs, 0, "rule")] == ["1B8A8F", "1F5FA8", "1B8A8F", "1F5FA8"]
    assert all(t.text_frame.paragraphs[0].alignment in (None, 1) for t in tiles)  # left aligned
    assert not [d for d in deck.diagnostics if d.level == "warning" and not d.rule.startswith("design")]


def test_a_tile_card_without_the_slide_word(tmp_path):
    _, prs = make(tmp_path, "\n# T\n@2\n## 9% {.kpi .tile}\ndesc\n## 8% {.kpi}\n7%\ncaption\n")
    assert len(shapes(prs, 0, "Tile")) == 1


def test_tiles_round_trip_through_import(tmp_path):
    make(tmp_path, "style: kpi.stripe=teal kpi.stripe.side=left\n" + TILES, "t")
    text, _ = import_pptx(tmp_path / "t.pptx")
    assert "## 2.1兆円 {.kpi .tile}" in text and "国内の市場は2.1兆円へ" in text
    build(text, tmp_path / "t2.pptx")
    assert len(shapes(Presentation(str(tmp_path / "t2.pptx")), 0, "Tile")) == 4


# --------------------------------------------------------------------------- 4. cover


COVER = (
    "style: cover.band_h=70% cover.band=none cover.bar=accent cover.bottom.bar=secondary "
    "cover.bottom.bar_h=0.3in cover.footer=off\nfooter: ACME\n"
)


def test_cover_bottom_bar_is_a_band_on_the_bottom_edge(tmp_path):
    deck, prs = make(tmp_path, COVER + "\n# 中期経営計画\n@cover bg=primary dark\n2027年4月\n")
    bars = [s for s in shapes(prs, 0, "rule") if s.top + s.height == prs.slide_height]
    assert len(bars) == 1
    assert fill_hex(bars[0]) == "1F5FA8" and bars[0].width == prs.slide_width
    assert abs(bars[0].height - Inches(0.3)) < 2000
    assert not [d for d in deck.diagnostics if d.level == "warning" and not d.rule.startswith("design")]
    bar = [s for s in shapes(prs, 0, "rule") if s.width < Inches(0.3) and s.height < prs.slide_height * 0.5]
    assert bar and fill_hex(bar[0]) == "B35C00"  # the bar beside the title is still there


def test_cover_bottom_bar_works_on_the_plain_cover_and_not_on_content_slides(tmp_path):
    md = "style: cover.bottom.bar=teal\n\n# Plan\nsub\n\n# Body\n- a\n- b\n"
    _, prs = make(tmp_path, md)
    assert [fill_hex(s) for s in shapes(prs, 0, "rule")] == ["1B8A8F"]
    assert not shapes(prs, 1, "rule")


def test_cover_bottom_bar_token_without_a_cover_is_reported(tmp_path):
    deck, _ = make(tmp_path, "style: cover.bottom.bar=teal\n\n# A\n- a\n- b\n\n# B\n- c\n- d\n")
    assert [d for d in warns(deck, "attr-ignored") if "cover.bottom.bar" in d.message]


# --------------------------------------------------------------------------- 5. dark statement panel

PANEL = (
    "\n# T\n@1:2\n## スタンダードプラン {fill=primary valign=middle}\n[9,800円]{size=32 bold=true}\n▼\n"
    "[10,800円]{size=32 bold=true color=accent}\n2027年10月から\n## 方針 {.plain}\n- a\n- b\n"
)


def test_a_dark_panel_holds_lines_of_different_sizes_in_a_readable_ink(tmp_path):
    deck, prs = make(tmp_path, PANEL)
    assert not [d for d in deck.diagnostics if d.rule in ("contrast", "overflow")]
    assert sizes(prs, 0, "9,800円") == [32.0] and sizes(prs, 0, "10,800円") == [32.0]
    ink = {
        str(r.font.color.rgb) for r in all_runs(prs, 0) if "9,800" in norm(r.text) or "スタンダード" in r.text
    }
    assert ink == {"FFFFFF"}  # heading and text turn light on the dark fill without a color=
    accent = [r for r in all_runs(prs, 0) if "10,800円" in norm(r.text)][0]
    assert str(accent.font.color.rgb) == "B35C00"  # a span colour wins


def test_a_box_with_its_own_fill_keeps_readable_text_without_color(tmp_path):
    deck, _ = make(tmp_path, "\n# T\n@2\n## A {fill=primary}\ntext\n## B {fill=#FFF3CD}\ntext\n")
    assert not warns(deck, "contrast")


def test_a_light_box_fill_changes_nothing(tmp_path):
    _, a = make(tmp_path, "\n# T\n@2\n## A {fill=#FFF3CD}\ntext\n## B\ntext\n", "a")
    cols = {str(r.font.color.rgb) for r in all_runs(a, 0)}
    assert cols == {"222B36"}  # the theme text colour on both boxes


# --------------------------------------------------------------------------- the example and fuzz


def test_example_25_builds_clean_and_states_its_design(tmp_path):
    md = Path(__file__).resolve().parent.parent / "examples" / "25-roundtrip-forms.md"
    deck = build(md, tmp_path / "e.pptx")
    assert not [d for d in deck.diagnostics if d.level == "warning"], [d.message for d in deck.diagnostics]
    prs = Presentation(str(tmp_path / "e.pptx"))
    assert len(prs.slides) == 7 and sizes(prs, 3, "420億円") == [30.0]
    assert len(shapes(prs, 2, "Tile")) == 4 and len(shapes(prs, 1, "Number text")) == 3


@settings(max_examples=60, deadline=None, suppress_health_check=[HealthCheck.function_scoped_fixture])
@given(
    size=st.text(alphabet="0123456789.-+eEpt nainfINF", max_size=8),
    color=st.sampled_from(["primary", "#zzz", "", "teal", "1e999", "none"]),
    body=st.text(alphabet="ab 12億円{}[]()*=_\\", min_size=0, max_size=12),
    flag=st.sampled_from(["true", "off", "maybe", ""]),
)
def test_span_attributes_never_raise(size, color, body, flag):
    for where in ("{b}", "- {b}", "## {b}\nx", "| a |\n|-|\n| {b} |"):
        text = where.replace("{b}", f"[{body}]{{size={size} bold={flag} color={color}}}")
        deck = parse(HEAD + "\n# T\n" + text + "\n")
        build_dir = TMP / "fz"
        build_dir.mkdir(exist_ok=True)
        build(deck, build_dir / "f.pptx")


@settings(max_examples=40, deadline=None)
@given(
    side=st.sampled_from(["top", "left", "right", "bottom", "middle", ""]),
    cols=st.sampled_from(["primary", "primary,teal", "nope", ",,", "#12", "a,b,c,d,e,f,g"]),
    h=st.sampled_from(["6pt", "0", "-3pt", "9in", "abc", "100%"]),
    which=st.sampled_from(["box", "item", "kpi"]),
)
def test_stripe_tokens_never_raise(side, cols, h, which):
    deck = parse(
        HEAD
        + f"style: {which}.stripe={cols} {which}.stripe_h={h} {which}.stripe.side={side}\n"
        + "\n# T\n@kpi tile\n## 1\nx\n## 2\ny\n\n# U\n@2 items num=text\n## A\n- a\n## B {stripe=teal}\n- b\n"
    )
    build(deck, TMP / "fs.pptx")
    assert deck_theme(deck) is not None
    assert re.fullmatch(r".*", "")


def test_a_panel_round_trip_keeps_each_line_size_and_the_odd_colour(tmp_path):
    make(tmp_path, PANEL, "p")
    text, _ = import_pptx(tmp_path / "p.pptx")
    body = text.split("# T")[1]
    assert "[9,800円]{size=32 bold=true}" in body
    assert "[10,800円]{size=32 bold=true color=accent}" in body
    assert "color=#FFFFFF" not in body  # the ink of the panel is the block's own colour: not written
    again = build(text, tmp_path / "p2.pptx")
    assert not [d for d in again.diagnostics if d.rule in ("contrast", "overflow")]
    assert sizes(Presentation(str(tmp_path / "p2.pptx")), 0, "9,800円") == [32.0]


def test_num_badges_take_a_colour_list_too(tmp_path):
    _, prs = make(
        tmp_path, "style: box.num.fill=primary,teal box.num.color=bg\n" + BOXES.replace("@3", "@3 num")
    )
    badges = [s for s in prs.slides[0].shapes if s.name.startswith("Num ")]
    assert [fill_hex(s) for s in badges] == ["142B4D", "1B8A8F", "142B4D"]
