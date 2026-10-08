"""Design wave 3, lane C (DL2 part 2): the decisions DESIGN_COVERAGE.md marked missing, plus ``h=`` / ``y=`` /
``w=`` on KPI cards in a row. Each form: parser (IR / theme), render (the .pptx is reopened, XML asserted),
importer round trip where the form leaves shapes behind, and fuzz-safety (a bad value is a diagnostic with a
hint, never a raise)."""

from __future__ import annotations

import tempfile
from pathlib import Path

from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import build, parse
from slidemark.importer import import_pptx
from slidemark.template import deck_theme

HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#222B36 primary=#142B4D secondary=#1F5FA8 accent=#E09F1F teal=#2A9D8F "
    "muted=#59626E surface=#EEF2F7 border=#C9D2DE\n"
)
EMU_IN = 914400


def make(tmp_path, text: str, name: str = "d"):
    deck = build(HEAD + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def shapes(prs, i, prefix=""):
    return [s for s in prs.slides[i].shapes if s.name.lower().startswith(prefix.lower())]


def warns(deck, rule):
    return [d for d in deck.diagnostics if d.rule == rule]


def txt(shape) -> str:
    return shape.text_frame.text.replace("\u2060", "").replace("\xa0", " ")


def chart_of(prs, i=0):
    return next(s for s in prs.slides[i].shapes if s.has_chart).chart


# --------------------------------------------------------------------------- 1. KPI divider rule

KPIS = (
    "\n# Goals\n> lead\n"
    "## Revenue {.kpi .hero}\n1,280\nvs last year +8%\n"
    "## Profit {.kpi}\n96\nvs last year +12%\n"
    "## Margin {.kpi}\n7.5%\n+0.3pt\n"
)


def test_kpi_rule_token_reaches_the_theme():
    d = parse(HEAD + "style: kpi.rule=border kpi.rule_h=2pt kpi.rule_w=60%\n" + KPIS)
    th, _ = deck_theme(d, ".")
    assert th.kpi_rule == "border" and th.kpi_rule_h == "2pt" and th.kpi_rule_w == "60%"
    assert not [x for x in d.diagnostics if x.rule == "bad-token"]


def test_kpi_rule_is_off_by_default(tmp_path):
    _, prs = make(tmp_path, KPIS)
    assert not shapes(prs, 0, "rule") and not shapes(prs, 0, "KPI ")
    assert len(shapes(prs, 0, "Text")) == 3  # number and caption share one frame, as before


def test_kpi_rule_splits_the_caption_and_draws_the_rule_between(tmp_path):
    deck, prs = make(tmp_path, "style: kpi.rule=border kpi.rule_h=1pt\n" + KPIS)
    assert not warns(deck, "bad-token") and not warns(deck, "overflow")
    values, caps, rules = shapes(prs, 0, "KPI value"), shapes(prs, 0, "KPI caption"), shapes(prs, 0, "rule")
    cards = shapes(prs, 0, "Card")
    assert len(values) == len(caps) == len(rules) == len(cards) == 3
    assert [txt(v) for v in values] == ["1,280", "96", "7.5%"]
    assert txt(caps[0]) == "vs last year +8%"
    for card, v, r, c in zip(cards, values, rules, caps, strict=True):
        assert r.height == 12700  # 1pt
        assert v.top + v.height <= r.top and r.top + r.height <= c.top  # value, rule, caption
        assert card.top <= v.top and c.top + c.height <= card.top + card.height + 2
        assert card.left <= r.left and r.left + r.width <= card.left + card.width
        assert str(r.fill.fore_color.rgb) == "C9D2DE"
    assert len({r.top for r in rules}) == 1  # one row: the rules line up across the cards
    assert len({c.top for c in caps}) == 1


def test_kpi_rule_width_and_height_tokens(tmp_path):
    _, prs = make(tmp_path, "style: kpi.rule=primary kpi.rule_h=3pt kpi.rule_w=50%\n" + KPIS)
    for r, v in zip(shapes(prs, 0, "rule"), shapes(prs, 0, "KPI value"), strict=True):
        assert r.height == 3 * 12700
        assert abs(r.width - v.width // 2) <= 2
        assert abs((r.left + r.width / 2) - (v.left + v.width / 2)) <= 2  # centred


def test_kpi_rule_skips_a_card_without_a_caption(tmp_path):
    _, prs = make(tmp_path, "style: kpi.rule=border\n\n# T\n## A {.kpi}\n5\n## B {.kpi}\n6\ncaption\n")
    assert len(shapes(prs, 0, "rule")) == 1


def test_kpi_rule_bad_value_warns_with_a_hint(tmp_path):
    deck, prs = make(tmp_path, "style: kpi.rule=nope-color kpi.rule_h=wide\n" + KPIS)
    bad = warns(deck, "bad-token")
    assert bad and all(d.hint for d in bad)
    assert prs.slides[0].shapes  # the deck still builds


def test_kpi_rule_round_trip_folds_the_boxes_back(tmp_path):
    make(tmp_path, "style: kpi.rule=border\n" + KPIS, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert text.count("{.kpi") == 3
    assert "vs last year +8%" in text and "+0.3pt" in text
    assert "KPI caption" not in text and "rule" not in text.lower().replace("kpi_rule", "")
    lines = text.splitlines()
    k = next(i for i, ln in enumerate(lines) if ln.startswith("## Revenue"))
    assert lines[k + 1 : k + 3] == ["1,280", "vs last year +8%"]  # number then caption, one block


# --------------------------------------------------------------------------- 2. exact table row height

TABLE = (
    "| 事業 | 2026 | 2027 |\n|-|-|-|\n| 冷凍 | 420 | 465 |\n| 調味料 | 310 | 325 |\n| 飲料 | 255 | 270 |\n"
)


def row_heights(prs, i=0):
    gf = next(s for s in prs.slides[i].shapes if s.has_table)
    return [r.height for r in gf.table.rows], gf


def test_rowh_parses_per_table_and_per_row():
    d = parse(HEAD + "\n# T\n{rowh=0.8in,1.05in}\n" + TABLE)
    t = next(e for e in d.slides[0].elements if e.type == "table")
    assert t.rowh == ["0.8in", "1.05in"]
    d2 = parse(HEAD + "\n# T\n{rowh=0.6in}\n" + TABLE)
    assert next(e for e in d2.slides[0].elements if e.type == "table").rowh == ["0.6in"]
    assert not [x for x in d.diagnostics + d2.diagnostics if x.rule == "bad-table-option"]


def test_rowh_pins_every_row_in_emu_and_the_frame(tmp_path):
    deck, prs = make(tmp_path, "\n# T\n{rowh=0.8in,1.05in}\n" + TABLE)
    assert not warns(deck, "overflow") and not warns(deck, "bad-table-option")
    rh, gf = row_heights(prs)
    assert rh == [round(0.8 * EMU_IN)] + [round(1.05 * EMU_IN)] * 3
    assert gf.height == sum(rh)


def test_rowh_one_value_pins_all_rows(tmp_path):
    _, prs = make(tmp_path, "\n# T\n{rowh=0.5in}\n" + TABLE)
    assert row_heights(prs)[0] == [round(0.5 * EMU_IN)] * 4


def test_unpinned_table_is_stretched_so_the_pin_is_a_real_difference(tmp_path):
    _, prs = make(tmp_path, "\n# T\n" + TABLE)
    assert any(h != round(0.5 * EMU_IN) for h in row_heights(prs)[0])


def test_pinned_rows_survive_every_stretch_and_growth_pass(tmp_path):
    pin = "{rowh=0.5in}\n"
    md = (
        "\n# lone sparse table\n"
        + pin
        + TABLE
        + "\n# with a conclusion bar\n> bar\n"
        + pin
        + TABLE
        + "\n# beside a chart\n"
        + pin
        + TABLE
        + "```column\n,a,b\nx,1,2\n```\n"
        + "\n# under chevrons\n@3 chevron\n## a\n## b\n## c\n"
        + pin
        + TABLE
        + "\n# beside cards\n## Left\n- text\n- text\n## Right\n- text\n"
        + pin
        + TABLE
    )
    deck, prs = make(tmp_path, md)
    assert not warns(deck, "overflow")
    for i in range(len(prs.slides)):
        got, gf = next(((r, g) for r, g in [row_heights(prs, i)] if r), (None, None))
        assert got == [round(0.5 * EMU_IN)] * 4, (i, got)
        assert gf.height == sum(got), i


def test_rowh_pinned_table_does_not_grow_its_text(tmp_path):
    base = "\n# T\n## Left\n- a\n- b\n" + "{rowh=1in}\n" + TABLE
    _, pinned = make(tmp_path, base, "p")
    _, free = make(tmp_path, base.replace("{rowh=1in}\n", ""), "f")

    def size(prs):
        gf = next(s for s in prs.slides[0].shapes if s.has_table)
        return gf.table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size.pt

    assert size(pinned) <= size(free)


def test_rowh_overflow_inside_a_pinned_row_warns(tmp_path):
    deck, prs = make(tmp_path, "\n# T\n{rowh=0.2in}\n| a | b |\n|-|-|\n| " + "long text " * 12 + " | 2 |\n")
    w = warns(deck, "overflow")
    assert len(w) == 1 and "rowh" in w[0].message and "raise rowh" in w[0].hint
    assert row_heights(prs)[0] == [round(0.2 * EMU_IN)] * 2  # still the author's heights


def test_rowh_taller_than_the_area_warns_instead_of_shrinking(tmp_path):
    deck, prs = make(tmp_path, "\n# T\n{rowh=3in}\n" + TABLE)
    assert any("pinned" in d.message for d in warns(deck, "overflow"))
    assert row_heights(prs)[0] == [3 * EMU_IN] * 4


def test_rowh_free_rows_are_not_pinned(tmp_path):
    _, prs = make(tmp_path, "\n# T\n{rowh=0.9in,auto}\n" + TABLE)
    rh = row_heights(prs)[0]
    assert rh[0] == round(0.9 * EMU_IN) and rh[1] != rh[0]


def test_rowh_bad_values_warn_never_raise(tmp_path):
    for bad in ("0.8", "wide", "50%", "0in", "99in"):
        deck, prs = make(tmp_path, f"\n# T\n{{rowh={bad}}}\n" + TABLE, "b")
        w = warns(deck, "bad-table-option")
        assert w and all(d.hint for d in w), bad
        assert next(s for s in prs.slides[0].shapes if s.has_table), bad
    deck, _ = make(tmp_path, "\n# T\n{rowh=0.5in,0.6in,0.7in,0.8in,0.9in}\n" + TABLE, "many")
    w = warns(deck, "bad-table-option")
    assert len(w) == 1 and "5 values for 4 rows" in w[0].message and "rowh=" in w[0].hint


def test_rowh_round_trip(tmp_path):
    make(tmp_path, "\n# T\n{rowh=0.8in,1.05in}\n" + TABLE, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "rowh=0.8in,1.05in" in text
    (tmp_path / "again.md").write_text(text, encoding="utf-8")
    build(str(tmp_path / "again.md"), tmp_path / "again.pptx")
    assert (
        row_heights(Presentation(str(tmp_path / "again.pptx")))[0]
        == [round(0.8 * EMU_IN)] + [round(1.05 * EMU_IN)] * 3
    )


def test_unpinned_tables_import_without_rowh(tmp_path):
    make(tmp_path, "\n# T\n" + TABLE, "plain")
    text, _ = import_pptx(tmp_path / "plain.pptx")
    assert "rowh" not in text


# --------------------------------------------------------------------------- 3. a colour per bar


BARS = "\n# a\n```column {{{opts}}}\n,A,B,C,D\nv,10,20,30,15\n```\n"


def point_fills(chart):
    ser = chart.plots[0].series[0]
    out = {}
    for dpt in ser._element.findall(qn("c:dPt")):
        out[int(dpt.find(qn("c:idx")).get("val"))] = dpt.find(".//" + qn("a:srgbClr")).get("val")
    return out


def test_one_series_column_takes_a_colour_per_bar(tmp_path):
    deck, prs = make(tmp_path, BARS.format(opts="colors=primary,secondary,teal,muted"))
    assert not warns(deck, "bad-chart-option")
    assert point_fills(chart_of(prs)) == {0: "142B4D", 1: "1F5FA8", 2: "2A9D8F", 3: "59626E"}


def test_one_series_bar_chart_and_hex_colours(tmp_path):
    md = BARS.format(opts="colors=#112233,#445566,#778899,#AABBCC").replace("```column", "```bar")
    _, prs = make(tmp_path, md)
    assert point_fills(chart_of(prs)) == {0: "112233", 1: "445566", 2: "778899", 3: "AABBCC"}


def test_hl_still_wins_over_a_per_bar_colour(tmp_path):
    _, prs = make(tmp_path, BARS.format(opts="colors=primary,secondary,teal,muted hl=B"))
    fills = point_fills(chart_of(prs))
    assert (
        fills[1] == "E09F1F" and fills[0] == "142B4D" and fills[2] == "2A9D8F"
    )  # B took the emphasis colour


def test_colour_count_other_than_the_bar_count_warns_and_keeps_one_colour(tmp_path):
    deck, prs = make(tmp_path, BARS.format(opts="colors=primary,teal"))
    w = warns(deck, "bad-chart-option")
    assert len(w) == 1 and "4 bars" in w[0].message and "exactly 4" in w[0].hint
    assert point_fills(chart_of(prs)) == {}


def test_one_colour_stays_the_series_colour_and_several_series_keep_series_colours(tmp_path):
    deck, prs = make(tmp_path, BARS.format(opts="colors=teal"))
    assert not warns(deck, "bad-chart-option") and point_fills(chart_of(prs)) == {}
    md = "\n# a\n```column {colors=primary,teal}\n,A,B\nx,1,2\ny,3,4\n```\n"
    deck, prs = make(tmp_path, md, "multi")
    ser = chart_of(prs).plots[0].series
    assert not warns(deck, "bad-chart-option") and not ser[0]._element.findall(qn("c:dPt"))


def test_inside_labels_on_per_bar_colours_read_on_each_fill(tmp_path):
    _, prs = make(tmp_path, BARS.format(opts="labels=inside colors=primary,#FFE9A8,primary,#FFE9A8"))
    ser = chart_of(prs).plots[0].series[0]
    inks = {}
    for dl in ser._element.find(qn("c:dLbls")).findall(qn("c:dLbl")):
        inks[int(dl.find(qn("c:idx")).get("val"))] = dl.find(".//" + qn("a:srgbClr")).get("val")
    assert inks[0] == "FFFFFF" and inks[1] != "FFFFFF" and inks[2] == "FFFFFF" and inks[3] != "FFFFFF"


def test_pie_colours_are_unchanged(tmp_path):
    md = "\n# a\n```pie {colors=primary,secondary,teal}\n,a,b,c\nv,1,2,3\n```\n"
    _, prs = make(tmp_path, md)
    assert point_fills(chart_of(prs)) == {0: "142B4D", 1: "1F5FA8", 2: "2A9D8F"}


def test_per_bar_colours_import_without_inventing_options(tmp_path):
    make(tmp_path, BARS.format(opts="colors=primary,secondary,teal,muted hl=B labels=on"), "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "```column" in text and "hl=B" in text


# --------------------------------------------------------------------------- 4. per-point label position


def label_positions(chart, si=0):
    ser = chart.plots[0].series[si]
    pos = {}
    for dl in ser._element.find(qn("c:dLbls")).findall(qn("c:dLbl")):
        pos[int(dl.find(qn("c:idx")).get("val"))] = dl.find(qn("c:dLblPos")).get("val")
    return pos


LINE = "\n# a\n```line {{{opts}}}\n,A,B,C,D\nx,10,20,30,15\ny,12,18,28,16\n```\n"


def test_labels_per_category_on_a_line_chart(tmp_path):
    deck, prs = make(tmp_path, LINE.format(opts="labels=above,below,above,right"))
    assert not warns(deck, "bad-chart-option")
    for si in (0, 1):
        assert label_positions(chart_of(prs), si) == {0: "t", 1: "b", 2: "t", 3: "r"}


def test_labels_per_category_overrides_the_automatic_collision_rule(tmp_path):
    close = "\n# a\n```line {{{opts}}}\n,A,B\nx,60,70\ny,58,69\n```\n"
    _, prs_a = make(tmp_path, close.format(opts="labels=on"), "auto")
    # the two close series alternate above / below on their own: series 1 is "below"
    assert set(label_positions(chart_of(prs_a), 1).values()) == {"b"}
    _, prs = make(tmp_path, close.format(opts="labels=above,above"))
    assert set(label_positions(chart_of(prs), 1).values()) == {"t"}


def test_labels_auto_keeps_the_collision_rule_for_that_category(tmp_path):
    _, prs = make(tmp_path, LINE.format(opts="labels=auto,auto,left,auto"))
    pos = label_positions(chart_of(prs), 0)
    assert pos[2] == "l"


def test_labels_per_category_on_a_column_chart(tmp_path):
    md = BARS.format(opts="labels=inside,outside,center,outside")
    deck, prs = make(tmp_path, md)
    assert not warns(deck, "bad-chart-option")
    assert label_positions(chart_of(prs)) == {0: "inEnd", 1: "outEnd", 2: "ctr", 3: "outEnd"}


def test_labels_wrong_count_warns_with_the_category_count(tmp_path):
    deck, prs = make(tmp_path, LINE.format(opts="labels=above,below"))
    w = warns(deck, "bad-chart-option")
    assert len(w) == 1 and "2 positions for 4 categories" in w[0].message and "exactly 4" in w[0].hint
    pos = label_positions(chart_of(prs), 0)
    assert pos[0] == "t" and pos[1] == "b" and 2 not in pos  # the rest stays automatic


def test_labels_bad_word_and_wrong_chart_kind_warn(tmp_path):
    deck, _ = make(tmp_path, BARS.format(opts="labels=above,below,above,above"))
    assert any("does not fit a column chart" in d.message and d.hint for d in warns(deck, "bad-chart-option"))
    pie = "\n# a\n```pie {labels=outside,inside,outside}\n,a,b,c\nv,1,2,3\n```\n"
    deck, prs = make(tmp_path, pie, "pie")
    w = warns(deck, "bad-chart-option")
    assert w and all(d.hint for d in w) and chart_of(prs)


def test_single_word_labels_are_unchanged(tmp_path):
    _, prs = make(tmp_path, LINE.format(opts="labels=below"))
    assert 'c:dLblPos val="b"' in chart_of(prs).plots[0]._element.xml


# --------------------------------------------------------------------------- 5. h= y= w= on KPI cards


ROW = (
    "\n# T\n> lead\n"
    "## A {{.kpi .hero{a}}}\n1,280\nvs +8%\n"
    "## B {{.kpi{b}}}\n96\nvs +12%\n"
    "## C {{.kpi{c}}}\n7.5%\n+0.3pt\n"
)


def cards(prs):
    return sorted(shapes(prs, 0, "Card"), key=lambda s: s.left)


def body_of(tmp_path):
    """(top, height) of the body in EMU: a row with h=100% fills it."""
    _, prs = make(tmp_path, ROW.format(a=" h=100%", b="", c=""), "body")
    c = cards(prs)[0]
    return c.top, c.height


def test_kpi_row_h_and_y_are_honoured(tmp_path):
    top, height = body_of(tmp_path)
    deck, prs = make(tmp_path, ROW.format(a=" h=55% y=24%", b="", c=""))
    cs = cards(prs)
    assert len(cs) == 3 and cs[0].left < cs[1].left < cs[2].left  # cards keep their columns
    assert {c.top for c in cs} == {cs[0].top} and {c.height for c in cs} == {cs[0].height}
    assert abs(cs[0].height - round(0.55 * height)) <= 2
    assert abs(cs[0].top - (top + round(0.24 * height))) <= 2
    assert not warns(deck, "overlap") and not warns(deck, "overflow")


def test_kpi_row_height_is_the_largest_h_given(tmp_path):
    top, height = body_of(tmp_path)
    _, prs = make(tmp_path, ROW.format(a=" h=30%", b=" h=50%", c=""))
    assert {c.height for c in cards(prs)} == {round(0.5 * height)}


def test_kpi_row_y_alone_moves_the_row_and_h_alone_keeps_the_top(tmp_path):
    top, height = body_of(tmp_path)
    _, moved = make(tmp_path, ROW.format(a=" y=10%", b="", c=""), "y")
    _, plain = make(tmp_path, ROW.format(a="", b="", c=""), "plain")
    assert cards(moved)[0].top == top + round(0.1 * height)
    assert 0 < cards(moved)[0].height <= cards(plain)[0].height  # y alone: the row keeps its natural height
    _, tall = make(tmp_path, ROW.format(a="", b=" h=70%", c=""), "h")
    assert cards(tall)[0].top == top and cards(tall)[0].height == round(0.7 * height)


def test_kpi_row_w_on_the_hero_sets_its_width(tmp_path):
    _, prs = make(tmp_path, ROW.format(a=" w=50%", b="", c=""))
    cs = cards(prs)
    body_w = cs[-1].left + cs[-1].width - cs[0].left
    assert abs(cs[0].width - 0.5 * body_w) <= 0.01 * body_w
    assert abs(cs[1].width - cs[2].width) <= 2 and cs[1].width < cs[0].width  # the rest share what is left
    assert cs[0].left < cs[1].left < cs[2].left


def test_kpi_pinned_row_text_stays_inside_the_cards(tmp_path):
    _, prs = make(tmp_path, ROW.format(a=" h=55% y=24%", b="", c=""))
    cs = cards(prs)
    for sh in shapes(prs, 0, "Text") + shapes(prs, 0, "Heading"):
        assert any(
            c.left <= sh.left
            and sh.left + sh.width <= c.left + c.width + 2
            and c.top <= sh.top
            and sh.top + sh.height <= c.top + c.height + 2
            for c in cs
        )


def test_kpi_row_bad_values_warn_never_raise(tmp_path):
    deck, prs = make(tmp_path, ROW.format(a=" h=abc y=-5 w=wide", b="", c=""))
    assert len(cards(prs)) == 3
    assert all(d.hint for d in warns(deck, "bad-length"))


def test_kpi_card_with_x_is_still_absolute(tmp_path):
    _, prs = make(tmp_path, ROW.format(a=" x=1in y=2in w=3in h=2in", b="", c=""))
    first = next(c for c in cards(prs) if c.width == 3 * EMU_IN)  # absolute: x / y are inside the body
    assert first.height == 2 * EMU_IN and first.left > EMU_IN and first.top > 2 * EMU_IN


# --------------------------------------------------------------------------- fuzz: bad values never raise

_TMP = Path(tempfile.mkdtemp(prefix="sm_dl2c_"))  # hypothesis needs a module-level directory, not tmp_path
_VALUE = st.text(
    alphabet=st.characters(blacklist_categories=("Cs", "Cc"), blacklist_characters="{}`\n\r"), max_size=20
)
_WORDS = st.lists(
    st.sampled_from(["above", "below", "left", "right", "center", "outside", "inside", "auto", "x"])
)


@settings(max_examples=30, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(rowh=_VALUE, colors=_VALUE, h=_VALUE, y=_VALUE, w=_VALUE, rule=_VALUE, words=_WORDS)
def test_new_options_never_raise(rowh, colors, h, y, w, rule, words):
    md = (
        HEAD
        + f"style: kpi.rule={rule} kpi.rule_h={rule} kpi.rule_w={rule}\n"
        + f"\n# T\n## A {{.kpi h={h} y={y} w={w}}}\n1\nn\n## B {{.kpi}}\n2\nm\n"
        + f"\n# T2\n{{rowh={rowh}}}\n"
        + TABLE
        + f"\n# T3\n```column {{labels={','.join(words)} colors={colors}}}\n,A,B\nv,1,2\n```\n"
        + f"\n# T4\n```line {{labels={','.join(words)}}}\n,A,B\nv,1,2\n```\n"
    )
    deck = build(md, _TMP / "fuzz.pptx")
    assert deck.slides and not [d for d in deck.diagnostics if d.rule == "layout-error"]
