"""Chart text grows with a big frame (lone chart) and stays at the theme size for a small chart."""

from __future__ import annotations

from pptx import Presentation

from slidemark import build
from slidemark.render.axis import chart_text_pt
from slidemark.theme import RenderTokens

EMU = 914400
BIG = """theme: jp-business
lang: ja

# 問い合わせ内容の内訳
```bar {labels=on}
,料金,解約,故障,住所変更,その他
件数,4200,2600,1900,1300,800
```
"""
SMALL = """theme: jp-business
lang: ja

# 内訳
@3
## 左
- 要点
```bar {labels=on}
,料金,解約,故障
件数,4200,2600,1900
```
## 右
- 要点
## 他
- 要点
"""


def _chart_pt(src: str, tmp_path) -> float:
    out = tmp_path / "d.pptx"
    build(src, out)
    prs = Presentation(str(out))
    sizes = []
    for shape in prs.slides[0].shapes:
        if shape.has_chart:
            sizes.append(shape.chart.font.size.pt)
    assert sizes
    return min(sizes)


def test_lone_big_chart_text_grows(tmp_path):
    assert _chart_pt(BIG, tmp_path) >= 13


def test_small_chart_text_is_unchanged(tmp_path):
    assert _chart_pt(SMALL, tmp_path) <= 11


def test_chart_text_pt_rules():
    rt = RenderTokens()
    assert chart_text_pt(10.5, 3 * EMU, 3 * EMU, rt) == 10.5  # never below the base
    assert chart_text_pt(10.5, 12 * EMU, 5.7 * EMU, rt) > 13
    assert chart_text_pt(10.5, 40 * EMU, 40 * EMU, rt) == rt.chart_text_max_pt  # capped
    assert chart_text_pt(10.5, 12 * EMU, 6 * EMU, RenderTokens(chart_text_ratio=0)) == 10.5  # off
