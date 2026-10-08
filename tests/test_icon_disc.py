# ruff: noqa: E501
"""Icon discs (`icon.disc`), card elevation / borderless shadowed cards, the conclusion bar icon.

Every check reopens the built .pptx: the disc is a native `Icon disc` shape under the glyph, the glyph is
centred in it, the shadow is a real `a:outerShdw`, a shadowed card has no outline unless the deck asked for one.
"""

from __future__ import annotations

import re

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import build, parse
from slidemark.theme import ELEVATIONS, Theme, apply_tokens

PT = 12700
HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#1F2933 primary=#0B2A3C secondary=#1B8A8F accent=#F2A33A surface=#EEF4F6 "
    "border=#D5DEE3\n"
)
CARDS = "\n# Three ways\n@3\n## One {icon=target}\nBody one.\n## Two {icon=users}\nBody two.\n## Three {icon=gear disc=accent}\nBody three.\n"  # noqa: E501


def make(tmp_path, text: str):
    out = tmp_path / "d.pptx"
    deck = build_text(tmp_path, text, out)
    return deck, Presentation(str(out))


def build_text(tmp_path, text: str, out):
    src = tmp_path / "d.md"
    src.write_text(text, encoding="utf-8")
    return build(src, out)


def shapes(prs, slide=0):
    return list(prs.slides[slide].shapes)


def named(prs, name, slide=0):
    """Shapes called ``name`` (a numbered family too: ``Card`` = ``Card 1``, ``Card 2`` ...)."""
    return [
        s for s in shapes(prs, slide) if s.name == name or re.fullmatch(re.escape(name) + r" \d+", s.name)
    ]


def fill_of(shp) -> str | None:
    el = shp._element.spPr.find(qn("a:solidFill"))
    return None if el is None else el.find(qn("a:srgbClr")).get("val")


def prst_of(shp) -> str | None:
    g = shp._element.spPr.find(qn("a:prstGeom"))
    return None if g is None else g.get("prst")


def outline_off(shp) -> bool:
    ln = shp._element.spPr.find(qn("a:ln"))
    return ln is None or ln.find(qn("a:noFill")) is not None


def shadow_of(shp):
    eff = shp._element.spPr.find(qn("a:effectLst"))
    return None if eff is None else eff.find(qn("a:outerShdw"))


def center(s):
    return s.left + s.width / 2, s.top + s.height / 2


# --------------------------------------------------------------------------- tokens


def test_tokens_parse_and_apply():
    th, diags = apply_tokens(
        Theme(name="none"),
        {
            "icon_disc": "secondary",
            "icon_disc_size": "0.7in",
            "icon_disc_shape": "rounded",
            "icon_color": "#FFFFFF",
            "card_elevation": "2",
            "conclusion_icon": "refresh",
        },
    )
    assert not diags
    assert (th.icon_disc, th.icon_disc_size, th.icon_disc_shape) == ("secondary", "0.7in", "rounded")
    assert (th.icon_color, th.card_elevation, th.conclusion_icon) == ("#FFFFFF", 2, "refresh")
    assert th.classes["card"].shadow == "0 3 12 #00000030"


def test_defaults_are_off():
    th = Theme(name="none")
    assert th.icon_disc is None and th.icon_color is None and th.conclusion_icon is None
    assert th.icon_disc_shape == "circle" and th.card_elevation is None
    assert th.render.shadow == "0 3 12 #00000030"
    assert th.layout.icon_disc_ratio == 2.2


def test_header_names_reach_the_theme():
    deck = parse(
        HEAD + "style: icon.disc=secondary icon.disc.size=0.7in icon.disc.shape=square icon.color=bg\n\n# T\n"
    )
    assert not [d for d in deck.diagnostics if d.rule in ("unknown-token", "bad-token")]
    assert {"icon_disc", "icon_disc_size", "icon_disc_shape", "icon_color"} <= set(deck.tokens)


@pytest.mark.parametrize(
    "level,want", [(0, False), (1, "0 1 4 #0000001F"), (2, "0 3 12 #00000030"), (3, "0 8 24 #00000040")]
)
def test_elevation_values(level, want):
    assert ELEVATIONS[level] == want
    th, diags = apply_tokens(Theme(name="none"), {"card_elevation": str(level)})
    assert not diags and th.classes["card"].shadow == want


def test_bad_values_are_diagnostics_with_hints():
    th, diags = apply_tokens(
        Theme(name="none"),
        {
            "icon_disc_shape": "sqare",
            "card_elevation": "9",
            "conclusion_icon": "refrsh",
            "icon_disc_size": "wide",
        },
    )
    assert [d.rule for d in diags] == ["bad-token"] * 4
    assert "did you mean 'square'" in diags[0].hint
    assert "0" in diags[1].hint and "3" in diags[1].hint
    assert th.icon_disc_shape == "circle" and th.icon_disc_size is None and th.conclusion_icon is None


def test_unknown_disc_shape_in_a_deck_never_raises(tmp_path):
    deck, prs = make(tmp_path, HEAD + "style: icon.disc=secondary icon.disc.shape=blob\n" + CARDS)
    assert any(
        d.rule == "bad-token" and "did you mean" in d.hint or d.rule == "bad-token" for d in deck.diagnostics
    )
    assert prst_of(named(prs, "Icon disc")[0]) == "ellipse"  # the default shape is kept


# --------------------------------------------------------------------------- discs on box headings


def test_heading_icon_sits_on_a_disc(tmp_path):
    _deck, prs = make(tmp_path, HEAD + "style: radius=8 icon.disc=secondary\n" + CARDS)
    discs = named(prs, "Icon disc")
    glyphs = [s for s in shapes(prs) if s.name.startswith("icon ")]
    assert len(discs) == len(glyphs) == 3
    assert [g.name for g in glyphs] == ["icon target", "icon users", "icon gear"]
    for d, g in zip(discs, glyphs, strict=True):
        assert prst_of(d) == "ellipse" and d.width == d.height
        assert abs(center(d)[0] - center(g)[0]) <= 2 and abs(center(d)[1] - center(g)[1]) <= 2  # centred
        assert g.width == pytest.approx(0.46 * d.width, abs=3)  # `layout.icon_disc_glyph`
        assert outline_off(d) and shadow_of(d) is None
        assert shapes(prs).index(d) + 1 == shapes(prs).index(g)  # the disc is right under its glyph
    assert [fill_of(d) for d in discs] == ["1B8A8F", "1B8A8F", "F2A33A"]  # `disc=accent` on the third
    # the glyph ink: white on teal, whichever reads on the amber (dark text colour, not white)
    assert fill_of(glyphs[0]) == "FFFFFF" and fill_of(glyphs[2]) != "FFFFFF"
    # the disc is about 2.2 x the bare glyph side (1.2 x the heading font)
    head = next(s for s in shapes(prs) if s.has_text_frame and s.text_frame.text == "One")
    size_pt = head.text_frame.paragraphs[0].runs[0].font.size.pt
    assert discs[0].width / PT == pytest.approx(2.2 * 1.2 * size_pt, rel=0.2)


def test_text_shift_follows_the_disc_not_the_glyph(tmp_path):
    _deck, prs = make(tmp_path, HEAD + "style: icon.disc=secondary\n" + CARDS)
    d = named(prs, "Icon disc")[0]
    head = next(s for s in shapes(prs) if s.has_text_frame and s.text_frame.text == "One")
    assert head.left >= d.left + d.width  # the heading text starts after the whole disc
    assert head.left - (d.left + d.width) < 0.6 * d.width


def test_disc_none_and_override_per_heading(tmp_path):
    md = (
        HEAD
        + "style: icon.disc=secondary\n\n# T\n@2\n## A {icon=bolt disc=none}\nx\n## B {icon=star disc=#FF0000}\ny\n"
    )
    _deck, prs = make(tmp_path, md)
    discs = named(prs, "Icon disc")
    assert len(discs) == 1 and fill_of(discs[0]) == "FF0000"
    assert [s.name for s in shapes(prs) if s.name.startswith("icon ")] == ["icon bolt", "icon star"]


def test_no_disc_token_leaves_bare_glyphs(tmp_path):
    _deck, prs = make(tmp_path, HEAD + CARDS)
    discs = named(prs, "Icon disc")
    assert [fill_of(d) for d in discs] == ["F2A33A"]  # only the heading that says `disc=accent` has one
    glyphs = [s for s in shapes(prs) if s.name.startswith("icon ")]
    assert [g.name for g in glyphs] == ["icon target", "icon users", "icon gear"]
    assert fill_of(glyphs[0]) == fill_of(glyphs[1]) == "1F2933"  # a bare glyph keeps the heading colour


def test_disc_shapes_and_size(tmp_path):
    for shape, prst in (("rounded", "roundRect"), ("square", "rect"), ("circle", "ellipse")):
        _d, prs = make(
            tmp_path,
            HEAD + f"style: icon.disc=primary icon.disc.shape={shape} icon.disc.size=0.6in\n" + CARDS,
        )
        for d in named(prs, "Icon disc"):
            assert prst_of(d) == prst
            assert d.width == d.height == pytest.approx(0.6 * 914400, abs=2)


def test_icon_color_token(tmp_path):
    _d, prs = make(tmp_path, HEAD + "style: icon.disc=primary icon.color=#FFFF00\n" + CARDS)
    glyphs = [s for s in shapes(prs) if s.name.startswith("icon ")]
    assert {fill_of(g) for g in glyphs} == {"FFFF00"}
    _d, prs = make(
        tmp_path, HEAD + "style: icon.color=#FF00FF\n\n# T\n@2\n## A {icon=bolt}\nx\n## B {icon=star}\ny\n"
    )
    assert {fill_of(s) for s in shapes(prs) if s.name.startswith("icon ")} == {
        "FF00FF"
    }  # no disc: bare glyphs


def test_bad_disc_attribute_is_a_diagnostic(tmp_path):
    deck, prs = make(
        tmp_path,
        HEAD + "style: icon.disc=secondary\n\n# T\n@2\n## A {icon=bolt disc=zzzz}\nx\n## B {icon=star}\ny\n",
    )
    assert any(d.rule == "bad-attr" and "disc=zzzz" in d.message for d in deck.diagnostics)
    assert len(named(prs, "Icon disc")) == 2  # falls back to the deck's disc


def test_disc_on_a_heading_band_stays_inside_the_card(tmp_path):
    md = (
        HEAD
        + "style: heading.band=primary icon.disc=secondary\n\n# T\n@3\n## A {icon=bolt}\nx\n## B {icon=star}\ny\n## C {icon=heart}\nz\n"
    )
    _d, prs = make(tmp_path, md)
    cards = named(prs, "Card")
    for d in named(prs, "Icon disc"):
        card = next(c for c in cards if c.left <= d.left < c.left + c.width)
        assert d.top >= card.top and d.top + d.height <= card.top + card.height
        assert d.left >= card.left and d.left + d.width <= card.left + card.width


def test_kpi_card_icon_disc(tmp_path):
    md = (
        HEAD
        + "style: icon.disc=secondary\n\n# T\n@3\n## Revenue {.kpi icon=chart}\n**1.2B**\n## Users {.kpi icon=users}\n**3.4K**\n## Growth {.kpi icon=rocket}\n**+18%**\n"
    )
    _d, prs = make(tmp_path, md)
    discs = named(prs, "Icon disc")
    assert len(discs) == 3
    for d in discs:
        card = next(c for c in named(prs, "Card") if c.left <= d.left < c.left + c.width)
        assert abs(center(d)[0] - center(card)[0]) <= 3  # centred above the label
        assert d.top >= card.top and d.top + d.height <= card.top + card.height


def test_iconlist_items_get_discs(tmp_path):
    md = (
        HEAD
        + "style: icon.disc=secondary\n\n# T\n@iconlist cols=2 fill=surface\n- icon=document **Text** Write.\n- icon=camera disc=accent **Image** See.\n- icon=headphones **Audio** Hear.\n- icon=play disc=none **Video** Watch.\n"
    )
    deck, prs = make(tmp_path, md)
    discs = named(prs, "Icon disc")
    assert [fill_of(d) for d in discs] == ["1B8A8F", "F2A33A", "1B8A8F"]  # the 4th said disc=none
    glyphs = [s for s in shapes(prs) if s.name.startswith("icon ")]
    assert [g.name for g in glyphs] == ["icon document", "icon camera", "icon headphones", "icon play"]
    texts = [s for s in shapes(prs) if s.has_text_frame and s.text_frame.text.startswith(("Text", "Image"))]
    for d, t in zip(discs[:2], texts, strict=False):
        assert t.text_frame.margin_left >= d.width  # the text column starts after the disc
    assert not [
        x for x in deck.diagnostics if x.level == "warning" and x.rule in ("unknown-icon", "bad-attr")
    ]


def test_steps_arrow_and_chevron_icons_with_discs(tmp_path):
    md = (
        HEAD
        + "style: icon.disc=secondary\n\n# T\n@steps head=arrow\n## One {icon=database}\nA\n## Two {icon=cpu}\nB\n## Three {icon=rocket}\nC\n\n# C\n@chevron\n## P {icon=bolt}\n## Q {icon=star}\n## R {icon=heart}\n"
    )
    _d, prs = make(tmp_path, md)
    for n in (0, 1):
        discs = named(prs, "Icon disc", n)
        glyphs = [s for s in shapes(prs, n) if s.name.startswith("icon ")]
        assert len(discs) == len(glyphs) == 3
        arrows = [
            s
            for s in shapes(prs, n)
            if s.has_text_frame
            and s._element.spPr.find(qn("a:prstGeom")) is not None
            and prst_of(s) in ("chevron", "homePlate")
        ]
        assert len(arrows) == 3
        for d, a in zip(discs, arrows, strict=True):
            assert a.left < d.left and d.left + d.width < a.left + a.width  # inside its arrow
            assert a.top <= d.top and d.top + d.height <= a.top + a.height


# --------------------------------------------------------------------------- shadows


def test_shadow_on_uses_the_soft_default(tmp_path):
    _d, prs = make(tmp_path, HEAD + "style: card.shadow=on\n\n# T\n@2\n## A\nx\n## B\ny\n")
    cards = named(prs, "Card")
    assert len(cards) == 2
    for c in cards:
        sh = shadow_of(c)
        assert sh is not None
        assert (sh.get("blurRad"), sh.get("dist"), sh.get("dir")) == (str(12 * PT), str(3 * PT), "5400000")
        assert sh.find(qn("a:srgbClr")).find(qn("a:alpha")).get("val") == str(round(0x30 / 255 * 100000))


def test_soft_shadow_matches_pptxgenjs(tmp_path):
    _d, prs = make(tmp_path, HEAD + 'style: card.shadow="0 2 8 #0000001F"\n\n# T\n@2\n## A\nx\n## B\ny\n')
    sh = shadow_of(named(prs, "Card")[0])
    assert (sh.get("blurRad"), sh.get("dist"), sh.get("dir")) == ("101600", "25400", "5400000")
    assert (sh.get("algn"), sh.get("rotWithShape")) == ("ctr", "0")
    alpha = int(sh.find(qn("a:srgbClr")).find(qn("a:alpha")).get("val"))
    assert abs(alpha - 12000) <= 300  # pptxgenjs writes alpha 12%


def test_shadowed_card_has_no_border_unless_asked(tmp_path):
    base = HEAD + "\n# T\n@2\n## A\nx\n## B\ny\n"
    _d, prs = make(tmp_path, base.replace("\n# T", "style: card.elevation=2\n\n# T", 1))
    assert all(outline_off(c) and shadow_of(c) is not None for c in named(prs, "Card"))
    _d, prs = make(tmp_path, base.replace("\n# T", "style: card.shadow=on card.line=border\n\n# T", 1))
    assert all(not outline_off(c) and shadow_of(c) is not None for c in named(prs, "Card"))
    _d, prs = make(
        tmp_path,
        base.replace("\n# T", "style: card.shadow=on\n\n# T", 1).replace("## B", "## B {line=accent}"),
    )
    a, b = named(prs, "Card")
    assert outline_off(a) and not outline_off(b)  # an explicit `{line=}` keeps that card's border
    _d, prs = make(tmp_path, base)
    assert all(
        not outline_off(c) and shadow_of(c) is None for c in named(prs, "Card")
    )  # no shadow: as before


def test_per_box_shadow_drops_the_implied_border(tmp_path):
    _d, prs = make(tmp_path, HEAD + "\n# T\n@2\n## A {shadow=on}\nx\n## B\ny\n")
    a, b = named(prs, "Card")
    assert outline_off(a) and shadow_of(a) is not None
    assert not outline_off(b) and shadow_of(b) is None


def test_elevation_reaches_every_card_kind(tmp_path):
    md = (
        HEAD
        + "style: card.elevation=3 icon.disc=secondary\n\n# T\n@3\n## A {.kpi}\n**1**\n## B {.kpi}\n**2**\n## C {.kpi}\n**3**\n\n# S\n@steps\n## One\nA\n## Two\nB\n\n# I\n@iconlist fill=surface\n- icon=bolt **A** x\n- icon=star **B** y\n"
    )
    _d, prs = make(tmp_path, md)
    kpis = named(prs, "Card", 0)
    assert len(kpis) == 3 and all(shadow_of(c) is not None and outline_off(c) for c in kpis)
    steps = [s for s in shapes(prs, 1) if s.name.endswith("card")]
    assert len(steps) == 2 and all(shadow_of(c) is not None and outline_off(c) for c in steps)
    cells = [s for s in shapes(prs, 2) if s.name.startswith("Iconlist")]
    assert len(cells) == 2 and all(shadow_of(c) is not None for c in cells)
    for c in kpis:
        sh = shadow_of(c)
        assert (sh.get("blurRad"), sh.get("dist")) == (str(24 * PT), str(8 * PT))


def test_rows_bars_take_a_shadow_only_when_rows_shadow_is_set(tmp_path):
    base = HEAD + "style: card.elevation=2\n\n# T\n@rows\n1. First\n2. Second\n3. Third\n"
    _d, prs = make(tmp_path, base)
    bars = named(prs, "Row")
    assert bars and all(shadow_of(b) is None for b in bars)
    _d, prs = make(tmp_path, base.replace("elevation=2", 'elevation=2 rows.shadow="0 3 12 #00000030"'))
    bars = named(prs, "Row")
    assert bars and all(
        shadow_of(b) is not None and outline_off(b) for b in bars
    )  # no default border under it


def test_elevation_zero_is_flat(tmp_path):
    _d, prs = make(tmp_path, HEAD + "style: card.elevation=0\n\n# T\n@2\n## A\nx\n## B\ny\n")
    assert all(shadow_of(c) is None and not outline_off(c) for c in named(prs, "Card"))


# --------------------------------------------------------------------------- conclusion bar icon


BAR = "\n# T\n@3\n## A\nx\n## B\ny\n## C\nz\n> The bottom line\n"


def bar_of(prs, slide=0):
    return next(s for s in shapes(prs, slide) if s.name == "Conclusion")


def test_conclusion_icon_without_disc_is_white_on_the_bar(tmp_path):
    _d, prs = make(tmp_path, HEAD + "style: conclusion.icon=refresh\n" + BAR)
    bar = bar_of(prs)
    ic = next(s for s in shapes(prs) if s.name == "icon refresh")
    assert bar.left < ic.left and ic.left + ic.width < bar.left + bar.width
    assert bar.top <= ic.top and ic.top + ic.height <= bar.top + bar.height
    assert abs(center(ic)[1] - center(bar)[1]) <= 3  # vertically centred
    assert fill_of(ic) == "FFFFFF" and named(prs, "Icon disc") == []
    assert bar.text_frame.margin_left >= ic.left + ic.width - bar.left  # the text starts right of it
    assert bar.height >= 0.55 * 914400 - 2  # `layout.conclusion_icon_h`


def test_conclusion_icon_on_a_disc(tmp_path):
    _d, prs = make(tmp_path, HEAD + "style: icon.disc=secondary conclusion.icon=refresh\n" + BAR)
    bar, ic, disc = (
        bar_of(prs),
        next(s for s in shapes(prs) if s.name == "icon refresh"),
        named(prs, "Icon disc"),
    )
    assert len(disc) == 1 and fill_of(disc[0]) == "1B8A8F"
    assert disc[0].top >= bar.top and disc[0].top + disc[0].height <= bar.top + bar.height
    assert abs(center(disc[0])[0] - center(ic)[0]) <= 2 and abs(center(disc[0])[1] - center(ic)[1]) <= 2


def test_conclusion_without_the_token_is_unchanged(tmp_path):
    _d, prs = make(tmp_path, HEAD + BAR)
    assert not [s for s in shapes(prs) if s.name.startswith("icon ")]


def test_conclusion_icon_unknown_name_never_raises(tmp_path):
    deck, prs = make(tmp_path, HEAD + "style: conclusion.icon=refrsh\n" + BAR)
    assert any(d.rule == "bad-token" and "did you mean" in d.hint for d in deck.diagnostics)
    assert bar_of(prs) is not None


# --------------------------------------------------------------------------- fuzz


_WORDS = st.sampled_from(
    [
        "none",
        "off",
        "on",
        "secondary",
        "#GGG",
        "#12345",
        "",
        "0",
        "-5",
        "1e9",
        "wide",
        "circle",
        "blob",
        "{}",
        "100%",
        "9in",
        "refresh",
        "nope.svg",
    ]
)


@settings(max_examples=25, deadline=None, suppress_health_check=[HealthCheck.function_scoped_fixture])
@given(disc=_WORDS, size=_WORDS, shape=_WORDS, color=_WORDS, elev=_WORDS, icon=_WORDS, attr=_WORDS)
def test_fuzz_never_raises(tmp_path_factory, disc, size, shape, color, elev, icon, attr):
    tmp = tmp_path_factory.mktemp("f")
    head = f'style: icon.disc="{disc}" icon.disc.size="{size}" icon.disc.shape="{shape}" icon.color="{color}" card.elevation="{elev}" conclusion.icon="{icon}"\n'
    body = f"\n# T\n@2\n## A {{icon=bolt disc={attr or 'x'}}}\nx\n## B {{icon=star}}\ny\n> bar\n"
    out = tmp / "f.pptx"
    deck = build_text(tmp, HEAD + head + body, out)
    assert out.exists() and Presentation(str(out)).slides
    assert all(d.hint for d in deck.diagnostics if d.rule in ("bad-token", "bad-attr"))


# --------------------------------------------------------------------------- import


ROUND = (
    HEAD
    + "style: radius=8 card.elevation=2 icon.disc=secondary conclusion.icon=refresh\n\n"
    + "# Ways\n@3\n## One {icon=target}\nBody one.\n## Two {icon=users disc=none}\nBody two.\n"
    + "## Three {icon=gear disc=accent}\nBody three.\n> The line\n\n"
    + "# List\n@iconlist cols=2 fill=surface\n- icon=document **Text** Write.\n"
    + "- icon=camera disc=accent **Image** See.\n- icon=headphones **Audio** Hear.\n- icon=play **Video** Watch.\n"
)


def test_import_folds_disc_and_glyph_into_one_icon(tmp_path):
    from slidemark.importer import import_pptx

    src = tmp_path / "a.pptx"
    build_text(tmp_path, ROUND, src)
    text, _diags = import_pptx(src, tmp_path)
    assert "{icon=target}" in text and "{icon=users disc=none}" in text and "{icon=gear disc=accent}" in text
    assert "- icon=camera disc=accent **Image**" in text and "- icon=document **Text**" in text
    assert "icon=disc" not in text  # the disc is not an icon of its own
    assert "icon_disc=secondary" in text.splitlines()[2] and "conclusion_icon=refresh" in text.splitlines()[2]


def test_roundtrip_keeps_discs_glyphs_and_fills(tmp_path):
    from slidemark.importer import import_pptx

    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    build_text(tmp_path, ROUND, a)
    text, _ = import_pptx(a, tmp_path)
    (tmp_path / "t.md").write_text(text, encoding="utf-8")
    build(tmp_path / "t.md", b)

    def facts(path):
        prs = Presentation(str(path))
        return [
            [(s.name, fill_of(s)) for s in sl.shapes if s.name == "Icon disc" or s.name.startswith("icon ")]
            for sl in prs.slides
        ]

    assert facts(a) == facts(b)


def test_second_import_equals_the_first(tmp_path):
    from slidemark.importer import import_pptx

    a, b = tmp_path / "a.pptx", tmp_path / "b.pptx"
    build_text(tmp_path, ROUND, a)
    text1, _ = import_pptx(a, tmp_path)
    (tmp_path / "t.md").write_text(text1, encoding="utf-8")
    build(tmp_path / "t.md", b)
    text2, _ = import_pptx(b, tmp_path)
    assert text1 == text2


def test_fold_pairs_a_glyph_with_the_disc_under_it_and_names_the_deck_disc():
    from slidemark.importer import icondisc
    from slidemark.importer.read import Item, SlideData

    def disc(x, color, prst="ellipse"):
        return Item(kind="shape", x=x, y=100, w=600, h=600, name="Icon disc", fill=color, prst=prst, uid=x)

    def glyph(x, name):
        return Item(kind="shape", x=x + 150, y=250, w=300, h=300, name=name, fill="FFFFFF", uid=x + 1)

    items = [
        disc(0, "1B8A8F"),
        glyph(0, "icon bolt"),
        disc(1000, "1B8A8F", "roundRect"),
        glyph(1000, "icon star"),
        disc(2000, "F2A33A"),
        glyph(2000, "icon gear"),
        glyph(3000, "icon heart"),
    ]
    data = SlideData(items=items)
    icondisc.fold(data)
    g = [i for i in data.items if i.name.startswith("icon ")]
    assert [x.disc for x in g] == ["1B8A8F", "1B8A8F", "F2A33A", None]
    assert [x.disc_shape for x in g] == ["circle", "rounded", "circle", None]
    assert {i.role for i in data.items if i.name == "Icon disc"} == {"decor"}
    assert icondisc.deck_disc([data]) == ("1B8A8F", "circle")  # 2 of 4 icons: the commonest, at least half
    assert icondisc.attr("F2A33A", "1B8A8F", {"accent": "#F2A33A"}) == "disc=accent"
    assert icondisc.attr(None, "1B8A8F", {}) == "disc=none" and icondisc.attr("1b8a8f", "1B8A8F", {}) == ""
