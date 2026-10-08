"""DL3b composition vocabulary, part 1: ``@timeline`` ``@vs`` ``@matrix`` ``@funnel`` ``@pyramid``
``@cycle`` ``@agenda`` ``@statement``. Every form: parser (word and secondary attributes reach the IR, bad
values warn with a hint), layout (geometry of the placed items), render (the .pptx is reopened: names, shapes,
tokens), importer round trip, fit line, ``attr-ignored``, and fuzz-safety (never raises, falls back to
ordinary blocks)."""

from __future__ import annotations

import re
import tempfile
from pathlib import Path

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from slidemark import build, forms, parse
from slidemark.importer import import_pptx
from slidemark.theme import normalize_token

ROOT = Path(__file__).resolve().parent.parent
_TMP = Path(tempfile.mkdtemp(prefix="vocab-"))
HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#1F2A44 primary=#1F3A8A secondary=#0F766E accent=#D97706 teal=#2A9D8F "
    "muted=#5B6475 surface=#F3F5FA border=#D5DBE7\n"
    "fonts: heading=Arial body=Arial\n"
)

TIMELINE = (
    "\n# T\n@timeline dir=h marks=on\n"
    "## 2026 Q1\nKickoff\n- pilot\n## 2026 Q2 {.accent}\nBeta\n## 2026 Q3\nGA\n## 2026 Q4\nScale\n"
)
TIMELINE_V = TIMELINE.replace("dir=h marks=on", "dir=v marks=num")
VS = "\n# T\n@vs\n## Build\n- Control\n- Slow\n## Buy {.hero}\n- Fast\n- Lock-in\n## Verdict\nBuy now\n"
MATRIX = (
    '\n# T\n@matrix x="low effort to high" y="low impact to high"\n'
    "## Plan\nbig rebuilds\n## Do first\nquick wins\n## Drop\nvanity\n## Fill in\nsmall fixes\n"
)
FUNNEL = "\n# T\n@funnel\n## Leads 100\nweb\n## Deals 40\ncalls\n## Won 10\nsigned\n"
PYRAMID = "\n# T\n@pyramid\n## Board\nsets course\n## Heads\nplan\n## Leads\nrun\n## Staff\ndo\n"
CYCLE = "\n# T\n@cycle\n## Plan\ngoals\n## Do\nship\n## Check\nmeasure\n## Act\nadjust\n"
AGENDA = "\n# T\n@agenda\n1. Context\n2. Options {.accent}\n   - three of them\n3. Decision\n"
STATEMENT = "\n# T\n@statement\n**+18%**\nrevenue growth year on year\n"
STAIRS = (
    "\n# T\n@stairs\n## Buy\nready now\n## Prompt\nwrite it\n## Retrieve\nyour data\n## Train\nown model\n"
)
NESTED = "\n# T\n@nested\n## AI\nthe field\n## ML\nlearns from data\n## DL\ndeep networks\n"
ALL = {
    "timeline": TIMELINE,
    "vs": VS,
    "matrix": MATRIX,
    "funnel": FUNNEL,
    "pyramid": PYRAMID,
    "cycle": CYCLE,
    "agenda": AGENDA,
    "statement": STATEMENT,
    "stairs": STAIRS,
    "nested": NESTED,
}


def make(tmp_path, text: str, head: str = HEAD, name: str = "d"):
    deck = build(head + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def named(prs, i: int, pattern: str):
    return [s for s in prs.slides[i].shapes if re.fullmatch(pattern, s.name)]


def one(prs, i: int, name: str):
    got = named(prs, i, re.escape(name))
    assert len(got) == 1, (name, [s.name for s in prs.slides[i].shapes])
    return got[0]


def fill_hex(shp) -> str:
    return str(shp.fill.fore_color.rgb)


def run_color(shp, k: int = 0) -> str:
    return str(shp.text_frame.paragraphs[k].runs[0].font.color.rgb)


def run_size(shp, k: int = 0) -> float:
    return shp.text_frame.paragraphs[k].runs[0].font.size.pt


def line_hex(shp) -> str:
    return str(shp.line.color.rgb)


def warns(deck, rule: str):
    return [d for d in deck.diagnostics if d.rule == rule]


def tx(shp) -> str:
    """The text of a shape without the renderer's orphan binding (NBSP) and range joiners (U+2060)."""
    return shp.text_frame.text.replace("\xa0", " ").replace("\u2060", "")


def inch(emu: int) -> float:
    return emu / 914400


# --------------------------------------------------------------------------- the contract (forms.py)


def test_every_form_has_keys_an_example_and_a_composer():
    from slidemark.layout import vocab

    assert set(forms.FORMS) | {"flowdisc"} == set(forms.KEYS) == set(forms.EXAMPLE) == set(vocab._FORMS)
    for form, example in forms.EXAMPLE.items():
        word = forms.WORD.get(form, form)
        assert example.startswith(f"@{word}") or f"@{word}" in example


@pytest.mark.parametrize("form", forms.FORMS)
def test_the_word_reaches_the_slide_and_nothing_warns(form):
    deck = parse(HEAD + ALL[form])
    assert forms.form_of(deck.slides[0]) == form
    assert not [d for d in deck.diagnostics if d.rule not in ("design-none", "design-slide")], (
        deck.diagnostics
    )


def test_secondary_attributes_reach_slide_attrs():
    deck = parse(HEAD + TIMELINE_V.replace("marks=num", "marks=off size=14 fill=teal,accent gap=0.1in"))
    a = deck.slides[0].attrs
    assert (a["dir"], a["marks"], a["size"], a["fill"]) == ("v", "off", "14", "teal,accent")
    assert a["gap"] == "0.1in"
    deck = parse(HEAD + MATRIX)
    assert deck.slides[0].attrs["x"] == "low effort to high" and deck.slides[0].attrs["y"].startswith(
        "low impact"
    )
    deck = parse(HEAD + STATEMENT.replace("@statement", "@statement align=left valign=top"))
    assert (deck.slides[0].attrs["align"], deck.slides[0].attrs["valign"]) == ("left", "top")


def test_a_form_slide_is_never_a_cover():
    deck = parse(HEAD + STATEMENT)
    assert deck.slides[0].layout is None and deck.slides[0].elements and deck.slides[0].subtitle is None
    deck = parse(HEAD + "\n# T\n@timeline\n## a\n## b\n")  # headings only: still a form slide, not a cover
    assert deck.slides[0].layout is None and len(deck.slides[0].elements) == 2


@pytest.mark.parametrize(
    "text, rule, hint",
    [
        ("\n# T\n@timeline dir=x\n## a\n## b\n", "bad-attr", "dir=h|v"),
        ("\n# T\n@timeline marks=big\n## a\n## b\n", "bad-attr", "marks=on|off|num"),
        ("\n# T\n@funnel dir=left\n## a\n## b\n", "bad-attr", "dir=down|up"),
        ("\n# T\n@cycle dir=around\n## a\n## b\n## c\n", "bad-attr", "dir=cw|ccw"),
        ("\n# T\n@statement valign=low\nbig\n", "bad-attr", "valign=top|middle|bottom"),
        ("\n# T\n@timeline size=huge\n## a\n## b\n", "bad-attr", "size=14"),
        ("\n# T\n@timeline fill=\n## a\n## b\n", "bad-attr", "fill=primary"),
        ("\n# T\n@vs marks=on\n## a\n## b\n", "attr-ignored", "@vs takes"),
        ("\n# T\n@3 dir=v\n## a\n## b\n## c\n", "attr-ignored", "belongs to"),
        ("\n# T\n@timeline vs\n## a\n## b\n", "form-conflict", "keep one"),
    ],
)
def test_bad_attributes_warn_with_a_hint(text, rule, hint):
    got = warns(parse(HEAD + text), rule)
    assert got, text
    assert hint in (got[0].hint or ""), got[0]


@pytest.mark.parametrize(
    "form, text",
    [
        ("timeline", "\n# T\n@timeline\n## only one\nx\n"),
        ("vs", "\n# T\n@vs\n## a\n## b\n## c\n## d\n"),
        ("matrix", "\n# T\n@matrix\n## a\n## b\n## c\n"),
        ("funnel", "\n# T\n@funnel\n## a\n"),
        ("pyramid", "\n# T\n@pyramid\n## a\ntext\n\nnot a box\n"),
        ("cycle", "\n# T\n@cycle\n## a\n## b\n"),
        ("agenda", "\n# T\n@agenda\n- a\n- b\n"),
        ("statement", "\n# T\n@statement\n- a\n- b\n"),
        ("stairs", "\n# T\n@stairs\n## a\n"),
        ("nested", "\n# T\n@nested\n## a\n## b\n"),
    ],
)
def test_a_slide_without_what_the_form_needs_warns_and_stays_ordinary(form, text, tmp_path):
    deck = parse(HEAD + text)
    got = warns(deck, f"{form}-skipped")
    assert got and "laid out as ordinary blocks" in got[0].hint
    built, prs = make(tmp_path, text)  # never raises, never composes
    assert not [
        s
        for s in prs.slides[0].shapes
        if re.match(r"(Timeline|VS|Matrix|Funnel|Pyramid|Cycle|Agenda|Statement|Stairs|Nested)", s.name)
    ]


def test_check_attrs_is_pure_and_total():
    assert forms.check_attrs("timeline", {"dir": "v", "marks": "num", "size": "14", "fill": "a"}) == []
    assert forms.check_attrs(None, {"gap": "0.1in"}) == []  # gap is a grid key on every slide
    assert forms.words({}, "timeline", "dir") == "h" and forms.words({"dir": "bad"}, "timeline", "dir") == "h"
    assert forms.words({"dir": "UP"}, "pyramid", "dir") == "up"
    assert forms.num({"size": "14pt"}, "size") == 14 and forms.num({"size": "x"}, "size") is None


# --------------------------------------------------------------------------- tokens


@pytest.mark.parametrize(
    "key, value, path",
    [
        ("timeline.line", "teal", "timeline_line"),
        ("timeline.line.w", "5pt", "timeline_line_w"),
        ("timeline.dot", "accent", "timeline_dot"),
        ("timeline.dot.size", "0.3in", "timeline_dot_size"),
        ("timeline.now", "#112233", "timeline_now"),
        ("timeline.date.color", "none", "timeline_date_color"),
        ("vs.badge.text", "対", "vs_badge_text"),
        ("vs.badge.fill", "accent", "vs_badge_fill"),
        ("vs.badge.size", "0.7in", "vs_badge_size"),
        ("vs.verdict.fill", "teal", "vs_verdict_fill"),
        ("vs.win.line", "accent", "vs_win_line"),
        ("vs.win.w", "5", "vs_win_w"),
        ("matrix.axis.color", "teal", "matrix_axis_color"),
        ("matrix.axis.size", "14", "matrix_axis_size"),
        ("matrix.fill", "primary,teal,accent,surface", "matrix_fill"),
        ("funnel.fill", "primary,teal", "funnel_fill"),
        ("funnel.gap", "0.1in", "funnel_gap"),
        ("funnel.taper", "0.5", "funnel_taper"),
        ("pyramid.fill", "primary", "pyramid_fill"),
        ("pyramid.share", "0.6", "pyramid_share"),
        ("cycle.fill", "primary,teal", "cycle_fill"),
        ("cycle.arrow", "accent", "cycle_arrow"),
        ("cycle.node.size", "1.2in", "cycle_node_size"),
        ("agenda.num.size", "40", "agenda_num_size"),
        ("agenda.num.text", "第{n}章", "agenda_num_text"),
        ("agenda.rule", "none", "agenda_rule"),
        ("agenda.dim", "none", "agenda_dim"),
        ("statement.size", "90pt", "statement_size"),
        ("statement.sub.color", "teal", "statement_sub_color"),
    ],
)
def test_every_token_form_is_accepted(key, value, path):
    from slidemark.theme import canonical_token

    got, hint = canonical_token("style", key)
    assert got == path, (key, hint)
    pairs = normalize_token(got, value, {"primary", "teal", "accent", "surface"})
    assert pairs and pairs[0][0] == path


@pytest.mark.parametrize(
    "line",
    [
        "style: funnel.taper=2",
        "style: pyramid.share=0.05",
        "style: vs.win.w=-1",
        "style: matrix.fill=a,b,c,d,e",
        "style: timeline.dot.size=big",
        "style: cycle.arrow=notacolor",
        "style: statement.size=huge",
    ],
)
def test_a_bad_token_is_a_warning_with_a_hint_and_never_raises(line, tmp_path):
    deck, _prs = make(tmp_path, "\n# T\n@timeline\n## a\n## b\n", head=HEAD + line + "\n")
    got = [d for d in deck.diagnostics if d.rule in ("bad-token", "unknown-token")]
    assert got and got[0].hint, line


def test_tokens_are_listed_by_slidemark_tokens():
    from slidemark.theme import schema_table

    paths = {p for p, _v in schema_table()}
    for p in ("timeline_dot_size", "vs_badge_text", "matrix_axis_color", "funnel_taper", "cycle_arrow"):
        assert p in paths
    assert {"layout.vocab_grow", "layout.timeline_stem", "layout.cycle_node_ratio"} <= paths


# --------------------------------------------------------------------------- @timeline


def test_timeline_h_is_a_line_with_dots_and_alternating_text(tmp_path):
    _deck, prs = make(tmp_path, TIMELINE)
    axis = one(prs, 0, "Timeline line")
    dots = [one(prs, 0, f"Timeline {n} dot") for n in range(1, 5)]
    texts = [named(prs, 0, rf"Timeline {n}( now)?")[0] for n in range(1, 5)]
    assert axis.width > axis.height * 50  # a thin horizontal rule
    assert [d.left for d in dots] == sorted(d.left for d in dots)
    mid = axis.top + axis.height // 2
    assert all(abs(d.top + d.height // 2 - mid) < 2000 for d in dots)
    above = [t.top + t.height <= mid for t in texts]
    assert above == [True, False, True, False]  # alternate above / below
    assert texts[1].name == "Timeline 2 now"  # {.accent} marks the current milestone
    assert len(named(prs, 0, r"Timeline \d stem")) == 4
    assert texts[0].text_frame.paragraphs[0].runs[0].font.bold  # the heading (date) leads


def test_timeline_v_puts_every_milestone_right_of_a_vertical_axis(tmp_path):
    _deck, prs = make(tmp_path, TIMELINE_V)
    axis = one(prs, 0, "Timeline line")
    assert axis.height > axis.width * 50
    texts = [named(prs, 0, rf"Timeline {n}( now)?")[0] for n in range(1, 5)]
    assert len({t.left for t in texts}) == 1 and all(t.left > axis.left for t in texts)
    assert [t.top for t in texts] == sorted(t.top for t in texts)
    dots = [one(prs, 0, f"Timeline {n} dot") for n in range(1, 5)]
    assert [d.text_frame.text for d in dots] == ["1", "2", "3", "4"]  # marks=num
    assert all(abs((d.left + d.width // 2) - (axis.left + axis.width // 2)) < 2000 for d in dots)


def test_timeline_marks_off_has_no_dots_and_on_has_plain_ones(tmp_path):
    _d, prs = make(tmp_path, TIMELINE.replace("marks=on", "marks=off"))
    assert not named(prs, 0, r"Timeline \d dot") and len(named(prs, 0, r"Timeline \d stem")) == 4
    _d, prs = make(tmp_path, TIMELINE)
    assert all(not s.text_frame.text for s in named(prs, 0, r"Timeline \d dot"))


def test_timeline_tokens_reach_the_shapes(tmp_path):
    head = (
        HEAD
        + "style: timeline.line=#112233 timeline.line.w=6pt timeline.dot=#AA0000 timeline.dot.size=0.5in "
        + "timeline.now=#00AA00 timeline.date.color=#0000AA\n"
    )
    _d, prs = make(tmp_path, TIMELINE, head=head)
    axis = one(prs, 0, "Timeline line")
    assert fill_hex(axis) == "112233" and axis.height == Emu(6 * 12700)
    d1, d2 = one(prs, 0, "Timeline 1 dot"), one(prs, 0, "Timeline 2 dot")
    assert fill_hex(d1) == "AA0000" and abs(inch(d1.width) - 0.5) < 0.01
    assert fill_hex(d2) == "00AA00"  # the {.accent} milestone
    assert run_color(one(prs, 0, "Timeline 1")) == "0000AA"
    assert run_color(one(prs, 0, "Timeline 2 now")) == "00AA00"


def test_timeline_fill_attribute_and_box_fill_color(tmp_path):
    text = TIMELINE.replace("dir=h marks=on", "dir=h marks=on fill=#AA0000,#00AA00").replace(
        "## 2026 Q4", "## 2026 Q4 {fill=#123456 color=#654321}"
    )
    _d, prs = make(tmp_path, text)
    assert [fill_hex(one(prs, 0, f"Timeline {n} dot")) for n in (1, 3)] == ["AA0000", "AA0000"]
    assert fill_hex(one(prs, 0, "Timeline 2 dot")) == "D97706"  # the current one keeps `timeline.now`
    assert fill_hex(one(prs, 0, "Timeline 4 dot")) == "123456"  # a box's own fill wins
    assert run_color(one(prs, 0, "Timeline 4")) == "654321"


def test_timeline_size_pins_the_text_and_ratios_are_checked(tmp_path):
    deck, prs = make(tmp_path, TIMELINE.replace("marks=on", "marks=on size=14"))
    t = named(prs, 0, r"Timeline 1")[0]
    assert run_size(t, 1) == 14  # body text pinned; the heading is larger
    assert run_size(t, 0) > 14
    deck, _prs = make(tmp_path, TIMELINE.replace("@timeline", "@1:2 timeline"))
    got = warns(deck, "form-grid")
    assert got and "4 ratios" in got[0].message
    deck, prs = make(tmp_path, TIMELINE.replace("@timeline", "@1:1:1:3 timeline"))
    assert not warns(deck, "form-grid")
    axis_x = [one(prs, 0, f"Timeline {n} dot").left for n in range(1, 5)]
    assert axis_x[3] - axis_x[2] > 1.5 * (axis_x[2] - axis_x[1])  # the wide last slot spaces its dot


def test_timeline_stems_are_lines_that_touch_axis_and_text(tmp_path):
    _d, prs = make(tmp_path, TIMELINE)
    axis = one(prs, 0, "Timeline line")
    for n in range(1, 5):
        stem, dot = one(prs, 0, f"Timeline {n} stem"), one(prs, 0, f"Timeline {n} dot")
        assert stem.shape_type == 9  # a native line
        assert min(abs(stem.top + stem.height - dot.top), abs(stem.top - dot.top - dot.height)) < 2000
    _d, prs = make(tmp_path, TIMELINE.replace("marks=on", "marks=off"), name="off")
    axis = one(prs, 0, "Timeline line")
    for n in range(1, 5):  # without dots the stem runs to the axis itself
        stem = one(prs, 0, f"Timeline {n} stem")
        assert min(abs(stem.top + stem.height - axis.top), abs(stem.top - axis.top - axis.height)) < 25000


# --------------------------------------------------------------------------- @vs


def test_vs_is_two_cards_a_badge_between_and_a_verdict_bar_below(tmp_path):
    _d, prs = make(tmp_path, VS)
    c1, c2 = one(prs, 0, "Card 1"), one(prs, 0, "Card 2")
    badge, bar = one(prs, 0, "VS badge"), one(prs, 0, "VS verdict")
    assert c1.left + c1.width <= badge.left and badge.left + badge.width <= c2.left
    assert c1.top == c2.top and abs(c1.height - c2.height) < 2000
    mid = c1.top + c1.height // 2
    assert abs(badge.top + badge.height // 2 - mid) < 5000
    assert bar.top >= c1.top + c1.height and bar.width > c1.width
    assert badge.text_frame.text == "vs"
    assert bar.text_frame.paragraphs[0].runs[0].font.bold  # the heading leads the bar
    assert "Buy now" in tx(bar) and tx(bar).startswith("Verdict")


def test_vs_without_a_verdict_and_with_ratios(tmp_path):
    _d, prs = make(tmp_path, VS.replace("## Verdict\nBuy now\n", "").replace("@vs", "@2:1 vs"))
    assert not named(prs, 0, "VS verdict")
    c1, c2 = one(prs, 0, "Card 1"), one(prs, 0, "Card 2")
    assert c1.width > 1.8 * c2.width  # 2:1


def test_vs_hero_marks_the_winner_with_win_line_tokens(tmp_path):
    head = HEAD + "style: vs.win.line=#CC0000 vs.win.w=5 vs.win.fill=#FFF4E5\n"
    _d, prs = make(tmp_path, VS, head=head)
    win, other = one(prs, 0, "Card 2"), one(prs, 0, "Card 1")
    assert line_hex(win) == "CC0000" and win.line.width == Emu(5 * 12700)
    assert fill_hex(win) == "FFF4E5" and line_hex(other) != "CC0000"


def test_vs_badge_and_verdict_tokens(tmp_path):
    head = (
        HEAD
        + 'style: vs.badge.text="対" vs.badge.fill=#AA0000 vs.badge.color=#FFFFFF vs.badge.size=0.8in '
        + "vs.verdict.fill=#00AA00 vs.verdict.color=#000000\n"
    )
    _d, prs = make(tmp_path, VS, head=head)
    badge, bar = one(prs, 0, "VS badge"), one(prs, 0, "VS verdict")
    assert badge.text_frame.text == "対" and fill_hex(badge) == "AA0000"
    assert abs(inch(badge.width) - 0.8) < 0.01 and run_color(badge) == "FFFFFF"
    assert fill_hex(bar) == "00AA00" and run_color(bar) == "000000"


def test_vs_fill_attribute_colours_the_cards(tmp_path):
    _d, prs = make(tmp_path, VS.replace("@vs", "@vs fill=#E8EEFB,#E7F5F2"))
    assert fill_hex(one(prs, 0, "Card 1")) == "E8EEFB" and fill_hex(one(prs, 0, "Card 2")) == "E7F5F2"


# --------------------------------------------------------------------------- @matrix


def test_matrix_is_a_2x2_with_axis_arrows_and_labels(tmp_path):
    _d, prs = make(tmp_path, MATRIX)
    cards = [one(prs, 0, f"Card {n}") for n in range(1, 5)]
    assert cards[0].left == cards[2].left and cards[1].left == cards[3].left
    assert cards[0].top == cards[1].top and cards[2].top == cards[3].top
    assert cards[1].left > cards[0].left and cards[2].top > cards[0].top
    xa, ya = one(prs, 0, "Matrix x axis"), one(prs, 0, "Matrix y axis")
    xl, yl = one(prs, 0, "Matrix x label"), one(prs, 0, "Matrix y label")
    assert xa.top >= cards[2].top + cards[2].height and xa.width >= cards[0].width * 2  # along the bottom
    assert ya.left + ya.width <= cards[0].left and ya.height >= cards[0].height * 2  # along the left
    assert tx(xl) == "low effort to high" and tx(yl).startswith("low impact")
    assert yl.text_frame._txBody.find(qn("a:bodyPr")).get("vert") == "vert270"  # reads bottom to top
    tails = [e for e in etree.fromstring(etree.tostring(xa._element)).iter(qn("a:tailEnd"))]
    assert tails and tails[0].get("type") == "triangle"  # an arrowhead


def test_matrix_without_labels_has_no_axes_and_ratios_split_columns(tmp_path):
    _d, prs = make(
        tmp_path,
        MATRIX.replace('x="low effort to high" y="low impact to high"', "").replace("@matrix", "@1:2 matrix"),
    )
    assert not named(prs, 0, "Matrix .*")
    c1, c2 = one(prs, 0, "Card 1"), one(prs, 0, "Card 2")
    assert c2.width > 1.8 * c1.width


def test_matrix_tokens_and_quadrant_fills(tmp_path):
    head = (
        HEAD
        + "style: matrix.axis.color=#AA0000 matrix.axis.size=14 matrix.axis.w=4pt "
        + "matrix.fill=#E8EEFB,#E7F5F2,#FBEFE3,#F3F5FA\n"
    )
    _d, prs = make(tmp_path, MATRIX, head=head)
    assert [fill_hex(one(prs, 0, f"Card {n}")) for n in range(1, 5)] == [
        "E8EEFB",
        "E7F5F2",
        "FBEFE3",
        "F3F5FA",
    ]
    assert line_hex(one(prs, 0, "Matrix x axis")) == "AA0000"
    assert one(prs, 0, "Matrix x axis").line.width == Emu(4 * 12700)
    assert (
        run_size(one(prs, 0, "Matrix x label")) == 14 and run_color(one(prs, 0, "Matrix x label")) == "AA0000"
    )
    _d, prs = make(
        tmp_path, MATRIX.replace("@matrix", "@matrix fill=#111111,#222222")
    )  # the slide attribute wins
    assert fill_hex(one(prs, 0, "Card 1")) == "111111" and fill_hex(one(prs, 0, "Card 2")) == "222222"


# --------------------------------------------------------------------------- @funnel / @pyramid


def _stage_widths(prs, name: str, n: int):
    """Width of the drawn trapezoid at its top and bottom edge (read from the custom geometry)."""
    out = []
    for k in range(1, n + 1):
        shp = (
            one(prs, 0, f"{name} {k}")
            if not named(prs, 0, rf"{name} {k} (up|down)")
            else named(prs, 0, rf"{name} {k} (up|down)")[0]
        )
        pts = [(int(p.get("x")), int(p.get("y"))) for p in shp._element.iter(qn("a:pt"))]
        top = [x for x, y in pts if y == 0]
        bot = [x for x, y in pts if y > 0]
        out.append((max(top) - min(top), max(bot) - min(bot)))
    return out


def test_funnel_narrows_downwards_with_bodies_to_the_right(tmp_path):
    _d, prs = make(tmp_path, FUNNEL)
    widths = _stage_widths(prs, "Funnel", 3)
    tops = [w[0] for w in widths]
    assert (
        tops == sorted(tops, reverse=True) and widths[0][0] > widths[-1][1]
    )  # the narrow end is at the bottom
    assert all(widths[i][1] <= widths[i][0] for i in range(3))
    shapes = [one(prs, 0, f"Funnel {k}") for k in (1, 2, 3)]
    assert [s.top for s in shapes] == sorted(s.top for s in shapes)
    assert all(s.shape_type == 5 for s in shapes)  # freeform (custom geometry)
    body = one(prs, 0, "Funnel 1 text")
    assert body.left >= shapes[0].left + shapes[0].width and body.text_frame.text == "web"
    assert shapes[0].text_frame.text.startswith("Leads")  # the heading is inside the shape


def test_pyramid_widens_downwards_and_funnel_dir_up_does_too(tmp_path):
    _d, prs = make(tmp_path, PYRAMID)
    w = _stage_widths(prs, "Pyramid", 4)
    assert [x[1] for x in w] == sorted(x[1] for x in w)  # widest at the bottom
    assert w[0][0] < w[-1][1]
    _d, prs = make(tmp_path, FUNNEL.replace("@funnel", "@funnel dir=up"))
    w = _stage_widths(prs, "Funnel", 3)
    assert [x[0] for x in w] == sorted(x[0] for x in w)
    assert named(prs, 0, r"Funnel \d up")  # the direction rides in the name (the importer reads it back)
    _d, prs = make(tmp_path, PYRAMID.replace("@pyramid", "@pyramid dir=down"))
    w = _stage_widths(prs, "Pyramid", 4)
    assert [x[0] for x in w] == sorted((x[0] for x in w), reverse=True)


def test_stage_fill_gap_taper_and_share_tokens(tmp_path):
    head = (
        HEAD
        + "style: funnel.fill=#AA0000,#00AA00 funnel.color=#FFFF00 funnel.gap=0.3in "
        + "funnel.taper=0.6 funnel.share=0.4\n"
    )
    _d, prs = make(tmp_path, FUNNEL, head=head)
    s = [one(prs, 0, f"Funnel {k}") for k in (1, 2, 3)]
    assert [fill_hex(x) for x in s] == ["AA0000", "00AA00", "AA0000"]  # the list cycles
    assert run_color(s[0]) == "FFFF00"
    gap = s[1].top - (s[0].top + s[0].height)
    assert abs(inch(gap) - 0.3) < 0.01
    full = s[0].width
    widths = _stage_widths(prs, "Funnel", 3)
    assert abs(widths[-1][1] / full - 0.6) < 0.02  # the narrow end is `taper` of the wide end
    _d, prs = make(tmp_path, FUNNEL)
    assert s[0].width < one(prs, 0, "Funnel 1").width  # share 0.4 < the default 0.5


def test_stage_without_bodies_uses_the_width_and_box_fill_wins(tmp_path):
    text = PYRAMID.replace("\nsets course", "").replace("\nplan", "").replace("\nrun", "").replace("\ndo", "")
    text = text.replace("## Staff", "## Staff {fill=#123456}")
    _d, prs = make(tmp_path, text)
    assert not named(prs, 0, r"Pyramid \d text")
    s = one(prs, 0, "Pyramid 4")
    assert fill_hex(s) == "123456" and s.width > 0.7 * prs.slide_width * 0.8


def test_stage_ratio_splits_shapes_from_bodies(tmp_path):
    _d, prs = make(tmp_path, FUNNEL.replace("@funnel", "@3:1 funnel"))
    s, b = one(prs, 0, "Funnel 1"), one(prs, 0, "Funnel 1 text")
    assert s.width > 2.5 * b.width


# --------------------------------------------------------------------------- @cycle


def test_cycle_nodes_sit_on_a_circle_with_curved_arrows_between(tmp_path):
    _d, prs = make(tmp_path, CYCLE)
    nodes = [one(prs, 0, f"Cycle {n}") for n in range(1, 5)]
    arrows = [one(prs, 0, f"Cycle {n} arrow") for n in range(1, 5)]
    cx = sum(n.left + n.width / 2 for n in nodes) / 4
    cy = sum(n.top + n.height / 2 for n in nodes) / 4
    radii = [((n.left + n.width / 2 - cx) ** 2 + (n.top + n.height / 2 - cy) ** 2) ** 0.5 for n in nodes]
    assert max(radii) - min(radii) < 3000 and min(radii) > nodes[0].width  # one ring, nodes apart
    assert len({(a.left, a.top, a.width, a.height) for a in arrows}) == 1  # arcs share the ring's box
    a = arrows[0]
    assert abs(a.left + a.width / 2 - cx) < 3000 and abs(a.width / 2 - radii[0]) < 3000
    geom = a._element.spPr.find(qn("a:prstGeom"))
    assert geom.get("prst") == "arc"
    adj = {g.get("name"): g.get("fmla") for g in geom.iter(qn("a:gd"))}
    assert set(adj) == {"adj1", "adj2"}
    assert a._element.spPr.find(qn("a:ln")).find(qn("a:tailEnd")).get("type") == "triangle"
    assert nodes[0].text_frame.text == "Plan" and one(prs, 0, "Cycle 1 text").text_frame.text == "goals"
    # the body text sits outside the ring: right of a right-hand node, left of a left-hand one
    t1, t3 = one(prs, 0, "Cycle 1 text"), one(prs, 0, "Cycle 3 text")
    assert t1.left >= nodes[0].left + nodes[0].width and t3.left + t3.width <= nodes[2].left


def test_cycle_counts_three_to_six_and_ccw_flips_the_arrows(tmp_path):
    for n in (3, 5, 6):
        text = "\n# T\n@cycle\n" + "".join(f"## S{k}\nbody {k}\n" for k in range(n))
        deck, prs = make(tmp_path, text)
        assert len(named(prs, 0, r"Cycle \d")) == n and len(named(prs, 0, r"Cycle \d arrow")) == n
        assert not [d for d in deck.diagnostics if d.rule in ("overlap", "overflow", "off-slide")], n
    _d, prs = make(tmp_path, CYCLE.replace("@cycle", "@cycle dir=ccw"))
    a = named(prs, 0, r"Cycle 1 arrow ccw")[0]
    ln = a._element.spPr.find(qn("a:ln"))
    assert ln.find(qn("a:headEnd")) is not None and ln.find(qn("a:tailEnd")) is None
    nodes = [one(prs, 0, f"Cycle {n}") for n in range(1, 5)]
    _d, prs_cw = make(tmp_path, CYCLE, name="cw")
    cw = [one(prs_cw, 0, f"Cycle {n}") for n in range(1, 5)]
    assert nodes[1].top != cw[1].top or nodes[1].left != cw[1].left  # the loop runs the other way


def test_cycle_tokens(tmp_path):
    head = (
        HEAD
        + "style: cycle.fill=#AA0000,#00AA00 cycle.color=#FFFF00 cycle.arrow=#0000AA "
        + "cycle.arrow.w=5pt cycle.node.size=1.2in\n"
    )
    _d, prs = make(tmp_path, CYCLE, head=head)
    n = [one(prs, 0, f"Cycle {k}") for k in range(1, 5)]
    assert [fill_hex(x) for x in n] == ["AA0000", "00AA00", "AA0000", "00AA00"]
    assert run_color(n[0]) == "FFFF00" and abs(inch(n[0].width) - 1.2) < 0.01
    a = one(prs, 0, "Cycle 1 arrow")
    assert line_hex(a) == "0000AA" and a.line.width == Emu(5 * 12700)


# --------------------------------------------------------------------------- @agenda


def test_agenda_rows_have_a_big_number_column_rules_and_a_current_item(tmp_path):
    _d, prs = make(tmp_path, AGENDA)
    nums = [named(prs, 0, rf"Agenda {n}( now)? num")[0] for n in (1, 2, 3)]
    rows = [named(prs, 0, rf"Agenda {n}( now)?")[0] for n in (1, 2, 3)]
    assert [n.text_frame.text for n in nums] == ["1", "2", "3"]
    assert run_size(nums[0]) > run_size(rows[0]) * 1.3  # a large number column
    assert rows[1].name == "Agenda 2 now" and "three of them" in tx(rows[1])
    assert all(r.left > nums[0].left + nums[0].width for r in rows)
    assert len(named(prs, 0, r"Agenda \d rule")) == 3  # between the rows and above the first
    assert [r.top for r in rows] == sorted(r.top for r in rows)
    assert "{.accent}" not in " ".join(r.text_frame.text for r in rows)


def test_agenda_tokens(tmp_path):
    head = (
        HEAD
        + 'style: agenda.num.size=44 agenda.num.color=#AA0000 agenda.num.text="第{n}章" agenda.rule=#00AA00 '
        + "agenda.rule.w=3pt agenda.now=#0000FF agenda.dim=#999999\n"
    )
    _d, prs = make(tmp_path, AGENDA, head=head)
    n1, n2 = named(prs, 0, r"Agenda 1 num")[0], named(prs, 0, r"Agenda 2 now num")[0]
    assert n1.text_frame.text == "第1章" and run_size(n1) == 44
    assert run_color(n1) == "999999"  # dimmed: another item is current
    assert run_color(n2) == "0000FF"
    r = named(prs, 0, r"Agenda 1 rule")[0]
    assert line_hex(r) == "00AA00" and r.line.width == Emu(3 * 12700)
    _d, prs = make(tmp_path, AGENDA.replace(" {.accent}", ""), head=head)  # nothing current: nothing dimmed
    assert run_color(named(prs, 0, r"Agenda 1 num")[0]) == "AA0000"
    _d, prs = make(tmp_path, AGENDA, head=HEAD + "style: agenda.rule=none\n")
    assert not named(prs, 0, r"Agenda \d rule")


def test_agenda_fill_attribute_and_size_pin(tmp_path):
    _d, prs = make(tmp_path, AGENDA.replace("@agenda", "@agenda size=20 fill=#EEEEEE"))
    assert fill_hex(one(prs, 0, "Agenda 1 fill")) == "EEEEEE"
    assert run_size(named(prs, 0, "Agenda 1")[0], 0) >= 20


# --------------------------------------------------------------------------- @statement


def test_statement_is_one_huge_centred_line_with_a_caption(tmp_path):
    _d, prs = make(tmp_path, STATEMENT)
    big, cap = one(prs, 0, "Statement"), one(prs, 0, "Statement caption")
    assert big.text_frame.text == "+18%" and cap.text_frame.text.startswith("revenue")
    assert run_size(big) > 3 * run_size(cap) and run_size(big) >= 100
    assert big.text_frame.paragraphs[0].alignment is not None
    assert cap.top >= big.top + big.height  # the caption is under the number
    mid = (big.left + big.width / 2) / prs.slide_width
    assert abs(mid - 0.5) < 0.01


def test_statement_align_valign_and_tokens(tmp_path):
    head = (
        HEAD
        + "style: statement.size=60 statement.color=#AA0000 statement.sub.size=20 "
        + "statement.sub.color=#00AA00\n"
    )
    _d, prs = make(tmp_path, STATEMENT.replace("@statement", "@statement align=left valign=top"), head=head)
    big, cap = one(prs, 0, "Statement"), one(prs, 0, "Statement caption")
    assert run_size(big) == 60 and run_color(big) == "AA0000"
    assert run_size(cap) == 20 and run_color(cap) == "00AA00"
    assert str(big.text_frame.paragraphs[0].alignment).startswith("LEFT")
    _d, prs2 = make(tmp_path, STATEMENT.replace("@statement", "@statement valign=bottom"), name="b")
    _d, prs3 = make(tmp_path, STATEMENT.replace("@statement", "@statement valign=top"), name="t")
    assert one(prs3, 0, "Statement").top < one(prs2, 0, "Statement").top


def test_statement_with_one_line_has_no_caption_and_size_attr_pins(tmp_path):
    _d, prs = make(tmp_path, "\n# T\n@statement size=80\n12 → 18%\n")
    assert not named(prs, 0, "Statement caption")
    assert run_size(one(prs, 0, "Statement")) == 80


# --------------------------------------------------------------------------- lint, XML, fit


@pytest.mark.parametrize("form", forms.FORMS)
def test_a_form_slide_builds_clean_and_valid(form, tmp_path):
    deck, _prs = make(tmp_path, ALL[form])
    bad = [
        d
        for d in deck.diagnostics
        if d.level in ("warning", "error") and d.rule not in ("design-none", "design-slide")
    ]
    assert not bad, [str(d) for d in bad]
    from slidemark.xsd import validate_pptx

    assert validate_pptx(tmp_path / "d.pptx") == []


@pytest.mark.parametrize("form", forms.FORMS)
def test_the_fit_line_names_the_form(form, tmp_path):
    from slidemark.fit import fit_lines
    from slidemark.layout import layout_slide
    from slidemark.template import deck_theme

    deck = parse(HEAD + ALL[form])
    theme, _ = deck_theme(deck, str(tmp_path))
    placed = [layout_slide(deck.slides[0], deck, theme, 0)]
    line = fit_lines(deck, placed, theme)[0]
    assert line.startswith("slide 1: " + form), line
    assert "text" in line or "big line" in line or "heading" in line


# --------------------------------------------------------------------------- attr-ignored


def test_style_tokens_need_their_form(tmp_path):
    for token in (
        "timeline.dot=#FF0000",
        "vs.badge.fill=#FF0000",
        "funnel.fill=#FF0000",
        "cycle.arrow=#FF0000",
    ):
        deck, _ = make(tmp_path, "\n# T\n- a\n- b\n", head=HEAD + f"style: {token}\n")
        got = warns(deck, "attr-ignored")
        assert got and f"style: {token.split('=')[0]}=" in got[0].message, token
    form = {"timeline": TIMELINE, "vs": VS, "funnel": FUNNEL, "cycle": CYCLE}
    for token, name in (
        ("timeline.dot=#FF0000", "timeline"),
        ("vs.badge.fill=#FF0000", "vs"),
        ("funnel.fill=#FF0000", "funnel"),
        ("cycle.arrow=#FF0000", "cycle"),
    ):
        deck, _ = make(tmp_path, form[name], head=HEAD + f"style: {token}\n")
        assert not warns(deck, "attr-ignored"), token


def test_a_box_attribute_a_form_cannot_honour_is_reported(tmp_path):
    deck, _ = make(tmp_path, FUNNEL.replace("## Deals 40", "## Deals 40 {size=40 fill=#123456}"))
    got = warns(deck, "attr-ignored")
    assert [d.message for d in got] == ["size= on a form box is not honoured"]
    assert "size= on the @ line" in got[0].hint
    deck, _ = make(tmp_path, TIMELINE.replace("## 2026 Q1", "## 2026 Q1 {color=#123456}"))
    assert not warns(deck, "attr-ignored")


# --------------------------------------------------------------------------- importer round trip


def _roundtrip(tmp_path, text: str) -> str:
    out = tmp_path / "rt.pptx"
    build(HEAD + text, out)
    md, _diags = import_pptx(out)
    return md


@pytest.mark.parametrize(
    "form, text, tokens, texts",
    [
        (
            "timeline",
            TIMELINE,
            "@timeline",
            ["## 2026 Q1", "Kickoff", "- pilot", "## 2026 Q2 {.accent}", "## 2026 Q4"],
        ),
        ("timeline", TIMELINE_V, "@timeline dir=v marks=num", ["## 2026 Q3", "GA"]),
        (
            "timeline",
            TIMELINE.replace("marks=on", "marks=off"),
            "@timeline marks=off",
            ["## 2026 Q2 {.accent}"],
        ),
        ("vs", VS, "@vs", ["## Build", "- Control", "## Buy {.hero}", "- Lock-in", "## Verdict", "Buy now"]),
        (
            "matrix",
            MATRIX,
            '@matrix x="low effort to high" y="low impact to high"',
            ["## Plan", "## Fill in", "small fixes"],
        ),
        ("funnel", FUNNEL, "@funnel", ["## Leads 100", "web", "## Won 10", "signed"]),
        ("funnel", FUNNEL.replace("@funnel", "@funnel dir=up"), "@funnel dir=up", ["## Deals 40"]),
        ("pyramid", PYRAMID, "@pyramid", ["## Board", "sets course", "## Staff"]),
        ("pyramid", PYRAMID.replace("@pyramid", "@pyramid dir=down"), "@pyramid dir=down", ["## Heads"]),
        ("cycle", CYCLE, "@cycle", ["## Plan", "goals", "## Act", "adjust"]),
        ("cycle", CYCLE.replace("@cycle", "@cycle dir=ccw"), "@cycle dir=ccw", ["## Check"]),
        ("agenda", AGENDA, "@agenda", ["1. Context", "1. Options {.accent}", "Decision"]),
        ("statement", STATEMENT, "@statement", ["+18%", "revenue growth year on year"]),
        (
            "statement",
            STATEMENT.replace("@statement", "@statement align=left"),
            "@statement align=left",
            ["+18%"],
        ),
    ],
)
def test_importer_folds_the_named_shapes_back(form, text, tokens, texts, tmp_path):
    md = _roundtrip(tmp_path, text)
    lines = md.splitlines()
    assert tokens in lines, (tokens, md)
    for t in texts:
        assert any(ln.strip() == t or ln.strip().endswith(t) for ln in lines), (t, md)
    again = parse(md)
    assert forms.form_of(again.slides[0]) == form
    assert not [d for d in again.diagnostics if d.rule in (f"{form}-skipped", "bad-attr", "attr-ignored")]


def test_a_reimported_form_builds_the_same_shapes(tmp_path):
    for form in ("timeline", "funnel", "cycle", "agenda", "vs", "matrix"):
        a = tmp_path / f"{form}-a.pptx"
        build(HEAD + ALL[form], a)
        md, _ = import_pptx(a)
        b = tmp_path / f"{form}-b.pptx"
        build(md, b)
        na = [
            s.name
            for s in Presentation(str(a)).slides[0].shapes
            if not s.name.startswith(("Footer", "Slide"))
        ]
        nb = [
            s.name
            for s in Presentation(str(b)).slides[0].shapes
            if not s.name.startswith(("Footer", "Slide"))
        ]
        assert na == nb, form


# --------------------------------------------------------------------------- fuzz


_WORD = st.text(
    alphabet=st.characters(blacklist_categories=("Cs", "Cc"), blacklist_characters="{}`\n\r\"'"), max_size=12
)


@settings(max_examples=40, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(
    form=st.sampled_from(forms.FORMS),
    n=st.integers(0, 9),
    a=_WORD,
    b=_WORD,
    key=st.sampled_from(["dir", "marks", "x", "y", "size", "fill", "gap", "align", "valign"]),
    val=_WORD,
    tok=st.sampled_from(
        ["timeline.dot.size", "funnel.taper", "cycle.node.size", "matrix.fill", "statement.size"]
    ),
    tv=_WORD,
)
def test_forms_never_raise_on_odd_input(form, n, a, b, key, val, tok, tv):
    boxes = "".join(f"## {a or 'h'}{i}\n{b}\n" for i in range(n))
    body = (
        boxes
        if form not in ("agenda", "statement")
        else ("1. a\n2. b\n" if form == "agenda" else f"{a}\n{b}\n")
    )
    text = f"{HEAD}style: {tok}={tv or '1'}\n\n# T\n@{form} {key}={val or 'x'}\n{body}"
    deck = build(text, _TMP / "fuzz.pptx")
    assert deck.slides and not [d for d in deck.diagnostics if d.rule in ("layout-error", "render-error")]


def test_non_finite_lengths_are_text_not_a_crash(tmp_path):
    deck, _prs = make(tmp_path, "\n# T\n## A {.kpi w=INF h=nan y=-inf}\n1\nn\n## B {.kpi}\n2\nm\n")
    assert not [d for d in deck.diagnostics if d.rule == "layout-error"]


def test_matrix_box_fill_wins_over_the_token_list(tmp_path):
    head = HEAD + "style: matrix.fill=#E8EEFB,#E7F5F2,#FBEFE3,#F3F5FA\n"
    _d, prs = make(tmp_path, MATRIX.replace("## Plan", "## Plan {fill=#123456}"), head=head)
    assert fill_hex(one(prs, 0, "Card 1")) == "123456" and fill_hex(one(prs, 0, "Card 2")) == "E7F5F2"
