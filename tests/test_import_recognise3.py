"""DL3d part 2, lane E: four hand-drawn forms of a foreign deck (``importer/recognise3.py``).

Each test draws loose python-pptx shapes the way the t2 deck does (a pentagon row with text cards and
``STEP n`` captions, centred KPI cards with a rule, header-band cards with ``■`` bullets, a takeaway bar that
ends just above the footer), imports the file and asserts the deck.md; negatives show where it must not fire.
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.importer.recognise3 import TH3

TASKS = ("Pricing", "Factory", "Launch", "Review")
NAVY, BLUE, TEAL, LIGHT, EDGE, GREY = "142B4D", "1F5FA8", "2A9D8F", "EEF2F7", "C9D2DE", "59626E"


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _rect(slide, x, y, w, h, rgb, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(rgb)
    if line:
        s.line.color.rgb = RGBColor.from_string(line)
    else:
        s.line.fill.background()
    return s


def _text(slide, x, y, w, h, text, size=16, bold=False, color="222B36", center=False, shape=None):
    tb = shape or slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if center:
            p.alignment = PP_ALIGN.CENTER
        parts = line if isinstance(line, list) else [(line, size, color)]
        for t, sz, c in parts:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(sz)
            r.font.bold = bold
            r.font.color.rgb = RGBColor.from_string(c)
    return tb


def _page(prs, title: str, n: int):
    """A content slide: a title text, a footer text and a page number (the usual chrome)."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, 0.6, 0.3, 12, 0.7, title, size=28, bold=True, color=NAVY)
    _text(s, 0.6, 7.05, 6, 0.3, "ACME Corp", size=10, color=GREY)
    _text(s, 11.7, 7.05, 1.0, 0.3, str(n), size=10, color=GREY)
    return s


def _import(prs, tmp_path: Path) -> str:
    path = tmp_path / "foreign.pptx"
    prs.save(path)
    text, diags = import_pptx(path)
    assert not [d for d in diags if d.level == "error"], diags
    return text


def _body(text: str, n: int) -> str:
    blocks = [b for b in re.split(r"(?m)^(?=# )", text) if b.startswith("# ")]
    return blocks[n - 1] if len(blocks) >= n else text


def _shapes(tmp_path: Path, text: str, n: int) -> dict:
    out = tmp_path / "back.pptx"
    build(text, out, base_dir=tmp_path)
    return {sh.name: sh for sh in Presentation(str(out)).slides[n - 1].shapes}


# --------------------------------------------------------------------------- 1. steps whose cards hold text


def _steps_slide(prs, captions=True, card_text=True, labels=("STEP 1", "STEP 2", "STEP 3", "STEP 4")):
    s = _page(prs, "Schedule", 1)
    names = ("Q1", "Q2", "Q3", "Q4")
    for i in range(4):
        x = 0.6 + i * 3.1
        _text(
            s,
            0,
            0,
            0,
            0,
            names[i],
            size=26,
            bold=True,
            color="FFFFFF",
            shape=_pent(s, x, 2.0, 2.9, 1.3, [NAVY, BLUE][i % 2]),
        )
        _rect(s, x, 3.6, 2.6, 2.6, LIGHT, line=EDGE)  # a card a little narrower than its arrow
        if card_text:
            _text(s, x + 0.1, 3.6, 2.4, 2.6, TASKS[i], size=24, bold=True, center=True)
        if captions:
            _text(s, x, 6.3, 2.6, 0.4, labels[i], size=14, bold=True, color=TEAL, center=True)
    return s


def _pent(slide, x, y, w, h, rgb):
    sh = _rect(slide, x, y, w, h, rgb, shape=MSO_SHAPE.PENTAGON)
    return sh


def test_steps_with_text_cards_and_step_captions(tmp_path):
    prs = _deck()
    _steps_slide(prs)
    text = _import(prs, tmp_path)
    body = _body(text, 1)
    assert "@4 steps num" in body, body
    assert re.search(r"steps-arrow\.fill=(primary|#142B4D),(secondary|#1F5FA8),", body)
    assert "steps.caption_color=#2A9D8F" in body and "steps.caption_size=14" in body
    assert "steps.caption=" not in body, "STEP {n} is the default pattern"
    assert "## Q1" in body and "Pricing" in body
    assert "STEP 1" not in body, "the caption is a token, not content"
    shapes = _shapes(tmp_path, text, 1)
    assert "Step 1 arrow" in shapes and "Step 4 card" in shapes


def test_steps_caption_pattern_is_stated_when_it_differs(tmp_path):
    prs = _deck()
    _steps_slide(prs, labels=("Phase 1", "Phase 2", "Phase 3", "Phase 4"))
    body = _body(_import(prs, tmp_path), 1)
    assert "@4 steps num" in body and 'steps.caption="Phase {n}"' in body


def test_steps_without_captions_have_no_num(tmp_path):
    prs = _deck()
    _steps_slide(prs, captions=False)
    body = _body(_import(prs, tmp_path), 1)
    assert "steps" in body.split("\n@", 1)[1].split("\n", 1)[0] and " num" not in body


def test_steps_with_a_missing_card_stay_a_chevron_row(tmp_path):
    prs = _deck()
    s = _steps_slide(prs)
    # drop the third card: not every arrow has one
    card = [
        sh
        for sh in s.shapes
        if sh.shape_type == 1 and sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == LIGHT
    ][2]
    card._element.getparent().remove(card._element)
    body = _body(_import(prs, tmp_path), 1)
    assert "steps num" not in body


# --------------------------------------------------------------------------- 2. centred KPI cards


def _kpi_slide(prs, rule=True, stripe=True, left=False, n=4):
    s = _page(prs, "Targets", 1)
    vals = [
        ("Sales", "1,280", "+8%"),
        ("Profit", "96", "+12%"),
        ("Margin", "7.5%", "+0.3pt"),
        ("ROE", "9.0%", "+0.8pt"),
    ]
    for i in range(n):
        x = 0.6 + i * 3.1
        _rect(s, x, 1.8, 2.8, 4.4, LIGHT, line=EDGE)
        if stripe:
            _rect(s, x, 1.8, 2.8, 0.1, BLUE)
        label, value, delta = vals[i]
        _text(s, x + 0.1, 2.2, 2.6, 0.6, label, size=20, bold=True, color=NAVY, center=not (left and i == 1))
        _text(s, x + 0.1, 3.2, 2.6, 1.2, value, size=34, bold=True, color=BLUE, center=True)
        if rule:
            _rect(s, x + 0.55, 4.95, 1.7, 0.03, EDGE)
        _text(s, x + 0.1, 5.2, 2.6, 0.6, delta, size=20, bold=True, color=TEAL, center=True)
    return s


def test_centred_kpi_cards_with_rule_and_stripe(tmp_path):
    prs = _deck()
    _kpi_slide(prs)
    text = _import(prs, tmp_path)
    body = _body(text, 1)
    assert "@kpi" in body
    for tok in ("kpi.rule=#C9D2DE", "kpi.rule_w=1.7in"):
        assert tok in body, (tok, body)
    assert re.search(r"kpi\.stripe=(secondary|primary|#1F5FA8)", body) and re.search(
        r"kpi\.color=(secondary|primary|#1F5FA8)", body
    )
    assert "kpi.label.size=20" in body and "kpi.note.size=20" in body and "kpi.note.color=#2A9D8F" in body
    assert "kpi.h=4.4in" in body and "kpi.line=#C9D2DE" in body and "sizes: kpi=34" in body
    assert "## Sales" in body and "1,280" in body and "+8%" in body
    assert "kpi.align=left" not in body, "centred is the default"
    shapes = _shapes(tmp_path, text, 1)
    assert any("KPI" in k or "Card" in k for k in shapes)


def test_kpi_cards_without_rule_or_stripe_state_neither(tmp_path):
    prs = _deck()
    _kpi_slide(prs, rule=False, stripe=False)
    body = _body(_import(prs, tmp_path), 1)
    assert "@kpi" in body and "kpi.rule" not in body and "kpi.stripe" not in body


def test_kpi_group_with_one_left_aligned_card_is_not_read(tmp_path):
    prs = _deck()
    _kpi_slide(prs, left=True)
    assert "kpi.rule" not in _body(_import(prs, tmp_path), 1)


def test_one_kpi_card_alone_is_not_a_group(tmp_path):
    prs = _deck()
    _kpi_slide(prs, n=1)
    assert "kpi.rule" not in _body(_import(prs, tmp_path), 1)


# --------------------------------------------------------------------------- 3. header-band cards


def _band_slide(prs, glyph="■ ", lines=2, footnote=True):
    s = _page(prs, "Strategy", 1)
    for i in range(3):
        x = 0.6 + i * 4.1
        _rect(s, x, 1.6, 3.8, 4.4, LIGHT, line=EDGE)
        head = _rect(s, x, 1.6, 3.8, 0.9, NAVY)
        _text(s, 0, 0, 0, 0, f"Theme {i + 1}", size=20, bold=True, color="FFFFFF", center=True, shape=head)
        body = [[(glyph, 18, TEAL), (f"point {k + 1} of theme {i + 1}", 22, "222B36")] for k in range(lines)]
        _text(s, x + 0.1, 2.7, 3.6, 3.0, body)
    if footnote:
        _text(s, 0.6, 6.2, 12, 0.4, "* see appendix", size=13)
    return s


def test_header_band_cards_with_glyph_bullets(tmp_path):
    prs = _deck()
    _band_slide(prs)
    text = _import(prs, tmp_path)
    body = _body(text, 1)
    assert "heading.band=primary" in body and "heading.align=center" in body, body
    assert "bullet=■" in body and "bullet.color=#2A9D8F" in body and "card.line=#C9D2DE" in body
    assert "## Theme 1" in body and "- point 1 of theme 1" in body
    assert "■" not in body.split("bullet=■", 1)[1], "the glyph is a bullet token, not text"
    assert "{size=" not in body and "[■]" not in body
    assert "sizes: heading=20! body=22!" in body
    assert "* see appendix" in body or "see appendix" in body
    shapes = _shapes(tmp_path, text, 1)
    assert "Card 1" in shapes or any("Card" in k for k in shapes)


def test_cards_with_a_light_header_are_not_header_bands(tmp_path):
    prs = _deck()
    s = _band_slide(prs)
    for sh in s.shapes:  # every header rectangle becomes light: no dark band
        if sh.shape_type == 1 and sh.fill.type == 1 and str(sh.fill.fore_color.rgb) == NAVY:
            sh.fill.fore_color.rgb = RGBColor.from_string("D9E2EF")
    assert "heading.band" not in _body(_import(prs, tmp_path), 1)


def test_header_band_cards_whose_text_has_no_glyph_still_read(tmp_path):
    prs = _deck()
    _band_slide(prs, glyph="")
    body = _body(_import(prs, tmp_path), 1)
    assert "heading.band=primary" in body and "bullet=" not in body


# --------------------------------------------------------------------------- 4. takeaway bar


def _bar_slide(prs, fill=BLUE, y=6.0, h=0.8, extra=False):
    s = _page(prs, "Sales", 1)
    _text(s, 0.6, 1.4, 12, 1.0, "body line", size=24)
    bar = _rect(s, 0.6, y, 12.1, h, fill)
    _text(
        s,
        0,
        0,
        0,
        0,
        "Frozen foods and exports lead",
        size=24,
        bold=True,
        color="FFFFFF",
        center=True,
        shape=bar,
    )
    if extra:
        _text(s, 0.6, 6.85, 12, 0.3, "late note", size=14)
    return s


def test_takeaway_bar_that_ends_above_the_footer_line(tmp_path):
    prs = _deck()
    _bar_slide(prs, y=6.0, h=0.8)  # ends at 6.8 in = 90.7% of the height (recognise2 stops at 90%)
    text = _import(prs, tmp_path)
    body = _body(text, 1)
    assert re.search(r"conclusion\.fill=(secondary|#1F5FA8)", body) and "conclusion.size=24" in body, body
    assert re.search(r"^> .*Frozen foods", body, re.M)
    assert "Conclusion" in _shapes(tmp_path, text, 1)


def test_light_bar_is_not_a_takeaway_bar(tmp_path):
    prs = _deck()
    _bar_slide(prs, fill="DDE6F2")
    assert "conclusion.fill" not in _body(_import(prs, tmp_path), 1)


def test_bar_with_content_below_it_is_not_a_takeaway_bar(tmp_path):
    prs = _deck()
    _bar_slide(prs, y=5.2, h=0.8, extra=True)  # a text under the bar, above the footer line
    assert "conclusion.fill" not in _body(_import(prs, tmp_path), 1)


# --------------------------------------------------------------------------- guards


def test_a_deck_slidemark_built_is_never_read_by_geometry(tmp_path):
    md = (
        "# KPI\n@kpi\n## Sales\n**1,280**\n\n+8%\n## Profit\n**96**\n\n+12%\n"
        "\n# Steps\n@3 steps num\n## A\nx\n## B\ny\n## C\nz\n> done\n"
    )
    out = tmp_path / "own.pptx"
    build(md, out, base_dir=tmp_path)
    text, _ = import_pptx(out)
    assert "{.kpi}" in text and "@3 steps" in text
    assert "kpi.rule" not in text and "heading.band" not in text and "conclusion.fill" not in text


def test_thresholds_are_one_table():
    assert TH3.card_w[0] < 1 < TH3.card_w[1] and TH3.bar_bottom > TH3.footer_y
