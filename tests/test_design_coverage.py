"""Design wave 3, lane 2: every design decision of the python-pptx deck has a short SlideMark form
(docs/DESIGN_COVERAGE.md). Each form: parser (token reaches the IR / theme), render (the .pptx is reopened)
and fuzz-safety (a bad value is a diagnostic with a hint, never an exception)."""

from __future__ import annotations

import re

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from slidemark import build, parse
from slidemark.importer import import_pptx
from slidemark.template import deck_theme

HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#222B36 primary=#142B4D secondary=#1F5FA8 accent=#E09F1F teal=#2A9D8F "
    "muted=#59626E surface=#EEF2F7 border=#C9D2DE\n"
)


def make(tmp_path, text: str, name: str = "d"):
    deck = build(HEAD + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def shapes(prs, i, prefix=""):
    return [s for s in prs.slides[i].shapes if s.name.lower().startswith(prefix.lower())]


def rules(prs, i):
    return [s for s in prs.slides[i].shapes if s.name == "rule"]


def fill_hex(shp) -> str:
    return str(shp.fill.fore_color.rgb)


def warns(deck, rule):
    return [d for d in deck.diagnostics if d.rule == rule]


# --------------------------------------------------------------------------- token routing


def test_role_tokens_reach_their_element_not_an_unused_class():
    d = parse(
        HEAD + "style: lead.bold=on conclusion.bold=on footnote.size=9 footer.color=accent "
        'kpi.border-top="6pt solid secondary" kpi.label.size=20 kpi.note.color=teal\n\n# T\n> x\n- a\n'
    )
    by = {r.selector: r.style for r in d.css}
    assert by[".lead"].bold is True and by[".conclusion"].bold is True
    assert by[".footnote"].font_size == 9 and by[".caption"].color == "accent"
    assert by[".kpi"].border_top == "6pt solid secondary"
    assert by[".kpi h2"].font_size == 20 and by[".kpi .caption"].color == "teal"
    assert not [k for k in d.tokens if k.startswith("classes.lead") or k.startswith("classes.conclusion")]


def test_real_class_tokens_keep_their_meaning():
    d = parse(HEAD + "style: kpi.color=secondary lead.color=secondary card.fill=#EEEEEE\n\n# T\n- a\n")
    assert d.tokens["classes.kpi.color"] == "secondary" and d.tokens["lead_color"] == "secondary"
    assert d.tokens["classes.card.fill"] == "#EEEEEE"


def test_footer_color_and_kpi_parts_render(tmp_path):
    _, prs = make(
        tmp_path,
        "footer: ACME\nstyle: footer.color=accent kpi.label.bold=on kpi.note.color=teal\n\n"
        "# T\n## A {.kpi}\n1\nnote\n",
    )
    foot = next(s for s in prs.slides[0].shapes if s.name == "Footer")
    assert str(foot.text_frame.paragraphs[0].runs[0].font.color.rgb) == "E09F1F"


def test_declared_color_names_beat_css_names(tmp_path):
    d = parse(HEAD + "style: cover.bar=teal\n\n# T\n- a\n")
    th, _ = deck_theme(d, ".")
    assert th.cover_bar == "teal" and th.colors["teal"] == "#2A9D8F"
    assert (
        parse("colors: teal=#2A9D8F\n```css\n.x { color: teal }\n```\n# T\n- a\n").css[0].style.color
        == "teal"
    )


# --------------------------------------------------------------------------- slide chrome


def test_top_and_bottom_bar_and_title_rule(tmp_path):
    deck, prs = make(
        tmp_path,
        "style: top.bar=primary top.bar_h=0.12in bottom.bar=accent title.rule=secondary title.rule_h=3pt\n\n"
        "# Cover\n\n# Content\n- a\n",
    )
    assert not warns(deck, "bad-token")
    assert rules(prs, 0) == []  # the cover keeps its full-bleed look
    got = {fill_hex(s): s for s in rules(prs, 1)}
    top, bottom, rule = got["142B4D"], got["E09F1F"], got["1F5FA8"]
    assert top.top == 0 and top.width == prs.slide_width and abs(top.height - Emu(int(0.12 * 914400))) < 5
    assert bottom.top + bottom.height == prs.slide_height
    assert rule.height == 38100 and rule.left == Emu(int(0.5 * 914400))


def test_chrome_tokens_bad_values_warn_never_raise(tmp_path):
    deck, _ = make(tmp_path, "style: top.bar=nope-color top.bar_h=wide title.rule=#GGG\n\n# T\n- a\n")
    assert warns(deck, "bad-token")
    assert all(d.hint for d in warns(deck, "bad-token"))


def test_kpi_stripe(tmp_path):
    _, prs = make(
        tmp_path,
        "style: kpi.stripe=secondary kpi.stripe_h=8pt\n\n# T\n## A {.kpi}\n1\nn\n## B {.kpi}\n2\nn\n",
    )
    st = rules(prs, 0)
    assert len(st) == 2 and all(fill_hex(s) == "1F5FA8" and s.height == 8 * 12700 for s in st)
    cards = shapes(prs, 0, "card")
    assert {c.top for c in cards} == {s.top for s in st}


# --------------------------------------------------------------------------- bullets


def test_bullet_glyph_and_color(tmp_path):
    _, prs = make(tmp_path, "style: bullet=■ bullet.color=teal\n\n# T\n- one\n- two\n1. num\n")
    xml = re.sub(r">\s+<", "><", prs.slides[0].shapes._spTree.xml)
    assert xml.count('<a:buChar char="■"/>') == 2 and "<a:buAutoNum" in xml
    assert '<a:buClr><a:srgbClr val="2A9D8F"/></a:buClr>' in xml


def test_default_bullet_unchanged(tmp_path):
    _, prs = make(tmp_path, "\n# T\n- one\n")
    assert '<a:buChar char="•"/>' in prs.slides[0].shapes._spTree.xml


# --------------------------------------------------------------------------- charts


def _chart(prs, i=0):
    return next(s for s in prs.slides[i].shapes if s.has_chart).chart


def test_chart_gap_marker_size_label_position(tmp_path):
    md = (
        "\n# a\n```column {labels=inside gap=30 size=12 colors=teal,secondary}\n,x,y\na,1,2\nb,3,4\n```\n"
        "\n# b\n```line {labels=below marker=14}\n,x,y\na,1,2\n```\n"
        "\n# c\n```pie {labels=outside}\n,a,b\nv,1,2\n```\n"
    )
    deck, prs = make(tmp_path, md)
    assert not warns(deck, "bad-chart-option")
    col = _chart(prs, 0)
    assert col.plots[0].gap_width == 30
    xml = col._chartSpace.xml
    assert 'val="inEnd"' in xml and 'sz="1200"' in xml and "2A9D8F" in xml
    line = _chart(prs, 1)._chartSpace.xml
    assert 'val="b"' in line and '<c:size val="14"/>' in line
    assert 'val="outEnd"' in _chart(prs, 2)._chartSpace.xml


def test_chart_options_bad_values_warn(tmp_path):
    md = "\n# a\n```bar {labels=above gap=900 marker=1 size=2 colors=nope}\n,x\na,1\n```\n"
    deck, _ = make(tmp_path, md)
    msgs = warns(deck, "bad-chart-option")
    assert len(msgs) == 5 and all(d.hint for d in msgs)


def test_chart_grid_token(tmp_path):
    _, prs = make(tmp_path, "style: render.chart_grid=accent\n\n# a\n```column\n,x,y\na,1,2\n```\n")
    assert "E09F1F" in _chart(prs)._chartSpace.xml


def test_hl_names_a_series_or_a_category(tmp_path):
    md = (
        "\n# a\n```column {hl=b colors=primary,secondary}\n,x,y\na,1,2\nb,3,4\n```\n"
        "\n# b\n```line {hl=nope}\n,x,y\na,1,2\nb,3,4\n```\n"
    )
    deck, prs = make(tmp_path, md)
    ser = _chart(prs, 0).plots[0].series
    fills = [s._element.find(qn("c:spPr")).find(qn("a:solidFill"))[0].get("val") for s in ser]
    assert fills[0] == "142B4D" and fills[1] == "E09F1F"  # series b took the emphasis color
    w = warns(deck, "chart-hl")
    assert len(w) == 1 and "series are: a, b" in w[0].hint


# --------------------------------------------------------------------------- steps


STEPS = "\n# S\n@4 steps\n" + "".join(f"## s{i}\n- t{i}\n" for i in range(4))


def test_steps_caption_cycle_and_card_selectors(tmp_path):
    md = (
        'style: steps-arrow.fill=primary,secondary steps.caption="STEP {n}" steps.caption_color=teal\n'
        + STEPS
        + "```css\n.steps-card:nth-child(even) { background: #FFDDEE }\n```\n"
    )
    _, prs = make(tmp_path, md)
    arrows = sorted(shapes(prs, 0, "Step "), key=lambda s: s.name)
    a = [s for s in arrows if s.name.endswith("arrow")]
    c = [s for s in arrows if s.name.endswith("card")]
    assert [fill_hex(s) for s in a] == ["142B4D", "1F5FA8", "142B4D", "1F5FA8"]
    assert [fill_hex(s) for s in c] == ["EEF2F7", "FFDDEE", "EEF2F7", "FFDDEE"]
    texts = [s.text_frame.text for s in shapes(prs, 0, "Heading")]  # (the card look: caption + heading)
    assert any(t.startswith("STEP 1") for t in texts)
    text, _ = import_pptx(tmp_path / "d.pptx")  # the caption is a token: it is not imported as content
    assert "STEP 1" not in text and "@4 steps" in text


def test_steps_num_flag_uses_the_default_caption(tmp_path):
    _, prs = make(tmp_path, STEPS.replace("@4 steps", "@4 steps num"))
    allt = " ".join(s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame)
    assert "STEP 1" in allt and "STEP 4" in allt


def test_chevron_shape_token_draws_pentagons_and_round_trips(tmp_path):
    _, prs = make(tmp_path, "style: render.chevron_shape=pentagon\n" + STEPS)
    arrows = [s for s in shapes(prs, 0, "Step ") if s.name.endswith("arrow")]
    assert len(arrows) == 4 and all("homePlate" in s._element.xml for s in arrows)
    text, _ = import_pptx(tmp_path / "d.pptx")
    assert "@4 steps" in text
    _, plain = make(tmp_path, STEPS, "plain")
    assert all(
        'prst="chevron"' in s._element.xml for s in shapes(plain, 0, "Step ") if s.name.endswith("arrow")
    )


def test_steps_bad_cycle_value_warns(tmp_path):
    deck, _ = make(tmp_path, "style: steps-arrow.fill=primary,notacolor\n" + STEPS)
    assert warns(deck, "bad-token")


# --------------------------------------------------------------------------- rows


ROWS = "\n# R\n@rows\n> lead\n1. one\n2. two\n3. three\n"


def test_rows_draws_numbered_bars(tmp_path):
    _, prs = make(tmp_path, "style: rows-num.fill=primary,secondary\n" + ROWS)
    bars = [s for s in shapes(prs, 0, "Row ") if not s.name.endswith("num")]
    nums = [s for s in shapes(prs, 0, "Row ") if s.name.endswith("num")]
    assert len(bars) == 3 and len(nums) == 3
    assert [n.text_frame.text for n in nums] == ["1", "2", "3"]
    assert [fill_hex(n) for n in nums] == ["142B4D", "1F5FA8", "142B4D"]
    assert bars[0].text_frame.text == "one" and bars[0].height == nums[0].height
    assert bars[1].top > bars[0].top + bars[0].height  # a gap between the bars
    assert nums[0].left == bars[0].left


def test_rows_token_applies_without_the_flag_and_off_by_default(tmp_path):
    plain = "\n# R\n1. one\n2. two\n"
    _, off = make(tmp_path, plain, "off")
    assert not shapes(off, 0, "Row ")
    _, on = make(tmp_path, "style: layout.rows=on\n" + plain, "on")
    assert len(shapes(on, 0, "Row ")) == 4


def test_rows_skipped_with_a_hint_when_the_slide_does_not_fit(tmp_path):
    deck, prs = make(tmp_path, "\n# R\n@rows\n1. one\n- bullet\n")
    d = warns(deck, "rows-skipped")
    assert len(d) == 1 and d[0].hint and d[0].level == "info"
    assert not shapes(prs, 0, "Row ")


def test_rows_round_trip(tmp_path):
    make(tmp_path, ROWS, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "@rows" in text and "1. one" in text and "## one" not in text


# --------------------------------------------------------------------------- cover


def test_cover_bar_rule_and_no_band_over_a_slide_bg(tmp_path):
    md = (
        "style: cover.band_h=61% cover.rule=accent cover.bar=teal cover.band=none title.band=primary\n\n"
        "# Title\n@cover bg=primary dark\n## sub\n\n# next\n- a\n"
    )
    deck, prs = make(tmp_path, md)
    r = {fill_hex(s): s for s in rules(prs, 0)}
    assert set(r) == {"E09F1F", "2A9D8F"}  # no band shape although title.band is set
    assert r["E09F1F"].width == prs.slide_width  # the rule is full-bleed over a bg
    bar = r["2A9D8F"]
    title = next(s for s in prs.slides[0].shapes if s.name == "Title")
    assert bar.left < title.left and bar.width > 0


# --------------------------------------------------------------------------- pinned sizes


def test_pinned_size_stays_while_others_grow(tmp_path):
    base = "sizes: heading=20 body=16\n\n# T\n## A\n- x\n## B\n- y\n"
    _, grown = make(tmp_path, base, "g")
    _, pinned = make(tmp_path, base.replace("heading=20", "heading=20!"), "p")

    def head(prs):
        return (
            next(s for s in prs.slides[0].shapes if s.name.startswith("Heading"))
            .text_frame.paragraphs[0]
            .runs[0]
            .font.size.pt
        )

    assert head(pinned) == 20.0
    d = parse(HEAD + "sizes: heading=20!\n\n# T\n- a\n")
    assert deck_theme(d, ".")[0].pinned == ["heading"]
    assert not [x for x in d.diagnostics if x.rule == "bad-token"]


def test_sizes_bad_pin_warns(tmp_path):
    deck, _ = make(tmp_path, "sizes: heading=big!\n\n# T\n- a\n")
    assert warns(deck, "bad-token")


# --------------------------------------------------------------------------- footer on @html slides


def test_html_slide_footer_is_opt_in(tmp_path):
    html = '\n# H\n@html\n```html\n<div style="height:100%;background:#eee"><h1>Hi</h1></div>\n```\n'
    base = "footer: ACME\nnum: on\n"
    _, off = make(tmp_path, base + html, "off")
    assert "Footer" not in {s.name for s in off.slides[0].shapes}
    _, on = make(tmp_path, base + "style: layout.html_footer=on\n" + html, "on")
    assert {"Footer", "Slide Number"} <= {s.name for s in on.slides[0].shapes}
