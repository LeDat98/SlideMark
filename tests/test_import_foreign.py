"""Import of decks that SlideMark did not build: free text boxes, loose cards, bus connectors, scoring."""

from __future__ import annotations

import runpy
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

from slidemark.importer import import_pptx

ROOT = Path(__file__).parent.parent


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs, prs.slides.add_slide(prs.slide_layouts[6])


def _text(slide, x, y, w, h, lines, size=16, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
    return tb


def _rect(slide, x, y, w, h, rgb, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(rgb)
    s.line.fill.background()
    return s


def _import(prs, tmp_path) -> str:
    path = tmp_path / "foreign.pptx"
    prs.save(path)
    text, diags = import_pptx(path)
    assert not [d for d in diags if d.level == "error"]
    return text


def test_rect_with_overlapping_textbox_is_one_box(tmp_path):
    prs, s = _deck()
    _text(s, 0.6, 0.3, 12, 0.8, "Plans", size=30, bold=True)
    for i, (h, body) in enumerate([("Basic", "Ten seats"), ("Pro", "Fifty seats")]):
        x = 0.6 + i * 6.2
        _rect(s, x, 1.6, 5.8, 3.0, "F1F4F9")
        _text(s, x + 0.2, 1.8, 5.4, 0.6, h, size=22, bold=True)
        _text(s, x + 0.2, 2.6, 5.4, 1.5, body, size=16)
    text = _import(prs, tmp_path)
    assert text.startswith("# Plans")
    assert "## Basic\nTen seats" in text
    assert "## Pro\nFifty seats" in text


def test_free_text_title_not_a_kpi_number(tmp_path):
    prs, s = _deck()
    _rect(s, 0, 0, 13.333, 1.1, "1F2A44")
    _text(s, 0.6, 0.15, 12, 0.8, "The problem", size=30, bold=True)
    _rect(s, 0.6, 1.9, 3.8, 4.0, "F1F4F9")
    _text(s, 0.8, 2.2, 3.4, 1.0, "20+ hrs", size=48, bold=True)
    _text(s, 0.9, 3.8, 3.2, 1.5, "Manual data entry every month.", size=18)
    text = _import(prs, tmp_path)
    assert text.startswith("# The problem")
    assert "20+ hrs" in text and "Manual data entry" in text


def test_small_bottom_text_is_a_footnote(tmp_path):
    prs, s = _deck()
    _text(s, 0.6, 0.3, 12, 0.8, "Results", size=30, bold=True)
    _text(s, 0.6, 1.8, 12, 3, ["- one", "- two"], size=18)
    _text(s, 0.6, 6.95, 12, 0.35, "Source: internal survey", size=10)
    text = _import(prs, tmp_path)
    assert "※ Source: internal survey" in text


def test_box_body_paragraphs_survive(tmp_path):
    prs, s = _deck()
    _text(s, 0.6, 0.3, 12, 0.8, "Cards", size=30, bold=True)
    card = _rect(s, 0.6, 1.8, 5.0, 2.0, "2E6FD8", MSO_SHAPE.ROUNDED_RECTANGLE)
    tf = card.text_frame
    tf.text = "Headline"
    tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.size = Pt(22)
    p = tf.add_paragraph()
    p.text = "Second line"
    p.runs[0].font.size = Pt(14)
    text = _import(prs, tmp_path)
    assert "Headline" in text and "Second line" in text


def test_status_pill_over_empty_cell_goes_into_the_table(tmp_path):
    prs, s = _deck()
    _text(s, 0.6, 0.3, 12, 0.8, "Status", size=30, bold=True)
    gf = s.shapes.add_table(3, 2, Inches(0.6), Inches(2), Inches(8), Inches(1.5))
    t = gf.table
    for r, (a, b) in enumerate([("Item", "State"), ("Alpha", ""), ("Beta", "")]):
        t.cell(r, 0).text = a
        t.cell(r, 1).text = b
    pill = _rect(s, 5.2, 3.02, 2.0, 0.3, "2E9E5B", MSO_SHAPE.ROUNDED_RECTANGLE)
    pill.text_frame.text = "Shipped"
    text = _import(prs, tmp_path)
    assert "| Beta | Shipped |" in text.replace("**", "")


def test_bus_connectors_become_links(tmp_path):
    prs, s = _deck()
    _text(s, 0.6, 0.3, 12, 0.8, "Org", size=30, bold=True)
    top = _rect(s, 5.0, 1.5, 3.3, 0.9, "1F2A44")
    top.text_frame.text = "Board"
    for i, n in enumerate(["North", "South", "East"]):
        b = _rect(s, 0.7 + i * 4.2, 3.6, 3.6, 0.9, "2E6FD8")
        b.text_frame.text = n
    for pts in [(6.65, 2.4, 6.65, 3.0), (2.5, 3.0, 10.9, 3.0)]:
        s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *(Inches(v) for v in pts))
    for i in range(3):
        s.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(2.5 + i * 4.2), Inches(3.0), Inches(2.5 + i * 4.2), Inches(3.6)
        )
    text = _import(prs, tmp_path)
    assert "a-b a-c a-d" in text


def test_import_fidelity_scores_one_generated_deck(tmp_path, monkeypatch):
    monkeypatch.syspath_prepend(str(ROOT / "bench"))
    monkeypatch.chdir(tmp_path)
    runpy.run_path(str(ROOT / "bench" / "answers" / "run-2-python-pptx" / "12-agenda.py"))
    orig = tmp_path / "12-agenda.pptx"
    assert orig.exists()
    sys.modules.pop("import_fidelity", None)
    import import_fidelity as fid

    rebuilt, _ = fid.roundtrip(orig, tmp_path)
    assert rebuilt is not None
    score = fid.score_pair(orig, rebuilt)
    assert score["slides"][0] == score["slides"][1]
    assert score["fidelity"] >= 0.9
    assert fid.score_pair(orig, orig)["fidelity"] == 1.0
