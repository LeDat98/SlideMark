"""Vertical balance of the body: spread the leftover, then center (``layout.body_valign``), lead slot,
card anchoring, type unification, callout padding and the chart / footnote gap."""

from __future__ import annotations

from pptx import Presentation

from slidemark.ir import Chart, Container, Deck, Series, Slide, Table, Text
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.theme import get_theme
from slidemark.units import to_emu

from .test_layout_policies import H, T, box, bullets, cards, cell, of


def _theme(**kw):
    th = get_theme("jp-business").model_copy(deep=True)
    for k, v in kw.items():
        setattr(th.layout, k, v)
    return th


def _lay(slide, **kw):
    return layout_slide(slide, Deck(slides=[slide]), _theme(**kw), 0)


def _table_slide(n=5, **kw):
    rows = [[cell("項目"), cell("値")]] + [[cell("a"), cell("12")] for _ in range(n)]
    return Slide(title=T("t", "title"), lead=T("lead", "lead"), elements=[Table(rows=rows)], **kw)


def _span(placed):
    body = [
        p
        for p in placed
        if p.element is not None and getattr(p.element, "role", "") not in ("title", "lead", "caption")
    ]
    body = [p for p in body if getattr(p.element, "id", None) != "band"]
    return min(p.y for p in body), max(p.y + p.h for p in body)


def test_sparse_table_rows_spread_and_block_is_centered():
    s = _table_slide()
    top = _lay(s, body_valign="top", table_text_step=0)  # ungrown text: rows stretch (no row cap)
    mid = _lay(s, table_text_step=0)
    (t0,), (t1,) = of(top, Table), of(mid, Table)
    assert t0.h < t1.h <= 1.4 * t0.h + 2  # rows grew by at most body_spread_max
    assert sum(t1.element.attrs["_row_h"]) == t1.h
    lead = next(p for p in mid if isinstance(p.element, Text) and p.element.role == "lead")
    y0, y1 = _span(mid)
    th = _theme()
    bottom = H - to_emu(th.margin_y)
    above, below = y0 - (lead.y + lead.h), bottom - y1
    assert abs(above - below) < 0.05 * H  # centered between the lead and the bottom


def test_valign_top_keeps_the_block_at_the_top_and_threshold_gates_auto():
    s = _table_slide()
    lead_bottom = lambda pl: next(p for p in pl if isinstance(p.element, Text) and p.element.role == "lead")  # noqa: E731
    top = _lay(s, body_valign="top", table_text_step=0)
    lb = lead_bottom(top)
    assert _span(top)[0] - (lb.y + lb.h) == to_emu(_theme().layout.top_gap)
    # a threshold above the leftover turns auto off: same as top
    off = _lay(s, body_free_max=0.95, table_text_step=0)
    assert _span(off) == _span(top)
    assert _span(_lay(s, body_valign="center", table_text_step=0)) != _span(top)


def test_full_slide_is_not_moved():
    dense = [f"line {i} " + "word " * 25 for i in range(9)]
    s = Slide(title=T("t", "title"), grid="2", elements=[box("a", *dense), box("b", *dense)])
    assert _span(_lay(s)) == _span(_lay(s, body_valign="top"))


def test_tree_rank_gap_grows_and_links_follow(tmp_path):
    from slidemark.ir import Link

    kids = [box(n, "x") for n in "abcd"]
    s = Slide(
        title=T("t", "title"),
        lead=T("lead", "lead"),
        elements=kids,
        links=[Link(src=0, dst=1), Link(src=1, dst=2), Link(src=1, dst=3)],
    )
    top, mid = _lay(s, body_valign="top"), _lay(s)
    c0, c1 = cards(top), cards(mid)
    ys0, ys1 = sorted({c.y for c in c0}), sorted({c.y for c in c1})
    assert ys1[1] - ys1[0] >= ys0[1] - ys0[0]
    assert all(a.h == b.h for a, b in zip(c0, c1, strict=True))  # cards keep hugging their text
    lines = [p for p in mid if p.element is not None and getattr(p.element, "shape", "") == "line"]
    card_bottoms = {c.y + c.h for c in c1}
    card_tops = {c.y for c in c1}
    assert any(abs(p.y - b) <= 2 for p in lines for b in card_bottoms)  # connectors start at a card bottom
    assert any(abs(p.y + p.h - t) <= 2 for p in lines for t in card_tops)  # ... and end at a card top
    deck = Deck(slides=[s])
    out = tmp_path / "d.pptx"
    render(deck, [mid], _theme(), out)
    assert Presentation(str(out)).slides[0].shapes


def test_rendered_table_rows_match_placed_geometry(tmp_path):
    s = _table_slide()
    placed = _lay(s)
    (tp,) = of(placed, Table)
    out = tmp_path / "d.pptx"
    render(Deck(slides=[s]), [placed], _theme(), out)
    gf = next(sh for sh in Presentation(str(out)).slides[0].shapes if sh.has_table)
    assert abs(gf.top - tp.y) <= 2
    assert abs(sum(r.height for r in gf.table.rows) - tp.h) <= len(gf.table.rows) + 2


def _chart():
    return Chart(kind="column", categories=["a", "b"], series=[Series(name="s", values=[1, 2])])


def test_card_beside_a_chart_is_top_anchored():
    s = Slide(
        title=T("t", "title"),
        grid="2",
        elements=[_chart(), box("Insight", "one", "two")],
    )
    placed = _lay(s, body_valign="top")
    (card,) = cards(placed)
    body = next(p for p in placed if isinstance(p.element, Text) and p.element.role == "body")
    head = next(p for p in placed if isinstance(p.element, Text) and p.element.role == "heading")
    assert body.y - (head.y + head.h) < 0.1 * card.h  # right under the heading, not centered in the card
    assert not get_theme("default").layout.center_beside


def test_table_text_is_not_smaller_than_box_text():
    rows = [[cell("KPI"), cell("v")]] + [[cell("x"), cell("1")] for _ in range(3)]
    s = Slide(
        title=T("t", "title"),
        grid="3",
        elements=[box("a", "x", "y"), box("b", "x", "y"), box("c", "x", "y"), Table(rows=rows)],
    )
    for unify in (True, False):
        placed = _lay(s, body_size_unify=unify)
        (tp,) = of(placed, Table)
        txt = [p for p in of(placed, Text) if p.element.role == "body"]
        box_pt = max(p.style.font_size * p.font_scale for p in txt)
        tab_pt = tp.style.font_size * tp.font_scale
        if unify:
            assert tab_pt >= box_pt - 0.01


def test_lead_slot_is_reserved_when_most_slides_have_a_lead():
    with_lead = Slide(title=T("t", "title"), lead=T("lead", "lead"), elements=[bullets("a", "b", "c")])
    without = Slide(title=T("t", "title"), elements=[bullets("a", "b", "c")])
    deck = Deck(slides=[with_lead, with_lead, without])
    th = _theme(body_valign="top")

    def body_y(i, th=th):
        pl = layout_slide(deck.slides[i], deck, th, i)
        return min(p.y for p in pl if isinstance(p.element, Text) and p.element.role == "body")

    assert body_y(2) == body_y(0)
    off = _theme(body_valign="top", reserve_lead="off")
    assert body_y(2, off) < body_y(0, off)
    few = Deck(slides=[with_lead, without, without])
    pl = layout_slide(few.slides[1], few, th, 1)
    assert min(p.y for p in pl if isinstance(p.element, Text) and p.element.role == "body") < body_y(0)


def test_callout_inside_a_card_has_padding_and_readable_text():
    note = T("note", "body", classes=["callout", "note"])
    s = Slide(
        title=T("t", "title"),
        grid="2",
        elements=[_chart(), Container(title=T("h", "heading"), children=[bullets("a"), note])],
    )
    placed = _lay(s, body_valign="top")
    (cp,) = [p for p in of(placed, Text) if "callout" in p.element.classes]
    for f in ("padding",):
        assert to_emu(getattr(cp.style, f)) >= to_emu("6pt") - 1
    assert cp.style.font_size * cp.font_scale >= get_theme("jp-business").sizes["footnote"] - 1e-6


def test_chart_keeps_the_footnote_gap():
    s = Slide(title=T("t", "title"), elements=[_chart()], footnotes=[T("出所: x", "footnote")])
    placed = _lay(s, body_valign="top")
    (ch,) = of(placed, Chart)
    fn = next(p for p in placed if isinstance(p.element, Text) and p.element.role == "footnote")
    assert fn.y - (ch.y + ch.h) >= to_emu(_theme().layout.footnote_gap) - 1
    wide = _lay(s, body_valign="top", footnote_gap="0.4in")
    (ch2,) = of(wide, Chart)
    assert fn.y - (ch2.y + ch2.h) >= to_emu("0.4in") - 1
