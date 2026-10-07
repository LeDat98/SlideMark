"""DL3d part 2 (lane F): exact ``@steps`` geometry, the quote card's bar / by-line / height, badge shapes
and card shadows: the tokens (parser), the render (reopen the .pptx), the importer (synthetic decks)."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.parser import parse
from slidemark.template import deck_theme

IN = 914400
NAVY, BLUE, TEAL, AMBER, LIGHT, GREY = "122B4A", "1F5FA8", "1B8A8F", "E08A1E", "EEF2F7", "6B7582"


def _shapes(tmp_path: Path, md: str, slide: int = 0) -> dict:
    out = tmp_path / "out.pptx"
    build(md, out, base_dir=tmp_path)
    return {sh.name: sh for sh in Presentation(str(out)).slides[slide].shapes}


def _diags(md: str):
    """Parse + layout + lint (what ``check`` runs): the diagnostics without writing a .pptx."""
    from slidemark.cli import _add_layout_diagnostics

    deck = parse(md)
    _add_layout_diagnostics(deck, "d.md")
    return deck.diagnostics


def _geom(sh) -> str:
    return sh._element.spPr.find(qn("a:prstGeom")).get("prst")


# --------------------------------------------------------------------------- tokens (parser)

STEPS = """style: steps-arrow.h=0.45in steps-card.h=2.7in steps.gap=0.2in conclusion.h=1.05in{extra}

# Plan
@3 steps
## 2027
- one
## 2028
- two
## 2029
- three
> done
"""


def test_tokens_parse_without_diagnostics():
    md = (
        "style: steps-arrow.h=0.45in steps-card.h=2.7in steps.gap=0.2in conclusion.h=1.05in "
        "quote.bar=#E08A1E "
        "quote.bar_w=0.15in quote.by.align=right quote.fill=#EEF2F7 quote.h=full rows-num.shape=circle "
        "box.num.shape=rounded\n\n# t\n"
    )
    deck = parse(md)
    assert not [d for d in deck.diagnostics if d.rule in ("bad-token", "unknown-token")]
    assert deck.tokens["steps_card_h"] == "2.7in" and deck.tokens["quote_h"] == "full"
    theme, _ = deck_theme(deck, ".")
    assert theme.steps_arrow_h == "0.45in" and theme.conclusion_h == "1.05in"
    assert theme.layout.steps_gap == "0.2in"  # `steps.gap=` is also the layout token every pass reads
    assert theme.rows_num_shape == "circle" and theme.box_num_shape == "rounded"
    assert theme.quote_by_align == "right" and theme.quote_bar == "#E08A1E"


@pytest.mark.parametrize(
    "bad",
    [
        "steps-arrow.h=tall",
        "steps-card.h=-2in",
        "steps.gap=wide",
        "conclusion.h=zz",
        "quote.bar=notacolour9",
        "quote.bar_w=thick",
        "quote.by.align=up",
        "quote.h=huge",
        "rows-num.shape=hexagon",
        "box.num.shape=star",
    ],
)
def test_bad_values_warn_with_a_hint(bad):
    bads = [d for d in parse(f"style: {bad}\n\n# t\n").diagnostics if d.rule == "bad-token"]
    assert len(bads) == 1 and bads[0].hint and bad.split("=")[0] in bads[0].message


# --------------------------------------------------------------------------- render: @steps geometry


def test_steps_arrow_card_gap_and_bar_heights_are_exact(tmp_path):
    sh = _shapes(tmp_path, STEPS.format(extra=""))
    for k in (1, 2, 3):
        assert sh[f"Step {k} arrow"].height == round(0.45 * IN)
        assert sh[f"Step {k} card"].height == round(2.7 * IN)
        gap = sh[f"Step {k} card"].top - (sh[f"Step {k} arrow"].top + sh[f"Step {k} arrow"].height)
        assert abs(gap - round(0.2 * IN)) <= 2
    assert sh["Conclusion"].height == round(1.05 * IN)
    assert (
        sh["Conclusion"].top >= sh["Step 1 card"].top + sh["Step 1 card"].height
    )  # the bar follows the cards


def test_without_the_tokens_the_cards_still_stretch(tmp_path):
    sh = _shapes(tmp_path, "# Plan\n@3 steps\n## a\n- one\n## b\n- two\n## c\n- three\n> done\n")
    assert sh["Step 1 card"].height > 1.5 * IN and sh["Step 1 card"].height != round(2.7 * IN)  # stretched


def test_a_pinned_card_height_beats_the_stretch_and_growth_passes(tmp_path):
    sh = _shapes(
        tmp_path, "style: steps-card.h=1.2in\n\n# Plan\n@3 steps\n## a\n- one\n## b\n- two\n## c\n- 3\n> x\n"
    )
    assert all(sh[f"Step {k} card"].height == round(1.2 * IN) for k in (1, 2, 3))


def test_steps_gap_alone_moves_the_cards(tmp_path):
    base = "# Plan\n@3 steps\n## a\n- one\n## b\n- two\n## c\n- 3\n"
    a = _shapes(tmp_path, "style: steps.gap=0.1in\n\n" + base)
    b = _shapes(tmp_path, "style: steps.gap=0.5in\n\n" + base)
    ga = a["Step 1 card"].top - (a["Step 1 arrow"].top + a["Step 1 arrow"].height)
    gb = b["Step 1 card"].top - (b["Step 1 arrow"].top + b["Step 1 arrow"].height)
    assert abs(ga - 0.1 * IN) <= 3 and abs(gb - 0.5 * IN) <= 3


def test_conclusion_h_is_exact_without_steps(tmp_path):
    sh = _shapes(tmp_path, "style: conclusion.h=0.8in\n\n# T\n- a\n- b\n> take away\n")
    assert sh["Conclusion"].height == round(0.8 * IN)


def test_steps_tokens_without_steps_say_attr_ignored():
    for key in ("steps-card.h=2in", "steps-arrow.h=0.5in", "steps.gap=0.2in"):
        got = [d for d in _diags(f"style: {key}\n\n# T\n- a\n") if d.rule == "attr-ignored"]
        assert got and "@steps" in got[0].hint, key


def test_conclusion_h_without_a_conclusion_says_attr_ignored():
    got = [d for d in _diags("style: conclusion.h=1in\n\n# T\n- a\n") if d.rule == "attr-ignored"]
    assert got and "conclusion" in got[0].hint
    assert not [d for d in _diags("style: conclusion.h=1in\n\n# T\n- a\n> bar\n") if d.rule == "attr-ignored"]


# --------------------------------------------------------------------------- render: quote card

QUOTE = """style: quote.bar=#E08A1E quote.bar_w=0.15in quote.by.align=right quote.h=3in quote.width=1

# Voice
@quote fill=#EEF2F7 size=30
> **Month-end close went from three days to half a day.**
>
> — Miura Works, Controller
"""


def test_quote_card_bar_by_line_and_height(tmp_path):
    sh = _shapes(tmp_path, QUOTE)
    panel, bar = sh["Quote panel"], sh["Quote bar"]
    assert bar.width == round(0.15 * IN) and bar.left == panel.left and bar.height == panel.height
    assert str(bar.fill.fore_color.rgb) == AMBER and _geom(bar) == "rect"
    assert panel.height == round(3 * IN)
    assert abs(panel.width - 12.33 * IN) < 0.3 * IN  # quote.width=1: the card spans the body
    algn = sh["Quote by"].text_frame.paragraphs[0].alignment
    assert algn is not None and "RIGHT" in str(algn)
    assert sh["Quote text"].left >= panel.left + bar.width  # the text starts right of the bar


def test_quote_h_full_and_quote_fill_token(tmp_path):
    md = QUOTE.replace("quote.h=3in", "quote.h=full quote.fill=#DDEEFF").replace(" fill=#EEF2F7", "")
    sh = _shapes(tmp_path, md)
    assert str(sh["Quote panel"].fill.fore_color.rgb) == "DDEEFF"  # the deck-wide token paints the card
    assert sh["Quote panel"].height > 4.5 * IN  # the whole body


def test_quote_without_the_tokens_is_unchanged(tmp_path):
    sh = _shapes(tmp_path, QUOTE.split("\n\n", 1)[1])
    assert "Quote bar" not in sh and "Quote panel" in sh
    assert sh["Quote by"].text_frame.paragraphs[0].alignment is None or "LEFT" in str(
        sh["Quote by"].text_frame.paragraphs[0].alignment
    )


def test_quote_tokens_without_a_quote_say_attr_ignored():
    for key in ("quote.bar=red", "quote.h=full", "quote.by.align=right"):
        got = [d for d in _diags(f"style: {key}\n\n# T\n- a\n") if d.rule == "attr-ignored"]
        assert got and "@quote" in got[0].hint, key


# --------------------------------------------------------------------------- render: badge shapes


def test_rows_badge_shapes(tmp_path):
    md = "style: rows-num.shape={shape}\n\n# T\n@rows\n1. one\n2. two\n"
    assert _geom(_shapes(tmp_path, md.format(shape="square"))["Row 1 num"]) == "rect"
    assert _geom(_shapes(tmp_path, md.format(shape="circle"))["Row 1 num"]) == "ellipse"
    assert _geom(_shapes(tmp_path, md.format(shape="rounded"))["Row 1 num"]) == "roundRect"
    assert _geom(_shapes(tmp_path, md.split("\n\n", 1)[1])["Row 1 num"]) == "rect"  # the default


def test_box_num_badge_shapes(tmp_path):
    md = "style: box.num.shape={shape}\n\n# T\n@2 num\n## a\n- x\n## b\n- y\n"
    assert _geom(_shapes(tmp_path, md.format(shape="square"))["Num 1"]) == "rect"
    assert _geom(_shapes(tmp_path, md.format(shape="rounded"))["Num 1"]) == "roundRect"
    assert _geom(_shapes(tmp_path, md.format(shape="circle"))["Num 1"]) == "ellipse"
    assert _geom(_shapes(tmp_path, md.split("\n\n", 1)[1])["Num 1"]) == "ellipse"  # the default


def test_badge_shape_tokens_without_badges_say_attr_ignored():
    for key in ("rows-num.shape=circle", "box.num.shape=square"):
        assert [d for d in _diags(f"style: {key}\n\n# T\n- a\n") if d.rule == "attr-ignored"], key


# --------------------------------------------------------------------------- render: card shadows


def _shadows(tmp_path: Path, md: str) -> dict[str, bool]:
    sh = _shapes(tmp_path, md)
    return {
        n: bool(s._element.spPr.xpath("./a:effectLst/a:outerShdw"))
        for n, s in sh.items()
        if re.match(r"(Card|Box)", n) or "card" in n.lower()
    }


def test_card_shadow_token_and_box_shadow_off(tmp_path):
    out = tmp_path / "s.pptx"
    build(
        'style: card.shadow="0 2 6 #00000040"\n\n# T\n@3\n## a\n- x\n## b {shadow=off}\n- y\n## c\n- z\n',
        out,
        base_dir=tmp_path,
    )
    xml = etree.tostring(Presentation(str(out)).slides[0]._element).decode()
    assert xml.count("<a:outerShdw") == 2  # two cards have it, the `shadow=off` one does not


# --------------------------------------------------------------------------- importer (synthetic decks)


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(s, x, y, w, h, rgb, shape=MSO_SHAPE.RECTANGLE, shadow=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(rgb)
    sp.line.fill.background()
    sp.shadow.inherit = False
    if shadow:  # (blur pt, dist pt, alpha %)
        eff = sp._element.spPr.find(qn("a:effectLst"))
        sh = etree.SubElement(
            eff,
            qn("a:outerShdw"),
            blurRad=str(shadow[0] * 12700),
            dist=str(shadow[1] * 12700),
            dir="5400000",
            rotWithShape="0",
        )
        clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
        etree.SubElement(clr, qn("a:alpha"), val=str(shadow[2] * 1000))
    return sp


def _text(s, x, y, w, h, paras, size=16, bold=False, color="222B36", align=None, shape=None):
    tb = shape or s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align is not None:
            p.alignment = align
        r = p.add_run()
        r.text = t
        r.font.size, r.font.bold = Pt(size), bold
        r.font.color.rgb = RGBColor.from_string(color)
    return tb


def _chrome(s, title):
    _rect(s, 0, 0, 13.333, 1.15, NAVY)
    _text(s, 0.6, 0.12, 12, 0.9, title, size=28, bold=True, color="FFFFFF")


def _import(prs, tmp_path: Path) -> str:
    path = tmp_path / "foreign.pptx"
    prs.save(path)
    text, diags = import_pptx(path)
    assert not [d for d in diags if d.level == "error"], diags
    return text


def _body(text: str) -> str:
    return text[text.index("\n# ") :]


def test_importer_states_the_steps_geometry(tmp_path):
    prs = _deck()
    s = _slide(prs)
    _chrome(s, "Product plan")
    for i, (lab, body) in enumerate([("2027", "Reconcile"), ("2028", "Forecast"), ("2029", "Templates")]):
        x = 0.6 + i * 4.1
        arrow = _rect(s, x, 1.55, 3.9, 0.45, (BLUE, TEAL, AMBER)[i], MSO_SHAPE.PENTAGON)
        _text(s, 0, 0, 0, 0, [lab], size=24, bold=True, color="FFFFFF", shape=arrow)
        _rect(s, x, 2.25, 3.9, 2.6, LIGHT)
        _text(s, x + 0.25, 2.4, 3.4, 2.3, body, size=26, bold=True, color=NAVY)
    _rect(s, 0.6, 5.2, 12.13, 0.9, NAVY)
    _text(s, 0.9, 5.2, 11.5, 0.9, "Big release in Q1", size=26, bold=True, color="FFFFFF")
    body = _body(_import(prs, tmp_path))
    assert "@3 steps" in body
    for tok in ("steps-arrow.h=0.45in", "steps-card.h=2.6in", "steps.gap=0.25in", "conclusion.h=0.9in"):
        assert tok in body, tok
    text = _import(prs, tmp_path)
    sh = _shapes(tmp_path, text)
    assert sh["Step 1 arrow"].height == round(0.45 * IN) and sh["Step 1 card"].height == round(2.6 * IN)
    assert sh["Conclusion"].height == round(0.9 * IN)


def test_importer_states_the_quote_bar_by_align_and_height(tmp_path):
    prs = _deck()
    s = _slide(prs)
    _chrome(s, "Voice")
    _rect(s, 0.6, 1.7, 12.13, 4.9, LIGHT)
    _rect(s, 0.6, 1.7, 0.15, 4.9, AMBER)
    _text(s, 1.0, 1.7, 1.5, 1.6, "“", size=96, bold=True, color=AMBER)
    _text(s, 1.6, 2.5, 10.1, 2.5, "Close went from three days to half a day.", size=38, bold=True, color=NAVY)
    from pptx.enum.text import PP_ALIGN

    _text(s, 1.6, 5.3, 10.1, 0.8, "— Miura Works, Controller", size=24, color=GREY, align=PP_ALIGN.RIGHT)
    text = _import(prs, tmp_path)
    body = _body(text)
    for tok in ("quote.bar=#E08A1E", "quote.bar_w=0.15in", "quote.h=4.9in", "quote.by.align=right"):
        assert tok in body, tok
    assert "quote.width=" in body and "@quote" in body
    sh = _shapes(tmp_path, text)
    assert sh["Quote bar"].width == round(0.15 * IN) and sh["Quote panel"].height == round(4.9 * IN)
    assert "RIGHT" in str(sh["Quote by"].text_frame.paragraphs[0].alignment)


def test_importer_writes_the_badge_shape_it_saw(tmp_path):
    prs = _deck()
    s = _slide(prs)
    _chrome(s, "Next steps")
    for i, col in enumerate((BLUE, TEAL, AMBER)):
        y = 1.6 + i * 1.3
        _rect(s, 0.6, y, 12.13, 1.12, LIGHT)
        badge = _rect(s, 0.6, y, 1.1, 1.12, col, MSO_SHAPE.OVAL)
        _text(s, 0, 0, 0, 0, [str(i + 1)], size=26, bold=True, color="FFFFFF", shape=badge)
        _text(s, 2.0, y, 10.3, 1.12, f"Approve item {i + 1}", size=30, bold=True, color=NAVY)
    body = _body(_import(prs, tmp_path))
    assert "@rows" in body and "rows-num.shape=circle" in body


def test_importer_square_badges_beside_a_panel_say_box_num_square(tmp_path):
    prs = _deck()
    s = _slide(prs)
    _chrome(s, "Pricing")
    _rect(s, 0.6, 1.6, 5.3, 5.1, NAVY)
    _text(s, 0.9, 1.8, 4.7, 0.5, "Standard plan", size=16, color="C9D6E6")
    _text(s, 0.9, 2.6, 4.7, 1.0, "9,800", size=40, bold=True, color="FFFFFF")
    _text(s, 0.9, 4.3, 4.7, 1.1, "10,800", size=48, bold=True, color=AMBER)
    for i, (lab, col) in enumerate([("Raise in October", BLUE), ("Existing customers wait", TEAL)]):
        y = 1.6 + i * 1.35
        _rect(s, 6.2, y, 6.5, 1.15, LIGHT)
        badge = _rect(s, 6.2, y, 0.8, 1.15, col)
        _text(s, 0, 0, 0, 0, [str(i + 1)], size=26, bold=True, color="FFFFFF", shape=badge)
        _text(s, 7.15, y, 5.4, 1.15, lab, size=21, bold=True, color=NAVY)
    body = _body(_import(prs, tmp_path))
    assert " num" in body and "box.num.shape=square" in body


def _cards(tmp_path, shadows):
    prs = _deck()
    s = _slide(prs)
    _chrome(s, "Cards")
    for i, shadow in enumerate(shadows):
        x = 0.6 + (i % 2) * 6.2
        y = 1.6 + (i // 2) * 2.7
        _rect(s, x, y, 5.9, 2.4, LIGHT, shadow=shadow)
        _text(s, x + 0.2, y + 0.2, 5.5, 2.0, f"Card {i + 1}", size=24, bold=True, color=NAVY)
    return _import(prs, tmp_path)


def test_a_shadow_most_cards_share_is_card_shadow_the_odd_one_says_off(tmp_path):
    sh = (6, 3, 25)
    text = _cards(tmp_path, [sh, sh, sh, None])
    head = text.split("\n\n# ")[0]
    assert re.search(r'(card\.)?shadow="0 3 6 #00000040"', head), head
    assert text.count("shadow=off") == 1
    assert "shadow=" not in text.split("# Cards", 1)[1].replace("shadow=off", "")
    out = tmp_path / "back.pptx"
    build(text, out, base_dir=tmp_path)
    xml = etree.tostring(Presentation(str(out)).slides[0]._element).decode()
    assert xml.count("<a:outerShdw") == 3


def test_differing_shadows_are_stated_per_box(tmp_path):
    text = _cards(tmp_path, [(6, 3, 25), (12, 6, 40)])
    assert "card.shadow" not in text.split("\n\n# ")[0]
    assert 'shadow="0 3 6 #00000040"' in text and 'shadow="0 6 12 #00000066"' in text


def test_cards_without_shadows_state_nothing(tmp_path):
    text = _cards(tmp_path, [None, None, None, None])
    assert "shadow" not in text


def test_importer_states_the_rows_geometry_and_badge_colours(tmp_path):
    prs = _deck()
    s = _slide(prs)
    _chrome(s, "Next steps")
    for i, col in enumerate((BLUE, TEAL, AMBER)):
        y = 1.6 + i * 1.4
        _rect(s, 0.6, y, 12.13, 1.12, LIGHT)
        badge = _rect(s, 0.6, y, 1.1, 1.12, col)
        _text(s, 0, 0, 0, 0, [str(i + 1)], size=26, bold=True, color="FFFFFF", shape=badge)
        _text(s, 2.0, y, 10.3, 1.12, f"Approve item {i + 1}", size=30, bold=True, color=NAVY)
    body = _body(_import(prs, tmp_path))
    assert "layout.rows_h=1.12in" in body and "layout.rows_gap=0.28in" in body
    assert "rows-num.shape" not in body  # square is the default: nothing to say


# --------------------------------------------------------------------------- fuzz: never raise


@pytest.mark.parametrize(
    "tokens",
    [
        "steps-card.h=50in steps-arrow.h=0.01in steps.gap=9in",
        "steps-card.h=0.1in conclusion.h=20in",
        "quote.h=50in quote.bar_w=50in quote.width=3",
        "quote.h=0.1pt quote.bar=none quote.fill=none",
        "steps-card.h=none steps-arrow.h=auto quote.by.align=none",
    ],
)
def test_odd_values_build_without_raising(tmp_path, tokens):
    out = tmp_path / "fuzz.pptx"
    steps = f"style: {tokens}\n\n# a\n@3 steps\n## x\n- one\n## y\n- two\n## z\n- three\n> bar\n"
    quote = f"style: {tokens}\n\n# q\n@quote fill=#EEF2F7\n> hello world\n>\n> — me\n"
    for md in (steps, quote):
        build(md, out, base_dir=tmp_path)
        assert Presentation(str(out)).slides
