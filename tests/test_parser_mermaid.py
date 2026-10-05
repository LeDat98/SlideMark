from pptx import Presentation

from slidemark.build import build
from slidemark.ir import Container, Raw
from slidemark.layout.grid import parse_spec
from slidemark.parser import parse

from .helpers import needs_soffice


def diagram(src: str):
    deck = parse(f"# T\n```mermaid\n{src}\n```\n")
    el = deck.slides[0].elements[0]
    return el, deck


def names(c: Container) -> list[str]:
    return [t.paragraphs[0].plain for t in c.children]


def edges(c: Container) -> list[tuple[str, str]]:
    n = names(c)
    return [(n[ln.src], n[ln.dst]) for ln in c.links]


def valid_grid(c: Container) -> None:
    gs = parse_spec(c.grid, len(c.children))
    assert gs is not None and not gs.errors
    assert gs.areas is not None and set(gs.areas) == set(range(len(c.children)))


def test_td_chain_is_rows():
    c, deck = diagram("graph TD\nA-->B-->C")
    assert isinstance(c, Container) and c.classes == ["diagram"]
    assert c.grid == "a/b/c"
    assert names(c) == ["A", "B", "C"]
    assert edges(c) == [("A", "B"), ("B", "C")]
    assert all(t.classes == ["node"] for t in c.children)
    assert not deck.diagnostics
    valid_grid(c)


def test_lr_chain_is_columns():
    c, _ = diagram("flowchart LR\nA --> B --> C")
    assert c.grid == "abc"
    valid_grid(c)


def test_shapes_labels_and_decision():
    c, _ = diagram(
        'graph TD\nA[Start] --> B{Ok?}\nB -->|yes| C(Done)\nB -- no --> D["Re try"]\nB -.-> E([Stadium])\nB ==> F((Dot))'  # noqa: E501
    )
    assert names(c) == ["Start", "Ok?", "Done", "Re try", "Stadium", "Dot"]
    assert c.children[1].classes == ["node", "decision"]
    assert c.children[2].classes == ["node", "round"]
    assert c.children[4].classes == ["node", "round"]
    labels = {(ln.src, ln.dst): ln.label for ln in c.links}
    assert labels[(1, 2)] == "yes" and labels[(1, 3)] == "no" and labels[(0, 1)] is None
    valid_grid(c)


def test_plain_line_has_no_arrow():
    c, _ = diagram("graph LR\nA --- B\nB --> C")
    assert [ln.arrow for ln in c.links] == [False, True]


def test_ampersand_and_semicolons_and_comments():
    c, _ = diagram("graph TD; %% comment\nA & B --> C; C --> D %% tail\n")
    assert names(c) == ["A", "B", "C", "D"]
    assert sorted(edges(c)) == [("A", "C"), ("B", "C"), ("C", "D")]
    assert c.grid == "ab/c./d."
    valid_grid(c)


def test_same_rank_siblings_are_centered_and_rows_equal_width():
    c, _ = diagram("graph TD\nA-->B\nA-->C\nA-->D")
    rows = c.grid.split("/")
    assert len({len(r) for r in rows}) == 1
    assert rows[0] == ".a." and rows[1] == "bcd"


def test_rank_is_longest_path():
    c, _ = diagram("graph TD\nA-->B-->C\nA-->C")
    assert c.grid == "a/b/c"


def test_cycle_does_not_hang():
    c, _ = diagram("graph TD\nA-->B\nB-->C\nC-->A")
    assert c.grid == "a/b/c"
    assert len(c.links) == 3
    valid_grid(c)


def test_bottom_up_and_right_to_left():
    bt, _ = diagram("graph BT\nA-->B")
    assert bt.grid == "b/a"
    rl, _ = diagram("graph RL\nA-->B")
    assert rl.grid.count("/") == 0 and len(rl.grid) == 2
    valid_grid(rl)
    valid_grid(bt)


def test_single_node():
    c, _ = diagram("graph TD\nA[Only]")
    assert c.grid is None and names(c) == ["Only"]


def test_ignored_lines_give_one_info():
    c, deck = diagram(
        "graph TD\nsubgraph S\nA-->B\nend\nclassDef x fill:#f00\nstyle A fill:#fff\nclick A call f()"
    )
    assert names(c) == ["A", "B"]
    infos = [d for d in deck.diagnostics if d.rule == "mermaid-ignored"]
    assert len(infos) == 1 and infos[0].level == "info" and infos[0].hint


def test_too_big_is_raw_with_warning():
    src = "graph TD\n" + "\n".join(f"N{i}-->N{i + 1}" for i in range(27))
    el, deck = diagram(src)
    assert isinstance(el, Raw) and el.kind == "mermaid"
    d = [x for x in deck.diagnostics if x.rule == "diagram-too-big"]
    assert d and d[0].level == "warning" and d[0].hint


def test_26_nodes_ok():
    src = "graph TD\n" + "\n".join(f"N{i}-->N{i + 1}" for i in range(25))
    c, _ = diagram(src)
    assert isinstance(c, Container) and len(c.children) == 26
    valid_grid(c)


def test_unsupported_diagrams_stay_raw():
    for src in (
        "sequenceDiagram\nA->>B: hi",
        "gantt\ntitle x",
        "graph TD\nA[unclosed --> B",
        'pie\n"a": 1',
        "",
    ):
        el, deck = diagram(src)
        assert isinstance(el, Raw) and el.kind == "mermaid", src
        d = [x for x in deck.diagnostics if x.rule == "mermaid-unsupported"]
        assert d and d[0].level == "info" and d[0].hint, src


def test_quoted_text_keeps_special_characters():
    c, _ = diagram('graph LR\nA["a --> b; (c)"] --> B')
    assert names(c)[0] == "a --> b; (c)"


def test_cjk_labels():
    c, _ = diagram("graph LR\nA[申請] -->|承認| B{判定}")
    assert names(c) == ["申請", "判定"] and c.links[0].label == "承認"


@needs_soffice
def test_end_to_end_connectors(tmp_path):
    md = (
        "# Flow\n```mermaid\ngraph TD\nA[Start] --> B{Ok?}\nB -->|yes| C(Done)\n"
        "B -->|no| D[Retry]\nD --> B\nC --> E[End]\n```\n"
    )
    out = tmp_path / "d.pptx"
    deck = build(md, out)
    assert not [d for d in deck.diagnostics if d.level == "error"]
    prs = Presentation(str(out))
    shapes = list(prs.slides[0].shapes)
    lines = [
        s
        for s in shapes
        if s.shape_type is not None and "LINE" in str(s.shape_type) or s.__class__.__name__ == "Connector"
    ]
    assert len(lines) >= 5
    texts = {s.text_frame.text for s in shapes if s.has_text_frame}
    assert {"Start", "Ok?", "Done", "Retry", "End"} <= texts
    from slidemark.preview import pptx_to_pngs

    pngs = pptx_to_pngs(out, tmp_path / "png")
    assert pngs and pngs[0].stat().st_size > 0
