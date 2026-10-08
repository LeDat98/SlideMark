"""Vector icons: the glyph set, custom-geometry rendering, and placement next to headings / kpi numbers."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import icons
from slidemark.ir import Container, Deck, Placed, Shape, Slide, Style, Text
from slidemark.layout import layout_slide
from slidemark.lint import lint
from slidemark.parser import parse
from slidemark.preview import pptx_to_pngs
from slidemark.render import render
from slidemark.theme import get_theme
from slidemark.units import slide_size

from .helpers import needs_soffice

EXPECTED = (
    "check x warning info user users building factory chart money yen target rocket lightbulb gear clock "
    "calendar document mail phone globe lock shield cloud database search star heart truck cart leaf "
    "arrow-up arrow-down arrow-right headphones battery music sparkles smile camera book graduation-cap "
    "home map-pin wifi code cpu bell flag gift coffee plane wrench key play trophy briefcase chat bolt tooth "
    "refresh tag"
).split()
E = 914400


def test_icon_set_is_complete():
    assert sorted(icons.names()) == sorted(EXPECTED)
    assert len(icons.names()) == 62


@pytest.mark.parametrize("name", EXPECTED)
def test_icon_paths_are_well_formed(name):
    layers = icons.commands(name)
    assert layers
    for layer in layers:
        ops = [op for op, _ in layer]
        assert ops[0] == "M" and ops[-1] == "Z"
        for _op, pts in layer:
            for x, y in pts:
                assert -0.01 <= x <= icons.GRID + 0.01 and -0.01 <= y <= icons.GRID + 0.01, (name, x, y)


def test_unknown_icon():
    assert icons.path("nope") is None
    assert icons.commands("nope") == []


def _render_icons(tmp_path, names, fill="primary"):
    th = get_theme("default")
    items = []
    for i, n in enumerate(names):
        c, r = i % 6, i // 6
        items.append(
            Placed(
                element=Shape(shape="icon", attrs={"icon": n}),
                x=int((0.6 + c * 1.55) * E),
                y=int((0.4 + r * 1.2) * E),
                w=int(0.7 * E),
                h=int(0.7 * E),
                style=Style(fill=fill),
            )
        )
    out = tmp_path / "icons.pptx"
    render(Deck(slides=[Slide()]), [items], th, out)
    return out, th


def test_every_icon_renders_as_custom_geometry(tmp_path):
    out, th = _render_icons(tmp_path, EXPECTED + ["not-an-icon"])
    shapes = list(Presentation(str(out)).slides[0].shapes)
    assert len(shapes) == len(EXPECTED)  # the unknown name draws nothing
    assert sorted(s.name for s in shapes) == sorted(f"icon {n}" for n in EXPECTED)
    want = th.color("primary").lstrip("#").upper()
    for s in shapes:
        sp_pr = s._element.spPr
        assert sp_pr.find(qn("a:prstGeom")) is None
        geom = sp_pr.find(qn("a:custGeom"))
        paths = geom.findall(f"{qn('a:pathLst')}/{qn('a:path')}")
        assert paths and all(p.get("w") == "24000" and p.get("h") == "24000" for p in paths)
        tags = {c.tag for p in paths for c in p}
        assert qn("a:moveTo") in tags and qn("a:close") in tags
        assert qn("a:lnTo") in tags or qn("a:cubicBezTo") in tags
        assert sp_pr.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == want
        assert sp_pr.find(qn("a:ln")).find(qn("a:noFill")) is not None
        assert s._element.find(qn("p:style")) is None


def test_icon_color_from_literal_style_color(tmp_path):
    out, _ = _render_icons(tmp_path, ["heart"], fill="#ff0000")
    shp = Presentation(str(out)).slides[0].shapes[0]
    assert shp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == "FF0000"


def test_curved_icons_use_beziers(tmp_path):
    out, _ = _render_icons(tmp_path, ["heart", "target"])
    for shp in Presentation(str(out)).slides[0].shapes:
        assert shp._element.spPr.xpath(".//a:cubicBezTo")


@needs_soffice
def test_all_icons_preview(tmp_path):
    out, _ = _render_icons(tmp_path, EXPECTED)
    assert pptx_to_pngs(out, tmp_path)


# --------------------------------------------------------------------------- layout


def lay(md: str, slide: int = 0):
    deck = parse(md)
    th = get_theme(deck.theme)
    placed = layout_slide(deck.slides[slide], deck, th, slide)
    return deck, th, placed


def icon_items(placed):
    return [p for p in placed if isinstance(p.element, Shape) and p.element.shape == "icon"]


def by_text(placed, text):
    return next(p for p in placed if isinstance(p.element, Text) and text in p.element.paragraphs[0].plain)


def card_of(placed, heading):
    return next(
        p
        for p in placed
        if isinstance(p.element, Container)
        and p.element.title is not None
        and heading in p.element.title.paragraphs[0].plain
    )


def inside(a, b, tol=2000):
    return (
        a.x >= b.x - tol
        and a.y >= b.y - tol
        and a.x + a.w <= b.x + b.w + tol
        and a.y + a.h <= b.y + b.h + tol
    )


BOXES = """# Boxes
@3
## Security {icon=shield}
- Encrypt
## Delivery {icon=truck}
- Ship
## Plain
- None
"""


def test_icon_left_of_heading():
    deck, th, placed = lay(BOXES)
    assert not [d for d in deck.diagnostics if d.level != "info"]
    ics = icon_items(placed)
    assert [p.element.attrs["icon"] for p in ics] == ["shield", "truck"]
    for ic, name in zip(ics, ("Security", "Delivery"), strict=True):
        head = by_text(placed, name)
        card = card_of(placed, name)
        assert ic.w == ic.h > 0
        assert ic.x + ic.w <= head.x  # heading text is shifted right of the icon
        assert abs((ic.y + ic.h / 2) - (head.y + ic.h / 2)) < ic.h  # same row as the heading
        assert inside(ic, card) and inside(head, card)
        assert ic.style.fill == head.style.color
        # about 1.2 x the heading font
        size = head.style.font_size * head.font_scale
        assert 1.0 * size * 12700 <= ic.h <= 1.4 * size * 12700
    plain = by_text(placed, "Plain")
    sec, sec_card = by_text(placed, "Security"), card_of(placed, "Security")
    assert (
        plain.x - card_of(placed, "Plain").x < sec.x - sec_card.x
    )  # a box without icon keeps its heading flush
    assert not lint(deck, [placed], th)


def test_icon_on_heading_band_is_clean():
    md = "theme: jp-business\n\n# T\n@2\n## セキュリティ {icon=shield}\n- 暗号化\n## 配送 {icon=truck}\n- 出荷\n"  # noqa: E501
    deck, th, placed = lay(md)
    ics = icon_items(placed)
    assert len(ics) == 2
    for ic in ics:
        head = by_text(placed, "セキュリティ" if ic.element.attrs["icon"] == "shield" else "配送")
        assert ic.x + ic.w <= head.x + 1
        assert ic.style.fill == "#FFFFFF"  # band text color
    assert not lint(deck, [placed], th)


KPI = """# KPI
@3
## 売上 {.kpi icon=chart}
**¥120M**
前年比 +12%
## 顧客 {.kpi icon=users}
**3,400**
## 目標 {.kpi}
**98%**
"""


def test_icon_above_kpi_number_centered():
    deck, th, placed = lay(KPI)
    assert not [d for d in deck.diagnostics if d.level != "info"]
    ics = icon_items(placed)
    assert len(ics) == 2
    for ic, label, num in zip(ics, ("売上", "顧客"), ("¥120M", "3,400"), strict=True):
        card = card_of(placed, label)
        number = by_text(placed, num)
        head = by_text(placed, label)
        assert inside(ic, card)
        assert ic.y + ic.h <= head.y <= number.y  # icon, label, number from top to bottom
        assert abs((ic.x + ic.w / 2) - (card.x + card.w / 2)) < 2000  # centered
        assert ic.style.fill == "primary"
        assert ic.w == pytest.approx(2 * head.style.font_size * 12700, rel=0.15)
    assert not lint(deck, [placed], th)


def test_kpi_icon_does_not_overflow_short_rows():
    deck, th, placed = lay(KPI)
    W, H = slide_size(deck.size)
    for p in placed:
        assert p.x >= 0 and p.y >= 0 and p.x + p.w <= W and p.y + p.h <= H


CHEV = """# Flow
@4 chevron
## Plan {icon=target}
## Build {icon=gear}
## Test
## Ship {icon=rocket}
"""


def test_chevron_icon_before_text():
    """The chevron carries the icon in its attributes; the renderer draws it from the final rectangle
    (a pass that moves the arrow, like `@steps` growth, never leaves the icon behind)."""
    deck, th, placed = lay(CHEV)
    assert icon_items(placed) == []
    chevrons = [p for p in placed if isinstance(p.element, Shape) and p.element.shape == "chevron"]
    assert len(chevrons) == 4
    with_icon = [c for c in chevrons if c.element.attrs.get("icon")]
    assert [c.element.attrs["icon"] for c in with_icon] == ["target", "gear", "rocket"]
    for ch in with_icon:
        assert ch.element.attrs["icon_inset"] > ch.element.attrs["icon_side"]  # text starts right of the icon
    assert not lint(deck, [placed], th)


def test_chevron_icon_reserves_left_inset_in_pptx(tmp_path):
    deck = parse(CHEV)
    th = get_theme(deck.theme)
    placed = [layout_slide(deck.slides[0], deck, th, 0)]
    out = tmp_path / "c.pptx"
    render(deck, placed, th, out)
    shapes = list(Presentation(str(out)).slides[0].shapes)
    icons_ = [s for s in shapes if s.name.startswith("icon ")]
    assert [s.name for s in icons_] == ["icon target", "icon gear", "icon rocket"]
    chev = [
        s
        for s in shapes
        if s._element.spPr.find(qn("a:prstGeom")) is not None
        and s.has_text_frame
        and s.text_frame.text in ("Plan", "Build", "Ship")
    ]
    for ic, ch in zip(icons_, chev, strict=True):  # each icon sits inside its chevron, left of the text
        assert ch.left < ic.left and ic.left + ic.width < ch.left + ch.width
        assert ch.top < ic.top and ic.top + ic.height < ch.top + ch.height
    plain = next(s for s in shapes if s.has_text_frame and s.text_frame.text == "Test")
    iconed = next(s for s in shapes if s.has_text_frame and s.text_frame.text == "Plan")
    assert iconed.text_frame.margin_left > plain.text_frame.margin_left


def test_unknown_icon_name_in_layout_is_ignored():
    md = "# T\n@2\n## A {icon=nope}\n- x\n## B\n- y\n"
    deck, th, placed = lay(md)
    assert icon_items(placed) == []
    _d, _t, plain = lay(md.replace(" {icon=nope}", ""))
    assert (by_text(placed, "A").x, by_text(placed, "A").w) == (by_text(plain, "A").x, by_text(plain, "A").w)


def test_icon_deck_builds_and_reopens(tmp_path):
    deck = parse(BOXES + "\n---\n\n" + KPI.replace("# KPI", "# K2"))
    th = get_theme(deck.theme)
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    out = tmp_path / "d.pptx"
    render(deck, placed, th, out)
    prs = Presentation(str(out))
    names = [s.name for sl in prs.slides for s in sl.shapes if s.name.startswith("icon ")]
    assert names == ["icon shield", "icon truck", "icon chart", "icon users"]


# --------------------------------------------------------------------------- elbow from a decision node

APPROVAL_SIDES = """# Approval
```mermaid
graph TD
A[Submit] --> B{Approved?}
B -->|yes| C[Finance pays]
B -->|no| D[Return]
B -->|escalate| E[Director]
```
"""


def test_elbow_from_diamond_is_glued_to_both_nodes(tmp_path):
    deck = parse(APPROVAL_SIDES)
    th = get_theme(deck.theme)
    placed = [layout_slide(deck.slides[0], deck, th, 0)]
    out = tmp_path / "a.pptx"
    render(deck, placed, th, out)
    sl = Presentation(str(out)).slides[0]
    ids = {s.shape_id: s for s in sl.shapes}
    cxns = [s for s in sl.shapes if s._element.tag == qn("p:cxnSp")]
    elbows = [c for c in cxns if c._element.spPr.find(qn("a:prstGeom")).get("prst") == "bentConnector3"]
    assert len(elbows) == 2  # yes / escalate; "no" is straight down
    for c in elbows:
        nv = c._element.find(qn("p:nvCxnSpPr")).find(qn("p:cNvCxnSpPr"))
        st, en = nv.find(qn("a:stCxn")), nv.find(qn("a:endCxn"))
        assert st is not None and en is not None
        assert ids[int(st.get("id"))].auto_shape_type is not None
        assert (st.get("idx"), en.get("idx")) == ("2", "0")  # bottom-middle -> top-middle


def test_parser_icon_names_in_sync():
    from slidemark.icons import names
    from slidemark.parser.attrs import ICON_NAMES

    assert set(names()) == set(ICON_NAMES)
