"""Compact header rows on stretched tables and multi-row merged headers."""

from __future__ import annotations

from pptx import Presentation
from pptx.enum.text import PP_ALIGN

from slidemark.layout import layout_slide
from slidemark.parser import parse
from slidemark.render import render
from slidemark.theme import get_theme

ROADMAP = """theme: jp-business
lang: ja

# ロードマップ
> 3か年の施策

{header=2}
| 施策 | 2027年度 | < | 2028年度 | < | 2029年度 |
|-|-|-|-|-|-|
| ^ | 上期 | 下期 | 上期 | 下期 | 通期 |
| IoT標準化 | 新機種へ搭載 | 全機種へ | 既設機の後付け | < | 接続率95% |
| データ基盤 | 要件定義 | 構築 | 予兆モデル運用 | < | 全機種展開 |
"""

STRETCHED = """theme: jp-business

# 体制
| 施策 | 2027 | 2028 |
|-|-|-|
| IoT | 新機種 | 全機種 |
| データ | 要件 | 構築 |
| 稼働 | 試行 | 販売 |
| 代理店 | 方針 | 契約 |
"""


def build(md, tmp_path):
    deck = parse(md)
    th = get_theme(deck.theme)
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    out = tmp_path / "o.pptx"
    render(deck, placed, th, out)
    prs = Presentation(str(out))
    return deck, next(s for s in prs.slides[0].shapes if s.has_table).table


def test_header_row_compact_on_stretched_table(tmp_path):
    _, t = build(STRETCHED, tmp_path)
    hs = [r.height for r in t.rows]
    assert hs[0] < min(hs[1:]) * 0.75
    assert hs[0] < 0.75 * 914400


def test_merged_header_xml_and_style(tmp_path):
    _, t = build(ROADMAP, tmp_path)

    def tc(r, c):
        return t.cell(r, c)._tc

    assert tc(0, 1).get("gridSpan") == "2" and tc(0, 2).get("hMerge") == "1"
    assert tc(0, 3).get("gridSpan") == "2" and tc(0, 4).get("hMerge") == "1"
    assert tc(0, 0).get("rowSpan") == "2" and tc(1, 0).get("vMerge") == "1"
    assert tc(2, 3).get("gridSpan") == "2"
    fills = {str(t.cell(r, c).fill.fore_color.rgb) for r in (0, 1) for c in range(6)}
    assert len(fills) == 1  # both header rows share the header fill
    assert str(t.cell(2, 1).fill.fore_color.rgb) not in fills
    assert t.cell(0, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert t.cell(1, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert t.cell(2, 3).text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER  # spanning body bar
    hs = [r.height for r in t.rows]
    assert max(hs[:2]) < min(hs[2:])


def test_odd_header_and_merge_combos_do_not_crash(tmp_path):
    for n in (0, 1, 2, 3, 9):
        _, t = build(ROADMAP.replace("{header=2}", f"{{header={n}}}"), tmp_path)
        assert len(t.rows) >= 1
    _, t = build("# T\n{header=2}\n| ^ | < |\n|-|-|\n| ^ | < |\n", tmp_path)
    assert len(t.rows) >= 1
