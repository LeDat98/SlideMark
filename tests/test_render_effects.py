"""Native fills and effects: CSS gradients -> a:gradFill, box-shadows -> a:outerShdw, #RRGGBBAA, opacity."""

from __future__ import annotations

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from slidemark.ir import Container, Deck, Paragraph, Run, Slide, Style, Text
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.render.util import hex6, parse_color
from slidemark.theme import Theme, get_theme

ORDER = ["xfrm", "custGeom", "prstGeom", "noFill", "solidFill", "gradFill", "ln", "effectLst", "scene3d"]


def T(s, role="body"):
    return Text(role=role, paragraphs=[Paragraph(runs=[Run(text=s)])])


def card_sppr(style: Style, tmp_path, theme="default"):
    s = Slide(
        title=T("t", "title"),
        elements=[Container(title=T("h", "heading"), children=[T("body")], style=style)],
    )
    deck = Deck(slides=[s])
    th = get_theme(theme)
    placed = [layout_slide(s, deck, th, 0)]
    out = tmp_path / "o.pptx"
    render(deck, placed, th, out)
    prs = Presentation(str(out))
    card = next(x for x in prs.slides[0].shapes if x.name.startswith("Card"))
    return card._element.spPr, deck


def children(sppr):
    return [etree.QName(c).localname for c in sppr]


def assert_order(sppr):
    idx = [ORDER.index(t) for t in children(sppr) if t in ORDER]
    assert idx == sorted(idx), children(sppr)


def test_linear_gradient(tmp_path):
    sppr, _ = card_sppr(Style(fill="linear-gradient(90deg, #FF0000 10%, primary, #0000FF)"), tmp_path)
    g = sppr.find(qn("a:gradFill"))
    assert g is not None and sppr.find(qn("a:solidFill")) is None
    gs = g.findall(qn("a:gsLst") + "/" + qn("a:gs"))
    assert [x.get("pos") for x in gs] == ["10000", "55000", "100000"]  # the middle stop is spread evenly
    assert gs[0][0].get("val") == "FF0000" and gs[2][0].get("val") == "0000FF"
    assert gs[1][0].get("val") == hex6(get_theme("default"), "primary")
    assert g.find(qn("a:lin")).get("ang") == "0"  # CSS 90deg = to right = DrawingML 0
    assert_order(sppr)


@pytest.mark.parametrize(
    "css,ang",
    [("0deg", 270), ("180deg", 90), ("135deg", 45), ("to right", 0), ("to bottom", 90), ("270deg", 180)],
)
def test_gradient_angle_conversion(tmp_path, css, ang):
    sppr, _ = card_sppr(Style(fill=f"linear-gradient({css}, #FFF, #000)"), tmp_path)
    assert sppr.find(qn("a:gradFill")).find(qn("a:lin")).get("ang") == str(ang * 60000)


def test_radial_gradient_with_alpha(tmp_path):
    sppr, _ = card_sppr(Style(fill="radial-gradient(circle, #FF000080, #0000FF)"), tmp_path)
    g = sppr.find(qn("a:gradFill"))
    assert g.find(qn("a:path")).get("path") == "circle"
    stops = g.findall(qn("a:gsLst") + "/" + qn("a:gs"))
    alpha = stops[0][0].find(qn("a:alpha"))
    assert alpha is not None and abs(int(alpha.get("val")) - 50196) <= 1
    assert stops[1][0].find(qn("a:alpha")) is None
    assert_order(sppr)


def test_opacity_applies_to_gradient_and_solid(tmp_path):
    sppr, _ = card_sppr(Style(fill="linear-gradient(#FFF, #000)", opacity=0.5), tmp_path)
    for gs in sppr.iter(qn("a:gs")):
        assert gs[0].find(qn("a:alpha")).get("val") == "50000"
    sppr, _ = card_sppr(Style(fill="#112233", opacity=0.25), tmp_path)
    assert sppr.find(qn("a:solidFill"))[0].find(qn("a:alpha")).get("val") == "25000"


def test_solid_fill_with_alpha_channel(tmp_path):
    sppr, _ = card_sppr(Style(fill="#11223380"), tmp_path)
    clr = sppr.find(qn("a:solidFill"))[0]
    assert clr.get("val") == "112233"
    assert abs(int(clr.find(qn("a:alpha")).get("val")) - 50196) <= 1


def test_shadow_string(tmp_path):
    sppr, _ = card_sppr(Style(fill="surface", line="border", shadow="0 8 24 #00000055"), tmp_path)
    sh = sppr.find(qn("a:effectLst")).find(qn("a:outerShdw"))
    assert sh.get("blurRad") == str(24 * 12700)
    assert sh.get("dist") == str(8 * 12700)
    assert sh.get("dir") == str(90 * 60000)  # straight down
    assert sh[0].get("val") == "000000"
    assert abs(int(sh[0].find(qn("a:alpha")).get("val")) - 33333) <= 1
    assert_order(sppr)
    assert children(sppr).index("ln") < children(sppr).index("effectLst")


def test_shadow_direction_and_spread(tmp_path):
    sppr, _ = card_sppr(Style(fill="surface", shadow="4 4 0 2 red"), tmp_path)
    sh = sppr.find(qn("a:effectLst")).find(qn("a:outerShdw"))
    assert sh.get("dir") == str(45 * 60000)
    assert int(sh.get("sx")) > 100000 and int(sh.get("sy")) > 100000
    assert sh[0].get("val") == "FF0000"


def test_shadow_true_uses_render_token(tmp_path):
    sppr, _ = card_sppr(Style(fill="surface", shadow=True), tmp_path)
    sh = sppr.find(qn("a:effectLst")).find(qn("a:outerShdw"))
    assert sh is not None  # "0 2 6 #00000040": 2pt down, 6pt blur
    assert sh.get("blurRad") == str(6 * 12700) and sh.get("dist") == str(2 * 12700)


def test_no_shadow_gives_empty_effect_list(tmp_path):
    sppr, _ = card_sppr(Style(fill="surface"), tmp_path)
    eff = sppr.find(qn("a:effectLst"))
    assert eff is not None and len(eff) == 0


def test_bad_gradient_falls_back_to_first_color_with_a_warning(tmp_path):
    sppr, deck = card_sppr(Style(fill="linear-gradient(banana, nope)"), tmp_path)
    assert sppr.find(qn("a:gradFill")) is None and sppr.find(qn("a:solidFill")) is not None
    d = [x for x in deck.diagnostics if x.rule == "bad-gradient"]
    assert d and d[0].hint


def test_bad_shadow_is_dropped_with_a_warning(tmp_path):
    sppr, deck = card_sppr(Style(fill="surface", shadow="wobbly"), tmp_path)
    assert len(sppr.find(qn("a:effectLst"))) == 0
    assert any(x.rule == "bad-shadow" and x.hint for x in deck.diagnostics)


def test_inset_shadow_warns(tmp_path):
    _, deck = card_sppr(Style(fill="surface", shadow="inset 0 2 4 #000"), tmp_path)
    assert any(x.rule == "bad-shadow" for x in deck.diagnostics)


@pytest.mark.parametrize(
    "value,expect",
    [
        ("#112233", ("112233", None)),
        ("#123", ("112233", None)),
        ("#11223380", ("112233", 0.5020)),
        ("#1238", ("112233", 0.5333)),
        ("rgb(255, 0, 0)", ("FF0000", None)),
        ("rgba(0, 0, 255, 0.5)", ("0000FF", 0.5)),
        ("linear-gradient(90deg, #AABBCC 0%, #000)", ("AABBCC", None)),
        ("transparent", ("FFFFFF", 0.0)),
    ],
)
def test_parse_color(value, expect):
    hx, al = parse_color(Theme(name="none"), value)
    assert hx == expect[0]
    assert (al is None) == (expect[1] is None) and (al is None or abs(al - expect[1]) < 1e-3)


@pytest.mark.parametrize(
    "value", ["", "not-a-color", "#12", "#GGGGGG", "rgba(", None, "linear-gradient(", "####"]
)
def test_color_resolution_never_raises(value):
    assert hex6(Theme(name="none"), value) == "000000"


# --------------------------------------------------------------------------- theme: none + tokens


def _none_deck(tokens):
    from slidemark.parser import parse

    deck = parse("theme: none\n\n# Title\n@2 a>b\n## A\nleft\n## B\nright\n")
    deck.tokens = tokens
    return deck


def _build_none(tmp_path, tokens):
    from slidemark.build import build_deck

    deck = _none_deck(tokens)
    out = tmp_path / "none.pptx"
    build_deck(deck, out)
    return Presentation(str(out)), deck


TOKENS = {
    "colors.primary": "#7C5CFF",
    "colors.bg": "#0B1020",
    "classes.card.fill": "linear-gradient(135deg, #141A2E, #1F2A4D)",
    "classes.card.shadow": "0 8 24 #00000055",
    "layout.top_gap": "0.25in",
    "render.connector_width": "4",
}


def test_theme_none_tokens_drive_colors_gradient_shadow_and_connector(tmp_path):
    prs, _ = _build_none(tmp_path, TOKENS)
    slide = prs.slides[0]
    cards = [s for s in slide.shapes if s.name.startswith("Card")]
    assert len(cards) == 2
    for c in cards:
        g = c._element.spPr.find(qn("a:gradFill"))
        assert g is not None
        assert [gs[0].get("val") for gs in g.iter(qn("a:gs"))] == ["141A2E", "1F2A4D"]
        assert c._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw")) is not None
    bg = slide.background.fill
    assert str(bg.fore_color.rgb) == "0B1020"
    lines = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]
    assert lines and all(abs(ln.line.width.pt - 4) < 1e-6 for ln in lines)
    assert str(lines[0].line.color.rgb) == "7C5CFF"  # connectors use the primary color


def test_theme_none_top_gap_token_moves_the_body(tmp_path):
    def body_top(gap):
        prs, _ = _build_none(tmp_path, {**TOKENS, "layout.top_gap": gap})
        return min(s.top for s in prs.slides[0].shapes if s.name.startswith("Card"))

    # a sparse slide moves its leftover down, so the shift is not exact: a bigger gap pushes the body down
    assert 0 < body_top("1in") - body_top("0.25in") <= 685800


def test_theme_none_has_no_preset_look(tmp_path):
    """Without tokens the deck is the neutral schema defaults: white background."""
    prs, deck = _build_none(tmp_path, {})
    assert str(prs.slides[0].background.fill.fore_color.rgb) == "FFFFFF"
    assert deck.theme == "none"
