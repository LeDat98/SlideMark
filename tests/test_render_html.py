"""Raw HTML blocks -> picture through headless Chromium (optional extra); placeholder without it."""

from __future__ import annotations

import http.server
import io
import sys
import threading

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidemark.build import build
from slidemark.render.htmlimg import ALT, HtmlRenderer
from slidemark.theme import Theme

SVG = '<svg width="200" height="100"><circle cx="50" cy="50" r="40" fill="red"/></svg>'


def _deck(body: str) -> str:
    return f"# Chart\n\n```html{{render=image}}\n{body}\n```\n"


def _pictures(path):
    prs = Presentation(str(path))
    return [s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


@pytest.fixture(scope="module")
def chromium():
    pytest.importorskip("playwright.sync_api")
    r = HtmlRenderer()
    try:
        png = r.render("<p>x</p>", 40, 20, Theme(name="probe"))
    finally:
        r.close()
    if png is None:
        pytest.skip("Chromium cannot be launched")


def test_html_block_becomes_picture(tmp_path, chromium):
    from PIL import Image

    out = tmp_path / "a.pptx"
    deck = build(_deck(SVG), out)
    pics = _pictures(out)
    assert len(pics) == 1
    pic = pics[0]
    assert pic.name == "html 1"
    assert pic._element.nvPicPr.cNvPr.get("descr") == ALT
    assert pic.width > 0 and pic.height > 0
    assert not [d for d in deck.diagnostics if d.rule == "html-image-unavailable"]
    img = Image.open(io.BytesIO(pic.image.blob)).convert("RGB")
    assert abs(img.width / 2 - pic.width / 9525) <= 1 and abs(img.height / 2 - pic.height / 9525) <= 1
    assert any(r > 200 and g < 60 and b < 60 for r, g, b in img.get_flattened_data())


def test_external_request_is_blocked(tmp_path, chromium):
    hits: list[str] = []

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            hits.append(self.path)
            self.send_response(404)
            self.end_headers()

        def log_message(self, *a):
            pass

    srv = http.server.HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    try:
        url = f"http://127.0.0.1:{srv.server_port}/x.png"
        out = tmp_path / "b.pptx"
        build(_deck(f'<svg width="50" height="50"><image href="{url}"/></svg><img src="{url}">'), out)
        assert len(_pictures(out)) == 1
    finally:
        srv.shutdown()
        srv.server_close()
    assert hits == []


def test_browser_is_reused_for_many_blocks(tmp_path, chromium):
    md = "# A\n\n" + "".join(f"```html{{render=image}}\n{SVG}\n```\n\n" for _ in range(3))
    out = tmp_path / "c.pptx"
    build(md, out)
    prs = Presentation(str(out))
    names = [s.name for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert names == ["html 1", "html 2", "html 3"]


def test_without_playwright_keeps_placeholder(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "playwright", None)
    monkeypatch.setitem(sys.modules, "playwright.sync_api", None)
    out = tmp_path / "d.pptx"
    deck = build(_deck(SVG), out)
    assert _pictures(out) == []
    unavailable = [d for d in deck.diagnostics if d.rule == "html-image-unavailable"]
    assert len(unavailable) == 1
    assert unavailable[0].level == "info"
    assert "slidemark[html]" in unavailable[0].hint
    texts = [s.text_frame.text for s in Presentation(str(out)).slides[0].shapes if s.has_text_frame]
    assert "[html]" in texts
