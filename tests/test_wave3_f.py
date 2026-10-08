"""Wave 2026-10-08 lane F: one step label, card headings that never wrap in a row, palette `@proscons`."""

from __future__ import annotations

import pytest
from pptx import Presentation

from slidemark import build, parse

HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#14272F primary=#0B2A3C secondary=#1B8A8F accent=#F2A33A surface=#EEF4F6 "
    "muted=#5A6F78\n"
    "fonts: heading=Arial body=Arial\n"
    "sizes: title=34 heading=20 body=16 caption=12\n"
)
STEPS = "# Steps\n@3 steps num\n## Một\n- a\n## Hai\n- b\n## Ba\n- c\n"


def _texts(path) -> list[str]:
    prs = Presentation(str(path))
    return [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame]


def _build(tmp_path, src: str, name: str = "d.pptx"):
    out = tmp_path / name
    deck = build(HEAD + src, out)
    return deck, out


# --------------------------------------------------------------------------- 1. one step label


def test_steps_num_draws_no_step_caption_on_the_card_look(tmp_path):
    _deck, out = _build(tmp_path, STEPS)
    assert not any("STEP" in t for t in _texts(out))
    names = [s.name for s in Presentation(str(out)).slides[0].shapes]
    assert "Step 1 arrow" in names and "Heading 1" in names


def test_steps_caption_token_states_the_caption_on_the_card_look(tmp_path):
    _deck, out = _build(tmp_path, 'style: steps.caption="Bước {n}"\n' + STEPS)
    assert any(t.startswith("Bước 1") for t in _texts(out))


def test_steps_num_with_head_arrow_draws_the_caption(tmp_path):
    _deck, out = _build(tmp_path, STEPS.replace("steps num", "steps head=arrow num"))
    assert any(t.startswith("STEP 1") for t in _texts(out))
    _deck, out2 = _build(tmp_path, "style: steps.head=arrow\n" + STEPS, "e.pptx")
    assert any(t.startswith("STEP 2") for t in _texts(out2))


def test_steps_num_is_fuzz_safe():
    for src in ("# T\n@steps num\n", "# T\n@2 steps num\n## a\n", "# T\n@3 steps num head=card\n## a\n## b\n"):
        parse(src)
