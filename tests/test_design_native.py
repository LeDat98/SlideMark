"""DL2 part 1: decisions of the python-pptx deck that needed a css fence or `@html` have a native form.

Forms: slide-scoped `sizes:` / `style:` lines, the text shorthand (`kpi.label="20 bold primary"`),
`fonts: font=`,
the `@kpi` slide flag, `kpi.h=4.4in`, the table option `hlcol=`. Each: parser, render (the .pptx is reopened),
importer where it can see it, and a bad value = a diagnostic with a hint (never an exception)."""

from __future__ import annotations

from pptx import Presentation

from slidemark import build, parse
from slidemark.importer import import_pptx
from slidemark.template import deck_theme

HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#222B36 primary=#142B4D secondary=#1F5FA8 accent=#E09F1F teal=#2A9D8F "
    "muted=#59626E surface=#EEF2F7 border=#C9D2DE\n"
)


def make(tmp_path, text: str, name: str = "d"):
    deck = build(HEAD + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def warns(deck, rule):
    return [d for d in deck.diagnostics if d.rule == rule]


def _nolines(x):
    if isinstance(x, dict):
        return {k: _nolines(v) for k, v in x.items() if k != "line"}
    return [_nolines(v) for v in x] if isinstance(x, list) else x


def run_sizes(prs, i, prefix):
    return [
        r.font.size.pt
        for s in prs.slides[i].shapes
        if s.name.startswith(prefix) and s.has_text_frame
        for p in s.text_frame.paragraphs
        for r in p.runs
        if r.font.size
    ]


# --------------------------------------------------------------------------- slide-scoped tokens

BOXES = "## A\n- x\n- y\n## B\n- z\n## C\n- w\n"


def test_slide_tokens_are_read_and_blanked():
    d = parse(
        HEAD
        + "sizes: body=22\n\n# One\nsizes: heading=26! body=24!\nstyle: kpi.size=54\n"
        + BOXES
        + "\n# Two\n- a\n"
    )
    assert d.slides[0].tokens == {"sizes.heading": "26!", "sizes.body": "24!", "classes.kpi.font_size": "54"}
    assert d.slides[1].tokens == {}
    assert d.tokens["sizes.body"] == "22"  # the header is untouched
    assert not [x for x in d.diagnostics if x.rule in ("bad-token", "unknown-token")]
    texts = [
        p.plain
        for e in d.slides[0].elements
        for c in getattr(e, "children", [])
        for p in getattr(c, "paragraphs", [])
    ]
    assert not any("sizes:" in t or "style:" in t for t in texts)


def test_slide_tokens_change_that_slide_only(tmp_path):
    md = "sizes: heading=20! body=22!\n\n# One\n" + BOXES + "\n# Two\nsizes: heading=26! body=24!\n" + BOXES
    deck, prs = make(tmp_path, md)
    assert not warns(deck, "bad-token")
    assert set(run_sizes(prs, 0, "Heading")) == {20.0} and set(run_sizes(prs, 0, "Text")) == {22.0}
    assert set(run_sizes(prs, 1, "Heading")) == {26.0} and set(run_sizes(prs, 1, "Text")) == {24.0}


def test_slide_style_line_carries_element_css_and_counts_as_a_choice(tmp_path):
    md = "\n# One\n@cover bg=primary dark\nstyle: subtitle.color=#D6E2F0\nA subtitle\n\n# Two\n- a\n"
    deck, prs = make(tmp_path, md)
    sub = next(s for s in prs.slides[0].shapes if s.name == "Subtitle")
    assert str(sub.text_frame.paragraphs[0].runs[0].font.color.rgb) == "D6E2F0"
    assert deck.slides[0].css and deck.attrs["design_stated"]


def test_slide_token_line_needs_key_value_words_so_prose_stays_text():
    d = parse(HEAD + "\n# T\nstyle: modern and clean\n- a\n")
    assert d.slides[0].tokens == {}
    flat = " ".join(p.plain for e in d.slides[0].elements for p in getattr(e, "paragraphs", []))
    assert "style: modern and clean" in flat


def test_slide_token_bad_values_warn_with_a_hint_on_the_slide():
    d = parse(HEAD + '\n# T\nsizes: body=huge heading=26!\nstyle: nonsense.key=1 kpi.label="20 zzz!"\n- a\n')
    bad = [x for x in d.diagnostics if x.rule in ("bad-token", "unknown-token")]
    assert len(bad) >= 2 and all(x.hint and x.slide == 1 for x in bad)
    assert d.slides[0].tokens.get("sizes.heading") == "26!"  # the good pair survives


# --------------------------------------------------------------------------- text shorthand


def test_text_shorthand_expands_to_the_long_tokens():
    d = parse(
        HEAD
        + 'style: kpi.label="20 bold primary" kpi.note="20 bold teal" steps-card="24 bold center" '
        + 'heading=center conclusion="24 bold" lead="bold secondary" title=primary\n\n# T\n- a\n'
    )
    by = {r.selector: r.style for r in d.css}
    assert (by[".kpi h2"].font_size, by[".kpi h2"].bold, by[".kpi h2"].color) == (20, True, "primary")
    assert (by[".kpi .caption"].font_size, by[".kpi .caption"].color) == (20, "teal")
    assert by["h2"].align == "center" and by[".conclusion"].font_size == 24 and by[".lead"].bold is True
    assert (
        d.tokens["classes.steps-card.font_size"] == "24" and d.tokens["classes.steps-card.align"] == "center"
    )
    assert d.tokens["lead_color"] == "secondary" and d.tokens["title_color"] == "primary"
    assert not [x for x in d.diagnostics if x.rule in ("bad-token", "unknown-token")]


def test_text_shorthand_leaves_other_tokens_alone():
    d = parse(HEAD + "style: radius=14 card.fill=#EEEEEE raduis=14 bullet=■\n\n# T\n- a\n")
    assert d.tokens["classes.card.radius"] == "14" and d.tokens["classes.card.fill"] == "#EEEEEE"
    typo = [x for x in d.diagnostics if x.rule == "unknown-token"]
    assert len(typo) == 1 and "raduis" in typo[0].message and typo[0].hint


def test_text_shorthand_bad_word_is_a_diagnostic_with_a_hint():
    d = parse(HEAD + 'style: kpi.label="20 bold 3in!" heading=""\n\n# T\n- a\n')
    bad = [x for x in d.diagnostics if x.rule == "bad-token"]
    assert bad and all(x.hint for x in bad)


def test_text_shorthand_renders(tmp_path):
    _, prs = make(
        tmp_path,
        'style: kpi.label="20 bold primary" kpi.note="18 bold teal"\n\n'
        "# T\n@kpi\n## Sales\n1,280\nvs last year\n",
    )
    label = next(s for s in prs.slides[0].shapes if s.name == "Heading 1").text_frame.paragraphs[0].runs[0]
    assert label.font.size.pt == 20 and label.font.bold and str(label.font.color.rgb) == "142B4D"


# --------------------------------------------------------------------------- fonts: font=


def test_fonts_font_sets_the_text_roles():
    d = parse('fonts: font="Yu Gothic" mono=Menlo\n\n# T\n- a\n')
    th, _ = deck_theme(d, ".")
    assert (th.fonts.heading, th.fonts.body, th.fonts.ea, th.fonts.mono) == (
        "Yu Gothic",
        "Yu Gothic",
        "Yu Gothic",
        "Menlo",
    )
    d2 = parse("fonts: font=Inter heading=Georgia\n\n# T\n- a\n")
    assert deck_theme(d2, ".")[0].fonts.heading == "Georgia"  # a later role wins


# --------------------------------------------------------------------------- @kpi


KPIS = "## A\n1\nn\n## B\n2\nn\n## C\n3\nn\n"


def test_at_kpi_makes_every_box_a_kpi_card():
    d = parse(HEAD + "\n# T\n@kpi\n" + KPIS)
    boxes = d.slides[0].elements
    assert len(boxes) == 3 and all("kpi" in b.classes for b in boxes)
    assert "kpi" not in d.slides[0].classes
    long = parse(
        HEAD
        + "\n# T\n"
        + KPIS.replace("## A", "## A {.kpi}").replace("## B", "## B {.kpi}").replace("## C", "## C {.kpi}")
    )
    assert _nolines([c.model_dump() for c in boxes]) == _nolines(
        [c.model_dump() for c in long.slides[0].elements]
    )


def test_at_kpi_renders_like_the_class_and_round_trips(tmp_path):
    _, prs = make(tmp_path, "style: kpi.stripe=secondary\n\n# T\n@kpi\n" + KPIS)
    assert len([s for s in prs.slides[0].shapes if s.name == "rule"]) == 3  # a stripe on every card
    text, _ = import_pptx(tmp_path / "d.pptx")
    assert "{.kpi}" in text or "@kpi" in text


# --------------------------------------------------------------------------- kpi.h


def test_kpi_h_is_the_card_height(tmp_path):
    _, prs = make(tmp_path, "style: kpi.h=4.4in\n\n# T\n> lead\n@kpi\n" + KPIS)
    cards = [s for s in prs.slides[0].shapes if s.name.startswith("Card")]
    assert len(cards) == 3 and all(abs(c.height - 4.4 * 914400) < 2000 for c in cards)
    _, other = make(tmp_path, "style: kpi.h=3.5in\n\n# T\n> lead\n@kpi\n" + KPIS, "two")
    assert all(
        abs(c.height - 3.5 * 914400) < 2000 for c in other.slides[0].shapes if c.name.startswith("Card")
    )


def test_kpi_h_bad_value_warns(tmp_path):
    deck, _ = make(tmp_path, "style: kpi.h=tall\n\n# T\n@kpi\n" + KPIS)
    d = warns(deck, "bad-token")
    assert d and d[0].hint


# --------------------------------------------------------------------------- hlcol

TABLE = (
    "| 事業 | 2026 | 2027計画 | 前年比 |\n|-|-|-|-|\n"
    "| 冷凍 | 420 | 465 | +11% |\n| 海外 | 200 | 220 | +10% |\n"
)


def _bold(prs, r):
    tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
    return [bool(c.text_frame.paragraphs[0].runs[0].font.bold) for c in tbl.rows[r].cells]


def test_hlcol_names_a_header_cell_or_a_number():
    for opt in ("hlcol=2027計画", "hlcol=3"):
        d = parse(HEAD + f"\n# T\n{{{opt}}}\n" + TABLE)
        t = d.slides[0].elements[0]
        assert t.attrs["hlcol"] == [3] and not [x for x in d.diagnostics if x.rule == "table-hl"]
    d = parse(HEAD + "\n# T\n{hlcol=2026,前年比}\n" + TABLE)
    assert d.slides[0].elements[0].attrs["hlcol"] == [2, 4]


def test_hlcol_renders_bold_tinted_body_cells_only(tmp_path):
    _, prs = make(tmp_path, "\n# T\n{hlcol=2027計画}\n" + TABLE)
    assert _bold(prs, 1) == [False, False, True, False] and _bold(prs, 2) == [False, False, True, False]
    tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
    fills = [str(tbl.cell(1, c).fill.fore_color.rgb) for c in range(4)]
    assert fills[2] != fills[1] and fills[0] == fills[1] == fills[3]  # only that column is tinted
    assert str(tbl.cell(0, 2).fill.fore_color.rgb) == str(
        tbl.cell(0, 1).fill.fore_color.rgb
    )  # header keeps its own


def test_hlcol_with_hl_rows_and_the_look_tokens(tmp_path):
    _, prs = make(tmp_path, "style: table.hl.fill=teal\n\n# T\n{hlcol=2 hl=海外}\n" + TABLE)
    tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
    tint = str(tbl.cell(2, 2).fill.fore_color.rgb)
    assert str(tbl.cell(2, 0).fill.fore_color.rgb) == tint  # the hl row
    assert str(tbl.cell(1, 1).fill.fore_color.rgb) == tint  # the hlcol column
    assert str(tbl.cell(1, 3).fill.fore_color.rgb) != tint


def test_hlcol_bad_values_warn_with_a_hint(tmp_path):
    deck, _ = make(tmp_path, "\n# T\n{hlcol=nope,9}\n" + TABLE)
    w = warns(deck, "table-hl")
    assert len(w) == 2 and all(x.hint for x in w) and "columns are 1..4" in w[1].hint
    d = parse(HEAD + "\n# T\n{hlcol=2027計}\n" + TABLE)
    assert "did you mean '2027計画'" in [x for x in d.diagnostics if x.rule == "table-hl"][0].hint


def test_hlcol_round_trips_through_the_importer(tmp_path):
    make(tmp_path, "\n# T\n{hlcol=2027計画}\n" + TABLE, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "hlcol=3" in text and "**465**" not in text  # the bold of an emphasised column is not content
    redo = tmp_path / "again.pptx"
    build(HEAD + text, redo)
    assert _bold(Presentation(str(redo)), 1) == [False, False, True, False]
