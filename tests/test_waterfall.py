"""Waterfall chart: parser markers, native stacked render, importer round trip, fuzz safety."""

from __future__ import annotations

import random

from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import build, parse
from slidemark.importer import import_pptx
from slidemark.render import waterfall as wf
from slidemark.theme import RenderTokens

MD = """# T
```waterfall {title="OP" labels=on}
,A,b,c,d,e
x,44,18,-33,12,=
```
"""


def _chart(path):
    prs = Presentation(str(path))
    return [s for s in prs.slides[0].shapes if s.has_chart][0].chart


def test_parse_markers_and_values():
    deck = parse(MD)
    ch = deck.slides[0].elements[0]
    assert ch.kind == "waterfall"
    assert ch.series[0].values == [44, 18, -33, 12, None]
    assert ch.options["totals"] == [4]
    assert not [d for d in deck.diagnostics if d.rule == "bad-number"]


def test_parse_bad_cell_and_extra_series():
    deck = parse("# T\n```waterfall\n,a,b,c\nx,1,zz,=\ny,2,3,4\n```\n")
    rules = [d.rule for d in deck.diagnostics]
    assert "bad-number" in rules and "waterfall-series" in rules
    assert len(deck.slides[0].elements[0].series) == 1
    assert all(d.hint for d in deck.diagnostics if d.rule == "waterfall-series")


def test_render_native_stacked_columns(tmp_path):
    (tmp_path / "a.md").write_text(MD, encoding="utf-8")
    build(tmp_path / "a.md", tmp_path / "a.pptx")
    chart = _chart(tmp_path / "a.pptx")
    sers = list(chart.plots[0].series)
    assert len(sers) == 6
    assert [s.name for s in sers][:4] == ["base (hidden)", "up", "down", "total"]
    assert sers[0]._element.find(qn("c:spPr")).find(qn("a:noFill")) is not None
    assert list(sers[0].values) == [0, 44, 29, 29, 0]
    assert list(sers[1].values)[1] == 18 and list(sers[2].values)[2] == 33
    assert list(sers[3].values)[0] == 44 and list(sers[3].values)[4] == 41
    fills = [s._element.find(qn("c:spPr")).find(qn("a:solidFill")) for s in sers[1:4]]
    assert len({f.find(qn("a:srgbClr")).get("val") for f in fills}) == 3
    assert not chart.has_legend
    assert chart.value_axis.minimum_scale == 0


def test_negative_running_total_splits_bar():
    bs = wf.bars([10, -25, 5, 30], [])
    cols = wf.columns(bs, 1)
    assert cols[wf.CROSS][1] == -15 and cols[wf.DOWN][1] == 10  # crosses zero: 10 above, 15 below
    assert wf.decode(cols) == ([10, -25, 5, 30], [])
    ax = wf.axis(bs, RenderTokens())
    assert ax and ax[0] < 0 < ax[1]


def test_label_text():
    assert wf.fmt_num(18, None, True) == "+18"
    assert wf.fmt_num(-33, None, True) == "-33"
    assert wf.fmt_num(1240, "#,##0", False) == "1,240"
    assert wf.fmt_num(1.5, '0.0"億"', True) == "+1.5億"


def test_round_trip(tmp_path):
    md = MD + "\n---\n# N\n```waterfall\n,s,a,b,c\nx,10,-25,5,=\n```\n"
    (tmp_path / "a.md").write_text(md, "utf-8")
    build(tmp_path / "a.md", tmp_path / "a.pptx")
    text, _ = import_pptx(tmp_path / "a.pptx")
    assert "```waterfall" in text and "x,44,18,-33,12,=" in text and "x,10,-25,5,=" in text
    assert "max=" not in text and "min=" not in text


def test_fuzz_never_crashes(tmp_path):
    rnd = random.Random(7)
    cells = ["1", "-3", "=", "", "abc", "1,5", "▲4", "0", "12%"]
    for _ in range(25):
        n = rnd.randint(0, 6)
        rows = [",".join(rnd.choice(cells) for _ in range(n)) for _ in range(rnd.randint(0, 3))]
        (tmp_path / "f.md").write_text(
            "# T\n```waterfall {labels=on}\n" + "\n".join(rows) + "\n```\n", "utf-8"
        )
        build(tmp_path / "f.md", tmp_path / "f.pptx")
