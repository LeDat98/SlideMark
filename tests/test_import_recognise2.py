"""DL3d lane C: geometry-based recognition of a foreign deck's design (``importer/recognise2.py``).

Each test draws loose python-pptx shapes the way a deck made elsewhere does, imports the file, asserts the
imported deck.md and (where the token has a look) rebuilds it and reopens the .pptx.
"""

from __future__ import annotations

import re
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.importer.recognise2 import T, has_page_numbers, is_page_number
from slidemark.parser import parse

NAVY, BLUE, TEAL, AMBER, LIGHT, GREY = "122B4A", "1F5FA8", "1B8A8F", "E08A1E", "EEF2F7", "6B7582"


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, x, y, w, h, rgb, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(rgb)
    s.line.fill.background()
    return s


def _text(slide, x, y, w, h, paras, size=16, bold=False, color="222B36", shape=None):
    """``paras``: a str, or a list of paragraphs, each a str or a list of (text, size, bold, color) runs."""
    tb = shape or slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, para in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        for t, sz, b, c in para if isinstance(para, list) else [(para, size, bold, color)]:
            r = p.add_run()
            r.text = t
            r.font.size = Pt(sz)
            r.font.bold = b
            r.font.color.rgb = RGBColor.from_string(c)
    return tb


def _chrome(slide, title, n, footer="ACME Corp board paper"):
    """A band with a rule under it, the title, a footer, a hand-drawn page number."""
    _rect(slide, 0, 0, 13.333, 1.15, NAVY)
    _rect(slide, 0, 1.15, 13.333, 0.06, AMBER)
    _text(slide, 0.6, 0.12, 12, 0.9, title, size=28, bold=True, color="FFFFFF")
    _text(slide, 0.6, 7.05, 8, 0.3, footer, size=10, color=GREY)
    _text(slide, 11.7, 7.05, 1.0, 0.3, str(n), size=10, color=GREY)


def _import(prs, tmp_path: Path) -> str:
    path = tmp_path / "foreign.pptx"
    prs.save(path)
    text, diags = import_pptx(path)
    assert not [d for d in diags if d.level == "error"], diags
    return text


def _head(text: str) -> str:
    return text.split("\n\n# ")[0]


def _slide_md(text: str, n: int) -> str:
    """Slide ``n`` (1-based) of an imported deck (blocks split on a blank line before a ``# `` title)."""
    return re.split(r"\n\n(?=# )", text)[n]


def _filler(prs, n: int):
    """A plain content slide with the deck chrome (page numbers need more than one slide to be a pattern)."""
    s = _slide(prs)
    _chrome(s, f"Slide {n}", n)
    _text(s, 0.6, 1.6, 12, 1.0, f"Point {n}", size=24)
    return s


# --------------------------------------------------------------------------- 4. backgrounds and chrome


def _cover_deck():
    prs = _deck()
    s = _slide(prs)
    _rect(s, 0, 0, 13.333, 7.5, NAVY)
    _rect(s, 0.9, 2.2, 0.12, 2.3, AMBER)
    _text(s, 1.25, 2.1, 11, 1.6, "Mid-term plan 2027-2029", size=42, bold=True, color="FFFFFF")
    _text(s, 1.25, 3.8, 11, 0.7, "ACME Corp board paper, April 2027", size=20, color="C9D6E6")
    _rect(s, 0, 6.9, 13.333, 0.6, BLUE)
    for n, title in ((2, "Message"), (3, "Market")):
        t = _slide(prs)
        _chrome(t, title, n)
        _text(t, 0.6, 1.6, 12, 1.0, f"Point {n}", size=24)
    return prs


def test_cover_background_bars_and_chrome(tmp_path):
    text = _import(_cover_deck(), tmp_path)
    head = _head(text)
    assert "@bg=primary dark" in text or re.search(r"@bg=#122B4A dark", text)
    assert "cover.bar=#E08A1E" in head and "cover.bar_w=0.12in" in head
    assert "cover.bottom_bar=#1F5FA8" in head and "cover.bottom_bar_h=0.6in" in head
    assert "cover.band=none" in head and re.search(r"cover\.band_h=\d+%", head)
    assert "title.rule=#E08A1E" in head and "title.rule_h=0.06in" in head
    assert "num: on" in head and "footer: ACME Corp board paper" in head
    assert not re.search(r"^\^ \d$", text, re.M), "the hand-drawn page number is not a footnote"
    out = tmp_path / "back.pptx"
    build(text, out, base_dir=tmp_path)  # the tokens build and the bars are drawn back
    cover = Presentation(str(out)).slides[0]
    fills = {
        str(sh.fill.fore_color.rgb)
        for sh in cover.shapes
        if sh.shape_type == 1 and sh.fill.type == 1  # autoshapes with a solid fill
    }
    assert {AMBER, BLUE} <= fills
    assert str(cover.background.fill.fore_color.rgb) == NAVY or NAVY in fills


def test_plain_white_background_rect_is_not_a_cover_background(tmp_path):
    prs = _deck()
    s = _slide(prs)
    _rect(s, 0, 0, 13.333, 7.5, "FFFFFF")
    _text(s, 1, 2, 11, 1.5, "Title only", size=40, bold=True)
    text = _import(prs, tmp_path)
    assert "bg=" not in text and "cover.band" not in text


def test_page_number_helpers():
    prs = _deck()
    s = _slide(prs)
    num = _text(s, 11.7, 7.05, 1.0, 0.3, "3", size=10)
    body = _text(s, 0.6, 3, 4, 0.5, "3", size=10)
    W, H = int(prs.slide_width), int(prs.slide_height)
    from slidemark.importer.read import ReadCtx, read_slide

    data = read_slide(s, ReadCtx(accent="000000"))
    by = {i.uid: i for i in data.items}
    assert any(is_page_number(i, 3, W, H) for i in by.values())
    assert not any(is_page_number(i, 4, W, H) for i in by.values())
    assert sum(is_page_number(i, 3, W, H) for i in by.values()) == 1  # (the one near the top is body text)
    assert num is not body
    assert not has_page_numbers([data], W, H)  # a single slide is the cover: no evidence


def test_thin_strips_on_edges_become_top_and_bottom_bar(tmp_path):
    prs = _deck()
    for n in (1, 2, 3):
        s = _slide(prs)
        _rect(s, 0, 0, 13.333, 0.12, TEAL)
        _rect(s, 0, 7.38, 13.333, 0.12, AMBER)
        _text(s, 0.6, 0.5, 12, 0.9, f"Slide {n}", size=28, bold=True)
        _text(s, 0.6, 1.8, 12, 1.0, "Body line", size=20)
    head = _head(_import(prs, tmp_path))
    assert "top.bar=#1B8A8F" in head and "top.bar_h=0.12in" in head
    assert "bottom.bar=#E08A1E" in head


def test_short_segment_under_the_title_is_rule2(tmp_path):
    prs = _deck()
    for n in (1, 2, 3):
        s = _slide(prs)
        _text(s, 0.6, 0.4, 12, 0.8, f"Slide {n}", size=28, bold=True)
        _rect(s, 0.6, 1.25, 1.6, 0.06, AMBER)  # a short accent segment, no full-width line
        _text(s, 0.6, 1.8, 12, 1.0, "Body line", size=20)
    head = _head(_import(prs, tmp_path))
    assert "title.rule2=#E08A1E" in head and "title.rule2_w=1.6in" in head
    assert "title.rule=" not in head


def test_title_band_height_and_cover_rule(tmp_path):
    prs = _deck()
    cover = _slide(prs)
    _rect(cover, 0, 0, 13.333, 7.5, NAVY)
    _text(cover, 1, 2.2, 11, 1.2, "Plan", size=42, bold=True, color="FFFFFF")
    _text(cover, 1, 3.4, 11, 0.6, "April", size=20, color="C9D6E6")
    _rect(cover, 0, 4.6, 13.333, 0.05, AMBER)  # a line across the cover: the edge of the band
    for n in (2, 3):
        _filler(prs, n)
    head = _head(_import(prs, tmp_path))
    assert "cover.rule=#E08A1E" in head and "cover.rule_h=0.05in" in head
    assert "cover.band_h=61%" in head, head  # 4.6in of 7.5in: the rule sits on the band edge
    assert re.search(r"cover\.pad=0\.[5-9]\d?in", head), head
    assert "title.height=0.95in" in head, head  # a 1.15in band less half the 0.4in margin


def test_alternating_plain_rows_are_a_zebra(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    s = _slide(prs)
    _chrome(s, "Plain table", 2)
    tbl = s.shapes.add_table(5, 2, Inches(0.6), Inches(1.6), Inches(12), Inches(4)).table
    for r in range(5):
        for c in range(2):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                NAVY if r == 0 else (LIGHT if r % 2 == 0 else "FFFFFF")
            )
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = "h" if r == 0 else f"v{r}{c}"
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else "222B36")
    _filler(prs, 3)
    body = _slide_md(_import(prs, tmp_path), 2)
    assert ".zebra" in body and "table.zebra.fill=" in body and "table.hl" not in body
    assert "table.body.fill" not in body  # the first body row is white: the default body fill
    assert "color=" not in body, body  # the ink of every cell is the deck's own ink: no span


def test_cover_top_bar_token_draws_a_strip(tmp_path):
    tokens = "cover.band_h=60% cover.top_bar=#1B8A8F cover.top_bar_h=0.3in cover.bottom_bar=#E08A1E"
    md = f"style: {tokens}\n\n# Cover\nSub\n"
    out = tmp_path / "c.pptx"
    build(md, out)
    boxes = {
        (round(sh.top / 914400, 2), round(sh.height / 914400, 2)): str(sh.fill.fore_color.rgb)
        for sh in Presentation(str(out)).slides[0].shapes
        if sh.shape_type == 1 and sh.fill.type == 1
    }
    assert boxes.get((0.0, 0.3)) == "1B8A8F"
    assert any(top > 7.0 and rgb == "E08A1E" for (top, _h), rgb in boxes.items())


# --------------------------------------------------------------------------- 1. chevron sequence


def _steps_slide(prs, cards=True, bar=True, bar_rows=False):
    s = _slide(prs)
    _chrome(s, "Product plan", 2)
    fills = [BLUE, TEAL, AMBER]
    for i, (label, body) in enumerate(
        [
            ("2027:", "Release reconciliation"),
            ("2028:", "Cash-flow forecast"),
            ("2029:", "Industry templates"),
        ]
    ):
        x = 0.6 + i * 4.1
        arrow = _rect(s, x, 1.55, 3.9, 0.75, fills[i], MSO_SHAPE.PENTAGON)
        _text(s, 0, 0, 0, 0, [[(label, 24, True, "FFFFFF")]], shape=arrow)
        if cards:
            _rect(s, x, 2.5, 3.9, 2.7, LIGHT)
            _text(s, x + 0.25, 2.65, 3.4, 2.4, body, size=26, bold=True, color=NAVY)
    if bar:
        _rect(s, 0.6, 5.5, 12.13, 1.05, NAVY)
        _text(
            s,
            0.9,
            5.5,
            11.5,
            1.05,
            "Big release in Q1, improvements in Q3",
            size=26,
            bold=True,
            color="FFFFFF",
        )
    if bar_rows:  # stacked light rows of one fill and width: not a takeaway
        for k in range(2):
            _rect(s, 0.6, 5.3 + k * 0.6, 12.13, 0.5, LIGHT)
            _text(s, 0.9, 5.3 + k * 0.6, 11.5, 0.5, f"row {k}", size=18)
    return s


def test_chevron_row_with_cards_and_bar_is_steps(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    _steps_slide(prs)
    _filler(prs, 3)
    text = _import(prs, tmp_path)
    body = _slide_md(text, 2)
    assert "@3 steps" in body
    assert re.search(r"steps-arrow\.fill=#1F5FA8,#1B8A8F,(#E08A1E|secondary|accent)", body)
    assert "steps-arrow.size=24" in body and "steps-card.size=26" in body
    assert "## 2027:" in body and "Release reconciliation" in body
    assert re.search(r"^> .*Q1", body, re.M), "the dark bar is the conclusion"
    assert "conclusion.size=26" in body
    assert "render.chevron_shape=pentagon" in _head(text)
    out = tmp_path / "back.pptx"
    build(text, out, base_dir=tmp_path)
    shapes = {sh.name: sh for sh in Presentation(str(out)).slides[1].shapes}
    got = [str(shapes[f"Step {k} arrow"].fill.fore_color.rgb) for k in (1, 2, 3)]
    assert got == [BLUE, TEAL, AMBER]
    assert "Conclusion" in shapes


def test_chevrons_without_cards_stay_a_chevron_row(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    _steps_slide(prs, cards=False, bar=False)
    _filler(prs, 3)
    body = _slide_md(_import(prs, tmp_path), 2)
    assert "steps" not in body


def test_stacked_rows_are_not_a_takeaway_bar(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    s = _slide(prs)
    _chrome(s, "Targets", 2)
    for k in range(4):
        _rect(s, 0.6, 1.5 + k * 1.04, 12.13, 0.92, "101010" if k == 3 else LIGHT)
        _text(s, 1.0, 1.5 + k * 1.04, 11, 0.92, f"row {k}", size=20)
    _rect(s, 0.6, 5.7, 12.13, 0.9, LIGHT)
    _text(s, 1.0, 5.7, 11, 0.9, "last row", size=20)
    _filler(prs, 3)
    assert "conclusion." not in _import(prs, tmp_path)


# --------------------------------------------------------------------------- 2. tables


def _table_slide(prs, tint=True):
    s = _slide(prs)
    _chrome(s, "Comparison", 2)
    gf = s.shapes.add_table(4, 3, Inches(0.6), Inches(1.6), Inches(12.1), Inches(4.6))
    tbl = gf.table
    rows = [
        ("", "Fee", "Setup"),
        ("ACME:", "9,800", "3 days"),
        ("Beta:", "12,000", "10 days"),
        ("Gamma:", "7,500", "5 days"),
    ]
    fills = [NAVY, "FDF1DE" if tint else "FFFFFF", LIGHT, "FFFFFF"]
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(fills[r])
            p = cell.text_frame.paragraphs[0]
            if r == 0 or c == 0:
                run = p.add_run()
                run.text = v
                run.font.size = Pt(18 if r == 0 else 20)
                run.font.bold = True
                run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else NAVY)
            else:  # label above value: two runs of other sizes and colours
                a = p.add_run()
                a.text = "Item "
                a.font.size = Pt(15)
                a.font.color.rgb = RGBColor.from_string(GREY)
                b = p.add_run()
                b.text = v
                b.font.size = Pt(26)
                b.font.bold = True
                b.font.color.rgb = RGBColor.from_string(AMBER if r == 1 else BLUE)
    return s


def test_table_highlight_row_and_cell_run_spans(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    _table_slide(prs)
    _filler(prs, 3)
    body = _slide_md(_import(prs, tmp_path), 2)
    assert "table.hl.fill=#FDF1DE" in body and "table.hl.strength=1" in body
    assert re.search(r"\{[^}\n]*hl=ACME:[^}\n]*\}", body)
    assert re.search(r"\[\*\*9,800\*\*\]\{size=26 color=#E08A1E\}", body), body
    assert re.search(r"\[\*\*12,000\*\*\]\{size=26 color=#1F5FA8\}", body), body
    assert re.search(r"\{[^}\n]*size=15[^}\n]*\}", body), "the table states its base text size"
    deck = parse("colors: primary=#122B4A\n\n" + body)
    assert not [d for d in deck.diagnostics if d.level == "error"]


def test_table_without_a_tinted_row_has_no_highlight(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    _table_slide(prs, tint=False)
    _filler(prs, 3)
    body = _slide_md(_import(prs, tmp_path), 2)
    assert "table.hl" not in body and "hl=" not in body


def test_arrow_glyph_cell_stays_text_with_its_colour(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    s = _slide(prs)
    _chrome(s, "Risks", 2)
    tbl = s.shapes.add_table(3, 3, Inches(0.6), Inches(1.6), Inches(12), Inches(3)).table
    for r in range(3):
        for c, (v, col) in enumerate(
            [("Risk" if r == 0 else f"Risk {r}", NAVY), ("▶", AMBER), ("Answer", NAVY)]
        ):
            run = tbl.cell(r, c).text_frame.paragraphs[0].add_run()
            run.text = "" if (r == 0 and c == 1) else v
            run.font.size = Pt(20)
            run.font.bold = c == 1
            run.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else col)
            tbl.cell(r, c).fill.solid()
            tbl.cell(r, c).fill.fore_color.rgb = RGBColor.from_string(NAVY if r == 0 else "FFFFFF")
    _filler(prs, 3)
    body = _slide_md(_import(prs, tmp_path), 2)
    assert "▶" in body and "{color=#E08A1E}" in body


# --------------------------------------------------------------------------- 3. charts and the panel


def _chart_slide(prs, colored=True, panel=True, tiles=False):
    s = _slide(prs)
    _chrome(s, "Segments", 2)
    cd = CategoryChartData()
    cd.categories = ["A", "B", "C", "D"]
    cd.add_series("Share", [31, 24, 18, 12])
    gf = s.shapes.add_chart(
        XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.5), Inches(8.2), Inches(5.3), cd
    )
    ch = gf.chart
    ch.has_legend = False
    pl = ch.plots[0]
    pl.gap_width = 45
    pl.has_data_labels = True
    pl.data_labels.font.size = Pt(18)
    pl.data_labels.font.bold = True
    if colored:
        ser = pl.series[0]
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = RGBColor.from_string(BLUE)
        ser.points[0].format.fill.solid()
        ser.points[0].format.fill.fore_color.rgb = RGBColor.from_string(AMBER)
    if panel:
        _rect(s, 9.1, 1.6, 3.6, 5.1, LIGHT)
        _text(s, 9.3, 1.75, 3.2, 0.5, "Share of sales", size=14, bold=True, color=GREY)
        for i, (c, v) in enumerate(zip("ABCD", (31, 24, 18, 12), strict=True)):
            _text(
                s,
                9.3,
                2.35 + i * 0.85,
                3.2,
                0.7,
                [[(c + " ", 20, False, "222B36"), (f"{v}%", 26, True, AMBER if i == 0 else BLUE)]],
            )
    if tiles:  # three small cards beside the chart are tiles, not a panel
        for k in range(3):
            _rect(s, 9.1, 1.6 + k * 1.7, 3.6, 1.5, LIGHT)
            _text(
                s, 9.3, 1.6 + k * 1.7, 3.2, 1.5, [[(f"T{k} ", 20, False, "222B36"), ("9%", 30, True, BLUE)]]
            )
    return s


def test_chart_point_colours_label_size_and_gap(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    _chart_slide(prs, panel=False)
    _filler(prs, 3)
    body = _slide_md(_import(prs, tmp_path), 2)
    fence = re.search(r"```bar \{([^}]*)\}", body)
    assert fence, body
    assert "colors=#E08A1E,#1F5FA8,#1F5FA8,#1F5FA8" in fence.group(1), "one color per bar"
    assert "size=18" in fence.group(1) and "gap=45" in fence.group(1)
    assert "labels.bold=on" in fence.group(1)
    out = tmp_path / "back.pptx"
    build(_import(prs, tmp_path), out, base_dir=tmp_path)
    chart = next(sh for sh in Presentation(str(out)).slides[1].shapes if sh.has_chart).chart
    ser = chart.plots[0].series[0]
    assert str(ser.points[0].format.fill.fore_color.rgb) == AMBER
    assert str(ser.points[1].format.fill.fore_color.rgb) == BLUE


def test_chart_series_colours_and_default_palette(tmp_path):
    prs = _deck()
    s = _slide(prs)
    cd = CategoryChartData()
    cd.categories = ["Y1", "Y2"]
    cd.add_series("Eng", [60, 80])
    cd.add_series("Other", [60, 70])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(1), Inches(1), Inches(6), Inches(5), cd)
    for ser, rgb in zip(gf.chart.plots[0].series, (BLUE, GREY), strict=True):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = RGBColor.from_string(rgb)
    plain = _slide(prs)  # no explicit fills: the build palette draws it, nothing to state
    cd2 = CategoryChartData()
    cd2.categories = ["Q1", "Q2"]
    cd2.add_series("S", [1, 2])
    plain.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(6), Inches(5), cd2)
    text = _import(prs, tmp_path)
    fences = re.findall(r"```(?:stacked-column|column) \{?([^\n]*)", text)
    assert re.search(r"colors=#1F5FA8,(#6B7582|muted)", fences[0]), fences[
        0
    ]  # (muted: the theme grey, 2 off)
    assert "colors" not in fences[1]


def test_panel_beside_chart_gets_spans_tiles_do_not(tmp_path):
    prs = _deck()
    _filler(prs, 1)
    _chart_slide(prs)
    _filler(prs, 3)
    body = _slide_md(_import(prs, tmp_path), 2)
    assert re.search(r"\[\*\*31%\*\*\]\{size=26 color=(#E08A1E|secondary|accent)\}", body), body
    assert re.search(r"\[\*\*24%\*\*\]\{size=26 color=(#1F5FA8|primary)\}", body), body
    assert re.search(r"^@2:1", body, re.M), body
    prs2 = _deck()
    _filler(prs2, 1)
    _chart_slide(prs2, panel=False, tiles=True)
    _filler(prs2, 3)
    assert "size=30" not in _slide_md(_import(prs2, tmp_path), 2)


def test_recognition_is_off_for_a_deck_with_its_own_design(tmp_path):
    """A SlideMark deck says its design itself: no geometry recognition (round trips stay as they were)."""
    src = "# T\n@3 steps\n## a\nx\n## b\ny\n## c\nz\n> done\n"
    pptx = tmp_path / "own.pptx"
    build(src, pptx)
    text, _ = import_pptx(pptx)
    assert "steps-arrow.fill" not in text and "conclusion.fill" not in text and "cover." not in text


def test_thresholds_are_one_table():
    assert T.bleed > 0.5 and T.strip_w > 0.9 and 0 < T.chrome_share < 1
