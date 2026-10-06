"""Decision slides: capped card text, ruled spread lists, conclusion bar shrink-before-wrap, bar vs footer."""

# ruff: noqa: E501
from __future__ import annotations

import random
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from slidemark import build
from slidemark.ir import Shape
from slidemark.layout import layout_slide, measure
from slidemark.parser import parse
from slidemark.template import deck_theme, resolve_theme
from slidemark.theme import LayoutTokens
from slidemark.units import to_emu

from .test_layout_l3_fill import EX, cards, inside, lay, role

CARDS = """# Decision
> Lead
## Left
1. Alpha item
2. Beta item
3. Gamma item
## Right
- One
- Two
- Three
{bar}
"""


def _rules(placed):
    return [p for p in placed if isinstance(p.element, Shape) and p.element.shape == "line"]


def _src(bar="> Approve now", head="", foot=""):
    return head + CARDS.format(bar=bar) + foot


def test_tokens_exist_and_parse_off():
    lt = LayoutTokens()
    assert lt.card_text_max == 18 and lt.card_spread_rules is True and lt.conclusion_min_scale == 0.8
    deck = parse("style: layout.card_spread_rules=off\n\n" + _src())
    theme, diags = deck_theme(deck, EX)
    assert theme.layout.card_spread_rules is False and not diags
    assert not _rules(layout_slide(deck.slides[0], deck, theme, 0))
    deck = parse(_src())
    assert _rules(layout_slide(deck.slides[0], deck, deck_theme(deck, EX)[0], 0))


@pytest.mark.parametrize(("name", "n"), [("19-vi-consulting-brand.md", 5), ("16-jp-strategy.md", 10)])
def test_card_text_is_capped_and_never_below_theme_size(name, n):
    placed, _, theme = lay(name, n)
    for c in cards(placed):
        (b,) = inside(placed, c)
        assert b.style.font_size * b.font_scale <= theme.layout.card_text_max + 0.01
        assert b.font_scale >= 1.0 - 1e-6  # not shrunk below the theme size


def test_card_text_max_zero_restores_growth():
    placed, _, _ = lay("19-vi-consulting-brand.md", 5, card_text_max=0)
    (b,) = inside(placed, cards(placed)[0])
    assert b.style.font_size * b.font_scale > 18.5


def test_explicit_size_bypasses_the_cap():
    src = _src().replace("## Left\n", "## Left\n{size=24}\n")
    placed, _, _ = lay("x", 1, src=src)
    sizes = [b.style.font_size * b.font_scale for c in cards(placed) for b in inside(placed, c)]
    assert max(sizes) >= 23.9


@pytest.mark.parametrize(
    ("name", "n", "k"), [("19-vi-consulting-brand.md", 5, 4), ("16-jp-strategy.md", 10, 5)]
)
def test_spread_items_get_rules_between_them(name, n, k):
    placed, _, theme = lay(name, n)
    rules = _rules(placed)
    assert len(rules) == k  # items - 1 per card
    for r in rules:
        assert r.h == 0 and r.style.line == "border" and r.style.line_width == 0.75
        c = next(c for c in cards(placed) if c.x <= r.x and r.x + r.w <= c.x + c.w and c.y < r.y < c.y + c.h)
        assert r.x > c.x and r.x + r.w < c.x + c.w  # inset by the card padding
        (b,) = inside(placed, c)
        assert abs((r.x - c.x) - (b.x - c.x)) <= to_emu("0.15in")
    for c in cards(placed):  # both cards of the row spread: the last item ends near the card bottom
        (b,) = inside(placed, c)
        end = b.y + measure.paragraphs_height(
            b.element.paragraphs, b.w, b.style, b.font_scale, gap=b.element.attrs.get("para_gap")
        )
        assert (c.y + c.h - end) / c.h < 0.3
        mine = sorted(r.y for r in rules if c.x < r.x < c.x + c.w and c.y < r.y < c.y + c.h)
        assert all(mine[i + 1] > mine[i] for i in range(len(mine) - 1))


def test_rules_sit_between_items():
    placed, _, _ = lay("19-vi-consulting-brand.md", 5)
    c = sorted(cards(placed), key=lambda p: p.x)[0]
    (b,) = inside(placed, c)
    ys = sorted(r.y for r in _rules(placed) if c.x < r.x < c.x + c.w and c.y < r.y < c.y + c.h)
    paras = b.element.paragraphs
    g = b.element.attrs["para_gap"]
    size = b.style.font_size * b.font_scale * 12700
    for i, y in enumerate(ys):
        top = b.y + measure.paragraphs_height(paras[: i + 1], b.w, b.style, b.font_scale, gap=g)
        assert top < y < top + g * size + 2


def test_rules_off_keeps_top_anchored_lists():
    placed, _, _ = lay("16-jp-strategy.md", 10, card_spread_rules=False)
    assert not _rules(placed)
    assert len({inside(placed, c)[0].y for c in cards(placed)}) == 1


def test_rule_shapes_render_as_native_connectors(tmp_path):
    out = tmp_path / "x.pptx"
    build(str(EX / "19-vi-consulting-brand.md"), str(out))
    slide = Presentation(str(out)).slides[4]
    conns = [sh for sh in slide.shapes if sh._element.tag.endswith("}cxnSp")]
    assert len(conns) == 4
    for cx in conns:
        assert cx.height == Emu(0)
        assert cx.line.width.pt == pytest.approx(0.75)
        assert str(cx.line.color.rgb) == "D9D2C5"  # the deck's border color


def _bar_size(placed):
    b = role(placed, "conclusion")[0]
    return b.style.font_size * b.font_scale, b


def test_bar_shrinks_before_it_wraps():
    short = "> Approve"
    mid = "> " + "Approving the first phase budget this month keeps the pilot on schedule for the board " * 1
    placed, _, theme = lay("x", 1, src=_src(bar=mid.strip()), conclusion_min_ratio=0.0)
    base = theme.sizes.get("conclusion", 18)
    size, bar = _bar_size(placed)
    one, _ = _bar_size(lay("x", 1, src=_src(bar=short))[0])
    assert one >= base
    assert size <= max(base, one) + 1e-6
    # a sentence too long even at the floor keeps the theme size and wraps
    long = "> " + "word " * 60
    s_long, b_long = _bar_size(lay("x", 1, src=_src(bar=long.strip()))[0])
    assert s_long == pytest.approx(base) and b_long.h > to_emu("0.5in")


def test_bar_floor_scale_is_token_driven():
    txt = "> " + "Approving the first phase budget this month keeps the pilot on track, " * 2
    p1, _, th = lay("x", 1, src=_src(bar=txt.strip()), conclusion_min_scale=1.0)
    p2, _, _ = lay("x", 1, src=_src(bar=txt.strip()), conclusion_min_scale=0.8)
    base = th.sizes.get("conclusion", 18)
    s1, _ = _bar_size(p1)
    s2, _ = _bar_size(p2)
    assert s1 == pytest.approx(base)  # no shrink allowed: theme size
    assert base * 0.8 - 1e-6 <= s2 <= base + 1e-6 and s2 >= 12


def test_bar_keeps_a_gutter_above_the_footer_without_footnote():
    placed, _, theme = lay("19-vi-consulting-brand.md", 5)
    bar = role(placed, "conclusion")[0]
    foot = [p for p in role(placed, "caption") if p.element.attrs.get("field") == "footer"][0]
    assert foot.y - (bar.y + bar.h) >= to_emu(theme.gap) - 2
    cs = cards(placed)
    last = max(c.y + c.h for c in cs)
    assert abs(bar.y - last - to_emu(theme.gap)) <= 12700  # still one gutter below the cards


def test_fuzz_never_raises():
    rnd = random.Random(7)
    words = ["a", "bb ccc", "日本語のテキスト", "x" * 40, "1.", "- item"]
    for _ in range(25):
        body = "\n".join(
            rnd.choice(["- ", "1. ", "", "> "]) + rnd.choice(words) for _ in range(rnd.randint(1, 7))
        )
        src = f"num: on\nfooter: f\n\n# T\n> L\n## A\n{body}\n## B\n- {rnd.choice(words)}\n- {rnd.choice(words)}\n> {rnd.choice(words) * 3}\n"
        deck = parse(src)
        theme, _ = resolve_theme(deck.theme, Path("."))
        layout_slide(deck.slides[0], deck, theme, 0)


def _rule_ys(placed, card):
    return sorted(
        r.y for r in _rules(placed) if card.x <= r.x < card.x + card.w and card.y <= r.y <= card.y + card.h
    )


def test_sibling_cards_share_item_rows_and_rules():
    placed, _, _ = lay("11-jp-consulting.md", 8)
    top = sorted((c for c in cards(placed) if c.y == min(c2.y for c2 in cards(placed))), key=lambda c: c.x)
    assert len(top) == 2
    (b0,), (b1,) = (inside(placed, c) for c in top)
    assert b0.y == b1.y  # first items start at one height
    assert b0.element.attrs["para_gap"] == b1.element.attrs["para_gap"]  # one pitch
    ys0, ys1 = (_rule_ys(placed, c) for c in top)
    assert len(ys0) == 2 and len(ys1) == 1
    assert ys0[0] == ys1[0]  # the first divider is level across the row


def test_matching_item_counts_align_every_rule():
    src = _src()
    deck = parse(
        src.replace("- One\n- Two\n- Three", "- One\n- A much longer second item that wraps maybe\n- Three")
    )
    theme, _ = deck_theme(deck, EX)
    placed = layout_slide(deck.slides[0], deck, theme, 0)
    cs = sorted((c for c in placed if c.element.__class__.__name__ == "Container"), key=lambda c: c.x)
    ys = [_rule_ys(placed, c) for c in cs]
    assert len(ys) == 2 and ys[0] and len(ys[0]) == len(ys[1])
    assert ys[0] == ys[1]


@pytest.mark.parametrize(
    "name", ["11-jp-consulting.md", "16-jp-strategy.md", "20-jp-retail-dense.md", "22-en-launch-plan.md"]
)
def test_sibling_cards_on_a_row_share_one_text_size(name):
    deck = parse((EX / name).read_text(encoding="utf-8"))
    theme, _ = deck_theme(deck, EX)
    for i, slide in enumerate(deck.slides):
        placed = layout_slide(slide, deck, theme, i)
        rows: dict[int, list] = {}
        for c in cards(placed):
            rows.setdefault(c.y, []).append(c)
        for row in rows.values():
            sizes = {
                round((p.style.font_size or 0) * p.font_scale, 1)
                for c in row
                for p in inside(placed, c)
                if "callout" not in p.element.classes
            }
            assert len(sizes) <= 1, (name, i + 1, sizes)
