"""DL3d: bench/roundtrip_compare.py and the title detection it exposed (badge digit, quote glyph)."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from slidemark.build import build
from slidemark.importer import import_pptx

from .helpers import needs_soffice

ROOT = Path(__file__).resolve().parent.parent
NAVY, LIGHT, ORANGE = RGBColor(0x12, 0x2B, 0x4A), RGBColor(0xEE, 0xF2, 0xF7), RGBColor(0xE0, 0x8A, 0x1E)
WHITE = RGBColor(255, 255, 255)
KEYS = {
    "per_slide",
    "mean",
    "max",
    "tolerance",
    "over_tolerance",
    "md_tokens",
    "build_py_tokens",
    "token_ratio",
    "import_warnings",
    "slides",
    "sheets",
}


def _tool():
    spec = importlib.util.spec_from_file_location(
        "roundtrip_compare", ROOT / "bench" / "roundtrip_compare.py"
    )
    mod = importlib.util.module_from_spec(spec)
    sys.modules["roundtrip_compare"] = mod
    spec.loader.exec_module(mod)
    return mod


def _box(slide, x, y, w, h, rgb=None, text=None, size=16, bold=False, color=NAVY, align=None, shape=True):
    if shape:
        s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        style = s._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
        if style is not None:
            s._element.remove(style)
        if rgb is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = rgb
        s.line.fill.background()
    else:
        s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if text is not None:
        p = s.text_frame.paragraphs[0]
        if align is not None:
            p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return s


def _new():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _foreign_deck(path: Path) -> Path:
    """Two slides as a python-pptx script draws them: a header band and title text box, then (1) numbered rows
    whose badge digit is a bigger-than-title shape added after the row text, (2) a quote card, 96 pt mark."""
    prs = _new()

    def header(s, title):
        _box(s, 0, 0, 13.333, 1.15, NAVY)
        _box(s, 0.6, 0.12, 12.13, 0.9, text=title, size=28, bold=True, color=WHITE, shape=False)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Requests to the board")
    for i, t in enumerate(["Approve the plan", "Approve the budget", "Approve the price change"]):
        y = 1.6 + 1.3 * i
        _box(s, 0.6, y, 12.13, 1.12, LIGHT)
        _box(s, 2.0, y, 10.33, 1.12, text=t, size=30, bold=True, shape=False)
        _box(s, 0.6, y, 1.1, 1.12, NAVY, str(i + 1), 32, True, WHITE, PP_ALIGN.CENTER)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    header(s, "Customer voice")
    _box(s, 0.6, 1.7, 12.13, 4.9, LIGHT)
    _box(s, 0.6, 1.7, 0.15, 4.9, ORANGE)
    _box(s, 1.0, 1.7, 1.5, 1.6, text="“", size=96, bold=True, color=ORANGE, shape=False)
    _box(
        s,
        1.6,
        2.5,
        10.13,
        2.5,
        text="Closing the month went from three days to half a day.",
        size=38,
        shape=False,
    )
    _box(s, 1.6, 5.3, 10.13, 0.8, text="- Miura Works, controller", size=24, shape=False, color=ORANGE)
    prs.save(str(path))
    return path


def _title_lines(text: str) -> list[str]:
    return [ln for ln in text.splitlines() if ln.startswith("# ")]


def test_badge_digit_and_quote_mark_are_not_titles(tmp_path):
    text, _ = import_pptx(_foreign_deck(tmp_path / "f.pptx"), tmp_path)
    assert _title_lines(text) == ["# Requests to the board", "# Customer voice"]


def test_a_year_alone_is_still_a_title(tmp_path):
    prs = _new()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(s, 0.6, 0.4, 12, 1.0, text="2029", size=40, bold=True, shape=False)
    prs.save(str(tmp_path / "y.pptx"))
    text, _ = import_pptx(tmp_path / "y.pptx", tmp_path)
    assert _title_lines(text) == ["# 2029"]


def test_title_is_the_top_wide_box_whatever_the_z_order(tmp_path):
    prs = _new()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _box(
        s, 0.6, 1.8, 5, 1.0, text="Side note under the title", size=26, shape=False
    )  # first in z-order, lower
    _box(s, 0.6, 0.3, 12, 1.0, text="The real title of the slide", size=30, bold=True, shape=False)
    prs.save(str(tmp_path / "z.pptx"))
    text, _ = import_pptx(tmp_path / "z.pptx", tmp_path)
    assert _title_lines(text) == ["# The real title of the slide"]


@needs_soffice
def test_tool_on_a_foreign_deck(tmp_path):
    mod = _tool()
    r = mod.compare(_foreign_deck(tmp_path / "f.pptx"), tmp_path / "out")
    assert KEYS <= set(r)
    assert r["slides"] == r["slides_imported"] == 2 and len(r["per_slide"]) == 2
    assert r["build_py_tokens"] is None and r["md_tokens"] > 0
    assert (tmp_path / "out" / "compare-1.png").is_file() and (tmp_path / "out" / "report.json").is_file()
    assert "mean diff" in mod.summary(r)
    assert "# Requests to the board" in (tmp_path / "out" / "deck.md").read_text(encoding="utf-8")


@needs_soffice
def test_own_deck_round_trips_to_about_zero(tmp_path):
    md = "# Plan\n\n- one\n- two\n\n---\n\n# Numbers\n\n## Sales\n\n42%\n\n## Cost\n\n17%\n"
    build(md, tmp_path / "own.pptx")
    r = _tool().compare(tmp_path / "own.pptx", tmp_path / "out")
    assert r["slides"] == 2 and r["mean"] <= 2.0, r["per_slide"]


def test_missing_original_without_regen_is_an_error(tmp_path):
    mod = _tool()
    with pytest.raises(FileNotFoundError):
        mod.compare(tmp_path / "nope.pptx", tmp_path / "out")
    assert mod.main([str(tmp_path / "nope.pptx"), "--out", str(tmp_path / "o")]) == 2
