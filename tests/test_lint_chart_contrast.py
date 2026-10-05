"""Chart text (title, legend, axis labels, data labels) is contrast-linted with the renderer's colors."""

from __future__ import annotations

import re
from pathlib import Path

from slidemark.layout import layout_slide
from slidemark.lint import lint
from slidemark.parser import parse
from slidemark.template import deck_theme
from slidemark.theme import Theme

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"
COL = "# T\n\n```column {labels=on title=Sales}\nq,a,b\nQ1,3,4\nQ2,5,6\n```\n"
PIE = "# T\n\n```pie {labels=on title=Share}\nAPAC,48\nEU,30\nUS,22\n```\n"
DONUT = PIE.replace("```pie", "```doughnut")
STACK = COL.replace("```column", "```stacked-column")


def _warnings(md: str) -> list:
    deck = parse(md)
    th, _ = deck_theme(deck, EXAMPLES)
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    return [d for d in lint(deck, placed, th) if d.rule == "contrast" and d.message.startswith("chart")]


def _kinds(got: list) -> list[str]:
    return sorted(re.match(r"chart (.+?) have", d.message).group(1) for d in got)  # type: ignore[union-attr]


def test_unreadable_text_is_grouped_per_kind_and_hints_fix_it():
    md = "colors: bg=#FFFFFF fg=#DDDDDD\n\n" + COL
    got = _warnings(md)
    assert _kinds(got) == ["axis labels", "data labels", "legend", "title"]  # one per kind, not per label
    fence = {d.message.split()[1]: d.hint for d in got}
    title = fence["title"]
    assert title.startswith("colors: fg=")
    ink = re.search(r"\{color=(#\w+)\}", fence["legend"]).group(1)  # type: ignore[union-attr]
    fixed = md.replace("fg=#DDDDDD", "fg=" + title.split("=")[1].split()[0]).replace(
        "{labels=on", f"{{color={ink} labels=on"
    )
    assert _warnings(fixed) == [], fixed


def test_chart_color_attribute_is_judged_on_the_card():
    md = "# T\n\n```column {color=#FFFFFF labels=on}\nq,a\nQ1,3\nQ2,5\n```\n"
    got = _warnings(md)
    assert _kinds(got) == ["axis labels", "data labels", "legend"]
    ink = re.search(r"\{color=(#\w+)\}", got[0].hint).group(1)  # type: ignore[union-attr]
    assert _warnings(md.replace("#FFFFFF", ink)) == []


def test_default_themes_and_chart_kinds_are_quiet():
    for head in ("", "theme: midnight\n\n", "theme: jp-business\n\n"):
        for src in (
            COL,
            PIE,
            DONUT,
            STACK,
            COL.replace("```column", "```bar"),
            COL.replace("```column", "```line"),
        ):
            assert _warnings(head + src) == [], (head, src)


def test_slice_label_ink_that_fails_on_the_slice_warns_and_labels_off_clears(monkeypatch):
    monkeypatch.setattr(Theme, "chart_label_ink", lambda self, fill, *backs: fill)  # ink == slice: unreadable
    got = _warnings(PIE)
    assert _kinds(got) == ["data labels"]
    assert "{labels=off}" in got[0].hint
    assert _warnings(PIE.replace("{labels=on", "{labels=off")) == []
    assert _kinds(_warnings(STACK)) == ["data labels"]


def test_examples_have_no_chart_warnings():
    for md in sorted(EXAMPLES.glob("*.md")):
        deck = parse(md.read_text(encoding="utf-8"))
        th, _ = deck_theme(deck, md.parent)
        placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
        bad = [str(d) for d in lint(deck, placed, th) if d.rule == "contrast" and "chart" in d.message]
        assert not bad, (md.name, bad)


def test_chart_palette_resolver_matches_render(tmp_path):
    """Lint and render take slice colors and label inks from the same Theme methods."""
    from pptx import Presentation

    from slidemark import build

    out = tmp_path / "c.pptx"
    build(PIE, out)
    chart = next(s for s in Presentation(out).slides[0].shapes if s.has_chart).chart
    th, _ = deck_theme(parse(PIE), EXAMPLES)
    pal = th.chart_palette(None)
    ser = chart.plots[0].series[0]
    for i in range(len(ser.values)):
        assert str(ser.points[i].format.fill.fore_color.rgb) == pal[i]
        assert str(ser.points[i].data_label.font.color.rgb) == th.chart_label_ink(pal[i], "FFFFFF")
