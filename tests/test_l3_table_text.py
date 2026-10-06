"""Stretched tables grow their text (vfill), rows stay <= table_row_max_em; badges get side padding."""

from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation

from slidemark.ir import Table, Text
from slidemark.layout import layout_slide
from slidemark.parser import parse
from slidemark.render import render
from slidemark.template import deck_theme

EMU_PT = 12700
EX = Path(__file__).resolve().parent.parent / "examples"


def table_of(path: Path, slide: int, vtext: float | None = None):
    deck = parse(path.read_text(encoding="utf-8"))
    theme, _ = deck_theme(deck)
    if vtext is not None:
        theme = copy.deepcopy(theme)
        theme.layout.table_vtext_max = vtext
    items = layout_slide(deck.slides[slide - 1], deck, theme, slide - 1)
    tab = next(p for p in items if isinstance(p.element, Table))
    foot = [p for p in items if isinstance(p.element, Text) and p.element.role == "footnote"]
    return tab, foot, theme


def size(p) -> float:
    return (p.style.font_size or 14) * p.font_scale


def test_roadmap_table_text_grows_and_rows_capped():
    off, _, _ = table_of(EX / "16-jp-strategy.md", 9, vtext=1.0)
    on, _, theme = table_of(EX / "16-jp-strategy.md", 9)
    assert size(on) > size(off) * 1.1
    cap = theme.layout.table_row_max_em * size(on) * EMU_PT
    assert max(on.element.attrs["_row_h"]) <= cap + 2


def test_table_below_chevrons_reaches_toward_footnote():
    off, foot, _ = table_of(EX / "11-jp-consulting.md", 4, vtext=1.0)
    on, _, theme = table_of(EX / "11-jp-consulting.md", 4)
    assert foot
    gap_off = foot[0].y - (off.y + off.h)
    gap_on = foot[0].y - (on.y + on.h)
    assert gap_on < gap_off * 0.6
    assert size(on) >= size(off)  # chevron text caps the table text: only the rows take the slack
    cap = max(theme.layout.table_row_max_em, theme.layout.table_vrow_max_em)
    assert max(on.element.attrs["_row_h"]) <= cap * size(on) * EMU_PT + 2
    assert on.y + on.h < foot[0].y  # never into the footnote


DENSE = "# Plan\n```table\nItem,Q1\n" + "".join(f"R{i},{i}\n" for i in range(9)) + "```\n"


def test_dense_table_unchanged():
    deck = parse(DENSE)
    theme, _ = deck_theme(deck)
    a = next(p for p in layout_slide(deck.slides[0], deck, theme, 0) if isinstance(p.element, Table))
    theme2 = copy.deepcopy(theme)
    theme2.layout.table_vtext_max = 1.0
    b = next(p for p in layout_slide(deck.slides[0], deck, theme2, 0) if isinstance(p.element, Table))
    assert a.font_scale == b.font_scale and a.element.attrs["_row_h"] == b.element.attrs["_row_h"]


def test_odd_tables_never_raise():
    for body in (
        "a|b\n-|-\n",
        "|x|\n|-|\n|" + "あ" * 200 + "|\n",
        "```table\n```\n",
        "```table\n,\n,\n```\n",
    ):
        deck = parse("# T\n" + body)
        theme, _ = deck_theme(deck)
        layout_slide(deck.slides[0], deck, theme, 0)


BADGE = "# T\n| 項目 | 状態 |\n|-|-|\n| A | 開始[完了]{.badge .success} |\n| B | Go [OK]{.badge .success} |\n"


def runs_of_badge_cells(tmp_path, md: str):
    deck = parse(md)
    theme, _ = deck_theme(deck)
    out = tmp_path / "b.pptx"
    render(deck, [layout_slide(s, deck, theme, i) for i, s in enumerate(deck.slides)], theme, out)
    tbl = next(s for s in Presentation(str(out)).slides[0].shapes if s.has_table).table
    return [[r for r in tbl.cell(row, 1).text_frame.paragraphs[0].runs] for row in (1, 2)]


def test_badge_run_padded_and_separated(tmp_path):
    cjk_cell, latin_cell = runs_of_badge_cells(tmp_path, BADGE)
    # CJK badge: full-width pad inside the highlighted run; a plain space ends the text before it
    assert [r.text for r in cjk_cell] == ["開始 ", "\u3000完了\u3000"]
    assert "highlight" in cjk_cell[1]._r.xml and "highlight" not in cjk_cell[0]._r.xml
    # Latin badge: no-break spaces inside the highlight; the text before already ends in a space
    assert [r.text for r in latin_cell] == ["Go ", "\u00a0\u00a0OK\u00a0\u00a0"]
