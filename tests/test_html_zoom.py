"""HTML block zoom: short content re-renders at a narrower viewport and scales back into the box."""

from __future__ import annotations

import pytest
from pptx import Presentation

from slidemark.build import build
from slidemark.htmlnative import html_to_placed
from slidemark.ir import Text
from slidemark.render.htmlimg import HtmlRenderer
from slidemark.theme import Theme

PX = 9525
THEME = Theme(name="t")
GRID = (
    '<div style="display:grid;grid-template-columns:1fr 1fr 1fr;gap:16px;font-family:sans-serif">'
    + "".join(
        f'<div style="background:#eef;padding:12px"><h3 style="margin:0">Card {i}</h3><p>Short text</p></div>'
        for i in range(3)
    )
    + "</div>"
)
TALL = '<div style="background:#eef;height:290px;font-size:16px;font-family:sans-serif">Fills the box</div>'


@pytest.fixture()
def renderer():
    pytest.importorskip("playwright.sync_api")
    r = HtmlRenderer()
    if r.render("<p>x</p>", 40, 20, THEME) is None:
        r.close()
        pytest.skip("Chromium cannot be launched")
    yield r
    r.close()


def _fs(items):
    return sorted({p.style.font_size for p in items if isinstance(p.element, Text)})


def _bottom(items):
    return max(p.y + p.h for p in items) / PX


def _run(renderer, html, **kw):
    res = html_to_placed(html, 0, 0, 900 * PX, 400 * PX, THEME, ".", renderer=renderer, **kw)
    assert res is not None
    return res[0]


def test_auto_zoom_scales_fonts_and_stays_in_box(renderer):
    base = _run(renderer, GRID)
    zoomed = _run(renderer, GRID, zoom="auto", zoom_max=1.5)
    assert _fs(zoomed)[0] == pytest.approx(_fs(base)[0] * 1.5, abs=0.6)
    assert _bottom(zoomed) > _bottom(base) * 1.3
    assert _bottom(zoomed) <= 400 * 0.93
    assert max(p.x + p.w for p in zoomed) <= 900 * PX + 2


def test_no_zoom_when_block_fills_its_box(renderer):
    base = _run(renderer, TALL)
    auto = _run(renderer, TALL, zoom="auto", zoom_max=1.5)
    assert _fs(auto) == _fs(base)
    assert _bottom(auto) == pytest.approx(_bottom(base))


def test_zoom_one_and_zoom_max_one_disable(renderer):
    base = _run(renderer, GRID)
    assert _fs(_run(renderer, GRID, zoom=1)) == _fs(base)
    assert _fs(_run(renderer, GRID, zoom="auto", zoom_max=1.0)) == _fs(base)


def test_explicit_zoom_number(renderer):
    base = _run(renderer, GRID)
    z = _run(renderer, GRID, zoom=1.2)
    assert _fs(z)[0] == pytest.approx(_fs(base)[0] * 1.2, abs=0.6)


def _deck(attr: str) -> str:
    return f"# T\n\n```html{attr}\n{GRID}\n```\n"


def _first_font(tmp_path, src: str) -> float:
    md = tmp_path / "d.md"
    md.write_text(src, encoding="utf-8")
    out = tmp_path / "d.pptx"
    build(md, out)
    sizes = [
        r.font.size.pt
        for s in Presentation(str(out)).slides[0].shapes
        if s.has_text_frame and s.text_frame.text.startswith("Card")
        for p in s.text_frame.paragraphs
        for r in p.runs
        if r.font.size
    ]
    assert sizes
    return max(sizes)


def test_block_in_slide_zooms_and_attr_disables(renderer, tmp_path):
    renderer.close()  # one sync Playwright per process
    auto = _first_font(tmp_path, _deck(" {render=native}"))
    off = _first_font(tmp_path, _deck(" {render=native zoom=1}"))
    assert auto > off * 1.2


def test_html_slide_is_not_zoomed(renderer, tmp_path):
    renderer.close()
    src = f"# T\n@html\n```html\n{GRID}\n```\n"
    got = _first_font(tmp_path, src)
    off = _first_font(tmp_path, _deck(" {render=native zoom=1}"))
    assert got == pytest.approx(off, abs=0.6)  # authored size, no zoom


def test_svg_data_uri_image_with_charset_param_is_kept(renderer):
    from slidemark.ir import Image

    src = (
        "data:image/svg+xml;utf8,%3Csvg xmlns='http://www.w3.org/2000/svg' width='40' height='40'%3E"
        "%3Crect width='40' height='40' fill='red'/%3E%3C/svg%3E"
    )
    items = _run(renderer, f'<img src="{src}" width="40" height="40"><p>caption</p>')
    imgs = [p.element for p in items if isinstance(p.element, Image)]
    assert imgs and imgs[0].src.endswith(".svg")
