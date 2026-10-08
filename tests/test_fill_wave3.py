"""Wave 3 lane E: sparse forms fill like cards. The icon list stacks its icon above the text in a row of
columns and takes ``iconlist.title.max`` / ``iconlist.text.max``; box cards with free height inside grow their
heading, text and icon disc (``card.heading.max`` / ``card.text.max``); the conclusion bar is never under the
body text of its slide (it wraps first, ``conclusion.floor``)."""

from __future__ import annotations

import random
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from slidemark import build
from slidemark.fit import fit_report
from slidemark.ir import Container, Shape, Text
from slidemark.layout import layout_slide
from slidemark.layout.grid import Rect
from slidemark.layout.l3fill import fill_box_cards
from slidemark.parser import parse
from slidemark.template import deck_theme

EX = Path(__file__).resolve().parent.parent / "examples"
SLIDE_W, SLIDE_H = 12192000, 6858000

HEAD = (
    "theme: none\ncolors: bg=#FFFFFF fg=#1D2B27 primary=#16423C secondary=#4E9F8E accent=#F29E4C"
    " surface=#EAF3F0 border=#CFE0DA muted=#5E7069\n"
    'fonts: heading="Cambria" body="Calibri"\nsizes: title=38! heading=22 body=16 lead=20 caption=12\n'
    "style: radius=10 title.band=none card.elevation=1 icon.disc=secondary{extra}\n\n"
)

ICONS = (
    "- icon=check **Có giám sát** Học từ dữ liệu đã có nhãn đúng. Ví dụ: phân loại thư rác\n"
    "- icon=search **Không giám sát** Tự tìm cấu trúc trong dữ liệu không nhãn. Ví dụ: gom nhóm khách hàng\n"
    "- icon=refresh **Tăng cường** Thử, sai và nhận phần thưởng để cải thiện. Ví dụ: điều khiển robot\n"
)
SHORT_CARDS = (
    "# Ba lợi ích\n@3\n## Nhanh {icon=bolt}\nPhản hồi trong vài giây.\n"
    "## Rẻ {icon=gear}\nChi phí thấp hơn nhiều.\n## Dễ {icon=check}\nKhông cần chuyên gia.\n"
)
BAR = "> Giá trị của AI đến từ cách dùng: chọn một bài toán, thử trong bốn tuần, rồi quyết định tiếp\n"


def lay(body: str, extra: str = "", n: int = 1, **layout):
    src = HEAD.format(extra=(" " + extra) if extra else "") + body
    deck = parse(src)
    deck.attrs["base_dir"] = str(EX)
    theme, _ = deck_theme(deck, EX)
    if layout:
        theme = theme.model_copy(update={"layout": theme.layout.model_copy(update=layout)})
    return layout_slide(deck.slides[n - 1], deck, theme, n - 1), deck, theme


def pt(p) -> float:
    return (p.style.font_size or 0) * p.font_scale


def texts(placed, role):
    return [p for p in placed if isinstance(p.element, Text) and p.element.role == role]


def inside(a, b) -> bool:
    return a.x >= b.x - 2 and a.y >= b.y - 2 and a.x + a.w <= b.x + b.w + 2 and a.y + a.h <= b.y + b.h + 2


def disjoint(a, b) -> bool:
    return a.x + a.w <= b.x + 2 or b.x + b.w <= a.x + 2 or a.y + a.h <= b.y + 2 or b.y + b.h <= a.y + 2


def in_slide(placed) -> bool:
    return all(
        p.x >= -2 and p.y >= -2 and p.x + p.w <= SLIDE_W + 2 and p.y + p.h <= SLIDE_H + 2 for p in placed
    )


# ---- part 1: the icon list ------------------------------------------------------------------------


def iconlist(extra: str = "", attrs: str = "cols=1", items: str = ICONS, **layout):
    return lay(f"# Ba cách\n@iconlist {attrs}\n{items}", extra, **layout)


def test_iconlist_rows_use_the_body_with_grown_type():
    placed, _, theme = iconlist()
    rows = [p for p in placed if isinstance(p.element, Text) and p.element.role == "body"]
    assert len(rows) == 3
    body_pt = theme.sizes["body"]
    first = rows[0].element.paragraphs
    assert pt(rows[0]) >= body_pt * 1.4  # the title: ~24pt
    assert first[1].style is not None and first[1].style.font_size is not None
    assert first[1].style.font_size * rows[0].font_scale >= body_pt * 1.15  # the text: ~20pt
    top, bottom = min(p.y for p in rows), max(p.y + p.h for p in rows)
    assert bottom - top >= 0.8 * 3.6 * 914400  # the rows spread over the body, no empty lower half
    assert in_slide(placed)


def test_iconlist_title_and_text_max_tokens():
    big, _, _ = iconlist()
    small, _, _ = iconlist("iconlist.title.max=20 iconlist.text.max=15")
    b, s = texts(big, "body")[0], texts(small, "body")[0]
    assert pt(s) <= 20 + 0.01
    assert s.element.paragraphs[1].style.font_size * s.font_scale <= 15 + 0.01
    assert pt(b) > pt(s)
    # a cap under the role size never shrinks the text below it
    tiny, _, theme = iconlist("iconlist.title.max=4 iconlist.text.max=4")
    assert pt(texts(tiny, "body")[0]) >= theme.sizes["body"] - 0.01


def test_iconlist_fill_off_keeps_the_compact_look():
    on, _, _ = iconlist()
    off, _, _ = iconlist("iconlist.fill=off")
    assert max(p.h for p in texts(off, "body")) < max(p.h for p in texts(on, "body"))


def test_iconlist_columns_stack_the_icon_above_the_text():
    placed, _, _ = iconlist(attrs="cols=3 fill=surface")
    cards = texts(placed, "body")
    icons = [p for p in placed if isinstance(p.element, Shape) and p.element.shape == "icon"]
    assert len(cards) == 3 and len(icons) == 3
    for c, ic in zip(sorted(cards, key=lambda p: p.x), sorted(icons, key=lambda p: p.x), strict=True):
        assert inside(ic, c)  # the disc sits inside its card ...
        pad_top = float(c.style.padding_top.removesuffix("pt")) * 12700
        assert pad_top >= ic.y - c.y + ic.h - 2  # ... and the text starts under it
        assert ic.x < c.x + c.w * 0.5  # icon at the left, text below it
    assert all(disjoint(a, b) for a, b in zip(cards, cards[1:], strict=False))
    assert in_slide(placed)


def test_iconlist_icon_pos_token():
    left, _, _ = iconlist("iconlist.icon.pos=left", attrs="cols=3 fill=surface")
    c = texts(left, "body")[0]
    ic = [p for p in left if isinstance(p.element, Shape) and p.element.shape == "icon"][0]
    assert float(c.style.padding_top.removesuffix("pt")) * 12700 < ic.h  # beside the text: no icon height
    top, _, _ = iconlist("iconlist.icon.pos=top", attrs="cols=1 fill=surface")
    c = texts(top, "body")[0]
    assert float(c.style.padding_top.removesuffix("pt")) * 12700 > 1.0 * 914400 * 0.5


def test_iconlist_cards_stretch_to_their_rows():
    placed, _, _ = iconlist(attrs="cols=1 fill=surface")
    cards = sorted(texts(placed, "body"), key=lambda p: p.y)
    assert len({p.h for p in cards}) == 1
    assert all(a.y + a.h <= b.y for a, b in zip(cards, cards[1:], strict=False))
    assert cards[0].h > 1.2 * 914400  # taller than the compact row (a text of ~0.8in)


@pytest.mark.parametrize("n", [1, 2, 3, 4, 5, 6])
@pytest.mark.parametrize("attrs", ["", "cols=1", "cols=2", "cols=3", "cols=3 fill=surface", "fill=surface"])
def test_iconlist_in_bounds_for_every_count(n, attrs):
    items = "".join((ICONS * 2).splitlines(True)[:n])
    placed, deck, _ = iconlist(attrs=attrs, items=items)
    assert in_slide(placed)
    assert not [d for d in deck.diagnostics if d.rule == "overflow"]
    cards = texts(placed, "body")
    assert len(cards) == n
    assert all(disjoint(a, b) for i, a in enumerate(cards) for b in cards[i + 1 :])  # no two items overlap


# ---- part 2: box cards grow into the free height inside them -----------------------------------------


def card_parts(placed):
    cards = sorted((p for p in placed if isinstance(p.element, Container)), key=lambda p: (p.y, p.x))
    out = []
    for c in cards:
        kids = [p for p in placed if p is not c and inside(p, c)]
        icon = [p for p in kids if isinstance(p.element, Shape)]
        head = [p for p in kids if isinstance(p.element, Text) and p.element.role == "heading"]
        body = [p for p in kids if isinstance(p.element, Text) and p.element.role == "body"]
        out.append((c, icon[0] if icon else None, head[0], body))
    return out


def test_box_cards_with_free_height_grow_heading_text_and_disc():
    off, _, _ = lay(SHORT_CARDS, **{"card_fill_max": 0.0})
    on, _, theme = lay(SHORT_CARDS)
    for (c0, i0, h0, b0), (c1, i1, h1, b1) in zip(card_parts(off), card_parts(on), strict=True):
        assert (c0.x, c0.y, c0.w, c0.h) == (c1.x, c1.y, c1.w, c1.h)  # the cards keep their size
        assert pt(h1) > pt(h0) - 1e-6 and pt(h1) <= theme.layout.card_heading_max_pt + 0.01
        assert pt(b1[0]) >= pt(b0[0]) - 1e-6 and pt(b1[0]) <= theme.layout.card_body_max_pt + 0.01
        assert i1.w > i0.w and i1.w == i1.h  # the disc follows the heading
        assert abs(i1.w / i0.w - pt(h1) / pt(h0)) < 0.03
        assert h1.y == i1.y - (h1.h - i1.h) // 2 or abs(h1.y + h1.h // 2 - (i1.y + i1.h // 2)) <= 2
        for p in (i1, h1, *b1):
            assert inside(p, c1)
        assert h1.y + h1.h <= b1[0].y + 2  # the heading band never overlaps the text
        assert b1[0].y + b1[0].h <= c1.y + c1.h  # the text stays inside, top-anchored
        assert b1[0].y - c1.y < c1.h * 0.7
    assert pt(card_parts(on)[0][2]) > pt(card_parts(off)[0][2])


def test_box_card_tokens_cap_the_growth():
    base, _, _ = lay(SHORT_CARDS)
    capped, _, _ = lay(SHORT_CARDS, "card.heading.max=28")
    h0, h1 = card_parts(base)[0][2], card_parts(capped)[0][2]
    assert pt(h1) <= pt(h0) + 1e-6
    # a cap under the size the layout reached never shrinks it
    low, _, _ = lay(SHORT_CARDS, "card.heading.max=10 card.text.max=10")
    off, _, _ = lay(SHORT_CARDS, **{"card_fill_max": 0.0})
    for a, b in zip(card_parts(low), card_parts(off), strict=True):
        assert pt(a[2]) == pytest.approx(pt(b[2])) and pt(a[3][0]) == pytest.approx(pt(b[3][0]))
    # the deck-wide ceiling still rules
    flat, _, _ = lay(SHORT_CARDS, **{"grow_max": 1.0})
    assert pt(card_parts(flat)[0][2]) <= 22 + 0.01


def test_pinned_sizes_stay():
    for extra in ("sizes: body=16!",):
        placed, _, theme = lay(SHORT_CARDS.replace("@3\n", "@3\n" + extra + "\n", 1))
        assert all(pt(b[0]) == pytest.approx(16.0) for _c, _i, _h, b in card_parts(placed))
    pinned = SHORT_CARDS.replace("Phản hồi trong vài giây.", "[Phản hồi trong vài giây.]{size=14}")
    placed, _, _ = lay(pinned)
    off, _, _ = lay(pinned, **{"card_fill_max": 0.0})
    assert pt(card_parts(placed)[0][3][0]) == pytest.approx(pt(card_parts(off)[0][3][0]))
    assert pt(card_parts(placed)[0][2]) == pytest.approx(
        pt(card_parts(off)[0][2])
    )  # (the card is left as it is)
    placed, _, _ = lay(SHORT_CARDS.replace("@3\n", "@3\nsizes: heading=22!\n", 1))
    assert pt(card_parts(placed)[0][2]) == pytest.approx(22.0)


def test_a_full_card_does_not_grow():
    full = (
        "# Hai thẻ\n@2\n## A {icon=bolt}\n"
        + "- một ý khá dài để thẻ đầy chiều cao\n" * 6
        + "## B {icon=gear}\n- ngắn\n"
    )
    placed, _, _ = lay(full)
    off, _, _ = lay(full, **{"card_fill_max": 0.0})
    for a, b in zip(card_parts(placed), card_parts(off), strict=True):
        assert pt(a[2]) == pytest.approx(pt(b[2])) and pt(a[3][0]) == pytest.approx(pt(b[3][0]))


def test_box_cards_pptx_reopens_with_the_grown_sizes(tmp_path):
    src = HEAD.format(extra="") + SHORT_CARDS
    out = tmp_path / "d.pptx"
    build(src, out)
    prs = Presentation(str(out))
    sizes = {
        sh.name: max(r.font.size.pt for para in sh.text_frame.paragraphs for r in para.runs if r.font.size)
        for sh in prs.slides[0].shapes
        if sh.has_text_frame and sh.name.startswith(("Heading", "Text"))
    }
    assert sizes
    assert all(v >= 22 for k, v in sizes.items() if k.startswith("Heading"))
    for sh in prs.slides[0].shapes:
        assert sh.left >= 0 and sh.top >= 0 and sh.left + sh.width <= prs.slide_width + Emu(2)


def test_box_card_pass_never_raises():
    rng = random.Random(3)
    words = ["AI", "Tiết kiệm thời gian", "データ分析", "x" * 40, "", "một hai ba bốn năm sáu bảy tám"]
    for _ in range(40):
        n = rng.randint(2, 5)
        form = rng.choice(["@2", "@3", "@4", "@2x2", ""])
        parts = []
        for _ in range(n):
            ic = rng.choice(["", " {icon=bolt}", " {icon=nope}"])
            body = rng.choice(["- " + rng.choice(words), rng.choice(words), "", "- a\n- b\n  - c"])
            parts.append(f"## {rng.choice(words) or 'H'}{ic}\n{body}\n")
        src = HEAD.format(extra="") + f"# T\n{form}\n" + "".join(parts) + rng.choice(["", BAR])
        placed, _deck, _ = lay(src)
        assert placed is not None  # (overflow is a diagnostic, never a crash)
    assert fill_box_cards([], Rect(0, 0, 10, 10), None) is None  # type: ignore[arg-type]


# ---- part 3: the conclusion bar -----------------------------------------------------------------------


def bar_and_body(body: str, extra: str = ""):
    placed, deck, theme = lay(body + BAR, extra)
    bar = texts(placed, "conclusion")[0]
    sizes = [pt(p) for p in texts(placed, "body")]
    return placed, bar, max(sizes), deck, theme


PROS = "# Hai mặt\n@proscons\n## Lợi\n- Nhanh hơn nhiều\n- Rẻ hơn\n## Hại\n- Có thể sai\n- Cần kiểm tra\n"


def test_bar_is_never_under_the_body_text_and_wraps_first():
    placed, bar, body_pt, _, _ = bar_and_body(PROS)
    assert pt(bar) >= body_pt - 0.01
    assert bar.h > 1.5 * pt(bar) * 12700  # two lines: the bar grew taller instead of shrinking the text
    assert in_slide(placed)
    off, bar_off, _, _, _ = bar_and_body(PROS, "conclusion.floor=0")
    assert pt(bar_off) < pt(bar)  # (floor off: the old one-line shrink)


def test_bar_that_fits_one_line_stays_one_line_at_the_body_size():
    short = "# Hai mặt\n@proscons\n## Lợi\n- Nhanh\n## Hại\n- Sai\n"
    placed = lay(short + "> Dùng có kiểm soát\n")[0]
    bar = texts(placed, "conclusion")[0]
    body_pt = max(pt(p) for p in texts(placed, "body"))
    assert pt(bar) >= body_pt - 0.01
    assert bar.h < 1.9 * pt(bar) * 12700  # one line (text plus the insets)


def test_explicit_bar_size_wins():
    placed, bar, _, _, _ = bar_and_body(PROS, "sizes: conclusion=18")
    assert pt(bar) == pytest.approx(18.0)


def test_fit_line_reports_the_bar_size():
    src = HEAD.format(extra="") + PROS + BAR
    deck = parse(src)
    deck.attrs["base_dir"] = str(EX)
    theme, _ = deck_theme(deck, EX)
    placed = [layout_slide(s, deck, theme, i) for i, s in enumerate(deck.slides)]
    lines, _ = fit_report(deck, placed, theme)
    bar = texts(placed[0], "conclusion")[0]
    assert f"bar {round(pt(bar))}pt" in lines[0]


def test_example_decks_keep_the_bar_at_the_body_size():
    for name, n in (("24-vocabulary-2.md", 5), ("26-ai-overview.md", 12)):
        deck = parse((EX / name).read_text(encoding="utf-8"))
        deck.attrs["base_dir"] = str(EX)
        theme, _ = deck_theme(deck, EX)
        placed = layout_slide(deck.slides[n - 1], deck, theme, n - 1)
        bar = texts(placed, "conclusion")[0]
        body_pt = max(pt(p) for p in texts(placed, "body"))
        assert pt(bar) >= min(body_pt, theme.layout.conclusion_max_pt) - 0.01, name
        assert in_slide(placed)
