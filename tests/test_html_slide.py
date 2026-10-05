"""Whole-slide HTML: @html slides, deck.html sections, shared tokens, native fidelity (Chromium skips)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidemark.build import build
from slidemark.htmlnative import html_to_placed
from slidemark.ir import Image, Shape, Text
from slidemark.parser import parse, parse_html
from slidemark.render.htmlimg import HtmlRenderer
from slidemark.theme import Theme

PX = 9525
THEME = Theme(name="t", colors={"bg": "#FFFFFF", "fg": "#111111", "primary": "#7C5CFF"})


def md(html: str, extra: str = "") -> str:
    return f"{extra}# Launch\n@html\n```html\n{html}\n```\n"


# --------------------------------------------------------------------------- parser


def test_at_html_takes_fence_body():
    deck = parse(md('<div style="color:red">x</div>'))
    s = deck.slides[0]
    assert s.html == '<div style="color:red">x</div>'
    assert s.title.paragraphs[0].plain == "Launch"
    assert s.elements == [] and "html" not in s.classes
    assert not [d for d in deck.diagnostics if d.level != "info"]


def test_at_html_without_fence_warns():
    deck = parse("# T\n@html\n\n- a\n")
    assert deck.slides[0].html is None
    d = [x for x in deck.diagnostics if x.rule == "html-slide-empty"]
    assert d and d[0].hint


def test_at_html_drops_other_blocks_with_warning():
    deck = parse(md("<p>x</p>") + "\n- extra bullet\n")
    assert deck.slides[0].html == "<p>x</p>" and deck.slides[0].elements == []
    assert [d for d in deck.diagnostics if d.rule == "dropped-content"]


SECTION_CSS = "<style>.hero{background:linear-gradient(red,blue);padding:40px}</style>"


def test_deck_plain_section_stays_structural():
    deck = parse_html(
        "<style>h1{font-size:40px}</style><section><h1>T</h1><ul><li>a</li></ul>"
        '<p style="color:#333;text-align:center">x</p></section>'
    )
    s = deck.slides[0]
    assert s.html is None and s.title is not None and s.elements


def test_deck_styled_sections_become_html():
    deck = parse_html(
        SECTION_CSS
        + '<section class="hero"><h2>Big</h2></section>'
        + "<section><h1>Plain</h1><p>x</p></section>"
        + '<section><h1>Svg</h1><svg width="5"></svg></section>'
        + '<section><h1>Inline</h1><div style="border-radius:8px;background:#eee">y</div></section>'
        + '<section data-render="html"><h1>Forced</h1></section>'
    )
    flags = [s.html is not None for s in deck.slides]
    assert flags == [True, False, True, True, True]
    first = deck.slides[0]
    assert first.title.paragraphs[0].plain == "Big"
    assert first.html.startswith("<style>") and '<section class="hero">' in first.html
    assert "</section>" in first.html


def test_deck_native_attribute_forces_html():
    deck = parse_html(
        '<html data-slidemark="native"><body><section><h1>T</h1><p>x</p></section></body></html>'
    )
    assert deck.slides[0].html and "<section>" in deck.slides[0].html


# --------------------------------------------------------------------------- Chromium


@pytest.fixture(scope="module")
def chromium():
    pytest.importorskip("playwright.sync_api")
    r = HtmlRenderer()
    try:
        ok = r.render("<p>x</p>", 40, 20, THEME) is not None
    finally:
        r.close()
    if not ok:
        pytest.skip("Chromium cannot be launched")


def _place(html: str, w: int = 600, h: int = 300, theme: Theme = THEME):
    r = HtmlRenderer()
    try:
        res = html_to_placed(html, 0, 0, w * PX, h * PX, theme, ".", renderer=r)
    finally:
        r.close()
    assert res is not None
    return res


def _shapes(items):
    return [p for p in items if isinstance(p.element, Shape)]


def test_theme_tokens_in_css_vars(chromium):
    items, _ = _place('<div style="background:var(--primary);width:100px;height:50px"></div>')
    assert _shapes(items)[0].style.fill == "#7C5CFF"


def test_inline_tokens_reach_html_slide(chromium, tmp_path):
    src = "colors: primary=#123456\n\n" + md(
        '<div style="width:1280px;height:720px;background:var(--primary);color:#fff">Hi</div>'
    )
    out = tmp_path / "a.pptx"
    deck = build(src, out)
    assert deck.slides[0].html
    shp = Presentation(str(out)).slides[0].shapes
    fills = {
        str(s.fill.fore_color.rgb)
        for s in shp
        if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s.fill.type == 1
    }
    assert "123456" in fills
    assert not [s for s in shp if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_font_and_size_vars(chromium):
    items, _ = _place('<p style="font-family:var(--font-heading);font-size:var(--size-title)">T</p>')
    t = [p for p in items if isinstance(p.element, Text)][0]
    assert t.style.font_size == pytest.approx(32.0, abs=0.1)


def test_gradient_shadow_opacity_ellipse(chromium):
    items, _ = _place(
        '<div style="width:100px;height:60px;background:linear-gradient(90deg,#f00,#00f);'
        'box-shadow:0 8px 24px rgba(0,0,0,.4);opacity:.5"></div>'
        '<div style="width:80px;height:80px;border-radius:50%;background:#0a0"></div>'
    )
    a, b = _shapes(items)[:2]
    assert a.style.fill.startswith("linear-gradient(") and a.style.shadow
    assert a.style.shadow.split()[:3] == ["0", "6", "18"]
    assert a.style.opacity == pytest.approx(0.5, abs=0.01)
    assert b.element.shape == "ellipse"


def test_rotation_moves_children_with_the_card(chromium):
    items, _ = _place(
        '<div style="margin:100px;width:200px;height:100px;background:#eee;transform:rotate(-10deg)">'
        '<p style="margin:0">Hello</p></div>'
    )
    card = _shapes(items)[0]
    text = [p for p in items if isinstance(p.element, Text)][0]
    assert card.style.rotation == pytest.approx(350.0, abs=0.1)
    assert text.style.rotation == pytest.approx(350.0, abs=0.1)
    assert card.w == 200 * PX and card.h == 100 * PX  # unrotated size


def test_inline_pill_is_a_rounded_rect(chromium):
    items, _ = _place(
        '<p>Status <span style="background:#fdd;border-radius:10px;padding:2px 10px">live</span></p>'
    )
    pill = [p for p in _shapes(items) if p.style.fill == "#FFDDDD"]
    assert pill and pill[0].element.shape == "rounded-rect"


def test_svg_shapes_become_native(chromium):
    items, diags = _place(
        '<svg width="300" height="100"><rect x="0" y="0" width="100" height="50" rx="8" fill="#1d4ed8"/>'
        '<circle cx="150" cy="25" r="20" fill="none" stroke="#0f766e" stroke-width="3"/>'
        '<line x1="0" y1="80" x2="200" y2="80" stroke="#333" stroke-width="2"/>'
        '<polygon points="200,70 210,80 200,90" fill="#333"/>'
        '<text x="10" y="40" fill="#fff" font-size="20">Hi</text></svg>'
    )
    kinds = [(p.element.shape if isinstance(p.element, Shape) else "text") for p in items]
    assert kinds == ["rounded-rect", "ellipse", "line", "triangle", "text"]
    assert not any(isinstance(p.element, Image) for p in items) and not diags


def test_unmappable_element_is_a_picture_with_info(chromium):
    items, diags = _place(
        '<p>keep</p><div style="filter:blur(2px);width:80px;height:40px;background:#f0f"></div>'
    )
    pics = [p for p in items if isinstance(p.element, Image)]
    assert len(pics) == 1 and "filter" in pics[0].element.alt
    assert [d for d in diags if d.rule == "html-element-image" and "filter" in d.message]
    assert [p for p in items if isinstance(p.element, Text)]  # the rest stays native


def test_html_slide_builds_native_and_reopens(chromium, tmp_path):
    out = tmp_path / "b.pptx"
    body = (
        '<div style="width:1280px;height:720px;background:linear-gradient(135deg,var(--primary),#00D1B2);'
        'padding:96px;color:#fff"><h1 style="font-size:72px;margin:0">Meet Aurora</h1></div>'
    )
    deck = build(md(body), out)
    assert not [d for d in deck.diagnostics if d.level == "error"]
    slide = Presentation(str(out)).slides[0]
    assert any(s.has_text_frame and "Meet Aurora" in s.text_frame.text for s in slide.shapes)
    assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]


def test_not_convertible_slide_is_one_picture(chromium, tmp_path):
    out = tmp_path / "c.pptx"
    deck = build(md('<canvas width="100" height="50"></canvas>'), out)
    assert [d for d in deck.diagnostics if d.rule == "html-slide-image"]
    pics = [s for s in Presentation(str(out)).slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1


def test_deck_html_file_mixed_build(chromium, tmp_path):
    src = tmp_path / "deck.html"
    src.write_text(
        "<style>.hero{background:#123;color:#fff;padding:40px;width:1280px;height:720px}</style>"
        '<section class="hero"><h1>Styled</h1></section>'
        "<section><h1>Plain</h1><ul><li>one</li></ul></section>",
        encoding="utf-8",
    )
    out = tmp_path / "d.pptx"
    deck = build(src, out)
    assert [s.html is not None for s in deck.slides] == [True, False]
    assert len(Presentation(str(out)).slides) == 2
