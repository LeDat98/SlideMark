"""``.pptx`` -> SlideMark text: round trip of the examples, a hand-made foreign deck, CLI."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from slidemark.build import build
from slidemark.cli import main
from slidemark.importer import import_pptx
from slidemark.ir import Chart, Container, Deck, Table
from slidemark.parser import parse

EXAMPLES = sorted((Path(__file__).parent.parent / "examples").glob("*.md"))


def _plain(paras) -> str:
    return " ".join(p.plain for p in paras).strip()


def _walk(els, boxes: list, tables: list, charts: list) -> None:
    for el in els:
        if isinstance(el, Container):
            if el.title is not None:
                boxes.append(_plain(el.title.paragraphs))
            _walk(el.children, boxes, tables, charts)
        elif isinstance(el, Table):
            tables.append([[_plain(c.paragraphs) for c in row] for row in el.rows])
        elif isinstance(el, Chart):
            series = [(s.name, [None if v is None else round(v, 6) for v in s.values]) for s in el.series]
            charts.append((el.kind, series))


def signature(deck: Deck) -> dict:
    """What must survive a round trip: titles, box headings, table text, chart values, notes."""
    sig: dict = {"n": len(deck.slides), "slides": []}
    for sl in deck.slides:
        boxes: list[str] = []
        tables: list = []
        charts: list = []
        _walk(sl.elements, boxes, tables, charts)
        sig["slides"].append(
            {
                "title": _plain(sl.title.paragraphs) if sl.title else "",
                "boxes": boxes,
                "tables": tables,
                "charts": charts,
                "notes": (sl.notes or "").strip(),
            }
        )
    return sig


def _roundtrip(tmp_path: Path, source: str):
    pptx1 = tmp_path / "a.pptx"
    build(source, pptx1)
    text1, d1 = import_pptx(pptx1, tmp_path)
    pptx2 = tmp_path / "b.pptx"
    build(text1, pptx2)
    text2, d2 = import_pptx(pptx2, tmp_path)
    return text1, text2, d1, d2


@pytest.mark.parametrize("path", EXAMPLES, ids=lambda p: p.stem)
def test_roundtrip_examples(path: Path, tmp_path: Path):
    if "html" in path.stem:
        pytest.importorskip("playwright")  # html fences come back from the design part; measuring needs Chromium
    src = path.read_text(encoding="utf-8")
    text1, text2, d1, _ = _roundtrip(tmp_path, src)
    assert not [d for d in d1 if d.level in ("error", "warning")]
    assert text1 == text2, "second import must equal the first"
    orig, imported = parse(src), parse(text1)
    assert not [d for d in imported.diagnostics if d.level in ("error", "warning")], imported.diagnostics
    a, b = signature(orig), signature(imported)
    assert a["n"] == b["n"]
    for i, (sa, sb) in enumerate(zip(a["slides"], b["slides"], strict=True), 1):
        assert sa == sb, f"slide {i} differs"


def test_header_keys(tmp_path: Path):
    text, _, _, _ = _roundtrip(
        tmp_path, "theme: jp-business\nlang: ja\nfooter: ACME\nnum: on\n\n# 題名\n- 項目\n"
    )
    head = text.split("\n\n")[0].splitlines()
    assert head == ["theme: jp-business", "lang: ja", "footer: ACME", "num: on"]


def test_hidden_and_notes(tmp_path: Path):
    text, _, _, _ = _roundtrip(tmp_path, "# A\n- x\n\n# B\n@hidden\n- y\n??? say this\nand that\n")
    deck = parse(text)
    assert deck.slides[1].hidden and deck.slides[1].notes == "say this\nand that"


def test_picture_exported(tmp_path: Path):
    from PIL import Image

    Image.new("RGB", (8, 8), "red").save(tmp_path / "p.png")
    pptx = tmp_path / "pic.pptx"
    build("# P\n![a red square](p.png)\n", pptx, base_dir=tmp_path)
    out = tmp_path / "o"
    text, diags = import_pptx(pptx, out)
    assert "![a red square](images/1-1.png)" in text
    assert (out / "images" / "1-1.png").is_file()
    _, diags2 = import_pptx(pptx)
    assert any(d.rule == "import-image" for d in diags2)


def _foreign(path: Path) -> None:
    prs = Presentation()  # 4:3 default template, python-pptx layouts
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap"
    tf = s.placeholders[1].text_frame
    tf.text = "First point"
    p = tf.add_paragraph()
    p.text = "Nested detail"
    p.level = 1
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "Bold and link"
    r.font.bold = True
    r.hyperlink.address = "https://example.com"
    s.notes_slide.notes_text_frame.text = "speaker words"

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Numbers"
    gf = s.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2))
    for r_i, row in enumerate([("Region", "Sales"), ("APAC", "8.2"), ("EU", "6.1")]):
        for c_i, v in enumerate(row):
            gf.table.cell(r_i, c_i).text = v

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Trend"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("2025", (10, 12, 15))
    cd.add_series("2026", (12, 16, 21))
    gfc = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(7), Inches(4), cd)
    gfc.chart.has_title = True
    gfc.chart.chart_title.text_frame.text = "Revenue"

    s = prs.slides.add_slide(prs.slide_layouts[3])  # two content
    s.shapes.title.text = "Compare"
    s.placeholders[1].text_frame.text = "Left side"
    s.placeholders[2].text_frame.text = "Right side"
    prs.save(str(path))


def test_foreign_deck(tmp_path: Path):
    f = tmp_path / "foreign.pptx"
    _foreign(f)
    text, diags = import_pptx(f, tmp_path)
    assert not [d for d in diags if d.level in ("error", "warning")]
    deck = parse(text)
    assert deck.diagnostics == [], deck.diagnostics
    assert deck.size == "4:3"
    assert [_plain(s.title.paragraphs) for s in deck.slides] == ["Roadmap", "Numbers", "Trend", "Compare"]
    s1 = deck.slides[0]
    paras = [p for el in s1.elements for p in el.paragraphs]
    assert [p.plain for p in paras] == ["First point", "Nested detail", "Bold and link"]
    assert [p.level for p in paras] == [0, 1, 0] and all(p.marker == "bullet" for p in paras)
    assert paras[2].runs[0].bold and paras[2].runs[0].link == "https://example.com"
    assert s1.notes == "speaker words"
    tbl = next(el for el in deck.slides[1].elements if isinstance(el, Table))
    assert [[_plain(c.paragraphs) for c in r] for r in tbl.rows] == [
        ["Region", "Sales"],
        ["APAC", "8.2"],
        ["EU", "6.1"],
    ]
    ch = next(el for el in deck.slides[2].elements if isinstance(el, Chart))
    assert ch.kind == "column" and ch.title == "Revenue" and ch.categories == ["Q1", "Q2", "Q3"]
    assert [(s.name, s.values) for s in ch.series] == [("2025", [10, 12, 15]), ("2026", [12, 16, 21])]
    assert len(deck.slides[3].elements) == 2  # two columns, no @ line needed
    assert "\n@" not in text.split("# Compare")[1]


def test_merged_cells_and_inline(tmp_path: Path):
    src = "# T\n| a | b | c |\n|-|-|-|\n| x | y | z |\n| w | < | ^ |\n\nH~2~O x^2^ ~~old~~ **b** *i*\n"
    text, _, _, _ = _roundtrip(tmp_path, src)
    orig = parse(src)
    imp = parse(text)
    ta = next(e for e in orig.slides[0].elements if isinstance(e, Table))
    tb = next(e for e in imp.slides[0].elements if isinstance(e, Table))
    assert [(c.colspan, c.rowspan) for r in ta.rows for c in r] == [
        (c.colspan, c.rowspan) for r in tb.rows for c in r
    ]
    runs = [
        r for e in imp.slides[0].elements if hasattr(e, "paragraphs") for p in e.paragraphs for r in p.runs
    ]
    assert any(r.sub for r in runs) and any(r.sup for r in runs) and any(r.strike for r in runs)


def test_cli_import(tmp_path: Path, capsys):
    f = tmp_path / "foreign.pptx"
    _foreign(f)
    assert main(["import", str(f)]) == 0
    assert "# Roadmap" in capsys.readouterr().out
    out = tmp_path / "out" / "deck.md"
    assert main(["import", str(f), "-o", str(out)]) == 0
    assert out.read_text(encoding="utf-8").startswith("size: 4:3")


def test_cli_import_unreadable(tmp_path: Path, capsys):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"not a zip")
    assert main(["import", str(bad)]) == 2
    assert main(["import", str(tmp_path / "missing.pptx")]) == 2
    assert "error" in capsys.readouterr().err


def test_never_raises_on_empty_deck(tmp_path: Path):
    f = tmp_path / "empty.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(f))
    text, diags = import_pptx(f)
    assert parse(text).slides and not [d for d in diags if d.level == "error"]


def _cards_deck(path: Path) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Plan"
    lead = s.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(9), Inches(0.4))
    lead.text_frame.text = "Three steps to ship"
    lead.text_frame.paragraphs[0].runs[0].font.size = Inches(0.22)
    for i, (head, body) in enumerate([("Plan", "scope"), ("Build", "code"), ("Ship", "release")]):
        x = Inches(0.5 + i * 3.1)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, Inches(2), Inches(2.9), Inches(3))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(0xEE, 0xEE, 0xEE)
        h = s.shapes.add_textbox(x + Inches(0.1), Inches(2.1), Inches(2.7), Inches(0.5))
        h.text_frame.text = head
        h.text_frame.paragraphs[0].runs[0].font.bold = True
        b = s.shapes.add_textbox(x + Inches(0.1), Inches(2.8), Inches(2.7), Inches(2))
        b.text_frame.text = body
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(5.6), Inches(9), Inches(0.6))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0x1E, 0x3A, 0x5F)
    bar.text_frame.text = "Ship by Friday"
    grp = s.shapes.add_group_shape()
    note = grp.shapes.add_textbox(Inches(0.5), Inches(6.9), Inches(9), Inches(0.3))
    note.text_frame.text = "Source: internal"
    note.text_frame.paragraphs[0].runs[0].font.size = Inches(0.12)
    prs.save(str(path))


def test_foreign_cards(tmp_path: Path):
    f = tmp_path / "cards.pptx"
    _cards_deck(f)
    text, diags = import_pptx(f)
    deck = parse(text)
    assert deck.diagnostics == [], (text, deck.diagnostics)
    sl = deck.slides[0]
    assert [_plain(e.title.paragraphs) for e in sl.elements if isinstance(e, Container)] == [
        "Plan",
        "Build",
        "Ship",
    ]
    assert sl.lead and sl.lead.paragraphs[0].plain == "Three steps to ship"
    assert sl.conclusion and sl.conclusion.paragraphs[0].plain == "Ship by Friday"
    assert [f.paragraphs[0].plain for f in sl.footnotes] == ["Source: internal"]
    assert "@" not in text  # three boxes: automatic arrangement


@pytest.mark.parametrize(
    "kind",
    ["bar", "column", "stacked-bar", "stacked-column", "line", "area", "pie", "doughnut", "scatter", "radar"],
)
def test_chart_kinds(kind: str, tmp_path: Path):
    csv = ",1,2,3\nA,4,5,6\nB,7,8,9" if kind not in ("pie", "doughnut") else ",x,y,z\nshare,4,5,6"
    src = f'# C\n```{kind} {{title="T" labels=on fmt=0.0 legend=right}}\n{csv}\n```\n'
    text, _, _, _ = _roundtrip(tmp_path, src)
    a = next(e for e in parse(src).slides[0].elements if isinstance(e, Chart))
    b = next(e for e in parse(text).slides[0].elements if isinstance(e, Chart))
    assert (a.kind, a.title, a.categories) == (b.kind, b.title, b.categories)
    assert [(s.name, s.values) for s in a.series] == [(s.name, s.values) for s in b.series]
    assert b.options.get("legend") == "right"
    if kind != "scatter":  # the renderer draws no data labels on XY charts
        assert b.options.get("labels") == a.options.get("labels")


def test_box_color_classes_survive_import(tmp_path):
    src = (
        "theme: jp-business\nlang: ja\n\n# 課題\n@3\n## 市場 {.muted}\n- 縮小\n- 競争\n"
        "## 課題 {.danger}\n- 粗利率\n## 方針 {.primary}\n1. 投入\n"
    )
    build(src, tmp_path / "a.pptx")
    md, _ = import_pptx(tmp_path / "a.pptx")
    assert "## 市場 {.muted}" in md and "## 課題 {.danger}" in md and "## 方針 {.primary}" in md
