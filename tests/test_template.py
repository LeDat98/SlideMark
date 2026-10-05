"""User templates: theme resolution from .pptx/.potx/.yaml and rendering on the template's layouts."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from slidemark.build import build
from slidemark.template import footer_top, resolve_theme, template_size
from slidemark.theme import get_theme

DECK = """---
title: T
footer: Acme
num: on
---

# Cover

---

# Content slide

- one
- two

---

# Section
"""


def _make_template(path: Path, *, potx: bool = False, with_slide: bool = True) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = 9144000, 6858000
    if with_slide:
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "left over"
    tmp = path.with_suffix(".tmp.pptx")
    prs.save(tmp)
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "ppt/theme/theme1.xml":
                x = data.decode("utf-8")
                for k, v in (("accent1", "AA0011"), ("accent2", "00AA22"), ("dk2", "334455")):
                    x = re.sub(rf"(<a:{k}>).*?(</a:{k}>)", rf'\1<a:srgbClr val="{v}"/>\2', x, flags=re.S)
                x = re.sub(r'(<a:majorFont><a:latin typeface=")[^"]*', r"\1Georgia", x)
                x = re.sub(r'(<a:minorFont><a:latin typeface=")[^"]*', r"\1Verdana", x)
                data = x.encode("utf-8")
            if potx and item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main", b"presentationml.template.main")
            zout.writestr(item, data)
    tmp.unlink()
    return path


@pytest.fixture
def tpl(tmp_path):
    return _make_template(tmp_path / "corp.pptx")


def test_resolve_builtin():
    theme, diags = resolve_theme("midnight", ".")
    assert theme.name == "midnight" and not diags


def test_resolve_pptx_colors_fonts_size(tpl, tmp_path):
    theme, diags = resolve_theme("corp.pptx", tmp_path)
    assert not diags
    assert theme.colors["primary"] == "#AA0011"
    assert theme.colors["secondary"] == "#00AA22"
    assert theme.colors["muted"] == "#334455"
    assert theme.fonts.heading == "Georgia" and theme.fonts.body == "Verdana"
    assert theme.template == str(tpl.resolve())
    assert template_size(theme) == "9144000emux6858000emu"


def test_resolve_yaml(tmp_path):
    (tmp_path / "t.yaml").write_text(
        "colors:\n  primary: '#123456'\nfonts:\n  body: Arial\n", encoding="utf-8"
    )
    theme, diags = resolve_theme("t.yaml", tmp_path)
    assert not diags and theme.colors["primary"] == "#123456" and theme.colors["fg"] == "#1F2937"
    assert theme.fonts.body == "Arial"


def test_yaml_validation_error(tmp_path):
    (tmp_path / "t.yaml").write_text("nonsense_field: 1\n", encoding="utf-8")
    theme, diags = resolve_theme("t.yaml", tmp_path)
    assert theme.name == "default" and diags[0].rule == "bad-theme" and diags[0].hint


@pytest.mark.parametrize("name", ["nope.pptx", "nope.yaml", "nope"])
def test_missing_file_falls_back(tmp_path, name):
    theme, diags = resolve_theme(name, tmp_path)
    assert theme.name == "default" and theme.template is None
    assert [d.rule for d in diags] == ["bad-theme"] and diags[0].hint


def test_corrupt_pptx(tmp_path):
    (tmp_path / "bad.pptx").write_bytes(b"not a zip")
    theme, diags = resolve_theme("bad.pptx", tmp_path)
    assert theme.name == "default" and diags[0].rule == "bad-theme"


def _build_with(theme_name: str, tmp_path: Path):
    src = tmp_path / "d.md"
    src.write_text(DECK.replace("title: T", f"title: T\ntheme: {theme_name}"), encoding="utf-8")
    out = tmp_path / "out.pptx"
    # the orchestrator's build.py does the same: resolve_theme + deck.size from the template
    deck = build(src, out)
    return deck, out


def test_render_uses_template(tpl, tmp_path):
    deck, out = _build_with(str(tpl), tmp_path)
    prs = Presentation(str(out))
    assert prs.slide_width == 9144000
    assert len(prs.slides) == 3  # the template's own slide is gone
    names = [s.slide_layout.name for s in prs.slides]
    assert names == ["Title Slide", "Title Only", "Section Header"]
    s2 = prs.slides[1]
    assert s2.shapes.title.text_frame.text == "Content slide"
    assert not any(
        sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text for sh in s2.placeholders
    )
    xml = s2._element.xml
    assert 'type="slidenum"' in xml and 'type="ftr"' in xml and "Acme" in xml
    assert "left over" not in "".join(
        sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame
    )


def test_potx_accepted(tmp_path):
    from slidemark.template import open_template

    p = _make_template(tmp_path / "t.potx", potx=True)
    prs = open_template(p)
    assert len(prs.slides) == 0 and len(prs.slide_layouts) >= 6


def test_no_template_unchanged(tmp_path):
    out = tmp_path / "o.pptx"
    build(DECK, out)
    prs = Presentation(str(out))
    assert [s.slide_layout.name for s in prs.slides][1] == "Title Only"
    assert get_theme("default").template is None


def test_unknown_layout_names_fallback(tmp_path):
    """Layouts renamed beyond recognition still get a title-bearing layout."""
    from slidemark.template import pick_layout

    prs = Presentation()
    for lo in prs.slide_layouts:
        lo._element.attrib.pop("type", None)
        lo._element.cSld.set("name", "X " + lo.name[::-1])
    assert any(
        ph.get("type") in ("title", "ctrTitle") for ph in pick_layout(prs, "title")._element.iter("{*}ph")
    )


def test_footer_top_reads_the_template_footer_zone(tpl, tmp_path):
    theme, _ = resolve_theme("corp.pptx", tmp_path)
    ft = footer_top(theme)
    assert ft is not None and 0.85 * 6858000 < ft < 6858000
    assert footer_top(get_theme("default")) is None


def test_content_stays_above_the_template_footer_zone(tpl, tmp_path):
    from slidemark.ir import Deck, Paragraph, Run, Slide, Text
    from slidemark.layout import layout_slide

    theme, _ = resolve_theme("corp.pptx", tmp_path)
    txt = lambda s, role: Text(role=role, paragraphs=[Paragraph(runs=[Run(text=s)])])  # noqa: E731
    slide = Slide(
        title=txt("t", "title"),
        elements=[txt("body", "body")],
        footnotes=[txt("note", "footnote")],
        conclusion=txt("so what", "conclusion"),
    )
    deck = Deck(slides=[slide], footer="Acme", slide_number=True, size="4:3")
    placed = layout_slide(slide, deck, theme, 0)
    ft = footer_top(theme)
    for p in placed:
        if isinstance(p.element, Text) and p.element.attrs.get("field"):
            continue  # the footer text itself (replaced by the template placeholder)
        assert p.y + p.h <= ft, (p.element.role, p.y + p.h, ft)
    note = next(p for p in placed if p.element.role == "footnote")
    assert note.y + note.h < ft  # a small gap, not touching
