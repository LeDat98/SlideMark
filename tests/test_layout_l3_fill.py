"""L3 fill: sparse rows of text cards and text panels beside a chart use the height they have."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from slidemark import build
from slidemark.ir import Chart, Container, Text
from slidemark.layout import engine, layout_slide
from slidemark.layout.grid import Rect
from slidemark.layout.l3fill import fill_panels, fill_row
from slidemark.parser import parse
from slidemark.template import resolve_theme
from slidemark.theme import LayoutTokens

EX = Path(__file__).resolve().parent.parent / "examples"


def lay(name: str, n: int, src: str | None = None, **tokens):
    text = src if src is not None else (EX / name).read_text(encoding="utf-8")
    deck = parse(text)
    theme, _ = resolve_theme(deck.theme, EX)
    if tokens:
        theme = theme.model_copy(update={"layout": theme.layout.model_copy(update=tokens)})
    placed = layout_slide(deck.slides[n - 1], deck, theme, n - 1)
    return placed, deck, theme


def role(placed, r):
    return [p for p in placed if isinstance(p.element, Text) and p.element.role == r]


def cards(placed):
    return [p for p in placed if isinstance(p.element, Container)]


def inside(placed, card, r="body"):
    return [
        p
        for p in role(placed, r)
        if card.x <= p.x and p.x + p.w <= card.x + card.w and card.y <= p.y <= card.y + card.h
    ]


def tail(p, card) -> int:
    return card.y + card.h - (p.y + engine._text_h(p))


@pytest.mark.parametrize(("name", "n"), [("16-jp-strategy.md", 10), ("11-jp-consulting.md", 9)])
def test_decision_row_is_top_anchored_tall_and_spread(name, n):
    placed, deck, _ = lay(name, n)
    H = 720 * 9525
    lead = role(placed, "lead")[0]
    foot = role(placed, "footnote")[0]
    cs = cards(placed)
    assert len(cs) == 2
    assert min(c.y for c in cs) - (lead.y + lead.h) <= 0.05 * H  # no empty band under the lead
    assert foot.y - max(c.y + c.h for c in cs) <= 0.12 * H  # cards reach down to the footnote band
    assert len({c.h for c in cs}) == 1 and len({c.y for c in cs}) == 1  # one shared height
    lt = LayoutTokens()
    off, _, _ = lay(name, n, l3_fill=False)
    for c, c0 in zip(cs, cards(off), strict=True):
        (body,) = inside(placed, c)
        (body0,) = inside(off, c0)
        extra = body.element.attrs["para_gap"] - (body0.element.attrs.get("para_gap") or 0.25)
        assert 0 < extra <= lt.l3_gap_extra_numbered + 0.02  # even, capped gaps
        # the text is distributed, not piled up: either it reaches the bottom or the cap was hit
        cap = lt.l3_gap_extra_numbered if body.element.paragraphs[0].marker == "number" else lt.l3_gap_extra
        assert tail(body, c) <= 0.15 * c.h or extra >= cap - 0.1
        assert body.y + body.h <= c.y + c.h  # never outside its card
        assert body.style.valign in (None, "top")  # never vertically centred
    assert not [d for d in deck.diagnostics if d.level == "error"]


@pytest.mark.parametrize(("name", "n"), [("16-jp-strategy.md", 3), ("11-jp-consulting.md", 3)])
def test_panel_beside_chart_spreads_and_anchors_its_note(name, n):
    placed, _, _ = lay(name, n)
    off, _, _ = lay(name, n, l3_fill=False)
    (panel,) = cards(placed)
    chart = next(p for p in placed if isinstance(p.element, Chart))
    assert panel.h == chart.h
    main, note = sorted(inside(placed, panel), key=lambda p: p.y)
    assert abs((panel.y + panel.h) - (note.y + note.h)) <= 0.03 * panel.h  # note on the panel bottom
    g0 = max(p.element.attrs.get("para_gap") or 0.25 for p in inside(off, cards(off)[0]))
    assert 0 < main.element.attrs["para_gap"] - g0 <= LayoutTokens().l3_gap_extra + 0.02
    assert [engine._text_h(p) for p in inside(off, cards(off)[0])][0] < 0.5 * panel.h  # it was the bug
    assert main.y + engine._text_h(main) < note.y  # no overlap


def test_render_reopens_with_tall_cards(tmp_path):
    out = tmp_path / "d.pptx"
    build(EX / "16-jp-strategy.md", out)
    prs = Presentation(out)
    sl = prs.slides[9]
    tall = [
        s for s in sl.shapes if s.height > 0.5 * prs.slide_height * 0.6 and s.width < prs.slide_width * 0.7
    ]
    assert len(tall) >= 2
    top = min(s.top for s in tall)
    assert top < 0.25 * prs.slide_height  # right under the lead
    assert max(s.top + s.height for s in tall) > 0.78 * prs.slide_height
    paras = [s for s in sl.shapes if s.has_text_frame and "2027" in s.text_frame.text]
    assert paras and any(
        p.space_before and p.space_before.pt > 8 for sh in paras for p in sh.text_frame.paragraphs[1:]
    )


def test_tokens_switch_it_off():
    off, _, _ = lay("16-jp-strategy.md", 10, l3_fill=False)
    assert cards(off)[0].y > 200 * 9525  # the old floating block
    top, _, _ = lay("16-jp-strategy.md", 10, body_valign="top")
    assert cards(top)[0].h < 300 * 9525  # an explicit top anchor keeps the hugging cards


def test_explicit_height_valign_and_top_win():
    base = "# T\n> lead\n@1:1\n## A\n- one\n- two\n- three\n## B\n- x\n- y\n- z\n"
    placed, deck, theme = lay("", 1, base)
    cs = cards(placed)
    plain_h = cs[0].h
    body = Rect(0, 0, 10, 10)
    assert fill_row([], body, theme.layout) == []  # nothing to do
    assert fill_panels([], Rect(0, 0, 0, 0), theme.layout) == []
    explicit = base.replace("## A", "## A {h=2in}")
    p2, _, _ = lay("", 1, explicit)
    assert not any(c.element.attrs.get("para_gap", 0) > 1.0 for c in p2 if isinstance(c.element, Text))
    assert plain_h > 0


def test_fuzz_never_raises():
    for src in (
        "# T\n## A\n",
        "# T\n> l\n## A\n## B\n- x\n",
        "# T\n```column\n,a\nx,1\n```\n## P\n> [!note] n\n",
    ):
        placed, deck, _ = lay("", 1, src)
        assert not [d for d in deck.diagnostics if d.rule == "layout-error"]


KPI = """theme: none
colors: bg=#FFFFFF fg=#1F2933 primary=#1E3A5F accent=#C8102E surface=#F3F5F8 border=#C9D1DB
lang: ja

# 再編の効果（2029年度）
## 営業利益 {.kpi}
+24億円
年間
## 赤字店舗 {.kpi}
45→8店
-37店
## ネット売上比率 {.kpi}
25%
+13pt
"""


def test_kpi_row_is_top_anchored_and_values_stay_centred():
    placed, _, _ = lay("", 1, KPI)
    title = role(placed, "title")[0]
    cs = cards(placed)
    assert len(cs) == 3 and len({c.h for c in cs}) == 1
    assert min(c.y for c in cs) - (title.y + title.h) <= 0.08 * 720 * 9525  # right under the title
    assert 0.78 * 720 * 9525 <= max(c.y + c.h for c in cs) <= 0.9 * 720 * 9525  # leftover = bottom band only
    for c in cs:
        (body,) = inside(placed, c)
        assert body.style.valign == "middle"  # the value is centred in its taller card
        assert body.y + body.h <= c.y + c.h
