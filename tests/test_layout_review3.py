"""Review pass 3: boxes beside visuals, top-aligned media, roomy sparse slides, table width cap."""

from __future__ import annotations

from pptx import Presentation

from slidemark.build import build_deck
from slidemark.ir import Chart, Container, Deck, Image, Link, Media, Series, Slide, Table
from slidemark.layout import layout_slide
from slidemark.layout.engine import TABLE_GROW_ROOMY
from slidemark.parser import parse
from slidemark.theme import get_theme
from slidemark.units import EMU_PER_INCH, slide_size, to_emu

from .test_layout_policies import T, box, bullets, cell, of

W, H = slide_size("16:9")


def lay(slide, theme="jp-business"):
    deck = Deck(slides=[slide])
    return layout_slide(slide, deck, get_theme(theme), 0), deck


def chart():
    return Chart(kind="column", categories=["a", "b"], series=[Series(name="s", values=[1, 2])])


def full_width(theme="jp-business"):
    return W - 2 * to_emu(get_theme(theme).margin_x)


def overflow(deck):
    return [d for d in deck.diagnostics if d.rule == "overflow"]


# ---- 1/4: a text box beside a visual takes its natural height, top-aligned


def test_box_beside_chart_takes_natural_height():
    s = Slide(title=T("t", "title"), grid="1:1", elements=[chart(), box("論点", "a", "b", "c")])
    placed, deck = lay(s)
    (ch,) = of(placed, Chart)
    (bx,) = of(placed, Container)
    assert bx.y == ch.y  # tops line up
    assert bx.h < 0.75 * ch.h  # not stretched over the whole chart
    assert bx.h >= 0.4 * ch.h  # ... but not tiny either
    assert not overflow(deck)


def test_long_box_beside_chart_fills_the_height():
    items = [f"item {i} " + "word " * 12 for i in range(14)]
    s = Slide(title=T("t", "title"), grid="1:1", elements=[chart(), box("論点", *items)])
    placed, _ = lay(s)
    (ch,) = of(placed, Chart)
    (bx,) = of(placed, Container)
    assert bx.y == ch.y and bx.h <= ch.h  # content decides, never beyond the visual


def test_box_beside_image_takes_natural_height():
    s = Slide(title=T("t", "title"), grid="1:1", elements=[Image(src="a.png"), box("Sites", "a", "b")])
    placed, _ = lay(s, "default")
    (im,) = of(placed, Image)
    (bx,) = of(placed, Container)
    assert bx.y == im.y and bx.h < im.h


# ---- 2: pictures / videos beside text are top-aligned


def test_media_beside_text_is_top_aligned():
    s = Slide(title=T("t", "title"), elements=[bullets("one", "two"), Media(src="demo.mp4", kind="video")])
    placed, _ = lay(s, "midnight")
    (md,) = of(placed, Media)
    assert md.style.valign == "top"


def test_lone_media_stays_centered():
    s = Slide(title=T("t", "title"), elements=[Media(src="demo.mp4", kind="video")])
    placed, _ = lay(s, "midnight")
    (md,) = of(placed, Media)
    assert md.style.valign is None


def test_video_top_matches_text_top_in_the_pptx(tmp_path):
    (tmp_path / "demo.mp4").write_bytes(b"\x00\x00\x00\x18ftypmp42fake")
    out = tmp_path / "v.pptx"
    deck = parse("# T\n- one\n- two\n\n![demo](demo.mp4)\n")
    build_deck(deck, out, base_dir=tmp_path)
    slide = Presentation(str(out)).slides[0]
    vid = next(sh for sh in slide.shapes if sh._element.xpath(".//a:videoFile"))
    txt = next(sh for sh in slide.shapes if sh.has_text_frame and "one" in sh.text_frame.text)
    assert abs(vid.top - txt.top) < 0.05 * EMU_PER_INCH


# ---- 3: sparse dense slides grow tables / trees modestly and never overflow


def _plain_table(nrows=4):
    rows = [[cell("項目"), cell("内容")]] + [[cell("あ"), cell("い")] for _ in range(nrows - 1)]
    return Table(rows=rows)


def test_roomy_table_rows_grow_up_to_2x_and_text_up_to_1_2x():
    s = Slide(title=T("t", "title"), elements=[_plain_table()])
    placed, deck = lay(s)
    (tp,) = of(placed, Table)
    assert tp.font_scale <= 1.2 + 1e-9
    nat = 1.2 * tp.style.font_size * tp.font_scale * 12700 + 2 * 45720
    assert max(tp.element.attrs["_row_h"]) <= TABLE_GROW_ROOMY * nat * 1.15
    assert sum(tp.element.attrs["_row_h"]) == tp.h
    assert not overflow(deck)


def test_busy_slide_is_not_stretched_past_the_body():
    rows = [[cell("a"), cell("b")]] + [[cell("x"), cell("y")] for _ in range(14)]
    s = Slide(title=T("t", "title"), elements=[Table(rows=rows)])
    placed, deck = lay(s)
    (tp,) = of(placed, Table)
    assert tp.y + tp.h <= H - to_emu(get_theme("jp-business").margin_y)
    assert not overflow(deck)


def test_dense_tree_boxes_stay_inside_the_slide():
    boxes = [box("A", "x", "y"), box("B", "x"), box("C", "x"), box("D", "x")]
    links = [Link(src=0, dst=1), Link(src=0, dst=2), Link(src=0, dst=3)]
    s = Slide(title=T("t", "title"), elements=boxes, links=links, classes=["dense"])
    placed, deck = lay(s)
    assert all(p.y + p.h <= H for p in placed)
    assert not overflow(deck)


# ---- 5: a table with a few short columns does not stretch over the whole width


def _figures():
    rows = [[cell("項目"), cell("2026"), cell("2029")]] + [
        [cell("配送コスト率"), cell("11.2%"), cell("9.2%")],
        [cell("積載率"), cell("68%"), cell("82%")],
    ]
    return Table(rows=rows)


def test_numeric_table_width_is_capped_and_left_aligned():
    s = Slide(title=T("t", "title"), elements=[_figures()])
    placed, _ = lay(s)
    (tp,) = of(placed, Table)
    assert tp.x == to_emu(get_theme("jp-business").margin_x)
    assert tp.w < 0.8 * full_width()
    assert sum(tp.element.attrs["_col_w"]) == tp.w


def test_text_table_keeps_full_width():
    placed, _ = lay(Slide(title=T("t", "title"), elements=[_plain_table()]))
    (tp,) = of(placed, Table)
    assert tp.w == full_width()


def test_table_with_explicit_column_widths_keeps_full_width():
    t = _figures().model_copy(update={"col_widths": ["1", "1", "1"]})
    placed, _ = lay(Slide(title=T("t", "title"), elements=[t]))
    (tp,) = of(placed, Table)
    assert tp.w == full_width()


def test_many_column_table_keeps_full_width():
    rows = [[cell(f"h{i}") for i in range(7)], [cell(str(i)) for i in range(7)]]
    placed, _ = lay(Slide(title=T("t", "title"), elements=[Table(rows=rows)]))
    (tp,) = of(placed, Table)
    assert tp.w == full_width()


def test_empty_cells_never_raise():
    t = Table(rows=[[cell(""), cell("")], [cell(""), cell("")]])
    placed, _ = lay(Slide(title=T("t", "title"), elements=[t]))
    assert placed


# ---- a numeric table next to a full-width block keeps the full width (no ragged edge)


def test_table_below_a_box_row_keeps_full_width():
    s = Slide(title=T("t", "title"), elements=[box("A", "x"), box("B", "y"), box("C", "z"), _figures()])
    placed, _ = lay(s)
    (tp,) = of(placed, Table)
    assert tp.w == full_width()


def test_table_below_a_chart_keeps_full_width():
    s = Slide(title=T("t", "title"), grid="1", elements=[chart(), _figures()])
    placed, _ = lay(s)
    (tp,) = of(placed, Table)
    assert tp.w == full_width()


def test_table_with_text_only_is_still_capped():
    s = Slide(title=T("t", "title"), elements=[T("a short remark"), _figures(), T("another remark")])
    placed, _ = lay(s)
    (tp,) = of(placed, Table)
    assert tp.w < 0.8 * full_width()


def test_kpi_row_group_above_a_table_keeps_full_width():
    md = (
        "theme: jp-business\n\n# t\n"
        "## A {.kpi}\n1\n## B {.kpi}\n2\n## C {.kpi}\n3\n## D {.kpi}\n4\n@end\n"
        "| 項目 | 2026 | 2029 |\n|-|-|-|\n| 配送 | 11.2% | 9.2% |\n| 積載 | 68% | 82% |\n"
    )
    deck = parse(md)
    placed = layout_slide(deck.slides[0], deck, get_theme(deck.theme), 0)
    (tp,) = of(placed, Table)
    assert tp.w == full_width()
