"""A badge in a table cell stays on the line of the word before it when the column can be widened."""

from __future__ import annotations

from slidemark.layout import measure, tables
from slidemark.layout.tables import column_widths, table_grid
from slidemark.parser import parse
from slidemark.theme import get_theme

MD = """# T
{header=2}
| Hạng mục | 2027 | < | 2028 |
|-|-|-|-|
| ^ | H1 | H2 | Cả năm |
| Kho thông minh | Thiết kế | Thí điểm | 12 kho [xong]{.badge .success} |
| Điều phối AI | Thu thập dữ liệu | Mô hình | Toàn quốc |
"""


def _table():
    from slidemark.ir import Table

    deck = parse(MD)
    return next(e for e in deck.slides[0].elements if isinstance(e, Table))


def test_badge_cell_reserves_word_plus_badge():
    measure.set_tokens(get_theme("default").layout)
    t = _table()
    nrows, ncols, anchors = table_grid(t)
    size = 18.0
    total = 9 * 914400
    cw = column_widths(t, ncols, anchors, total, size)
    cell = next(c for r, k, c in anchors if k == 3 and r == 2)
    keep = tables._para_em(cell.paragraphs[0], False, None)[1]
    assert keep >= tables._badge_em("xong")  # the badge is one unit, with its padding
    col_em = cw[3] / 12700 / size - 2 * measure.cell_pad()[0] / 12700 / size
    assert col_em >= keep  # word + badge fit on one line


def test_badge_is_never_split_and_pad_is_counted():
    measure.set_tokens(get_theme("default").layout)
    assert tables._badge_em("xong") > measure.text_em("xong", bold=True)
    assert tables._badge_em("完了") > measure.text_em("完了", bold=True)


def test_rendered_badge_shares_the_paragraph(tmp_path):
    from pptx import Presentation

    from slidemark import build

    out = tmp_path / "b.pptx"
    build(MD, out)
    tbl = next(sh.table for sh in Presentation(out).slides[0].shapes if sh.has_table)
    para = tbl.cell(2, 3).text_frame.paragraphs[0]
    txt = "".join(r.text for r in para.runs)
    assert "12 kho" in txt and "xong" in txt and len(tbl.cell(2, 3).text_frame.paragraphs) == 1
