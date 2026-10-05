from __future__ import annotations

import shutil
import subprocess

import pytest
from pptx import Presentation

from slidemark.ir import Cell, Container, Deck, Image, Paragraph, Run, Slide, Table, Text
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.theme import get_theme


def T(s, role="body", **kw):
    return Text(role=role, paragraphs=[Paragraph(runs=[Run(text=s)])], **kw)


def c(t):
    return Cell(paragraphs=[Paragraph(runs=[Run(text=t)])])


def build(deck, tmp_path, theme="default"):
    th = get_theme(theme)
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    out = tmp_path / "out.pptx"
    render(deck, placed, th, out)
    return Presentation(str(out)), out


def cell_fill(cell) -> str:
    return str(cell.fill.fore_color.rgb)


def table_of(prs):
    return next(s for s in prs.slides[0].shapes if s.has_table).table


def test_table_header_fill_borders_and_zebra(tmp_path):
    rows = [[c("h1"), c("h2")]] + [[c(f"a{i}"), c(f"b{i}")] for i in range(4)]
    d = Deck(slides=[Slide(title=T("t", "title"), elements=[Table(rows=rows, classes=["zebra"])])])
    prs, _ = build(d, tmp_path, "jp-business")
    th = get_theme("jp-business")
    tbl = table_of(prs)
    assert cell_fill(tbl.cell(0, 0)) == th.color("primary").lstrip("#")  # header fill from theme
    fills = [cell_fill(tbl.cell(r, 0)) for r in range(1, 5)]
    assert fills[0] == fills[2] and fills[1] == fills[3] and fills[0] != fills[1]
    xml = tbl.cell(1, 1)._tc.xml
    assert all(t in xml for t in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"))
    assert 'w="6350"' in xml  # thin borders
    assert "tableStyleId" not in tbl._tbl.xml  # explicit fills only


def test_table_without_zebra_has_uniform_body(tmp_path):
    rows = [[c("h")]] + [[c(f"a{i}")] for i in range(4)]
    prs, _ = build(Deck(slides=[Slide(elements=[Table(rows=rows)])]), tmp_path)
    tbl = table_of(prs)
    assert len({cell_fill(tbl.cell(r, 0)) for r in range(1, 5)}) == 1


def test_image_cover_crops_preserving_aspect_and_contain_centers(tmp_path):
    from PIL import Image as PI

    PI.new("RGB", (400, 100), "red").save(tmp_path / "wide.png")
    PI.new("RGB", (100, 400), "blue").save(tmp_path / "tall.png")
    els = [
        Image(src="wide.png", fit="cover", alt="w"),
        Image(src="tall.png", fit="cover", alt="t"),
        Image(src="wide.png", fit="contain", alt="c"),
    ]
    d = Deck(slides=[Slide(layout="blank", grid="3", elements=els)], attrs={"base_dir": str(tmp_path)})
    prs, _ = build(d, tmp_path)
    pics = [s for s in prs.slides[0].shapes if s.shape_type == 13]
    assert len(pics) == 3
    for pic, (iw, ih) in zip(pics[:2], [(400, 100), (100, 400)], strict=True):
        assert "a:srcRect" in pic._element.xml
        shown = (1 - pic.crop_left - pic.crop_right) * iw / ((1 - pic.crop_top - pic.crop_bottom) * ih)
        assert shown == pytest.approx(pic.width / pic.height, rel=0.01)
    assert pics[0].crop_left > 0 and pics[0].crop_top == 0
    assert pics[1].crop_top > 0 and pics[1].crop_left == 0
    ct = pics[2]
    assert ct.crop_left == ct.crop_top == 0
    assert ct.width / ct.height == pytest.approx(4.0, rel=0.01)


def test_chevron_box_with_table_renders_both(tmp_path):
    tbl = Table(rows=[[c("h1"), c("h2")], [c("x"), c("y")]])
    s = Slide(
        title=T("Flow", "title"),
        classes=["chevron"],
        grid="2",
        elements=[
            Container(title=T("Step 1", "heading"), children=[T("plan"), tbl]),
            Container(title=T("Step 2", "heading"), children=[T("do")]),
        ],
    )
    d = Deck(slides=[s])
    prs, _ = build(d, tmp_path)
    shapes = list(prs.slides[0].shapes)
    chev = [x for x in shapes if x.shape_type == 1 and "CHEVRON" in str(x.auto_shape_type)]
    assert len(chev) == 2
    assert any(x.has_table for x in shapes)
    assert "plan" in chev[0].text_frame.text
    pPr = chev[0].text_frame.paragraphs[0]._p.pPr
    assert pPr.get("eaLnBrk") == "1" and pPr.get("hangingPunct") == "0"
    assert not [x for x in d.diagnostics if x.rule == "dropped-content"]


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_deck_opens_in_libreoffice(tmp_path):
    from PIL import Image as PI

    jp = "新規顧客の獲得数は前年同期比で十二パーセント増加し（速報値）、「解約率」も改善した。" * 4
    rows = [[c("項目"), c("内容")]] + [[c(f"行{i}"), c(jp[:30])] for i in range(6)]
    PI.new("RGB", (400, 100), "red").save(tmp_path / "a.png")
    s1 = Slide(
        title=T("日本語の密なスライド", "title"),
        grid="2",
        elements=[
            Container(title=T("現状", "heading"), children=[T(jp)]),
            Table(rows=rows, classes=["zebra"]),
        ],
        footnotes=[T("※ 出所：社内資料（2026年）", "footnote")],
    )
    s2 = Slide(
        title=T("Flow", "title"),
        classes=["chevron"],
        grid="2",
        elements=[
            Container(title=T("Step 1", "heading"), children=[T("plan"), Table(rows=rows[:3])]),
            Container(title=T("Step 2", "heading"), children=[T("do")]),
        ],
    )
    s3 = Slide(layout="blank", elements=[Image(src="a.png", fit="cover")])
    d = Deck(slides=[s1, s2, s3], attrs={"base_dir": str(tmp_path)})
    _, out = build(d, tmp_path, "jp-business")
    r = subprocess.run(
        [
            "soffice",
            "--headless",
            "-env:UserInstallation=file:///tmp/lo-b2",
            "--convert-to",
            "pdf",
            "--outdir",
            str(tmp_path),
            str(out),
        ],
        capture_output=True,
        timeout=180,
    )
    pdf = tmp_path / "out.pdf"
    assert pdf.exists() and pdf.stat().st_size > 1000, r.stderr
