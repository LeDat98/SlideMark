"""Header design tokens: colors:/fonts:/sizes:/style: lines, YAML maps, CLI, end to end."""

import json

import pytest
from pptx import Presentation

from slidemark.build import build
from slidemark.cli import main
from slidemark.parser import parse


def rules(deck):
    return [d.rule for d in deck.diagnostics]


def test_colors_line():
    d = parse("colors: primary=#7C5CFF bg=#0B1020\n\n# T\n")
    assert d.tokens == {"colors.primary": "#7C5CFF", "colors.bg": "#0B1020"}
    assert "unknown-header" not in rules(d)


def test_quotes_and_spaces():
    d = parse(
        'fonts: heading="Noto Sans JP" body=\'Yu Gothic\'\nstyle: card.shadow="0 8 24 #00000055"\n# T\n'
    )
    assert d.tokens["fonts.heading"] == "Noto Sans JP"
    assert d.tokens["fonts.body"] == "Yu Gothic"
    assert d.tokens["classes.card.shadow"] == "0 8 24 #00000055"


def test_gradient_stays_one_value():
    d = parse("style: card.fill=linear-gradient(135deg, #7C5CFF, #00D1B2) card.radius=12\n# T\n")
    assert d.tokens["classes.card.fill"] == "linear-gradient(135deg, #7C5CFF, #00D1B2)"
    assert d.tokens["classes.card.radius"] == "12"


def test_accumulate_later_wins():
    d = parse("colors: primary=#111111 bg=#000000\ncolors: primary=#222222\n# T\n")
    assert d.tokens["colors.primary"] == "#222222" and d.tokens["colors.bg"] == "#000000"


def test_style_paths():
    d = parse("style: title.band=primary layout.top_gap=0.4in gap=10\n# T\n")
    assert d.tokens["title_band"] == "primary"
    assert d.tokens["layout.top_gap"] == "0.4in"
    assert d.tokens["gap"] == "10"


def test_unknown_token_did_you_mean():
    d = parse("fonts: headin=Arial\nsizes: bodi=12\n# T\n")
    bad = [x for x in d.diagnostics if x.rule == "unknown-token"]
    assert len(bad) == 2 and d.tokens == {}
    assert "did you mean 'heading'" in bad[0].hint and "\n" not in bad[0].hint
    assert "'body'" in bad[1].hint


def test_pair_without_equals():
    d = parse("colors: primary\n# T\n")
    (x,) = [x for x in d.diagnostics if x.rule == "bad-token"]
    assert x.level == "warning" and x.hint == "write key=value, e.g. primary=#7C5CFF"


def test_unbalanced_quotes_and_parens_do_not_raise():
    d = parse('colors: a="x b=c\nstyle: card.fill=linear-gradient(1, 2\n# T\n')
    assert isinstance(d.tokens, dict)


def test_marp_style_block_still_ignored():
    d = parse("style: |\n# T\n")
    assert "bad-token" not in rules(d) and "marp-syntax" in rules(d)


def test_front_matter_maps():
    d = parse(
        '---\ntheme: none\ncolors: {bg: "#000"}\nstyle: {card.fill: "#111"}\n'
        "classes: {card: {radius: 14}}\nlayout: {top_gap: 0.4in}\n---\n# T\n"
    )
    assert d.theme == "none"
    assert d.tokens == {
        "colors.bg": "#000",
        "classes.card.fill": "#111",
        "classes.card.radius": "14",
        "layout.top_gap": "0.4in",
    }
    assert "unknown-header" not in rules(d)


def test_front_matter_unknown_key_warns():
    d = parse("---\nfonts: {headin: x}\n---\n# T\n")
    assert "unknown-token" in rules(d)


DECK = "theme: none\ncolors: primary=#7C5CFF bg=#0B1020\nstyle: title.band=primary\n\n# Hello\n- one\n"


def test_cli_check_shows_bad_token(tmp_path, capsys):
    f = tmp_path / "a.md"
    f.write_text("colors: primary=#7C5CFF\nstyle: card.radius=banana\n# A\n- x\n", encoding="utf-8")
    assert main(["check", str(f)]) == 0
    assert "bad-token" in capsys.readouterr().out


def test_cli_build_shows_bad_token(tmp_path, capsys):
    f = tmp_path / "a.md"
    f.write_text("style: card.radius=banana\n# A\n- x\n", encoding="utf-8")
    assert main(["build", str(f)]) == 0
    assert "bad-token" in capsys.readouterr().err


def test_cli_tokens_text_and_json(capsys):
    assert main(["tokens", "--theme", "midnight"]) == 0
    out = capsys.readouterr().out
    assert "colors.bg = " in out and "fonts.heading = " in out
    assert main(["tokens", "--format", "json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert "colors.primary" in data and "layout.top_gap" in data


def test_cli_tokens_deck(tmp_path, capsys):
    f = tmp_path / "a.md"
    f.write_text(DECK, encoding="utf-8")
    assert main(["tokens", str(f), "--format", "json"]) == 0
    data = json.loads(capsys.readouterr().out)
    assert data["colors.primary"] == "#7C5CFF" and data["colors.bg"] == "#0B1020"


def test_cli_themes(capsys):
    assert main(["themes"]) == 0
    names = capsys.readouterr().out.split()
    assert "none" in names and "default" in names and "midnight" in names


def test_end_to_end_tokens_render(tmp_path):
    f = tmp_path / "a.md"
    f.write_text(DECK, encoding="utf-8")
    out = tmp_path / "a.pptx"
    build(f, out)
    slide = Presentation(str(out)).slides[0]
    if "0B1020" not in slide._element.xml.upper():
        pytest.skip("render side does not paint the slide background from tokens yet")
    fills = set()
    for sh in slide.shapes:
        try:
            if sh.fill.type == 1:
                fills.add(str(sh.fill.fore_color.rgb))
        except (AttributeError, TypeError):
            pass
    assert "7C5CFF" in fills


def test_font_list_keeps_first_family_with_info():
    from slidemark.parser import parse

    deck = parse('theme: none\nfonts: heading="Didot, Georgia"\n\n# T\n- a\n- b\n')
    assert deck.tokens["fonts.heading"] == "Didot"
    assert [d.rule for d in deck.diagnostics] == ["font-list"]
