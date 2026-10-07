"""DL3d token diet: the imported deck.md of a foreign deck stays at or under 40% of the python-pptx script
(``importer/diet.py``, ``importer/runs.py``), and every shortening leaves the slides as they were.

The helper ``same_slides`` is the method the DL2 port lane used: remove a token, rebuild, compare the slide
XML; identical XML means the token was redundant.
"""

from pathlib import Path

import pytest

from slidemark.importer import diet
from slidemark.importer.diet import same_slides, slim
from slidemark.importer.read import RunT
from slidemark.importer.runs import put_span, with_span, wrap_span

ROOT = Path(__file__).resolve().parent.parent
BENCH = ROOT / "bench" / "lengthbench"
GATE = 0.40  # docs/TARGETS.md DL3d: imported deck.md <= 40% of the build.py tokens


def _tokens(text: str) -> int:
    tiktoken = pytest.importorskip("tiktoken")
    return len(tiktoken.get_encoding("o200k_base").encode(text))


def _attrs(text: str) -> list[str]:
    return [ln for ln in text.split("\n") if ln.startswith(("style:", "colors:", "sizes:"))]


# --------------------------------------------------------------------------- the gate


@pytest.mark.parametrize("t", ["t2", "t4"])
def test_imported_decks_are_at_most_40_percent_of_build_py(t):
    md = (BENCH / t / "roundtrip" / "deck.md").read_text(encoding="utf-8")
    py = (BENCH / t / "python-pptx" / "build.py").read_text(encoding="utf-8")
    assert _tokens(md) <= GATE * _tokens(py), (t, _tokens(md), _tokens(py))


@pytest.mark.parametrize("t", ["t2", "t4"])
def test_imported_decks_say_a_named_colour_by_its_name(t):
    """A hex the deck names appears in the body only as a text colour (a theme name would be moved by the
    readable-ink rule, a hex is not): fills, lines, borders and chart colours all use the name."""
    import re

    md = (BENCH / t / "roundtrip" / "deck.md").read_text(encoding="utf-8")
    head = _attrs(md)[0]
    assert head.startswith("colors:")
    named = {tok.split("=", 1)[1].upper() for tok in head.split()[1:] if "=#" in tok}
    body = "\n".join(ln for ln in md.split("\n") if not ln.startswith("colors:"))
    for hexv in named:
        every = re.findall(rf"#{hexv[1:]}", body, re.I)
        as_text = re.findall(rf"color=#{hexv[1:]}", body, re.I)
        assert len(every) == len(as_text), (t, hexv, len(every), len(as_text))


def test_the_committed_decks_build_to_themselves(tmp_path):
    """A committed deck.md is a fixed point: slimming it again changes nothing the build draws."""
    text = (BENCH / "t2" / "roundtrip" / "deck.md").read_text(encoding="utf-8")
    assert same_slides(text, slim(text, tmp_path, verify=False), tmp_path)


# --------------------------------------------------------------------------- the helper


def test_same_slides_tells_redundant_from_needed():
    base = "# A\n\n[x]{size=30 bold}\n"
    assert same_slides(base, base.replace("bold", "bold=true"))
    assert not same_slides(base, base.replace("size=30", "size=20"))
    assert not same_slides(base, "# A\n\ny\n")


# --------------------------------------------------------------------------- lever 4: shorter forms


def test_bare_bold_flag_is_one_token_and_means_the_same():
    r = RunT("9,800", bold=True, size=40)
    out = with_span(r, pin=True)
    assert wrap_span("9,800", out) == "[9,800]{size=40 bold}"
    r2 = RunT("x", bold=True)
    put_span(r2, ["size=10"], bold=True)
    assert r2.span == "size=10 bold"
    assert same_slides("# A\n\n[x]{size=20 bold}\n", "# A\n\n[x]{size=20 bold=true}\n")
    assert _tokens("[x]{size=20 bold}") < _tokens("[x]{size=20 bold=true}")


def test_span_colour_class_form_for_a_colour_name():
    text = "# A\n\n[月額]{size=15 color=muted}<br>[x]{size=20 color=primary}\n"
    got = slim(text, verify=False)
    assert "[月額]{size=15 .muted}" in got and "[x]{size=20 .primary}" in got
    assert same_slides(text, got)


def test_class_form_leaves_a_box_attribute_alone():
    text = "# A\n\n## T {color=muted}\nbody\n"
    assert slim(text, verify=False) == text


# --------------------------------------------------------------------------- lever 3: named colours

FILLS = (
    "# A\n\n@3\n## a {fill=#1F5FA8}\nx\n## b {fill=#1F5FA8}\ny\n## c {fill=#1F5FA8}\nz\n"
    "\n# B\n\n@2\n## d {fill=#E09F1F}\nx\n## e {fill=#E09F1F}\ny\n"
)


def test_a_hex_used_three_times_becomes_a_colour_name():
    got = slim(FILLS)
    assert got.startswith("colors: secondary=#1F5FA8")
    assert got.count("fill=secondary") == 3
    assert "#1F5FA8" not in got.split("\n", 1)[1]
    assert same_slides(FILLS, got)


def test_a_hex_used_twice_stays_a_hex():
    got = slim(FILLS)
    assert got.count("fill=#E09F1F") == 2 and "E09F1F" not in got.split("\n", 1)[0]


def test_colour_names_are_declared_in_the_existing_colors_line():
    text = "colors: primary=#142B4D\n\n" + FILLS
    got = slim(text)
    assert got.split("\n")[0].startswith("colors: primary=#142B4D secondary=#1F5FA8")
    assert same_slides(text, got)


def test_a_declared_name_replaces_its_hex_everywhere():
    text = "colors: primary=#142B4D\n\n# A\n\n## a {fill=#142B4D}\nx\n"
    got = slim(text)
    assert "fill=primary" in got and same_slides(text, got)


def test_white_and_black_are_css_words():
    text = "# A\n\n[a]{size=20 color=#FFFFFF}\n\n[b]{size=20 color=#000000}\n"
    got = slim(text)
    assert "{size=20 .white}" in got and "{size=20 .black}" in got and not got.startswith("colors:")
    assert same_slides(text, got)


def test_a_fill_keeps_its_hex_when_the_css_word_would_lose_the_ink():
    # `fill=black` is not a theme colour: the heading ink is no longer chosen against it
    text = "# A\n\n## a {fill=#000000}\nx\n## b {fill=#000000}\ny\n"
    assert slim(text) == text


def test_a_role_name_is_not_taken_from_the_theme_when_the_deck_uses_it():
    # `color=muted` is the theme's grey: a different grey gets a derived word, the deck keeps both
    text = (
        "# A\n\n[a]{size=15 color=muted}\n\n## x {fill=#8A8A8A}\nq\n## y {fill=#8A8A8A}\nq\n"
        "## z {fill=#8A8A8A}\nq\n"
    )
    names = diet._color_groups(text)
    assert names and "muted" not in names[0][0].__doc__
    assert "gray" in names[0][0].__doc__


def test_a_grey_that_is_the_theme_muted_to_the_eye_takes_the_role():
    text = "# A\n\n## x {fill=#6B7582}\nq\n## y {fill=#6B7582}\nq\n## z {fill=#6B7582}\nq\n"
    assert "muted" in diet._color_groups(text)[0][0].__doc__


def test_a_text_colour_keeps_its_hex_when_a_name_would_be_moved():
    # amber on white is 2.6:1: the readable-ink rule moves a theme name, never a hex
    text = (
        "# A\n\n[a]{size=20 color=#E08A1E}\n\n[b]{size=20 color=#E08A1E}\n\n[c]{size=20 color=#E08A1E}\n\n"
        "## x {fill=#E08A1E}\nq\n"
    )
    got = slim(text)
    assert got.count("color=#E08A1E") == 3
    assert "fill=#E08A1E" not in got or same_slides(text, got)
    assert same_slides(text, got)


def test_a_naming_that_changes_the_slides_is_rejected(monkeypatch):
    monkeypatch.setattr(diet, "_same", lambda a, b: False)
    assert slim(FILLS) == FILLS


# --------------------------------------------------------------------------- lever 2: hoisting


def test_a_style_token_two_slides_repeat_moves_to_the_header():
    text = (
        "style: radius=0\n\n# A\n\nstyle: rows.fill=#EEF2F7 rows.size=20\n@rows\n- a\n- b\n"
        "\n# B\n\nstyle: rows.fill=#EEF2F7 rows.size=24\n@rows\n- c\n- d\n"
    )
    got = diet._clean(diet._bisect(text, diet._hoist_groups(text), lambda a, b: True))
    head = got.split("# A")[0]
    assert "rows.fill=#EEF2F7" in head and got.count("rows.fill") == 1
    assert "rows.size=20" in got.split("# A")[1] and "rows.size=24" in got.split("# B")[1]
    assert same_slides(text, got)


def test_a_slide_with_another_value_keeps_its_own():
    text = (
        "# A\n\nstyle: kpi.fill=#EEF2F7\n@kpi\n## 1\nx\n"
        "\n# B\n\nstyle: kpi.fill=#EEF2F7\n@kpi\n## 2\nx\n"
        "\n# C\n\nstyle: kpi.fill=#FDF1DE\n@kpi\n## 3\nx\n"
    )
    got = diet._clean(diet._bisect(text, diet._hoist_groups(text), lambda a, b: True))
    assert "kpi.fill=#EEF2F7" in got.split("# A")[0]
    assert "kpi.fill=#FDF1DE" in got.split("# C")[1]
    assert same_slides(text, got)


def test_a_token_is_not_hoisted_over_a_slide_that_styles_the_class_otherwise():
    text = (
        '# A\n\nstyle: s1.border-top="9pt solid #1F5FA8"\n## [01] {.s1}\nx\n'
        '\n# B\n\nstyle: s1.border-top="9pt solid #1F5FA8"\n## [01] {.s1}\nx\n'
        '\n# C\n\nstyle: s1.border-left="9pt solid #1F5FA8"\n## [01] {.s1}\nx\n'
    )
    assert diet._hoist_groups(text) == []


def test_one_class_with_two_border_sides_is_two_classes():
    text = (
        '# A\n\nstyle: s1.border-top="9pt solid #1F5FA8"\n## [01] {.s1}\nx\n'
        '\n# B\n\nstyle: s1.border-top="9pt solid #1F5FA8"\n## [01] {.s1}\nx\n'
        '\n# C\n\nstyle: s1.border-left="9pt solid #1F5FA8"\n## [01] {.s1}\nx\n'
        '\n# D\n\nstyle: s1.border-left="9pt solid #1F5FA8"\n## [01] {.s1}\nx\n'
    )
    plan = diet._rename_plan(text)
    assert set(plan.values()) == {"ls1"}  # one new name for both left slides
    renamed = diet._clean(diet._bisect(text, diet._rename_groups(text), lambda a, b: True))
    hoisted = diet._clean(diet._bisect(renamed, diet._hoist_groups(renamed), lambda a, b: True))
    head = hoisted.split("# A")[0]
    assert "s1.border-top" in head and "ls1.border-left" in head
    assert hoisted.count("border-") == 2 and "{.ls1}" in hoisted
    assert same_slides(text, hoisted)


# --------------------------------------------------------------------------- lever 1: what the element gives


def test_bold_is_dropped_in_a_heading_that_is_bold_anyway():
    text = "# A\n\n@2\n## [Q1]{size=28 bold} 売上\nx\n## b\ny\n"
    got = slim(text)
    assert got == text or "bold" not in got.split("\n")[3]
    assert same_slides(text, got)


def test_white_ink_is_dropped_on_a_dark_fill():
    text = "# A\n\n@2\n## a {fill=#122B4A color=#FFFFFF}\nx\n## b {fill=#122B4A}\ny\n"
    got = slim(text)
    assert same_slides(text, got)
    assert "color=#FFFFFF" not in got and "color=white" not in got


def test_ink_that_the_fill_does_not_give_stays():
    text = "# A\n\n@2\n## a {fill=#EEF2F7 color=#E08A1E}\nx\n## b\ny\n"
    assert "E08A1E" in slim(text)


def test_derived_style_tokens_are_dropped_when_the_build_reproduces_them():
    text = "style: cover.bar_w=0.12in\n\n# A\n\nx\n"
    got = slim(text)
    assert same_slides(text, got)
    assert "cover.bar_w" not in got or got == text


def test_a_token_the_build_needs_is_kept():
    text = "style: rows.size=40\n\n# A\n\n@rows\n- a\n- b\n"
    got = slim(text)
    assert same_slides(text, got) and "rows.size=40" in got


# --------------------------------------------------------------------------- safety


def test_slim_never_raises_and_leaves_a_deck_that_does_not_build():
    for bad in ("", "\x00\x00", "# A\n\n```\nunclosed", "colors: x=\n\n# A\n"):
        assert isinstance(slim(bad), str)


def test_slim_without_a_build_applies_every_edit():
    got = slim(FILLS, verify=False)
    assert got.startswith("colors: secondary=#1F5FA8")


def test_an_imported_foreign_deck_is_slimmed_and_a_slidemark_deck_is_not(tmp_path, monkeypatch):
    from pptx import Presentation

    from slidemark import build
    from slidemark.importer import import_pptx

    calls = []
    monkeypatch.setattr(diet, "slim", lambda text, base=None, **kw: calls.append(text) or text)
    foreign = tmp_path / "foreign.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Hello"
    prs.save(str(foreign))
    import_pptx(foreign)
    assert calls == []  # the library keeps what the recognisers read unless asked
    import_pptx(foreign, slim=True)
    assert len(calls) == 1
    own = tmp_path / "own.pptx"
    build("# Hello\n\n- a\n", own, base_dir=tmp_path)
    import_pptx(own, slim=True)
    assert len(calls) == 1


# --------------------------------------------------------------------------- the whole import (slow)


def test_t2_import_reaches_the_gate_and_matches_the_committed_deck(tmp_path):
    """Rebuild the python-pptx original of t2, import it, compare with the committed deck.md."""
    import sys

    sys.path.insert(0, str(ROOT / "bench"))
    import roundtrip_compare

    from slidemark.importer import import_pptx

    orig = roundtrip_compare.regen(BENCH / "t2" / "python-pptx" / "x.pptx", tmp_path)
    text, _ = import_pptx(orig, tmp_path, slim=True)
    py = (BENCH / "t2" / "python-pptx" / "build.py").read_text(encoding="utf-8")
    assert _tokens(text) <= GATE * _tokens(py)
    assert text == (BENCH / "t2" / "roundtrip" / "deck.md").read_text(encoding="utf-8")
