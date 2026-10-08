"""Wave 2026-10-08 lane B: the grown-up composition forms.

``@cycle`` (nodes that follow the ring, ``icon=`` in the node, a centre label), ``@steps`` (the fit line reads
the body text, cards that grew say so, icon arrows fill the body like plain ones), ``@stairs``, ``@nested``
and ``@flow disc``. Every form: parser (word, keys, tokens), render (the .pptx is reopened: names, geometry,
ink), importer round trip, fit line, ``attr-ignored``, design facts and fuzz-safety (a diagnostic with a
hint, never an exception). Vietnamese headings carry diacritics on purpose."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st
from pptx import Presentation

from slidemark import build, forms, parse
from slidemark.design import stated_groups
from slidemark.importer import import_pptx
from slidemark.theme import normalize_token

ROOT = Path(__file__).resolve().parent.parent
HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#14272F primary=#0B2A3C secondary=#1B8A8F accent=#F2A33A teal=#2A9D8F "
    "muted=#5A6F78 surface=#EEF4F6 border=#CFDDE2\n"
    "fonts: heading=Arial body=Arial\n"
    "sizes: title=30 heading=20 body=16 caption=12\n"
)

CYCLE4 = (
    "\n# Vòng lặp agent\n@cycle\n"
    "## Quan sát\nĐọc yêu cầu và công cụ\n## Lập kế hoạch\nChia thành bước nhỏ\n"
    "## Hành động\nGọi công cụ, chạy mã\n## Phản hồi\nKiểm tra kết quả\n"
)
CYCLE_ICONS = (
    '\n# Vòng lặp agent\n@cycle center="Vòng lặp agent"\n'
    "## Quan sát {icon=search}\nĐọc yêu cầu và công cụ\n## Lập kế hoạch {icon=target}\nChia thành bước nhỏ\n"
    "## Hành động {icon=bolt}\nGọi công cụ, chạy mã\n## Phản hồi {icon=check}\nKiểm tra kết quả\n"
)
STAIRS4 = (
    "\n# Bốn cách đưa AI vào sản phẩm\n@stairs\n"
    "## Dùng sẵn\nMua dịch vụ có sẵn\n## Lời nhắc\nViết lời nhắc và mẫu\n"
    "## Truy xuất\nNối với tài liệu riêng\n## Huấn luyện\nTinh chỉnh mô hình riêng\n"
)
NESTED3 = (
    "\n# AI, học máy và học sâu\n@nested\n"
    "## Trí tuệ nhân tạo {icon=globe}\nMáy làm việc cần trí thông minh\n"
    "## Học máy {icon=chart}\nMáy tự học quy luật từ dữ liệu\n"
    "## Học sâu {icon=cpu}\nMạng nơ-ron nhiều tầng\n"
)
FLOW5 = (
    "\n# Quy trình RAG\n@flow disc\n"
    "## Kho tài liệu {.above icon=database}\nPDF, wiki, email\n"
    "## Câu hỏi {icon=chat}\nNgười dùng hỏi\n## Truy xuất {icon=search}\nTìm đoạn liên quan\n"
    "## Ghép ngữ cảnh {icon=document}\nĐưa vào lời nhắc\n## Trả lời {icon=bolt}\nMô hình trả lời\n"
)
FLOW3 = "\n# Ba bước\n@flow disc\n## Một\nviệc đầu\n## Hai\nviệc giữa\n## Ba\nviệc cuối\n"
STEPS_NUM = (
    "\n# Bốn bước\n@4 steps num\n"
    "## Token hóa {icon=code}\n- Cắt văn bản thành token\n- Một từ có thể thành vài token\n"
    "## Embedding {icon=database}\n- Mỗi token thành vector\n- Vector gần nhau nghĩa gần nhau\n"
    "## Attention {icon=search}\n- Cân nhắc token liên quan\n- Hiểu ngữ cảnh của cả câu\n"
    "## Dự đoán {icon=bolt}\n- Chọn token kế tiếp\n- Lặp lại đến khi xong\n"
    "@end\n\n> Một token mỗi lần\n"
)


def make(tmp_path, text: str, head: str = HEAD, name: str = "d"):
    deck = build(head + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def shapes(prs, i: int = 0, pattern: str = ".*"):
    return [s for s in prs.slides[i].shapes if re.fullmatch(pattern, s.name)]


def one(prs, i: int, name: str):
    got = shapes(prs, i, re.escape(name))
    assert len(got) == 1, (name, [s.name for s in prs.slides[i].shapes])
    return got[0]


def sizes(shp) -> list[float]:
    return [r.font.size.pt for p in shp.text_frame.paragraphs for r in p.runs if r.font.size]


def ink(shp) -> str:
    r = next(r for p in shp.text_frame.paragraphs for r in p.runs)
    return str(r.font.color.rgb)


def fill_hex(shp) -> str:
    return str(shp.fill.fore_color.rgb)


def lum(hexv: str) -> float:
    r, g, b = (int(hexv[i : i + 2], 16) / 255 for i in (0, 2, 4))

    def f(c: float) -> float:
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4

    return 0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b)


def centre(shp) -> tuple[float, float]:
    return shp.left + shp.width / 2, shp.top + shp.height / 2


def warns(deck, rule: str):
    return [d for d in deck.diagnostics if d.rule == rule]


def clean(deck):
    return [
        d
        for d in deck.diagnostics
        if d.level in ("warning", "error") and d.rule not in ("design-none", "design-slide")
    ]


NBSP, WJ = chr(0xA0), chr(0x2060)


def tx(shp) -> str:
    """A shape's text without the orphan binding (NBSP) and the range joiners (U+2060)."""
    return shp.text_frame.text.replace(NBSP, " ").replace(WJ, "")


def pt(par) -> str:
    """A paragraph's text, normalised like ``tx``."""
    return par.text.replace(NBSP, " ").replace(WJ, "")


# --------------------------------------------------------------------------- parser and contract


def test_the_new_words_are_forms_and_keys_reach_the_slide():
    assert {"stairs", "nested"} <= set(forms.FORMS)
    d = parse(HEAD + STAIRS4.replace("@stairs", "@stairs dir=down"))
    assert forms.form_of(d.slides[0]) == "stairs" and d.slides[0].attrs["dir"] == "down"
    d = parse(HEAD + NESTED3.replace("@nested", "@nested side=right"))
    assert forms.form_of(d.slides[0]) == "nested" and d.slides[0].attrs["side"] == "right"
    d = parse(HEAD + CYCLE_ICONS)
    assert d.slides[0].attrs["center"] == "Vòng lặp agent"
    d = parse(HEAD + FLOW5)
    assert forms.form_of(d.slides[0]) == "flowdisc" and "disc" in d.slides[0].classes
    assert d.slides[0].elements[0].classes == ["above"]
    assert not clean(d)


@pytest.mark.parametrize(
    "text, rule, hint",
    [
        ("\n# T\n@stairs dir=sideways\n## a\n## b\n", "bad-attr", "dir=up|down"),
        ("\n# T\n@nested side=middle\n## a\n## b\n## c\n", "bad-attr", "side=left|right"),
        ("\n# T\n@flow disc above=maybe\n## a\n## b\n", "bad-attr", "above=on|off|1|0"),
        ("\n# T\n@stairs marks=on\n## a\n## b\n", "attr-ignored", "@stairs takes dir="),
        ("\n# T\n@3 above=1\n## a\n## b\n## c\n", "attr-ignored", "@flow disc"),
        ("\n# T\n@3 disc\n## a\n## b\n## c\n", "attr-ignored", "@flow disc"),
    ],
)
def test_bad_attributes_warn_with_a_hint(text, rule, hint):
    got = warns(parse(HEAD + text), rule)
    assert got, text
    assert hint in (got[0].hint or ""), got[0]


@pytest.mark.parametrize(
    "form, rule, text",
    [
        ("stairs", "stairs-skipped", "\n# T\n@stairs\n## only\nx\n"),
        ("stairs", "stairs-skipped", "\n# T\n@stairs\n" + "".join(f"## s{k}\nx\n" for k in range(7))),
        ("nested", "nested-skipped", "\n# T\n@nested\n## only\nx\n"),
        ("nested", "nested-skipped", "\n# T\n@nested\n" + "".join(f"## r{k}\nx\n" for k in range(5))),
        ("nested", "nested-skipped", "\n# T\n@nested\n## a\nx\n## b\n\ntext not in a box\n"),
        ("flow", "flow-skipped", "\n# T\n@flow disc\n## only\nx\n"),
        ("flow", "flow-skipped", "\n# T\n@flow disc\n" + "".join(f"## s{k}\nx\n" for k in range(8))),
    ],
)
def test_a_slide_without_what_the_form_needs_is_skipped_with_a_hint_and_never_raises(
    form, rule, text, tmp_path
):
    deck = parse(HEAD + text)
    got = warns(deck, rule)
    assert got and "laid out as ordinary blocks" in got[0].hint and "## " in got[0].hint
    built, prs = make(tmp_path, text)  # never raises, never composes
    assert not [s for s in prs.slides[0].shapes if re.match(r"(Stairs|Nested|Flow \d)", s.name)]


# --------------------------------------------------------------------------- tokens


@pytest.mark.parametrize(
    "key, value, path",
    [
        ("cycle.center", '"Vòng lặp"', "cycle_center"),
        ("cycle.center.size", "30", "cycle_center_size"),
        ("cycle.center.color", "accent", "cycle_center_color"),
        ("cycle.center.fill", "surface", "cycle_center_fill"),
        ("cycle.head.color", "primary", "cycle_head_color"),
        ("cycle.icon.ratio", "0.5", "cycle_icon_ratio"),
        ("cycle.icon.color", "accent", "cycle_icon_color"),
        ("stairs.fill", "primary,teal", "stairs_fill"),
        ("stairs.color", "bg", "stairs_color"),
        ("stairs.step", "0.9in", "stairs_step"),
        ("stairs.gap", "0.2in", "stairs_gap"),
        ("stairs.low", "0.5", "stairs_low"),
        ("nested.fill", "teal,secondary,primary", "nested_fill"),
        ("nested.color", "bg", "nested_color"),
        ("nested.size", "24", "nested_size"),
        ("nested.core", "0.5", "nested_core"),
        ("nested.share", "0.45", "nested_share"),
        ("nested.line", "bg", "nested_line"),
        ("nested.line.w", "2pt", "nested_line_w"),
        ("nested.list.title", "off", "nested_list_title"),
        ("flow.disc.size", "1.2in", "flow_disc_size"),
        ("flow.disc.fill", "primary,teal", "flow_disc_fill"),
        ("flow.disc.color", "bg", "flow_disc_color"),
        ("flow.line", "accent", "flow_line"),
        ("flow.line.w", "4pt", "flow_line_w"),
    ],
)
def test_the_form_tokens_are_valid_style_keys(key, value, path):
    from slidemark.theme import apply_tokens, canonical_token, get_theme

    got, hint = canonical_token("style", key)
    assert got == path, (key, got, hint)
    pairs = normalize_token(path, value, None)
    assert pairs and pairs[0][0] == path
    th, diags = apply_tokens(get_theme("default"), {path: value})
    assert not diags, diags


@pytest.mark.parametrize(
    "token",
    [
        "stairs.fill=notacolor",
        "nested.size=huge",
        "flow.disc.size=wide",
        "nested.list.title=maybe",
        "cycle.icon.ratio=9",
    ],
)
def test_a_bad_token_value_is_a_bad_token_warning_with_a_hint(token, tmp_path):
    deck, _ = make(tmp_path, STAIRS4, head=HEAD + f"style: {token}\n")
    got = warns(deck, "bad-token")
    assert got and got[0].hint, token


# --------------------------------------------------------------------------- @cycle


def test_cycle_node_text_is_never_smaller_than_the_body(tmp_path):
    _d, prs = make(tmp_path, CYCLE4)
    nodes = [one(prs, 0, f"Cycle {k}") for k in (1, 2, 3, 4)]
    assert all(min(sizes(n)) >= 16 for n in nodes), [sizes(n) for n in nodes]
    # Vietnamese syllables bound by the orphan control ("kế hoạch") still fit the node
    node = one(prs, 0, "Cycle 2")
    assert node.width == node.height and tx(node) == "Lập kế hoạch"


def test_cycle_node_size_still_pins_the_diameter_and_text_may_shrink(tmp_path):
    _d, prs = make(tmp_path, CYCLE4, head=HEAD + "style: cycle.node.size=0.7in\n")
    node = one(prs, 0, "Cycle 1")
    assert abs(node.width / 914400 - 0.7) < 0.01


def test_cycle_icon_sits_in_the_node_and_the_heading_moves_beside_it(tmp_path):
    deck, prs = make(tmp_path, CYCLE_ICONS)
    assert not clean(deck), [str(d) for d in clean(deck)]
    node = one(prs, 0, "Cycle 1")
    assert node.text_frame.text == "" and node.width == node.height
    icon = one(prs, 0, "icon search")
    (nx, ny), (ix, iy) = centre(node), centre(icon)
    assert abs(nx - ix) < 2000 and abs(ny - iy) < 2000  # centred
    assert abs(icon.width / node.width - 0.45) < 0.02  # ~45% of the diameter
    assert str(icon.fill.fore_color.rgb) == "FFFFFF"  # readable on the dark node
    text = one(prs, 0, "Cycle 1 text")
    paras = text.text_frame.paragraphs
    assert pt(paras[0]) == "Quan sát" and paras[0].runs[0].font.bold
    assert paras[0].runs[0].font.size.pt > paras[1].runs[0].font.size.pt  # heading above the body
    assert pt(paras[1]).startswith("Đọc yêu cầu")
    assert text.left > node.left + node.width or text.left + text.width < node.left  # beside, not inside


CYCLE_LONG = (
    "\n# Năm nút chữ dài\n@cycle\n"
    "## Thu thập yêu cầu của khách hàng\nPhỏng vấn và khảo sát\n## Phân tích\nTìm điểm đau\n"
    "## Thiết kế giải pháp\nPhác thảo\n## Xây dựng nguyên mẫu\nLàm nhanh\n"
    "## Kiểm thử cùng người dùng\nLấy phản hồi\n"
)


def test_cycle_headings_a_node_cannot_hold_at_the_body_size_go_beside_it_under_a_number(tmp_path):
    deck, prs = make(tmp_path, CYCLE_LONG)
    assert not clean(deck), [str(d) for d in clean(deck)]
    node = one(prs, 0, "Cycle 1")
    assert tx(node) == "1" and min(sizes(node)) >= 16  # never below the body size
    text = one(prs, 0, "Cycle 1 text")
    paras = text.text_frame.paragraphs
    assert pt(paras[0]) == "Thu thập yêu cầu của khách hàng" and paras[0].runs[0].font.bold
    assert pt(paras[1]).startswith("Phỏng vấn")
    # short headings keep the old look: the heading in the node
    _d, prs = make(tmp_path, CYCLE4, name="short")
    assert tx(one(prs, 0, "Cycle 1")) == "Quan sát"
    # a pinned node size is the author's: the heading stays in the node
    _d, prs = make(tmp_path, CYCLE_LONG, head=HEAD + "style: cycle.node.size=0.9in\n", name="pin")
    assert tx(one(prs, 0, "Cycle 1")).startswith("Thu thập")


def test_the_numbered_cycle_round_trips_through_the_importer(tmp_path):
    out = tmp_path / "n.pptx"
    build(HEAD + CYCLE_LONG, out)
    md, _ = import_pptx(out)
    lines = md.splitlines()
    assert "## Thu thập yêu cầu của khách hàng" in lines and "Phỏng vấn và khảo sát" in lines
    assert "## 1" not in lines


def test_cycle_without_icons_keeps_the_heading_in_the_node(tmp_path):
    _d, prs = make(tmp_path, CYCLE4)
    assert tx(one(prs, 0, "Cycle 1")) == "Quan sát"
    assert not [s for s in prs.slides[0].shapes if s.name.startswith("icon ")]
    assert not shapes(prs, 0, "Cycle center")


def test_cycle_center_label_at_the_ring_centre_and_its_tokens(tmp_path):
    _d, prs = make(tmp_path, CYCLE_ICONS)
    lab = one(prs, 0, "Cycle center")
    assert tx(lab) == "Vòng lặp agent" and lab.text_frame.paragraphs[0].runs[0].font.bold
    nodes = [centre(one(prs, 0, f"Cycle {k}")) for k in (1, 2, 3, 4)]
    cx = sum(n[0] for n in nodes) / 4
    cy = sum(n[1] for n in nodes) / 4
    assert abs(centre(lab)[0] - cx) < 20000 and abs(centre(lab)[1] - cy) < 20000
    head = HEAD + "style: cycle.center.size=30 cycle.center.color=#AA0000\n"
    _d, prs = make(tmp_path, CYCLE_ICONS, head=head, name="t")
    lab = one(prs, 0, "Cycle center")
    assert sizes(lab) == [30] and ink(lab) == "AA0000"
    # the token alone (no `center=` on the @ line) writes the label too, the @ line wins
    plain = CYCLE_ICONS.replace(' center="Vòng lặp agent"', "")
    _d, prs = make(tmp_path, plain, head=HEAD + 'style: cycle.center="Ở giữa"\n', name="u")
    assert tx(one(prs, 0, "Cycle center")) == "Ở giữa"
    _d, prs = make(tmp_path, CYCLE_ICONS, head=HEAD + 'style: cycle.center="Ở giữa"\n', name="v")
    assert tx(one(prs, 0, "Cycle center")) == "Vòng lặp agent"


def test_cycle_center_fill_draws_a_disc_with_readable_ink(tmp_path):
    _d, prs = make(tmp_path, CYCLE_ICONS, head=HEAD + "style: cycle.center.fill=primary\n")
    lab = one(prs, 0, "Cycle center")
    assert fill_hex(lab) == "0B2A3C" and ink(lab) == "FFFFFF"


def test_cycle_arrows_keep_their_heads_and_names(tmp_path):
    _d, prs = make(tmp_path, CYCLE_ICONS)
    arcs = [s for s in prs.slides[0].shapes if re.fullmatch(r"Cycle \d arrow", s.name)]
    assert len(arcs) == 4
    for a in arcs:
        xml = a._element.xml
        assert "tailEnd" in xml or "headEnd" in xml, a.name
    _d, prs = make(tmp_path, CYCLE_ICONS.replace("@cycle", "@cycle dir=ccw"), name="c")
    assert len([s for s in prs.slides[0].shapes if re.fullmatch(r"Cycle \d arrow ccw", s.name)]) == 4


def test_cycle_center_on_a_ring_too_small_says_so_and_still_builds(tmp_path):
    deck, prs = make(tmp_path, CYCLE_ICONS, head=HEAD + "style: cycle.node.size=3.4in\n")
    assert warns(deck, "cycle-center")
    assert not shapes(prs, 0, "Cycle center")


def test_cycle_icon_unknown_name_is_reported_and_the_heading_stays_in_the_node(tmp_path):
    text = CYCLE4.replace("## Quan sát", "## Quan sát {icon=nosuchicon}")
    deck, prs = make(tmp_path, text)
    assert warns(deck, "unknown-icon")
    assert tx(one(prs, 0, "Cycle 1")) == "Quan sát"


def test_cycle_icon_is_honoured_on_the_box_but_not_on_other_forms(tmp_path):
    deck, _ = make(tmp_path, CYCLE_ICONS)
    assert not warns(deck, "attr-ignored")
    deck, _ = make(tmp_path, "\n# T\n@funnel\n## a {icon=bolt}\nx\n## b\ny\n")
    assert [d.message for d in warns(deck, "attr-ignored")] == ["icon= on a form box is not honoured"]


# --------------------------------------------------------------------------- @steps


def test_steps_fit_line_reads_the_body_text_not_the_caption(tmp_path):
    from slidemark.fit import fit_lines
    from slidemark.layout import layout_slide
    from slidemark.template import deck_theme

    parsed = parse(HEAD + STEPS_NUM)
    theme, _ = deck_theme(parsed, str(tmp_path))
    placed = [layout_slide(parsed.slides[0], parsed, theme, 0)]
    fit = fit_lines(parsed, placed, theme)[0]
    assert "card text 16pt" in fit or "card text 16->" in fit
    assert "10pt" not in fit and "shrunk" not in fit, fit  # the STEP n caption (10 pt) is not the card text
    assert "cards grown to fill" in fit, fit


def test_steps_text_does_not_shrink_while_the_body_is_free(tmp_path):
    _d, prs = make(tmp_path, STEPS_NUM)
    body = [s for s in prs.slides[0].shapes if re.fullmatch(r"Text \d", s.name)]
    assert len(body) == 4
    for t in body:
        assert max(sizes(t)) >= 16, sizes(t)  # asked 16: grown, never shrunk
    cards = shapes(prs, 0, r"Step \d card")
    bar = one(prs, 0, "Conclusion")
    assert all(0 <= bar.top - (c.top + c.height) < 914400 * 0.4 for c in cards)  # the cards reach the bar


def test_steps_with_icons_fill_the_body_like_steps_without(tmp_path):
    sparse = (
        "\n# Ba bước\n> Lộ trình ngắn cho một nhóm nhỏ\n@3 steps\n"
        "## Khám phá {icon=search}\n- Liệt kê các việc lặp lại\n"
        "## Thử nghiệm {icon=bolt}\n- Dùng trợ lý AI có sẵn\n"
        "## Mở rộng {icon=rocket}\n- Nhân rộng cách làm tốt\n"
    )
    deck, prs = make(tmp_path, sparse)
    plain = re.sub(r" \{icon=\w+\}", "", sparse)
    _d2, prs2 = make(tmp_path, plain, name="p")
    card = lambda p: one(p, 0, "Step 1 card")  # noqa: E731
    assert card(prs).height > 914400 * 1.5, card(prs).height  # grew with the body, not left at its content
    assert abs(card(prs).height - card(prs2).height) < 914400 * 0.2
    arrow, icon = one(prs, 0, "Step 1 arrow"), one(prs, 0, "icon search")
    assert (
        arrow.top <= icon.top and icon.top + icon.height <= arrow.top + arrow.height
    )  # the icon rides its arrow
    assert max(sizes(one(prs, 0, "Text 1"))) > 16  # the text grew
    assert not [d for d in deck.diagnostics if d.rule in ("overlap", "overflow")], [
        str(d) for d in clean(deck)
    ]


def test_the_step_caption_is_not_a_card_item(tmp_path):
    sparse = "\n# Ba\n@3 steps num\n## A\n- một\n## B\n- hai\n## C\n- ba\n"
    _d, prs = make(tmp_path, sparse)
    t = one(prs, 0, "Text 1")
    assert max(sizes(t)) > 16  # the sparse composition ran although the card holds caption + one bullet


# --------------------------------------------------------------------------- @stairs


def test_stairs_cards_rise_left_to_right_with_bottoms_aligned(tmp_path):
    deck, prs = make(tmp_path, STAIRS4)
    assert not clean(deck)
    cards = [one(prs, 0, f"Stairs {k}") for k in (1, 2, 3, 4)]
    bottoms = {c.top + c.height for c in cards}
    assert len(bottoms) == 1
    hs = [c.height for c in cards]
    assert hs == sorted(hs) and len(set(hs)) == 4
    steps = [b - a for a, b in zip(hs, hs[1:], strict=False)]
    assert max(steps) - min(steps) < 5000  # an even rise
    assert cards[0].left < cards[1].left < cards[2].left < cards[3].left
    assert [tx(c).splitlines()[0] for c in cards] == ["Dùng sẵn", "Lời nhắc", "Truy xuất", "Huấn luyện"]


def test_stairs_the_last_card_reaches_the_body_top_and_dir_down_reverses(tmp_path):
    _d, prs = make(tmp_path, STAIRS4)
    cards = [one(prs, 0, f"Stairs {k}") for k in (1, 2, 3, 4)]
    top = min(c.top for c in cards)
    assert cards[3].top == top
    _d, prs = make(tmp_path, STAIRS4.replace("@stairs", "@stairs dir=down"), name="down")
    down = [one(prs, 0, f"Stairs {k}") for k in (1, 2, 3, 4)]
    assert [c.height for c in down] == [c.height for c in cards][::-1]
    assert down[0].top == top and len({c.top + c.height for c in down}) == 1


def test_stairs_one_text_size_heading_bold_and_readable_ink(tmp_path):
    _d, prs = make(tmp_path, STAIRS4)
    cards = [one(prs, 0, f"Stairs {k}") for k in (1, 2, 3, 4)]
    all_sizes = {round(s, 1) for c in cards for s in sizes(c)}
    assert len(all_sizes) == 1, all_sizes  # one size for every card
    for c in cards:
        runs = [r for p in c.text_frame.paragraphs for r in p.runs]
        assert runs[0].font.bold and not runs[-1].font.bold
        f = fill_hex(c)
        ratio = (
            (lum(f) + 0.05) / (lum(ink(c)) + 0.05)
            if lum(f) > lum(ink(c))
            else (lum(ink(c)) + 0.05) / (lum(f) + 0.05)
        )
        assert ratio >= 4.5, (f, ink(c), ratio)


def test_stairs_default_fills_run_from_light_to_dark(tmp_path):
    _d, prs = make(tmp_path, STAIRS4)
    ls = [lum(fill_hex(one(prs, 0, f"Stairs {k}"))) for k in (1, 2, 3, 4)]
    assert ls == sorted(ls, reverse=True)  # lighter first, darker last
    assert ink(one(prs, 0, "Stairs 4")) != ink(one(prs, 0, "Stairs 1"))  # the ink follows the fill


def test_stairs_tokens(tmp_path):
    head = (
        HEAD + "style: stairs.fill=#AA0000,#00AA00 stairs.color=#FFFF00 stairs.gap=0.3in stairs.step=0.8in\n"
    )
    _d, prs = make(tmp_path, STAIRS4, head=head)
    cards = [one(prs, 0, f"Stairs {k}") for k in (1, 2, 3, 4)]
    assert [fill_hex(c) for c in cards] == ["AA0000", "00AA00", "AA0000", "00AA00"]
    assert {ink(c) for c in cards} == {"FFFF00"}
    gap = cards[1].left - (cards[0].left + cards[0].width)
    assert abs(gap / 914400 - 0.3) < 0.01
    step = cards[1].height - cards[0].height
    assert abs(step / 914400 - 0.8) < 0.02
    # a box's own fill wins and keeps its place in the list
    own = STAIRS4.replace("## Lời nhắc", "## Lời nhắc {fill=#112233}")
    _d, prs = make(tmp_path, own, head=head, name="own")
    assert fill_hex(one(prs, 0, "Stairs 2")) == "112233"


def test_stairs_with_two_and_six_cards_and_a_lead(tmp_path):
    two = "\n# T\n> Lead line\n@stairs\n## Một\na\n## Hai\nb\n"
    deck, prs = make(tmp_path, two)
    assert not clean(deck) and shapes(prs, 0, "Lead")
    assert one(prs, 0, "Stairs 2").height > one(prs, 0, "Stairs 1").height
    six = "\n# T\n@stairs\n" + "".join(f"## Bậc {k}\nnội dung {k}\n" for k in range(1, 7))
    deck, prs = make(tmp_path, six, name="six")
    assert not clean(deck) and len(shapes(prs, 0, r"Stairs \d")) == 6


def test_stairs_text_never_overflows_its_card(tmp_path):
    long = "\n# T\n@stairs\n" + "".join(
        f"## Bước {k}\n" + "Nội dung khá dài của bước này, có dấu tiếng Việt đầy đủ. " * 2 + "\n"
        for k in range(1, 5)
    )
    deck, _ = make(tmp_path, long)
    assert not [d for d in deck.diagnostics if d.rule in ("overflow", "overlap")], [
        str(d) for d in deck.diagnostics
    ]


# --------------------------------------------------------------------------- @nested


def test_nested_rings_are_concentric_and_shrink_inward(tmp_path):
    deck, prs = make(tmp_path, NESTED3)
    assert not clean(deck), [str(d) for d in clean(deck)]
    rings = [one(prs, 0, f"Nested {k}") for k in (1, 2, 3)]
    cs = [centre(r) for r in rings]
    assert all(abs(c[0] - cs[0][0]) < 200 and abs(c[1] - cs[0][1]) < 200 for c in cs)
    ws = [r.width for r in rings]
    assert ws == sorted(ws, reverse=True) and all(r.width == r.height for r in rings)
    names = [s.name for s in prs.slides[0].shapes]
    assert (
        names.index("Nested 1") < names.index("Nested 2") < names.index("Nested 3")
    )  # the big one is behind
    assert rings[0].left + rings[0].width < prs.slide_width * 0.55  # drawn on the left half


def test_nested_headings_sit_at_the_top_of_their_ring_and_the_innermost_in_the_middle(tmp_path):
    _d, prs = make(tmp_path, NESTED3)
    rings = [one(prs, 0, f"Nested {k}") for k in (1, 2, 3)]
    heads = [one(prs, 0, f"Nested {k} text") for k in (1, 2, 3)]
    assert [tx(h) for h in heads] == ["Trí tuệ nhân tạo", "Học máy", "Học sâu"]
    for r, h in zip(rings[:2], heads[:2], strict=True):
        assert r.top <= h.top and h.top + h.height < r.top + r.height * 0.35  # in the top band
        assert abs(centre(h)[0] - centre(r)[0]) < 200  # centred
        assert h.text_frame.paragraphs[0].runs[0].font.bold
    assert abs(centre(heads[2])[1] - centre(rings[2])[1]) < 1500  # the innermost heading is in the middle
    assert len({round(s, 1) for h in heads for s in sizes(h)}) == 1  # one heading size


def test_nested_ink_is_readable_on_every_ring(tmp_path):
    _d, prs = make(tmp_path, NESTED3)
    for k in (1, 2, 3):
        f, i = fill_hex(one(prs, 0, f"Nested {k}")), ink(one(prs, 0, f"Nested {k} text"))
        hi, lo = max(lum(f), lum(i)), min(lum(f), lum(i))
        assert (hi + 0.05) / (lo + 0.05) >= 4.5, (k, f, i)
    ls = [lum(fill_hex(one(prs, 0, f"Nested {k}"))) for k in (1, 2, 3)]
    assert ls == sorted(ls, reverse=True)  # tints of `secondary`, light (outer) to dark (inner)


def test_nested_bodies_are_an_icon_list_on_the_other_half(tmp_path):
    _d, prs = make(tmp_path, NESTED3)
    rows = [one(prs, 0, f"Nested {k} list") for k in (1, 2, 3)]
    ring = one(prs, 0, "Nested 1")
    assert all(r.left > ring.left + ring.width for r in rows)  # right of the rings
    assert [r.top for r in rows] == sorted(r.top for r in rows)
    first = rows[0].text_frame.paragraphs
    assert pt(first[0]) == "Trí tuệ nhân tạo" and first[0].runs[0].font.bold  # title = the ring heading
    assert pt(first[1]).startswith("Máy làm việc")
    icons = sorted((s for s in prs.slides[0].shapes if s.name.startswith("icon ")), key=lambda s: s.top)
    assert [i.name for i in icons] == ["icon globe", "icon chart", "icon cpu"]
    for ic, row in zip(icons, rows, strict=True):
        assert ic.left + ic.width <= row.left and row.top <= centre(ic)[1] <= row.top + row.height


def test_nested_side_right_flips_the_halves(tmp_path):
    _d, prs = make(tmp_path, NESTED3.replace("@nested", "@nested side=right"))
    ring = one(prs, 0, "Nested 1")
    rows = [one(prs, 0, f"Nested {k} list") for k in (1, 2, 3)]
    assert ring.left > prs.slide_width * 0.4
    assert all(r.left + r.width < ring.left for r in rows)


def test_nested_tokens(tmp_path):
    head = (
        HEAD
        + "style: nested.fill=#AA0000,#00AA00,#0000AA nested.color=#FFFF00 nested.size=26 nested.core=0.5 "
        + "nested.line=#FFFFFF nested.list.title=off nested.share=0.45\n"
    )
    _d, prs = make(tmp_path, NESTED3, head=head)
    rings = [one(prs, 0, f"Nested {k}") for k in (1, 2, 3)]
    assert [fill_hex(r) for r in rings] == ["AA0000", "00AA00", "0000AA"]
    assert abs(rings[2].width / rings[0].width - 0.5) < 0.01
    heads = [one(prs, 0, f"Nested {k} text") for k in (1, 2, 3)]
    assert all(sizes(h) == [26] and ink(h) == "FFFF00" for h in heads)
    assert str(rings[0].line.color.rgb) == "FFFFFF"
    first = one(prs, 0, "Nested 1 list").text_frame.paragraphs
    assert pt(first[0]).startswith("Máy làm việc") and len(first) == 1  # no title line
    assert abs(rings[0].width / prs.slide_width - 0.0) < 1  # (the share limits the ring half)
    assert rings[0].left + rings[0].width <= prs.slide_width * 0.45 + 914400 * 0.6


def test_nested_four_rings_and_a_box_own_fill(tmp_path):
    four = NESTED3 + "## Mô hình lớn\nHàng tỷ tham số\n"
    four = four.replace("## Học sâu {icon=cpu}", "## Học sâu {icon=cpu fill=#445566}")
    deck, prs = make(tmp_path, four)
    assert not clean(deck)
    assert len(shapes(prs, 0, r"Nested \d")) == 4
    assert fill_hex(one(prs, 0, "Nested 3")) == "445566"
    assert [r.width for r in (one(prs, 0, f"Nested {k}") for k in (1, 2, 3, 4))] == sorted(
        (one(prs, 0, f"Nested {k}").width for k in (1, 2, 3, 4)), reverse=True
    )


def test_nested_without_icons_has_no_icon_column(tmp_path):
    _d, prs = make(tmp_path, re.sub(r" \{icon=\w+\}", "", NESTED3))
    assert not [s for s in prs.slides[0].shapes if s.name.startswith("icon ")]
    assert (
        one(prs, 0, "Nested 1 list").left
        < one(prs, 0, "Nested 1").left + one(prs, 0, "Nested 1").width + 914400 * 2
    )


# --------------------------------------------------------------------------- @flow disc


def test_flow_disc_is_a_row_of_discs_joined_by_lines_with_heads(tmp_path):
    deck, prs = make(tmp_path, FLOW3)
    assert not clean(deck), [str(d) for d in clean(deck)]
    discs = [one(prs, 0, f"Flow {k} disc") for k in (1, 2, 3)]
    assert all(d.width == d.height for d in discs) and len({d.top for d in discs}) == 1
    lines = [s for s in prs.slides[0].shapes if re.fullmatch(r"Flow \d line", s.name)]
    assert len(lines) == 2
    for ln in lines:
        assert "tailEnd" in ln._element.xml or "headEnd" in ln._element.xml  # an arrowhead
    assert [d.left for d in discs] == sorted(d.left for d in discs)
    # the number badge stands in for a missing icon
    assert [tx(d) for d in discs] == ["1", "2", "3"]


def test_flow_disc_heading_is_bold_under_the_disc_and_the_body_follows(tmp_path):
    _d, prs = make(tmp_path, FLOW3)
    for k, head in ((1, "Một"), (2, "Hai"), (3, "Ba")):
        d, t = one(prs, 0, f"Flow {k} disc"), one(prs, 0, f"Flow {k} text")
        assert t.top >= d.top + d.height
        paras = t.text_frame.paragraphs
        assert pt(paras[0]) == head and paras[0].runs[0].font.bold and not paras[1].runs[0].font.bold
        assert abs(centre(t)[0] - centre(d)[0]) < 200


def test_flow_disc_icons_replace_the_number(tmp_path):
    _d, prs = make(tmp_path, FLOW5)
    assert tx(one(prs, 0, "Flow 2 disc")) == ""
    icons = [s.name for s in prs.slides[0].shapes if s.name.startswith("icon ")]
    assert set(icons) == {"icon database", "icon chat", "icon search", "icon document", "icon bolt"}


def test_flow_disc_first_box_above_sits_over_the_second_disc_joined_by_a_vertical_line(tmp_path):
    deck, prs = make(tmp_path, FLOW5)
    assert not clean(deck)
    store, second = one(prs, 0, "Flow 1 disc"), one(prs, 0, "Flow 3 disc")
    row = one(prs, 0, "Flow 2 disc")
    assert store.top + store.height < row.top  # above the row
    assert abs(centre(store)[0] - centre(second)[0]) < 200 or abs(centre(store)[0] - centre(row)[0]) < 200
    vline = one(prs, 0, "Flow 1 line")
    assert (
        vline.width == 0
        and vline.top >= store.top + store.height - 10
        and vline.top + vline.height <= row.top + 10
    )
    assert len([s for s in prs.slides[0].shapes if re.fullmatch(r"Flow \d disc", s.name)]) == 5
    # the row discs share one top; the store text is beside its disc
    tops = {one(prs, 0, f"Flow {k} disc").top for k in (2, 3, 4, 5)}
    assert len(tops) == 1
    t = one(prs, 0, "Flow 1 text")
    assert t.left >= store.left + store.width


def test_flow_disc_above_attribute_on_the_at_line(tmp_path):
    text = FLOW5.replace("{.above icon=database}", "{icon=database}").replace(
        "@flow disc", "@flow disc above=1"
    )
    _d, prs = make(tmp_path, text)
    assert one(prs, 0, "Flow 1 disc").top < one(prs, 0, "Flow 2 disc").top


def test_flow_disc_tokens(tmp_path):
    head = (
        HEAD
        + "style: flow.disc.size=1.1in flow.disc.fill=#AA0000,#00AA00 flow.disc.color=#FFFF00 "
        + "flow.line=#0000AA flow.line.w=5pt\n"
    )
    _d, prs = make(tmp_path, FLOW3, head=head)
    discs = [one(prs, 0, f"Flow {k} disc") for k in (1, 2, 3)]
    assert [fill_hex(d) for d in discs] == ["AA0000", "00AA00", "AA0000"]
    assert all(abs(d.width / 914400 - 1.1) < 0.01 for d in discs)
    assert {ink(d) for d in discs} == {"FFFF00"}
    ln = one(prs, 0, "Flow 1 line")
    assert str(ln.line.color.rgb) == "0000AA" and abs(ln.line.width.pt - 5) < 0.1


def test_flow_disc_with_two_and_seven_boxes_and_long_vietnamese_text(tmp_path):
    two = "\n# T\n@flow disc\n## Đầu vào\nnhận dữ liệu\n## Đầu ra\ntrả kết quả\n"
    deck, prs = make(tmp_path, two)
    assert not clean(deck) and len(shapes(prs, 0, r"Flow \d disc")) == 2
    seven = "\n# T\n@flow disc\n" + "".join(
        f"## Bước {k}\nMô tả khá dài của bước {k} bằng tiếng Việt có dấu\n" for k in range(1, 8)
    )
    deck, prs = make(tmp_path, seven, name="seven")
    assert not [d for d in deck.diagnostics if d.rule in ("overflow", "overlap")], [
        str(d) for d in deck.diagnostics
    ]
    assert len(shapes(prs, 0, r"Flow \d disc")) == 7


def test_a_plain_flow_still_draws_cards_and_arrows(tmp_path):
    _d, prs = make(tmp_path, "\n# T\n@3 flow\n## a\nx\n## b\ny\n## c\nz\n")
    assert not shapes(prs, 0, r"Flow \d disc")


# --------------------------------------------------------------------------- build feedback


@pytest.mark.parametrize(
    "text, start",
    [
        (STAIRS4, "slide 1: stairs 4 (up)"),
        (NESTED3, "slide 1: nested 3 rings (left)"),
        (FLOW5, "slide 1: flow of 5 discs, side node above"),
        (FLOW3, "slide 1: flow of 3 discs"),
        (CYCLE_ICONS, "slide 1: cycle 4 nodes (cw), 4 icons in nodes, centre label"),
    ],
)
def test_the_fit_line_names_the_form(text, start, tmp_path):
    from slidemark.fit import fit_lines
    from slidemark.layout import layout_slide
    from slidemark.template import deck_theme

    deck = parse(HEAD + text)
    theme, _ = deck_theme(deck, str(tmp_path))
    placed = [layout_slide(deck.slides[0], deck, theme, 0)]
    line = fit_lines(deck, placed, theme)[0]
    assert line.startswith(start), line


@pytest.mark.parametrize("text", [STAIRS4, NESTED3, FLOW5, CYCLE_ICONS])
def test_the_design_facts_count_the_word_as_a_decided_form(text):
    deck = parse(HEAD + text)
    assert stated_groups(deck.slides[0])["form"]


@pytest.mark.parametrize(
    "token, needs",
    [("stairs.fill=#FF0000", STAIRS4), ("nested.fill=#FF0000", NESTED3), ("flow.line=#FF0000", FLOW3)],
)
def test_style_tokens_need_their_form(token, needs, tmp_path):
    deck, _ = make(tmp_path, "\n# T\n- a\n- b\n", head=HEAD + f"style: {token}\n")
    got = warns(deck, "attr-ignored")
    assert got and f"style: {token.split('=')[0]}=" in got[0].message, token
    deck, _ = make(tmp_path, needs, head=HEAD + f"style: {token}\n", name="ok")
    assert not warns(deck, "attr-ignored"), token


def test_a_box_attribute_a_new_form_cannot_honour_is_reported(tmp_path):
    deck, _ = make(tmp_path, STAIRS4.replace("## Lời nhắc", "## Lời nhắc {size=40 rotate=5}"))
    got = warns(deck, "attr-ignored")
    assert {d.message for d in got} == {
        "size= on a form box is not honoured",
        "rotate= on a form box is not honoured",
    }


# --------------------------------------------------------------------------- importer round trip


def _roundtrip(tmp_path, text: str, name: str = "rt") -> str:
    out = tmp_path / f"{name}.pptx"
    build(HEAD + text, out)
    md, _diags = import_pptx(out)
    return md


@pytest.mark.parametrize(
    "text, tokens, lines",
    [
        (STAIRS4, "@stairs", ["## Dùng sẵn", "Mua dịch vụ có sẵn", "## Huấn luyện"]),
        (
            STAIRS4.replace("@stairs", "@stairs dir=down"),
            "@stairs dir=down",
            ["## Lời nhắc", "Viết lời nhắc và mẫu"],
        ),
        (
            NESTED3,
            "@nested",
            [
                "## Trí tuệ nhân tạo {icon=globe}",
                "## Học máy {icon=chart}",
                "## Học sâu {icon=cpu}",
                "Mạng nơ-ron nhiều tầng",
            ],
        ),
        (NESTED3.replace("@nested", "@nested side=right"), "@nested side=right", ["## Học máy {icon=chart}"]),
        (
            CYCLE_ICONS,
            '@cycle center="Vòng lặp agent"',
            ["## Quan sát {icon=search}", "## Phản hồi {icon=check}", "Kiểm tra kết quả"],
        ),
        (CYCLE4, "@cycle", ["## Lập kế hoạch", "Chia thành bước nhỏ"]),
        (
            FLOW5,
            "@flow disc",
            [
                "## Kho tài liệu {.above icon=database}",
                "## Câu hỏi {icon=chat}",
                "## Trả lời {icon=bolt}",
                "Mô hình trả lời",
            ],
        ),
        (FLOW3, "@flow disc", ["## Một", "việc đầu", "## Ba"]),
    ],
)
def test_importer_folds_the_named_shapes_back(text, tokens, lines, tmp_path):
    md = _roundtrip(tmp_path, text)
    got = md.splitlines()
    assert tokens in got, (tokens, md)
    for ln in lines:
        assert any(g.strip() == ln for g in got), (ln, md)
    again = parse(md)
    assert not [
        d for d in again.diagnostics if d.rule.endswith("-skipped") or d.rule in ("bad-attr", "attr-ignored")
    ]


@pytest.mark.parametrize("text", [STAIRS4, NESTED3, CYCLE_ICONS, FLOW5, FLOW3])
def test_a_reimported_form_builds_the_same_shapes(text, tmp_path):
    a = tmp_path / "a.pptx"
    build(HEAD + text, a)
    md, _ = import_pptx(a)
    b = tmp_path / "b.pptx"
    build(md, b)

    def names(p):
        return [
            s.name
            for s in Presentation(str(p)).slides[0].shapes
            if not s.name.startswith(("Footer", "Slide"))
        ]

    assert names(a) == names(b)


# --------------------------------------------------------------------------- the example deck


EXAMPLE = ROOT / "examples" / "26-ai-overview.md"


def test_the_ai_overview_example_builds_with_no_warnings_and_every_form(tmp_path):
    deck = build(EXAMPLE.read_text(encoding="utf-8"), tmp_path / "ex.pptx", base_dir=EXAMPLE.parent)
    bad = [d for d in deck.diagnostics if d.level in ("warning", "error")]
    assert not bad, [str(d) for d in bad]
    assert len(deck.slides) == 15
    words = {w for s in deck.slides for w in s.classes}
    assert {
        "nested",
        "timeline",
        "vs",
        "iconlist",
        "cycle",
        "flow",
        "disc",
        "proscons",
        "stairs",
        "agenda",
    } <= words
    assert sum(1 for s in deck.slides if any("steps" in getattr(e, "classes", []) for e in s.elements)) == 2
    assert deck.slides[0].background == "primary" and deck.slides[-1].background == "primary"
    assert deck.slides[8].attrs["center"] == "Vòng lặp agent"


def test_the_ai_overview_example_round_trips_every_composed_form(tmp_path):
    out = tmp_path / "ex.pptx"
    build(EXAMPLE.read_text(encoding="utf-8"), out, base_dir=EXAMPLE.parent)
    md, _ = import_pptx(out)
    for token in ("@nested", "@cycle", "@stairs", "@flow disc", "@timeline", "@vs", "@agenda"):
        assert any(ln.startswith(token) for ln in md.splitlines()), token
    again = parse(md)
    assert len(again.slides) == 15


# --------------------------------------------------------------------------- fuzz

_WORD = st.text(
    alphabet=st.characters(min_codepoint=0x20, max_codepoint=0x24F, blacklist_characters="#{}`|*_\\<>[]\n\r"),
    min_size=0,
    max_size=24,
)


@settings(
    max_examples=25,
    deadline=None,
    derandomize=True,
    suppress_health_check=[HealthCheck.too_slow, HealthCheck.function_scoped_fixture],
)
@given(
    word=st.sampled_from(
        [
            "@stairs",
            "@stairs dir=down",
            "@nested",
            "@nested side=right",
            "@flow disc",
            "@flow disc above=1",
            '@cycle center="x"',
        ]
    ),
    heads=st.lists(_WORD, min_size=0, max_size=9),
    body=_WORD,
    icon=st.sampled_from(["", " {icon=bolt}", " {icon=nope}", " {.above}", " {fill=#12345}"]),
)
def test_fuzz_a_form_slide_never_raises(word, heads, body, icon, tmp_path_factory):
    text = f"\n# T\n{word}\n" + "".join(f"## {h or 'x'}{icon}\n{body}\n" for h in heads)
    out = tmp_path_factory.mktemp("fz") / "d.pptx"
    deck = build(HEAD + text, out)  # never raises: a diagnostic with a hint instead
    assert out.exists()
    for d in deck.diagnostics:
        if d.level in ("warning", "error"):
            assert d.hint or d.rule.startswith("design"), str(d)
