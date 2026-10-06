"""Chart values written with thousands separators keep them in the labels (#,##0), in any locale."""

from __future__ import annotations

from pptx import Presentation

from slidemark import build
from slidemark.parser import parse


def _chart(src: str):
    deck = parse(src)
    return deck.slides[0].elements[0]


def test_grouped_values_set_the_option():
    assert _chart('# T\n```line\n,a,b\nX,"1,240","2,500"\n```\n').options.get("grouped")
    assert _chart("lang: vi\n\n# T\n```line\n,T4,T5\nĐơn,2.100,12.600\n```\n").options.get("grouped")


def test_plain_and_decimal_values_do_not():
    assert not _chart("# T\n```line\n,a,b\nX,1240,2500\n```\n").options.get("grouped")
    assert not _chart('lang: vi\n\n# T\n```column\n,a,b\nX,"3,6","3,9"\n```\n').options.get("grouped")
    assert not _chart("# T\n```column\n,a,b\nX,12%,30%\n```\n").options.get("grouped")


def test_render_uses_thousands_format(tmp_path):
    src = tmp_path / "c.md"
    src.write_text(
        "lang: vi\n\n# Đơn hàng\n```line {labels=on}\n,T4,T5\nĐơn,2.100,12.600\n```\n", encoding="utf-8"
    )
    build(src, tmp_path / "c.pptx")
    chart = next(s.chart for s in Presentation(tmp_path / "c.pptx").slides[0].shapes if s.has_chart)
    assert list(chart.plots[0].series[0].values) == [2100.0, 12600.0]
    assert chart.plots[0].data_labels.number_format == "#,##0"
    assert chart.value_axis.tick_labels.number_format == "#,##0"


def test_bar_axes_start_at_zero(tmp_path):
    src = tmp_path / "z.md"
    src.write_text(
        "# T\n```column\n,a,b\nX,3.6,3.9\n```\n\n# L\n```line\n,a,b\nX,3.6,3.9\n```\n"
        "\n# M\n```column {min=3}\n,a,b\nX,3.6,3.9\n```\n",
        encoding="utf-8",
    )
    build(src, tmp_path / "z.pptx")
    slides = Presentation(tmp_path / "z.pptx").slides
    axes = [next(s.chart for s in sl.shapes if s.has_chart).value_axis for sl in slides]
    assert axes[0].minimum_scale == 0
    assert axes[1].minimum_scale != 0  # lines may zoom
    assert axes[2].minimum_scale == 3  # explicit min wins


def test_zero_base_does_not_come_back_as_min(tmp_path):
    from slidemark.importer import import_pptx

    src = tmp_path / "z.md"
    src.write_text("# T\n```column\n,a,b\nX,3.6,3.9\n```\n", encoding="utf-8")
    build(src, tmp_path / "z.pptx")
    assert "min=" not in import_pptx(tmp_path / "z.pptx")[0]
