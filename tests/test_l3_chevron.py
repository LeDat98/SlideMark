"""Chevron rows: a lone row is tall and sits at the optical center; above a table its text grows."""

from __future__ import annotations

from pptx import Presentation

from slidemark import build, parse
from slidemark.ir import Placed, Shape, Table
from slidemark.layout import measure
from slidemark.layout.engine import layout_slide
from slidemark.template import deck_theme
from slidemark.theme import LayoutTokens

EMU = 914400
HEAD = """theme: none
colors: bg=#FBF8F3 fg=#1F2A30 primary=#0F5257 accent=#E4572E surface=#FFFFFF border=#D9D2C5 muted=#5E6B70
fonts: heading="Montserrat" body="Inter"
lang: vi
footer: Công ty CP Logistics Sao Mai · Nội bộ
num: on

"""
LONE = (
    HEAD
    + """# Lộ trình ba bước
> Kết nối dữ liệu trước, tối ưu sau, mở rộng hệ sinh thái cuối cùng
@chevron
## Bước 1 Kết nối
- Số hóa 100% đơn hàng
## Bước 2 Tối ưu
- Điều phối xe bằng AI
## Bước 3 Mở rộng
- Nền tảng cho đối tác
※ Mỗi bước kéo dài khoảng 12 tháng
"""
)
TABLE = """theme: jp-business
lang: ja

# 施策の全体像
> 3 つの施策を順に展開する
@3 chevron
## Step 1 基盤整備
WMS 刷新・データ統合
## Step 2 最適化
AI 配車・需要予測
## Step 3 連携
共同配送・荷主連携
@end
| 施策 | 主な内容 | 投資額 |
|-|-|-|
| WMS 刷新 | クラウド WMS、ハンディ端末 | 8.0 億円 |
| AI 配車 | 配車最適化エンジン | 5.5 億円 |
| 需要予測 | SKU 別の週次予測 | 3.0 億円 |
"""


def lay(src: str, n: int) -> list[Placed]:
    deck = parse(src)
    theme, _ = deck_theme(deck)
    measure.set_tokens(theme.layout)
    return layout_slide(deck.slides[n - 1], deck, theme, n)


def chevrons(placed: list[Placed]) -> list[Placed]:
    return [p for p in placed if isinstance(p.element, Shape) and p.element.shape == "chevron"]


def size(p: Placed) -> float:
    return (p.style.font_size or 0) * p.font_scale


def body_of(placed: list[Placed]) -> tuple[int, int]:
    lead = next(p for p in placed if getattr(p.element, "role", None) == "lead")
    foot = next(p for p in placed if getattr(p.element, "role", None) == "footnote")
    top = lead.y + lead.h + round(0.25 * EMU)
    return top, foot.y


def test_lone_row_is_tall_and_near_the_optical_center():
    placed = lay(LONE, 1)
    chev = chevrons(placed)
    top, bottom = body_of(placed)
    body = bottom - top
    assert len(chev) == 3
    assert chev[0].h >= 0.45 * body
    center = chev[0].y + chev[0].h / 2
    assert 0.4 * body <= center - top <= 0.5 * body  # in the middle band, slightly above the middle
    assert chev[0].y - top >= 0.05 * body  # not top-packed


def test_lone_row_height_is_a_token():
    lt = LayoutTokens()
    assert 0 < lt.chevron_lone_h <= 1 and 0 <= lt.chevron_lone_top <= 1


def test_lone_row_text_is_large_and_keeps_words_whole():
    placed = lay(LONE, 1)
    chev = chevrons(placed)
    assert size(chev[0]) >= 20
    assert all(size(c) == size(chev[0]) for c in chev)
    from slidemark.layout.engine import _chevron_text_w, _longest_word_em
    from slidemark.layout.grid import Rect

    for c in chev:
        avail = (
            _chevron_text_w(Rect(c.x, c.y, c.w, c.h), c.style, c.element, c.element.attrs.get("adj")) / 12700
        )
        assert _longest_word_em(c.element.paragraphs) * size(c) <= avail  # no word breaks


def test_chevron_above_table_text_grows_and_table_follows():
    placed = lay(TABLE, 1)
    chev = chevrons(placed)
    (tbl,) = [p for p in placed if isinstance(p.element, Table)]
    base = LayoutTokens(chevron_table_fill=0.0)
    assert base.chevron_table_fill == 0.0
    assert size(chev[0]) >= 16  # was ~14pt
    assert size(tbl) <= size(chev[0]) + 0.01  # table text stays <= chevron text
    assert tbl.y >= chev[0].y + chev[0].h  # no overlap
    assert chev[0].h <= 1.6 * EMU


def test_chevron_above_table_grows_text_vs_switched_off():
    on = size(chevrons(lay(TABLE, 1))[0])
    deck = parse(TABLE.replace("lang: ja\n", "lang: ja\nstyle: layout.chevron_table_fill=0\n"))
    theme, _ = deck_theme(deck)
    measure.set_tokens(theme.layout)
    off = size(chevrons(layout_slide(deck.slides[0], deck, theme, 1))[0])
    assert on > off


def test_render_reopens_with_chevrons_inside_the_slide(tmp_path):
    out = tmp_path / "c.pptx"
    build(LONE, out)
    prs = Presentation(str(out))
    shapes = [s for s in prs.slides[0].shapes if s._element.xpath('.//a:prstGeom[@prst="chevron"]')]
    assert len(shapes) == 3
    for s in shapes:
        assert s.top + s.height <= prs.slide_height
        assert s.text_frame.text.strip()


JP4 = """theme: jp-business
lang: ja

# 今後のスケジュール
@4 chevron
## 11月
FAQ公開
## 12月
チャット試行
## 2027年1月
システム更新
## 2027年3月
効果検証
@end
"""


def test_lone_row_of_short_text_is_capped_by_width_and_text():
    chev = chevrons(lay(JP4, 1))
    lt = LayoutTokens()
    assert len(chev) == 4
    assert chev[0].h <= lt.chevron_lone_cap_aspect * min(c.w for c in chev) + 2  # no giant arrows
    assert chev[0].h <= 0.3 * 7.5 * EMU
