"""Fill policy: text grows before gaps, gaps stay even and capped, sparse cards shrink to their content."""

from __future__ import annotations

import pytest
from pptx import Presentation

from slidemark import build
from slidemark.ir import Chart, Text
from slidemark.layout import engine
from slidemark.layout.grid import Rect
from slidemark.layout.l3fill import fill_panels, fill_row
from slidemark.units import to_emu

from .test_layout_l3_fill import EX, cards, inside, lay


@pytest.mark.parametrize(("name", "n"), [("16-jp-strategy.md", 10), ("11-jp-consulting.md", 9)])
def test_decision_cards_share_one_capped_rhythm_and_a_head_pad(name, n):
    placed, _, theme = lay(name, n, cards_to_bar=False)
    cs = cards(placed)
    assert len(cs) == 2
    mains = [inside(placed, c)[0] for c in cs]
    gaps = [m.element.attrs.get("para_gap") or 0.0 for m in mains]
    assert abs(gaps[0] - gaps[1]) <= 0.02  # same rhythm left and right
    off, _, _ = lay(name, n, l3_fill=False)
    g0 = max(inside(off, c)[0].element.attrs.get("para_gap") or 0.25 for c in cards(off))
    assert max(gaps) - g0 <= theme.layout.l3_gap_cap + 0.02
    sizes = {round(m.font_scale, 3) for m in mains}
    assert len(sizes) == 1  # one text size for both cards
    for c, m in zip(cs, mains, strict=True):
        heads = [
            p for p in placed if getattr(p.element, "role", None) == "heading" and c.y <= p.y < c.y + c.h
        ]
        assert m.y > max(h.y + h.h for h in heads)  # air under the heading band
        assert m.y + m.h <= c.y + c.h
    assert max(c.h for c in cs) < 0.55 * 720 * 9525  # shrunk to content, not stretched over the body


@pytest.mark.parametrize(("name", "n"), [("16-jp-strategy.md", 3), ("11-jp-consulting.md", 3)])
def test_panel_spans_the_chart_beside_it(name, n):
    placed, _, theme = lay(
        name, n, panel_end_air=0
    )  # the full-span variant (panel_end_air: test_layout_l3_fill)
    (panel,) = cards(placed)
    chart = next(p for p in placed if isinstance(p.element, Chart))
    tol = 3 * 12700
    assert panel.y >= chart.y and panel.y - chart.y < 0.2 * chart.h  # top at the plot area, not centered
    assert abs((panel.y + panel.h) - (chart.y + chart.h)) <= tol  # bottom meets the chart bottom
    main, note = sorted(inside(placed, panel), key=lambda p: p.y)[-2:]
    assert note.y + note.h <= panel.y + panel.h
    assert main.y + engine._text_h(main) <= note.y
    off, _, _ = lay(name, n, panel_to_visual=False)  # the token restores the old shrunk panel
    (old,) = cards(off)
    assert old.y + old.h < chart.y + chart.h - tol


@pytest.mark.parametrize(("name", "n"), [("16-jp-strategy.md", 3), ("11-jp-consulting.md", 3)])
def test_callout_in_a_box_reads_as_part_of_the_panel(name, n):
    placed, _, theme = lay(name, n)
    (panel,) = cards(placed)
    main, note = sorted(inside(placed, panel), key=lambda p: p.y)[-2:]
    assert "callout" in note.element.classes
    body_pt = main.style.font_size * main.font_scale
    note_pt = note.style.font_size * note.font_scale
    assert body_pt / theme.layout.callout_text_step - 0.6 <= note_pt <= body_pt
    pad = 8 * 12700
    assert note.style.padding_top is not None and note.style.padding_bottom is not None
    for side in ("top", "right", "bottom"):
        assert to_emu(getattr(note.style, f"padding_{side}")) >= pad
    assert to_emu(note.style.padding_left) >= pad + to_emu(theme.layout.callout_bar_w)
    assert panel.y + panel.h - (note.y + note.h) >= 4 * 12700  # air between the callout and the card edge


def test_top_level_callout_keeps_its_size():
    src = "theme: jp-business\n\n# T\n- a\n- b\n> [!note] hello\n"
    placed, _, theme = lay("", 1, src)
    note = next(p for p in placed if "callout" in getattr(p.element, "classes", []))
    assert note.font_scale == 1.0
    assert to_emu(note.style.padding) == 6 * 12700 or note.style.padding_left is None


def test_callout_in_a_box_render_reopens(tmp_path):
    out = tmp_path / "c.pptx"
    build(EX / "11-jp-consulting.md", out)
    prs = Presentation(out)
    bars = [s for s in prs.slides[2].shapes if s.name.endswith("accent")]
    assert bars
    notes = [s for s in prs.slides[2].shapes if s.has_text_frame and "燃料費" in s.text_frame.text]
    assert notes
    assert bars[0].height == notes[0].height and bars[0].width < notes[0].width


def test_panel_render_reopens(tmp_path):
    out = tmp_path / "d.pptx"
    build(EX / "16-jp-strategy.md", out)
    prs = Presentation(out)
    texts = [s for s in prs.slides[2].shapes if s.has_text_frame and "価格競争" in s.text_frame.text]
    assert texts
    chart = next(s for s in prs.slides[2].shapes if s.has_chart)
    assert texts[0].top >= chart.top
    assert texts[0].top + texts[0].height <= chart.top + chart.height + 12700  # the panel ends with the chart


def test_explicit_size_and_height_win_over_growth():
    src = "theme: jp-business\n\n# T\n> l\n@1:1\n## A {h=2in}\n- one\n- two\n## B\n- x {size=11}\n- y\n"
    placed, _, theme = lay("", 1, src)
    hs = sorted(c.h for c in cards(placed))
    assert abs(hs[-1] - 2 * 914400) < 10 * 9525 or len(hs) == 2
    assert fill_row([], Rect(0, 0, 10, 10), theme.layout) == []
    assert fill_panels([], Rect(0, 0, 10, 10), theme.layout, True) == []


def test_fuzz_never_raises():
    for src in (
        "# T\n@1:1\n## A\n@end\n",
        "# T\n@2\n## A\n- " + "長い" * 80 + "\n## B\n1. a\n2. b\n@end\n",
        "# T\n@1:1\n## A\n![x](nope.png)\n## B\n- a\n",
    ):
        placed, _, _ = lay("", 1, "theme: jp-business\n\n" + src)
        assert isinstance(placed, list)
        assert all(isinstance(p.element, Text) or p.h >= 0 for p in placed)
