"""Attribute-only blocks: `{x= y= w= h= shape= fill=}` with no content after it is a drawn shape."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE

from slidemark.build import build
from slidemark.parser import parse

HEAD = "theme: jp-business\nlang: en\n"
CHEV = "{x=24.4% y=36% w=2.4% h=8% shape=chevron fill=accent}"


def _shapes(md: str, tmp: Path):
    out = tmp / "s.pptx"
    deck = build(HEAD + md, out)
    return Presentation(out).slides[0].shapes, deck


def test_attribute_only_line_is_a_shape_element():
    deck = parse(HEAD + f"# T\n@free\n{CHEV}\n{CHEV}\n{CHEV}\n")
    els = deck.slides[0].elements
    assert [e.type for e in els] == ["shape"] * 3
    assert els[0].style.shape == "chevron" and els[0].style.fill == "accent"
    assert not [d for d in deck.diagnostics if d.rule == "dangling-attrs"]


def test_a_block_with_content_after_it_is_not_a_shape():
    deck = parse(HEAD + "# T\n@free\n{x=10% y=30% w=10% h=10% fill=accent}\ntext\n")
    assert [e.type for e in deck.slides[0].elements] == ["text"]


def test_shape_blocks_work_on_any_slide_and_inside_a_box(tmp_path):
    md = "# T\n- x\n- y\n{x=60% y=10% w=10% h=10% fill=primary}\n"
    shapes, _ = _shapes(md, tmp_path)
    assert any(s.name == "Shape 1" for s in shapes)
    md = HEAD + "# T\n## a\n- x\n{x=1in y=1in w=1in h=1in fill=primary}\n@end\n## b\n- y\n"
    box = parse(md).slides[0].elements[0]
    assert [c.type for c in box.children] == ["text", "shape"]
    build(md, tmp_path / "bx.pptx")
    # (an attribute line right before a `##` heading styles that box, as before)
    deck = parse(HEAD + "# T\n## a\n- x\n{fill=primary}\n## b\n- y\n")
    assert [e.type for e in deck.slides[0].elements] == ["container", "container"]


def test_three_chevrons_render_as_native_shapes_named_shape_n(tmp_path):
    shapes, deck = _shapes(
        f"# T\n@free\n{CHEV}\n{CHEV.replace('24.4', '47.9')}\n{CHEV.replace('24.4', '71.4')}\n", tmp_path
    )
    chev = [s for s in shapes if s.shape_type == 1 and s.auto_shape_type == MSO_SHAPE.CHEVRON]
    assert [s.name for s in chev] == ["Shape 1", "Shape 2", "Shape 3"]
    assert len({s.left for s in chev}) == 3
    xml = chev[0]._element.xml
    assert 'prst="chevron"' in xml


def test_look_keys_are_honoured(tmp_path):
    look = "shape=ellipse fill=#336699 line=#FF0000 line.w=3pt opacity=0.5 rotate=30 shadow=on"
    md = f"# T\n@free\n{{x=10% y=20% w=20% h=20% {look}}}\n"
    shapes, _ = _shapes(md, tmp_path)
    s = next(s for s in shapes if s.name == "Shape 1")
    assert s.auto_shape_type == MSO_SHAPE.OVAL
    assert s.rotation == 30
    xml = s._element.xml
    assert "336699" in xml and "FF0000" in xml and 'w="38100"' in xml
    assert "alpha" in xml and "outerShdw" in xml


def test_default_shape_is_a_rectangle_and_radius_rounds_it(tmp_path):
    shapes, _ = _shapes(
        "# T\n@free\n{x=10% y=20% w=20% h=20% fill=primary}\n"
        "{x=40% y=20% w=20% h=20% fill=primary radius=12}\n",
        tmp_path,
    )
    kinds = [s.auto_shape_type for s in shapes if s.name.startswith("Shape")]
    assert kinds == [MSO_SHAPE.RECTANGLE, MSO_SHAPE.ROUNDED_RECTANGLE]


def test_a_bare_line_is_a_connector(tmp_path):
    md = (
        "# T\n@free\n{x=10% y=50% w=30% h=0 shape=line line=accent line.w=2pt}\n"
        "{x=50% y=30% w=0 h=30% shape=line line=primary head=arrow}\n"
    )
    shapes, _ = _shapes(md, tmp_path)
    lines = [s for s in shapes if s.shape_type == 9]
    assert len(lines) == 2
    assert lines[0].height == 0 and lines[1].width == 0
    assert 'w="25400"' in lines[0]._element.xml
    assert "tailEnd" in lines[1]._element.xml


def test_line_without_a_color_warns_and_bad_values_never_raise(tmp_path):
    deck = parse(
        HEAD
        + "# T\n@free\n{x=1in y=1in w=1in h=0 shape=line}\n"
        + "{x=1in y=1in w=1in h=1in fill=primary line.w=thick shape=nope head=x}\n"
    )
    rules = {d.rule for d in deck.diagnostics}
    assert "bad-attr" in rules
    build(HEAD + "# T\n@free\n{x=1in y=1in w=1in h=0 shape=line}\n", tmp_path / "b.pptx")


def test_fit_line_counts_shapes_and_prints_the_free_area(tmp_path):
    md = HEAD + f"# T\n@free\n{CHEV}\n{{x=10% y=50% w=20% h=20%}}\ntext\n"
    out = tmp_path / "f.pptx"
    deck = build(md, out)
    lines = deck.attrs.get("_fit") if hasattr(deck, "attrs") else None
    from slidemark.cli import main

    md_path = tmp_path / "f.md"
    md_path.write_text(md, encoding="utf-8")
    import contextlib
    import io

    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        main(["build", str(md_path), "-o", str(tmp_path / "g.pptx")])
    text = buf.getvalue()
    m = re.search(
        r"free 2 blocks, 2 pinned, 1 shape, free area (\d+\.\d),(\d+\.\d) (\d+\.\d)x(\d+\.\d)in", text
    )
    assert m, text
    assert float(m.group(3)) > 11 and float(m.group(4)) > 4
    del lines


def test_opens_as_valid_ooxml(tmp_path):
    out = tmp_path / "v.pptx"
    build(HEAD + f"# T\n@free\n{CHEV}\n{{x=10% y=50% w=30% h=0 shape=line line=accent}}\n", out)
    with zipfile.ZipFile(out) as z:
        assert z.testzip() is None


@pytest.mark.parametrize(
    "line",
    [
        "{x= y= w= h= shape=chevron}",
        "{x=1in shape=}",
        "{w=1e999 h=-5 fill=}",
        "{x=10% y=10% w=10% h=10% shape=line}",
        "{x=10% y=10% w=0 h=0 shape=line line=primary}",
        "{x=10% y=10% w=-1in h=1in fill=accent}",
        "{x=10% y=10% w=1in h=1in fill=accent line.w=}",
        "{x=10% line.w=2 head=arrow}",
        "{shape=hexagon}",
    ],
)
def test_odd_shape_blocks_never_raise(line, tmp_path):
    md = f"# T\n@free\n{line}\n"
    deck = parse(md)
    assert deck is not None
    build(md, tmp_path / "o.pptx")
