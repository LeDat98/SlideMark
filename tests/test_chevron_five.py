"""A row of >= 5 chevrons keeps every bold heading ("Tháng 11") on one line."""

from __future__ import annotations

from pptx import Presentation

from slidemark import build
from slidemark.layout import measure

SRC = """theme: none
colors: bg=#FFF7ED fg=#431407 primary=#EA580C accent=#C2410C
fonts: heading="Lora" body="Inter"
lang: vi

# Lộ trình ra mắt
@chevron
## Tháng 11
- Thử nghiệm vị
## Tháng 12
- Sản xuất lô đầu
## Tháng 1
- Ra mắt tại TP.HCM
## Tháng 3
- Mở rộng toàn quốc
## Tháng 6
- Đánh giá kết quả
"""


def test_chevron_headings_stay_on_one_line(tmp_path):
    src = tmp_path / "c.md"
    src.write_text(SRC, encoding="utf-8")
    build(src, tmp_path / "c.pptx")
    shapes = [s for s in Presentation(tmp_path / "c.pptx").slides[0].shapes if s.has_text_frame]
    chev = [s for s in shapes if s.text_frame.text.startswith("Tháng")]
    assert len(chev) == 5
    for s in chev:
        head = s.text_frame.paragraphs[0]
        size = head.runs[0].font.size.pt
        adj = s.adjustments[0] if len(s.adjustments) else 0.3
        text_w = (s.width - 2 * adj * min(s.width, s.height)) / 12700  # pt, both point depths removed
        need = measure.text_em(head.text, bold=True) * size * 1.15  # a fallback font runs ~15% wider
        assert need <= text_w, (head.text, need, text_w)
