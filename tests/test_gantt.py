"""``{.gantt}`` tables: filled body cells become native bars over the table."""

from __future__ import annotations

from pptx import Presentation

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.ir import Shape, Table
from slidemark.layout import css, layout_slide, measure
from slidemark.layout.gantt import table_boxes
from slidemark.parser import parse
from slidemark.theme import get_theme

GFM = """# Plan

{.gantt header=2}
| 施策 | 2027 | < | 2028 | < |
|-|-|-|-|-|
| ^ | 上期 | 下期 | 上期 | 下期 |
| IoT標準化 | 新機種へ搭載 | < | 既設機へ | ― |
| AI保全 | - | PoC | 本番 | < |
| 海外 | | | 北米 | 欧州 |
"""

FENCE = """# Plan

{.gantt}
```table
施策,上期,下期
IoT,新機種,<
AI,―,PoC
```
"""


def _layout(md: str, theme: str = "default"):
    deck = parse(md)
    return layout_slide(deck.slides[0], deck, get_theme(theme), 0), deck


def _parts(md: str, theme: str = "default"):
    placed, deck = _layout(md, theme)
    table = next(p for p in placed if isinstance(p.element, Table))
    bars = [p for p in placed if isinstance(p.element, Shape) and "gantt" in p.element.classes]
    return table, bars, deck


def test_class_accepted_on_both_syntaxes():
    for md in (GFM, FENCE):
        deck = parse(md)
        t = next(e for e in deck.slides[0].elements if isinstance(e, Table))
        assert "gantt" in t.classes
        assert not [d for d in deck.diagnostics if d.level == "error"]


def test_bar_geometry_spans_merged_cells():
    table, bars, _ = _parts(GFM)
    cw, rh = table_boxes(table)
    # bars: 新機種へ搭載 (2 cols), 既設機へ, PoC, 本番 (2 cols), 北米, 欧州
    assert len(bars) == 6
    first = bars[0]
    left = table.x + cw[0]
    assert first.x > left
    assert first.x + first.w < left + cw[1] + cw[2]
    assert first.w > 1.7 * min(cw[1], cw[2]) - 2 * 38100  # spans both period columns
    assert all(isinstance(b.element, Shape) and b.element.shape == "rounded-rect" for b in bars)
    # centered in the row, 60-70 percent of its height
    row0 = table.y + sum(rh[:2])
    assert abs((first.y + first.h / 2) - (row0 + rh[2] / 2)) < 2000
    assert 0.55 <= first.h / rh[2] <= 0.8
    # bar text is the cell text, bars sit above the table in z-order
    assert "".join(r.text for r in first.element.paragraphs[0].runs) == "新機種へ搭載"


def test_no_bar_for_empty_and_dash_cells_and_cells_are_emptied():
    table, bars, _ = _parts(GFM)
    texts = {"".join(p.plain for p in b.element.paragraphs) for b in bars}
    assert texts == {"新機種へ搭載", "既設機へ", "PoC", "本番", "北米", "欧州"}
    rows = [["".join(p.plain for p in c.paragraphs) for c in row] for row in table.element.rows]
    assert rows[2][1:] == ["", "", "", "―"]  # bar cells emptied, the mark stays
    assert rows[2][0] == "IoT標準化"  # first column stays text
    assert rows[1][1] == "上期"  # header rows untouched
    assert rows[3][1] == "-"  # a "no value" mark stays


def test_table_before_bars_in_z_order():
    placed, _ = _layout(GFM)
    idx = [i for i, p in enumerate(placed) if isinstance(p.element, Table)]
    bars = [i for i, p in enumerate(placed) if isinstance(p.element, Shape) and "gantt" in p.element.classes]
    assert idx and bars and max(idx) < min(bars)


def test_bar_style_tokens_and_css():
    md = GFM.replace("# Plan", "# Plan\n```css\n.gantt { background: #FF0000; border-radius: 9px }\n```", 1)
    _, bars, _ = _parts(md)
    assert bars[0].style.fill == "#FF0000"
    placed, _ = _layout(GFM)
    t = next(p for p in placed if isinstance(p.element, Table))
    assert not t.style.fill  # the bar look never leaks onto the table body


def test_every_preset_has_the_gantt_class():
    for name in ("default", "jp-business", "midnight"):
        assert "gantt" in get_theme(name).classes


def test_long_bar_text_shrinks_or_warns():
    md = GFM.replace("PoC", "非常に長い説明文がここに入りますがバーは狭いです" * 8)
    table, bars, deck = _parts(md)
    long = next(b for b in bars if "非常に" in b.element.paragraphs[0].plain)
    ph, pv = css.inset_hv(long.style)
    need = measure.paragraphs_height(long.element.paragraphs, long.w - ph, long.style, long.font_scale) + pv
    assert need <= long.h * 1.05 or any(d.rule == "gantt-text" for d in deck.diagnostics)


def test_render_bars_are_native_shapes_and_cells_empty(tmp_path):
    out = tmp_path / "g.pptx"
    build(GFM, out)
    prs = Presentation(str(out))
    slide = prs.slides[0]
    shapes = list(slide.shapes)
    tbl = next(s for s in shapes if s.has_table)
    bars = [
        s
        for s in shapes
        if s.shape_type is not None
        and not s.has_table
        and s.has_text_frame
        and s.text_frame.text in {"新機種へ搭載", "PoC", "北米", "欧州", "既設機へ", "本番"}
    ]
    assert len(bars) == 6
    assert shapes.index(tbl) < shapes.index(bars[0])
    cell_texts = [c.text for row in tbl.table.rows for c in row.cells]
    assert "新機種へ搭載" not in cell_texts and "PoC" not in cell_texts
    assert "IoT標準化" in cell_texts and "上期" in cell_texts
    for b in bars:
        assert b.left >= tbl.left and b.left + b.width <= tbl.left + tbl.width


def test_import_round_trip(tmp_path):
    out = tmp_path / "g.pptx"
    build(GFM, out)
    text, _diags = import_pptx(out)
    assert "{.gantt header=2}" in text
    assert "| IoT標準化 | 新機種へ搭載 | < | 既設機へ | ― |" in text
    assert "| 海外 | | | 北米 | 欧州 |" in text


def test_fuzz_safe():
    for md in (
        "# T\n{.gantt}\n| a |\n|-|\n| b |\n",
        "# T\n{.gantt}\n```table\n```\n",
        "# T\n{.gantt header=5}\n| a | b |\n|-|-|\n| < | ^ |\n",
        "# T\n{.gantt}\n| a | b | c |\n|-|-|-|\n| x | ^ | < |\n| ^ | y | |\n",
    ):
        deck = parse(md)
        layout_slide(deck.slides[0], deck, get_theme("default"), 0)


def test_style_token_restyles_bars():
    from slidemark.template import deck_theme

    deck = parse("style: gantt.fill=#00AA00 gantt.radius=9\n\n" + GFM)
    theme, _ = deck_theme(deck)
    placed = layout_slide(deck.slides[0], deck, theme, 0)
    bars = [p for p in placed if isinstance(p.element, Shape) and "gantt" in p.element.classes]
    assert bars and all(b.style.fill == "#00AA00" and b.style.radius == 9 for b in bars)
