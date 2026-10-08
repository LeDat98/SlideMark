"""DL3d lane B: geometry-based recognition of forms drawn by hand in a foreign (python-pptx) deck.

Each test draws one form with plain rectangles and text boxes (python-pptx default names), imports the deck
and asserts that the imported SlideMark text holds the form; a slide that does not match stays plain.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.parser import parse

NAVY, BLUE, TEAL, ORANGE, GREY, SURFACE = "122B4A", "1F5FA8", "1B8A8F", "E08A1E", "6B7582", "EEF2F7"


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


class Deck:
    """A 13.33 x 7.5 in deck whose slides are made of unnamed rectangles and text boxes."""

    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.s = None

    def slide(self, title: str) -> Deck:
        self.s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.rect(0, 0, 13.333, 1.15, NAVY)
        self.text(0.6, 0.12, 12.1, 0.9, [(title, 28, True, "FFFFFF")])
        return self

    def rect(self, x, y, w, h, fill, text=None, shape=MSO_SHAPE.RECTANGLE):
        sh = self.s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = _rgb(fill)
        sh.line.fill.background()
        sh.shadow.inherit = False
        if text:
            self._paras(sh.text_frame, [text] if isinstance(text, tuple) else text)
        return sh

    def text(self, x, y, w, h, paras):
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tb.text_frame.word_wrap = True
        self._paras(tb.text_frame, paras)
        return tb

    @staticmethod
    def _paras(tf, paras):
        """``paras``: tuples ``(text, size, bold, color)`` or lists of such runs (one paragraph)."""
        first = True
        for para in paras:
            runs = [para] if isinstance(para, tuple) else para
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            for text, size, bold, color in runs:
                r = p.add_run()
                r.text = text
                r.font.size, r.font.bold = Pt(size), bold
                r.font.color.rgb = _rgb(color)

    def save(self, path: Path) -> Path:
        self.prs.save(str(path))
        return path


def _import(deck: Deck, tmp_path: Path) -> str:
    text, diags = import_pptx(deck.save(tmp_path / "foreign.pptx"))
    assert not [d for d in diags if d.level == "error"], diags
    return text


def _body(text: str) -> str:
    """The slide text after the header lines."""
    return text[text.index("\n# ") :]


# --------------------------------------------------------------------------- number tiles


def _tiles(d: Deck, figs, colors=(BLUE, TEAL, ORANGE), size=48, cap=20):
    for i, (fig, note) in enumerate(figs):
        x = 0.6 + i * 4.2
        d.rect(x, 1.8, 3.9, 2.6, SURFACE)
        d.rect(x, 1.8, 0.12, 2.6, colors[i])
        d.text(x + 0.35, 1.9, 3.4, 1.0, [(fig, size, True, colors[i])])
        d.text(x + 0.35, 3.0, 3.4, 1.2, [(note, cap, True, NAVY)])


def test_number_tiles_become_a_kpi_slide(tmp_path):
    d = Deck().slide("市場環境")
    _tiles(d, [("2.1兆円", "市場は2.1兆円へ"), ("38%", "導入率は38%"), ("+22%", "前年比+22%")])
    text = _import(d, tmp_path)
    body = _body(text)
    assert "# 市場環境" in body and "@kpi" in body
    assert body.count("\n##") == 3 and "\n2.1兆円\n市場は2.1兆円へ" in body
    assert "color=#1F5FA8" in body and "s2.border-left=" in body  # figure colour and the stripe of each card
    assert "sizes: kpi=48" in body and "kpi.note.size=20" in body
    deck = parse(text)
    cards = [e for e in deck.slides[0].elements if hasattr(e, "classes") and "kpi" in e.classes]
    assert len(cards) == 3
    build(text, tmp_path / "back.pptx")  # and it builds


def test_tiles_stay_plain_when_one_card_is_no_figure(tmp_path):
    d = Deck().slide("市場環境")
    _tiles(d, [("2.1兆円", "市場は2.1兆円へ"), ("導入が進む", "導入率は高い"), ("+22%", "前年比+22%")])
    assert "@kpi" not in _import(d, tmp_path)


def test_a_tile_is_not_a_title(tmp_path):
    d = Deck().slide("市場環境")
    _tiles(d, [("2.1兆円", "市場は2.1兆円へ"), ("38%", "導入率は38%")])
    text = _import(d, tmp_path)
    assert "# 市場環境" in text and "# 2.1兆円" not in text


# --------------------------------------------------------------------------- number headings


def test_numbered_cards_keep_a_big_coloured_number_heading(tmp_path):
    d = Deck().slide("社長メッセージ")
    for i, col in enumerate((BLUE, TEAL, ORANGE)):
        x = 0.6 + i * 4.2
        d.rect(x, 1.8, 3.9, 4.0, SURFACE)
        d.rect(x, 1.8, 3.9, 0.12, col)
        d.text(x + 0.3, 2.1, 3.3, 0.9, [(f"0{i + 1}", 44, True, col)])
        d.text(x + 0.3, 3.2, 3.3, 2.9, [(f"方針{i + 1}を実行する", 30, True, NAVY)])
    text = _import(d, tmp_path)
    assert "# 社長メッセージ" in text
    assert "sizes: heading=44! body=30!" in text
    assert "## [01]{color=#1F5FA8} {.s1}" in text and "## [03]{color=#E08A1E} {.s3}" in text
    assert 's1.border-top="9pt solid #1F5FA8"' in text
    build(text, tmp_path / "back.pptx")


# --------------------------------------------------------------------------- badges


def _bars(d: Deck, labels, colors, w=12.13, badge_w=1.1):
    for i, (lab, col) in enumerate(zip(labels, colors, strict=True)):
        y = 1.6 + i * 1.3
        d.rect(0.6, y, w, 1.12, SURFACE)
        d.rect(0.6, y, badge_w, 1.12, col, (str(i + 1), 32, True, "FFFFFF"))
        d.text(0.6 + badge_w + 0.3, y, w - badge_w - 0.5, 1.12, [(lab, 30, True, NAVY)])


def test_badge_bars_alone_become_rows(tmp_path):
    d = Deck().slide("取締役会への依頼事項")
    _bars(d, ["計画の承認", "投資枠の承認", "価格改定の承認"], (BLUE, TEAL, ORANGE))
    text = _import(d, tmp_path)
    body = _body(text)
    assert body.startswith("\n# 取締役会への依頼事項")  # the badge "1" is never the title
    assert "@rows" in body and "plain" not in body
    assert "rows-num.fill=#1F5FA8,#1B8A8F,#E08A1E" in body
    assert "1. **計画の承認**" in body and "**価格改定の承認**" in body
    build(text, tmp_path / "back.pptx")


def test_badge_bars_beside_a_panel_become_numbered_boxes(tmp_path):
    d = Deck().slide("価格改定の方針")
    d.rect(0.6, 1.6, 5.3, 5.1, NAVY)
    d.text(0.9, 1.8, 4.7, 0.5, [("現行プラン", 16, False, "C9D6E6")])
    d.text(0.9, 2.6, 4.7, 1.0, [("9,800円", 40, True, "FFFFFF")])
    d.text(0.9, 3.6, 4.7, 0.7, [("▼", 28, False, ORANGE)])
    d.text(0.9, 4.3, 4.7, 1.1, [("10,800円", 48, True, ORANGE)])
    for i, (lab, col) in enumerate([("10月から値上げ", BLUE), ("既存顧客は据え置き", TEAL)]):
        y = 1.6 + i * 1.35
        d.rect(6.2, y, 6.5, 1.15, SURFACE)
        d.rect(6.2, y, 0.8, 1.15, col, (str(i + 1), 26, True, "FFFFFF"))
        d.text(7.15, y, 5.4, 1.15, [(lab, 21, True, NAVY)])
    text = _import(d, tmp_path)
    body = _body(text)
    assert body.startswith("\n# 価格改定の方針") and " num" in body
    assert "box.num.fill=#1F5FA8" in body
    assert "## 10月から値上げ" in body and "## 既存顧客は据え置き" in body
    # the dark panel: label heading, lines with their own sizes (mixed-size span syntax)
    assert "[現行プラン]{size=16" in body and "[9,800円]{size=40" in body and "[10,800円]{size=48" in body
    assert "fill=#122B4A" in body and ".kpi" in body
    build(text, tmp_path / "back.pptx")


def test_circles_above_card_bodies_become_num(tmp_path):
    d = Deck().slide("進め方")
    for i, col in enumerate((BLUE, TEAL, ORANGE)):
        x = 0.6 + i * 4.2
        d.rect(x, 1.8, 3.9, 3.0, SURFACE)
        d.rect(x + 0.25, 2.0, 0.7, 0.7, col, (str(i + 1), 24, True, "FFFFFF"), shape=MSO_SHAPE.OVAL)
        d.text(x + 0.25, 2.9, 3.4, 1.6, [(f"手順{i + 1}を進める", 22, True, NAVY)])
    text = _import(d, tmp_path)
    assert text.count("\n##") == 3 and "num" in _body(text) and "rows" not in _body(text)
    assert "# 進め方" in text and "# 1" not in text


# --------------------------------------------------------------------------- quote card


def test_quote_card(tmp_path):
    d = Deck().slide("お客様の声")
    d.rect(0.6, 1.7, 12.13, 4.9, SURFACE)
    d.rect(0.6, 1.7, 0.15, 4.9, ORANGE)
    d.text(1.0, 1.7, 1.5, 1.6, [("“", 96, True, ORANGE)])
    d.text(1.6, 2.5, 10.1, 2.5, [("締め作業が3日から半日になった。", 38, True, NAVY)])
    d.text(1.6, 5.3, 10.1, 0.8, [("— 株式会社三浦製作所 経理部長", 24, False, GREY)])
    text = _import(d, tmp_path)
    body = _body(text)
    assert body.startswith("\n# お客様の声")  # the quote glyph is never the title
    assert "@quote" in body and "fill=#EEF2F7" in body and "size=38" in body
    assert "quote.mark.color=#E08A1E" in body
    assert "> — 株式会社三浦製作所 経理部長" in body
    build(text, tmp_path / "back.pptx")


# --------------------------------------------------------------------------- big-number rows


def _bignum(d: Deck, rows, height=0.92):
    for i, (label, num, note) in enumerate(rows):
        y = 1.5 + i * (height + 0.12)
        d.rect(0.6, y, 12.13, height, SURFACE)
        d.rect(0.6, y, 0.12, height, BLUE)
        runs = [(label, 22, True, NAVY), (num, 34, True, BLUE)]
        if note:
            runs.append((note, 18, False, GREY))
        d.text(1.0, y, 11.3, height, [runs])


def test_big_number_rows_alone_become_plain_rows_with_a_stripe(tmp_path):
    d = Deck().slide("数値目標")
    _bignum(
        d,
        [
            ("売上高 ", "420億円", "（前期 210億円）"),
            ("利益率 ", "20%", "（同 11%）"),
            ("ARR ", "380億円", ""),
        ],
    )
    text = _import(d, tmp_path)
    body = _body(text)
    assert "@rows plain" in body and "rows.stripe=#1F5FA8" in body and "rows.size=22" in body
    assert "[420億円]{size=34 color=#1F5FA8 bold}" in body
    assert "[（前期 210億円）]{size=18 color=#6B7582}" in body
    build(text, tmp_path / "back.pptx")


def test_big_number_rows_beside_a_chart_become_boxes_with_a_stripe_class(tmp_path):
    d = Deck().slide("採用計画")
    _bignum(d, [("2027年度 ", "120名", ""), ("2028年度 ", "150名", "")])
    d.text(7.2, 1.5, 5.5, 4.0, [("右の図: 採用数の推移", 20, False, NAVY)])
    text = _import(d, tmp_path)
    body = _body(text)
    assert "@rows" not in body
    assert 's1.border-left="9pt solid #1F5FA8"' in body and "sizes: heading=22!" in body
    assert "{.s1}" in body and "[120名]{size=34" in body


# --------------------------------------------------------------------------- conservative


def test_plain_cards_are_left_alone(tmp_path):
    d = Deck().slide("現状")
    for i in range(2):
        x = 0.6 + i * 6.2
        d.rect(x, 1.8, 5.9, 3.0, SURFACE)
        d.text(x + 0.3, 2.0, 5.3, 0.6, [("見出し", 22, True, NAVY)])
        d.text(x + 0.3, 2.8, 5.3, 1.8, [("本文の説明が続きます。", 16, False, NAVY)])
    text = _import(d, tmp_path)
    assert "rows" not in text and "@kpi" not in text and "num" not in _body(text) and "{size=" not in text


def test_a_deck_the_library_drew_is_read_by_its_names(tmp_path):
    """Shapes named by SlideMark itself (`Card 1`, `Text 2`) are never read by geometry: the slide of a
    built deck imports as before (a KPI tile with a label first stays a plain `.kpi` card)."""
    src = "# 売上\n@kpi\n## 売上高\n1,280億円\n前年比 +8%\n## 営業利益\n96億円\n前年比 +12%\n"
    out = tmp_path / "k.pptx"
    build(src, out)
    text, _ = import_pptx(out)
    assert "## 売上高 {.kpi}" in text and "{size=" not in text and "sizes:" not in text


@pytest.mark.parametrize("n", [1, 12])
def test_recognition_never_raises_on_odd_groups(tmp_path, n):
    d = Deck().slide("数字")
    for i in range(n):
        d.rect(0.3 + (i % 6) * 2.1, 1.6 + (i // 6) * 2.5, 1.9, 2.3, SURFACE)
        d.text(0.4 + (i % 6) * 2.1, 1.7 + (i // 6) * 2.5, 1.7, 0.8, [(f"{i}%", 40, True, BLUE)])
        d.text(0.4 + (i % 6) * 2.1, 2.6 + (i // 6) * 2.5, 1.7, 1.0, [("説明", 12, False, NAVY)])
    text, diags = import_pptx(d.save(tmp_path / "odd.pptx"))
    assert text and not [x for x in diags if x.level == "error"]
    assert Emu(0) == 0
