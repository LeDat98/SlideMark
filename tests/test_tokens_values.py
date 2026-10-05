"""Design-token values: typed validation, CSS-like shorthands, never silent."""

from __future__ import annotations

import pytest
from pptx import Presentation

from slidemark.build import build
from slidemark.parser import parse
from slidemark.theme import DEFAULT, apply_tokens

BOXES = "# T\n@2\n## Alpha\nfirst body\n## Beta\nsecond body\n"


def ap(**tokens):
    return apply_tokens(DEFAULT, {k.replace("__", "."): v for k, v in tokens.items()})


def card(tokens: dict[str, str]):
    th, diags = apply_tokens(DEFAULT, tokens)
    assert not diags, diags
    return th.classes["card"]


# --------------------------------------------------------------------------- colors


@pytest.mark.parametrize(
    ("raw", "want"),
    [
        ("#b08d57", "#B08D57"),
        ("#abc", "#AABBCC"),
        ("#abcd", "#AABBCCDD"),
        ("#B08D5780", "#B08D5780"),
        ("rgb(176, 141, 87)", "#B08D57"),
        ("rgba(0, 0, 0, 0.5)", "#00000080"),
        ("crimson", "#DC143C"),
        ("transparent", "#00000000"),
        ("none", "#00000000"),
        ("muted", "muted"),
    ],
)
def test_color_forms(raw, want):
    th, diags = apply_tokens(DEFAULT, {"classes.card.color": raw, "colors.accent": raw})
    assert not diags and th.classes["card"].color == want and th.colors["accent"] == want


def test_declared_color_name_is_valid():
    th, diags = apply_tokens(DEFAULT, {"classes.card.fill": "brand", "colors.brand": "#FF5A1F"})
    assert not diags and th.classes["card"].fill == "brand"


@pytest.mark.parametrize(
    "path",
    [
        "colors.primary",
        "classes.card.color",
        "classes.card.fill",
        "heading_color",
        "table_header_fill",
        "title_band",
        "render.ink_dark",
        "render.highlight",
        "palette",
    ],
)
def test_bad_color_is_one_warning_with_example(path):
    th, diags = apply_tokens(DEFAULT, {path: "0.75pt #B08D57"})
    assert len(diags) == 1 and diags[0].rule == "bad-token" and diags[0].level == "warning"
    assert "\n" not in diags[0].hint and "#" in diags[0].hint
    assert th.model_dump() == DEFAULT.model_dump()


def test_theme_color_fields_and_optional_none():
    th, diags = apply_tokens(
        DEFAULT, {"table_header_fill": "rgb(1,2,3)", "title_band": "primary", "heading_band": "none"}
    )
    assert not diags and th.table_header_fill == "#010203" and th.title_band == "primary"
    assert th.heading_band is None


def test_palette_items_validated():
    th, diags = apply_tokens(DEFAULT, {"palette": "#fff, crimson, primary"})
    assert not diags and th.palette == ["#FFFFFF", "#DC143C", "primary"]
    _, diags = apply_tokens(DEFAULT, {"palette": "#fff, banana"})
    assert [d.rule for d in diags] == ["bad-token"]


def test_fill_gradient_and_url():
    c = card({"classes.card.fill": "linear-gradient(135deg, #7c5cff, #00d1b2)"})
    assert c.fill == "linear-gradient(135deg, #7C5CFF, #00D1B2)"
    assert card({"classes.card.fill": "radial-gradient(#fff, #000)"}).fill.startswith("radial-gradient(")
    assert card({"classes.card.fill": "url(bg.png)"}).fill == "url(bg.png)"
    _, diags = apply_tokens(DEFAULT, {"classes.card.fill": "linear-gradient(#fff)"})
    assert [d.rule for d in diags] == ["bad-token"]
    _, diags = apply_tokens(DEFAULT, {"classes.card.fill": "conic-gradient(#fff, #000)"})
    assert len(diags) == 1


def test_fill_none_is_transparent():
    assert card({"classes.card.fill": "none"}).fill == "#00000000"


# --------------------------------------------------------------------------- lengths


def test_length_forms():
    c = card({"classes.card.font_size": "14", "classes.card.radius": "0.5in", "classes.card.padding": "8px"})
    assert c.font_size == 14 and c.radius == 36 and c.padding == "8px"
    th, diags = apply_tokens(
        DEFAULT, {"margin_x": "1cm", "gap": "12", "layout.top_gap": "0.3in", "sizes.body": "12pt"}
    )
    assert not diags and th.margin_x == "1cm" and th.gap == 12 and th.layout.top_gap == "0.3in"
    assert th.sizes["body"] == 12


def test_font_size_dash_alias_and_unit():
    d = parse("style: card.font-size=14 hero.font-size=0.25in\n# T\n")
    assert d.tokens["classes.card.font_size"] == "14" and not d.diagnostics
    th, diags = apply_tokens(DEFAULT, d.tokens)
    assert not diags and th.classes["card"].font_size == 14 and th.classes["hero"].font_size == 18


@pytest.mark.parametrize("path", ["classes.card.font_size", "classes.card.padding", "gap", "sizes.body"])
def test_bad_length_is_one_warning(path):
    _, diags = apply_tokens(DEFAULT, {path: "wide"})
    assert len(diags) == 1 and diags[0].rule == "bad-token" and "12pt" in diags[0].hint


# --------------------------------------------------------------------------- shorthands


def test_border_shorthand_any_order():
    for raw in ("0.75pt solid #B08D57", "#B08D57 0.75pt solid", "solid #b08d57 0.75pt"):
        c = card({"classes.card.line": raw})
        assert (c.line, c.line_width, c.line_dash) == ("#B08D57", 0.75, "solid")


def test_border_dash_and_px():
    c = card({"classes.card.line": "2px dashed crimson"})
    assert (c.line, c.line_width, c.line_dash) == ("#DC143C", 1.5, "dash")


def test_border_width_alone_and_color_alone():
    base = DEFAULT.classes["card"]
    c = card({"classes.card.line": "2pt"})
    assert c.line == base.line and c.line_width == 2
    c = card({"classes.card.line": "#FF0000"})
    assert c.line == "#FF0000" and c.line_width == base.line_width


def test_border_none():
    c = card({"classes.card.line": "none"})
    assert c.line_width == 0


def test_border_via_aliases_and_new_class():
    d = parse('style: border="1pt dotted #112233" hero.border="3pt #FF0000"\n# T\n')
    assert not d.diagnostics
    th, diags = apply_tokens(DEFAULT, d.tokens)
    assert not diags
    assert (th.classes["card"].line, th.classes["card"].line_dash) == ("#112233", "dot")
    assert (th.classes["hero"].line, th.classes["hero"].line_width) == ("#FF0000", 3)


def test_bad_border_is_one_warning():
    th, diags = apply_tokens(DEFAULT, {"classes.card.line": "thick-ish banana"})
    assert len(diags) == 1 and diags[0].rule == "bad-token" and "0.75pt" in diags[0].hint
    assert th.classes["card"] == DEFAULT.classes["card"]


def test_border_sides():
    d = parse(
        'style: card.border-top="2pt #FF0000" card.border-left=none card.border-bottom="dashed navy"\n# T\n'
    )
    assert not d.diagnostics
    c = card(d.tokens)
    assert c.border_top == "2pt solid #FF0000" and c.border_left == "none"
    assert c.border_bottom == "2.25pt dash #000080"
    _, diags = apply_tokens(DEFAULT, {"classes.card.border_top": "2 zigzag"})
    assert len(diags) == 1


def test_shadow_forms():
    assert card({"classes.card.shadow": "none"}).shadow is False
    assert card({"classes.card.shadow": "on"}).shadow is True
    assert card({"classes.card.shadow": "0 8 24 #00000055"}).shadow == "0 8 24 #00000055"
    assert card({"classes.card.shadow": "0 2pt 6pt rgba(0,0,0,.25)"}).shadow == "0 2 6 #00000040"
    _, diags = apply_tokens(DEFAULT, {"classes.card.shadow": "soft glow"})
    assert [d.rule for d in diags] == ["bad-token"] and "0 2 6" in diags[0].hint


def test_global_shadow_alias_still_maps_to_card():
    d = parse("style: shadow=none\n# T\n")
    assert d.tokens == {"classes.card.shadow": "none"}
    assert card(d.tokens).shadow is False


# --------------------------------------------------------------------------- parse time (line numbers)


def test_check_reports_bad_value_with_line():
    d = parse("theme: none\nstyle: card.fill=banana\n# T\n")
    (x,) = [x for x in d.diagnostics if x.rule == "bad-token"]
    assert x.line == 2 and "banana" in x.message and "\n" not in x.hint
    assert "classes.card.fill" not in d.tokens


def test_forward_declared_color_name_is_not_flagged():
    d = parse("theme: none\nstyle: card.fill=brand\ncolors: brand=#FF5A1F\n# T\n")
    assert not d.diagnostics
    th, diags = apply_tokens(DEFAULT, d.tokens)
    assert not diags and th.classes["card"].fill == "brand"


def test_undeclared_bare_word_is_flagged_once_with_its_line():
    d = parse("theme: none\nstyle: card.fill=brand\n# T\n")
    (x,) = [x for x in d.diagnostics if x.rule == "bad-token"]
    assert x.line == 2 and "classes.card.fill" not in d.tokens
    _, diags = apply_tokens(DEFAULT, d.tokens)
    assert not diags


def test_each_bad_value_exactly_one_diagnostic_end_to_end(tmp_path):
    md = "style: card.fill=banana gap=wide card.border=nope\n" + BOXES
    deck = build(md, tmp_path / "o.pptx")
    bad = [d for d in deck.diagnostics if d.rule == "bad-token"]
    assert len(bad) == 3


# --------------------------------------------------------------------------- header CSS with theme: none


def test_header_css_fence_with_theme_none(tmp_path):
    md = "theme: none\n\n```css\n.box {border: 2px solid #FF8800}\n```\n\n" + BOXES
    deck = build(md, tmp_path / "o.pptx")
    assert deck.theme == "none"
    assert not [d for d in deck.diagnostics if d.level in ("warning", "error")]
    prs = Presentation(str(tmp_path / "o.pptx"))
    sp = next(s for s in prs.slides[0].shapes if s.name.startswith("Card"))
    assert sp._element.spPr.xpath("./a:ln/a:solidFill/a:srgbClr")[0].get("val") == "FF8800"


# --------------------------------------------------------------------------- the eval example


def test_eval_example_border_is_drawn(tmp_path):
    md = 'theme: none\nstyle: card.border="0.75pt #B08D57"\n\n' + BOXES
    deck = build(md, tmp_path / "o.pptx")
    assert not [d for d in deck.diagnostics if d.rule == "bad-token"]
    prs = Presentation(str(tmp_path / "o.pptx"))
    cards = [s for s in prs.slides[0].shapes if s.name.startswith("Card")]
    assert cards
    for sp in cards:
        ln = sp._element.spPr.xpath("./a:ln")[0]
        assert ln.get("w") == str(round(0.75 * 12700))
        assert ln.xpath("./a:solidFill/a:srgbClr")[0].get("val") == "B08D57"
