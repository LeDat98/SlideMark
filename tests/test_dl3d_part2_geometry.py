"""DL3d part 2 (lane I): geometry leftovers of the foreign-deck round trip, as tokens.

``box.h`` / ``box.anchor`` (tall top-aligned header-band cards), the deck margin and body top
(``margin`` / ``layout.top_gap`` written by the importer), ``steps-arrow.point`` (the preset ``adj``) and the
plain table row heights (``rowh=`` after the stated widths). Parser, render (reopen the .pptx), importer
(synthetic foreign decks)."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.parser import parse
from slidemark.template import deck_theme

IN = 914400
NAVY, BLUE, TEAL, AMBER, LIGHT = "122B4A", "1F5FA8", "1B8A8F", "E08A1E", "EEF2F7"


def _shapes(tmp_path: Path, md: str, slide: int = 0) -> dict:
    out = tmp_path / "out.pptx"
    build(md, out, base_dir=tmp_path)
    return {sh.name: sh for sh in Presentation(str(out)).slides[slide].shapes}


def _diags(md: str):
    from slidemark.cli import _add_layout_diagnostics

    deck = parse(md)
    _add_layout_diagnostics(deck, "d.md")
    return deck.diagnostics


BOXES = """{style}

# Plan
## North
- one
- two
## South
- three
## East
- four
"""


def _boxes(tmp_path: Path, style: str) -> dict:
    return _shapes(tmp_path, BOXES.format(style=f"style: {style}" if style else ""))


# --------------------------------------------------------------------------- tokens (parser)


def test_tokens_parse_without_diagnostics():
    md = "style: box.h=4.2in box.anchor=top steps-arrow.point=0.5 margin=0.6in layout.top_gap=0.45in\n\n# t\n"
    deck = parse(md)
    assert not [d for d in deck.diagnostics if d.rule in ("bad-token", "unknown-token")]
    theme, _ = deck_theme(deck, ".")
    assert theme.box_h == "4.2in" and theme.box_anchor == "top" and theme.steps_arrow_point == 0.5
    assert theme.margin_x == "0.6in" and theme.layout.top_gap == "0.45in"
    # the point is the one depth of every arrow: the layout's search over flatter points is off
    assert theme.layout.chevron_adj == 0.5 and theme.layout.chevron_adj_min == 0.5


def test_box_anchor_aliases_and_percent_point():
    theme, _ = deck_theme(parse("style: box.anchor=middle steps-arrow.point=40%\n\n# t\n"), ".")
    assert theme.box_anchor == "center" and theme.steps_arrow_point == 0.4


@pytest.mark.parametrize(
    "bad",
    [
        "box.h=tall",
        "box.h=-2in",
        "box.h=0in",
        "box.anchor=side",
        "steps-arrow.point=2",
        "steps-arrow.point=0",
        "steps-arrow.point=wide",
    ],
)
def test_bad_values_warn_with_a_hint(bad):
    bads = [d for d in parse(f"style: {bad}\n\n# t\n").diagnostics if d.rule == "bad-token"]
    assert len(bads) == 1 and bads[0].hint and bad.split("=")[0] in bads[0].message


def test_box_tokens_without_a_box_say_so():
    got = [
        d for d in _diags("style: box.h=4in box.anchor=top\n\n# t\nplain text\n") if d.rule == "attr-ignored"
    ]
    assert {d.message for d in got} >= {
        "style: box.h= has no effect here",
        "style: box.anchor= has no effect here",
    }
    assert all(d.hint for d in got)


# --------------------------------------------------------------------------- render: box height and anchor


def _cards(sh: dict) -> list:
    return [sh[f"Card {k}"] for k in (1, 2, 3)]


def test_box_h_is_exact_on_every_card(tmp_path):
    sh = _boxes(tmp_path, "box.h=4.2in")
    assert all(c.height == round(4.2 * IN) for c in _cards(sh))
    assert len({c.top for c in _cards(sh)}) == 1  # one row


def test_without_box_h_the_cards_keep_their_own_height(tmp_path):
    sh = _boxes(tmp_path, "")
    assert all(c.height != round(4.2 * IN) for c in _cards(sh))


def test_box_anchor_top_center_bottom_place_the_row_in_the_body(tmp_path):
    top = _boxes(tmp_path, "box.h=3in box.anchor=top")["Card 1"]
    mid = _boxes(tmp_path, "box.h=3in box.anchor=center")["Card 1"]
    bot = _boxes(tmp_path, "box.h=3in box.anchor=bottom")["Card 1"]
    assert top.top < mid.top < bot.top
    assert top.top < round(2.0 * IN)  # right under the title
    assert bot.top + bot.height <= round(7.5 * IN) and bot.top + bot.height > round(6.3 * IN)
    mid_gap_top = mid.top - top.top
    mid_gap_bottom = bot.top - mid.top
    assert abs(mid_gap_top - mid_gap_bottom) <= 0.03 * IN  # centred: equal room above and below


def test_card_children_follow_the_card(tmp_path):
    sh = _boxes(tmp_path, "box.h=4.4in box.anchor=top")
    card, head, text = sh["Card 1"], sh["Heading 1"], sh["Text 1"]
    base = _boxes(tmp_path, "")
    assert head.top - card.top == base["Heading 1"].top - base["Card 1"].top  # same place under the top edge
    assert text.top > head.top + head.height - 2
    assert text.top + text.height <= card.top + card.height + 2  # the text box ends inside the card
    assert text.top + text.height > card.top + card.height - 0.4 * IN  # and keeps its distance to the bottom


def test_anchor_alone_keeps_the_layouts_card_height(tmp_path):
    base = _boxes(tmp_path, "")["Card 1"]
    top = _boxes(tmp_path, "box.anchor=top")["Card 1"]
    assert top.height == base.height and top.top <= base.top


def test_two_rows_keep_their_gap(tmp_path):
    md = "style: box.h=2.2in box.anchor=top\n\n# t\n@2x2\n## A\n- a\n## B\n- b\n## C\n- c\n## D\n- d\n"
    sh = _shapes(tmp_path, md)
    cards = [sh[f"Card {k}"] for k in (1, 2, 3, 4)]
    assert all(c.height == round(2.2 * IN) for c in cards)
    assert cards[0].top == cards[1].top and cards[2].top == cards[3].top
    assert cards[2].top >= cards[0].top + cards[0].height  # no overlap


def test_a_box_with_its_own_height_is_left_alone(tmp_path):
    md = "style: box.h=4in\n\n# t\n## A {h=2in}\n- a\n## B\n- b\n"
    sh = _shapes(tmp_path, md)
    assert sh["Card 1"].height == round(2 * IN) and sh["Card 2"].height == round(4 * IN)


def test_kpi_and_steps_cards_are_not_boxes(tmp_path):
    sh = _shapes(tmp_path, "style: box.h=1in\n\n# Plan\n@3 steps\n## a\n- one\n## b\n- two\n## c\n- three\n")
    assert sh["Step 1 card"].height != round(1 * IN)


def test_a_box_note_under_the_cards_follows_them(tmp_path):
    md = "style: box.h=4.2in box.anchor=top\n\n# t\n@3\n## A\n- a\n## B\n- b\n## C\n- c\n@end\nA note\n"
    out = tmp_path / "o.pptx"
    build(md, out, base_dir=tmp_path)
    shapes = list(Presentation(str(out)).slides[0].shapes)
    cards = [s for s in shapes if s.name.startswith("Card")]
    note = next(s for s in shapes if s.has_text_frame and "A note" in s.text_frame.text)
    assert note.top >= max(c.top + c.height for c in cards) - 2


# --------------------------------------------------------------------------- render: body offset and margin


def test_top_gap_and_margin_move_the_body_and_the_chrome(tmp_path):
    md = "style: margin=0.6in layout.top_gap=0.45in\nfooter: x\n\n# Plan\n## A\n- a\n## B\n- b\n"
    base = _shapes(tmp_path, "footer: x\n\n# Plan\n## A\n- a\n## B\n- b\n")
    sh = _shapes(tmp_path, md)
    assert sh["Title"].left == round(0.6 * IN) and base["Title"].left == round(0.5 * IN)
    assert sh["Footer"].left == round(0.6 * IN)
    assert sh["Card 1"].left == round(0.6 * IN)


def test_pinned_steps_sit_at_the_body_top(tmp_path):
    md = (
        "style: steps-arrow.h=0.75in steps-card.h=2.7in steps.gap=0.2in layout.top_gap=0.45in\n\n"
        "# Plan\n@3 steps\n## a\n- one\n## b\n- two\n## c\n- three\n"
    )
    sh = _shapes(tmp_path, md)
    title_bottom = sh["Title"].top + sh["Title"].height
    assert abs(sh["Step 1 arrow"].top - (title_bottom + round(0.45 * IN))) <= 2


# --------------------------------------------------------------------------- render: point depth


def _prst(shape) -> str:
    geom = shape._element.spPr.find(qn("a:prstGeom"))
    return geom.get("prst") if geom is not None else ""


def _adj(shape) -> float:
    gd = shape._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
    return int(gd[0].get("fmla").split()[1]) / 100000


@pytest.mark.parametrize("point", ["0.5", "0.2", "0.35"])
def test_steps_arrow_point_is_the_presets_adj(tmp_path, point):
    md = f"style: steps-arrow.point={point}\n\n# Plan\n@3 steps\n## a\n- one\n## b\n- two\n## c\n- three\n"
    sh = _shapes(tmp_path, md)
    assert all(abs(_adj(sh[f"Step {k} arrow"]) - float(point)) < 0.001 for k in (1, 2, 3))


def test_steps_arrow_point_reaches_the_pinned_geometry_and_the_chevron_strip(tmp_path):
    md = (
        "style: steps-arrow.point=0.5 steps-arrow.h=0.6in steps-card.h=2in\n\n"
        "# Plan\n@3 steps\n## a\n- one\n## b\n- two\n## c\n- three\n"
    )
    sh = _shapes(tmp_path, md)
    assert abs(_adj(sh["Step 2 arrow"]) - 0.5) < 0.001 and sh["Step 2 arrow"].height == round(0.6 * IN)
    strip = _shapes(tmp_path, "style: steps-arrow.point=0.5\n\n# Plan\n@3 chevron\n## a\n## b\n## c\n")
    arrows = [s for s in strip.values() if _prst(s) in ("chevron", "homePlate")]
    assert len(arrows) == 3 and all(abs(_adj(a) - 0.5) < 0.001 for a in arrows)


def test_default_point_is_unchanged(tmp_path):
    sh = _shapes(tmp_path, "# Plan\n@3 steps\n## a\n- one\n## b\n- two\n## c\n- three\n")
    assert _adj(sh["Step 1 arrow"]) < 0.31


# --------------------------------------------------------------------------- importer (synthetic decks)


def _deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def _slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(s, x, y, w, h, rgb, shape=MSO_SHAPE.RECTANGLE, adj=None):
    sp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(rgb)
    sp.line.fill.background()
    sp.shadow.inherit = False
    if adj is not None:
        sp.adjustments[0] = adj
    return sp


def _text(s, x, y, w, h, paras, size=16, bold=False, color="222B36", shape=None):
    tb = shape or s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(paras if isinstance(paras, list) else [paras]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = t
        r.font.size, r.font.bold = Pt(size), bold
        r.font.color.rgb = RGBColor.from_string(color)
    return tb


def _chrome(s, title):
    _rect(s, 0, 0, 13.333, 1.15, NAVY)
    _text(s, 0.6, 0.12, 12.133, 0.9, title, size=28, bold=True, color="FFFFFF")


def _import(prs, tmp_path: Path) -> str:
    path = tmp_path / "foreign.pptx"
    prs.save(path)
    text, diags = import_pptx(path)
    assert not [d for d in diags if d.level == "error"], diags
    return text


def _band_slide(prs, title, top=1.6, h=4.4, labels=("North", "South", "East")):
    s = _slide(prs)
    _chrome(s, title)
    for i, lab in enumerate(labels):
        x = 0.6 + i * 4.1167
        _rect(s, x, top, 3.8833, h, LIGHT)
        head = _rect(s, x, top, 3.8833, 0.9, NAVY)
        _text(s, 0, 0, 0, 0, lab, size=24, bold=True, color="FFFFFF", shape=head)
        _text(s, x + 0.15, top + 1.1, 3.6, h - 1.3, ["■ first point", "■ second point"], size=20)
    return s


def _header_margin_deck(prs):
    for k in range(3):
        _band_slide(prs, f"Slide {k + 1}")


def test_importer_states_card_height_anchor_margin_and_body_top(tmp_path):
    prs = _deck()
    _header_margin_deck(prs)
    text = _import(prs, tmp_path)
    assert "box.h=4.4in" in text and "box.anchor=top" in text
    assert "margin=0.6in" in text and "layout.top_gap=0.45in" in text
    sh = _shapes(tmp_path, text)
    assert sh["Card 1"].height == round(4.4 * IN) and sh["Card 1"].left == round(0.6 * IN)
    assert (
        abs(sh["Card 1"].top - round(1.6 * IN)) <= 0.02 * IN
    )  # the card top lands where the original had it


def test_importer_leaves_the_default_margin_and_top_alone(tmp_path):
    prs = _deck()
    for k in range(3):
        s = _slide(prs)
        _rect(s, 0, 0, 13.333, 1.15, NAVY)
        _text(s, 0.5, 0.12, 12.333, 0.9, f"Slide {k + 1}", size=28, bold=True, color="FFFFFF")
        _text(s, 0.5, 1.4, 12.3, 1.0, "text at the default margin", size=20)
        _rect(s, 0.5, 2.8, 12.3, 2.0, LIGHT)
    text = _import(prs, tmp_path)
    assert "margin=" not in text and "layout.top_gap=" not in text


def test_importer_anchor_follows_the_cards_vertical_place(tmp_path):
    prs = _deck()
    _band_slide(prs, "Slide 1")
    _band_slide(prs, "Slide 2")
    _band_slide(prs, "Slide 3", top=2.3, h=3.0)  # a row that floats: no edge, not the middle
    text = _import(prs, tmp_path)
    third = text[text.index("# Slide 3") :]
    assert "box.h=3in" in third and "box.anchor" not in third
    assert "box.anchor=top" in text[: text.index("# Slide 3")]


def test_importer_reads_the_arrows_point_depth(tmp_path):
    prs = _deck()
    for k in range(3):
        s = _slide(prs)
        _chrome(s, f"Plan {k + 1}")
        for i, lab in enumerate(["One", "Two", "Three"]):
            x = 0.6 + i * 4.1
            arrow = _rect(s, x, 1.6, 3.9, 0.75, (BLUE, TEAL, AMBER)[i], MSO_SHAPE.PENTAGON, adj=0.5)
            _text(s, 0, 0, 0, 0, [lab], size=24, bold=True, color="FFFFFF", shape=arrow)
            _rect(s, x, 2.55, 3.9, 2.5, LIGHT)
    text = _import(prs, tmp_path)
    assert "steps-arrow.point=0.5" in text
    out = tmp_path / "rebuilt.pptx"
    build(text, out, base_dir=tmp_path)
    arrow = next(s for s in Presentation(str(out)).slides[0].shapes if s.name == "Step 1 arrow")
    assert abs(_adj(arrow) - 0.5) < 0.001


def test_importer_states_nothing_for_the_builds_own_point(tmp_path):
    prs = _deck()
    for k in range(2):
        s = _slide(prs)
        _chrome(s, f"Plan {k + 1}")
        for i, lab in enumerate(["One", "Two", "Three"]):
            x = 0.6 + i * 4.1
            arrow = _rect(s, x, 1.6, 3.9, 0.75, BLUE, MSO_SHAPE.PENTAGON, adj=0.3)
            _text(s, 0, 0, 0, 0, [lab], size=24, bold=True, color="FFFFFF", shape=arrow)
            _rect(s, x, 2.55, 3.9, 2.5, LIGHT)
    assert "steps-arrow.point" not in _import(prs, tmp_path)


def _plain_table(prs, title, rows, row_h, widths):
    s = _slide(prs)
    _chrome(s, title)
    shape = s.shapes.add_table(
        len(rows), len(widths), Inches(0.6), Inches(1.6), Inches(sum(widths)), Inches(row_h * len(rows))
    )
    tbl = shape.table
    for j, w in enumerate(widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(18)
    return s


def test_importer_pins_the_row_heights_of_a_plain_table_with_a_narrow_column(tmp_path):
    prs = _deck()
    rows = [
        ["Risk", "", "Answer"],
        ["Price war", ">", "Differentiate with features"],
        ["Hiring", ">", "Referral hiring up to 40% and a new pay table for engineers"],
        ["Security", ">", "Certification by 2028"],
    ]
    _plain_table(prs, "Risks", rows, 1.1, [3.0, 0.8, 8.3])
    text = _import(prs, tmp_path)
    assert "rowh=" in text
    out = tmp_path / "rebuilt.pptx"
    build(text, out, base_dir=tmp_path)
    tbl = next(s for s in Presentation(str(out)).slides[0].shapes if s.has_table).table
    heights = [r.height for r in tbl.rows]
    assert all(abs(h - round(1.1 * IN)) <= 0.12 * 1.1 * IN for h in heights), [h / IN for h in heights]
