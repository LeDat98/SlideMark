"""Orphan control (no-break space before the last word) and the `.muted` box look."""

from __future__ import annotations

from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import build
from slidemark.build import build_deck
from slidemark.contrast import ratio
from slidemark.ir import Paragraph, Run, Style
from slidemark.layout import measure
from slidemark.parser import parse
from slidemark.theme import get_theme

NB = " "


def test_bind_vi_and_en():
    assert (
        measure.bind_last_words("Ngân sách giai đoạn 1: 120 tỷ đồng")
        == f"Ngân sách giai đoạn 1: 120 tỷ{NB}đồng"
    )
    assert (
        measure.bind_last_words("Plain English sentence ends here") == f"Plain English sentence ends{NB}here"
    )


def test_bind_skips():
    assert measure.bind_last_words("two words") == "two words"  # fewer than 3 words
    assert measure.bind_last_words("a very Internationalization-heavy") == "a very Internationalization-heavy"
    assert measure.bind_last_words("see https://example.com/x/y now") == "see https://example.com/x/y now"
    assert measure.bind_last_words("これは 日本語 の 文章です") == "これは 日本語 の 文章です"  # CJK words
    assert measure.bind_last_words("a b" + NB + "c d") == "a b" + NB + "c d"  # already bound
    assert measure.bind_last_words("x y  z") == "x y  z"  # double space


def test_bind_across_runs_and_code():
    runs = [Run(text="alpha beta gamma "), Run(text="delta", bold=True)]
    assert "".join(measure.bound_texts(runs)) == f"alpha beta gamma{NB}delta"
    code = [Run(text="alpha beta "), Run(text="gamma", code=True)]
    assert measure.bound_texts(code) == ["alpha beta ", "gamma"]
    badge = [Run(text="alpha beta gamma "), Run(text="x", highlight="primary")]
    assert measure.bound_texts(badge) == ["alpha beta gamma ", "x"]


def test_nbsp_is_one_unbreakable_space_in_measure():
    assert len(measure._units0([("aaaa bbbb", False, False)])) == 2
    assert len(measure._units0([("aaaa" + NB + "bbbb", False, False)])) == 1  # one unbreakable unit
    assert abs(measure.text_em("a" + NB + "b") - measure.text_em("a b")) < 1e-9


def test_paragraph_segments_use_bound_text():
    p = Paragraph(runs=[Run(text="one two three four")])
    assert measure.para_segments(p)[0][0] == f"one two three{NB}four"


def test_render_has_nbsp_and_import_restores_space(tmp_path):
    md = "theme: none\n\n# T\n## Box\n1. Chiến lược chuyển đổi số 2027–2029 cho toàn công ty\n"
    out = tmp_path / "a.pptx"
    build(md, out)
    texts = [sh.text_frame.text for sh in Presentation(out).slides[0].shapes if sh.has_text_frame]
    assert any(f"toàn công{NB}ty" in t or f"công{NB}ty" in t for t in texts)
    from slidemark.importer import import_pptx

    res = import_pptx(out)
    assert NB not in str(res)


def _muted_deck(tmp_path, extra=""):
    md = (
        "theme: jp-business\n"
        + extra
        + "\n# T\n@1:1\n## Strong\nbody text here\n\n## Weak {.muted}\nbody text here\n"
    )
    out = tmp_path / "m.pptx"
    build_deck(parse(md), out, tmp_path)
    return Presentation(str(out))


def _fill_of(shape):
    f = shape._element.spPr.find(qn("a:solidFill"))
    return None if f is None else f[0].get("val")


def _body_colors(prs):
    return {
        str(r.font.color.rgb)
        for s in prs.slides[0].shapes
        if s.has_text_frame and s.text_frame.text.startswith("body")
        for p in s.text_frame.paragraphs
        for r in p.runs
    }


def test_muted_box_band_and_body_ink(tmp_path):
    prs = _muted_deck(tmp_path)
    by = {
        sh.text_frame.text.strip(): sh
        for sh in prs.slides[0].shapes
        if sh.has_text_frame and sh.text_frame.text.strip()
    }
    assert _fill_of(by["Weak"]) == "5B6573"  # the band takes the muted color
    assert _fill_of(by["Strong"]) == "1E3A5F"
    assert _body_colors(prs) == {"1F2937"}  # the body keeps the normal ink
    assert str(by["Weak"].text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"


def test_muted_box_grey_body_restored_by_token(tmp_path):
    assert "5B6573" in _body_colors(_muted_deck(tmp_path, "style: muted-box.color=muted"))


def test_muted_text_stays_grey():
    th = get_theme("default")
    assert th.classes["muted"] == Style(color="muted")
    assert th.heading_band_for(["muted"]) == (None, th.heading_band_color)  # no band theme: unchanged


def test_dark_muted_band_ink_is_readable():
    th = get_theme("midnight").model_copy(update={"heading_band": "primary"})
    fill, ink = th.heading_band_for(["muted"])
    assert ratio(th.hexval(ink), th.hexval(fill)) >= 4.5
