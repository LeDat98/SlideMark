"""Auto value-axis max / major unit and bar gap width (reopened .pptx)."""

from __future__ import annotations

import math

from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark.ir import Chart, Deck, Paragraph, Run, Series, Slide, Text
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.render.axis import auto_axis
from slidemark.theme import RenderTokens, get_theme


def build_chart(tmp_path, kind, cats, series, **opts):
    title = Text(role="title", paragraphs=[Paragraph(runs=[Run(text="T")])])
    ch = Chart(kind=kind, categories=cats, series=series, options=opts)
    d = Deck(slides=[Slide(title=title, elements=[ch])])
    th = get_theme("default")
    out = tmp_path / "c.pptx"
    render(d, [layout_slide(s, d, th, i) for i, s in enumerate(d.slides)], th, out)
    prs = Presentation(str(out))
    return [s for s in prs.slides[0].shapes if s.name.startswith("Chart")][0].chart


def _stack():
    return [
        Series(name="a", values=[38, 40, 43]),
        Series(name="b", values=[31, 34, 37]),
        Series(name="c", values=[9, 11, 13]),
        Series(name="d", values=[22, 23, 24]),
    ]


def test_auto_axis_is_nice_for_stacked_totals(tmp_path):
    chart = build_chart(tmp_path, "stacked-bar", ["2024", "2025", "2026"], _stack())
    va = chart.value_axis
    assert va.maximum_scale == 125  # largest sum 117 (+ headroom), not the auto 140
    assert va.major_unit == 25
    assert va.minimum_scale == 0
    assert 4 <= va.maximum_scale / va.major_unit <= 6


def test_explicit_max_wins_and_has_no_major_unit(tmp_path):
    chart = build_chart(tmp_path, "column", ["a", "b"], [Series(name="s", values=[3, 9])], max="20")
    assert chart.value_axis.maximum_scale == 20
    assert chart._chartSpace.find(".//" + qn("c:majorUnit")) is None


def test_negative_values_get_a_two_sided_axis(tmp_path):
    chart = build_chart(tmp_path, "column", ["a", "b"], [Series(name="s", values=[-3, 9])])
    assert chart.value_axis.minimum_scale < -3 and chart.value_axis.maximum_scale >= 9


def test_nan_and_empty_series_do_not_crash(tmp_path):
    nan = float("nan")
    build_chart(tmp_path, "column", ["a", "b"], [Series(name="s", values=[nan, None])])
    build_chart(tmp_path, "stacked-column", [], [])
    build_chart(tmp_path, "column", ["a"], [Series(name="s", values=[0, 0])])


def test_line_chart_stays_auto(tmp_path):
    chart = build_chart(tmp_path, "line", ["a", "b"], [Series(name="s", values=[3, 9])])
    assert chart.value_axis.maximum_scale is None


def test_nice_axis_function_properties():
    rt = RenderTokens()
    for top in (0.7, 3.6, 9, 21, 117, 480, 1234, 98765):
        mx, unit = auto_axis("column", [[top]], rt)
        assert mx >= top * (1 + rt.chart_axis_headroom) - 1e-9
        assert rt.chart_axis_lines_min <= round(mx / unit) <= rt.chart_axis_lines_max
        mant = unit / 10 ** math.floor(math.log10(unit))
        assert any(abs(mant - m) < 1e-9 for m in (1, 2, 2.5, 5))
    assert auto_axis("column", [[]], rt) is None


def test_few_categories_get_a_wider_gap(tmp_path):
    ser = [Series(name="s", values=[1, 2, 3, 4, 5])]
    few = build_chart(tmp_path, "column", ["a", "b", "c"], [Series(name="s", values=[1, 2, 3])])
    many = build_chart(tmp_path, "column", list("abcde"), ser)
    rt = RenderTokens()
    assert few.plots[0].gap_width == rt.chart_gap_few
    assert many.plots[0].gap_width == rt.chart_gap
    opt = build_chart(tmp_path, "column", ["a", "b"], [Series(name="s", values=[1, 2])], gap_width="40")
    assert opt.plots[0].gap_width == 40


def test_negative_bar_axis_has_label_room(tmp_path):
    ch = build_chart(tmp_path, "bar", ["a", "b", "c"], [Series(name="s", values=[4.8, 1.2, -2.6])])
    va = ch.value_axis
    assert va.minimum_scale <= -2.6 - 0.05 * 7.4  # room below the lowest bar for its label
    assert va.maximum_scale >= 4.8


def test_negative_column_axis(tmp_path):
    ch = build_chart(
        tmp_path,
        "column",
        ["Q1", "Q2"],
        [Series(name="a", values=[12, -8.5]), Series(name="b", values=[-3, 4])],
    )
    assert ch.value_axis.minimum_scale < -8.5 and ch.value_axis.maximum_scale > 12


def test_neg_axis_edges():
    from slidemark.render.axis import neg_axis

    rt = RenderTokens()
    assert neg_axis("bar", [[1, 2]], rt) is None
    assert neg_axis("line", [[-1, 2]], rt) is None
    assert neg_axis("bar", [[]], rt) is None
    assert neg_axis("bar", [[None, "x", -1.0]], rt) is not None
    assert neg_axis("bar", [[-5, -3]], rt)[1] == 0.0
