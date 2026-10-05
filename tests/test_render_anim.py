"""Day 7: build animations, transitions, sections, slide jumps, hidden."""

from __future__ import annotations

from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation

from slidemark.layout import layout_slide
from slidemark.parser import parse
from slidemark.preview import have_soffice, pptx_to_pngs
from slidemark.render import render
from slidemark.theme import get_theme

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
}


def xp(el, expr):
    return etree._Element.xpath(el, expr, namespaces=NS)


def make(md: str, tmp_path: Path):
    deck = parse(md)
    th = get_theme("default")
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    out = tmp_path / "o.pptx"
    render(deck, placed, th, out)
    return deck, out, Presentation(str(out))


def clicks(slide) -> list:
    return slide._element.xpath("p:timing//p:cTn[@nodeType='mainSeq']/p:childTnLst/p:par")


# --------------------------------------------------------------------------- parser


def test_parse_build_transition_sections():
    d = parse("sections: off\n\n# A\n@build t=fade:0.5\n- a\n- b\n")
    assert d.attrs["sections"] == "off"
    assert "build" in d.slides[0].classes
    assert d.slides[0].transition == "fade:0.5"
    assert not [x for x in d.diagnostics if x.rule in ("unknown-token", "unknown-header")]


def test_parse_bad_values_warn():
    d = parse("sections: maybe\n\n# A\n@t=spin\n- a\n\n# B\n@t=fade:99\n- a\n")
    rules = [x.rule for x in d.diagnostics]
    assert rules.count("bad-transition") == 2
    assert "bad-header" in rules
    assert all(x.hint for x in d.diagnostics if x.rule in ("bad-transition", "bad-header"))
    assert d.slides[0].transition is None


def test_parse_element_build_class():
    d = parse("# A\n- a {.build}\n")
    assert not [x for x in d.diagnostics if x.rule == "unknown-attr"]


# --------------------------------------------------------------------------- animations


def test_build_by_paragraph(tmp_path):
    _, _, prs = make("# A\n@build\n- one\n  - sub\n- two\n- three\n", tmp_path)
    s = prs.slides[0]
    assert len(clicks(s)) == 3
    assert s._element.xpath("p:timing//p:bldP[@build='p']")
    ranges = s._element.xpath("p:timing//p:pRg")
    assert [(r.get("st"), r.get("end")) for r in ranges] == [("0", "1"), ("2", "2"), ("3", "3")]
    assert s._element.xpath("p:timing//p:cTn[@presetClass='entr' and @presetID='1']")
    ids = [int(i) for i in s._element.xpath("p:timing//p:cTn/@id")]
    assert len(ids) == len(set(ids))
    spids = {int(i) for i in s._element.xpath("p:timing//p:spTgt/@spid")}
    assert spids <= {sh.shape_id for sh in s.shapes}


def test_build_blocks_and_cards(tmp_path):
    md = (
        "# A\n@build 2\n## Card one\n- x\n- y\n\n## Card two\n- z\n\n---\n"
        "# B\n@build\n| a | b |\n|---|---|\n| 1 | 2 |\n\n```py\nx=1\n```\n"
    )
    _, _, prs = make(md, tmp_path)
    assert len(clicks(prs.slides[0])) == 2  # one click per card
    assert len(clicks(prs.slides[1])) == 2  # table, code
    # a card appears as a group: later targets of a click are withEffect
    assert prs.slides[0]._element.xpath("p:timing//p:cTn[@nodeType='withEffect']")


def test_element_build_on_plain_slide(tmp_path):
    _, _, prs = make("# A\n- a\n- b\n\n{.build}\n| x | y |\n|---|---|\n| 1 | 2 |\n", tmp_path)
    s = prs.slides[0]
    assert s._element.xpath("p:timing")
    for c in clicks(s):
        assert xp(c, ".//p:spTgt")


def test_no_build_no_timing(tmp_path):
    _, _, prs = make("# A\n- a\n- b\n", tmp_path)
    assert not prs.slides[0]._element.xpath("p:timing")


def test_static_parts_not_animated(tmp_path):
    md = "# Title\n> lead\n@build\n- a\n- b\n\n> conclusion\n"
    _, _, prs = make(md, tmp_path)
    s = prs.slides[0]
    animated = {int(i) for i in s._element.xpath("p:timing//p:spTgt/@spid")}
    names = {sh.shape_id: sh.name for sh in s.shapes}
    assert all(names[i] not in ("Title", "Lead", "Conclusion") for i in animated)


# --------------------------------------------------------------------------- transitions


@pytest.mark.parametrize(
    "name,tag",
    [("fade", "fade"), ("push", "push"), ("wipe", "wipe"), ("split", "split"), ("cover", "cover")]
    + [("zoom", "zoom")],
)
def test_transitions(tmp_path, name, tag):
    _, _, prs = make(f"# A\n@t={name}:0.5\n- a\n\n# B\n@t={name}\n- b\n", tmp_path)
    x0 = prs.slides[0]._element
    assert xp(x0, f"mc:AlternateContent/mc:Choice/p:transition[@p14:dur='500']/p:{tag}")
    assert xp(x0, f"mc:AlternateContent/mc:Fallback/p:transition/p:{tag}")
    assert xp(prs.slides[1]._element, f"p:transition/p:{tag}")


def test_morph_alternate_content_and_order(tmp_path):
    _, _, prs = make("# A\n@build t=morph\n- a\n- b\n", tmp_path)
    x = prs.slides[0]._element
    assert xp(x, "mc:AlternateContent/mc:Choice[@Requires='p159']/p:transition/p159:morph")
    assert xp(x, "mc:AlternateContent/mc:Fallback/p:transition/p:fade")
    tags = [c.tag.split("}")[1] for c in x]
    assert tags[:4] == ["cSld", "clrMapOvr", "AlternateContent", "timing"]


# --------------------------------------------------------------------------- sections


def test_sections(tmp_path):
    md = "# Deck\n\n---\n# Intro\n- a\n\n---\n# Part 1\n\n---\n# P1 a\n- x\n\n---\n# P1 b\n- y\n"
    deck, _, prs = make(md, tmp_path)
    secs = xp(prs.part._element, "p:extLst/p:ext/p14:sectionLst/p14:section")
    names = [s.get("name") for s in secs]
    assert names == ["Deck", "Part 1"] or names[0] in ("Deck", "Default")
    ids = [e.get("id") for e in prs.slides._sldIdLst]
    flat = [i.get("id") for s in secs for i in xp(s, "p14:sldIdLst/p14:sldId")]
    assert flat == ids
    assert secs[-1].get("name") == "Part 1"
    assert len(xp(secs[-1], "p14:sldIdLst/p14:sldId")) == 3


def test_sections_off_or_absent(tmp_path):
    _, _, prs = make("sections: off\n\n# A\n- a\n\n---\n# S\n\n---\n# B\n- b\n", tmp_path)
    assert not xp(prs.part._element, "p:extLst/p:ext/p14:sectionLst")
    _, _, prs = make("# A\n- a\n\n---\n# B\n- b\n", tmp_path)
    assert not xp(prs.part._element, "p:extLst/p:ext/p14:sectionLst")


# --------------------------------------------------------------------------- jumps / hidden


def test_jumps_by_number_and_id_and_bad(tmp_path):
    md = "# A\n- [n](#2) [i](#target) [bad](#nope)\n\n---\n# B\n@id=target\n- b\n\n---\n# C\n- c\n@hidden\n"
    deck, _, prs = make(md, tmp_path)
    s0 = prs.slides[0]
    xml = s0._element.xml
    assert xml.count("ppaction://hlinksldjump") == 2
    rels = [r for r in s0.part.rels.values() if r.reltype.endswith("/slide")]
    assert {r.target_part.partname for r in rels} == {prs.slides[1].part.partname}
    bad = [d for d in deck.diagnostics if d.rule == "bad-jump"]
    assert len(bad) == 1 and bad[0].level == "info" and bad[0].hint
    assert prs.slides[2]._element.get("show") == "0"
    assert prs.slides[0]._element.get("show") is None


@pytest.mark.skipif(not have_soffice(), reason="needs LibreOffice")
def test_opens_in_libreoffice(tmp_path):
    md = (
        "# Deck\n\n---\n# Part\n\n---\n# Body\n@build t=fade:0.5\n- one\n- two\n\n---\n"
        "# Cards\n@build 2\n## A\n- x\n\n## B\n- y\n"
    )
    _, out, _ = make(md, tmp_path)
    pngs = pptx_to_pngs(out, tmp_path / "png")
    assert len(pngs) == 4
