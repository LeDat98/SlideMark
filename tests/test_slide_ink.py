"""Dark slides flip their text ink; cover `##` lines; task lists; katakana words stay whole."""

from __future__ import annotations

from pptx import Presentation

from slidemark.build import build_deck
from slidemark.layout import measure
from slidemark.parser import parse
from slidemark.theme import get_theme, slide_ink

HEAD = "theme: none\ncolors: bg=#FFFFFF fg=#1F1530 primary=#7A1FA2\n\n"


def built(md, tmp_path):
    deck = parse(md)
    out = tmp_path / "o.pptx"
    build_deck(deck, out, tmp_path)
    return Presentation(str(out)), deck


def run_colors(prs, slide=0):
    out = {}
    for sh in prs.slides[slide].shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            r = sh.text_frame.paragraphs[0].runs[0]
            out[sh.text_frame.text.split("\n")[0]] = str(r.font.color.rgb)
    return out


def rules(deck):
    return {d.rule for d in deck.diagnostics}


# ---- pure helper


def test_slide_ink_flips_on_dark_bg_and_keeps_light():
    th = get_theme("none")
    assert slide_ink(th, "#2E7D32", [])["color"] == th.render.ink_light
    assert slide_ink(th, "#0B1020", [])["color"] == th.render.ink_light
    assert slide_ink(th, "#FFF3E0", []) == {}  # the theme fg already reads
    assert slide_ink(th, "linear-gradient(135deg, #1A0B2E, #7A1FA2)", [])["color"] == th.render.ink_light
    assert slide_ink(th, None, ["dark"])["color"] == th.render.ink_light  # class always works
    assert slide_ink(th, "#0B1020", ["light"])["color"] == th.render.ink_dark
    assert slide_ink(th, None, []) == {}
    assert slide_ink(th, "url(x.png)", []) == {}  # a picture cannot be judged


def test_muted_ink_stays_readable():
    from slidemark.lint import _hex, contrast_ratio

    th = get_theme("none")
    ink = slide_ink(th, "#0B1020", [])
    assert ink["muted"] != ink["color"]
    assert contrast_ratio(_hex(ink["muted"], th), _hex("#0B1020", th)) >= 4.5


# ---- end to end: no css written by the author


def test_dark_bg_cover_and_gradient_cover_are_readable(tmp_path):
    md = HEAD + '# Aurelia\n@cover bg="linear-gradient(135deg, #1A0B2E, #7A1FA2)"\nHear more\nLaunch 2026\n'
    prs, deck = built(md, tmp_path)
    cols = run_colors(prs)
    assert cols["Aurelia"] == "FFFFFF"
    assert "contrast" not in rules(deck)
    md = "theme: none\ncolors: bg=#FFFFFF fg=#1B1B1B\n\n# Summit\n@cover bg=#2E7D32 dark\nGear\n"
    prs, deck = built(md, tmp_path)
    assert run_colors(prs)["Summit"] == "FFFFFF"
    assert "contrast" not in rules(deck)


def test_dark_class_without_theme_class_and_content_slide(tmp_path):
    md = HEAD + "# Night\n@bg=#101820 dark\n> lead line\n- one\n- two\n"
    prs, deck = built(md, tmp_path)
    cols = run_colors(prs)
    assert cols["Night"] == "FFFFFF"
    assert cols["one"] == "FFFFFF"
    assert "contrast" not in rules(deck)


def test_cards_keep_their_own_ink_on_a_dark_slide(tmp_path):
    md = HEAD + "# Night\n@bg=#101820\n@2\n## A\n- one\n## B\n- two\n"
    prs, deck = built(md, tmp_path)
    cols = run_colors(prs)
    assert cols["Night"] == "FFFFFF"
    assert cols["one"] == "1F1530"  # on the light card
    assert "contrast" not in rules(deck)


def test_declared_colors_win(tmp_path):
    md = HEAD + "# Night\n@cover bg=#101820\n```css\nh1 { color: #FFD700 }\n```\nsub line\n"
    prs, _ = built(md, tmp_path)
    cols = run_colors(prs)
    assert cols["Night"] == "FFD700"  # the slide's own rule
    assert cols["sub line"] != "1F1530"  # the rest still flips (to the muted light ink)
    md = HEAD + "# Night\n@cover bg=#101820\n```css\nslide { color: #FFD700 }\n```\nsub line\n"
    prs, _ = built(md, tmp_path)
    assert run_colors(prs)["Night"] == "FFD700"


def test_deck_wide_h1_color_yields_to_slide_bg(tmp_path):
    md = HEAD + "```css\nh1 { color: #7A1FA2 }\n```\n\n# Night\n@cover bg=#101820\nsub\n"
    prs, deck = built(md, tmp_path)
    assert run_colors(prs)["Night"] == "FFFFFF"
    assert "contrast" not in rules(deck)


def test_light_slide_is_untouched(tmp_path):
    md = HEAD + "# Day\n@bg=#FFF3E0\n- one\n"
    prs, _ = built(md, tmp_path)
    assert run_colors(prs)["one"] == "1F1530"


def test_slide_ink_is_idempotent_and_never_raises():
    from slidemark.theme import apply_slide_ink

    deck = parse(HEAD + "# T\n@cover bg=#000000\nx\n")
    th = get_theme("none")
    apply_slide_ink(deck, th)
    n = len(deck.slides[0].css)
    apply_slide_ink(deck, th)
    assert len(deck.slides[0].css) == n
    apply_slide_ink(object(), th)  # junk: no exception


# ---- `##` on a cover is a subtitle line


def test_cover_heading_is_a_subtitle_not_a_card():
    deck = parse("# Deck\n@cover\n## Sub line\nmore\n")
    s = deck.slides[0]
    assert s.elements == []
    assert [p.plain for p in s.subtitle.paragraphs] == ["Sub line", "more"]
    d = next(d for d in deck.diagnostics if d.rule == "cover-heading")
    assert d.level == "info" and d.hint and "\n" not in d.hint
    deck = parse("# Deck\nbody line\n## Sub line\n")  # inferred cover
    assert deck.slides[0].layout == "cover" and deck.slides[0].elements == []


def test_content_slide_boxes_stay_boxes():
    deck = parse("# Deck\n## A\n- x\n- y\n## B\n- z\n")
    assert len(deck.slides[0].elements) == 2
    assert not any(d.rule == "cover-heading" for d in deck.diagnostics)
    deck = parse("# Deck\n## Only\n- x\n")  # a lone box with a body is not a cover
    assert deck.slides[0].elements


# ---- task lists


def test_task_list_items_get_box_glyphs(tmp_path):
    deck = parse("# T\n- [ ] open\n- [x] done\n- [X] done too\n- [x]{.badge} badge stays\n- plain\n")
    ps = deck.slides[0].elements[0].paragraphs
    assert ps[0].plain.startswith("☐") and ps[0].plain.endswith("open") and ps[0].marker is None
    assert ps[1].plain.startswith("☑") and ps[2].plain.startswith("☑")
    assert ps[3].marker == "bullet" and not ps[3].plain.startswith("[")
    assert ps[4].marker == "bullet"
    prs, _ = built("# T\n- [ ] open\n- [x] done\n", tmp_path)
    text = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "☐" in text and "☑" in text and "[ ]" not in text


# ---- katakana words are not cut inside


def test_katakana_run_is_one_wrap_unit():
    text = "拠点とアプリで展開"
    seg = [(text, False, False)]
    # 5.5 em: char-by-char wrapping needs 2 lines (拠点とアプ|リで展開); a whole word does not fit the rest
    assert measure.count_lines(seg, 5.5 * 10, 10) == 3
    assert measure.count_lines(seg, 6.2 * 10, 10) == 2  # 拠点とアプリ|で展開
    assert measure.count_lines([("アプリケーション", False, False)], 4 * 10, 10) == 2  # too long: cut anyway
    # Latin glued to katakana is not one word; kanji next to katakana still wraps per character
    assert measure.count_lines([("Webアプリ", False, False)], 3.5 * 10, 10) >= 2
