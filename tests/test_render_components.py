"""Renderer: connectors, callouts, badges, heading bands, KPI cards, leftover tables (reopens the .pptx)."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

from slidemark.ir import Cell, Container, Deck, Link, Paragraph, Run, Slide, Table, Text
from slidemark.layout import layout_slide
from slidemark.preview import pptx_to_pngs
from slidemark.render import render
from slidemark.theme import get_theme

from .helpers import needs_soffice


def T(s, role="body", **kw):
    return Text(role=role, paragraphs=[Paragraph(runs=[Run(text=s)])], **kw)


def box(h, body="text", **kw):
    return Container(title=T(h, "heading"), children=[T(body)], **kw)


def build(deck, tmp_path, theme="default"):
    th = get_theme(theme)
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    out = tmp_path / "out.pptx"
    render(deck, placed, th, out)
    return Presentation(str(out)), out, th


def connectors(prs):
    return [s for s in prs.slides[0].shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]


def test_connectors_are_real_lines_with_arrow_heads(tmp_path):
    s = Slide(
        title=T("t", "title"),
        grid="3",
        elements=[box("a"), box("b"), box("c")],
        links=[Link(src=0, dst=1), Link(src=1, dst=2, arrow=False), Link(src=2, dst=0)],
    )
    prs, _, th = build(Deck(slides=[s]), tmp_path)
    shapes = list(prs.slides[0].shapes)
    cards = [x for x in shapes if x.name.startswith("Card")]
    cx = connectors(prs)
    assert len(cx) == 3
    assert shapes.index(cx[0]) > max(shapes.index(c) for c in cards)  # on top of the blocks
    a, b, c = cards
    assert (cx[0].begin_x, cx[0].end_x) == (a.left + a.width, b.left)
    assert cx[0].begin_y == cx[0].end_y == a.top + a.height // 2
    assert "a:tailEnd" in cx[0]._element.xml and 'type="triangle"' in cx[0]._element.xml
    assert "a:tailEnd" not in cx[1]._element.xml
    assert cx[0].line.width == Pt(1.5)
    assert str(cx[0].line.color.rgb) == th.color("primary").lstrip("#")
    assert cx[2].begin_x > cx[2].end_x  # c -> a runs right to left (flipH)
    assert "p:style" not in cx[0]._element.xml  # no theme shadow


def test_vertical_connector_flips(tmp_path):
    s = Slide(
        grid="1x2",
        layout="blank",
        elements=[box("a"), box("b")],
        links=[Link(src=1, dst=0)],
    )
    prs, _, _ = build(Deck(slides=[s]), tmp_path)
    (cx,) = connectors(prs)
    assert cx.begin_y > cx.end_y and cx.begin_x == cx.end_x
    assert 'flipV="1"' in cx._element.xml


def test_container_links_render(tmp_path):
    inner = Container(classes=["plain"], grid="2", children=[box("x"), box("y")], links=[Link(src=0, dst=1)])
    s = Slide(elements=[Container(title=T("h", "heading"), children=[inner])])
    prs, _, _ = build(Deck(slides=[s]), tmp_path)
    assert len(connectors(prs)) == 1


@pytest.mark.parametrize("kind", ["note", "tip", "warn", "caution"])
def test_callout_card_has_tint_border_and_accent_bar(tmp_path, kind):
    s = Slide(title=T("t", "title"), elements=[T("Careful", classes=["callout", kind])])
    prs, _, th = build(Deck(slides=[s]), tmp_path)
    shapes = list(prs.slides[0].shapes)
    card = next(x for x in shapes if x.has_text_frame and x.text_frame.text == "Careful")
    bar = next(x for x in shapes if x.name.endswith("accent"))
    line = th.classes[kind].line
    assert str(card.line.color.rgb) == th.color(line).lstrip("#")
    assert str(card.fill.fore_color.rgb) != th.color(line).lstrip("#")
    assert str(bar.fill.fore_color.rgb) == th.color(line).lstrip("#")
    assert bar.left == card.left and bar.top == card.top and bar.height == card.height
    assert bar.width == Pt(4)
    assert shapes.index(bar) > shapes.index(card)
    assert card.text_frame.margin_left >= Pt(4)


def test_callout_inside_a_box(tmp_path):
    inner = [T("one", classes=["callout", "tip"]), T("two", classes=["callout", "warn"])]
    s = Slide(elements=[Container(title=T("h", "heading"), children=inner)])
    prs, _, _ = build(Deck(slides=[s]), tmp_path, "midnight")
    assert len([x for x in prs.slides[0].shapes if x.name.endswith("accent")]) == 2


def test_badge_run_is_highlighted_bold_with_theme_color(tmp_path):
    runs = [
        Run(text="Status "),
        Run(text="OK", highlight="success", color="bg", bold=True),
        Run(text=" "),
        Run(text="HOT", highlight="accent"),
    ]
    s = Slide(elements=[Text(paragraphs=[Paragraph(runs=runs)])])
    prs, _, th = build(Deck(slides=[s]), tmp_path)
    tb = next(x for x in prs.slides[0].shapes if x.has_text_frame)
    r_ok, r_hot = tb.text_frame.paragraphs[0].runs[1], tb.text_frame.paragraphs[0].runs[3]
    rpr = r_ok._r.rPr
    hl = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}highlight")
    assert hl is not None and hl[0].get("val") == th.color("success").lstrip("#")
    assert r_ok.font.bold and str(r_ok.font.color.rgb) == th.color("bg").lstrip("#")
    # schema order: solidFill, highlight, latin, ea
    names = [c.tag.split("}")[1] for c in rpr]
    assert names.index("solidFill") < names.index("highlight") < names.index("latin") < names.index("ea")
    # no explicit color: bold, readable text color picked automatically
    assert r_hot.font.bold and r_hot.font.color.rgb is not None
    assert not tb.text_frame.paragraphs[0].runs[0].font.bold


def test_heading_band_drawn_on_jp_business(tmp_path):
    s = Slide(title=T("t", "title"), elements=[box("Heading", "content")])
    prs, _, th = build(Deck(slides=[s]), tmp_path, "jp-business")
    shapes = list(prs.slides[0].shapes)
    card = next(x for x in shapes if x.name.startswith("Card"))
    band = next(x for x in shapes if x.has_text_frame and x.text_frame.text == "Heading")
    assert str(band.fill.fore_color.rgb) == th.color(th.heading_band).lstrip("#")
    assert (band.left, band.top, band.width) == (card.left, card.top, card.width)
    run = band.text_frame.paragraphs[0].runs[0]
    assert run.font.bold and str(run.font.color.rgb) == th.color(th.heading_band_color).lstrip("#")
    assert shapes.index(band) > shapes.index(card)
    body = next(x for x in shapes if x.has_text_frame and x.text_frame.text == "content")
    assert body.top >= band.top + band.height


def test_kpi_card_text(tmp_path):
    kpi = Container(
        classes=["kpi"],
        title=T("Revenue", "heading"),
        children=[
            Text(paragraphs=[Paragraph(runs=[Run(text="12.4M")]), Paragraph(runs=[Run(text="+8% YoY")])])
        ],
    )
    s = Slide(title=T("t", "title"), grid="3", elements=[kpi, kpi.model_copy(), kpi.model_copy()])
    prs, _, th = build(Deck(slides=[s]), tmp_path)
    tb = next(x for x in prs.slides[0].shapes if x.has_text_frame and x.text_frame.text.startswith("12.4M"))
    big, cap = tb.text_frame.paragraphs
    assert big.runs[0].font.size == Pt(th.classes["kpi"].font_size)
    assert big.runs[0].font.bold and str(big.runs[0].font.color.rgb) == th.color("primary").lstrip("#")
    assert cap.runs[0].font.size < big.runs[0].font.size
    assert str(cap.runs[0].font.color.rgb) == th.color("muted").lstrip("#")
    assert big.alignment == cap.alignment and big.alignment is not None


def test_leftover_table_has_five_columns(tmp_path):
    rows = [[Cell(paragraphs=[Paragraph(runs=[Run(text=t)])]) for t in "abcde"]]
    rows.append(
        [Cell(paragraphs=[Paragraph(runs=[Run(text=t)])]) for t in "xy"]
        + [Cell(paragraphs=[Paragraph(runs=[Run(text="m")])], colspan=3), Cell(), Cell()]
    )
    s = Slide(
        title=T("t", "title"),
        grid="4",
        classes=["chevron"],
        elements=[box(f"P{i}") for i in range(4)] + [Table(rows=rows)],
    )
    prs, _, _ = build(Deck(slides=[s]), tmp_path)
    tbl = next(x for x in prs.slides[0].shapes if x.has_table)
    assert len(tbl.table.columns) == 5
    assert tbl.table.cell(1, 2).span_width == 3
    chev = [
        x for x in prs.slides[0].shapes if x.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and "Shape" in x.name
    ]
    assert tbl.top >= max(c.top + c.height for c in chev)


@needs_soffice
@pytest.mark.parametrize("theme", ["default", "jp-business", "midnight"])
def test_components_deck_opens_in_libreoffice(tmp_path, theme):
    kpi = Container(
        classes=["kpi"],
        title=T("ARR", "heading"),
        children=[Text(paragraphs=[Paragraph(runs=[Run(text="$4.2M")]), Paragraph(runs=[Run(text="+12%")])])],
    )
    runs = [Run(text="Badge "), Run(text="NEW", highlight="success", color="bg")]
    s = Slide(
        title=T("Components", "title"),
        grid="aab/aac",
        elements=[
            Container(
                title=T("Flow", "heading"),
                grid="3",
                children=[box("a"), box("b"), box("c")],
                links=[Link(src=0, dst=1), Link(src=1, dst=2)],
            ),
            Container(
                title=T("Notes", "heading"),
                children=[
                    T("watch out", classes=["callout", "warn"]),
                    Text(paragraphs=[Paragraph(runs=runs)]),
                ],
            ),
            kpi,
            T("Tail", classes=["callout", "tip"]),
        ],
    )
    _, out, _ = build(Deck(slides=[s]), tmp_path, theme)
    assert len(pptx_to_pngs(out, tmp_path / "png")) == 1
