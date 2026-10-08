"""Video / audio: parser, layout, render (reopened with python-pptx), lint, LibreOffice, XSD."""

from __future__ import annotations

import pytest
from pptx import Presentation

from slidemark.build import build_deck
from slidemark.ir import Container, Media
from slidemark.parser import parse
from slidemark.parser.html import parse_html
from slidemark.xsd import _load_schema, validate_pptx

from .helpers import needs_soffice

RT_MEDIA = "http://schemas.microsoft.com/office/2007/relationships/media"
RT_VIDEO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/video"
RT_AUDIO = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/audio"


def _media(deck) -> list[Media]:
    out: list[Media] = []

    def walk(els):
        for e in els:
            if isinstance(e, Media):
                out.append(e)
            elif isinstance(e, Container):
                walk(e.children)

    for s in deck.slides:
        walk(s.elements)
    return out


@pytest.fixture
def files(tmp_path):
    (tmp_path / "demo.mp4").write_bytes(b"\x00\x00\x00\x18ftypmp42fake")
    (tmp_path / "song.mp3").write_bytes(b"ID3fake")
    return tmp_path


# ------------------------------------------------------------------ parser


@pytest.mark.parametrize(
    ("src", "kind"),
    [
        ("a.mp4", "video"),
        ("a.MOV", "video"),
        ("a.webm?x=1", "video"),
        ("a.m4v#t=3", "video"),
        ("dir/a.wmv", "video"),
        ("a.avi", "video"),
        ("a.mp3", "audio"),
        ("a.WAV", "audio"),
        ("a.m4a?dl=1", "audio"),
        ("a.aac", "audio"),
        ("a.wma", "audio"),
    ],
)
def test_extension_detection(src, kind):
    ms = _media(parse(f"# T\n![alt]({src})\n"))
    assert len(ms) == 1 and ms[0].kind == kind and ms[0].src == src


def test_image_stays_image():
    deck = parse("# T\n![alt](a.png)\n")
    assert not _media(deck) and deck.slides[0].elements[-1].type == "image"


def test_attrs():
    (m,) = _media(parse("# T\n![demo](a.mp4){poster=p.png autoplay loop w=300 x=10}\n"))
    assert m.poster == "p.png" and m.autoplay and m.loop
    assert m.box is not None and m.box.w == 300 and m.box.x == 10


def test_attr_flag_false():
    (m,) = _media(parse("# T\n![demo](a.mp4){autoplay=false}\n"))
    assert not m.autoplay and not m.loop


def test_inside_box():
    deck = parse("# T\n## Box\ntext\n![clip](a.mp4)\n")
    (m,) = _media(deck)
    assert m.alt == "clip"


def test_html_video_and_audio():
    deck = parse_html(
        '<section><h1>T</h1><video src="a.mp4" poster="p.png" autoplay loop title="clip"></video>'
        '<audio><source src="b.mp3"></audio></section>'
    )
    a, b = _media(deck)
    assert (a.kind, a.src, a.poster, a.autoplay, a.loop, a.alt) == (
        "video",
        "a.mp4",
        "p.png",
        True,
        True,
        "clip",
    )
    assert (b.kind, b.src) == ("audio", "b.mp3")


def test_json_roundtrip():
    deck = parse("# T\n![d](a.mp4){loop}\n")
    again = type(deck).model_validate_json(deck.model_dump_json())
    assert _media(again)[0].loop


# ------------------------------------------------------------------ render


def _shapes_with_media(path):
    prs = Presentation(str(path))
    out = []
    for s in prs.slides:
        for shp in s.shapes:
            if shp._element.xpath(".//a:videoFile | .//a:audioFile"):
                out.append((s, shp))
    return out


def test_video_render(files):
    out = files / "v.pptx"
    md = "# T\n![demo](demo.mp4){poster=none.png autoplay loop}\n"
    deck = parse(md)
    build_deck(deck, out, base_dir=files)
    ((slide, shp),) = _shapes_with_media(out)
    rels = [r.reltype for r in slide.part.rels.values()]
    assert RT_MEDIA in rels and RT_VIDEO in rels
    assert any(r.reltype.endswith("/image") for r in slide.part.rels.values())  # generated poster
    assert shp._element.xpath(".//a:videoFile") and shp.shape_type is not None
    assert shp._element.xpath(".//p:cNvPr/@descr") == ["demo"]
    xml = slide._element.xml
    assert "playFrom(0.0)" in xml and 'repeatCount="indefinite"' in xml
    assert any(d.rule == "missing-media" for d in deck.diagnostics)  # poster not found
    ids = slide._element.xpath("//p:cTn/@id")
    assert len(ids) == len(set(ids))


def test_video_without_flags_has_no_autoplay(files):
    out = files / "v.pptx"
    build_deck(parse("# T\n![demo](demo.mp4)\n"), out, base_dir=files)
    ((slide, _),) = _shapes_with_media(out)
    assert "playFrom" not in slide._element.xml and "<p:video>" in slide._element.xml.replace(
        ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"', ""
    )


def test_audio_render(files):
    out = files / "a.pptx"
    build_deck(parse("# T\n![song](song.mp3){loop}\n"), out, base_dir=files)
    ((slide, shp),) = _shapes_with_media(out)
    rels = [r.reltype for r in slide.part.rels.values()]
    assert RT_AUDIO in rels and RT_MEDIA in rels and RT_VIDEO not in rels
    assert shp._element.xpath(".//a:audioFile") and not shp._element.xpath(".//a:videoFile")
    assert shp.width == shp.height <= 1_100_000  # icon-sized


def test_two_videos_and_build_animation(files):
    out = files / "b.pptx"
    md = "# T {.build}\n- a\n- b\n\n![x](demo.mp4){autoplay}\n![y](song.mp3)\n"
    build_deck(parse(md), out, base_dir=files)
    found = _shapes_with_media(out)
    assert len(found) == 2
    xml = found[0][0]._element
    assert len(xml.xpath("//p:timing")) == 1
    ids = xml.xpath("//p:cTn/@id")
    assert len(ids) == len(set(ids))


def test_missing_file_placeholder_and_diagnostic(tmp_path):
    out = tmp_path / "m.pptx"
    deck = parse("# T\n![clip](nope.mp4)\n")
    build_deck(deck, out, base_dir=tmp_path)
    prs = Presentation(str(out))
    texts = [shp.text_frame.text for shp in prs.slides[0].shapes if shp.has_text_frame]
    assert any("[video: clip]" in t for t in texts)
    (d,) = [d for d in deck.diagnostics if d.rule == "missing-media"]
    assert d.hint and "nope.mp4" in d.message


def test_url_source_is_a_placeholder(tmp_path):
    out = tmp_path / "u.pptx"
    deck = parse("# T\n![clip](https://example.com/a.mp4)\n")
    build_deck(deck, out, base_dir=tmp_path)
    assert any(d.rule == "missing-media" for d in deck.diagnostics)


def test_unreadable_media_never_raises(tmp_path):
    (tmp_path / "dir.mp4").mkdir()  # a directory is not a file: placeholder
    out = tmp_path / "d.pptx"
    build_deck(parse("# T\n![clip](dir.mp4)\n"), out, base_dir=tmp_path)
    assert out.exists()


def test_inside_box_renders(files):
    out = files / "bx.pptx"
    build_deck(parse("# T\n## Box\ntext\n![clip](demo.mp4)\n## Two\ntext\n"), out, base_dir=files)
    assert len(_shapes_with_media(out)) == 1


# ------------------------------------------------------------------ lint, LibreOffice, XSD


def test_lint_alt(files):
    deck = parse("# T\n![](demo.mp4)\n")
    build_deck(deck, files / "l.pptx", base_dir=files)
    rules = {d.rule for d in deck.diagnostics}
    assert "image-alt" in rules and "alt" in rules


@needs_soffice
def test_opens_in_libreoffice(files):
    from slidemark.preview import pptx_to_pngs

    out = files / "lo.pptx"
    build_deck(parse("# T\n![demo](demo.mp4){autoplay loop}\n"), out, base_dir=files)
    pngs = pptx_to_pngs(out, files / "png")
    assert pngs and pngs[0].stat().st_size > 0


@pytest.mark.skipif(_load_schema() is None, reason="ECMA-376 schemas not available")
def test_xsd_valid(files):
    out = files / "x.pptx"
    build_deck(parse("# T\n![demo](demo.mp4){autoplay loop}\n![s](song.mp3)\n"), out, base_dir=files)
    assert validate_pptx(out) == []


@pytest.mark.skipif(_load_schema() is None, reason="ECMA-376 schemas not available")
def test_xsd_valid_without_autoplay(files):
    # no autoplay and no build animation: an empty main sequence would be a schema error
    out = files / "y.pptx"
    build_deck(parse("# T\n![demo](demo.mp4)\n- a\n"), out, base_dir=files)
    assert validate_pptx(out) == []


# ------------------------------------------------------------------ importer


def test_import_roundtrip(files):
    from slidemark.importer import import_pptx

    out = files / "i.pptx"
    build_deck(parse("# T\n![demo](demo.mp4)\n![s](song.mp3)\n"), out, base_dir=files)
    md = import_pptx(out, out_dir=files / "imp")
    text = md if isinstance(md, str) else md[0]
    assert "](images/1-1.mp4)" in text and "](images/1-2.mp3)" in text
    assert (files / "imp" / "images" / "1-1.mp4").read_bytes() == (files / "demo.mp4").read_bytes()
    assert len(_media(parse(text))) == 2
