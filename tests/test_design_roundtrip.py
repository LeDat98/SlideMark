"""DF5 round trip: the deck's design (theme, tokens, css fences, @html slides) survives build -> import."""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.ir import Style
from slidemark.parser import parse
from slidemark.parser.css import parse_declarations
from slidemark.parser.ctx import Ctx
from slidemark.render.design_part import read_design_part, style_to_css

ROOT = Path(__file__).resolve().parent.parent
EXAMPLES = ROOT / "examples"
DESIGNED = ["13-brand-aurora", "14-brand-terracotta", "15-html-mixed"]
MIXED = EXAMPLES / "15-html-mixed.md"


def _rules(rs):
    return [(r.selector, r.style.model_dump(exclude_none=True)) for r in rs]


@pytest.mark.parametrize("name", DESIGNED)
def test_design_survives_round_trip(tmp_path, name):
    text = (EXAMPLES / f"{name}.md").read_text(encoding="utf-8")
    a, b = tmp_path / "A.pptx", tmp_path / "B.pptx"
    build(text, a, base_dir=EXAMPLES)
    t1, _ = import_pptx(a, tmp_path)
    build(t1, b, base_dir=tmp_path)
    t2, _ = import_pptx(b, tmp_path)
    d0, d1 = parse(text), parse(t1)
    assert d0.theme == d1.theme
    assert d0.tokens == d1.tokens
    assert _rules(d0.css) == _rules(d1.css)
    assert len(d0.slides) == len(d1.slides)
    for x, y in zip(d0.slides, d1.slides, strict=True):
        assert x.html == y.html
        assert _rules(x.css) == _rules(y.css)
    assert t1 == t2  # stable: the second import is identical to the first


def test_html_slide_is_a_fence_not_loose_boxes(tmp_path):
    a = tmp_path / "A.pptx"
    build(MIXED.read_text(encoding="utf-8"), a)
    t, diags = import_pptx(a)
    assert "@html" in t and "```html" in t
    assert not [d for d in diags if d.rule == "import-html-edited"]


def test_token_lines_use_shortest_groups(tmp_path):
    a = tmp_path / "A.pptx"
    src = (
        'theme: none\ncolors: primary=#7C5CFF\nstyle: card.radius=14 card.shadow="0 8 24 #00000066"\n\n# T\n'
    )
    src += "- a\n"
    build(src, a)
    t, _ = import_pptx(a)
    head = t.split("\n\n")[0].splitlines()
    assert head[0] == "theme: none"
    assert "colors: primary=#7C5CFF" in head
    style = next(h for h in head if h.startswith("style:"))
    assert 'shadow="0 8 24 #00000066"' in style and "radius=14" in style


def test_no_design_no_part(tmp_path):
    from pptx import Presentation

    a = tmp_path / "A.pptx"
    build("# T\n- a\n", a)
    assert read_design_part(Presentation(str(a))) is None
    assert not [n for n in zipfile.ZipFile(a).namelist() if n.startswith("customXml/")]


def test_part_is_a_customxml_item(tmp_path):
    a = tmp_path / "A.pptx"
    build(MIXED.read_text(encoding="utf-8"), a)
    z = zipfile.ZipFile(a)
    names = z.namelist()
    assert "customXml/item1.xml" in names and "customXml/itemProps1.xml" in names
    assert b"urn:slidemark:design" in z.read("customXml/item1.xml")
    assert b"/customXml/itemProps1.xml" in z.read("[Content_Types].xml")


def _edit_first_slide(src: Path, dst: Path, old: str, new: str) -> None:
    from pptx import Presentation

    prs = Presentation(str(src))
    hit = False
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if old in r.text:
                        r.text = r.text.replace(old, new)
                        hit = True
    assert hit
    prs.save(str(dst))


def test_edited_html_slide_prefers_the_edited_shapes(tmp_path):
    a, e = tmp_path / "A.pptx", tmp_path / "E.pptx"
    build(MIXED.read_text(encoding="utf-8"), a)
    _edit_first_slide(a, e, "Light-based computing", "Photon-based computing")
    t, diags = import_pptx(e)
    assert "```html" not in t and "Photon-based computing" in t
    assert any(d.rule == "import-html-edited" and d.level == "info" and d.slide == 1 for d in diags)


def test_reordered_slides_keep_their_design(tmp_path):
    from pptx import Presentation

    a, r = tmp_path / "A.pptx", tmp_path / "R.pptx"
    build(MIXED.read_text(encoding="utf-8"), a)
    prs = Presentation(str(a))
    lst = prs.slides._sldIdLst
    first = lst[0]
    lst.remove(first)
    lst.append(first)  # the html slide moves to the end
    prs.save(str(r))
    t, _ = import_pptx(r)
    d = parse(t)
    assert d.slides[-1].html is not None and d.slides[0].html is None


def test_corrupt_part_is_ignored(tmp_path):
    a, c = tmp_path / "A.pptx", tmp_path / "C.pptx"
    build(MIXED.read_text(encoding="utf-8"), a)
    with zipfile.ZipFile(a) as zi, zipfile.ZipFile(c, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            data = zi.read(item.filename)
            if item.filename == "customXml/item1.xml":
                data = data.replace(b'{"v":1', b'{"v":')
            zo.writestr(item, data)
    t, diags = import_pptx(c)
    assert t.strip() and not [d for d in diags if d.level == "error"]


SAMPLES = [
    "color: #7C5CFF; font-size: 54pt; letter-spacing: 1pt; font-weight: bold; font-style: italic",
    "background: linear-gradient(135deg, #6D4AFF, #127A86); border-radius: 14px; box-shadow: 0 8px 24px #0006",  # noqa: E501
    "border: 2px dashed #00D1B2; opacity: 0.8; padding: 8px 16px; margin: 4px; gap: 12px",
    "border-top: 3px solid #00D1B2; border-bottom: none; padding-left: 20px",
    "text-align: center; vertical-align: middle; text-transform: uppercase; line-height: 1.4",
    "text-decoration: underline line-through; transform: rotate(-2deg); font-family: 'Noto Sans JP'",
    "grid-template-columns: 1fr 2fr",
    "grid-template-columns: repeat(3, 1fr)",
    'grid-template-areas: "a a b" "a a c"',
    "box-shadow: none; background: none; font-weight: normal",
    "border-width: 0; text-decoration: none",
    "background: radial-gradient(circle at 20% 30%, #2E2380, #0B1020 70%)",
]


@pytest.mark.parametrize("decl", SAMPLES)
def test_style_to_css_is_exact(decl):
    names = {"primary"}
    s1, _ = parse_declarations(decl, 1, Ctx([]), names)
    css = style_to_css(s1)
    s2, _ = parse_declarations(css, 1, Ctx([]), names)
    assert s2 == s1, css
    assert style_to_css(s2) == css


def test_empty_style_is_empty_css():
    assert style_to_css(Style()) == ""


def test_token_class_on_a_box_is_recovered(tmp_path):
    text = "style: hero.fill=#112233 hero.color=#FFFFFF\n\n# T\n## A {.hero}\n- x\n## B\n- y\n"
    a = tmp_path / "A.pptx"
    build(text, a)
    t, _ = import_pptx(a, tmp_path)
    assert "## A {.hero}" in t
    assert "## B\n" in t or t.rstrip().endswith("## B\n- y")
    assert parse(t).slides[0].elements[0].classes == ["hero"]


def test_css_id_and_class_on_a_box_are_recovered(tmp_path):
    text = (
        "```css\n#top { color: #C00 }\n.note { color: #080 }\n```\n\n# T\n## A {#top .note}\n- x\n## B\n- y\n"
    )
    a = tmp_path / "A.pptx"
    build(text, a)
    t, _ = import_pptx(a, tmp_path)
    box = parse(t).slides[0].elements[0]
    assert box.id == "top"
    assert "note" in box.classes


def test_template_theme_path_follows_the_imported_file(tmp_path):
    text = (EXAMPLES / "10-template.md").read_text(encoding="utf-8")
    a = tmp_path / "A.pptx"
    build(text, a, base_dir=EXAMPLES)
    t, diags = import_pptx(a, tmp_path)
    assert t.startswith("theme: A.pptx")
    assert any(d.rule == "import-template" for d in diags)
