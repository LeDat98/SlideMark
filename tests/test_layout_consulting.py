"""Consulting-grade layout: card heights, balanced bands, table rows, title gap, stacked boxes, chevrons."""

from __future__ import annotations

from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark.ir import Chart, Container, Deck, Shape, Slide, Table, Text
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.theme import get_theme

from .test_layout_policies import H, T, box, bullets, cards, cell, lay, of


def _title(placed):
    return next(p for p in placed if isinstance(p.element, Text) and p.element.role == "title")


def _gap_after_title(placed):
    t = _title(placed)
    top = min(
        p.y for p in placed if p is not t and p.y >= t.y + t.h - 1 and p.h > 0 and p.element is not None
    )
    return top - (t.y + t.h)


def test_sparse_cards_follow_content_not_the_body():
    s = Slide(title=T("t", "title"), grid="3", elements=[box("a", "x"), box("b", "x"), box("c", "x")])
    placed, _ = lay(s, "default")
    nat = 18 * 1.2 * 12700 * 2 + 2 * 10 * 12700  # heading + one bullet, rough
    cs = cards(placed)
    assert cs[0].h < 0.45 * H and cs[0].h >= 0.8 * nat


def test_two_by_two_with_conclusion_has_no_big_empty_band():
    s = Slide(
        title=T("t", "title"),
        grid="2x2",
        elements=[box(n, "x", "y") for n in "abcd"],
        conclusion=T("so what", "conclusion"),
    )
    placed, _ = lay(s, "jp-business")
    cs = cards(placed)
    bar = next(p for p in placed if isinstance(p.element, Text) and p.element.role == "conclusion")
    band = bar.y - max(c.y + c.h for c in cs)
    assert band < 0.12 * H
    assert min(c.y for c in cs) - (_title(placed).y + _title(placed).h) < 0.12 * H


def test_title_to_body_gap_is_identical_across_slide_kinds():
    dense = [f"line {i} " + "word " * 25 for i in range(9)]
    slides = [
        Slide(title=T("t", "title"), elements=[Chart(kind="line")]),
        Slide(title=T("t", "title"), grid="2", elements=[box("a", *dense), box("b", *dense)]),
        Slide(title=T("t", "title"), elements=[bullets(*dense)]),
        Slide(title=T("t", "title"), lead=T("lead", "lead"), elements=[Chart(kind="line")]),
    ]
    for theme in ("jp-business", "default"):
        gaps = [_gap_after_title(lay(s, theme)[0]) for s in slides[:3]]
        assert max(gaps) - min(gaps) <= 2, gaps
        assert gaps[0] > 0.15 * 914400  # a real gap, not touching the title band
    lead_gap = _gap_after_title(lay(slides[3], "jp-business")[0])
    assert lead_gap > 0


def test_flow_boxes_above_a_chart_take_natural_height():
    s = Slide(
        title=T("t", "title"),
        grid="flow",
        elements=[*(box(n, "one", "two") for n in "abc"), Chart(kind="line")],
    )
    placed, _ = lay(s, "jp-business")
    cs = cards(placed)
    (chart,) = of(placed, Chart)
    assert cs[0].h < 0.3 * H  # the box row hugs its text
    assert chart.h > 1.5 * cs[0].h  # the chart gets the rest


def test_stacked_boxes_share_height_by_content_and_align_with_the_tall_box():
    left = Container(title=T("L", "heading"), children=[Chart(kind="column")])
    short = box("s", "one")
    long = Container(title=T("l", "heading"), children=[bullets(*[f"item {i}" for i in range(5)])])
    s = Slide(title=T("t", "title"), grid="aab/aac", elements=[left, short, long])
    placed, _ = lay(s, "jp-business")
    cs = {id(p.element): p for p in cards(placed)}
    a, b, c = cs[id(left)], cs[id(short)], cs[id(long)]
    assert c.h > b.h  # more content -> more height
    assert abs((c.y + c.h) - (a.y + a.h)) <= 2 and b.y == a.y  # fills the column, bottoms align


def test_chevron_lists_have_no_bullet_markers(tmp_path):
    boxes = [Container(title=T(f"S{i}", "heading"), children=[bullets("alpha", "beta")]) for i in range(3)]
    s = Slide(title=T("t", "title"), grid="chevron", elements=boxes)
    th = get_theme("jp-business")
    deck = Deck(slides=[s])
    placed = layout_slide(s, deck, th, 0)
    chev = [p for p in placed if isinstance(p.element, Shape) and p.element.shape == "chevron"]
    assert len(chev) == 3
    assert all(par.marker is None for p in chev for par in p.element.paragraphs)
    out = tmp_path / "c.pptx"
    render(deck, [placed], th, out)
    prs = Presentation(str(out))
    for shp in prs.slides[0].shapes:
        if shp.has_text_frame and "alpha" in shp.text_frame.text:
            for para in shp.text_frame.paragraphs:
                ppr = para._p.pPr
                assert ppr is None or ppr.find(qn("a:buChar")) is None
                assert ppr is None or ppr.find(qn("a:buAutoNum")) is None


def test_table_rows_are_capped_near_text_height():
    rows = [[cell("項目"), cell("値")]] + [[cell("a"), cell("1")] for _ in range(3)]
    s = Slide(title=T("t", "title"), elements=[Table(rows=rows)])
    placed, _ = lay(s, "jp-business")
    (tp,) = of(placed, Table)
    size = (tp.style.font_size or 12) * tp.font_scale
    assert max(tp.element.attrs["_row_h"]) <= 2.0 * (1.8 * size * 1.2 + 7.3) * 12700  # TABLE_GROW_ROOMY


# ---- top anchoring, chevron / tree heights, table column widths (review pass 2)


def _body_top(placed):
    t = _title(placed)
    skip = {"title", "lead", "subtitle"}
    return min(
        p.y
        for p in placed
        if p.y >= t.y + t.h - 1
        and p.h > 0
        and not (isinstance(p.element, Text) and p.element.role in skip)
        and not (isinstance(p.element, Shape) and p.element.id == "band")
    )


def test_non_sparse_slide_is_top_anchored_at_title_gap():
    from slidemark.units import EMU_PER_INCH

    kpis = [
        Container(title=T(f"k{i}", "heading"), classes=["kpi"], children=[T("1.0\n2.0")]) for i in range(3)
    ]
    rows = [[cell("a"), cell("b")]] + [[cell("x"), cell("1")] for _ in range(3)]
    s = Slide(title=T("t", "title"), lead=T("lead", "lead"), grid="3", elements=[*kpis, Table(rows=rows)])
    placed, _ = lay(s, "jp-business")
    lead = next(p for p in placed if isinstance(p.element, Text) and p.element.role == "lead")
    assert _body_top(placed) - (lead.y + lead.h) == round(0.25 * EMU_PER_INCH)


def test_dense_two_boxes_slide_is_top_anchored_and_fuller():
    s = Slide(
        title=T("t", "title"),
        lead=T("lead", "lead"),
        grid="2",
        elements=[box("a", "one", "two", "three"), box("b", "one", "two", "three", "four")],
        classes=["dense"],
    )
    placed, _ = lay(s, "jp-business")
    lead = next(p for p in placed if isinstance(p.element, Text) and p.element.role == "lead")
    assert _body_top(placed) - (lead.y + lead.h) == round(0.25 * 914400)
    cs = cards(placed)
    assert cs[0].h >= 0.4 * H * 0.6  # not shrunk to its text


def test_very_sparse_slide_may_shift_by_at_most_a_third_of_the_leftover():
    s = Slide(title=T("t", "title"), grid="3", elements=[box("a", "x"), box("b", "x"), box("c", "x")])
    placed, _ = lay(s, "default")
    gap = _body_top(placed) - (_title(placed).y + _title(placed).h)
    assert gap > round(0.25 * 914400)  # shifted down ...
    assert gap < 0.25 * 914400 + 0.34 * H  # ... but never by more than a third of the body


def test_chevron_row_height_is_text_plus_padding():
    boxes = [Container(title=T(f"S{i}", "heading"), children=[T("short")]) for i in range(3)]
    s = Slide(title=T("t", "title"), grid="chevron", elements=boxes)
    placed, _ = lay(s, "jp-business")
    chev = [p for p in placed if isinstance(p.element, Shape) and p.element.shape == "chevron"]
    assert len(chev) == 3
    assert 0.7 * 914400 - 1 <= chev[0].h <= 1.3 * 914400 + 1
    size = (chev[0].style.font_size or 0) * chev[0].font_scale
    assert size >= 11  # at least the theme body size


def test_chevron_text_is_not_smaller_than_the_table_below():
    boxes = [Container(title=T(f"S{i}", "heading"), children=[T("short")]) for i in range(3)]
    rows = [[cell("a"), cell("b")]] + [[cell("x"), cell("1")] for _ in range(3)]
    s = Slide(title=T("t", "title"), grid="chevron", elements=[*boxes, Table(rows=rows)])
    placed, _ = lay(s, "jp-business")
    chev = next(p for p in placed if isinstance(p.element, Shape) and p.element.shape == "chevron")
    (tp,) = of(placed, Table)
    assert chev.style.font_size * chev.font_scale >= tp.style.font_size * tp.font_scale - 0.01
    assert chev.h <= 1.3 * 914400 + 1


def test_org_tree_boxes_hug_their_text():
    from slidemark.ir import Link

    kids = [box(n, "x") for n in "abcd"]
    s = Slide(
        title=T("t", "title"),
        grid=None,
        elements=kids,
        links=[Link(src=0, dst=1), Link(src=1, dst=2), Link(src=1, dst=3)],
    )
    placed, _ = lay(s, "jp-business")
    cs = cards(placed)
    assert len(cs) == 4
    assert max(c.h for c in cs) < 0.2 * H  # one heading + one line, not 30% of the body
    ys = sorted({c.y for c in cs})
    assert len(ys) == 3
    gap = ys[1] - (ys[0] + cs[0].h)
    assert 0.3 * 914400 <= gap <= 0.55 * 914400


def test_numeric_columns_stay_close_to_content_width():
    from slidemark.layout.tables import column_widths, table_grid

    rows = [[cell("Region"), cell("Revenue"), cell("Growth")]] + [
        [cell("APAC region long"), cell("8.2M"), cell("+48%")] for _ in range(3)
    ]
    t = Table(rows=rows)
    _, nc, anchors = table_grid(t)
    w = column_widths(t, nc, anchors, 12000000, 11.0)
    assert sum(w) == 12000000
    assert w[0] > w[1] and w[0] > w[2]  # the text column takes the larger share of the spare width
    assert w[1] >= 0.9 * 914400 * 0.5  # "Revenue" never wraps mid-word
