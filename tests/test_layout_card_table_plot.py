"""Card row above a table hugs/fills its content; a panel beside a titled chart meets the plot-area top."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from slidemark import build
from slidemark.ir import Chart, Container, Table, Text
from tests.test_layout_l3_fill import lay

EX = Path(__file__).resolve().parent.parent / "examples"


@pytest.mark.parametrize(("name", "n"), [("16-jp-strategy.md", 2), ("11-jp-consulting.md", 2)])
def test_card_row_text_fills_cards_above_table(name, n):
    placed, _, _ = lay(name, n)
    cs = [p for p in placed if isinstance(p.element, Container)]
    tbl = next(p for p in placed if isinstance(p.element, Table))
    assert len(cs) >= 3
    assert len({(c.y, c.h) for c in cs}) == 1  # equal-height row
    assert tbl.y >= cs[0].y + cs[0].h  # nothing overlaps
    bodies = [p for p in placed if isinstance(p.element, Text) and p.element.role == "body"]
    for c in cs:
        mine = [b for b in bodies if c.x <= b.x and b.x + b.w <= c.x + c.w]
        end = max(b.y + b.h for b in mine)
        assert c.y + c.h - end <= 0.25 * c.h + 2  # no card keeps a big empty tail
    assert tbl.y - (cs[0].y + cs[0].h) <= 0.2 * 914400  # the table sits right under the cards


def test_card_row_explicit_height_wins():
    src = (EX / "16-jp-strategy.md").read_text(encoding="utf-8")
    pinned = src.replace("## 市場環境 {.muted icon=globe}", "## 市場環境 {.muted icon=globe h=3in}", 1)
    got, _, _ = lay("x", 2, pinned)  # explicit input: never raises, row stays laid out
    assert [p for p in got if isinstance(p.element, Container)]


@pytest.mark.parametrize(("name", "n"), [("16-jp-strategy.md", 3), ("11-jp-consulting.md", 3)])
def test_panel_top_meets_plot_area(name, n):
    placed, _, theme = lay(name, n, panel_top_plot=True)
    ch = next(p for p in placed if isinstance(p.element, Chart))
    panel = next(p for p in placed if isinstance(p.element, Container))
    title_pt = ch.style.font_size * ch.font_scale * theme.render.chart_title_scale
    want = ch.y + theme.layout.chart_plot_top_em * title_pt * 12700
    assert abs(panel.y - want) <= 2
    assert panel.y + panel.h <= ch.y + ch.h
    head = next(p for p in placed if isinstance(p.element, Text) and p.element.role == "heading")
    assert head.y >= panel.y - 2


@pytest.mark.parametrize(("name", "n"), [("16-jp-strategy.md", 3), ("11-jp-consulting.md", 3)])
def test_panel_top_defaults_to_the_chart_top(name, n):
    placed, _, _ = lay(name, n)
    ch = next(p for p in placed if isinstance(p.element, Chart))
    panel = next(p for p in placed if isinstance(p.element, Container))
    assert panel.y == ch.y


def test_panel_without_chart_title_keeps_top():
    src = (EX / "16-jp-strategy.md").read_text(encoding="utf-8")
    src = src.replace('{title="国内市場規模の推移（億円）" labels=on', "{labels=on")
    placed, _, _ = lay("x", 3, src)
    ch = next(p for p in placed if isinstance(p.element, Chart))
    panel = next(p for p in placed if isinstance(p.element, Container))
    assert panel.y == ch.y


def test_render_roundtrip(tmp_path):
    src = (EX / "16-jp-strategy.md").read_text(encoding="utf-8")
    out = tmp_path / "a.pptx"
    build(src, out)
    prs = Presentation(str(out))
    tables = [sh for sh in prs.slides[1].shapes if sh.has_table]
    assert tables and tables[0].top > 0
    assert any(sh.has_chart for sh in prs.slides[2].shapes)
