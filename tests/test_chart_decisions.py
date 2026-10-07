"""DL2 lane E: the chart decisions a python-pptx agent makes that a fence must be able to state
(`labels.bold` / `labels.color`, `size=16,14`, `legend.size`, `overlap`, `step`, `totals`, `labels=...+name`,
`slice.line`, `render.chart_pie_line`). Each form: parser, render (the .pptx is reopened and its chart XML
asserted), a bad value (a diagnostic with a hint, never an exception) and the importer round trip."""

from __future__ import annotations

import re

import pytest
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import build, parse
from slidemark.importer import import_pptx

HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#222B36 primary=#142B4D secondary=#1F5FA8 accent=#E09F1F teal=#2A9D8F "
    "muted=#59626E surface=#EEF2F7 border=#C9D2DE\n"
)
COL = ",Q1,Q2,Q3\n売上,1280,900,1100\n利益,300,280,270\n"
PIE = ",北米,欧州,日本\n売上,45,30,25\n"
STACK = ",x,y\na,1,2\nb,3,4\n"


def fence(kind: str, attrs: str, body: str) -> str:
    return f"\n# s\n```{kind} {{{attrs}}}\n{body}```\n"


def make(tmp_path, md: str, name: str = "d"):
    deck = build(HEAD + md, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def chart_of(prs, i=0):
    return next(s for s in prs.slides[i].shapes if s.has_chart).chart


def xml(prs, i=0) -> str:
    return etree.tostring(chart_of(prs, i)._chartSpace).decode()


def opts(md: str, i: int = 0) -> dict:
    d = parse(HEAD + md)
    return next(e for e in d.slides[i].elements if e.type == "chart").options


def bad(md: str) -> list:
    d = parse(HEAD + md)
    return [x for x in d.diagnostics if x.rule == "bad-chart-option"]


# --------------------------------------------------------------------------- parser


def test_dotted_attribute_keys_parse():
    o = opts(fence("column", "labels=outside labels.bold=on labels.color=primary legend.size=14", COL))
    assert o["label_bold"] is True and o["label_color"] == "primary" and o["legend_size"] == 14
    assert o["labels"] == "on" and o["label_pos"] == "outside"


def test_size_takes_one_or_two_values():
    assert opts(fence("column", "size=14", COL))["size"] == 14
    o = opts(fence("column", "size=16,14", COL))
    assert o["size"] == 16 and o["label_size"] == 16 and o["tick_size"] == 14
    assert "label_size" not in opts(fence("column", "size=14", COL))


def test_overlap_step_totals_and_slice_line_reach_the_ir():
    o = opts(fence("column", "overlap=-5 step=200", COL))
    assert o["overlap"] == -5 and o["step"] == 200
    assert opts(fence("stacked-column", "labels=on totals=off", STACK))["totals"] == "off"
    assert opts(fence("stacked-column", "labels=on totals=on", STACK))["totals"] == "on"
    o = opts(fence("pie", "labels=outside+name slice.line=bg,2.5", PIE))
    assert o["label_name"] is True and o["label_pos"] == "outside" and o["labels"] == "on"
    assert o["slice_line"] == "bg" and o["slice_line_w"] == 2.5
    assert opts(fence("pie", "labels=percent+name slice.line=none", PIE))["labels"] == "percent"
    assert opts(fence("pie", "labels=value+name", PIE))["labels"] == "value"


def test_totals_is_accepted_without_a_warning():
    d = parse(HEAD + fence("stacked-bar", "labels=on totals=off", STACK))
    assert not [x for x in d.diagnostics if x.rule in ("bad-chart-option",)]


@pytest.mark.parametrize(
    "kind,attrs,body",
    [
        ("column", "labels=on labels.bold=maybe", COL),
        ("column", "labels=on labels.color=nope", COL),
        ("column", "size=16,14,12", COL),
        ("column", "size=16,2", COL),
        ("column", "legend.size=900", COL),
        ("column", "overlap=300", COL),
        ("column", "overlap=abc", COL),
        ("stacked-column", "overlap=-5", STACK),
        ("column", "step=0", COL),
        ("column", "step=-10", COL),
        ("pie", "step=5", PIE),
        ("column", "totals=off", COL),
        ("stacked-column", "totals=maybe", STACK),
        ("column", "labels=on+name", COL),
        ("pie", "labels=on+size", PIE),
        ("column", "slice.line=bg", COL),
        ("pie", "slice.line=nope", PIE),
        ("pie", "slice.line=bg,99", PIE),
        ("column", "labels.bold=on", COL),  # no labels shown: a silent no-op otherwise
        ("pie", "labels.color=primary", PIE),
        ("pie", "labels.bld=on", PIE),  # did-you-mean
    ],
)
def test_bad_values_are_diagnostics_with_a_hint(tmp_path, kind, attrs, body):
    md = fence(kind, attrs, body)
    msgs = bad(md)
    assert msgs and all(m.hint for m in msgs), attrs
    build(HEAD + md, tmp_path / "x.pptx")  # never raises


# --------------------------------------------------------------------------- render: labels


def test_labels_bold_and_color_reach_every_data_label(tmp_path):
    md = fence("column", "labels=outside labels.bold=on labels.color=primary fmt=#,##0", COL)
    _, prs = make(tmp_path, md)
    x = xml(prs)
    dl = re.search(r"<c:dLbls>.*?</c:dLbls>", x, re.S).group(0)
    assert 'b="1"' in dl and "142B4D" in dl


def test_labels_bold_off_unbolds_a_pie(tmp_path):
    _, on = make(tmp_path, fence("pie", "labels=on", PIE), "on")
    _, off = make(tmp_path, fence("pie", "labels=on labels.bold=off", PIE), "off")
    assert 'b="1"' in xml(on) and 'b="1"' not in xml(off) and 'b="0"' in xml(off)


def test_labels_bold_on_stacked_segments_and_the_total_carrier(tmp_path):
    _, prs = make(tmp_path, fence("stacked-column", "labels=on labels.bold=on labels.color=#FFFFFF", STACK))
    x = xml(prs)
    blocks = re.findall(r"<c:dLbls>.*?</c:dLbls>", x, re.S)
    assert blocks and all('b="1"' in b and "FFFFFF" in b for b in blocks)
    carrier = re.findall(r"<c:dLbl>.*?</c:dLbl>", x, re.S)  # the stack totals carry their own text
    assert carrier and all('b="1"' in c and "FFFFFF" in c for c in carrier)


def test_labels_bold_on_a_waterfall(tmp_path):
    md = fence("waterfall", "labels=on labels.bold=off", ",a,b,c\nx,10,-4,=\n")
    _, prs = make(tmp_path, md)
    x = xml(prs)
    runs = re.findall(r"<a:rPr[^>]*>", x)
    assert runs and all('b="0"' in r for r in runs if "sz=" in r)


# --------------------------------------------------------------------------- render: sizes


def test_size_pair_sets_labels_and_value_axis_numbers(tmp_path):
    md = fence("column", "labels=outside size=16,14 axis=on", COL)
    _, prs = make(tmp_path, md)
    ch = chart_of(prs)
    x = etree.tostring(ch._chartSpace).decode()
    assert 'sz="1600"' in re.search(r"<c:dLbls>.*?</c:dLbls>", x, re.S).group(0)
    assert 'sz="1400"' in re.search(r"<c:valAx>.*?</c:valAx>", x, re.S).group(0)
    assert 'sz="1600"' in re.search(r"<c:txPr>.*?</c:txPr>", x, re.S).group(0) or ch.font.size.pt == 16


def test_size_pair_keeps_the_value_axis_on_a_sparse_labelled_chart(tmp_path):
    _, prs = make(tmp_path, fence("column", "labels=outside size=16,14", COL))
    assert chart_of(prs).value_axis.visible  # asking for the axis numbers' size asks for the axis


def test_one_size_keeps_its_old_meaning(tmp_path):
    _, prs = make(tmp_path, fence("column", "labels=outside size=14 axis=on", COL))
    ch = chart_of(prs)
    assert ch.font.size.pt == 14
    assert "<c:txPr>" not in re.search(r"<c:valAx>.*?</c:valAx>", ch._chartSpace.xml, re.S).group(0)


def test_pie_label_size_is_exact_with_a_pair(tmp_path):
    _, prs = make(tmp_path, fence("pie", "labels=on size=15,12", PIE))
    assert 'sz="1500"' in xml(prs) and 'sz="2100"' not in xml(prs)  # not 1.4 x


def test_legend_size_is_exact(tmp_path):
    _, prs = make(tmp_path, fence("column", "legend=top legend.size=14", COL))
    leg = re.search(r"<c:legend>.*?</c:legend>", xml(prs), re.S).group(0)
    assert 'sz="1400"' in leg and '<c:legendPos val="t"/>' in leg


# --------------------------------------------------------------------------- render: overlap, step


def test_overlap_and_step_reach_the_axis_xml(tmp_path):
    md = fence("column", "labels=outside overlap=-5 step=200 min=0 max=1400", COL)
    _, prs = make(tmp_path, md)
    x = xml(prs)
    assert '<c:overlap val="-5"/>' in x
    assert re.search(r'<c:majorUnit val="200(\.0)?"/>', x)
    assert chart_of(prs).value_axis.visible


def test_step_rounds_the_automatic_max_onto_a_gridline(tmp_path):
    _, prs = make(tmp_path, fence("column", "step=300", COL))
    va = chart_of(prs).value_axis
    assert va.major_unit == 300 and (va.maximum_scale / 300) == round(va.maximum_scale / 300)


def test_step_that_would_draw_hundreds_of_gridlines_is_ignored_with_a_warning(tmp_path):
    deck, prs = make(tmp_path, fence("column", "step=1 min=0 max=5000", COL))
    assert chart_of(prs).value_axis.major_unit != 1
    assert [d for d in deck.diagnostics if d.rule == "bad-chart-option" and d.hint]


def test_overlap_default_is_untouched(tmp_path):
    _, prs = make(tmp_path, fence("column", "labels=on", COL))
    assert "<c:overlap" not in xml(prs)


# --------------------------------------------------------------------------- render: totals


def _series_names(prs, i=0):
    return [s.name for s in chart_of(prs, i).plots[0].series]


def test_totals_off_draws_no_totals_and_on_forces_them(tmp_path):
    _, off = make(tmp_path, fence("stacked-column", "labels=on totals=off", STACK), "off")
    _, dflt = make(tmp_path, fence("stacked-column", "labels=on", STACK), "dflt")
    _, forced = make(tmp_path, fence("stacked-column", "totals=on", STACK), "forced")
    assert len(_series_names(off)) == 2
    assert len(_series_names(dflt)) == 3  # the hidden carrier of the stack totals
    assert len(_series_names(forced)) == 3  # forced without segment labels


def test_totals_off_on_a_stacked_bar(tmp_path):
    _, prs = make(tmp_path, fence("stacked-bar", "labels=on totals=off", STACK))
    assert len(_series_names(prs)) == 2


# --------------------------------------------------------------------------- render: pie


def test_pie_name_and_value_label(tmp_path):
    _, prs = make(tmp_path, fence("pie", "labels=outside+name", PIE))
    x = xml(prs)
    pts = re.findall(r"<c:dLbl>.*?</c:dLbl>", x, re.S)
    assert len(pts) == 3 and all('<c:showCatName val="1"/>' in p and "<c:separator>" in p for p in pts)
    blk = re.search(r"<c:dLbls>(?:(?!</c:dLbls>).)*</c:dLbls>", x, re.S).group(0)
    assert '<c:showCatName val="1"/>' in blk and "<c:separator>\n</c:separator>" in blk
    assert 'val="outEnd"' in x
    # schema order: showBubbleSize, separator, showLeaderLines
    assert re.search(r'showBubbleSize val="0"/><c:separator>\n</c:separator>', x)


def test_pie_percent_and_name_on_a_doughnut(tmp_path):
    _, prs = make(tmp_path, fence("doughnut", "labels=percent+name", PIE))
    x = xml(prs)
    assert '<c:showCatName val="1"/>' in x and '<c:showPercent val="1"/>' in x


def test_pie_value_and_name_keep_raw_numbers(tmp_path):
    _, prs = make(tmp_path, fence("pie", "labels=value+name", PIE))
    x = xml(prs)
    assert (
        '<c:showVal val="1"/>' in x
        and '<c:showCatName val="1"/>' in x
        and '<c:showPercent val="1"/>' not in x
    )


def test_plain_pie_labels_have_no_category_name(tmp_path):
    _, prs = make(tmp_path, fence("pie", "labels=on", PIE))
    assert '<c:showCatName val="1"/>' not in xml(prs) and "<c:separator>" not in xml(prs)


def _wedge_lines(prs):
    out = []
    for dpt in chart_of(prs).plots[0]._element.iter(qn("c:dPt")):
        ln = dpt.find(qn("c:spPr")).find(qn("a:ln"))
        out.append(ln)
    return out


def test_slice_line_colour_and_width(tmp_path):
    _, prs = make(tmp_path, fence("pie", "slice.line=#112233,2.5", PIE))
    lns = _wedge_lines(prs)
    assert len(lns) == 3
    assert all(
        ln.get("w") == "31750" and ln.find(".//" + qn("a:srgbClr")).get("val") == "112233" for ln in lns
    )


def test_slice_line_named_colour_and_none(tmp_path):
    _, a = make(tmp_path, fence("doughnut", "slice.line=accent", PIE), "a")
    assert all(ln.find(".//" + qn("a:srgbClr")).get("val") == "E09F1F" for ln in _wedge_lines(a))
    _, n = make(tmp_path, fence("pie", "slice.line=none", PIE), "n")
    assert all(ln.find(qn("a:noFill")) is not None for ln in _wedge_lines(n))


def test_pie_default_outline_is_the_page_colour(tmp_path):
    _, prs = make(tmp_path, fence("pie", "labels=on", PIE))
    lns = _wedge_lines(prs)
    assert all(ln.find(".//" + qn("a:srgbClr")).get("val") == "FFFFFF" and ln.get("w") is None for ln in lns)


def test_chart_pie_line_tokens(tmp_path):
    md = "style: render.chart_pie_line=accent render.chart_pie_line_width=3\n" + fence(
        "pie", "labels=on", PIE
    )
    _, prs = make(tmp_path, md)
    lns = _wedge_lines(prs)
    assert all(
        ln.get("w") == "38100" and ln.find(".//" + qn("a:srgbClr")).get("val") == "E09F1F" for ln in lns
    )
    # the fence option wins over the token
    _, prs = make(tmp_path, md.replace("{labels=on}", "{slice.line=#000000,1}"), "w")
    assert all(ln.find(".//" + qn("a:srgbClr")).get("val") == "000000" for ln in _wedge_lines(prs))


def test_chart_pie_line_token_bad_value_is_a_diagnostic(tmp_path):
    d = parse(HEAD + "style: render.chart_pie_line=notacolor!!\n" + fence("pie", "labels=on", PIE))
    assert [x for x in d.diagnostics if x.rule == "bad-token"]


# --------------------------------------------------------------------------- lint reads the override


def test_labels_color_is_judged_by_the_contrast_lint(tmp_path):
    md = fence("column", "labels=outside labels.color=#FAFAFA", COL)
    deck, _ = make(tmp_path, md)
    assert [d for d in deck.diagnostics if d.rule == "contrast" and "data labels" in d.message]


# --------------------------------------------------------------------------- importer round trip


def test_round_trip_keeps_the_decisions(tmp_path):
    md = (
        fence("column", "labels=outside labels.bold=on overlap=-5 step=300 min=0 max=1500", COL)
        + fence("stacked-column", "labels=on totals=off", STACK)
        + fence("pie", "labels=percent+name", PIE)
    )
    make(tmp_path, md, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "labels.bold=on" in text and "overlap=-5" in text and "step=300" in text
    assert "totals=off" in text and "labels=percent+name" in text
    # and the imported deck builds to the same decisions
    deck = build(text, tmp_path / "rt2.pptx")
    prs = Presentation(str(tmp_path / "rt2.pptx"))
    assert not [d for d in deck.diagnostics if d.rule == "bad-chart-option"]
    x0 = xml(prs, 0)
    assert '<c:overlap val="-5"/>' in x0 and re.search(r'<c:majorUnit val="300', x0)
    assert len(_series_names(prs, 1)) == 2
    assert '<c:showCatName val="1"/>' in xml(prs, 2)


def test_round_trip_adds_no_decisions_to_a_plain_chart(tmp_path):
    make(tmp_path, fence("column", "labels=outside", COL), "plain")
    text, _ = import_pptx(tmp_path / "plain.pptx")
    for k in ("labels.bold", "overlap", "step", "+name"):
        assert k not in text
