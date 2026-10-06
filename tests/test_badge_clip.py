"""CJK badge text is not synthetic-bold (LibreOffice smears and clips it); Latin badges stay bold."""

from __future__ import annotations

from pptx import Presentation

from slidemark import build

MD = "# T\n\n[首位]{.badge .success} and [OK]{.badge}\n"


def _runs(tmp_path):
    out = tmp_path / "b.pptx"
    build(MD, out)
    slide = Presentation(str(out)).slides[0]
    runs = [r for sh in slide.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    return {r.text.strip("　 "): r for r in runs}


def test_cjk_badge_not_bold_but_padded_and_highlighted(tmp_path):
    cjk = _runs(tmp_path)["首位"]
    assert cjk.font.bold is None
    assert cjk.text.startswith("　") and cjk.text.endswith("　")
    assert cjk._r.find(".//{*}highlight") is not None


def test_latin_badge_stays_bold(tmp_path):
    assert _runs(tmp_path)["OK"].font.bold is True
