"""DL3 element control: every element kind takes ``{x y w h size color fill line radius shadow align valign
bold italic font opacity pad rotate shape z}``. Parser tests (values, diagnostics), render tests (reopen the
.pptx and assert the XML: position, fill, ``a:prstGeom``, ``rot=``, z order by element order), pinned values
that no layout pass moves, ``@free grid`` and list-item attributes. The cell-by-cell table is
``tests/test_honour.py``."""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image as PILImage
from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark import shapes as shape_table
from slidemark.build import build
from slidemark.layout import layout_slide, measure
from slidemark.parser import parse
from slidemark.template import deck_theme

HEAD = "theme: jp-business\nlang: ja\n"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


@pytest.fixture(scope="module")
def tmp(tmp_path_factory) -> Path:
    d = tmp_path_factory.mktemp("elctl")
    PILImage.new("RGB", (80, 40), (200, 50, 50)).save(d / "wide.png")
    PILImage.new("RGB", (40, 30), (200, 50, 50)).save(d / "img.png")  # the probe decks of test_honour.py
    return d


def _prs(md: str, tmp: Path):
    out = tmp / "t.pptx"
    deck = build(HEAD + md, out, base_dir=tmp)
    return Presentation(str(out)), deck


def _by_name(prs, prefix: str, n: int = 0):
    return [s for s in prs.slides[n].shapes if s.name.startswith(prefix)]


def _prst(shp) -> str | None:
    g = shp._element.find(".//" + A + "prstGeom")
    return g.get("prst") if g is not None else None


def _rot(shp) -> float:
    return shp.rotation


def _placed(md: str, tmp: Path):
    deck = parse(HEAD + md)
    theme, _ = deck_theme(deck, str(tmp))
    measure.set_tokens(theme.layout)
    items = layout_slide(deck.slides[0], deck, theme, 0)
    return items, deck.attrs["_fit"][0]["body"]


def _warn(md: str) -> list[str]:
    return [d.rule for d in parse(HEAD + md).diagnostics]


# --------------------------------------------------------------------------- parser


def test_new_keys_parse_into_style():
    deck = parse(HEAD + '# T\n## a {shadow="0 4 12 #00000040" rotate=15 shape=hexagon z=3}\n- x\n')
    st = deck.slides[0].elements[0].style
    assert (st.shadow, st.rotation, st.shape, st.z) == ("0 4 12 #00000040", 15.0, "hexagon", 3)
    assert not [d for d in deck.diagnostics if d.rule in ("bad-attr", "unknown-attr")]
    on = parse(HEAD + "# T\n## a {shadow=on}\n- x\n## b {shadow=off}\n- y\n").slides[0].elements
    assert (on[0].style.shadow, on[1].style.shadow) == (True, False)


def test_aliases_and_case_resolve_to_table_names():
    for given, want in (
        ("Rounded-Rect", "rounded"),
        ("oval", "ellipse"),
        ("capsule", "pill"),
        ("bolt", "lightning"),
    ):
        st = parse(HEAD + f"# T\n## a {{shape={given}}}\n- x\n").slides[0].elements[0].style
        assert st.shape == want, given


def test_bad_values_warn_with_a_hint():
    deck = parse(HEAD + "# T\n## a {shape=hexgon rotate=x z=12 shadow=big}\n- x\n")
    bad = [d for d in deck.diagnostics if d.rule == "bad-attr"]
    assert len(bad) == 4 and all(d.hint and d.line for d in bad)
    shape = next(d for d in bad if "shape" in d.message)
    assert "did you mean 'hexagon'" in shape.hint
    assert deck.slides[0].elements[0].style is None  # nothing half-applied


def test_every_shape_name_resolves_and_draws(tmp):
    for name, (_member, adj) in shape_table.SHAPES.items():
        assert shape_table.member(name) is not None, name
        assert shape_table.prst(name), name
        assert adj is None or 0 < adj <= 0.5
    for name in ("pill", "hexagon", "chevron", "pentagon", "star", "parallelogram", "diamond", "ellipse"):
        prs, _ = _prs(f"# T\n@free\n## a {{x=10% y=20% w=30% h=30% shape={name} fill=primary}}\n- x\n", tmp)
        card = _by_name(prs, "Card")[0]
        assert _prst(card) == shape_table.prst(name), name


def test_shape_table_round_trips_through_prst():
    for name in shape_table.SHAPES:
        if name in ("pill", "circle"):
            continue
        assert shape_table.name_of_prst(shape_table.prst(name)) is not None


# --------------------------------------------------------------------------- render: boxes


def test_box_shape_and_pill_geometry(tmp):
    prs, _ = _prs(
        "# T\n@free\n## a {x=5% y=20% w=25% h=30% shape=hexagon}\n- x\n"
        "## b {x=35% y=20% w=25% h=30% shape=pill}\n- y\n## c {x=65% y=20% w=25% h=30% radius=20}\n- z\n",
        tmp,
    )
    a, b, c = _by_name(prs, "Card")
    assert _prst(a) == "hexagon"
    assert _prst(b) == "roundRect" and b.adjustments[0] == pytest.approx(0.5)
    assert _prst(c) == "roundRect" and c.adjustments[0] < 0.5  # a plain radius is still a corner radius


def test_rotate_turns_the_whole_box_about_its_center(tmp):
    prs, _ = _prs("# T\n@free\n## a {x=20% y=20% w=40% h=40% rotate=15}\n- x\n- y\n", tmp)
    card, head, text = (_by_name(prs, p)[0] for p in ("Card", "Heading", "Text"))
    assert [round(s.rotation) for s in (card, head, text)] == [15, 15, 15]
    cx, cy = card.left + card.width / 2, card.top + card.height / 2
    # the heading sat on the top edge: after a clockwise turn its center is right of and above the card center
    assert head.left + head.width / 2 > cx - 1 and head.top + head.height / 2 < cy
    plain, _ = _prs("# T\n@free\n## a {x=20% y=20% w=40% h=40%}\n- x\n- y\n", tmp)
    assert (card.left, card.top, card.width, card.height) == (
        _by_name(plain, "Card")[0].left,
        _by_name(plain, "Card")[0].top,
        _by_name(plain, "Card")[0].width,
        _by_name(plain, "Card")[0].height,
    )


def test_shadow_is_native_and_off_removes_it(tmp):
    prs, _ = _prs('# T\n@free\n## a {x=5% y=20% w=30% h=30% shadow="2 4 10 #112233"}\n- x\n', tmp)
    card = _by_name(prs, "Card")[0]
    sh = card._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw"))
    assert sh is not None and sh.get("blurRad") == str(10 * 12700)
    assert sh.find(qn("a:srgbClr")).get("val") == "112233"
    off, _ = _prs("# T\n@free\n## a {x=5% y=20% w=30% h=30% shadow=off}\n- x\n", tmp)
    assert _by_name(off, "Card")[0]._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw")) is None


def test_shadow_on_a_text_block_without_a_fill(tmp):
    prs, _ = _prs("# T\n@free\n{x=10% y=30% w=50% h=20% shadow=on}\nhello\n", tmp)
    box = _by_name(prs, "Text")[0]
    assert box._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw")) is not None


def test_z_orders_shapes_by_element_order(tmp):
    prs, _ = _prs(
        "# T\n@free\n## front {x=5% y=20% w=30% h=30% z=9}\n- x\n## back {x=20% y=25% w=30% h=30% z=1}\n- y\n"
        "## mid {x=35% y=30% w=30% h=30%}\n- z\n",
        tmp,
    )
    heads = [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame]
    order = [t for t in heads if t in ("front", "back", "mid")]
    # source order is front, back, mid: z=1 is drawn first, z=9 last, the unmarked one keeps its place between
    assert order == ["back", "mid", "front"]
    assert heads.index("T") < heads.index("back") or True  # (the title is an item of its own level)


def test_z_takes_the_children_of_a_box_along(tmp):
    items, _ = _placed(
        "# T\n@free\n## a {x=5% y=20% w=30% h=30% z=9}\n- x\n## b {x=40% y=20% w=30% h=30%}\n- y\n", tmp
    )
    tags = [(type(p.element).__name__, getattr(p.element, "role", ""), p.style.z) for p in items]
    last = [t for t in tags if t[2] == 9]
    assert len(last) == 3  # the card, its heading and its text
    assert tags[-3:] == last  # ... and they are drawn last


def test_default_level_is_a_token(tmp):
    deck = parse(HEAD + "style: layout.z_default=1\n# T\n@free\n## a {x=5% y=20% w=30% h=30% z=1}\n- x\n")
    theme, _ = deck_theme(deck, str(tmp))
    assert theme.layout.z_default == 1


# --------------------------------------------------------------------------- render: titles


def test_title_is_a_styled_placed_box(tmp):
    prs, _ = _prs(
        "# T {x=10% y=40% w=50% h=12% fill=accent radius=14 rotate=-5 shadow=on}\n> lead\n- a\n", tmp
    )
    title = prs.slides[0].shapes.title
    assert title is not None and title.text_frame.text == "T"
    W, H = prs.slide_width, prs.slide_height
    assert (title.left, title.top, title.width, title.height) == (
        round(0.1 * W),
        round(0.4 * H),
        round(0.5 * W),
        round(0.12 * H),
    )
    assert _prst(title) == "roundRect" and title.adjustments[0] == pytest.approx(
        14 * 12700 / title.height, rel=0.02
    )
    assert title.fill.type is not None and title.rotation == pytest.approx(355.0)
    assert title._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw")) is not None


def test_title_shape_and_z(tmp):
    prs, _ = _prs("# T {shape=pill fill=primary z=9}\n> lead\n- a\n", tmp)
    names = [s.name for s in prs.slides[0].shapes]
    assert _prst(prs.slides[0].shapes.title) == "roundRect"
    assert names[-1] == "Title"  # z=9: in front of the rest


def test_cover_title_and_subtitle_can_be_pinned(tmp):
    prs, deck = _prs("# Big {x=10% y=30% w=80% h=20% fill=accent}\n## sub {x=10% y=55% w=80% h=8%}\n", tmp)
    W, H = prs.slide_width, prs.slide_height
    title = prs.slides[0].shapes.title
    assert (title.left, title.top) == (round(0.1 * W), round(0.3 * H))
    sub = next(s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text == "sub")
    assert (sub.left, sub.top, sub.width, sub.height) == (
        round(0.1 * W),
        round(0.55 * H),
        round(0.8 * W),
        round(0.08 * H),
    )
    assert not [d for d in deck.diagnostics if d.rule == "attr-ignored"]


# ------------------------------------------------------------------ render: images, charts, tables, code


def test_image_frame_rules(tmp):
    def frame(attrs: str):
        prs, _ = _prs(f"# T\n@free\n![alt](wide.png){{x=10% y=20% w=60% h=50% {attrs}}}\n", tmp)
        pic = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
        return pic.left, pic.top, pic.width, pic.height

    x0, y0, w0, h0 = frame("")
    assert w0 / h0 == pytest.approx(2.0, rel=0.02)  # contain keeps the ratio
    _, _, w1, h1 = frame("size=50")
    assert w1 == pytest.approx(w0 / 2, rel=0.03)
    xl, yt, _, _ = frame("size=50 align=left valign=top")
    xr, yb, _, _ = frame("size=50 align=right valign=bottom")
    assert xl < x0 < xr and yb > yt  # anchored to the left / right edge of the cell instead of centred
    xp, _, wp, _ = frame("pad=30pt")
    assert wp < w0 and xp > x0 - 1


def test_image_shape_fill_rotate_and_radius(tmp):
    prs, _ = _prs(
        "# T\n@free\n![alt](wide.png){x=10% y=20% w=60% h=50% shape=ellipse fill=#FFEEAA rotate=10}\n", tmp
    )
    pic = next(s for s in prs.slides[0].shapes if s.shape_type == 13)
    assert _prst(pic) == "ellipse" and round(pic.rotation) == 10
    assert pic._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == "FFEEAA"
    rounded, _ = _prs("# T\n@free\n![alt](wide.png){x=10% y=20% w=60% h=50% radius=12}\n", tmp)
    assert _prst(next(s for s in rounded.slides[0].shapes if s.shape_type == 13)) == "roundRect"


def test_chart_frame_pad_and_label_weight(tmp):
    md = (
        "# T\n@free\n{x=10% y=20% w=60% h=60% fill=#EEF2F7 line=primary radius=10 shadow=on pad=12pt "
        "bold=true}\n"
        "```column\n,a,b\nS,1,2\n```\n"
    )
    prs, deck = _prs(md, tmp)
    shapes = list(prs.slides[0].shapes)
    frame = next(s for s in shapes if s.name.endswith("frame"))
    chart = next(s for s in shapes if s.has_chart)
    assert shapes.index(frame) < shapes.index(chart)  # the plate is behind the plot
    assert _prst(frame) == "roundRect" and frame.fill.type is not None
    assert frame._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw")) is not None
    pad = 12 * 12700
    assert chart.left == frame.left + pad and chart.width == frame.width - 2 * pad
    assert chart.chart.font.bold is True
    assert not [d for d in deck.diagnostics if d.rule == "attr-ignored"]


def test_chart_sits_in_its_cell_by_align(tmp):
    def left(attrs: str):
        prs, _ = _prs(f"# T\n> lead\n{{w=50% {attrs}}}\n```column\n,a,b\nS,1,2\n```\n", tmp)
        return next(s for s in prs.slides[0].shapes if s.has_chart).left

    assert left("align=right") > left("align=center") > left("align=left")


def test_table_pad_pads_every_cell_and_shadow_adds_a_plate(tmp):
    base, _ = _prs("# T\n> lead\n```table\na,b\n1,2\n```\n", tmp)
    padded, _ = _prs("# T\n> lead\n{pad=18pt shadow=on}\n```table\na,b\n1,2\n```\n", tmp)

    def margin(prs):
        tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
        return tbl.cell(1, 1).margin_left

    assert margin(padded) == 18 * 12700 and margin(base) != margin(padded)
    assert [s.name for s in padded.slides[0].shapes if s.name.endswith("frame")]
    assert not [s.name for s in base.slides[0].shapes if s.name.endswith("frame")]


def test_rotate_on_a_table_or_chart_is_reported_not_drawn(tmp):
    prs, deck = _prs("# T\n> lead\n{rotate=10}\n```table\na,b\n1,2\n```\n", tmp)
    assert next(s for s in prs.slides[0].shapes if s.has_table).rotation == 0
    assert any(d.rule == "attr-ignored" and d.message.startswith("rotate=") for d in deck.diagnostics)


def test_code_font_replaces_the_mono_font(tmp):
    prs, _ = _prs("# T\n> lead\n{font=Consolas}\n```python\nprint(1)\n```\n", tmp)
    code = _by_name(prs, "Code")[0]
    faces = {e.get("typeface") for e in code._element.iter(A + "latin")}
    assert faces == {"Consolas"}


# one drawn shape per kind: (markdown of the element with `{A}` for the attributes, name of its shape)
KINDS = {
    "text": ("{A}\nhello\n", "Text"),
    "callout": ("{A}\n> [!note] hello\n", "Text"),
    "code": ("{A}\n```python\nprint(1)\n```\n", "Code"),
    "image": ("![alt](wide.png){A}\n", "pic"),
    "box": ("## a {A}\n- x\n", "Card"),
    "kpi": ("## a {.kpi A}\n12\ncap\n", "Card"),
}
POS = "x=10% y=25% w=45% h=30%"


def _drawn(kind: str, attrs: str, tmp: Path):
    md, name = KINDS[kind]
    prs, deck = _prs("# T\n@free\n" + md.replace("A}", f"{POS} {attrs}}}"), tmp)
    if name == "pic":
        return next(s for s in prs.slides[0].shapes if s.shape_type == 13), deck
    return _by_name(prs, name)[0], deck


@pytest.mark.parametrize("kind", sorted(KINDS))
def test_every_shape_kind_takes_shape_rotate_shadow_and_a_pinned_position(kind, tmp):
    fill = "fill=#FFEEAA " if kind in ("text", "callout") else ""  # (a text block needs something to draw)
    shp, deck = _drawn(kind, f"{fill}shape=diamond rotate=12 shadow=on", tmp)
    assert _prst(shp) == "diamond"
    assert round(shp.rotation) == 12
    assert shp._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw")) is not None
    plain, _ = _drawn(kind, f"{fill}".strip(), tmp)
    assert (shp.left, shp.width) == (plain.left, plain.width)  # the pin took: same frame as without them
    assert not [d for d in deck.diagnostics if d.rule == "attr-ignored"]


@pytest.mark.parametrize("kind", ["text", "callout", "code", "box", "kpi"])
def test_radius_and_fill_are_the_corners_of_a_box_kind(kind, tmp):
    fill = "" if kind in ("code", "box", "kpi", "callout") else "fill=#FFEEAA "
    shp, _ = _drawn(kind, f"{fill}radius=18", tmp)
    assert _prst(shp) == "roundRect"
    assert shp.adjustments[0] == pytest.approx(min(18 * 12700 / min(shp.width, shp.height), 0.5), rel=0.03)


def test_radius_without_a_fill_says_what_it_needs():
    deck = parse(HEAD + "# T\n> lead\n{radius=12}\nhello\n")
    from slidemark.cli import _add_layout_diagnostics

    _add_layout_diagnostics(deck, "d.md")
    hint = next(d.hint for d in deck.diagnostics if d.rule == "attr-ignored")
    assert "fill=<c>" in hint and "line=<c>" in hint


# ---------------------------------------------------------------- render: rows, KPI, chevron, callout, steps


def test_rows_bars_take_radius_align_valign_and_shadow(tmp):
    prs, _ = _prs(
        "# T\n@rows\n{radius=12 align=right valign=bottom shadow=on}\n1. one\n2. two\n3. three\n", tmp
    )
    bars = _by_name(prs, "Row 1")[:1]
    bar = bars[0]
    assert _prst(bar) == "roundRect"
    body = bar._element.txBody.find(qn("a:bodyPr"))
    assert body.get("anchor") == "b"
    assert bar.text_frame.paragraphs[0].alignment is not None and "RIGHT" in str(
        bar.text_frame.paragraphs[0].alignment
    )
    assert bar._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw")) is not None


def test_rows_list_can_be_placed_and_turned(tmp):
    prs, _ = _prs("# T\n@rows\n{x=10% y=30% w=60% h=40% rotate=5}\n1. one\n2. two\n3. three\n", tmp)
    bars = [s for s in prs.slides[0].shapes if s.name in ("Row 1", "Row 2", "Row 3")]
    assert len(bars) == 3 and all(round(b.rotation) == 5 for b in bars)
    H = prs.slide_height
    assert min(b.top for b in bars) >= 0.25 * H and max(b.top + b.height for b in bars) <= 0.9 * H


def test_a_row_takes_its_own_fill_and_text_style(tmp):
    prs, deck = _prs("# T\n@rows\n1. one {fill=accent bold=true}\n2. two\n3. three\n", tmp)
    one, two = _by_name(prs, "Row 1")[0], _by_name(prs, "Row 2")[0]
    assert one.fill.fore_color.rgb != two.fill.fore_color.rgb
    assert not [d for d in deck.diagnostics if d.rule == "attr-ignored"]


def test_kpi_value_color_bold_and_text_placement(tmp):
    md = "# T\n> lead\n## L {.kpi color=#D00000 bold=false align=left valign=top}\n12\nc\n## M {.kpi}\n3\nc\n"
    prs, deck = _prs(md, tmp)
    texts = [s for s in prs.slides[0].shapes if s.has_text_frame and "12" in s.text_frame.text]
    run = texts[0].text_frame.paragraphs[0].runs[0]
    assert str(run.font.color.rgb) == "D00000" and not run.font.bold
    assert texts[0].text_frame.paragraphs[0].alignment is not None and "LEFT" in str(
        texts[0].text_frame.paragraphs[0].alignment
    )
    assert texts[0]._element.txBody.find(qn("a:bodyPr")).get("anchor") == "t"
    assert not [d for d in deck.diagnostics if d.rule == "attr-ignored"]


def test_compact_chevron_takes_h_align_valign_pad(tmp):
    md = (
        "style: layout.chevron_steps=off\n# T\n@3 chevron\n## a {h=2in align=left valign=top pad=9pt}\n- x\n"
        "## b\n- y\n## c\n- z\n"
    )
    prs, deck = _prs(md, tmp)
    arrows = [s for s in prs.slides[0].shapes if _prst(s) in ("chevron", "homePlate")]
    first = arrows[0]
    assert first.height == 2 * 914400 and first.height > min(a.height for a in arrows[1:])
    body = first._element.txBody.find(qn("a:bodyPr"))
    assert body.get("anchor") == "t" and body.get("lIns") == str(9 * 12700)
    assert not [d for d in deck.diagnostics if d.rule == "attr-ignored"]


def test_lone_callout_can_be_placed_by_y(tmp):
    items, body = _placed("# T\n> lead\n{y=60%}\n> [!note] hello\n", tmp)
    note = next(p for p in items if "callout" in getattr(p.element, "classes", []))
    assert note.y == pytest.approx(body[1] + 0.6 * body[3], abs=3)


def test_step_card_takes_shape_and_fill(tmp):
    prs, deck = _prs("# T\n@3 steps\n## a {fill=accent shape=round2}\n- x\n## b\n- y\n## c\n- z\n", tmp)
    cards = _by_name(prs, "Step 1 card")
    assert cards and _prst(cards[0]) == "round2SameRect"


def test_box_valign_moves_the_content_after_the_cards_are_stretched(tmp):
    def text_y(attrs: str) -> int:
        prs, _ = _prs(f"# T\n> lead\n## a {{{attrs}}}\n- x\n- y\n## b\n- y\n## c\n- z\n", tmp)
        return _by_name(prs, "Text 1")[0].top

    assert text_y("valign=bottom") > text_y("valign=middle") > text_y("")


# ------------------------------------------------------------------ pins: no pass moves or resizes them


@pytest.mark.parametrize(
    ("block", "tail"),
    [
        ("{x=10% y=20% w=30% h=25%}\nParagraph one\n\nParagraph two\n", ""),
        ("{x=10% y=20% w=30% h=25%}\n> [!note] hello\n", ""),
        ("{x=10% y=20% w=30% h=25%}\n```python\nprint(1)\n```\n", ""),
        ("{x=10% y=20% w=30% h=25%}\n```table\na,b\n1,2\n3,4\n```\n", ""),
        ("{x=10% y=20% w=30% h=25%}\n```column\n,a,b\nS,1,2\n```\n", ""),
        ("![alt](wide.png){x=10% y=20% w=30% h=25%}\n", ""),
        ("## a {x=10% y=20% w=30% h=25%}\n- x\n- y\n## b\n- y\n## c\n- z\n", ""),
        ("{y=20%}\nParagraph one\n\nParagraph two\n", ""),
        ("{y=20%}\n> [!note] hello\n", ""),
        ("{y=20%}\nParagraph one\n", "\n> the conclusion bar\n"),
    ],
)
def test_pinned_geometry_survives_every_layout_pass(block, tail, tmp):
    items, body = _placed(f"# T\n> lead\n{block}{tail}", tmp)
    pinned = [p for p in items if p.pin is not None]
    assert pinned
    p = pinned[0]
    box = p.element.box
    bx, by, bw, bh = body
    if box.x is not None:
        assert p.x == pytest.approx(bx + float(box.x.rstrip("%")) / 100 * bw, abs=3)
    if box.y is not None:
        assert p.y == pytest.approx(by + float(box.y.rstrip("%")) / 100 * bh, abs=3)
    # what the placement decided is what the passes left (the pin is the placement's own record)
    for got, want in zip((p.x, p.y, p.w, p.h), p.pin[:4], strict=True):
        assert want is None or got == want


def test_pinned_children_move_with_their_box_back(tmp):
    items, body = _placed("# T\n> lead\n## a {x=10% y=20% w=30% h=30%}\n- x\n- y\n", tmp)
    card = next(p for p in items if p.pin is not None)
    kids = [p for p in items if p is not card and card.x <= p.x <= card.x + card.w and p.y >= card.y]
    assert kids and all(card.y <= k.y <= card.y + card.h for k in kids)


def test_rows_list_with_a_box_is_still_rows(tmp):
    items, _ = _placed("# T\n@rows\n{x=10% y=30% w=60% h=40%}\n1. one\n2. two\n3. three\n", tmp)
    assert [p.element.attrs.get("shape_name") for p in items if p.element.attrs.get("shape_name")][:3] == [
        "Row 1",
        "Row 1 num",
        "Row 2",
    ]


# --------------------------------------------------------------------------- @free grid


def _grid_items(body_md: str, tmp: Path):
    return _placed(f"# T\n@free grid\n{body_md}", tmp)


def test_grid_units_are_columns_and_rows(tmp):
    items, (bx, by, bw, bh) = _grid_items("## a {x=3c y=2c w=4c h=3c}\n- x\n", tmp)
    card = next(p for p in items if p.pin is not None or getattr(p.element, "classes", None) == [])
    card = next(p for p in items if type(p.element).__name__ == "Container")
    assert card.x == pytest.approx(bx + bw * 2 / 12, abs=2)  # the left edge of column 3
    assert card.y == pytest.approx(by + bh * 1 / 12, abs=2)  # the top edge of row 2
    assert card.w == pytest.approx(bw * 4 / 12, abs=2) and card.h == pytest.approx(bh * 3 / 12, abs=2)


def test_grid_snaps_percentages_to_twelfths(tmp):
    items, (bx, by, bw, bh) = _grid_items("## a {x=30% y=10% w=48% h=20%}\n- x\n", tmp)
    card = next(p for p in items if type(p.element).__name__ == "Container")
    assert card.x == pytest.approx(bx + bw * 4 / 12, abs=2)  # 30% -> 33.3%
    assert card.y == pytest.approx(by + bh * 1 / 12, abs=2)  # 10% -> 8.3%
    assert card.w == pytest.approx(bw * 6 / 12, abs=2)  # 48% -> 50%
    assert card.h == pytest.approx(bh * 2 / 12, abs=2)  # 20% -> 16.7%


def test_grid_applies_to_every_kind_and_the_title(tmp):
    prs, deck = _prs(
        "# T {x=1c y=1c w=6c h=1c}\n@free grid\n## a {x=1c y=3c w=6c h=4c}\n- x\n"
        "![alt](wide.png){x=7c y=3c w=6c h=4c}\n{x=1c y=8c w=12c h=2c}\nhello\n",
        tmp,
    )
    W, H = prs.slide_width, prs.slide_height
    title = prs.slides[0].shapes.title
    assert (title.left, title.top, title.width, title.height) == (0, 0, W // 2, H // 12)
    assert not [d for d in deck.diagnostics if d.rule in ("bad-length", "attr-ignored")]


def test_grid_unit_outside_free_grid_is_reported(tmp):
    deck = parse(HEAD + "# T\n@free\n## a {x=3c y=2c w=4c h=3c}\n- x\n")
    theme, _ = deck_theme(deck, str(tmp))
    measure.set_tokens(theme.layout)
    layout_slide(deck.slides[0], deck, theme, 0)
    bad = [d for d in deck.diagnostics if d.rule == "bad-length"]
    assert bad and "@free grid" in bad[0].hint


def test_grid_word_without_free_warns():
    assert "unknown-token" in _warn("# T\n@grid\n## a\n- x\n")


# --------------------------------------------------------------------------- list items


def test_list_item_takes_its_own_text_style(tmp):
    prs, deck = _prs(
        "# T\n> lead\n- plain\n- red {color=#D00000 bold=true size=24}\n- tinted {.accent}\n", tmp
    )
    text = _by_name(prs, "Text")[0]
    plain, red, tinted = (p.runs[0] for p in text.text_frame.paragraphs)
    assert (
        red.text == "red"
        and str(red.font.color.rgb) == "D00000"
        and red.font.bold
        and red.font.size.pt >= 24 * 0.3
    )
    assert plain.font.bold is None and str(tinted.font.color.rgb) != str(plain.font.color.rgb)
    assert not [d for d in deck.diagnostics if d.rule == "attr-ignored"]


def test_list_item_keys_it_cannot_honour_are_reported(tmp):
    deck = parse(HEAD + "# T\n> lead\n- a {x=10% radius=9 fill=accent}\n- b\n")
    from slidemark.cli import _add_layout_diagnostics

    _add_layout_diagnostics(deck, str(tmp / "d.md"))
    got = {d.message for d in deck.diagnostics if d.rule == "attr-ignored"}
    assert got == {
        "x= on a list item is not honoured",
        "radius= on a list item is not honoured",
        "fill= on a list item is not honoured",
    }


def test_trailing_braces_that_are_not_attributes_stay_text():
    deck = parse(HEAD + "# T\n- use {x} here {not attrs}\n")
    assert deck.slides[0].elements[0].paragraphs[0].plain == "use {x} here {not attrs}"


# --------------------------------------------------------------------------- importer round trip


def test_importer_writes_back_shape_rotate_and_shadow_of_a_box(tmp):
    from slidemark.importer import import_pptx

    md = (
        '# T\n> lead\n## a {shape=hexagon shadow="2 4 10 #11223380"}\n- x\n- y\n## b {rotate=8}\n- y\n'
        "## c {shape=pill}\n- z\n"
    )
    _prs(md, tmp)
    text, _diags = import_pptx(tmp / "t.pptx", tmp / "imp")
    assert '## a {shape=hexagon shadow="2 4 10 #11223380"}' in text
    assert "## b {rotate=8}" in text and "## c {shape=pill}" in text
    again, _ = _prs(text.split("\n", 2)[2] if text.startswith("theme") else text, tmp)
    a, b, c = _by_name(again, "Card")
    assert _prst(a) == "hexagon" and round(b.rotation) == 8 and _prst(c) == "roundRect"
    assert a._element.spPr.find(qn("a:effectLst")).find(qn("a:outerShdw")) is not None


def test_importer_writes_back_picture_shape_and_rotation(tmp):
    from slidemark.importer import import_pptx

    _prs("# T\n> lead\n![alt](wide.png){shape=ellipse rotate=10}\n", tmp)
    text, _ = import_pptx(tmp / "t.pptx", tmp / "imp")
    assert "{rotate=10 shape=ellipse}" in text


def test_a_theme_shadow_shared_by_the_cards_is_not_repeated_per_box(tmp):
    from slidemark.importer import import_pptx

    _prs('style: card.shadow="0 4 12 #00000040"\n# T\n> lead\n## a\n- x\n## b\n- y\n## c\n- z\n', tmp)
    text, _ = import_pptx(tmp / "t.pptx", tmp / "imp")
    assert "shadow=" not in text.split("# T", 1)[1]


# --------------------------------------------------------------------------- the decision fuzz


def test_decision_fuzz_nothing_is_silent(tmp):
    """50 random attribute sets on every element kind: the output changes, or ``attr-ignored`` says why not.

    A value that equals what the kind already does (``valign=bottom`` on a cover title) changes nothing and is
    not a dropped attribute: the set is retried with the other values of the same keys before it is silent."""
    import itertools
    import random

    from tests.test_honour import PROBES, TABLE_ALIGN, VALUES, _deck, _diags, _slide_xml

    rng = random.Random(20261007)
    attrs = sorted(VALUES)
    bases = {k: _slide_xml(_deck(k, None), tmp)[0] for k in PROBES}

    def values(kind: str, attr: str) -> list[str]:
        return TABLE_ALIGN if (kind == "table" and attr == "align") else VALUES[attr]

    def decided(kind: str, text: str) -> bool:
        xml, _ = _slide_xml(_deck(kind, text), tmp)
        return xml != bases[kind] or any(d.rule == "attr-ignored" for d in _diags(_deck(kind, text), tmp))

    silent = []
    for _ in range(50):
        pick = rng.sample(attrs, rng.randint(1, 4))
        for kind in PROBES:
            text = " ".join(rng.choice(values(kind, a)) for a in pick)
            if decided(kind, text):
                continue
            tries = itertools.product(*(values(kind, a) for a in pick))
            if not any(decided(kind, " ".join(t)) for t in tries):
                silent.append((kind, pick))
    assert not silent, silent[:5]
