from pptx import Presentation

from slidemark.preview import pptx_to_pngs

from .helpers import needs_soffice


@needs_soffice
def test_pptx_to_pngs(tmp_path):
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.slides.add_slide(prs.slide_layouts[6])
    src = tmp_path / "d.pptx"
    prs.save(src)
    pngs = pptx_to_pngs(src, tmp_path / "png")
    assert [p.name for p in pngs] == ["slide-01.png", "slide-02.png"]
    assert all(p.stat().st_size > 0 for p in pngs)
