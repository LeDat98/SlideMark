"""Table cells are contrast-linted with the renderer's own fills and text colors."""

from __future__ import annotations

import re
from pathlib import Path

from slidemark.layout import layout_slide
from slidemark.lint import lint
from slidemark.parser import parse
from slidemark.template import deck_theme

EXAMPLES = Path(__file__).resolve().parent.parent / "examples"
TABLE = "# T\n\n| A | B |\n|---|---|\n| 1 | 2 |\n| 3 | 4 |\n| 5 | 6 |\n"
ZEBRA = TABLE.replace("\n| A", "\n{.zebra}\n| A", 1)


def _warnings(md: str) -> list:
    deck = parse(md)
    th, _ = deck_theme(deck, EXAMPLES)
    placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
    return [d for d in lint(deck, placed, th) if d.rule == "contrast" and d.message.startswith("table")]


def _kind(d) -> str:
    return d.message.split(" text")[0]


def _edits(hints: list[str]) -> str:
    """The hints' ``style: token=value`` edits as one deck header line."""
    return "style: " + " ".join(re.match(r"style: (\S+)", h).group(1) for h in hints) + "\n\n"  # type: ignore[union-attr]


def test_header_white_on_light_primary_warns_once_and_hint_fixes_it():
    md = "colors: primary=#DDEEFF\nstyle: table.header.fill=primary table.header.color=#FFFFFF\n\n" + TABLE
    got = _warnings(md)
    assert [_kind(d) for d in got] == ["table header"]  # one per table and kind, not per cell
    assert got[0].hint.startswith("style: table.header.color=")
    fixed = md.replace(
        "table.header.color=#FFFFFF", got[0].hint.split()[1].removeprefix("table.header.color=")
    )
    assert "#FFFFFF" not in fixed.splitlines()[1]
    assert _warnings(fixed) == []


def test_banded_and_body_rows_are_separate_kinds():
    md = "style: td.color=#C8C8C8\n\n" + ZEBRA
    got = _warnings(md)
    assert sorted(_kind(d) for d in got) == ["table banded rows", "table body"]
    assert all(d.hint.startswith("style: td.color=") for d in got)
    # one pasted edit per kind passes on every body fill, so either one clears both
    for d in got:
        fixed = re.sub(r"td.color=#\w+", d.hint.split()[1], md)
        assert _warnings(fixed) == [], fixed


def test_css_header_fill_and_color_are_judged():
    md = "```css\nth { background: #FFFF99; color: #FFFFFF }\n```\n\n" + TABLE
    got = _warnings(md)
    assert [_kind(d) for d in got] == ["table header"]
    assert got[0].hint.startswith("style: th.color=")
    ink = got[0].hint.split("=")[1].split()[0]
    assert _warnings(md.replace("color: #FFFFFF", f"color: {ink}")) == []


def test_run_color_gets_a_run_hint():
    md = "# T\n\n| A | B |\n|---|---|\n| [x]{color=#FFFFFF} | 2 |\n"
    got = _warnings(md)
    assert [_kind(d) for d in got] == ["table body"]
    assert "change the run's color to {color=#" in got[0].hint
    ink = re.search(r"color=(#\w+)\}", got[0].hint).group(1)  # type: ignore[union-attr]
    assert _warnings(md.replace("#FFFFFF", ink)) == []


def test_default_and_preset_tables_are_quiet():
    for head in ("", "style: preset=midnight\n\n", "style: preset=jp-business\n\n"):
        assert _warnings(head + TABLE) == []
        assert _warnings(head + ZEBRA) == []


def test_examples_have_no_table_warnings():
    for md in sorted(EXAMPLES.glob("*.md")):
        deck = parse(md.read_text(encoding="utf-8"))
        th, _ = deck_theme(deck, md.parent)
        placed = [layout_slide(s, deck, th, i) for i, s in enumerate(deck.slides)]
        bad = [str(d) for d in lint(deck, placed, th) if d.rule == "contrast" and "table" in d.message]
        assert not bad, (md.name, bad)


def test_table_resolver_matches_render(tmp_path):
    """The shared resolver gives the fills and header ink the .pptx really carries."""
    from pptx import Presentation

    from slidemark import build

    out = tmp_path / "t.pptx"
    build(
        "colors: primary=#DDEEFF\nstyle: table.header.fill=primary table.header.color=#FFFFFF\n\n" + ZEBRA,
        out,
    )
    tbl = next(s for s in Presentation(out).slides[0].shapes if s.has_table).table
    hdr = tbl.cell(0, 0)
    assert str(hdr.fill.fore_color.rgb) == "DDEEFF"
    assert str(hdr.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"
    assert str(tbl.cell(1, 0).fill.fore_color.rgb) != str(tbl.cell(2, 0).fill.fore_color.rgb)  # banded
