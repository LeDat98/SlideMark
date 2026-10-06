"""Importer: connectors become ``@`` link tokens, drawn flowcharts become mermaid fences."""

from __future__ import annotations

import re
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.parser import parse

EXAMPLES = sorted((Path(__file__).parent.parent / "examples").glob("*.md"))
ARROW = {"a:tailEnd": "triangle"}


def _connectors(path: Path) -> list[int]:
    return [  # card rules (decor, named "Rule") are not links: whether a card has room for them is layout
        sum(1 for sh in s.shapes if sh._element.tag.endswith("cxnSp") and sh.name.lower() != "rule")
        for s in Presentation(str(path)).slides
    ]


def _import(tmp_path: Path, src: str) -> tuple[str, list]:
    pptx = tmp_path / "a.pptx"
    build(src, pptx)
    return import_pptx(pptx, tmp_path)


def _links(text: str) -> list[str]:
    toks = [t for line in text.splitlines() if line.startswith("@") for t in line[1:].split()]
    return [t for t in toks if ">" in t or "-" in t]


def test_slide_links_are_imported(tmp_path: Path):
    src = "# T\n@a>b a-c\n## A\n- x\n## B\n- y\n## C\n- z\n"
    text, diags = _import(tmp_path, src)
    assert sorted(_links(text)) == ["a-c", "a>b"]
    assert not [d for d in diags if d.rule == "import-skipped"]
    build(text, tmp_path / "b.pptx")
    assert _connectors(tmp_path / "a.pptx") == _connectors(tmp_path / "b.pptx") == [2]


def test_links_keep_block_texts_and_survive_grid_tokens(tmp_path: Path):
    src = "# T\n@aab/aac a>b a>c\n## Big\n- 1\n- 2\n## Top\n- t\n## Bottom\n- u\n"
    text, _ = _import(tmp_path, src)
    assert {"a>b", "a>c"} <= set(_links(text))
    again = parse(text)
    assert [len(s.links) for s in again.slides] == [2]


def test_box_links_are_imported(tmp_path: Path):
    src = "# T\n## Outer\n@a>b\n### One\n- x\n### Two\n- y\n"
    text, _ = _import(tmp_path, src)
    assert "@a>b" in text.split("## Outer", 1)[1]
    build(text, tmp_path / "b.pptx")
    assert _connectors(tmp_path / "a.pptx") == _connectors(tmp_path / "b.pptx") == [1]


LINKED = [p for p in EXAMPLES if re.search(r"^@.*\b\w+[>-]\w+\b", p.read_text(encoding="utf-8"), re.M)]


@pytest.mark.parametrize("path", LINKED, ids=lambda p: p.stem)
def test_examples_keep_their_connectors(path: Path, tmp_path: Path):
    text, _ = _import(tmp_path, path.read_text(encoding="utf-8"))
    build(text, tmp_path / "b.pptx")
    assert _connectors(tmp_path / "a.pptx") == _connectors(tmp_path / "b.pptx")
    assert _links(text)


MERMAID = (
    "# Flow\n```mermaid\nflowchart LR\nA[Upload] --> B[OCR]\nB --> C{Confident?}\n"
    "C -->|yes| D[Post]\nC -->|no| E(Review)\nE --> D\n```\n"
)


def test_mermaid_is_rebuilt(tmp_path: Path):
    text, diags = _import(tmp_path, MERMAID)
    assert "```mermaid" in text and "flowchart LR" in text
    assert "A[Upload] --> B[OCR]" in text
    assert "B --> C{Confident?}" in text
    assert "-->|yes|" in text and "-->|no|" in text
    assert "(Review)" in text
    assert not [d for d in diags if d.level in ("warning", "error")]
    build(text, tmp_path / "b.pptx")
    assert _connectors(tmp_path / "a.pptx") == _connectors(tmp_path / "b.pptx") == [7]
    # same node texts after the round trip
    slide = Presentation(str(tmp_path / "b.pptx")).slides[0]
    texts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
    for word in ("Upload", "OCR", "Confident?", "Post", "Review", "yes", "no"):
        assert word in texts


def test_mermaid_top_down_and_lines(tmp_path: Path):
    src = "# T\n```mermaid\ngraph TD\nA[Top] --> B[Left]\nA --> C[Right]\nB --- D[Bottom]\nC --> D\n```\n"
    text, _ = _import(tmp_path, src)
    assert "flowchart TD" in text
    assert re.search(r"B ---(\|.*\|)? D|D --- B", text)
    assert text.count("-->") == 3
    build(text, tmp_path / "b.pptx")
    assert _connectors(tmp_path / "a.pptx") == _connectors(tmp_path / "b.pptx")


def test_mermaid_example_round_trips(tmp_path: Path):
    src = (Path(__file__).parent.parent / "examples" / "09-midnight-tech.md").read_text(encoding="utf-8")
    text, _ = _import(tmp_path, src)
    build(text, tmp_path / "b.pptx")
    assert _connectors(tmp_path / "a.pptx") == _connectors(tmp_path / "b.pptx")
    assert "```mermaid" in text


def _foreign(tmp_path: Path, connectors: list[tuple], glue: bool = False) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    boxes = []
    for i, (x, y) in enumerate([(1, 2.5), (6, 2.5), (1, 5)]):
        b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(3), Inches(1.2))
        b.text_frame.text = f"Box {i + 1}"
        boxes.append(b)
    for x0, y0, x1, y1 in connectors:
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x0), Inches(y0), Inches(x1), Inches(y1))
        if glue:
            c.begin_connect(boxes[0], 3)
            c.end_connect(boxes[1], 1)
    path = tmp_path / "f.pptx"
    prs.save(str(path))
    return path


def test_foreign_connector_resolved_by_geometry(tmp_path: Path):
    path = _foreign(tmp_path, [(4.0, 3.1, 6.0, 3.1)])
    text, _ = import_pptx(path)
    assert "a-b" in _links(text)  # no arrowhead: a plain line


def test_foreign_glued_connector(tmp_path: Path):
    path = _foreign(tmp_path, [(4.0, 3.1, 6.0, 3.1)], glue=True)
    text, _ = import_pptx(path)
    assert _links(text)


def test_stray_connector_is_ignored(tmp_path: Path):
    path = _foreign(tmp_path, [(11.0, 6.0, 12.5, 7.0), (0.1, 6.9, 0.3, 7.2)])
    text, diags = import_pptx(path)
    assert text.strip()
    assert not _links(text)
    assert not [d for d in diags if d.level == "error"]
