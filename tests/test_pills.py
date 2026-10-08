"""Table badges become native rounded pills over the cell (``layout.table_pills``)."""

from __future__ import annotations

from pptx import Presentation

from slidemark.build import build, deck_theme
from slidemark.importer import import_pptx
from slidemark.layout import layout_slide
from slidemark.layout.gantt import table_boxes
from slidemark.lint import lint
from slidemark.parser import parse
from slidemark.theme import get_theme

MD = """# 状態

| 項目 | 状態 | 備考 |
|-|-|-|
| A | [計画中]{.badge .muted} | 4.2 万台 [首位]{.badge .success} |
| B | [交渉中]{.badge} | 38 拠点 [課題]{.badge .danger} |
| C | [完了]{.badge .success} | 普通 |
"""
STATUS = ("計画中", "交渉中", "完了")


def _parts(md: str = MD, theme: str = "jp-business"):
    deck = parse(md)
    deck.theme = theme
    th, _ = deck_theme(deck, ".")
    placed = layout_slide(deck.slides[0], deck, th, 0)
    table = next(p for p in placed if p.element.type == "table")
    pills = [p for p in placed if p.element.type == "shape" and "pill" in p.element.classes]
    return table, pills, deck, placed


def _txt(p) -> str:
    return "".join(r.text for q in p.element.paragraphs for r in q.runs)


def test_pill_geometry_inside_its_cell_and_centered():
    table, pills, _, _ = _parts()
    cw, rh = table_boxes(table)
    assert len(pills) == 5
    status = [p for p in pills if _txt(p) in STATUS]
    assert len(status) == 3
    x0 = table.x + cw[0]
    for p in status:
        assert p.x >= x0 and p.x + p.w <= x0 + cw[1]
        assert any(
            abs(table.y + sum(rh[:r]) + rh[r] / 2 - (p.y + p.h / 2)) < 3000 and p.h <= rh[r]
            for r in range(1, 4)
        )
        assert p.element.shape == "rounded-rect"
        assert p.style.radius and p.style.fill


def test_pills_of_a_column_share_one_width():
    _, pills, _, _ = _parts()
    assert len({p.w for p in pills if _txt(p) in STATUS}) == 1
    tags = [p for p in pills if _txt(p) in ("首位", "課題")]
    assert len({p.w for p in tags}) == 1 and len({p.x for p in tags}) == 1  # one width, one left edge


def test_cell_text_is_kept_before_a_trailing_pill_and_emptied_otherwise():
    table, _, _, _ = _parts()
    rows = [["".join(p.plain for p in c.paragraphs) for c in row] for row in table.element.rows]
    assert rows[1][1] == "" and rows[1][2] == "4.2 万台"
    assert rows[3][2] == "普通"


def test_pill_text_ink_and_bold():
    _, pills, _, _ = _parts()
    assert next(p for p in pills if _txt(p) == "完了").style.bold is False  # synthetic bold smears CJK
    latin = _parts(MD.replace("完了", "Done"))[1]
    assert next(p for p in latin if _txt(p) == "Done").style.bold is True
    th = get_theme("jp-business")
    for p in pills:
        assert p.style.color == th.badge_ink(p.style.fill)


def test_token_off_keeps_the_old_text_highlight():
    _, pills, _, placed = _parts("style: layout.table_pills=off\n\n" + MD)
    assert not pills
    table = next(p for p in placed if p.element.type == "table")
    assert any(r.highlight for row in table.element.rows for c in row for q in c.paragraphs for r in q.runs)


def test_gantt_tables_keep_their_badges():
    # badge-only bar: text highlight; badge after a label: a pill inside the bar (test_gantt_pill.py)
    md = "# P\n\n{.gantt}\n| 施策 | 上期 |\n|-|-|\n| A | [済]{.badge} |\n"
    assert all("gantt" in p.element.classes for p in _parts(md)[1])


def test_too_narrow_cell_keeps_the_highlight():
    md = "# P\n\n| a | b |\n|-|-|\n| x | [ものすごくながいぶんしょうのバッジですよこれは本当に]{.badge} |\n"
    table, pills, _, _ = _parts(md)
    cw, _ = table_boxes(table)
    assert all(p.w <= cw[1] for p in pills)


def test_css_pill_rule_applies():
    md = "# P\n```css\n.pill { background: #FF0000 }\n```\n\n" + MD[MD.index("|") :]
    _, pills, _, _ = _parts(md)
    assert pills and all(p.style.fill == "#FF0000" for p in pills)


def test_lint_no_overlap_and_contrast_covers_pills():
    theme = get_theme("jp-business")
    _, _, deck, placed = _parts()
    assert not [d for d in lint(deck, [placed], theme) if d.rule in ("overlap", "contrast")]
    _, pills, deck2, placed2 = _parts(MD.replace("[交渉中]{.badge}", "[交渉中]{.badge .muted}"))
    bad = [p for p in placed2 if p.element.type == "shape" and "pill" in p.element.classes]
    assert bad  # contrast lint reads the pill's own fill / ink (shared resolver): readable here
    from slidemark.contrast import ratio

    for p in bad:
        assert ratio(theme.hexval(p.style.color), theme.hexval(p.style.fill)) >= 4.5


def test_render_roundrect_with_text_and_fill(tmp_path):
    out = tmp_path / "p.pptx"
    build(MD, out)
    prs = Presentation(str(out))
    pills = [s for s in prs.slides[0].shapes if s.name.startswith("Pill")]
    assert len(pills) == 5
    xml = pills[0]._element.xml
    assert 'prst="roundRect"' in xml and "<a:solidFill>" in xml and 'wrap="none"' in xml
    assert {p.text_frame.text for p in pills} == {"計画中", "交渉中", "完了", "首位", "課題"}


def test_roundtrip_returns_to_badge_syntax(tmp_path):
    a = tmp_path / "a.pptx"
    build(MD, a)
    text, _ = import_pptx(a)
    assert "[計画中]{.badge" in text and "[首位]{.badge" in text
    assert "4.2 万台 [首位]" in text
    b = tmp_path / "b.pptx"
    build(text, b)
    text2, _ = import_pptx(b)
    assert text2 == text


def test_fuzz_safe():
    for md in (
        "# T\n| a |\n|-|\n| [x]{.badge} |\n",
        "# T\n| a | b |\n|-|-|\n| [x]{.badge} [y]{.badge .danger} | [ ]{.badge} |\n",
        "# T\n| a | b |\n|-|-|\n| 1\n2 [x]{.badge} | [x]{.badge} 後ろ |\n",
        "# T\n| a | b |\n|-|-|\n| | [x]{.badge}[y]{.badge} |\n",
    ):
        _parts(md)
