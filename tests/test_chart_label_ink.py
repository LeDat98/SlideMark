"""Data labels drawn on a slice / bar get an ink readable on that point's own fill."""

from __future__ import annotations

from pptx import Presentation

from slidemark import build
from slidemark.contrast import ratio

HEAD = "theme: none\ncolors: bg=#FFFFFF fg=#1F2937 primary=#0F766E accent=#F59E0B\nlang: vi\n\n# Kênh\n"


def _chart_xml(tmp_path, body):
    src = tmp_path / "d.md"
    src.write_text(HEAD + body, encoding="utf-8")
    build(src, tmp_path / "d.pptx")
    for sh in Presentation(tmp_path / "d.pptx").slides[0].shapes:
        if sh.has_chart:
            return sh.chart._chartSpace
    raise AssertionError("no chart")


def _fills(cs):
    return [sp.xpath("string(.//a:srgbClr/@val)") for sp in cs.xpath(".//c:ser/c:dPt")]


def _inks(cs):
    return {
        int(d.xpath("string(c:idx/@val)")): d.xpath("string(c:txPr//a:srgbClr/@val)")
        for d in cs.xpath(".//c:ser/c:dLbls/c:dLbl")
    }


def test_doughnut_label_ink_readable_per_slice(tmp_path):
    cs = _chart_xml(tmp_path, "```doughnut {labels=on}\n,A,B,C\nT,86,9,5\n```\n")
    fills, inks = _fills(cs), _inks(cs)
    assert len(fills) == 3 and set(inks) == {0, 1, 2}
    for i, f in enumerate(fills):
        assert ratio("#" + inks[i], "#" + f) >= 4.5, (i, f, inks[i])
    assert inks[0] != "1F2937"  # not the dark text colour on the dark teal slice


def test_pie_percent_labels_keep_format_and_ink(tmp_path):
    cs = _chart_xml(tmp_path, "```pie {labels=percent}\n,A,B,C\nT,86,9,5\n```\n")
    d = cs.xpath(".//c:ser/c:dLbls/c:dLbl")
    assert d and all(x.xpath("string(c:showPercent/@val)") == "1" for x in d)
    assert all(x.xpath("string(c:numFmt/@formatCode)") == "0%" for x in d)
    tags = [t.tag.split("}")[1] for t in d[0]]
    assert tags.index("idx") < tags.index("numFmt") < tags.index("txPr") < tags.index("showLegendKey")
    fills, inks = _fills(cs), _inks(cs)
    for i, f in enumerate(fills):
        assert ratio("#" + inks[i], "#" + f) >= 4.5


def test_stacked_bar_label_ink_per_series(tmp_path):
    cs = _chart_xml(tmp_path, "```stacked-column {labels=on}\n,Q1,Q2\nA,5,6\nB,3,4\n```\n")
    sers = cs.xpath(".//c:ser")
    assert len(sers) == 2
    for ser in sers:
        fill = ser.xpath("string(c:spPr//a:srgbClr/@val)")
        ink = ser.xpath("string(c:dLbls/c:txPr//a:srgbClr/@val)")
        assert ratio("#" + ink, "#" + fill) >= 4.5
