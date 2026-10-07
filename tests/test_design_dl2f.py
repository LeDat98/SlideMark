"""DL2 lane F: box, row, cover and chrome decisions of the 5- and 25-slide python-pptx ports
(docs/DESIGN_COVERAGE.md, "t1 and t3"). Every form: parser (token / word reaches the IR), render (the .pptx is
reopened and its shapes asserted), importer round trip, and fuzz-safety (a bad value is a ``bad-token``
warning with a hint, never an exception)."""

from __future__ import annotations

import re

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Emu, Inches

from slidemark import build, parse
from slidemark.importer import import_pptx
from slidemark.template import deck_theme

HEAD = (
    "theme: none\n"
    "colors: bg=#FFFFFF fg=#222B36 primary=#142B4D secondary=#1F5FA8 accent=#E09F1F teal=#2A9D8F "
    "muted=#59626E surface=#EEF2F7 border=#C9D2DE\n"
)


def make(tmp_path, text: str, name: str = "d"):
    deck = build(HEAD + text, tmp_path / f"{name}.pptx")
    return deck, Presentation(str(tmp_path / f"{name}.pptx"))


def shapes(prs, i, prefix=""):
    return [s for s in prs.slides[i].shapes if s.name.lower().startswith(prefix.lower())]


def exact(prs, i, pattern: str):
    """Shapes whose whole name matches ``pattern`` (side-border lines carry the card's name plus a suffix)."""
    return [s for s in prs.slides[i].shapes if re.fullmatch(pattern, s.name)]


def fill_hex(shp) -> str:
    return str(shp.fill.fore_color.rgb)


def warns(deck, rule):
    return [d for d in deck.diagnostics if d.rule == rule]


def runs(shp):
    return [r for p in shp.text_frame.paragraphs for r in p.runs]


def text_in(prs, i, shp):
    """The text box whose centre lies inside ``shp`` (an item card or bar holds its text separately)."""
    cx, cy = shp.left + shp.width // 2, shp.top + shp.height // 2
    for t in prs.slides[i].shapes:
        if t.has_text_frame and t.text_frame.text.strip() and t is not shp:
            if t.left <= cx <= t.left + t.width and t.top <= cy <= t.top + t.height:
                return t
    return None


BOXES = (
    "\n# B\n{ITEMS}\n"
    "## 収益改善\n- 価格改定\n- SKU削減\n"
    "## 海外展開\n- タイ工場\n- 北米発売\n"
    "## DX\n- AI化\n- 電子化\n"
)


def boxes(word: str = "@3 items") -> str:
    return BOXES.replace("{ITEMS}", word)


# --------------------------------------------------------------------------- 1. @items


def test_items_word_reaches_the_slide_and_the_token_the_deck():
    d = parse(HEAD + boxes())
    assert "items" in d.slides[0].classes
    d2 = parse(HEAD + "style: box.items=cards\n" + boxes("@3"))
    assert d2.tokens["box_items"] == "cards" and not warns(d2, "bad-token")
    th, _ = deck_theme(d2, ".")
    assert th.box_items == "cards"


def test_items_draw_one_card_per_bullet_with_the_item_look(tmp_path):
    md = (
        'style: item.fill=surface item.line=teal item.border-left="5pt solid secondary" '
        "item.size=14 item.bold=off\n"
    )
    deck, prs = make(tmp_path, md + boxes())
    cards = exact(prs, 0, r"Item \d+")
    assert len(cards) == 6 and not warns(deck, "attr-ignored")
    assert {fill_hex(c) for c in cards} == {"EEF2F7"}
    assert str(cards[0].line.color.rgb) == "2A9D8F"
    texts = [text_in(prs, 0, c) for c in cards]
    assert [t.text_frame.text for t in texts[:2]] == ["価格改定", "SKU削減"]
    assert all(r.font.size.pt == 14 and not r.font.bold for t in texts for r in runs(t))
    # the left stripe is drawn as a line on the card's left edge
    xml = prs.slides[0].shapes._spTree.xml
    assert xml.count('w="63500"') >= 6  # 5pt in EMU, one per card
    # each box keeps its heading; the bullets are no list paragraphs any more
    heads = shapes(prs, 0, "Heading")
    assert [h.text_frame.text for h in heads] == ["収益改善", "海外展開", "DX"]
    assert not any("buChar" in t._element.xml for t in prs.slides[0].shapes if t.has_text_frame)


def test_items_text_is_centred_by_default_and_valign_is_honoured(tmp_path):
    _, mid = make(tmp_path, boxes(), "mid")
    _, top = make(tmp_path, "style: item.valign=top\n" + boxes(), "top")
    for prs, anchor in ((mid, MSO_ANCHOR.MIDDLE), (top, MSO_ANCHOR.TOP)):
        card = exact(prs, 0, r"Item \d+")[0]
        assert text_in(prs, 0, card).text_frame.vertical_anchor == anchor


def test_items_token_applies_without_the_word_and_off_by_default(tmp_path):
    _, off = make(tmp_path, boxes("@3"), "off")
    assert not exact(off, 0, r"Item \d+")
    _, on = make(tmp_path, "style: box.items=cards\n" + boxes("@3"), "on")
    assert len(exact(on, 0, r"Item \d+")) == 6


def test_item_class_on_sub_boxes_is_honoured_too(tmp_path):
    """`item.size/bold/valign` on `### x {.item}` used to be accepted and ignored."""
    md = (
        "style: heading.band=primary item.size=13 item.bold=off\n"
        "\n# S\n## Strategy\n@1x2\n### 価格改定 {.item}\n### SKU削減 {.item size=15}\n## Other\n- x\n"
    )
    deck, prs = make(tmp_path, md)
    cards = exact(prs, 0, r"Item \d+")
    assert len(cards) == 2 and not warns(deck, "attr-ignored")
    sizes = [text_in(prs, 0, c).text_frame.paragraphs[0].runs[0].font for c in cards]
    assert [f.size.pt for f in sizes] == [13, 15] and not any(f.bold for f in sizes)
    band = [s for s in prs.slides[0].shapes if s.name.startswith("Heading") and "価格" in s.text_frame.text]
    assert not band  # no band behind the item text: it is not a heading


def test_items_with_other_blocks_stay_bullets_with_a_hint(tmp_path):
    deck, prs = make(tmp_path, "\n# S\n@2 items\n## A\ntext first\n\n- x\n- y\n## B\n- z\n")
    got = warns(deck, "items-skipped")
    assert len(got) == 1 and got[0].level == "info" and got[0].hint
    assert len(exact(prs, 0, r"Item \d+")) == 1  # B is cards, A stays a list


def test_items_round_trip(tmp_path):
    make(tmp_path, boxes(), "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "items" in next(ln for ln in text.splitlines() if ln.startswith("@"))
    assert "- 価格改定" in text and "- 北米発売" in text and "###" not in text


# --------------------------------------------------------------------------- 2. @num


def test_num_draws_a_numbered_circle_on_every_box_heading(tmp_path):
    md = "style: box.num.fill=accent box.num.color=#000000 box.num.size=20\n" + boxes("@3 num")
    deck, prs = make(tmp_path, md)
    badges = sorted(exact(prs, 0, r"Num \d+"), key=lambda s: s.left)
    assert [b.name for b in badges] == ["Num 1", "Num 2", "Num 3"]
    assert [b.text_frame.text for b in badges] == ["1", "2", "3"]
    assert all('prst="ellipse"' in b._element.xml for b in badges)
    assert all(fill_hex(b) == "E09F1F" for b in badges)
    assert all(r.font.size.pt == 20 and str(r.font.color.rgb) == "000000" for b in badges for r in runs(b))
    assert all(b.width == b.height for b in badges) and badges[0].width > Emu(Inches(0.3))
    heads = sorted(shapes(prs, 0, "Heading"), key=lambda s: s.left)
    assert all(h.left >= b.left + b.width for h, b in zip(heads, badges, strict=True))  # text beside it
    assert not warns(deck, "attr-ignored")


def test_num_defaults_and_band_headings(tmp_path):
    deck, prs = make(tmp_path, boxes("@3 num"))
    badges = exact(prs, 0, r"Num \d+")
    assert len(badges) == 3 and {fill_hex(b) for b in badges} == {"142B4D"}  # box.num.fill = primary
    assert all(str(r.font.color.rgb) == "FFFFFF" for b in badges for r in runs(b))
    # on a heading band the circle takes the heading ink, so that it shows on the band
    deck, band = make(tmp_path, "style: heading.band=primary\n" + boxes("@3 num"), "band")
    on_band = exact(band, 0, r"Num \d+")
    assert len(on_band) == 3 and {fill_hex(b) for b in on_band} == {"FFFFFF"}
    assert all(str(r.font.color.rgb) == "142B4D" for b in on_band for r in runs(b))
    heads = sorted(shapes(band, 0, "Heading"), key=lambda s: s.left)
    assert all(
        h.left >= b.left + b.width for h, b in zip(heads, sorted(on_band, key=lambda s: s.left), strict=True)
    )
    assert not [
        d for d in deck.diagnostics if d.level == "warning" and not (d.rule or "").startswith("design-")
    ]


def test_num_is_off_for_steps_kpi_and_sub_boxes(tmp_path):
    _, prs = make(tmp_path, "\n# S\n@3 steps num\n## a\n- x\n## b\n- y\n## c\n- z\n")
    assert not exact(prs, 0, r"Num \d+")  # `num` on @steps is the STEP caption
    md = "\n# S\n@2 num\n## A\n@1x2\n### a\n- x\n### b\n- y\n## B\n- z\n"
    _, prs2 = make(tmp_path, md, "sub")
    assert [s.name for s in exact(prs2, 0, r"Num \d+")] == ["Num 1", "Num 2"]  # the ### boxes carry none


def test_num_round_trip_drops_the_badge_back_to_the_word(tmp_path):
    make(tmp_path, boxes("@3 num"), "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    at = next(ln for ln in text.splitlines() if ln.startswith("@"))
    assert "num" in at.lstrip("@").split()
    assert "## 収益改善" in text
    assert not any(ln.strip() in ("1", "2", "3", "- 1") for ln in text.splitlines())


# --------------------------------------------------------------------------- 3. heading.rule


def rules(prs, i):
    return [s for s in prs.slides[i].shapes if s.name == "rule"]


def test_heading_rule_under_plain_headings_and_bands(tmp_path):
    base = boxes("@3")
    _, plain = make(tmp_path, "style: heading.rule=teal heading.rule_h=3pt\n" + base, "plain")
    r = rules(plain, 0)
    assert len(r) == 3 and {fill_hex(s) for s in r} == {"2A9D8F"} and all(s.height == Emu(38100) for s in r)
    heads = sorted(shapes(plain, 0, "Heading"), key=lambda s: s.left)
    assert all(
        s.top >= h.top + h.height - 10 for s, h in zip(sorted(r, key=lambda s: s.left), heads, strict=True)
    )
    _, band = make(
        tmp_path, "style: heading.band=primary heading.rule=accent heading.rule_h=0.06in\n" + base, "band"
    )
    br = rules(band, 0)
    cards = sorted(shapes(band, 0, "Card"), key=lambda s: s.left)
    assert len(br) == 3 and all(
        s.left == c.left and s.width == c.width
        for s, c in zip(sorted(br, key=lambda s: s.left), cards, strict=True)
    )
    assert all(s.height == Emu(Inches(0.06)) for s in br)


def test_heading_rule_skips_sub_boxes_and_kpi_cards(tmp_path):
    md = (
        "style: heading.rule=teal\n"
        "\n# S\n## Top\n@1x2\n### a\n- x\n### b\n- y\n## Other\n- z\n"
        "\n# K\n## A {.kpi}\n1\nc\n## B {.kpi}\n2\nc\n"
    )
    _, prs = make(tmp_path, md)
    assert len(rules(prs, 0)) == 2  # the two ## boxes, not the two ### boxes
    assert not rules(prs, 1)


def test_heading_rule_makes_room_and_round_trips(tmp_path):
    _, none = make(tmp_path, boxes("@3"), "none")
    _, ruled = make(tmp_path, "style: heading.rule=teal heading.rule_h=6pt\n" + boxes("@3"), "ruled")
    first = lambda prs: min(  # noqa: E731
        (t for t in prs.slides[0].shapes if t.has_text_frame and t.text_frame.text.startswith("価格改定")),
        key=lambda t: t.top,
    )
    assert first(ruled).top > first(none).top  # the body starts below the rule
    text, _ = import_pptx(tmp_path / "ruled.pptx")
    assert "## 収益改善" in text and "- 価格改定" in text  # the rule is decor, not content


# --------------------------------------------------------------------------- 4. @rows plain


ROWS = "\n# R\n@rows plain\n- one\n- two\n- three\n"


def test_rows_plain_draws_unnumbered_bars_with_glyph_and_stripe(tmp_path):
    md = "style: rows.glyph=■ rows.glyph.color=accent rows.stripe=primary,secondary\n" + ROWS
    deck, prs = make(tmp_path, md)
    bars = [s for s in shapes(prs, 0, "Row ") if s.name in ("Row 1", "Row 2", "Row 3")]
    assert len(bars) == 3 and not shapes(prs, 0, "Row 1 num")
    stripes = [s for s in prs.slides[0].shapes if s.name.endswith("stripe")]
    assert [fill_hex(s) for s in stripes] == ["142B4D", "1F5FA8", "142B4D"]
    assert all(s.left == b.left and s.height == b.height for s, b in zip(stripes, bars, strict=True))
    glyphs = [s for s in prs.slides[0].shapes if s.name.endswith("glyph")]
    assert [g.text_frame.text for g in glyphs] == ["■", "■", "■"]
    assert all(str(r.font.color.rgb) == "E09F1F" for g in glyphs for r in runs(g))
    assert all(g.left > s.left + s.width for g, s in zip(glyphs, stripes, strict=True))
    assert bars[0].text_frame.text == "one"
    assert bars[0].text_frame.margin_left > glyphs[0].left - bars[0].left  # the text starts after the glyph
    assert not [
        d for d in deck.diagnostics if d.level == "warning" and not (d.rule or "").startswith("design-")
    ]


def test_rows_plain_takes_bullet_as_the_glyph(tmp_path):
    _, prs = make(tmp_path, "style: bullet=● bullet.color=teal\n" + ROWS)
    glyphs = [s for s in prs.slides[0].shapes if s.name.endswith("glyph")]
    assert [g.text_frame.text for g in glyphs] == ["●"] * 3
    assert all(str(r.font.color.rgb) == "2A9D8F" for g in glyphs for r in runs(g))
    _, bare = make(tmp_path, ROWS, "bare")  # no glyph token, no bullet token: bars only
    assert not [s for s in bare.slides[0].shapes if s.name.endswith("glyph")]


def test_rows_plain_accepts_numbered_lists_and_numbered_rows_keep_their_badges(tmp_path):
    _, prs = make(tmp_path, "\n# R\n@rows plain\n1. a\n2. b\n")
    assert len(shapes(prs, 0, "Row ")) == 2 and not shapes(prs, 0, "Row 1 num")
    _, num = make(tmp_path, "style: rows.stripe=teal rows.glyph=■\n\n# R\n@rows\n1. a\n2. b\n", "num")
    assert len([s for s in num.slides[0].shapes if s.name.endswith("num")]) == 2  # still numbered
    assert len([s for s in num.slides[0].shapes if s.name.endswith("stripe")]) == 2
    assert not [s for s in num.slides[0].shapes if s.name.endswith("glyph")]  # a badge is the marker


def test_rows_bullets_without_plain_are_skipped_with_a_hint(tmp_path):
    deck, _ = make(tmp_path, "\n# R\n@rows\n- a\n- b\n")
    assert warns(deck, "rows-skipped")


def test_rows_plain_round_trip(tmp_path):
    make(tmp_path, "style: rows.glyph=■ rows.stripe=primary,secondary\n" + ROWS, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "@rows plain" in text and "- one" in text and "■" not in text.split("```")[0].split("# R")[1]


# --------------------------------------------------------------------------- 5. cover


COVER = "\n# Title\n@cover bg=primary dark\n## sub line\n\n# Body\n- x\n"
CHROME = "style: cover.band=none cover.band_h=60% cover.rule=accent cover.rule_h=0.06in "


def test_cover_bar_on_the_slide_edge(tmp_path):
    _, prs = make(tmp_path, CHROME + "cover.bar=accent@edge cover.bar_w=0.35in\n" + COVER)
    bar = next(s for s in rules(prs, 0) if fill_hex(s) == "E09F1F" and s.width == Emu(Inches(0.35)))
    assert bar.left == 0 and bar.top == 0 and bar.height == prs.slide_height
    title = next(s for s in prs.slides[0].shapes if s.name == "Title")
    assert title.left > bar.width  # the text keeps right of the bar
    _, beside = make(tmp_path, CHROME + "cover.bar=accent cover.bar_w=0.35in\n" + COVER, "beside")
    side = next(s for s in rules(beside, 0) if s.width == Emu(Inches(0.35)))
    assert side.height < beside.slide_height  # without @edge it stays beside the title block


def test_cover_short_rule_above_or_below_the_title(tmp_path):
    base = CHROME + "cover.rule_w=6in "
    _, below = make(tmp_path, base + "cover.rule_pos=below\n" + COVER, "below")
    _, above = make(tmp_path, base + "cover.rule_pos=above\n" + COVER, "above")
    for prs in (below, above):
        r = [s for s in rules(prs, 0) if s.width == Emu(Inches(6))]
        assert len(r) == 1 and r[0].height == Emu(Inches(0.06))
    title = next(s for s in above.slides[0].shapes if s.name == "Title")
    sub = next(s for s in above.slides[0].shapes if s.name == "Subtitle")
    ra = next(s for s in rules(above, 0) if s.width == Emu(Inches(6)))
    rb = next(s for s in rules(below, 0) if s.width == Emu(Inches(6)))
    assert ra.top + ra.height <= title.top + title.height  # above the title text
    ttl_b = next(s for s in below.slides[0].shapes if s.name == "Title")
    sub_b = next(s for s in below.slides[0].shapes if s.name == "Subtitle")
    assert (
        ttl_b.top + ttl_b.height <= rb.top and rb.top + rb.height <= sub_b.top
    )  # between title and subtitle
    assert sub is not None


def test_cover_rule_stays_full_width_without_rule_w(tmp_path):
    _, prs = make(tmp_path, CHROME + "\n" + COVER.lstrip("\n"))
    assert any(s.width == prs.slide_width and s.height == Emu(Inches(0.06)) for s in rules(prs, 0))


def test_cover_stripes_at_x_positions_and_the_title_keeps_clear(tmp_path):
    md = CHROME + "cover.stripes=#1C3A68@8.9in,secondary@10.2in\n" + COVER
    _, prs = make(tmp_path, md)
    st = sorted((s for s in rules(prs, 0) if s.height == prs.slide_height), key=lambda s: s.left)
    assert [fill_hex(s) for s in st] == ["1C3A68", "1F5FA8"]
    assert [round(s.left / 914400, 2) for s in st] == [8.9, 10.2]
    assert all(s.left + s.width == prs.slide_width for s in st)
    title = next(s for s in prs.slides[0].shapes if s.name == "Title")
    assert title.left + title.width <= st[0].left  # the text never runs into a stripe
    z = [s.name for s in prs.slides[0].shapes]
    assert z.index("Title") > z.index("rule")  # stripes are behind the text


def test_cover_decoration_is_drawn_only_on_the_cover(tmp_path):
    md = CHROME + "cover.bar=accent@edge cover.stripes=secondary@10in\n" + COVER
    _, prs = make(tmp_path, md)
    assert not any(s.height == prs.slide_height for s in rules(prs, 1))


def test_cover_decoration_round_trips_as_a_cover(tmp_path):
    md = CHROME + "cover.bar=accent@edge cover.stripes=secondary@10in cover.rule_w=2in\n" + COVER
    make(tmp_path, md, "rt")
    text, _ = import_pptx(tmp_path / "rt.pptx")
    assert "# Title" in text and "sub line" in text
    assert "# Body" in text and "- x" in text


# --------------------------------------------------------------------------- 6. title.rule2


def test_title_rule2_is_a_short_segment_on_the_left_end(tmp_path):
    md = "style: title.rule=primary title.rule_h=3pt title.rule2=accent title.rule2_w=1.6in\n\n# T\n- a\n"
    _, prs = make(tmp_path, md)
    r = sorted(rules(prs, 0), key=lambda s: -s.width)
    assert [fill_hex(s) for s in r] == ["142B4D", "E09F1F"]
    long, short = r
    assert short.width == Emu(Inches(1.6)) and short.left == long.left and short.top == long.top
    assert short.height == long.height == Emu(38100)
    order = [s.shape_id for s in prs.slides[0].shapes]
    assert order.index(short.shape_id) > order.index(long.shape_id)  # drawn over it


def test_title_rule2_alone_still_draws(tmp_path):
    _, prs = make(tmp_path, "style: title.rule2=teal\n\n# T\n- a\n")
    assert [fill_hex(s) for s in rules(prs, 0)] == ["2A9D8F"]


# --------------------------------------------------------------------------- 7. chevron list


STEPS = "\n# S\n@4 steps\n" + "".join(f"## s{i}\n- t{i}\n" for i in range(4))


def arrows(prs):
    return sorted((s for s in shapes(prs, 0, "Step ") if s.name.endswith("arrow")), key=lambda s: s.left)


def test_chevron_shape_list_first_pentagon_rest_chevrons(tmp_path):
    _, prs = make(tmp_path, "style: render.chevron_shape=pentagon,chevron\n" + STEPS)
    kinds = ["homePlate" if "homePlate" in s._element.xml else "chevron" for s in arrows(prs)]
    assert kinds == ["homePlate", "chevron", "chevron", "chevron"]
    _, all_p = make(tmp_path, "style: render.chevron_shape=pentagon\n" + STEPS, "all")
    assert all("homePlate" in s._element.xml for s in arrows(all_p))
    _, rev = make(tmp_path, "style: render.chevron_shape=chevron,pentagon\n" + STEPS, "rev")
    kinds = ["homePlate" if "homePlate" in s._element.xml else "chevron" for s in arrows(rev)]
    assert kinds == ["chevron"] + ["homePlate"] * 3  # the last word repeats


def test_chevron_shape_list_on_a_compact_row_and_round_trip(tmp_path):
    md = (
        "style: render.chevron_shape=pentagon,chevron layout.chevron_steps=off\n"
        "\n# S\n@3 chevron\n## a\n- x\n## b\n- y\n## c\n- z\n"
    )
    _, prs = make(tmp_path, md)
    xs = sorted(
        (s for s in prs.slides[0].shapes if s.shape_type is not None and "Shape" in s.name),
        key=lambda s: s.left,
    )
    assert sum("homePlate" in s._element.xml for s in xs) == 1 and "homePlate" in xs[0]._element.xml
    _, steps = make(tmp_path, "style: render.chevron_shape=pentagon,chevron\n" + STEPS, "st")
    text, _ = import_pptx(tmp_path / "st.pptx")
    assert "@4 steps" in text


# --------------------------------------------------------------------------- 8. table.num_pad, kpi.fill


TABLE = "\n# T\n```table {align=lrr}\na,b,c\nx,12,34\ny,56,78\n```\n"


def test_table_num_pad_insets_the_right_aligned_cells(tmp_path):
    _, prs = make(tmp_path, "style: table.num_pad=0.6in\n" + TABLE)
    tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
    assert tbl.cell(1, 1).margin_right == Emu(Inches(0.6)) and tbl.cell(2, 2).margin_right == Emu(Inches(0.6))
    assert tbl.cell(0, 1).margin_right == Emu(Inches(0.6))  # the header of a numeric column too
    assert tbl.cell(1, 0).margin_right != Emu(Inches(0.6))  # the label column keeps the global inset
    _, plain = make(tmp_path, TABLE, "plain")
    t2 = next(s for s in plain.slides[0].shapes if s.has_table).table
    assert t2.cell(1, 1).margin_right != Emu(Inches(0.6))


KPI = "\n# K\n## A {.kpi}\n1\ncap\n## B {.kpi}\n2\ncap\n"


def test_kpi_fill_and_line_are_the_card_look(tmp_path):
    deck, prs = make(tmp_path, "style: kpi.fill=accent kpi.line=teal\n" + KPI)
    cards = shapes(prs, 0, "Card")
    assert len(cards) == 2 and {fill_hex(c) for c in cards} == {"E09F1F"}
    assert all(str(c.line.color.rgb) == "2A9D8F" for c in cards)
    assert not warns(deck, "attr-ignored")
    _, base = make(tmp_path, KPI, "base")
    assert {fill_hex(c) for c in shapes(base, 0, "Card")} == {"EEF2F7"}  # the card class otherwise


def test_kpi_fill_does_not_leak_into_the_number(tmp_path):
    _, prs = make(tmp_path, "style: kpi.fill=accent\n" + KPI)
    texts = [s for s in prs.slides[0].shapes if s.has_text_frame and s.text_frame.text.strip()]
    assert texts and not [s for s in texts if s.fill.type == 1]  # only the two cards carry the fill


# --------------------------------------------------------------------------- 9. contrast on a tall kpi stripe


def test_label_on_a_tall_kpi_stripe_is_judged_against_the_stripe(tmp_path):
    ok = "style: kpi.stripe=primary kpi.stripe_h=0.7in kpi.label.color=#FFFFFF kpi.label.size=14\n"
    deck, _ = make(tmp_path, ok + KPI)
    assert not warns(deck, "contrast"), [d.message for d in warns(deck, "contrast")]
    # a label that really is unreadable on the stripe still warns (dark ink on a dark stripe)
    bad = "style: kpi.stripe=primary kpi.stripe_h=0.7in kpi.label.color=#142B4D kpi.label.size=14\n"
    deck2, _ = make(tmp_path, bad + KPI, "bad")
    assert warns(deck2, "contrast")


def test_a_thin_stripe_does_not_change_the_verdict(tmp_path):
    deck, _ = make(tmp_path, "style: kpi.stripe=primary kpi.stripe_h=6pt\n" + KPI)
    assert not warns(deck, "contrast")


# --------------------------------------------------------------------------- fuzz-safety

BAD = [
    ("heading.rule=notacolor", "heading.rule"),
    ("heading.rule_h=thick", "heading.rule_h"),
    ("box.items=maybe", "box.items"),
    ("box.num.fill=nope", "box.num.fill"),
    ("box.num.color=nope", "box.num.color"),
    ("box.num.size=big", "box.num.size"),
    ("rows.stripe=primary,nope", "rows.stripe"),
    ("rows.stripe=", "rows.stripe"),
    ("rows.glyph.color=nope", "rows.glyph.color"),
    ("cover.bar=accent@middle", "cover.bar"),
    ("cover.bar=nope@edge", "cover.bar"),
    ("cover.rule_w=wide", "cover.rule_w"),
    ("cover.rule_pos=left", "cover.rule_pos"),
    ("cover.stripes=navy", "cover.stripes"),
    ("cover.stripes=primary@far", "cover.stripes"),
    ("title.rule2=nope", "title.rule2"),
    ("title.rule2_w=long", "title.rule2_w"),
    ("render.chevron_shape=triangle", "render.chevron_shape"),
    ("render.chevron_shape=pentagon,star", "render.chevron_shape"),
    ("table.num_pad=wide", "table.num_pad"),
]


@pytest.mark.parametrize(("token", "key"), BAD)
def test_bad_values_are_a_diagnostic_with_a_hint(token, key, tmp_path):
    deck = parse(HEAD + f"style: {token}\n" + boxes("@3 items num") + "\n# C\n@cover\n## s\n")
    got = [d for d in deck.diagnostics if d.rule == "bad-token" and f"{key}=" in d.message]
    assert len(got) == 1, (token, [d.message for d in deck.diagnostics])
    assert got[0].hint and len(got[0].hint) < 160
    built = build(HEAD + f"style: {token}\n" + boxes("@3 items num"), tmp_path / "x.pptx")  # never raises
    assert not [d for d in built.diagnostics if d.level == "error"]


def test_good_values_build_without_warnings(tmp_path):
    md = (
        "style: heading.rule=accent heading.rule_h=2pt box.items=cards box.num.fill=teal box.num.size=18 "
        "rows.glyph=■ rows.stripe=primary title.rule2=accent title.rule2_w=1in table.num_pad=0.3in "
        "render.chevron_shape=pentagon,chevron\n" + boxes("@3 num") + ROWS + STEPS + TABLE
    )
    deck, _ = make(tmp_path, md)
    bad = [d for d in deck.diagnostics if d.level == "warning" and not (d.rule or "").startswith("design-")]
    assert not bad, [(d.rule, d.message) for d in bad]


def test_items_and_num_on_odd_slides_never_crash(tmp_path):
    md = (
        "\n# A\n@items num\n\n# B\n@items\n## only a title\n"
        "\n# C\n@2 items num\n## x\n1. one\n2. two\n## y\n> note\n"
        "\n# D\n@rows plain\n\n# E\n@num\n- loose bullets\n"
    )
    deck, prs = make(tmp_path, md)
    assert len(prs.slides) == 5 and not [d for d in deck.diagnostics if d.level == "error"]
