"""Wave 2A: a lone block after a full `@N` row spans the width; hollow cards beside a chart spread (ruled);
a chevron row above a table with as many columns follows the table column edges."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu

from slidemark import build
from slidemark.ir import Chart, Code, Container, Shape, Table
from slidemark.layout import layout_slide
from slidemark.parser import parse
from slidemark.template import deck_theme
from slidemark.theme import get_theme

EX = Path(__file__).resolve().parent.parent / "examples"
COVER = "# Cover\nsub\n\n"
KPIS = "@3\n## A {.kpi}\n1\n## B {.kpi}\n2\n## C {.kpi}\n3\n@end\n"
LINE = "```line\n,a,b\ns,1,2\n```\n"


def lay(md: str, idx: int = 1, **tokens):
    th = get_theme("jp-business").model_copy(deep=True)
    th.layout = th.layout.model_copy(update=tokens)
    deck = parse(COVER + md)
    return layout_slide(deck.slides[idx], deck, th, idx), deck, th


def of(placed, cls):
    return [p for p in placed if isinstance(p.element, cls)]


# ------------------------------------------------------------------ 1. lone block after a box row


def test_lone_chart_after_row_spans_full_width():
    placed, _, _ = lay("# T\n" + KPIS + LINE)
    cards, (chart,) = of(placed, Container), of(placed, Chart)
    assert len(cards) == 3
    assert chart.x <= cards[0].x + 2 and chart.x + chart.w >= cards[-1].x + cards[-1].w - 2
    assert chart.y >= max(c.y + c.h for c in cards)


def test_lone_code_spans_but_short_row_keeps_cell():
    placed, _, _ = lay("# T\n" + KPIS + "```py\nx = 1\n```\n")
    code = of(placed, Code)[0]
    assert code.w > 0.9 * 12.33 * 914400  # spans the row width
    two = "@3\n## A {.kpi}\n1\n## B {.kpi}\n2\n" + LINE  # row not full: the chart is the third cell
    p2, _, _ = lay("# T\n" + two)
    c2 = of(p2, Chart)[0]
    assert c2.w < 0.4 * 12.33 * 914400


def test_lone_chart_after_row_renders_full_width(tmp_path):
    out = tmp_path / "a.pptx"
    build(COVER + "# T\n" + KPIS + LINE, out)
    prs = Presentation(str(out))
    gf = [s for s in prs.slides[1].shapes if s.has_chart]
    assert gf and gf[0].width > 0.85 * prs.slide_width


# ------------------------------------------------------------------ 2. stacked cards beside a chart


def lay20(idx: int, **tokens):
    deck = parse((EX / "20-jp-retail-dense.md").read_text(encoding="utf-8"))
    th, _ = deck_theme(deck, EX)
    th = th.model_copy(deep=True)
    th.layout = th.layout.model_copy(update=tokens)
    return layout_slide(deck.slides[idx], deck, th, idx)


def _ruled(placed):
    return [p for p in placed if isinstance(p.element, Shape) and p.element.shape == "line"]


def test_hollow_stacked_cards_beside_chart_spread_with_rules():
    placed = lay20(5)  # `@aab/aac`: chart left, two stacked cards right
    cards = sorted(of(placed, Container), key=lambda c: c.y)
    assert len(cards) == 2
    for c in cards:
        rules = [r for r in _ruled(placed) if c.y < r.y < c.y + c.h]
        assert len(rules) == 2 and all(c.x < r.x and r.x + r.w < c.x + c.w for r in rules)
    assert not _ruled(lay20(5, card_spread_fill=0.0))
    assert not _ruled(lay20(5, card_spread_rules=False))


def test_full_stacked_cards_are_not_spread():
    """Cards already filled to card_spread_fill keep their layout (no rules)."""
    assert not _ruled(lay20(5, card_spread_fill=0.01))


# ------------------------------------------------------------------ 3. chevrons follow table columns

CHEV = (
    "# T\n> lead\n@4 chevron\n## One\na\n## Two\nb\n## Three\nc\n## Four\nd\n@end\n"
    "| A | B | C | D |\n|-|-|-|-|\n| 1 | 2 long long long | 3 | 4 |\n| 1 | 2 | 3 | 4 |\n"
)


def _chevs(placed):
    return sorted(
        (p for p in placed if isinstance(p.element, Shape) and p.element.shape == "chevron"),
        key=lambda p: p.x,
    )


def test_chevrons_follow_table_columns():
    placed, _, _ = lay(CHEV)
    chevs, (tbl,) = _chevs(placed), of(placed, Table)
    cw, x = tbl.element.attrs["_col_w"], tbl.x
    assert len(chevs) == len(cw) == 4
    for i, (c, w) in enumerate(zip(chevs, cw, strict=True)):
        assert c.x == x  # chevron i starts at column i
        assert c.x + c.w <= x + w + 1
        if i == 3:
            assert abs(c.x + c.w - (x + w)) <= 1
        x += w
    assert chevs[-1].x + chevs[-1].w == tbl.x + tbl.w
    off, _, _ = lay(CHEV, chevron_table_align=False)
    assert [c.x for c in _chevs(off)] != [c.x for c in chevs]


def test_chevron_count_differs_from_columns_unchanged():
    five = CHEV.replace("| A | B | C | D |\n|-|-|-|-|", "| A | B | C | D | E |\n|-|-|-|-|-|").replace(
        "| 3 | 4 |", "| 3 | 4 | 5 |"
    )
    a, _, _ = lay(five)
    b, _, _ = lay(five, chevron_table_align=False)
    assert [(c.x, c.w) for c in _chevs(a)] == [(c.x, c.w) for c in _chevs(b)]


def test_chevron_table_renders_aligned(tmp_path):
    out = tmp_path / "c.pptx"
    build(COVER + CHEV, out)
    prs = Presentation(str(out))
    shapes = list(prs.slides[1].shapes)
    tbl = next(s for s in shapes if s.has_table)
    lefts = [tbl.left]
    for col in tbl.table.columns:
        lefts.append(lefts[-1] + col.width)
    chev = sorted(
        (s for s in shapes if s.shape_type == 1 and s.auto_shape_type == MSO_SHAPE.CHEVRON),
        key=lambda s: s.left,
    )
    assert len(chev) == 4
    for s, left in zip(chev, lefts, strict=False):
        assert abs(s.left - left) <= Emu(2000)


def test_fuzz_never_raises():
    bad = [
        "# T\n@3\n## A\n1\n## B\n2\n## C\n3\n@end\n```mermaid\n```\n",
        "# T\n@4 chevron\n## A\n@end\n| a |\n|-|\n",
        "# T\n@aab/aac\n## A\n## B\n```line\n```\n",
        "# T\n@3\n@end\n![x](missing.png)\n",
    ]
    for md in bad:
        deck = parse(COVER + md)
        th = get_theme("jp-business")
        layout_slide(deck.slides[1], deck, th, 1)
