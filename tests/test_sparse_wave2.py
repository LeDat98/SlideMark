"""Design wave 2: sparse slides use the body (short lists, card rows, lone chevrons / tables) and the
conclusion bar attaches to a lone table. Every behaviour has a token switch (``layout.*``)."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from slidemark import build
from slidemark.ir import Container, Table
from slidemark.layout import measure

from .test_layout_l3_fill import lay, role

H = 720 * 9525
BAR = "> 冷凍食品と海外が成長をけん引\n"
HEAD = "theme: jp-business\nlang: ja\n\n"

LIST = (
    HEAD
    + "# 価格改定の方針\n> 原材料高を価格に反映する\n- 主力3ブランドを平均6%改定\n- 改定は2027年5月出荷分から\n- 容量変更は行わない\n- 販促費を前年比10%削減\n"  # noqa: E501
)
CARDS = HEAD + (
    "# 重点地域\n## タイ\n- 工場の生産能力を1.5倍に\n- ASEAN向け輸出拠点\n"
    "## 北米\n- 冷凍餃子を発売\n- 大手スーパー3社で展開\n## 台湾\n- 調味料の現地生産\n- コンビニ向けPB\n"
)
CHEV = (
    HEAD
    + "# 年間スケジュール\n@chevron\n## 4–6月\n- 価格改定\n## 7–9月\n- タイ工場着工\n## 10–12月\n- 北米発売\n"
)
CHEV_PLAIN = HEAD + "# 年間スケジュール\n@chevron\n## 4–6月\n## 7–9月\n## 10–12月\n"
TABLE = HEAD + (
    "# 事業別の売上計画\n```table {align=lrrr}\n事業,2026実績,2027計画,前年比\n冷凍食品,420,465,+11%\n"
    "調味料,310,325,+5%\n飲料,255,270,+6%\n海外,200,220,+10%\n```\n"
)


def _pt(p) -> float:
    return (p.style.font_size or 0) * p.font_scale


def _bodies(placed):
    return [p for p in role(placed, "body")]


# ---- 1a: a short list grows with the free height --------------------------------------------------


def test_short_list_grows_and_leaves_no_floating_block():
    placed, _, _ = lay("", 1, src=LIST, grow_max=3.0)  # (the deck-wide ceiling, grow_max: test_grow_max.py)
    off, _, _ = lay("", 1, src=LIST, list_fill=False, grow_max=3.0)
    (b,), (b0,) = _bodies(placed), _bodies(off)
    assert _pt(b) > _pt(b0)
    assert _pt(b) <= 28 + 0.01  # list_text_max_pt
    lead = role(placed, "lead")[0]
    assert b.y > lead.y + lead.h  # under the lead, never over it
    assert b.y + b.h < H
    assert b.element.attrs["para_gap"] >= measure.para_gap()
    block = b.h
    assert block >= 0.4 * (H - (lead.y + lead.h))  # covers a real share of the body


def test_short_list_text_cap_is_a_token():
    off, _, _ = lay("", 1, src=LIST, list_fill=False, grow_max=3.0)
    placed, _, _ = lay("", 1, src=LIST, list_text_max_pt=20, grow_max=3.0)
    (b,), (b0,) = _bodies(placed), _bodies(off)
    assert _pt(b) <= max(20, _pt(b0)) + 0.01  # the pass never grows past its cap (and never shrinks)


def test_a_lone_paragraph_and_a_pinned_size_are_not_grown():
    cover = HEAD + "# 事業計画\n## 青葉フーズ\n\n"
    one = (
        cover
        + "# 方針\n> 原材料高を価格に反映する\n原材料高を価格に反映し、販売数量を維持する。主力ブランドは平均6%の改定を行い、容量は変更しない。\n"  # noqa: E501
    )
    placed, _, _ = lay("", 2, src=one)
    off, _, _ = lay("", 2, src=one, list_fill=False)
    assert _pt(_bodies(placed)[0]) == pytest.approx(_pt(_bodies(off)[0]))
    pinned = cover + LIST.replace(HEAD, "").replace("- 主力3", "{size=14}\n- 主力3")
    placed, _, _ = lay("", 2, src=pinned)
    off, _, _ = lay("", 2, src=pinned, list_fill=False)
    assert [(p.y, p.h) for p in placed] == [(p.y, p.h) for p in off]


# ---- 1b: cards stretch down the body --------------------------------------------------------------


def test_cards_alone_stretch_down_the_body():
    placed, _, _ = lay(
        "", 1, src=CARDS, grow_max=3.0
    )  # (a ceiling that keeps the text small keeps the cards compact)
    off, _, _ = lay("", 1, src=CARDS, cards_to_body=False, grow_max=3.0)
    cs = [p for p in placed if isinstance(p.element, Container)]
    cs0 = [p for p in off if isinstance(p.element, Container)]
    assert len(cs) == 3 and len({c.h for c in cs}) == 1
    assert cs[0].h > 1.3 * cs0[0].h
    assert cs[0].y + cs[0].h <= H  # never past the slide
    assert cs[0].y + cs[0].h > 0.6 * H  # no longer only the top third
    for c in cs:  # every text stays inside its card
        for p in _bodies(placed):
            if c.x <= p.x < c.x + c.w:
                assert p.y + p.h <= c.y + c.h + 2


def test_cards_text_cap_is_a_token():
    placed, _, _ = lay("", 1, src=CARDS, card_fill_text_max_pt=12)
    assert max(_pt(p) for p in _bodies(placed)) <= 21.5  # no growth beyond what the slide had


def test_cards_with_content_that_fills_them_are_left_alone():
    lines = "".join(f"- 項目{j}の説明文をここに書きます。少し長めの文章です\n" for j in range(9))
    dense = HEAD + "# 重点\n" + "".join(f"## 施策{i}\n" + lines for i in range(3))
    a, _, _ = lay("", 1, src=dense)
    b, _, _ = lay("", 1, src=dense, cards_to_body=False)
    assert [(p.x, p.y, p.w, p.h) for p in a] == [(p.x, p.y, p.w, p.h) for p in b]


# ---- 1c: @chevron with bodies is built as @steps --------------------------------------------------


def test_chevron_with_short_bodies_becomes_steps():
    placed, deck, _ = lay("", 1, src=CHEV)
    assert (
        len([p for p in placed if isinstance(p.element, Container) and "steps-card" in p.element.classes])
        == 3
    )
    arrows = [
        p for p in placed if str(getattr(p.element, "attrs", {}).get("shape_name", "")).endswith(" arrow")
    ]
    assert len(arrows) == 3
    assert deck.slides[0].classes == ["chevron"]  # the source IR is untouched
    off, _, _ = lay("", 1, src=CHEV, chevron_steps=False)
    assert not [p for p in off if isinstance(p.element, Container) and "steps-card" in p.element.classes]


def test_plain_chevron_and_long_bodies_stay_chevrons():
    for src in (
        CHEV_PLAIN,
        CHEV.replace("- 価格改定\n", "- 価格改定\n- 説明1\n- 説明2\n- 説明3\n"),
        CHEV + "@end\n```table\na,b\n1,2\n```\n",
    ):
        placed, _, _ = lay("", 1, src=src)
        assert not [
            p for p in placed if isinstance(p.element, Container) and "steps-card" in p.element.classes
        ]


def test_chevron_steps_render_as_arrow_and_card_shapes(tmp_path: Path):
    out = tmp_path / "c.pptx"
    build(CHEV, out)
    names = [s.name for s in Presentation(out).slides[0].shapes]
    assert [f"Step {i} arrow" for i in (1, 2, 3)] == [n for n in names if n.endswith(" arrow")]
    assert [f"Step {i} card" for i in (1, 2, 3)] == [n for n in names if n.endswith(" card")]


# ---- 2 + 3: lone table grows, the bar attaches ----------------------------------------------------


def _table(placed):
    (t,) = [p for p in placed if isinstance(p.element, Table)]
    return t


def test_lone_table_text_grows_up_to_the_token_max():
    placed, _, _ = lay("", 1, src=TABLE, grow_max=3.0)
    off, _, _ = lay("", 1, src=TABLE, table_free=False, grow_max=3.0)
    t, t0 = _table(placed), _table(off)
    assert _pt(t) > _pt(t0)
    assert _pt(t) <= 22 + 0.01
    assert t.h > t0.h
    capped, _, _ = lay("", 1, src=TABLE, table_free_text_max_pt=16, grow_max=3.0)
    assert _pt(_table(capped)) <= max(16, _pt(t0)) + 0.01


@pytest.mark.parametrize("mode", ["grow", "move"])
def test_bar_attaches_to_the_table(mode):
    src = TABLE + BAR
    placed, _, _ = lay("", 1, src=src, bar_attach=mode)
    off, _, _ = lay("", 1, src=src, bar_attach="off", table_free=False)
    t, bar = _table(placed), role(placed, "conclusion")[0]
    t0, bar0 = _table(off), role(off, "conclusion")[0]
    gap = bar.y - (t.y + t.h)
    assert 0 <= gap < bar0.y - (t0.y + t0.h)  # closer than before
    assert gap <= 0.3 * 914400 * 1.0 + 1  # about bar_attach_gap (0.15in) plus rounding
    assert t.y + t.h <= bar.y and bar.y + bar.h <= H
    if mode == "grow":
        assert t.h >= t0.h


def test_bar_attach_off_keeps_the_bar_at_the_bottom():
    placed, _, _ = lay("", 1, src=TABLE + BAR, bar_attach="off", table_free=False)
    on, _, _ = lay("", 1, src=TABLE + BAR, bar_attach="off")
    assert role(placed, "conclusion")[0].y == role(on, "conclusion")[0].y


def test_dense_tables_keep_their_size():
    rows = "".join(f"行{i},{i},{i * 2},{i * 3}\n" for i in range(12))
    src = HEAD + "# 一覧\n```table\n項目,A,B,C\n" + rows + "```\n"
    a, _, _ = lay("", 1, src=src)
    b, _, _ = lay("", 1, src=src, table_free=False)
    assert _pt(_table(a)) == pytest.approx(_pt(_table(b)))


def test_table_text_never_wraps_more_cjk_lines():
    placed, _, _ = lay("", 1, src=TABLE)
    t = _table(placed)
    rh = t.element.attrs["_row_h"]
    assert sum(rh) == t.h
    one_line = _pt(t) * 12700 * 1.6
    assert all(h >= one_line for h in rh[1:])


def test_table_pptx_font_size_follows_the_layout(tmp_path: Path):
    out = tmp_path / "t.pptx"
    build(TABLE.replace(HEAD, HEAD + "style: grow.max=3\n") + BAR, out)  # (the ceiling: test_grow_max.py)
    slide = Presentation(out).slides[0]
    tbl = next(s for s in slide.shapes if s.has_table).table
    size = tbl.cell(1, 0).text_frame.paragraphs[0].runs[0].font.size.pt
    assert 20 <= size <= 22.01
    gf = next(s for s in slide.shapes if s.has_table)
    bar = next(s for s in slide.shapes if s.has_text_frame and "成長" in s.text_frame.text)
    assert Emu(0) <= bar.top - (gf.top + gf.height) <= Emu(int(0.4 * 914400))


# ---- KPI rows and steps fill the body (small-body themes) ------------------------------------------------

KPI = HEAD + (
    "# 目標\n> 売上・利益ともに過去最高を目指す\n## 売上高 {.kpi .hero}\n1,280億円\n前年比 +8%\n"
    "## 営業利益 {.kpi}\n96億円\n前年比 +12%\n## ROE {.kpi}\n9.0%\n+0.8pt\n"
)
KPI_FLAT = KPI.replace(" .hero", "")


def _value_pt(p) -> float:
    return p.element.paragraphs[0].style.font_size * p.font_scale


def _kpi_cards(placed):
    return [p for p in placed if isinstance(p.element, Container) and "kpi" in p.element.classes]


def test_lone_kpi_row_stretches_down_the_body_and_the_hero_stays_larger():
    placed, _, _ = lay("", 1, src=KPI)
    off, _, _ = lay("", 1, src=KPI, kpi_to_body=False)
    cs, cs0 = _kpi_cards(placed), _kpi_cards(off)
    assert len({(c.y, c.h) for c in cs}) == 1
    assert cs[0].h > 1.15 * cs0[0].h
    assert cs[0].h <= 0.72 * H  # kpi_to_body_h of the body at most
    assert cs[0].y + cs[0].h < H
    vals = sorted((p.x, _value_pt(p)) for p in _bodies(placed))
    assert vals[0][1] > vals[1][1]  # the hero keeps its larger value
    flat, _, _ = lay("", 1, src=KPI_FLAT)
    assert len({round(_value_pt(p), 1) for p in _bodies(flat)}) == 1  # no hero invented: one weight


def test_kpi_to_body_share_is_a_token():
    a, _, _ = lay("", 1, src=KPI, kpi_to_body_h=0.8)
    b, _, _ = lay("", 1, src=KPI, kpi_to_body_h=0.6)
    assert _kpi_cards(a)[0].h > _kpi_cards(b)[0].h


def test_kpi_text_stays_inside_the_card_and_in_one_line_cjk():
    placed, _, _ = lay("", 1, src=KPI)
    for c in _kpi_cards(placed):
        for p in _bodies(placed) + role(placed, "heading"):
            if c.x <= p.x < c.x + c.w:
                assert c.y <= p.y and p.y + p.h <= c.y + c.h + 2


STEPS = HEAD + "# 年間\n@steps\n## 4–6月\n- 価格改定\n## 7–9月\n- タイ工場着工\n## 10–12月\n- 北米発売\n"


def _step_cards(placed):
    return [p for p in placed if isinstance(p.element, Container) and "steps-card" in p.element.classes]


def test_sparse_steps_fill_the_body_without_a_bar():
    placed, _, _ = lay("", 1, src=STEPS)
    off, _, _ = lay("", 1, src=STEPS, steps_to_body=False)
    cs, cs0 = _step_cards(placed), _step_cards(off)
    assert len(cs) == 3
    assert cs[0].y + cs[0].h > cs0[0].y + cs0[0].h + 0.05 * H
    assert cs[0].y + cs[0].h <= H
    assert cs[0].h <= 2.0 * cs[0].w + 1  # steps_to_body_aspect
    for c in cs:
        (t,) = [p for p in _bodies(placed) if c.x <= p.x < c.x + c.w]
        assert c.y <= t.y and t.y + t.h <= c.y + c.h + 2
    chev, _, _ = lay("", 1, src=STEPS.replace("@steps", "@chevron"))
    assert [(p.x, p.y, p.h) for p in _step_cards(chev)] == [(p.x, p.y, p.h) for p in cs]  # same either way


def test_steps_with_a_bar_still_stretch_to_the_bar():
    placed, _, _ = lay("", 1, src=STEPS + BAR)
    bar = role(placed, "conclusion")[0]
    assert max(c.y + c.h for c in _step_cards(placed)) <= bar.y


def test_kpi_and_steps_pptx_geometry(tmp_path: Path):
    out = tmp_path / "k.pptx"
    build(KPI + STEPS.replace(HEAD, "\n"), out)
    prs = Presentation(out)
    k, st = prs.slides[0], prs.slides[1]
    box = max(
        (s for s in k.shapes if s.height > 0.3 * prs.slide_height and s.width < prs.slide_width),
        key=lambda s: s.height,
    )
    assert box.height >= 0.45 * prs.slide_height
    step_cards = [s for s in st.shapes if s.name.endswith(" card")]
    assert len(step_cards) == 3
    assert max(s.top + s.height for s in step_cards) > 0.8 * prs.slide_height
