"""Sparse / labelled charts: segment labels, auto-dropped value axis, fewer gridlines, line axis."""

from __future__ import annotations

from pptx import Presentation
from pptx.oxml.ns import qn

from slidemark.ir import Chart, Deck, Paragraph, Run, Series, Slide, Text
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.theme import get_theme


def build_chart(tmp_path, kind, cats, series, **opts):
    title = Text(role="title", paragraphs=[Paragraph(runs=[Run(text="T")])])
    ch = Chart(kind=kind, categories=cats, series=series, options=opts)
    d = Deck(slides=[Slide(title=title, elements=[ch])])
    th = get_theme("default")
    out = tmp_path / "c.pptx"
    render(d, [layout_slide(s, d, th, i) for i, s in enumerate(d.slides)], th, out)
    prs = Presentation(str(out))
    return [s for s in prs.slides[0].shapes if s.name.startswith("Chart")][0].chart


def _label_pts(chart, n):
    return {i: s.data_labels.font.size.pt for i, s in enumerate(list(chart.plots[0].series)[:n])}


def _deleted(ser):
    d = ser._element.find(qn("c:dLbls"))
    return [
        int(x.find(qn("c:idx")).get("val"))
        for x in d.findall(qn("c:dLbl"))
        if x.find(qn("c:delete")) is not None
    ]


def test_segment_labels_are_not_smaller_than_the_chart_text(tmp_path):
    th = get_theme("default")
    base = th.sizes["table"]
    ch = build_chart(
        tmp_path,
        "stacked-column",
        ["a", "b"],
        [Series(name="x", values=[60, 80]), Series(name="y", values=[40, 20])],
        labels="on",
    )
    pts = _label_pts(ch, 2)
    assert pts and all(p >= base for p in pts.values())


def test_label_that_cannot_fit_its_segment_is_hidden(tmp_path):
    ch = build_chart(
        tmp_path,
        "stacked-column",
        ["a", "b"],
        [Series(name="x", values=[190, 100]), Series(name="y", values=[1, 90])],
        labels="on",
    )
    sers = list(ch.plots[0].series)
    assert _deleted(sers[1]) == [0]  # the 1-of-191 sliver
    assert _deleted(sers[0]) == []


def test_few_categories_with_labels_drop_the_value_axis(tmp_path):
    cols = [Series(name="s", values=[4, 12, 21])]
    ch = build_chart(tmp_path, "column", ["a", "b", "c"], cols, labels="on")
    assert not ch.value_axis.visible and not ch.value_axis.has_major_gridlines
    ch = build_chart(tmp_path, "column", ["a", "b", "c"], cols, labels="on", axis="on")
    assert ch.value_axis.visible and ch.value_axis.has_major_gridlines
    ch = build_chart(tmp_path, "column", ["a", "b", "c"], cols)  # no labels: the axis carries the numbers
    assert ch.value_axis.visible
    many = [Series(name="s", values=[1, 2, 3, 4, 5])]
    ch = build_chart(tmp_path, "column", list("abcde"), many, labels="on")
    assert ch.value_axis.visible


def test_sparse_chart_value_labels_are_larger(tmp_path):
    th = get_theme("default")
    cols = [Series(name="s", values=[4, 12, 21])]
    ch = build_chart(tmp_path, "column", ["a", "b", "c"], cols, labels="on")
    sz = ch.plots[0].data_labels.font.size.pt
    assert sz > th.sizes["table"] * th.render.chart_label_scale


def test_sparse_axis_has_fewer_gridlines(tmp_path):
    ch = build_chart(tmp_path, "column", ["a", "b", "c"], [Series(name="s", values=[4, 12, 21])])
    va = ch.value_axis
    assert 3 <= round(va.maximum_scale / va.major_unit) <= 6
    ch = build_chart(tmp_path, "column", list("abcde"), [Series(name="s", values=[4, 12, 21, 5, 7])])
    assert round(ch.value_axis.maximum_scale / ch.value_axis.major_unit) >= 4


def test_line_chart_far_from_zero_starts_above_it(tmp_path):
    ser = [
        Series(name="a", values=[7.4, 7.6, 7.8, 8.0, 8.1]),
        Series(name="b", values=[5.1, 4.8, 4.6, 4.3, 4.1]),
    ]
    cats = ["2022", "2023", "2024", "2025", "2026"]
    ch = build_chart(tmp_path, "line", cats, ser)
    assert ch.value_axis.minimum_scale > 0
    ch = build_chart(tmp_path, "line", cats, ser, min="0")
    assert ch.value_axis.minimum_scale == 0
    ch = build_chart(tmp_path, "line", ["a", "b", "c"], [Series(name="a", values=[1, 8, 3])])
    assert ch.value_axis.minimum_scale in (None, 0)
