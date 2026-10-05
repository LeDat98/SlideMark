"""SVG pictures: Office svgBlip with PNG fallback, sanitizer, placeholders, import round trip."""

from __future__ import annotations

import re
import subprocess
import zipfile

import pytest
from pptx import Presentation

from slidemark.build import build
from slidemark.importer import import_pptx
from slidemark.ir import Image
from slidemark.parser import parse
from slidemark.render import htmlimg, svg
from slidemark.render.svg import sanitize_svg, svg_size
from slidemark.theme import Theme
from slidemark.xsd import _load_schema, validate_pptx

from .helpers import needs_soffice

SVG = (
    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100">'
    "<title>Logo</title>"
    '<circle cx="50" cy="50" r="40" fill="#e11"/></svg>'
)
EVIL = (
    '<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" viewBox="0 0 10 10" '
    'onload="alert(1)"><script>alert(1)</script>'
    "<foreignObject><div>x</div></foreignObject>"
    '<image xlink:href="http://evil.example/a.png" width="5" height="5"/>'
    '<image href="data:image/png;base64,AAAA" width="5" height="5"/>'
    '<rect width="5" height="5" onclick="x()" style="fill:url(http://evil.example/p)"/></svg>'
)


@pytest.fixture(scope="module")
def chromium():
    pytest.importorskip("playwright.sync_api")
    r = htmlimg.HtmlRenderer()
    try:
        png = r.render("<p>x</p>", 40, 20, Theme(name="probe"))
    finally:
        r.close()
    if png is None:
        pytest.skip("Chromium cannot be launched")


def _parts(path):
    with zipfile.ZipFile(path) as zf:
        return {n: zf.read(n) for n in zf.namelist()}


def _svg_parts(path):
    return {n: b for n, b in _parts(path).items() if n.endswith(".svg")}


def test_parse_file_and_fence():
    deck = parse(f"# T\n\n![logo](logo.svg)\n\n```svg\n{SVG}\n```\n")
    imgs = [e for e in deck.slides[0].elements if isinstance(e, Image)]
    assert imgs[0].src == "logo.svg" and imgs[0].alt == "logo"
    assert imgs[1].src.startswith("data:image/svg+xml,") and imgs[1].alt == "Logo"


def test_fence_alt_attr():
    deck = parse(f'# T\n\n```svg {{alt="Mark"}}\n{SVG}\n```\n')
    assert [e for e in deck.slides[0].elements if isinstance(e, Image)][0].alt == "Mark"


def test_sanitizer_strips_active_content():
    clean, why = sanitize_svg(EVIL)
    assert clean is not None, why
    low = clean.lower()
    assert "<script" not in low and "foreignobject" not in low
    assert "onload" not in low and "onclick" not in low
    assert "evil.example" not in low
    assert "data:image/png" in low


def test_sanitizer_rejects_garbage():
    assert sanitize_svg("<svg><g></svg>")[0] is None
    assert sanitize_svg("<html/>")[0] is None


def test_sanitizer_adds_namespace_and_size():
    clean, _ = sanitize_svg('<svg width="40" height="20"><rect width="1" height="1"/></svg>')
    assert clean and "http://www.w3.org/2000/svg" in clean
    assert svg_size(clean) == (40, 20)
    assert svg_size(SVG) == (200, 100)


def test_render_svg_blip(tmp_path, chromium):
    (tmp_path / "logo.svg").write_text(EVIL, encoding="utf-8")
    out = tmp_path / "a.pptx"
    d = build(f"# T\n\n![logo](logo.svg)\n\n```svg\n{SVG}\n```\n", out, base_dir=tmp_path)
    assert not [x for x in d.diagnostics if x.level == "error"]
    prs = Presentation(str(out))
    pics = [s for s in prs.slides[0].shapes if s.shape_type == 13]
    assert len(pics) == 2
    for pic in pics:
        xml = pic._element.xml
        assert "asvg:svgBlip" in xml and "{96DAC541-7B7A-43D3-8B79-37D633B846F1}" in xml
        rid = re.search(r'<asvg:svgBlip[^>]*r:embed="(\w+)"', xml).group(1)
        part = prs.slides[0].part.related_part(rid)
        assert part.content_type == "image/svg+xml"
        assert part.blob.startswith(b"<svg")
        assert pic.image.blob[:4] == b"\x89PNG"  # fallback
    blobs = list(_svg_parts(out).values())
    assert len(blobs) == 2
    assert all(b"<script" not in b and b"onload" not in b for b in blobs)
    assert b"image/svg+xml" in _parts(out)["[Content_Types].xml"]


def test_png_fallback_is_not_blank(tmp_path, chromium):
    import io

    from PIL import Image as PILImage

    out = tmp_path / "a.pptx"
    build(f"# T\n\n```svg\n{SVG}\n```\n", out)
    pic = [s for s in Presentation(str(out)).slides[0].shapes if s.shape_type == 13][0]
    im = PILImage.open(io.BytesIO(pic.image.blob)).convert("RGBA")
    assert im.getchannel("A").getextrema()[1] > 0
    assert abs(im.width / im.height - 2) < 0.1


def test_missing_file(tmp_path):
    out = tmp_path / "a.pptx"
    d = build("# T\n\n![x](nope.svg)\n", out, base_dir=tmp_path)
    assert any(x.rule == "image-missing" for x in d.diagnostics)
    assert out.exists() and not _svg_parts(out)


def test_invalid_svg_never_raises(tmp_path):
    out = tmp_path / "a.pptx"
    d = build("# T\n\n```svg\n<svg><g>\n```\n", out)
    assert any(x.rule == "svg-invalid" for x in d.diagnostics)


def test_too_large(tmp_path, monkeypatch):
    monkeypatch.setattr(svg, "MAX_BYTES", 50)
    out = tmp_path / "a.pptx"
    d = build(f"# T\n\n```svg\n{SVG}\n```\n", out)
    assert any(x.rule == "svg-too-large" for x in d.diagnostics)


def test_no_chromium_placeholder(tmp_path, monkeypatch):
    monkeypatch.setattr(svg, "_rasterize", lambda *a, **k: None)
    out = tmp_path / "a.pptx"
    d = build(f"# T\n\n```svg\n{SVG}\n```\n", out)
    assert any(x.rule == "svg-fallback" and x.level == "info" for x in d.diagnostics)
    pic = [s for s in Presentation(str(out)).slides[0].shapes if s.shape_type == 13][0]
    assert pic.image.blob[:4] == b"\x89PNG" and "svgBlip" in pic._element.xml
    assert len(_svg_parts(out)) == 1


def test_xsd_valid(tmp_path, monkeypatch):
    if _load_schema() is None:
        pytest.skip("ECMA-376 schemas not available")
    monkeypatch.setattr(svg, "_rasterize", lambda *a, **k: None)
    out = tmp_path / "a.pptx"
    build(f"# T\n\n```svg\n{SVG}\n```\n", out)
    assert validate_pptx(out) == []


@needs_soffice
def test_opens_in_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr(svg, "_rasterize", lambda *a, **k: None)
    out = tmp_path / "a.pptx"
    build(f"# T\n\n```svg\n{SVG}\n```\n", out)
    r = subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp_path), str(out)],
        capture_output=True,
        timeout=180,
    )
    assert r.returncode == 0 and (tmp_path / "a.pdf").stat().st_size > 0


def test_import_round_trip(tmp_path, monkeypatch):
    monkeypatch.setattr(svg, "_rasterize", lambda *a, **k: None)
    out = tmp_path / "a.pptx"
    build(f"# T\n\n```svg {{alt=Mark}}\n{SVG}\n```\n", out)
    md, _ = import_pptx(out, tmp_path / "o")
    m = re.search(r"!\[[^\]]*\]\((images/[^)]+\.svg)\)", md)
    assert m, md
    assert b"<circle" in (tmp_path / "o" / m.group(1)).read_bytes()
