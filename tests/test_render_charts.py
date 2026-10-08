"""Native charts: type, series, legend, labels, number formats, axes, colors, fonts (reopened .pptx)."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.oxml.ns import qn

from slidemark.ir import Chart, Deck, Paragraph, Run, Series, Slide, Text
from slidemark.layout import layout_slide
from slidemark.render import render
from slidemark.theme import get_theme

KINDS = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked-bar": XL_CHART_TYPE.BAR_STACKED,
    "stacked-column": XL_CHART_TYPE.COLUMN_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "radar": XL_CHART_TYPE.RADAR_MARKERS,
}


def chart_of(tmp_path: Path, ch: Chart):
    title = Text(role="title", paragraphs=[Paragraph(runs=[Run(text="T")])])
    d = Deck(slides=[Slide(title=title, elements=[ch])])
    th = get_theme("default")
    out = tmp_path / "c.pptx"
    render(d, [layout_slide(s, d, th, i) for i, s in enumerate(d.slides)], th, out)
    prs = Presentation(str(out))
    gf = [s for s in prs.slides[0].shapes if s.name.startswith("Chart")][0]
    return gf.chart, th


def mk(kind="column", series=None, cats=None, **opts):
    if series is None:
        series = [Series(name="a", values=[1, 2, 3]), Series(name="b", values=[3, None, 1])]
    return Chart(
        kind=kind,
        categories=cats if cats is not None else ["Q1", "Q2", "Q3"],
        series=series,
        options=opts,
    )


ONE = [Series(name="s", values=[1, 2, 3])]


@pytest.mark.parametrize("kind", list(KINDS))
def test_every_kind(tmp_path, kind):
    chart, _ = chart_of(tmp_path, mk(kind))
    assert chart.chart_type == KINDS[kind]
    plot = chart.plots[0]
    assert [s.name for s in plot.series] == (["a"] if kind == "pie" else ["a", "b"])  # pie: one series
    if kind != "scatter":
        assert list(plot.categories) == ["Q1", "Q2", "Q3"]
        assert list(plot.series[-1].values) == ([1, 2, 3] if kind == "pie" else [3, None, 1])  # None = gap
    else:
        assert len(list(plot.series[1].values)) == 2


def test_scatter_lines_hidden(tmp_path):
    chart, _ = chart_of(tmp_path, mk("scatter", cats=["1", "2", "3"]))
    ln = chart.plots[0].series[0]._element.find(qn("c:spPr")).find(qn("a:ln"))
    assert ln.find(qn("a:noFill")) is not None


def test_legend_defaults_and_positions(tmp_path):
    c, _ = chart_of(tmp_path, mk())
    assert c.has_legend and c.legend.position == XL_LEGEND_POSITION.BOTTOM
    c, _ = chart_of(tmp_path, mk(series=ONE))
    assert not c.has_legend
    c, _ = chart_of(tmp_path, mk("pie", series=ONE))
    assert c.legend.position == XL_LEGEND_POSITION.RIGHT
    c, _ = chart_of(tmp_path, mk(legend="top"))
    assert c.legend.position == XL_LEGEND_POSITION.TOP
    c, _ = chart_of(tmp_path, mk(legend="none"))
    assert not c.has_legend
    c, _ = chart_of(tmp_path, mk(series=ONE, legend="left"))
    assert c.legend.position == XL_LEGEND_POSITION.LEFT


def test_labels_fmt_and_axis_bounds(tmp_path):
    c, _ = chart_of(tmp_path, mk(labels="on", fmt="#,##0.0", min=0, max=10, axis="on"))
    dl = c.plots[0].data_labels
    assert dl.show_value and dl.number_format == "#,##0.0" and not dl.number_format_is_linked
    assert dl.font.size.pt < 18
    va = c.value_axis
    assert va.minimum_scale == 0 and va.maximum_scale == 10
    assert va.tick_labels.number_format == "#,##0.0" and not va.tick_labels.number_format_is_linked


def test_percent_flag_default_format(tmp_path):
    c, _ = chart_of(tmp_path, mk(labels="on", percent=True, axis="on"))
    assert c.plots[0].data_labels.number_format == '0"%"'
    assert c.value_axis.tick_labels.number_format == '0"%"'
    c, _ = chart_of(tmp_path, mk(labels="on", percent=True, fmt="0.0"))
    assert c.plots[0].data_labels.number_format == "0.0"


def test_pie_percent_labels_and_doughnut_hole(tmp_path):
    c, _ = chart_of(tmp_path, mk("pie", series=ONE, labels="percent"))
    dl = c.plots[0].data_labels
    assert dl.show_percentage and not dl.show_value
    c, _ = chart_of(tmp_path, mk("doughnut", series=ONE))
    assert c.plots[0]._element.find(qn("c:holeSize")).get("val") == "55"


def test_axis_off(tmp_path):
    c, _ = chart_of(tmp_path, mk(axis="off"))
    assert not c.value_axis.visible
    assert not c.value_axis.has_major_gridlines


def test_colors_series_and_points(tmp_path):
    c, th = chart_of(tmp_path, mk(colors=["#112233", "accent"]))
    fills = [str(s.format.fill.fore_color.rgb) for s in c.plots[0].series]
    assert fills[0] == "112233"
    assert fills[1] == th.color("accent").lstrip("#").upper()
    c, _ = chart_of(tmp_path, mk("pie", series=ONE, colors=["#AA0000", "#00BB00", "#0000CC"]))
    pts = [str(c.plots[0].series[0].points[i].format.fill.fore_color.rgb) for i in range(3)]
    assert pts == ["AA0000", "00BB00", "0000CC"]


def test_fonts_latin_ea_and_cjk(tmp_path):
    c, th = chart_of(tmp_path, mk(cats=["東京", "大阪", "名古屋"]))
    d = c._chartSpace.find(qn("c:txPr")).find(qn("a:p")).find(qn("a:pPr")).find(qn("a:defRPr"))
    assert d.find(qn("a:latin")).get("typeface") == th.fonts.body
    assert d.find(qn("a:ea")).get("typeface") == th.fonts.ea
    assert d.get("sz") == str(int(th.sizes["table"] * 100))
    assert list(c.plots[0].categories) == ["東京", "大阪", "名古屋"]


@pytest.mark.parametrize("kind", list(KINDS))
def test_never_raises_on_odd_input(tmp_path, kind):
    odd = [
        mk(kind, series=[]),
        mk(kind, series=[], cats=[]),
        mk(kind, series=[Series(name="x", values=[None, None, None])]),
        mk(kind, series=[Series(name="x", values=[1])], cats=["a", "b", "c"]),
        mk(kind, series=[Series(name="x", values=[1, 2, 3, 4, 5])], cats=["a"]),
        mk(kind, min="abc", max=None, fmt=5, colors="nonsense", legend=7, labels="percent", axis=False),
    ]
    for ch in odd:
        chart, _ = chart_of(tmp_path, ch)
        assert chart.chart_type == KINDS[kind]


def test_pie_rows_as_slices(tmp_path):
    rows = [Series(name=n, values=[v]) for n, v in (("APAC", 48), ("EU", 30), ("US", 22))]
    c, _ = chart_of(tmp_path, mk("doughnut", series=rows, cats=["Share"]))
    assert list(c.plots[0].categories) == ["APAC", "EU", "US"]
    assert list(c.plots[0].series[0].values) == [48, 30, 22]


@pytest.mark.parametrize("kind", ["bar", "column", "stacked-column"])
def test_negative_values_put_category_labels_low(tmp_path, kind):
    neg = [Series(name="s", values=[4, -2, 3])]
    chart, _ = chart_of(tmp_path, mk(kind, neg))
    pos = chart.category_axis._element.find(qn("c:tickLblPos"))
    assert pos is not None and pos.get("val") == "low"
    chart, _ = chart_of(tmp_path, mk(kind, ONE))
    pos = chart.category_axis._element.find(qn("c:tickLblPos"))
    assert pos is None or pos.get("val") != "low"


# ---- stack totals (hidden carrier series with per-point sum labels)

STACK = [Series(name="a", values=[38, 40]), Series(name="b", values=[62.6, 68])]


def _carrier_texts(ser):
    return [t.text for t in ser._element.iter(qn("a:t"))]


@pytest.mark.parametrize("kind", ["stacked-bar", "stacked-column"])
def test_stack_totals_carrier(tmp_path, kind):
    chart, _ = chart_of(tmp_path, mk(kind, series=STACK, cats=["x", "y"], labels="on"))
    sers = list(chart.plots[0].series)
    assert len(sers) == 3
    car = sers[-1]
    assert car._element.find(qn("c:spPr")).find(qn("a:noFill")) is not None  # hidden
    assert _carrier_texts(car) == ["100.6", "108"]  # the stack sums
    assert all(s._element.find(qn("c:dLbls")) is not None for s in sers[:2])
    assert chart.value_axis.maximum_scale > 108


def test_stack_totals_options(tmp_path):
    off, _ = chart_of(tmp_path, mk("stacked-column", series=STACK, labels="on", totals="off"))
    assert len(list(off.plots[0].series)) == 2
    on, _ = chart_of(tmp_path, mk("stacked-column", series=STACK, totals="on", cats=["x", "y"]))
    assert len(list(on.plots[0].series)) == 3
    none, _ = chart_of(tmp_path, mk("stacked-column", series=STACK))  # no labels: no totals
    assert len(list(none.plots[0].series)) == 2
    neg, _ = chart_of(tmp_path, mk("stacked-column", series=[Series(name="a", values=[-1, 2])], labels="on"))
    assert len(list(neg.plots[0].series)) == 1  # a negative value: the sum is not the stack's end


def test_stack_totals_fmt_and_legend(tmp_path):
    chart, _ = chart_of(
        tmp_path, mk("stacked-column", series=STACK, cats=["x", "y"], labels="on", fmt="#,##0")
    )
    assert _carrier_texts(list(chart.plots[0].series)[-1]) == ["101", "108"]
    leg = chart._chartSpace.find(".//" + qn("c:legend"))
    assert [e.find(qn("c:idx")).get("val") for e in leg.findall(qn("c:legendEntry"))] == ["2"]
