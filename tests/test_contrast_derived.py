"""Contrast-safe derived defaults (auto ink), the fix-in-one-edit lint hints and dark-brand surfaces."""

from __future__ import annotations

import re

import pytest

from slidemark.contrast import ratio
from slidemark.layout import layout_slide
from slidemark.lint import lint
from slidemark.parser import parse
from slidemark.template import deck_theme
from slidemark.theme import available, derive_ink, get_theme

BRAND = (
    "theme: none\ncolors: bg=#FFFFFF fg=#1F2937 primary=#84CC16 accent=#14532D\n"
    "fonts: heading=Arial body=Arial\n"
)
BODY = """
# Cover
Subtitle

# Results
## Revenue {.kpi}
12.4B
+15%
## New customers {.kpi}
3,200
+22%
> Best month of the year

# Regions
- Da Nang has [new]{.badge} stores
"""


def _build(src: str):
    deck = parse(src)
    theme, _ = deck_theme(deck, ".")
    placed = [layout_slide(s, deck, theme, i) for i, s in enumerate(deck.slides)]
    return deck, theme, placed, [d for d in lint(deck, placed, theme) if d.rule == "contrast"]


def _hex(theme, v):
    return theme.hexval(v)


def test_repro_deck_has_no_contrast_warning_and_readable_colors():
    deck, theme, placed, diags = _build(BRAND + BODY)
    assert diags == []
    surface, bg = _hex(theme, "surface"), _hex(theme, "bg")
    kpi = [
        q.style.color
        for p in placed[1]
        for q in getattr(p.element, "paragraphs", [])
        if q.style and q.style.bold
    ]
    assert kpi and all(ratio(_hex(theme, c), surface) >= 3.0 for c in kpi)
    bar = next(p for p in placed[1] if p.style.fill == "primary")
    assert ratio(_hex(theme, bar.style.color), _hex(theme, "primary")) >= 4.5
    badge = theme.badge_ink("primary")
    assert ratio(_hex(theme, badge), _hex(theme, "primary")) >= 4.5
    assert theme.colors["primary"] == "#84CC16"  # fills keep the brand color
    assert bg == "#FFFFFF"


def test_rendered_pptx_uses_the_derived_colors(tmp_path):
    from pptx import Presentation

    from slidemark import build

    out = tmp_path / "d.pptx"
    build(BRAND + BODY, out)
    colors = {}
    for shape in Presentation(out).slides[1].shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.color and run.font.color.type is not None:
                        colors[run.text] = str(run.font.color.rgb)
    assert colors["12.4B"] != "84CC16" and ratio("#" + colors["12.4B"], "#F3F4F6") >= 3.0
    assert ratio("#" + colors["Best month of the year"], "#84CC16") >= 4.5


FULL = BODY + "\n# Detail\n> lead line\n## Box\ntext here\n^ footnote text\n"


@pytest.mark.parametrize(
    "line,token",
    [
        ("style: kpi.color=#84CC16", "kpi.color"),
        ("style: conclusion.color=#FFFFFF", "conclusion.color"),
        ("style: badge.color=#FFFFFF", "badge.color"),
        ("style: lead.color=#DDDDDD", "lead.color"),
        ("style: title.color=#DDDDDD", "title.color"),
        ("style: p.color=#DDDDDD", "p.color"),
        ("style: colors.muted=#DDDDDD", "colors.muted"),
        ("style: title.band=primary\nstyle: title.band.color=#FFFFFF", "title.band.color"),
        ("style: heading.band=primary\nstyle: heading.band.color=#FFFFFF", "heading.band.color"),
    ],
)
def test_explicit_bad_color_is_reported_and_the_hint_fixes_it(line, token):
    _, _, _, diags = _build(BRAND + line + "\n" + FULL)
    suggestions = {d.hint.split("style: ", 1)[1] for d in diags if f"style: {token}=" in (d.hint or "")}
    assert suggestions, [d.hint for d in diags]
    pasted = "".join(f"style: {x}\n" for x in suggestions)
    _, _, _, after = _build(BRAND + line + "\n" + pasted + FULL)
    assert after == []  # one paste clears every warning of that color


def test_hint_format_names_token_color_and_ratio():
    _, _, _, diags = _build(BRAND + "style: kpi.color=#84CC16\n" + BODY)
    d = diags[0]
    assert re.fullmatch(
        r"set the color to #[0-9A-F]{6} \(passes [\d.]+:1 on #[0-9A-F]{6}\): style: kpi\.color=#[0-9A-F]{6}",
        d.hint,
    )
    _, _, _, diags = _build(BRAND + "style: conclusion.color=#FFFFFF\n" + BODY)
    assert diags[0].hint == "text on primary #84CC16: style: conclusion.color=#1F2937"


def test_run_level_colors_are_judged_and_names_are_made_legible():
    src = BRAND + "\n# T\n- plain [lime text]{color=#84CC16} and ==marked== and [x]{.primary} y\n"
    deck, theme, placed, diags = _build(src)
    assert any("{color=" in (d.hint or "") for d in diags)  # explicit hex run: reported, never changed
    assert theme.legible("accent", 18) == "accent"  # dark green already reads
    assert ratio(theme.hexval(theme.legible("primary", 18)), "#F3F4F6") >= 4.5


def test_ink_auto_off_restores_the_plain_defaults():
    _, theme, _, diags = _build(BRAND + "style: ink.auto=off\n" + BODY)
    assert theme.classes["kpi"].color == "primary" and theme.conclusion_color == "bg"
    assert diags  # the unreadable defaults are now reported instead


def test_dark_brand_cards_are_visible():
    src = "theme: none\ncolors: bg=#0B1F3A fg=#FFFFFF primary=#FF6B57\n" + BODY
    _, theme, _, diags = _build(src)
    bg, surface, border = (theme.hexval(k) for k in ("bg", "surface", "border"))
    assert surface != bg and ratio(surface, bg) >= 1.08
    assert ratio(border, bg) >= 1.2
    assert ratio(theme.hexval("fg"), surface) >= 4.5
    assert diags == []


def test_declared_surface_and_border_are_kept():
    src = "theme: none\ncolors: bg=#0B1F3A fg=#FFFFFF surface=#112244 border=#445566\n" + BODY
    _, theme, _, _ = _build(src)
    assert theme.colors["surface"] == "#112244" and theme.colors["border"] == "#445566"


@pytest.mark.parametrize("name", available())
def test_presets_keep_their_look(name):
    theme = get_theme(name)
    assert derive_ink(theme) == theme
